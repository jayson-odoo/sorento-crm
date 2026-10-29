"""Compare a purchasing handover email with the OI Excel it attached (OI-EMAIL-PURCHASING-FOCUS).

Read-only. Finds the newest `order_inquiry_handover` email on record (a Notification of type
`automation_email`) whose body names the given S/O number, and its Email Outbox row (the email
System Management > Email Outbox shows). Reads the outbox row's line table, fetches the attached
OI Excel from storage and prints both side by side, then PASS when every email line is an Excel
row and nothing else is.

`--after` is when the Confirm (or resend) was pressed, in UTC: `EmailOutbox.created_at` is stored
as naive UTC (`server_default=func.now()` on a UTC database). A time with an offset
("2026-10-26 14:05+08:00") is converted. An outbox row from before it is an older email, so that
is a FAIL, never a pass on stale evidence. `--expect-none` (with `--after`) checks the opposite: a
Confirm that only saved Use own location / Use BRW decisions wrote NO handover email for that S/O.
`--save-excel PATH` writes the attached OI Excel to PATH, to hand to the owner.

    venv/bin/python scripts/compare_oi_handover_email_excel.py SO394881 --after "2026-10-26 06:05" \
        --save-excel /tmp/SO394881.xlsx
    venv/bin/python scripts/compare_oi_handover_email_excel.py SO380097 --after "..." --expect-none
    venv/bin/python scripts/compare_oi_handover_email_excel.py --latest
"""
from __future__ import annotations

import argparse
import io
import sys
from datetime import date, datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import List, Optional, Sequence, Tuple

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

#: The email's line-table headings and the Excel heading each one is in the sheet. Both the
#: current layout (`oihr_0010`) and the older wide one (`oihr_0004`) are read.
_EMAIL_TO_EXCEL = {
    "SO DATE": "SO DATE",
    "S/O NO": "S/O NO",
    "LOCATION": "LOCATION",
    "STOCK LOCATION": "LOCATION",
    "ITEM CODE": "ITEM CODE",
    "NEW ITEM CODE": "NEW ITEM CODE",
    "QTY": "QTY",
    "NEW QTY": "QTY CHANGE TO",
    "QTY CHANGE TO": "QTY CHANGE TO",
    "DELIVERY": "DELIVERY DATE",
    "DELIVERY DATE": "DELIVERY DATE",
    "NEW DATE": "DELIVERY DATE CHANGE TO",
    "DELIVERY DATE CHANGE TO": "DELIVERY DATE CHANGE TO",
    "PROJECT/CUSTOMER": "PROJECT/CUSTOMER",
    "REMARK": "REMARK",
}


class _Tables(HTMLParser):
    """Every `<table>` of a document as rows of cell texts."""

    def __init__(self) -> None:
        super().__init__()
        self.tables: List[List[List[str]]] = []
        self._row: Optional[List[str]] = None
        self._cell: Optional[List[str]] = None

    def handle_starttag(self, tag, attrs):  # noqa: ANN001
        if tag == "table":
            self.tables.append([])
        elif tag == "tr":
            self._row = []
        elif tag in ("td", "th"):
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):  # noqa: ANN001
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(" ".join("".join(self._cell).split()))
            self._cell = None
        elif tag == "tr" and self._row is not None and self.tables:
            self.tables[-1].append(self._row)
            self._row = None

    def handle_data(self, data):  # noqa: ANN001
        if self._cell is not None:
            self._cell.append(data)


def _text(value) -> str:  # noqa: ANN001
    if value is None:
        return ""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float):
        return f"{value:f}".rstrip("0").rstrip(".")
    return " ".join(str(value).split())


def email_rows(body_html: str) -> Tuple[List[str], List[Tuple[str, ...]]]:
    parser = _Tables()
    parser.feed(body_html or "")
    for table in parser.tables:
        if table and "ITEM CODE" in table[0] and "LOCATION" in table[0]:
            head = table[0]
            # A one-cell row is the "S/O NO ..." group heading, not a line.
            rows = [row for row in table[1:] if len(row) == len(head)]
            return head, [tuple(_text(cell) for cell in row) for row in rows]
    return [], []


def excel_rows(data: bytes, email_head: Sequence[str]) -> Tuple[List[str], List[Tuple[str, ...]]]:
    import openpyxl

    sheet = openpyxl.load_workbook(io.BytesIO(data)).active
    rows = list(sheet.iter_rows(values_only=True))
    head = [h for h in rows[1] if h is not None]
    index = {heading: position for position, heading in enumerate(head)}
    out = []
    for row in rows[2:]:
        if not any(v is not None for v in row):
            continue
        out.append(
            tuple(
                _text(row[index[_EMAIL_TO_EXCEL[h]]]) if _EMAIL_TO_EXCEL.get(h) in index else "?"
                for h in email_head
            )
        )
    return head, out


def _utc(text: str) -> datetime:
    """`--after` as the naive UTC datetime `EmailOutbox.created_at` is stored in."""
    moment = datetime.fromisoformat(text)
    if moment.tzinfo is not None:
        moment = moment.astimezone(timezone.utc).replace(tzinfo=None)
    return moment


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("so_number", nargs="?")
    parser.add_argument("--latest", action="store_true", help="the newest handover email")
    parser.add_argument(
        "--after",
        help="when the Confirm was pressed, UTC (YYYY-MM-DD HH:MM, or with +08:00); older = FAIL",
    )
    parser.add_argument(
        "--expect-none", action="store_true",
        help="PASS only if NO handover email for the S/O was written after --after",
    )
    parser.add_argument("--save-excel", metavar="PATH", help="write the attached OI Excel here")
    args = parser.parse_args(argv)
    if not args.so_number and not args.latest:
        parser.error("give an S/O number or --latest")
    if args.expect_none and not (args.after and args.so_number):
        parser.error("--expect-none needs an S/O number and --after")
    after = _utc(args.after) if args.after else None

    from app.database import SessionLocal
    from app.models.email_outbox import EmailOutbox
    from app.models.notification import Notification
    from app.services import storage_router

    db = SessionLocal()
    try:
        query = db.query(Notification).filter(
            Notification.type == "automation_email",
            Notification.data["trigger_type"].astext == "order_inquiry_handover",
        )
        if args.so_number:
            query = query.filter(Notification.data["body_html"].astext.contains(args.so_number))
        if args.expect_none:
            notes = [n.id for n in query.all()]
            written = (
                db.query(EmailOutbox)
                .filter(
                    EmailOutbox.metadata_json["notification_id"].astext.in_(
                        [str(i) for i in notes]
                    ),
                    EmailOutbox.created_at > after,
                )
                .count()
                if notes
                else 0
            )
            if written:
                print(f"FAIL: {written} handover Email Outbox row(s) for it after {args.after}.")
                return 1
            print(f"PASS: no handover email for {args.so_number} after {args.after} UTC.")
            return 0
        note = query.order_by(Notification.created_at.desc()).first()
        if note is None:
            print("No handover email on record for that.")
            return 2
        data = note.data or {}
        outbox = (
            db.query(EmailOutbox)
            .filter(EmailOutbox.metadata_json["notification_id"].astext == str(note.id))
            .order_by(EmailOutbox.created_at.desc())
            .first()
        )
        if outbox is None:
            print(f"FAIL: the email of {note.created_at} has no Email Outbox row yet.")
            return 1
        print(f"Email Outbox row of {outbox.created_at} ({outbox.status}): {outbox.subject}")
        if after is not None and outbox.created_at <= after:
            print(f"FAIL: that email is older than the Confirm at {args.after}.")
            return 1
        head, emailed = email_rows(outbox.body_html or data.get("body_html") or "")
        excel = [
            a for a in data.get("extra_attachments") or [] if a.get("filename", "").endswith(".xlsx")
        ]
        if not excel:
            print("FAIL: the email attached no OI Excel.")
            return 1
        workbook = storage_router.get_backend(excel[0]["storage_provider"]).download_file(
            excel[0]["storage_key"]
        )
        sheet_head, sheet = excel_rows(workbook, head)
        if args.save_excel:
            Path(args.save_excel).write_bytes(workbook)
            print(f"Saved the attached Excel to {args.save_excel}")
    finally:
        db.close()

    print(f"\nEMAIL ({len(emailed)} lines): {' | '.join(head)}")
    for row in emailed:
        print("  " + " | ".join(row))
    print(f"\nEXCEL {excel[0]['filename']} ({len(sheet)} rows): {' | '.join(sheet_head)}")
    for row in sheet:
        print("  " + " | ".join(row))
    if sorted(emailed) == sorted(sheet):
        print("\nPASS: the Excel rows are the email's lines.")
        return 0
    print("\nFAIL: they differ.")
    for row in sorted(set(emailed) - set(sheet)):
        print("  only in the email: " + " | ".join(row))
    for row in sorted(set(sheet) - set(emailed)):
        print("  only in the Excel: " + " | ".join(row))
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
