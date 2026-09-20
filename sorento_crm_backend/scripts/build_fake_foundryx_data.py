"""Builds full-size fake-FoundryX data files for the AutoCount pull lane's browser
pass (PLAN-autocount-pull-review.md "Local stack for the browser pass", SR4 fix round).
READ-ONLY against the database - never writes a row, never runs a migration.

    venv/bin/python scripts/build_fake_foundryx_data.py \
        --database-url postgresql://user:pass@host:5432/dbname \
        --company-code SRT \
        --xlsm /path/to/stock_export.xlsm \
        --out /path/to/fake-foundryx-data

Writes `<company-code>-products.json` (every product of the company, canonical row
shape - `source_ref`/`code`/`name`/`description`/`category_code`/`brand_code`
(omitted when the product has none)/`list_price`/`is_active`) and
`<company-code>-stock_balances.json` (the xlsm's `Master` sheet, summed to one row per
(item_code, location_code) pair, `On Hand Qty > 0` only) - the exact shape
`tests/support/fake_foundryx.py` reads when `FAKE_FOUNDRYX_DATA` points at `--out`.

`--database-url` is REQUIRED and used exactly as given - this never reads `DATABASE_URL`
from the environment. A worktree's own `.env` overrides whatever the shell exports
(LESSONS-LEARNT ".env overrides env"), so trusting the environment here risks silently
reading a different lane's database than the one the caller meant.

Table names checked against the models (`grep __tablename__`) before writing the SQL
below - they differ from the model class names in this codebase: `products`,
`product_categories`, `brands`, `companies`.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from sqlalchemy import create_engine, text

_PRODUCTS_SQL = text(
    """
    SELECT p.product_code, p.product_name, p.description, p.list_price, p.is_active,
           c.category_code, b.brand_code
    FROM products p
    JOIN companies co ON co.id = p.company_id
    JOIN product_categories c ON c.id = p.category_id
    LEFT JOIN brands b ON b.id = p.brand_id
    WHERE co.code = :company_code
    ORDER BY p.product_code
    """
)

#: The four columns the xlsm's `Master` sheet must carry - the header row's position
#: is not assumed to be row 1, so it is located by matching these names.
_STOCK_HEADER_COLUMNS = ("Item Code", "Item Description", "Location", "On Hand Qty")


def _build_products(engine, company_code: str) -> list[dict]:
    with engine.connect() as conn:
        rows = conn.execute(_PRODUCTS_SQL, {"company_code": company_code}).mappings().all()

    products: list[dict] = []
    for row in rows:
        code = row["product_code"]
        name = row["product_name"]
        canonical = {
            "source_ref": f"AED_SORENTO:{code}",
            "code": code,
            "name": name,
            # Derived from the stored row, not re-parsed: description when the
            # product has one, else the name (D24-style: no synthetic content).
            "description": row["description"] or name,
            "category_code": row["category_code"],
            "list_price": str(row["list_price"]),
            "is_active": bool(row["is_active"]),
        }
        if row["brand_code"]:
            canonical["brand_code"] = row["brand_code"]
        products.append(canonical)
    return products


def _find_header_row(worksheet) -> tuple[int, dict[str, int]]:
    """The 1-based row index and 0-based column index of each of the four expected
    headers. Scans the first 50 rows - the sheet is expected to carry the header
    somewhere near the top, not necessarily row 1 (a title/branding row is common)."""
    for row_idx, row in enumerate(worksheet.iter_rows(min_row=1, max_row=50, values_only=True), start=1):
        cells = {str(value).strip(): i for i, value in enumerate(row) if value is not None}
        if all(header in cells for header in _STOCK_HEADER_COLUMNS):
            return row_idx, {header: cells[header] for header in _STOCK_HEADER_COLUMNS}
    raise SystemExit(
        f"Could not find a header row carrying all of {_STOCK_HEADER_COLUMNS!r} in the "
        f"'Master' sheet (checked the first 50 rows)."
    )


def _cell(row: tuple, columns: dict[str, int], header: str):
    index = columns[header]
    return row[index] if index < len(row) else None


def _build_stock(xlsm_path: Path) -> list[dict]:
    import openpyxl

    workbook = openpyxl.load_workbook(str(xlsm_path), read_only=True, data_only=True)
    if "Master" not in workbook.sheetnames:
        raise SystemExit(f"'Master' sheet not found in {xlsm_path} (have: {workbook.sheetnames})")
    worksheet = workbook["Master"]
    header_row, columns = _find_header_row(worksheet)

    # (item_code, location_code) -> the accumulating row - one row per pair, summed,
    # per the plan ("one row per (item_code, location_code) summed").
    totals: dict[tuple[str, str], dict] = {}
    for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        raw_item_code = _cell(row, columns, "Item Code")
        if raw_item_code is None or str(raw_item_code).strip() == "":
            continue
        item_code = str(raw_item_code).strip()

        raw_location = _cell(row, columns, "Location")
        location = str(raw_location).strip() if raw_location is not None else ""
        if not location:
            continue

        raw_description = _cell(row, columns, "Item Description")
        description = str(raw_description).strip() if raw_description is not None else ""

        raw_qty = _cell(row, columns, "On Hand Qty")
        try:
            qty = float(raw_qty) if raw_qty is not None else 0.0
        except (TypeError, ValueError):
            continue
        if qty <= 0:
            continue

        key = (item_code, location)
        if key in totals:
            totals[key]["qty"] += qty
        else:
            totals[key] = {
                "source_ref": f"AED_SORENTO:{item_code}|{location}",
                "item_code": item_code,
                "item_description": description,
                "location_code": location,
                "qty": qty,
            }

    stock = list(totals.values())
    for entry in stock:
        # Every source cell reads as a float; AutoCount on-hand quantities are whole
        # units, matching the committed fixture's own shape (`"qty": 672`, not `672.0`).
        entry["qty"] = int(entry["qty"]) if float(entry["qty"]).is_integer() else entry["qty"]
    return stock


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--database-url", required=True, help="Full SQLAlchemy/psycopg URL - never read from env")
    parser.add_argument("--company-code", required=True, help="companies.code, e.g. SRT")
    parser.add_argument("--xlsm", required=True, type=Path, help="Path to the AutoCount stock export .xlsm")
    parser.add_argument("--out", required=True, type=Path, help="Directory to write the two JSON files into")
    args = parser.parse_args()

    engine = create_engine(args.database_url)
    try:
        products = _build_products(engine, args.company_code)
    finally:
        engine.dispose()
    stock = _build_stock(args.xlsm)

    args.out.mkdir(parents=True, exist_ok=True)
    products_path = args.out / f"{args.company_code}-products.json"
    stock_path = args.out / f"{args.company_code}-stock_balances.json"
    products_path.write_text(json.dumps(products, indent=2))
    stock_path.write_text(json.dumps(stock, indent=2))

    print(f"products: {len(products)} rows -> {products_path}")
    print(f"stock_balances: {len(stock)} rows -> {stock_path}")


if __name__ == "__main__":
    main()
