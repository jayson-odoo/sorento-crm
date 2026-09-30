"""Local simulation of the delivery-orders "Compare with my Excel" tab (DO-COMPARE-SIM).

Runs the SAME comparison code the route runs (`app.services.autocount_pull_compare`) on the
checker's two macro workbooks against a raw AutoCount DO snapshot, then breaks the
only-in-Excel / only-in-AutoCount / differing rows down by cause. Read-only: no database,
no network. The owner's workbooks are passed by path and never copied into the repo.

    venv/bin/python scripts/simulate_do_compare.py \
        --lines  "<Order Listing ... .xlsm>" \
        --headers "<Order Tracking ... .xlsm>" \
        --ac snapshot.json --from 2026-09-01 --to 2026-09-03 [--sheet template|master|auto]

`--ac` is a JSON array of raw vendor DO records (`DocNo`, `DocDate`, `Cancelled`,
`Details[]`), e.g. exported from a local ss `ac_pull_snapshot_row.payload_json`.
`--sheet` picks the workbook sheet the way the browser would: `template` reproduces the
pre-fix browser rule (any `.xlsm` -> sheet `Template`), `master` reads sheet `Master`,
`auto` uses the browser's current rule (`resolveCompareSheetName` in lib/excel-utils.ts).
"""
from __future__ import annotations

import argparse
import collections
import datetime as dt
import json
import os
import re
import sys
from typing import Any

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.services.autocount_pull_compare import (  # noqa: E402
    compare_delivery_order_headers,
    compare_delivery_orders,
    window_excel_rows,
)

_EPOCH = dt.datetime(1899, 12, 30)


def _cell(value: Any) -> Any:
    """What SheetJS `sheet_to_json` hands the browser: a date cell is an Excel serial."""
    if isinstance(value, dt.datetime):
        return (value - _EPOCH).total_seconds() / 86400
    if isinstance(value, dt.date):
        return (dt.datetime.combine(value, dt.time()) - _EPOCH).days
    if isinstance(value, dt.time):
        return (value.hour * 3600 + value.minute * 60 + value.second) / 86400
    return value


def sheet_rows(path: str, sheet: str) -> list[dict]:
    """`sheet_to_json` on one sheet: first row is the header, a repeated header gets `_1`,
    `_2`... appended, blank cells are left out and blank rows skipped."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = wb[sheet].iter_rows(values_only=True)
    header_raw = next(rows)
    seen: dict[str, int] = {}
    header: list[str | None] = []
    for h in header_raw:
        if h is None:
            header.append(None)
            continue
        name = str(h)
        if name in seen:
            seen[name] += 1
            header.append(f"{name}_{seen[name]}")
        else:
            seen[name] = 0
            header.append(name)
    out = []
    for row in rows:
        obj = {header[i]: _cell(v) for i, v in enumerate(row)
               if i < len(header) and header[i] and v is not None and v != ""}
        if obj:
            out.append(obj)
    return out


def pick_sheet(path: str, mode: str, source: str) -> str:
    import openpyxl

    names = openpyxl.load_workbook(path, read_only=True).sheetnames
    if mode == "template":
        if "Template" not in names:
            raise SystemExit(
                f"{os.path.basename(path)}: Macro workbook must contain a 'Template' sheet "
                f"(found sheets: {', '.join(names)})"
            )
        return "Template"
    if mode == "master":
        return "Master"
    # auto: the browser's DO compare rule - sheet Master by name, case-insensitive.
    for n in names:
        if n.lower() == "master":
            return n
    return names[0]


def _doc_shape(doc_no: str) -> str:
    return re.sub(r"\d", "9", doc_no or "") or "(blank)"


def _prefix(doc_no: str) -> str:
    m = re.match(r"^([A-Z-]*?)(\d{6}-\d+)$", (doc_no or "").upper())
    return m.group(1) if m else "(other)"


def breakdown_lines(result: dict, excel_rows: list[dict], ac: list[dict]) -> None:
    ac_docs = {str(r.get("DocNo") or "").strip().upper(): r for r in ac}
    excel_docs = {str(r.get("Doc No") or "").strip().upper() for r in excel_rows}

    def cause_only_ac(label: str) -> str:
        doc, item, loc = label.split("|")
        rec = ac_docs.get(doc.upper())
        if rec and str(rec.get("Cancelled") or "").upper() == "T":
            return "cancelled DO (AutoCount Cancelled=T)"
        if doc.upper() not in excel_docs:
            return f"whole DO missing from Excel (prefix {_prefix(doc) or 'none'})"
        return "DO in both, line key differs (item/location)"

    def cause_only_excel(label: str) -> str:
        doc, item, loc = label.split("|")
        if doc.upper() not in ac_docs:
            return f"whole DO missing from AutoCount (prefix {_prefix(doc) or 'none'})"
        if not loc:
            return "Excel line has no Location"
        return "DO in both, line key differs (item/location)"

    print("  only in AutoCount by cause:")
    for k, v in collections.Counter(map(cause_only_ac, result["only_in_pull"])).most_common():
        print(f"    {v:6d}  {k}")
    print("  only in Excel by cause:")
    for k, v in collections.Counter(map(cause_only_excel, result["only_in_excel"])).most_common():
        print(f"    {v:6d}  {k}")
    print("  differences by field:")
    for k, v in collections.Counter(d["field"] for d in result["differences"]).most_common():
        print(f"    {v:6d}  {k}")
    samples: dict[str, list] = collections.defaultdict(list)
    for d in result["differences"]:
        if len(samples[d["field"]]) < 4:
            samples[d["field"]].append((d["doc_no"], d["item_code"], d["excel"], d["pull"]))
    for f, s in samples.items():
        print(f"    e.g. {f}: {s}")


def breakdown_headers(result: dict, ac: list[dict]) -> None:
    print("  only in AutoCount by DocNo shape:")
    for k, v in collections.Counter(map(_doc_shape, result["only_in_pull"])).most_common(8):
        print(f"    {v:6d}  {k}")
    print("  only in Excel by DocNo shape:")
    for k, v in collections.Counter(map(_doc_shape, result["only_in_excel"])).most_common(8):
        print(f"    {v:6d}  {k}")
    print("  differences by field:")
    for k, v in collections.Counter(d["field"] for d in result["differences"]).most_common():
        print(f"    {v:6d}  {k}")
    for d in result["differences"][:6]:
        print(f"    e.g. {d['doc_no']} {d['field']}: excel={d['excel']} pull={d['pull']}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--lines")
    ap.add_argument("--headers")
    ap.add_argument("--ac", required=True)
    ap.add_argument("--from", dest="from_day")
    ap.add_argument("--to", dest="to_day")
    ap.add_argument("--sheet", choices=("template", "master", "auto"), default="auto")
    args = ap.parse_args()
    with open(args.ac) as fh:
        ac = json.load(fh)
    print(f"AutoCount: {len(ac)} documents, "
          f"{sum(len(r.get('Details') or []) for r in ac)} detail rows, window {args.from_day}..{args.to_day}")

    for source, path in (("lines", args.lines), ("headers", args.headers)):
        if not path:
            continue
        try:
            sheet = pick_sheet(path, args.sheet, source)
        except SystemExit as exc:
            print(f"\n[{source}] PARSE ERROR: {exc}")
            continue
        rows = sheet_rows(path, sheet)
        kept, ignored = window_excel_rows(rows, args.from_day, args.to_day)
        fn = compare_delivery_order_headers if source == "headers" else compare_delivery_orders
        result = fn(kept, ac)
        s = result["summary"]
        print(f"\n[{source}] sheet {sheet!r}: {len(rows)} rows, {len(kept)} in window, {ignored} ignored")
        print(f"  {s['matched']} of {s['total']} match. {s['different']} differ "
              f"({len(result['differences'])} differences), {len(result['only_in_excel'])} only in Excel, "
              f"{len(result['only_in_pull'])} only in AutoCount")
        if source == "lines":
            breakdown_lines(result, kept, ac)
        else:
            breakdown_headers(result, ac)


if __name__ == "__main__":
    main()
