"""RED tests for SR3 of the AutoCount pull + review lane (issue #1047).

Covers: AC-RV-3, AC-RV-5, AC-CM-2, AC-CM-4, AC-CM-5 (plus the owner/permission
gate on the compare route), AC-PC-1..5. AC-CM-1/AC-CM-3/AC-CM-6 are FE/stock -
out of scope here (SR4 does stock).

Plan: documentation/plans/autocount/PLAN-autocount-pull-review.md
UAC:  documentation/plans/autocount/autocount-pull-review-acceptance-criteria.md

Nothing under test here exists yet - the rows/download/compare/confirm routes,
`app/services/autocount_pull_compare.py`, `app/tasks/autocount_pull_tasks.
apply_autocount_pull`, and `MasterIngestService`'s `stamp_user_id` kwarg are ALL
the coder's SR3 deliverable. Every one of those is imported (or invoked) INSIDE
a test/fixture body, never at module import time, so a missing piece reds one
test, not collection.

Substrate reused from SR1 (`tests/test_autocount_pull_sr1.py`) by import, not
copy - `env` (HTTP route driver over `blank_session()`), `task_db` (the apply
task's own `SessionLocal`, patched to a private connection over the shared
blank schema), `_FakeFoundryX` / `_patch_foundryx`, `_header` / `_stored` /
`_content_hash`, `_seed_pull_job` / `_job_row` / `_job_count` / `_job_rows`,
`_fixture`, `_canonical_row`. Both `env` and `task_db` are plain
`@pytest.fixture` functions in that module - importing the NAME into this
module's globals is enough for pytest to discover them here too (verified
empirically before this file was written: a throwaway smoke test importing
just `env`/`task_db` collected and ran green). See that file's own module
docstring for why each shape exists.

PC-5 (the parity gate) instead follows `tests/test_ingest_parity_s1_products.
py`'s OWN pattern - a plain `blank_session()`, two `Company` rows in one
session, `set_company_scope` toggled between them, `ProductService(db).
bulk_import_products` for the manual/Excel side and `MasterIngestService(db,
company_id=...)` for the pull side - proven to work in this repo (that file's
own `TestAcP19BulkImportVsEsbParity` passes today, confirmed by running it
before this file was written) rather than driving two full
`apply_autocount_pull` job cycles across two companies for no extra coverage
PC-2/PC-3/PC-4 do not already give. Said explicitly in the tester report.
"""
from __future__ import annotations

import io
import uuid
from decimal import Decimal
from unittest.mock import MagicMock

import openpyxl
import pytest
from sqlalchemy import text

# MUST be the first app import - resolves the circular import in
# app.modules.runtime.guards (repo convention, see test_ingest_deletions.py).
from app.main import app  # noqa: E402,F401

from app.models.base import set_company_scope
from app.services.company_scope import DEFAULT_COMPANY_ID

from tests._pg_fixture import blank_session, unique_code
from tests.test_autocount_pull_sr1 import (  # noqa: F401 - env/task_db are fixtures
    API_KEY,
    BASE_URL,
    FIXTURE_DIR,
    PULLS_URL,
    _FakeFoundryX,
    _canonical_row,
    _content_hash,
    _fixture,
    _header,
    _job_count,
    _job_row,
    _job_rows,
    _patch_foundryx,
    _seed_pull_job,
    _stored,
    env,
    task_db,
)

MARKER = "ZZTAP3"


# ============================================================== row shaping


def _expected_view_row(row: dict) -> dict:
    """The `/rows` shape AC-RV-3 pins, computed from a raw snapshot row so the
    test does not just restate whatever the route happens to return.

    Captain ruling (SR3 amend round, contract fix): `desc_2` must ROUND-TRIP
    through the manual join formula (`f"{description} {desc2}".strip()`), so
    it is the remainder of `description` after `name` with EXACTLY ONE
    separator space removed - any further leading whitespace (a real double
    space in the source data) is KEPT, never `.strip()`-ped away, because that
    one extra space is exactly what the join's own inserted space needs to
    reproduce. `description == name` -> `desc_2 == ""`. `description` not even
    starting with `name` (a shape no fixture row hits, but the rule must still
    answer something) -> `desc_2 == ""` and the `description` CELL becomes the
    full raw `description`, not `name` - there is no boundary to split on.
    """
    name = row["name"]
    description = row.get("description") or ""
    if description == name:
        view_description, desc2 = name, ""
    elif not description.startswith(name):
        view_description, desc2 = description, ""
    else:
        remainder = description[len(name):]
        if remainder.startswith(" "):
            remainder = remainder[1:]
        view_description, desc2 = name, remainder
    return {
        "item_code": row["code"],
        "description": view_description,
        "desc_2": desc2,
        "item_group": row.get("category_code"),
        "item_brand": row.get("brand_code"),
        "price": row["list_price"],
        "is_active": row["is_active"],
    }


def _template_row(row: dict) -> dict:
    """The manual-import template shape, built FROM a raw snapshot row - the
    captain's CM-4/PC-5 construction (Description = row `name`, Desc 2 = the
    computed remainder)."""
    view = _expected_view_row(row)
    return {
        "Item Code": view["item_code"],
        "Description": view["description"],
        "Desc 2": view["desc_2"],
        "Item Group": view["item_group"] or "",
        "Item Brand": view["item_brand"] or "",
        "Price": str(view["price"]),
        "Is Active": "TRUE" if view["is_active"] else "FALSE",
    }


def _num(value) -> Decimal:
    return Decimal(str(value))


def _seed_review_job(env, *, owner, entity="products", phase="review", rows=None):
    """Seeds a pull job in `phase` and points the fake FoundryX's `/rows` at
    `rows` (default: the committed 10-row products fixture)."""
    rows = rows if rows is not None else _fixture("products-rows-page1.json")["rows"]
    snapshot_id = f"{MARKER}-snap-{uuid.uuid4().hex[:8]}"
    job_type = "autocount_products_pull" if entity == "products" else "autocount_stock_pull"
    job_id = _seed_pull_job(
        env.db, job_type=job_type, user_id=owner["id"], company_id=env.company_a,
        entity=entity, company_code=env.company_a_code, snapshot_id=snapshot_id,
        phase=phase, header=_stored(_header(rows)),
    )
    env.fake.rows = {
        1: (200, {"snapshotId": snapshot_id, "page": 1, "pageSize": 1000, "totalPages": 1,
                  "recordCount": len(rows), "rows": rows})
    }
    return job_id, rows


# ======================================================================= RV3


class TestRowsRoute:
    def test_rv_3a_shape_key_order_and_pagination(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, rows = _seed_review_job(env, owner=owner)

        resp = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"page": 2, "limit": 3})

        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["pagination"] == {"page": 2, "limit": 3, "total": 10, "total_pages": 4}
        assert len(body["data"]) == 3

        expected_slice = rows[3:6]  # page 2, limit 3 -> 0-indexed rows 3,4,5
        for got, raw in zip(body["data"], expected_slice):
            assert list(got.keys()) == [
                "item_code", "description", "desc_2", "item_group", "item_brand",
                "price", "is_active",
            ], got
            expected = _expected_view_row(raw)
            assert got["item_code"] == expected["item_code"]
            assert got["description"] == expected["description"]
            assert got["desc_2"] == expected["desc_2"]
            assert got["item_group"] == expected["item_group"]
            assert got["item_brand"] == expected["item_brand"]
            # Captain ruling (Phase 3 fix round, F-9): `price` must be a real number, not a
            # string the FE has to parse - the string-blind `_num()` compare alone would not
            # catch a type regression here.
            assert isinstance(got["price"], (int, float)), got["price"]
            assert _num(got["price"]) == _num(expected["price"])
            assert got["is_active"] == expected["is_active"]

    def test_rv_3b_desc2_remainder_double_space_byte_exact_and_equal_case_empty(self, env):
        """The double-space item keeps ONE leading space in `desc_2` - the
        rule's whole point: `name + " " + desc_2` (the manual join formula)
        must reproduce the pull's raw double-spaced `description` byte for
        byte, and that only works if `desc_2` itself carries the extra space
        rather than having it stripped away."""
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, rows = _seed_review_job(env, owner=owner)

        resp = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"limit": 50})
        assert resp.status_code == 200, resp.text
        by_code = {r["item_code"]: r for r in resp.json()["data"]}

        double_space = next(r for r in rows if r["code"] == "SRTSH9112-GM")
        assert by_code["SRTSH9112-GM"]["description"] == double_space["name"]
        assert by_code["SRTSH9112-GM"]["desc_2"] == " WITH LED LIGHT"
        # The round-trip itself, not just the literal.
        rejoined = f'{by_code["SRTSH9112-GM"]["description"]} {by_code["SRTSH9112-GM"]["desc_2"]}'.strip()
        assert rejoined == double_space["description"]

        equal_case = next(r for r in rows if r["code"] == "A611")
        assert equal_case["description"] == equal_case["name"]  # the fixture's own premise
        assert by_code["A611"]["desc_2"] == ""

    def test_rv_3b_description_not_starting_with_name_keeps_the_full_text(self, env):
        """A shape no committed fixture row hits, but the rule must still
        answer something: when `description` does not even start with `name`,
        there is no boundary to split on - `desc_2` is empty and the
        `description` CELL carries the full raw text, not the (wrong) `name`."""
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        hand_row = _canonical_row(
            f"{MARKER}-NOMATCH", name="Foo", description="Completely different text",
        )
        job_id, _rows = _seed_review_job(env, owner=owner, rows=[hand_row])

        resp = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"limit": 50})
        assert resp.status_code == 200, resp.text
        row = resp.json()["data"][0]
        assert row["description"] == "Completely different text"
        assert row["desc_2"] == ""

    def test_rv_3c_query_filters_by_item_code_case_insensitive(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, rows = _seed_review_job(env, owner=owner)

        resp = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"query": "acc", "limit": 50})

        assert resp.status_code == 200, resp.text
        codes = {r["item_code"] for r in resp.json()["data"]}
        assert codes == {"ACC-SRT8001", "ACC-SRT9013", "ACC-SRT1001"}

    def test_rv_3d_non_owner_404(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_review_job(env, owner=owner)

        other = env.user("master_data.products.autocount_pull")
        env.as_user(other)
        resp = env.client.get(f"{PULLS_URL}/{job_id}/rows")
        assert resp.status_code == 404
        # A route that does not exist AT ALL also answers a bare Starlette 404
        # (`{"detail": "Not Found"}`) - checked, empirically, before this test
        # was written. Pinning the AppException shape (`code`) is what makes
        # this fail now (route missing) and pass only once the real
        # owner-check 404 exists, not a moment before.
        assert resp.json().get("code") == "NOT_FOUND", resp.json()

    def test_rv_3d_products_pull_read_by_stock_only_owner_is_403(self, env):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_review_job(env, owner=owner, entity="products")

        resp = env.client.get(f"{PULLS_URL}/{job_id}/rows")
        assert resp.status_code == 403

    def test_rv_3d_building_pull_is_409(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_review_job(env, owner=owner, phase="building")

        resp = env.client.get(f"{PULLS_URL}/{job_id}/rows")
        assert resp.status_code == 409


# ======================================================================= RV5


class TestDownloadRoute:
    def test_rv_5a_header_row_count_and_values_equal_rows(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, rows = _seed_review_job(env, owner=owner)

        rows_resp = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"limit": 50})
        assert rows_resp.status_code == 200, rows_resp.text
        api_rows = rows_resp.json()["data"]
        assert len(api_rows) == 10

        resp = env.client.get(f"{PULLS_URL}/{job_id}/download.xlsx")
        assert resp.status_code == 200, resp.text

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert header == [
            "Item Code", "Description", "Desc 2", "Item Group", "Item Brand",
            "Price", "Is Active", "UOM",
        ]

        body_rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(body_rows) == 10
        # openpyxl reads a cell written with an empty string back as None (a
        # library quirk, not a data difference) - normalise BOTH sides (the
        # JSON `/rows` side can itself be None, e.g. MOCHA's item_brand)
        # before comparing.
        for sheet_row, api_row in zip(body_rows, api_rows):
            assert (sheet_row[0] or "") == (api_row["item_code"] or "")
            assert (sheet_row[1] or "") == (api_row["description"] or "")
            assert (sheet_row[2] or "") == (api_row["desc_2"] or "")
            assert (sheet_row[3] or "") == (api_row["item_group"] or "")
            assert (sheet_row[4] or "") == (api_row["item_brand"] or "")
            assert _num(sheet_row[5]) == _num(api_row["price"])
            assert bool(sheet_row[6]) == bool(api_row["is_active"])
            assert sheet_row[7] in (None, ""), "products add a trailing BLANK UOM column"

    def test_rv_5b_content_type_and_filename_present(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_review_job(env, owner=owner)

        resp = env.client.get(f"{PULLS_URL}/{job_id}/download.xlsx")

        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        disposition = resp.headers.get("content-disposition", "")
        assert "filename" in disposition
        # D3 (small-fix track): the products download name is untouched by the D3 fix -
        # only the stock download/archive names changed.
        assert "autocount-products-pull.xlsx" in disposition, disposition


# ======================================================================= CM2


class TestComparePureFunction:
    """`app.services.autocount_pull_compare.compare_products(excel_rows,
    pull_rows)` - a pure function, no database (plan: "two pure functions over
    lists of dicts"). `excel_rows` carry the manual-template column names (what
    the browser posts, parsed from the checker's own file); `pull_rows` carry
    the RAW snapshot-row shape (`code`/`name`/`description`/`category_code`/
    `brand_code`/`list_price`/`is_active`) - AC-CM-2 itself says the join is
    "compared to the pull's `description`", the raw field, never the mapped
    `/rows` shape.

    Working definition of the return shape (the route's own body minus the
    storage-only `filename`/`compared_at`, per the captain's brief): a `summary`
    dict with `total`/`matched`/`different` COUNTS, and top-level `differences`
    / `only_in_excel` / `only_in_pull`. `total`/`matched` are over the
    INTERSECTION (items present in both sides) - AC-CM-4's "matched == total"
    only makes sense read that way.
    """

    def test_cm_2a_key_trim_and_case_insensitive_match(self):
        from app.services.autocount_pull_compare import compare_products

        pull_row = _canonical_row(
            "ABC-1", name="Item", description="Item", category_code="CAT",
            brand_code="BR", list_price="10.00", is_active=True,
        )
        excel_row = {
            "Item Code": " abc-1 ", "Description": "Item", "Desc 2": "",
            "Item Group": "CAT", "Item Brand": "BR", "Price": "10.00", "Is Active": "TRUE",
        }

        result = compare_products([excel_row], [pull_row])

        assert result["differences"] == []
        assert result["only_in_excel"] == []
        assert result["only_in_pull"] == []
        assert result["summary"]["matched"] == 1
        assert result["summary"]["total"] == 1

    def test_cm_2b_description_join_matches_pull_description_double_space_byte_exact(self):
        from app.services.autocount_pull_compare import compare_products

        pull_row = _canonical_row(
            "SRTSH9112-GM",
            name="SORENTO SERENE CEILING MOUNTED RAIN SHOWER HEAD SRTSH9112-GM",
            description=(
                "SORENTO SERENE CEILING MOUNTED RAIN SHOWER HEAD SRTSH9112-GM  WITH LED LIGHT"
            ),
        )
        excel_row = {
            "Item Code": "SRTSH9112-GM",
            # Trailing space on Description + a clean Desc 2 reproduces the SAME
            # double space the manual join formula (`f"{d} {d2}".strip()`) would
            # have produced from the real raw AutoCount fields.
            "Description": "SORENTO SERENE CEILING MOUNTED RAIN SHOWER HEAD SRTSH9112-GM ",
            "Desc 2": "WITH LED LIGHT",
            "Item Group": pull_row["category_code"], "Item Brand": pull_row["brand_code"],
            "Price": pull_row["list_price"], "Is Active": "TRUE",
        }

        result = compare_products([excel_row], [pull_row])

        description_diffs = [d for d in result["differences"] if d["field"] == "description"]
        assert description_diffs == [], result["differences"]

    def test_cm_2b2_single_space_excel_vs_double_space_pull_is_a_strict_difference(self):
        """Captain ruling (SR3 amend round): compare is STRICT, never
        whitespace-normalising. An Excel file whose Description/Desc 2 join to
        a SINGLE space where the pull's raw `description` carries a double
        space is a real, reportable difference - the coder's own added
        normalisation on the description compare is overruled."""
        from app.services.autocount_pull_compare import compare_products

        pull_row = _canonical_row(
            "SRTSH9112-GM",
            name="SORENTO SERENE CEILING MOUNTED RAIN SHOWER HEAD SRTSH9112-GM",
            description=(
                "SORENTO SERENE CEILING MOUNTED RAIN SHOWER HEAD SRTSH9112-GM  WITH LED LIGHT"
            ),
        )
        excel_row = {
            "Item Code": "SRTSH9112-GM",
            # No trailing space on Description this time - the join formula's
            # own single separator space is the ONLY space between the two
            # halves, one short of the pull's real double space.
            "Description": "SORENTO SERENE CEILING MOUNTED RAIN SHOWER HEAD SRTSH9112-GM",
            "Desc 2": "WITH LED LIGHT",
            "Item Group": pull_row["category_code"], "Item Brand": pull_row["brand_code"],
            "Price": pull_row["list_price"], "Is Active": "TRUE",
        }

        result = compare_products([excel_row], [pull_row])

        description_diffs = [d for d in result["differences"] if d["field"] == "description"]
        assert len(description_diffs) == 1, result["differences"]
        assert description_diffs[0]["item_code"] == "SRTSH9112-GM"

    def test_cm_2c_item_group_and_item_brand_reported_as_differences(self):
        from app.services.autocount_pull_compare import compare_products

        pull_row = _canonical_row("ITM-1", category_code="CATX", brand_code="BRX")
        excel_row = {
            "Item Code": "ITM-1", "Description": pull_row["name"], "Desc 2": "",
            "Item Group": "CATY", "Item Brand": "BRY",
            "Price": pull_row["list_price"], "Is Active": "TRUE",
        }

        result = compare_products([excel_row], [pull_row])

        fields = {d["field"] for d in result["differences"]}
        assert "item_group" in fields, result["differences"]
        assert "item_brand" in fields, result["differences"]
        group_diff = next(d for d in result["differences"] if d["field"] == "item_group")
        assert group_diff["excel"] == "CATY"
        assert group_diff["pull"] == "CATX"

    def test_cm_2d_price_numeric_compare_string_forms_and_negative_clamp(self):
        from app.services.autocount_pull_compare import compare_products

        matching_row = _canonical_row("PRC-1", list_price="63.00")
        excel_matching = {
            "Item Code": "PRC-1", "Description": matching_row["name"], "Desc 2": "",
            "Item Group": matching_row["category_code"], "Item Brand": matching_row["brand_code"],
            "Price": "63", "Is Active": "TRUE",  # "63" == "63.00" numerically
        }
        clamp_row = _canonical_row("PRC-2", list_price="0.0")
        excel_clamp = {
            "Item Code": "PRC-2", "Description": clamp_row["name"], "Desc 2": "",
            "Item Group": clamp_row["category_code"], "Item Brand": clamp_row["brand_code"],
            "Price": "-1", "Is Active": "TRUE",  # the file's own negative reads as 0
        }

        result = compare_products([excel_matching, excel_clamp], [matching_row, clamp_row])

        price_diffs = [d for d in result["differences"] if d["field"] == "price"]
        assert price_diffs == [], result["differences"]

    def test_cm_2e_is_active_by_manual_truthy_rule(self):
        """`bulk_import_products`'s own rule: is_active is False only for the
        spellings F/FALSE/0/N/NO (case-insensitive); everything else, including
        blank, is active."""
        from app.services.autocount_pull_compare import compare_products

        active_pull = _canonical_row("ACT-1", is_active=True)
        active_excel = {
            "Item Code": "ACT-1", "Description": active_pull["name"], "Desc 2": "",
            "Item Group": active_pull["category_code"], "Item Brand": active_pull["brand_code"],
            "Price": active_pull["list_price"], "Is Active": "Y",  # not a falsy spelling
        }
        inactive_pull = _canonical_row("ACT-2", is_active=False)
        inactive_excel = {
            "Item Code": "ACT-2", "Description": inactive_pull["name"], "Desc 2": "",
            "Item Group": inactive_pull["category_code"], "Item Brand": inactive_pull["brand_code"],
            "Price": inactive_pull["list_price"], "Is Active": "N",  # a falsy spelling
        }

        result = compare_products([active_excel, inactive_excel], [active_pull, inactive_pull])

        active_diffs = [d for d in result["differences"] if d["field"] == "is_active"]
        assert active_diffs == [], result["differences"]

    def test_cm_2f_one_difference_entry_per_differing_field(self):
        from app.services.autocount_pull_compare import compare_products

        pull_row = _canonical_row("MULTI-1", category_code="CATX", list_price="10.00")
        excel_row = {
            "Item Code": "MULTI-1", "Description": pull_row["name"], "Desc 2": "",
            "Item Group": "CATY", "Item Brand": pull_row["brand_code"],
            "Price": "20.00", "Is Active": "TRUE",
        }

        result = compare_products([excel_row], [pull_row])

        by_field = {d["field"]: d for d in result["differences"] if d["item_code"] == "MULTI-1"}
        assert set(by_field) == {"item_group", "price"}, result["differences"]
        assert len(result["differences"]) == 2

    def test_cm_2g_only_in_excel_and_only_in_pull(self):
        from app.services.autocount_pull_compare import compare_products

        pull_row = _canonical_row("PULL-ONLY")
        excel_row = {
            "Item Code": "EXCEL-ONLY", "Description": "X", "Desc 2": "",
            "Item Group": "CAT", "Item Brand": "BR", "Price": "1.00", "Is Active": "TRUE",
        }

        result = compare_products([excel_row], [pull_row])

        assert result["only_in_excel"] == ["EXCEL-ONLY"]
        assert result["only_in_pull"] == ["PULL-ONLY"]


# =================================================================== CM route


class TestCompareRoute:
    def test_cm_4_file_built_from_fixture_rows_is_a_perfect_match(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, rows = _seed_review_job(env, owner=owner)
        template_rows = [_template_row(r) for r in rows]

        resp = env.client.post(
            f"{PULLS_URL}/{job_id}/compare",
            json={"filename": "check.xlsx", "rows": template_rows},
        )

        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["summary"]["matched"] == 10, body
        assert body["summary"]["total"] == 10, body
        assert body["differences"] == [], body
        assert body["only_in_excel"] == []
        assert body["only_in_pull"] == []

    def test_cm_5_summary_stored_rows_and_differences_are_not(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, rows = _seed_review_job(env, owner=owner)
        template_rows = [_template_row(r) for r in rows]

        before_products = env.db.execute(text("SELECT count(*) FROM products")).scalar()
        before_refs = env.db.execute(text("SELECT count(*) FROM integration_references")).scalar()

        resp = env.client.post(
            f"{PULLS_URL}/{job_id}/compare",
            json={"filename": "check.xlsx", "rows": template_rows},
        )
        assert resp.status_code == 200, resp.text

        after_products = env.db.execute(text("SELECT count(*) FROM products")).scalar()
        after_refs = env.db.execute(text("SELECT count(*) FROM integration_references")).scalar()
        assert after_products == before_products
        assert after_refs == before_refs

        get_resp = env.client.get(f"{PULLS_URL}/{job_id}")
        assert get_resp.status_code == 200, get_resp.text
        compare = get_resp.json()["compare"]
        assert compare["filename"] == "check.xlsx"
        assert compare["matched"] == 10
        assert compare["total"] == 10
        assert "compared_at" in compare

        row = _job_row(env.db, job_id)
        stored_compare = row["metadata"]["autocount_pull"]["compare"]
        assert "differences" not in stored_compare, stored_compare
        assert "rows" not in stored_compare, stored_compare
        assert stored_compare["matched"] == 10

    def test_cm_route_non_owner_404(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_review_job(env, owner=owner)

        other = env.user("master_data.products.autocount_pull")
        env.as_user(other)
        resp = env.client.post(
            f"{PULLS_URL}/{job_id}/compare", json={"filename": "x.xlsx", "rows": []},
        )
        assert resp.status_code == 404
        # See TestRowsRoute.test_rv_3d_non_owner_404 - a missing route ALSO
        # answers a bare 404, so the AppException `code` is what pins this.
        assert resp.json().get("code") == "NOT_FOUND", resp.json()

    def test_cm_route_wrong_entity_permission_403(self, env):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_review_job(env, owner=owner, entity="products")

        resp = env.client.post(
            f"{PULLS_URL}/{job_id}/compare", json={"filename": "x.xlsx", "rows": []},
        )
        assert resp.status_code == 403


# ======================================================================= PC1


class TestConfirmRoute:
    def test_pc_1a_confirm_creates_apply_job_enqueues_once_marks_confirmed(self, env, monkeypatch):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, rows = _seed_review_job(env, owner=owner)

        captured = []

        def _enqueue(func, *a, **k):
            captured.append({"func": func, "args": a, "kwargs": k})
            return MagicMock(id=str(uuid.uuid4()))

        monkeypatch.setattr("app.services.autocount_pull_service.enqueue_job", _enqueue)

        before_count = _job_count(env.db)
        resp = env.client.post(f"{PULLS_URL}/{job_id}/confirm")

        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["phase"] == "confirmed"
        apply_job_id = body["apply_job_id"]
        assert apply_job_id

        assert _job_count(env.db) == before_count + 1
        apply_row = _job_row(env.db, apply_job_id)
        assert apply_row is not None
        assert apply_row["job_type"] == "autocount_products_apply"
        assert str(apply_row["company_id"]) == env.company_a
        assert apply_row["user_id"] == owner["id"]
        apply_meta = apply_row["metadata"]["autocount_apply"]
        assert apply_meta["pull_job_id"] == str(job_id)
        assert apply_meta["entity"] == "products"
        assert apply_meta["snapshot_id"]

        assert len(captured) == 1
        assert captured[0]["kwargs"].get("queue_name", "imports") == "imports"
        assert getattr(captured[0]["func"], "__name__", "") == "apply_autocount_pull"
        assert env.fake.calls == [], "confirm itself makes no FoundryX call - the apply task does"

    def test_pc_1b_second_confirm_is_idempotent(self, env, monkeypatch):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, rows = _seed_review_job(env, owner=owner)

        captured = []
        monkeypatch.setattr(
            "app.services.autocount_pull_service.enqueue_job",
            lambda *a, **k: captured.append(1) or MagicMock(id=str(uuid.uuid4())),
        )

        first = env.client.post(f"{PULLS_URL}/{job_id}/confirm")
        second = env.client.post(f"{PULLS_URL}/{job_id}/confirm")

        assert first.status_code == 200, first.text
        assert second.status_code == 200, second.text
        assert second.json()["apply_job_id"] == first.json()["apply_job_id"]
        assert len(captured) == 1

    @pytest.mark.parametrize("phase", ["building", "previewing", "failed", "expired"])
    def test_pc_1c_non_review_phase_is_409(self, env, phase):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, rows = _seed_review_job(env, owner=owner, phase=phase)

        resp = env.client.post(f"{PULLS_URL}/{job_id}/confirm")
        assert resp.status_code == 409, resp.text

    def test_pc_1d_non_owner_404(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, rows = _seed_review_job(env, owner=owner)

        other = env.user("master_data.products.autocount_pull")
        env.as_user(other)
        resp = env.client.post(f"{PULLS_URL}/{job_id}/confirm")
        assert resp.status_code == 404
        # See TestRowsRoute.test_rv_3d_non_owner_404 - a missing route ALSO
        # answers a bare 404, so the AppException `code` is what pins this.
        assert resp.json().get("code") == "NOT_FOUND", resp.json()

    def test_pc_1d_wrong_entity_permission_403(self, env):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        job_id, rows = _seed_review_job(env, owner=owner, entity="products")

        resp = env.client.post(f"{PULLS_URL}/{job_id}/confirm")
        assert resp.status_code == 403


# ================================================================= apply task


def _seed_apply_job(db, *, user_id, company_id, entity, snapshot_id, pull_job_id=None):
    from app.models.job import ImportJob, JobStatus

    job_type = "autocount_products_apply" if entity == "products" else "autocount_stock_apply"
    meta = {
        "autocount_apply": {
            "pull_job_id": pull_job_id or str(uuid.uuid4()),
            "snapshot_id": snapshot_id,
            "entity": entity,
        }
    }
    job = ImportJob(
        id=uuid.uuid4(), job_id=str(uuid.uuid4()), job_type=job_type,
        status=JobStatus.QUEUED.value, user_id=user_id, company_id=company_id,
        job_metadata=meta,
    )
    db.add(job)
    db.commit()
    return job.id


def _run_apply(monkeypatch, factory, job_id):
    monkeypatch.setattr("app.tasks.autocount_pull_tasks.SessionLocal", factory, raising=False)
    from app.tasks.autocount_pull_tasks import apply_autocount_pull

    apply_autocount_pull(job_id)


def _set_rows_page(fake: _FakeFoundryX, snapshot_id: str, rows: list[dict]) -> None:
    fake.rows = {
        1: (200, {"snapshotId": snapshot_id, "page": 1, "pageSize": 1000, "totalPages": 1,
                  "recordCount": len(rows), "rows": rows})
    }


class TestApplyTaskGuards:
    """AC-PC-2: `apply_autocount_pull` re-fetches the SAME snapshot through
    `fetch_verified_snapshot` (SR1's own function, unchanged) and re-checks the
    AC-PP-1 guards before writing anything for real."""

    def test_pc_2a_snapshot_expired_410_fails_and_creates_nothing(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        job_user = str(uuid.uuid4())
        code = f"{MARKER}-PC2A"
        snapshot_id = f"{MARKER}-snap-pc2a"
        fake.status = (
            410,
            {"code": "SNAPSHOT_EXPIRED", "message": "expired", "companyCode": "SRT",
             "entity": "products"},
        )
        job_id = _seed_apply_job(
            db, user_id=job_user, company_id=DEFAULT_COMPANY_ID, entity="products",
            snapshot_id=snapshot_id,
        )

        _run_apply(monkeypatch, factory, job_id)

        row = _job_row(db, job_id)
        assert row["status"] == "failed", row
        blob = (row["error"] or "") + str(row["metadata"])
        assert "pull again" in blob.lower(), blob
        count = db.execute(
            text("SELECT count(*) FROM products WHERE product_code = :c"), {"c": code}
        ).scalar()
        assert count == 0

    def test_pc_2b_unknown_snapshot_404_fails_and_creates_nothing(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        job_user = str(uuid.uuid4())
        code = f"{MARKER}-PC2B"
        snapshot_id = f"{MARKER}-snap-pc2b"
        fake.status = (
            404,
            {"code": "UNKNOWN_SNAPSHOT", "message": "not found", "companyCode": "SRT",
             "entity": "products"},
        )
        job_id = _seed_apply_job(
            db, user_id=job_user, company_id=DEFAULT_COMPANY_ID, entity="products",
            snapshot_id=snapshot_id,
        )

        _run_apply(monkeypatch, factory, job_id)

        row = _job_row(db, job_id)
        assert row["status"] == "failed", row
        blob = (row["error"] or "") + str(row["metadata"])
        assert "pull again" in blob.lower(), blob
        count = db.execute(
            text("SELECT count(*) FROM products WHERE product_code = :c"), {"c": code}
        ).scalar()
        assert count == 0

    def test_pc_2c_guards_rechecked_incomplete_header_fails_and_writes_nothing(
        self, task_db, monkeypatch
    ):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        job_user = str(uuid.uuid4())
        code = f"{MARKER}-PC2C"
        rows = [_canonical_row(code)]
        snapshot_id = f"{MARKER}-snap-pc2c"
        fake.status = (200, {**_header(rows, complete=False), "snapshotId": snapshot_id})
        _set_rows_page(fake, snapshot_id, rows)
        job_id = _seed_apply_job(
            db, user_id=job_user, company_id=DEFAULT_COMPANY_ID, entity="products",
            snapshot_id=snapshot_id,
        )

        _run_apply(monkeypatch, factory, job_id)

        row = _job_row(db, job_id)
        assert row["status"] == "failed", row
        assert (row["error"] or "").strip() != ""
        assert _job_rows(db, job_id) == []
        count = db.execute(
            text("SELECT count(*) FROM products WHERE product_code = :c"), {"c": code}
        ).scalar()
        assert count == 0


class TestApplyTaskRealIngest:
    def test_pc_3_ten_fixture_rows_against_an_empty_company_all_create(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        job_user = str(uuid.uuid4())
        rows = _fixture("products-rows-page1.json")["rows"]
        snapshot_id = f"{MARKER}-snap-pc3"
        fake.status = (200, {**_header(rows), "snapshotId": snapshot_id})
        _set_rows_page(fake, snapshot_id, rows)
        job_id = _seed_apply_job(
            db, user_id=job_user, company_id=DEFAULT_COMPANY_ID, entity="products",
            snapshot_id=snapshot_id,
        )

        _run_apply(monkeypatch, factory, job_id)

        row = _job_row(db, job_id)
        assert row["status"] == "finished", row["error"]

        persisted = db.execute(
            text("SELECT count(*) FROM products WHERE product_code = ANY(:codes)"),
            {"codes": [r["code"] for r in rows]},
        ).scalar()
        assert persisted == 10

        written = _job_rows(db, job_id)
        created = [r for r in written if r["outcome"] == "created"]
        assert len(created) == 10, written
        assert not [r for r in written if r["outcome"] == "failed"], written


class TestApplyTaskStamping:
    """PC-4: `stamp_user_id` (a new, optional `MasterIngestService.__init__`
    keyword, per the plan) sets `created_by` + `updated_by` on `_insert` and
    only `updated_by` on `_update`; the existing push path (no `stamp_user_id`
    at all) stamps neither (PC-4c) - a regression guard, not new behaviour, so
    it may already be green."""

    def test_pc_4a_created_products_carry_the_confirming_user(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        job_user = str(uuid.uuid4())
        code = f"{MARKER}-PC4A"
        rows = [_canonical_row(code)]
        snapshot_id = f"{MARKER}-snap-pc4a"
        fake.status = (200, {**_header(rows), "snapshotId": snapshot_id})
        _set_rows_page(fake, snapshot_id, rows)
        job_id = _seed_apply_job(
            db, user_id=job_user, company_id=DEFAULT_COMPANY_ID, entity="products",
            snapshot_id=snapshot_id,
        )

        _run_apply(monkeypatch, factory, job_id)

        assert _job_row(db, job_id)["status"] == "finished"
        created_by, updated_by = db.execute(
            text("SELECT created_by, updated_by FROM products WHERE product_code = :c"), {"c": code}
        ).first()
        assert str(created_by) == job_user
        assert str(updated_by) == job_user

    def test_pc_4b_updated_product_keeps_its_original_created_by(self, task_db, monkeypatch):
        from app.models.product import Product, ProductCategory, UnitOfMeasure

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        job_user = str(uuid.uuid4())
        original_creator = str(uuid.uuid4())
        code = f"{MARKER}-PC4B"

        cat = ProductCategory(category_code=unique_code(MARKER), category_name="cat")
        uom = UnitOfMeasure(uom_code=unique_code(MARKER)[:20], uom_name="unit")
        db.add_all([cat, uom])
        db.flush()
        db.add(
            Product(
                product_code=code, product_name=code, category_id=cat.id, base_uom_id=uom.id,
                list_price=Decimal("1.00"), company_id=DEFAULT_COMPANY_ID,
                created_by=original_creator,
            )
        )
        db.commit()

        rows = [_canonical_row(code, list_price="99.00")]
        snapshot_id = f"{MARKER}-snap-pc4b"
        fake.status = (200, {**_header(rows), "snapshotId": snapshot_id})
        _set_rows_page(fake, snapshot_id, rows)
        job_id = _seed_apply_job(
            db, user_id=job_user, company_id=DEFAULT_COMPANY_ID, entity="products",
            snapshot_id=snapshot_id,
        )

        _run_apply(monkeypatch, factory, job_id)

        assert _job_row(db, job_id)["status"] == "finished"
        created_by, updated_by = db.execute(
            text("SELECT created_by, updated_by FROM products WHERE product_code = :c"), {"c": code}
        ).first()
        # A raw text() SELECT hands back a native uuid.UUID for a postgres
        # uuid column (see LESSONS-LEARNT "uuid-id stack" / master_ingest_
        # service._value_changed's own note on this) - stringify before
        # comparing to the plain str id this test seeded.
        assert str(created_by) == original_creator, "an update must never touch created_by"
        assert str(updated_by) == job_user

    def test_pc_4c_push_path_without_stamp_user_id_stamps_neither(self, task_db):
        db, factory = task_db
        from app.services.master_ingest_service import MasterIngestService

        code = f"{MARKER}-PC4C"
        svc = MasterIngestService(db, company_id=DEFAULT_COMPANY_ID)
        result = svc.ingest("products", [_canonical_row(code)])
        assert not result.records[0].errors, result.records[0].errors

        created_by, updated_by = db.execute(
            text("SELECT created_by, updated_by FROM products WHERE product_code = :c"), {"c": code}
        ).first()
        assert created_by is None
        assert updated_by is None


# ================================================================= parity gate


class TestParityGate:
    """AC-PC-5 (R10). Two scratch companies, one shared `blank_session()` (the
    `test_ingest_parity_s1_products.py` pattern - proven to work in this repo,
    confirmed by running that file's own equivalent test before writing this
    one): the 10 fixture items through `ProductService.bulk_import_products`
    (manual/Excel path, template rows built FROM the fixture - Description =
    row `name`, Desc 2 = the computed remainder) into company X, and the SAME
    10 raw fixture rows through `MasterIngestService(db, company_id=...,
    stamp_user_id=...)` (the pull/apply path, called directly rather than
    through two full `apply_autocount_pull` job cycles across two companies -
    that would add FoundryX-fetch machinery PC-2/PC-3/PC-4 already cover for no
    extra parity coverage; said in the tester report) into company I. Compared
    by raw SQL, joined to category/brand/uom by CODE (each company gets its OWN
    category/brand/uom row for the same code text - only `sales_agents` is a
    shared table, D11, irrelevant here).
    """

    @pytest.fixture()
    def db(self):
        with blank_session() as session:
            yield session

    @pytest.fixture()
    def companies(self, db):
        from app.models.company import Company

        x = Company(id=str(uuid.uuid4()), name=f"{MARKER} X", code=f"{MARKER}X{uuid.uuid4().hex[:8]}")
        i = Company(id=str(uuid.uuid4()), name=f"{MARKER} I", code=f"{MARKER}I{uuid.uuid4().hex[:8]}")
        db.add_all([x, i])
        db.flush()
        return str(x.id), str(i.id)

    def _seed_supplier(self, db, company_id: str, code: str) -> str:
        from app.models.procurement import Supplier

        supplier = Supplier(
            id=str(uuid.uuid4()), supplier_code=code, supplier_name=code, company_id=company_id
        )
        db.add(supplier)
        db.flush()
        return supplier.id

    def _product_row(self, db, company_id: str, code: str):
        row = db.execute(
            text(
                "SELECT p.product_code, p.product_name, p.description, p.list_price, "
                "p.is_active, p.is_discontinued, p.dimensions_length, p.dimensions_width, "
                "p.dimensions_height, p.brand_id, p.base_uom_id, "
                "c.category_code, b.brand_code, u.uom_code "
                "FROM products p "
                "JOIN product_categories c ON c.id = p.category_id "
                "LEFT JOIN brands b ON b.id = p.brand_id "
                "JOIN units_of_measure u ON u.id = p.base_uom_id "
                "WHERE p.company_id = :cid AND p.product_code = :code"
            ),
            {"cid": company_id, "code": code},
        ).mappings().first()
        return dict(row) if row else None

    def _supplier_link(self, db, company_id: str, code: str):
        row = db.execute(
            text(
                "SELECT s.supplier_code, ps.standard_lead_time_days "
                "FROM product_suppliers ps "
                "JOIN products p ON p.id = ps.product_id "
                "JOIN suppliers s ON s.id = ps.supplier_id "
                "WHERE p.company_id = :cid AND p.product_code = :code"
            ),
            {"cid": company_id, "code": code},
        ).mappings().first()
        return dict(row) if row else None

    def test_pc_5_manual_and_pull_paths_land_identical_products(self, db, companies):
        from app.services.master_ingest_service import MasterIngestService
        from app.services.product_service import ProductService

        company_x, company_i = companies
        supplier_code = unique_code(MARKER)[:20]
        self._seed_supplier(db, company_x, supplier_code)
        self._seed_supplier(db, company_i, supplier_code)
        db.commit()

        rows = _fixture("products-rows-page1.json")["rows"]
        excel_rows = [_template_row(r) for r in rows]
        excel_user = str(uuid.uuid4())
        pull_user = str(uuid.uuid4())

        set_company_scope(db, frozenset({company_x}))
        ProductService(db).bulk_import_products(excel_rows, user_id=excel_user)

        set_company_scope(db, frozenset({company_i}))
        ingest = MasterIngestService(db, company_id=company_i, stamp_user_id=pull_user)
        result = ingest.ingest("products", rows)
        assert all(not r.errors for r in result.records), [r.errors for r in result.records]

        compare_fields = [
            "product_name", "description", "list_price", "is_active", "is_discontinued",
            "dimensions_length", "dimensions_width", "dimensions_height",
            "category_code", "uom_code",
        ]
        for raw in rows:
            code = raw["code"]
            x_row = self._product_row(db, company_x, code)
            i_row = self._product_row(db, company_i, code)
            assert x_row is not None, f"{code} missing on the manual side"
            assert i_row is not None, f"{code} missing on the pull side"
            diff = {f: (x_row[f], i_row[f]) for f in compare_fields if x_row[f] != i_row[f]}
            assert diff == {}, (code, diff)

            x_supplier = self._supplier_link(db, company_x, code)
            i_supplier = self._supplier_link(db, company_i, code)
            assert x_supplier is not None and i_supplier is not None, code
            assert x_supplier["supplier_code"] == i_supplier["supplier_code"] == supplier_code
            assert x_supplier["standard_lead_time_days"] == i_supplier["standard_lead_time_days"]

        # Named sub-asserts (R10).
        assert self._product_row(db, company_x, "ACC-SRT8001")["is_discontinued"] is True
        assert self._product_row(db, company_i, "ACC-SRT8001")["is_discontinued"] is True

        dims_x = self._product_row(db, company_x, "AP4842")
        dims_i = self._product_row(db, company_i, "AP4842")
        assert (
            (dims_x["dimensions_length"], dims_x["dimensions_width"], dims_x["dimensions_height"])
            == (dims_i["dimensions_length"], dims_i["dimensions_width"], dims_i["dimensions_height"])
        )
        assert dims_x["dimensions_length"] == Decimal("480.00")

        join_x = self._product_row(db, company_x, "SRTSH9112-GM")["description"]
        join_i = self._product_row(db, company_i, "SRTSH9112-GM")["description"]
        assert join_x == join_i
        assert "  " in join_x, "the double space is real, not collapsed"

        for code in ("ACC-SRT9013", '1/2" ULTRA CIRCULAR'):
            assert self._product_row(db, company_x, code)["list_price"] == Decimal("0.00")
            assert self._product_row(db, company_i, code)["list_price"] == Decimal("0.00")

        assert self._product_row(db, company_x, "B2155-BLUE-DIY") is not None
        assert self._product_row(db, company_i, "B2155-BLUE-DIY") is not None

    def test_pc_5c_reconfirm_leaves_a_hand_set_brand_and_uom_excel_clears_them(self, db, companies):
        """AC-PC-5's two named differences (blank-brand row, UOM) are UPDATE-
        time behaviour, not create-time: on a first CREATE both paths
        independently land brand=NULL and uom=default for the MOCHA row (no
        `brand_code` / no `uom_code` on either input - proven for the whole 10
        by the parity test above), so there is nothing to observe yet. The
        difference only becomes OBSERVABLE once something else has set a
        brand/UOM on the row and the SAME row is re-synced: `bulk_import_
        products` unconditionally overwrites `brand_id`/`base_uom_id` on every
        update (`product_service.py`, `existing.brand_id = brand_id` /
        `existing.base_uom_id = uom_id`, both unconditional); the ingest
        service (D14) only ever touches a column that is actually present in
        the incoming payload, and MOCHA's payload never carries `brand_code` /
        `uom_code` at all - so a re-pull leaves a hand-set value alone. Modelled
        here as a re-sync after a hand edit - the realistic trigger for this
        difference to matter, and the only way to make it assertable rather
        than restating the plan's prose."""
        from app.models.product import Brand, UnitOfMeasure
        from app.services.master_ingest_service import MasterIngestService
        from app.services.product_service import ProductService

        company_x, company_i = companies
        mocha = next(
            r for r in _fixture("products-rows-page1.json")["rows"] if r["code"] == "MWT2800-N/H"
        )
        assert "brand_code" not in mocha and "uom_code" not in mocha  # the fixture's own premise
        excel_row = _template_row(mocha)
        excel_user = str(uuid.uuid4())
        pull_user = str(uuid.uuid4())

        set_company_scope(db, frozenset({company_x}))
        ProductService(db).bulk_import_products([excel_row], user_id=excel_user)
        set_company_scope(db, frozenset({company_i}))
        MasterIngestService(db, company_id=company_i, stamp_user_id=pull_user).ingest(
            "products", [mocha]
        )
        db.commit()

        # Hand-set a brand + a non-default UOM on BOTH created rows, simulating
        # an edit AutoCount does not know about.
        old_brand_x = Brand(
            id=str(uuid.uuid4()), brand_code=unique_code(MARKER)[:20], brand_name="Old",
            company_id=company_x,
        )
        old_brand_i = Brand(
            id=str(uuid.uuid4()), brand_code=unique_code(MARKER)[:20], brand_name="Old",
            company_id=company_i,
        )
        old_uom_x = UnitOfMeasure(
            id=str(uuid.uuid4()), uom_code=f"{MARKER}BX{uuid.uuid4().hex[:10]}", uom_name="Box",
            company_id=company_x,
        )
        old_uom_i = UnitOfMeasure(
            id=str(uuid.uuid4()), uom_code=f"{MARKER}BI{uuid.uuid4().hex[:10]}", uom_name="Box",
            company_id=company_i,
        )
        db.add_all([old_brand_x, old_brand_i, old_uom_x, old_uom_i])
        db.flush()
        db.execute(
            text(
                "UPDATE products SET brand_id = :b, base_uom_id = :u "
                "WHERE company_id = :cid AND product_code = :code"
            ),
            {"b": old_brand_x.id, "u": old_uom_x.id, "cid": company_x, "code": mocha["code"]},
        )
        db.execute(
            text(
                "UPDATE products SET brand_id = :b, base_uom_id = :u "
                "WHERE company_id = :cid AND product_code = :code"
            ),
            {"b": old_brand_i.id, "u": old_uom_i.id, "cid": company_i, "code": mocha["code"]},
        )
        db.commit()

        # Re-sync the SAME row through both paths again.
        set_company_scope(db, frozenset({company_x}))
        ProductService(db).bulk_import_products([excel_row], user_id=excel_user)
        set_company_scope(db, frozenset({company_i}))
        MasterIngestService(db, company_id=company_i, stamp_user_id=pull_user).ingest(
            "products", [mocha]
        )
        db.commit()

        x_row = self._product_row(db, company_x, mocha["code"])
        i_row = self._product_row(db, company_i, mocha["code"])

        assert x_row["brand_id"] is None, "Excel clears a brand the row no longer states"
        assert str(i_row["brand_id"]) == old_brand_i.id, "the pull leaves a brand it never sent"

        assert x_row["uom_code"] != old_uom_x.uom_code, "Excel re-stamps the default UOM every re-import"
        assert i_row["uom_code"] == old_uom_i.uom_code, "the pull leaves a UOM it never sent"
