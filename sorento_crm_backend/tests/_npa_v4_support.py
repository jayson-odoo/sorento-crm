"""Shared helpers for the NEW-PRODUCT-APP v4 (paper layout) suites, staff and portal.

Test support only: no implementation of the feature lives here. Storage is faked (no real
S3/R2), images are real bytes made with Pillow, and the workbook is read back with openpyxl.
"""
from __future__ import annotations

import io
import uuid
from datetime import datetime
from urllib.parse import unquote

# What the paper form calls each row, in the order AC-NPA-77 fixes. Column A of every sheet.
SHEET_LABELS = [
    "NEW PRODUCT APPLICATION",
    "REFF:",
    "DATE:",
    "SALES PERSON:",
    "PRODUCT CODE:",
    "PRODUCT TYPE:",
    "DIMENSION:",
    "MATERIAL:",
    "THICKNESS:",
    "PACKING",
    "BRAND:",
    "IMAGE:",
    "CUSTOMER:",
    "PROJECT:",
    "QTY:",
    "TARGET DELIVERY:",
    "REMARK:",
    "PURCHASING FEEDBACK",
    "COST:",
    "LEADTIME:",
]

#: Keys the v4 write body refuses (AC-NPA-70). The value is one the OLD contract accepted
#: (a real row id for the references), so a red run proves the key itself is refused.
REFUSED_LINE_KEYS = [
    "category_id",
    "base_uom_id",
    "item_type_id",
    "supplier_id",
    "supplier_item_code",
    "cost_price",
    "selling_price",
    "further_description",
    "remark",
]


def refused_value(key: str, refs) -> object:
    """`refs` carries real ids: category, uom, item_type, supplier."""
    return {
        "category_id": refs.category,
        "base_uom_id": refs.uom,
        "item_type_id": refs.item_type,
        "supplier_id": refs.supplier,
        "supplier_item_code": "XY1234",
        "cost_price": "12.50",
        "selling_price": "99.00",
        "further_description": "more words",
        "remark": "line remark",
    }[key]


IMAGE_FORMATS = [
    ("PNG", "image/png", "photo.png"),
    ("JPEG", "image/jpeg", "photo.jpg"),
    ("WEBP", "image/webp", "photo.webp"),
]


def image_bytes(fmt: str = "PNG", size=(16, 12)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", size, (200, 30, 30)).save(buf, format=fmt)
    return buf.getvalue()


class FakeStorage:
    """The storage backend: bytes by key, nothing leaves the process."""

    def __init__(self) -> None:
        self.files: dict[str, bytes] = {}
        self.fetched: list[str] = []

    def upload_file(self, *args, **kwargs):
        if args:
            content, key = args[0], args[1]
        else:
            content, key = kwargs.get("file_content"), kwargs.get("file_path")
        self.files[key] = content
        return key, f"https://cdn.test/{key}"

    def download_file(self, key: str) -> bytes:
        self.fetched.append(key)
        return self.files[key]

    def get_signed_url(self, key: str, expires_in: int = 3600) -> str:
        return f"https://cdn.test/{key}?sig=1"

    def file_exists(self, key: str) -> bool:
        return key in self.files

    def delete_file(self, key: str) -> bool:
        return self.files.pop(key, None) is not None


def install_storage(monkeypatch) -> FakeStorage:
    import app.services.storage_router as storage_router

    fake = FakeStorage()
    monkeypatch.setattr(storage_router, "default_provider", lambda: "s3")
    monkeypatch.setattr(storage_router, "get_backend", lambda provider: fake)
    monkeypatch.setattr(storage_router, "cdn_base_url", lambda provider, key: f"https://cdn.test/{key}")
    return fake


def seed_attachment_type(db) -> None:
    from app.models.resources import AttachmentType

    if db.query(AttachmentType).filter(AttachmentType.code == "response_attachment").first():
        return
    db.add(
        AttachmentType(
            id=str(uuid.uuid4()),
            code="response_attachment",
            type_name="Response Attachment",
            allowed_extensions="jpg,jpeg,png,webp,gif,pdf,xlsx,txt",
            max_file_size_mb=10,
        )
    )
    db.flush()


def seed_attachment(db, store: FakeStorage, entity_type: str, entity_id: str, name: str, data: bytes, mime: str, size=None):
    """A file already attached to a record (bypasses every route). Returns (attachment_id, link_id)."""
    from app.services.entity_attachment_service import EntityAttachmentService

    seed_attachment_type(db)
    key = f"seed/{entity_type}/{entity_id}/{uuid.uuid4()}-{name}"
    store.files[key] = data
    link = EntityAttachmentService(db).create_attachment_and_link(
        entity_type=entity_type,
        entity_id=entity_id,
        file_url=key,
        file_name=name,
        file_size_bytes=size if size is not None else len(data),
        attachment_type_code="response_attachment",
        storage_provider="s3",
        mime_type=mime,
    )
    db.flush()
    return str(link.attachment_id), str(link.id)


def noon(year=2028, month=1, day=17) -> datetime:
    """Midday: the same calendar day in UTC and in MYT, so a date assertion cannot flip."""
    return datetime(year, month, day, 12, 0, 0)


# ------------------------------------------------------------------ workbook readers


def open_workbook(content: bytes):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(content))


def label_rows(ws) -> list[tuple[str, int]]:
    """(label, row) for every row that holds text: column A first, else column B."""
    out = []
    for row in range(1, ws.max_row + 1):
        for col in (1, 2):
            value = ws.cell(row=row, column=col).value
            if isinstance(value, str) and value.strip():
                out.append((value.strip(), row))
                break
    return out


def sheet_values(ws) -> dict:
    """label -> the column B cell on that row."""
    return {label: ws.cell(row=row, column=2) for label, row in label_rows(ws)}


def text_of(cell) -> str:
    value = cell.value
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    return "" if value is None else str(value)


def downloaded_name(res) -> str:
    header = res.headers.get("content-disposition", "")
    marker = "filename*=UTF-8''"
    assert marker in header, header
    return unquote(header.split(marker, 1)[1].split(";", 1)[0])


MB = 1024 * 1024


def bad_upload(case: str) -> tuple[str, str, bytes]:
    """(file name, declared content type, bytes) for an upload that must be refused with 422."""
    import io

    from PIL import Image

    if case == "over_10mb":
        return "big.png", "image/png", image_bytes("PNG") + b"\0" * (10 * MB)  # real image, trailing padding
    if case == "not_an_image":
        return "photo.png", "image/png", b"hello"
    if case == "png_declared_jpeg":
        return "photo.jpg", "image/jpeg", image_bytes("PNG")
    if case == "over_40mp":
        buf = io.BytesIO()
        Image.new("L", (7000, 6000), 128).save(buf, format="PNG")
        return "huge.png", "image/png", buf.getvalue()
    raise KeyError(case)


BAD_UPLOAD_CASES = ["over_10mb", "not_an_image", "png_declared_jpeg", "over_40mp"]


def record_upload_reads(monkeypatch) -> list:
    """Patch starlette's UploadFile.read to record the size each call asked for."""
    from starlette.datastructures import UploadFile

    sizes: list = []
    original = UploadFile.read

    async def spy(self, size: int = -1):
        sizes.append(size)
        return await original(self, size)

    monkeypatch.setattr(UploadFile, "read", spy)
    return sizes
