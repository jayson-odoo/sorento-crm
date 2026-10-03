"""Round 2 red tests from the security and code reviews (CONTACT-BRAND-SCOPE).

1. order analytics, 2. the remaining `filtered` REST routes, 3. cross-company brand ids,
4. audit of the brands PUT, 6. low stock workbook and 7. sales analysis workbook as the
WORKER builds them (a fresh session, no hand stamping).
"""
from __future__ import annotations

import io
import uuid
from decimal import Decimal

import pytest
from sqlalchemy.orm import Session

from tests._brand_scope_seed import api, db, leaks, world  # noqa: F401  (fixtures by name)

BASE = "/api/v1"
ANALYTICS = f"{BASE}/order-management/orders/analytics"


def _ghost_id() -> str:
    return str(uuid.uuid4())


# --------------------------------------------------------------------- 1. order analytics


def test_analytics_total_value_is_the_in_scope_figure(api, world) -> None:
    """Reviewer measured scoped 1116.0 where the MOCHA DO line is 100.0."""
    scoped = api.get(ANALYTICS, world.scoped, metric="total_value").json()
    assert scoped["groups"][0]["value"] == 100.0, scoped
    assert scoped["total"]["value"] == 100.0, scoped
    plain = api.get(ANALYTICS, world.unscoped, metric="total_value").json()
    assert plain["total"]["value"] == 1116.0, plain


def test_analytics_by_product_rows_and_total_come_from_in_scope_lines(api, world) -> None:
    scoped = api.get(ANALYTICS, world.scoped, metric="total_value", group_by="product").json()
    assert [g["group_key"] for g in scoped["groups"]] == [world.codes["mocha"]], scoped
    assert scoped["total"]["value"] == 100.0, scoped
    plain = api.get(ANALYTICS, world.unscoped, metric="total_value", group_by="product").json()
    assert {g["group_key"] for g in plain["groups"]} == set(world.codes.values())


def test_analytics_count_of_an_out_of_scope_product_reads_as_an_unknown_product(api, world) -> None:
    sid, ghost = str(world.p_sorento.id), _ghost_id()
    out = api.get(ANALYTICS, world.scoped, metric="count", product_ids=sid)
    miss = api.get(ANALYTICS, world.scoped, metric="count", product_ids=ghost)
    assert out.status_code == miss.status_code == 200
    assert out.json()["total"]["value"] == 0
    assert out.text.replace(sid, "<X>") == miss.text.replace(ghost, "<X>")
    plain = api.get(ANALYTICS, world.unscoped, metric="count", product_ids=sid).json()
    assert plain["total"]["value"] == 2


# --------------------------------------------------------------------- 2. other filtered routes

PO = f"{BASE}/procurement/purchase-orders"
OTHER_PATHS = [
    ("po_placed", f"{PO}/placed", {"include_summary": "true"}),
    ("po_last_cost", f"{PO}/last-cost", {"top_n": 5}),
    ("promotion_products", f"{BASE}/marketing/promotion-products", {}),
    ("complaints", f"{BASE}/complaints-management/complaints/", {}),
    ("complaint_analytics_by_product", f"{BASE}/complaints-management/complaints/analytics",
     {"group_by": "product"}),
]
OTHER_IDS = [p[0] for p in OTHER_PATHS]


@pytest.mark.parametrize(("name", "path", "params"), OTHER_PATHS, ids=OTHER_IDS)
def test_scoped_contact_sees_in_scope_products_only(api, world, name, path, params) -> None:
    resp = api.get(path, world.scoped, **params)
    assert resp.status_code == 200, resp.text
    seen = {k: v for k, v in leaks(resp, world).items()}
    # complaint analytics lower-cases its group keys
    low = resp.text.lower()
    seen = {k: (world.codes[k] in resp.text or world.codes[k].lower() in low) for k in world.codes}
    assert seen == {"mocha": True, "sorento": False, "null": False}, (name, seen)


@pytest.mark.parametrize(("name", "path", "params"), OTHER_PATHS, ids=OTHER_IDS)
def test_unscoped_contact_sees_all_three(api, world, name, path, params) -> None:
    resp = api.get(path, world.unscoped, **params)
    assert resp.status_code == 200, resp.text
    low = resp.text.lower()
    assert all(world.codes[k].lower() in low for k in world.codes), name


def test_po_placed_summary_counts_in_scope_lines_only(api, world) -> None:
    body = api.get(f"{PO}/placed", world.scoped, include_summary="true").json()
    assert body["summary"]["po_placed_qty"] == 10, body.get("summary")
    plain = api.get(f"{PO}/placed", world.unscoped, include_summary="true").json()
    assert plain["summary"]["po_placed_qty"] == 60, plain.get("summary")


def test_complaints_without_an_in_scope_product_are_not_returned(api, world) -> None:
    body = api.get(f"{BASE}/complaints-management/complaints/", world.scoped).json()
    numbers = {c["complaint_number"] for c in body["data"]}
    assert numbers == {world.complaints["mixed"].complaint_number}, numbers
    mixed = body["data"][0]
    assert world.codes["sorento"] not in (mixed.get("product_code") or "")


def test_complaint_analytics_totals_count_in_scope_complaints_only(api, world) -> None:
    scoped = api.get(f"{BASE}/complaints-management/complaints/analytics", world.scoped).json()
    assert scoped["total"]["value"] == 1, scoped
    plain = api.get(f"{BASE}/complaints-management/complaints/analytics", world.unscoped).json()
    assert plain["total"]["value"] == 3


# --------------------------------------------------------------------- 6/7. worker built files


class _Worker:
    """What the RQ worker does: a fresh session on the SAME rolled-back connection, no scope
    stamped by hand, plus a storage backend that keeps the uploaded bytes."""

    def __init__(self, db, monkeypatch) -> None:
        self.db, self.monkeypatch, self.uploads = db, monkeypatch, []

    def patch(self, module) -> None:
        bind = self.db.connection()

        class _Fresh(Session):
            def close(self):  # the test owns the connection
                return None

        worker = self
        monkeypatch = self.monkeypatch
        monkeypatch.setattr(
            module, "SessionLocal", lambda: _Fresh(bind=bind, join_transaction_mode="create_savepoint")
        )

        class _Backend:
            def upload_file(self, *, file_content, file_path, content_type):
                worker.uploads.append(file_content)
                return (file_path, None)

        monkeypatch.setattr(module, "default_provider", lambda: "s3")
        monkeypatch.setattr(module, "get_backend", lambda provider: _Backend())


def _job_kwargs(kwargs: dict) -> dict:
    """The job's own keyword arguments: what `enqueue_job` forwards minus its queue plumbing."""
    return {k: v for k, v in kwargs.items() if k not in ("queue_name", "job_timeout", "depends_on")}


def _cells(xlsx: bytes) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx))
    return [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None]


def test_low_stock_workbook_built_by_the_worker_holds_in_scope_rows_only(api, world, db, monkeypatch) -> None:
    """The route is called as the scoped contact; the queued export job is then run the way the
    worker runs it. The run or the job carries the brand ids, so the file needs no manual stamp."""
    import asyncio
    from datetime import date, datetime

    from app.api.v1.scm import low_stock_report as route
    from app.services import queue_service
    from app.models.scm import OrderSummaryRow
    from app.tasks import export_tasks

    queued: list[tuple] = []
    monkeypatch.setattr(queue_service, "enqueue_job", lambda fn, *a, **k: queued.append((fn, a, _job_kwargs(k))) or object())

    async def _timeout(*_a, **_k):
        raise asyncio.TimeoutError

    monkeypatch.setattr(route, "_await_download", _timeout)
    resp = api.get(f"{BASE}/scm/low-stock-report", world.scoped, dry_run="true")
    assert resp.status_code == 200 and resp.json().get("run_id"), resp.text
    run_id = resp.json()["run_id"]
    for prod in world.products:  # what the run job would have frozen
        db.add(OrderSummaryRow(
            id=str(uuid.uuid4()), run_id=run_id, product_id=prod.id, as_of=date(2026, 9, 10),
            computed_at=datetime(2026, 9, 10, 6, 0, 0), pool_on_hand=10, reorder_level=100, suggested_qty=0,
        ))
    db.commit()
    export = next(q for q in queued if q[0] is export_tasks.generate_low_stock_report)
    worker = _Worker(db, monkeypatch)
    worker.patch(export_tasks)
    result = export[0](*export[1], **export[2])
    assert result["status"] == "ready", result
    cells = _cells(worker.uploads[0])
    assert world.codes["mocha"] in cells
    assert world.codes["sorento"] not in cells and world.codes["null"] not in cells


def test_sales_analysis_workbook_built_by_the_worker_totals_in_scope_lines_only(api, world, db, monkeypatch) -> None:
    from app.services import queue_service
    from app.tasks import report_export_tasks

    queued: list[tuple] = []
    monkeypatch.setattr(queue_service, "enqueue_job", lambda fn, *a, **k: queued.append((fn, a, _job_kwargs(k))) or object())
    body = api.get(f"{BASE}/sales/analysis", world.scoped, basis="ordered").json()
    assert Decimal(str(body["totals"]["total"])) == Decimal("200.00"), body
    job = next(q for q in queued if q[0] is report_export_tasks.generate_report_xlsx)
    worker = _Worker(db, monkeypatch)
    worker.patch(report_export_tasks)
    result = job[0](*job[1], **job[2])
    assert result["status"] == "ready", result
    numbers = set()
    for c in _cells(worker.uploads[0]):
        try:
            numbers.add(round(float(c), 2))
        except ValueError:
            pass
    assert 200.0 in numbers, sorted(numbers)
    assert 2225.0 not in numbers, "the file holds the unscoped total"


# --------------------------------------------------------------------- 3/4. brands API

from fastapi import Depends  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from sqlalchemy import text  # noqa: E402

from app.database import get_db  # noqa: E402
from app.dependencies import get_current_user  # noqa: E402
from app.main import app  # noqa: E402
from app.models.base import set_company_scope  # noqa: E402
from app.models.product import Brand  # noqa: E402
from app.services.company_scope import DEFAULT_COMPANY_ID  # noqa: E402
from app.services.company_scope_resolver import apply_company_scope  # noqa: E402

from tests._brand_scope_seed import brand, contact  # noqa: E402
from tests._mc_lookup_seed import MOCHA_ID, seed_mocha  # noqa: E402
from tests._pg_fixture import unique_code  # noqa: E402

CONTACTS = f"{BASE}/user-management/contacts"


@pytest.fixture
def sorento_caller(db, monkeypatch):
    """A caller holding every user-management slug, whose company scope is Sorento ONLY."""
    from app.services.user_service import UserPermissionService

    principal = {"id": str(uuid.uuid4()), "email": "cbs-r2@zzt.test"}

    def _override_db():
        yield db

    def _scope(_db=Depends(get_db)):
        scope = frozenset({DEFAULT_COMPANY_ID})
        set_company_scope(_db, scope)
        return scope

    app.dependency_overrides[get_db] = _override_db
    app.dependency_overrides[get_current_user] = lambda: principal
    app.dependency_overrides[apply_company_scope] = _scope
    monkeypatch.setattr(UserPermissionService, "check_user_has_permission", lambda self, uid, slug: True)
    try:
        yield TestClient(app)
    finally:
        app.dependency_overrides.clear()


def _foreign_brand(db) -> Brand:
    seed_mocha(db)
    row = Brand(
        id=str(uuid.uuid4()), brand_code=unique_code("FOR")[:50], brand_name="ZZT OTHER COMPANY BRAND",
        company_id=MOCHA_ID,
    )
    db.add(row)
    db.flush()
    return row


def test_put_rejects_a_brand_of_a_company_outside_the_callers_scope(db, sorento_caller) -> None:
    foreign = _foreign_brand(db)
    c = contact(db)
    db.commit()
    resp = sorento_caller.put(f"{CONTACTS}/{c.id}/brands", json={"brand_ids": [foreign.id]})
    assert resp.status_code == 422, resp.text
    stored = db.execute(text("SELECT brand_ids FROM respond_contacts WHERE id = :i"), {"i": c.id}).scalar()
    assert stored is None


def test_put_still_accepts_a_brand_of_the_callers_own_company(db, sorento_caller) -> None:
    own = brand(db, "MOCHA")
    c = contact(db)
    db.commit()
    resp = sorento_caller.put(f"{CONTACTS}/{c.id}/brands", json={"brand_ids": [own.id]})
    assert resp.status_code == 200, resp.text


def test_get_never_returns_another_companys_brand_name(db, sorento_caller) -> None:
    foreign = _foreign_brand(db)
    c = contact(db)
    db.execute(text("UPDATE respond_contacts SET brand_ids = ARRAY[CAST(:b AS uuid)] WHERE id = :i"),
               {"b": foreign.id, "i": c.id})
    db.commit()
    resp = sorento_caller.get(f"{CONTACTS}/{c.id}/brands")
    assert resp.status_code == 200, resp.text
    assert "OTHER COMPANY BRAND" not in resp.text


@pytest.fixture
def audited(db):
    from app.services.audit_service import register_audit_listeners

    register_audit_listeners()
    yield


def test_put_brands_writes_an_audit_row_with_before_and_after_ids(db, sorento_caller, audited) -> None:
    from app.models.audit import AuditLog

    m, s = brand(db, "MOCHA"), brand(db, "SORENTO")
    c = contact(db, brand_ids=[m.id])
    db.commit()
    resp = sorento_caller.put(f"{CONTACTS}/{c.id}/brands", json={"brand_ids": [s.id]})
    assert resp.status_code == 200, resp.text
    rows = db.query(AuditLog).filter(AuditLog.entity_id == str(c.id), AuditLog.action == "UPDATE").all()
    assert rows, "no audit row for the brands change"
    changed = [r for r in rows if "brand_ids" in (r.new_values or {})]
    assert changed, [r.new_values for r in rows]
    row = changed[-1]
    assert [str(x) for x in row.new_values["brand_ids"]] == [str(s.id)]
    assert [str(x) for x in row.old_values["brand_ids"]] == [str(m.id)]
