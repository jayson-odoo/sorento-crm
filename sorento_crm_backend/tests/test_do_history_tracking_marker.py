"""DO-HISTORY slice 1: the `orders.tracking_assumed_at` marker (owner Q2 a, 13 Oct 2026).

Plan: documentation/plans/autocount/PLAN-do-history-backfill-13oct.md section 4.2.
UAC:  documentation/plans/autocount/do-history-backfill-13oct-acceptance-criteria.md
(AC-DH15, AC-DH16).

The marker says "delivered by assumption, not by a tracking upload". It is bookkeeping no
business writer owns: only the DO ingest's assumed-delivery path writes it (never through the
AutoCount header write), and a real Order Tracking row for the DO clears it so it never claims
an assumption the sheet has since replaced.
"""
from __future__ import annotations

from datetime import date, datetime, timezone

from app.models.order import Order, OrderStatus
from app.services.order_field_ownership import (
    AUTOCOUNT_OWNED_ORDER_COLUMNS,
    ORDER_TRACKING_OWNED_ORDER_COLUMNS,
    SYSTEM_ORDER_COLUMNS,
    assert_autocount_writes,
)
from app.services.order_service import OrderService
from tests.test_order_field_ownership import (  # noqa: F401 - db is a fixture
    MASTER_HEADER,
    _USER,
    _autocount_order,
    _master_workbook,
    db,
)

import pytest


def test_marker_is_a_system_column_on_the_model():
    """AC-DH16: one owner, the bookkeeping set."""
    assert "tracking_assumed_at" in {c.name for c in Order.__table__.columns}
    assert "tracking_assumed_at" in SYSTEM_ORDER_COLUMNS
    assert "tracking_assumed_at" not in AUTOCOUNT_OWNED_ORDER_COLUMNS
    assert "tracking_assumed_at" not in ORDER_TRACKING_OWNED_ORDER_COLUMNS


def test_ingest_header_may_not_carry_the_marker():
    """AC-DH16: the DO ingest's own write guard refuses it."""
    with pytest.raises(RuntimeError, match="tracking_assumed_at"):
        assert_autocount_writes({"order_number", "tracking_assumed_at"})


def test_tracking_upload_clears_the_marker(db):
    """AC-DH15: an Overall Tracking row for an assumed DO replaces the assumption."""
    order = _autocount_order(db)
    order.tracking_assumed_at = datetime(2026, 10, 13, tzinfo=timezone.utc)
    db.commit()
    data = _master_workbook(MASTER_HEADER, [
        ["ZZDO-OG-1", date(2026, 9, 1), "01/09/2026 10:00", "999", "Sheet", "S-AG", "N", "cs",
         "RMA"],
    ])
    result = OrderService(db).import_excel_tracking(data, _USER)
    assert not result.get("errors"), result.get("errors")
    db.refresh(order)
    assert order.actual_delivery_date is not None
    assert order.tracking_assumed_at is None


def test_tracking_upload_leaves_other_rows_marker_alone(db):
    """AC-DH15: only the DO the sheet names loses its marker."""
    named = _autocount_order(db, "ZZDO-OG-1")
    other = Order(order_number="ZZDO-OG-9", doc_key=990109, source_book="db1",
                  order_date=date(2026, 9, 27), debtor_name="AutoCount Name")
    db.add(other)
    stamp = datetime(2026, 10, 13, tzinfo=timezone.utc)
    other.tracking_assumed_at = stamp
    db.commit()
    data = _master_workbook(MASTER_HEADER, [
        ["ZZDO-OG-1", date(2026, 9, 1), "01/09/2026 10:00", "999", "Sheet", "S-AG", "N", "cs",
         "RMA"],
    ])
    result = OrderService(db).import_excel_tracking(data, _USER)
    assert not result.get("errors"), result.get("errors")
    db.refresh(other)
    assert other.tracking_assumed_at == stamp
    assert named.id != other.id


def test_master_reupload_keeps_an_assumed_delivery(db):
    """AC-DH36 (review blocker 1): a Master row with no Tracking row for an assumed DO must not
    reset it to new; the marker counts as a delivery there, like a date."""
    delivered = db.query(OrderStatus).filter(OrderStatus.status_code == "DELIVERED").one()
    order = _autocount_order(db)
    order.order_status_id = delivered.id
    order.tracking_assumed_at = datetime(2026, 10, 13, tzinfo=timezone.utc)
    db.commit()
    data = _master_only_workbook(["ZZDO-OG-1", date(2026, 9, 1), "01/09/2026 10:00", "999",
                                  "Sheet", "S-AG", "N", "cs", "RMA"])
    result = OrderService(db).import_excel_tracking(data, _USER)
    assert not result.get("errors"), result.get("errors")
    db.refresh(order)
    assert str(order.order_status_id) == str(delivered.id)
    assert order.tracking_assumed_at is not None
    assert order.actual_delivery_date is None


def _master_only_workbook(row: list) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    master = wb.create_sheet("Master")
    master.append(MASTER_HEADER)
    master.append(row)
    tracking = wb.create_sheet("Overall Tracking")
    tracking.append(["Doc Number", "Date", "Time", "Checker", "Transporter", "Driver Name",
                     "Lorry Plate", "Customer", "Remarks CS", "Remarks", "Salesman", "Trips",
                     "W/H"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
