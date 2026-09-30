"""RED tests for lane DO-PULL-CRM: the `delivery_orders` pull entity (BL-SS-286).

Plan: documentation/plans/autocount/PLAN-autocount-do-pull-crm-30sep.md
UAC:  documentation/plans/autocount/autocount-do-pull-crm-30sep-acceptance-criteria.md

A third entity on the existing pull machinery. Preview and apply both run
`AutocountDocIngestService` (the DO ingest, contract 2.7) - one writer - so every test here
seeds the masters that ingest resolves (the DO ingest test's own recipe) and drives the
SAME preview / apply tasks and routes SR1 / SR3 / SR4 already pin for products and stock.

Substrate reused BY IMPORT from `tests/test_autocount_pull_sr1.py` (`env`, `task_db`,
`_FakeFoundryX`, `_patch_foundryx`, `_header`, `_stored`, `_seed_pull_job`, `_job_row`,
`_job_rows`, `_job_count`, `_run_preview`) and `tests/test_autocount_pull_sr3.py`
(`_run_apply`, `_set_rows_page`). Importing a `@pytest.fixture` function's NAME is enough
for pytest to discover it here (the SR3/SR4 technique).

Snapshot rows are the raw vendor DO dict verbatim plus `source_ref` `{book}:DO:{DocKey}`
(DO-PULL-SS contract, relayed 30 Sep) - built from the committed DO ingest fixture
`tests/fixtures/autocount/do_live_sample.json`, so the pull and the push are proven against
the same documents. The book is read off the rows' `source_ref` prefix.

Every name this lane adds (`AUTOCOUNT_RETRYABLE`, the DO maps, `map_delivery_order_rows`,
`compare_delivery_orders`, the DO task branches, the migration) is imported INSIDE a test
body, never at module import, so a missing piece reds one test and never collection.
"""
from __future__ import annotations

import copy
import importlib.util
import io
import json
import sys
import uuid
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import MagicMock

import openpyxl
import pytest
import sqlalchemy as sa
from alembic.migration import MigrationContext
from alembic.operations import Operations
from sqlalchemy import text

# MUST be the first app import - resolves the circular import in
# app.modules.runtime.guards (repo convention, see test_ingest_deletions.py).
from app.main import app  # noqa: E402,F401

from app.services.company_scope import DEFAULT_COMPANY_ID

from tests._pg_fixture import unique_code
from tests.test_autocount_pull_sr1 import (  # noqa: F401 - env/task_db are fixtures
    PULLS_URL,
    _FakeFoundryX,
    _header,
    _job_count,
    _job_row,
    _job_rows,
    _patch_foundryx,
    _run_preview,
    _seed_pull_job,
    _stored,
    env,
    task_db,
)
from tests.test_autocount_pull_sr3 import _run_apply, _set_rows_page  # noqa: F401

MARKER = "ZZTDOP"
DO_FIXTURE = Path(__file__).resolve().parent / "fixtures" / "autocount" / "do_live_sample.json"
VERSIONS = (Path(__file__).resolve().parent / ".." / "alembic" / "versions").resolve()

ENTITY = "delivery_orders"
SLUG = "order_management.orders.autocount_pull"
PULL_JOB_TYPE = "autocount_delivery_orders_pull"
APPLY_JOB_TYPE = "autocount_delivery_orders_apply"
BOOK = "db1"
MIGRATION = "do_pull_0001_perm"


# ================================================================== rows and headers


def _do_rows(book: str | None = BOOK) -> list[dict]:
    """The two live-shape DO records, each carrying the contract's `source_ref`."""
    records = json.loads(DO_FIXTURE.read_text())["records"]
    rows = []
    for rec in records:
        row = copy.deepcopy(rec)
        if book is not None:
            row["source_ref"] = f"{book}:DO:{rec['DocKey']}"
        rows.append(row)
    return rows


def _do_header(rows: list[dict], **overrides) -> dict:
    """A ready DELIVERY ORDERS header valid for `rows` (SR1's `_header`, re-labelled). No
    `book` key: the book travels on the rows (contract), and the one test that wants it
    on the header sets it explicitly."""
    header = _header(rows, **overrides)
    header["entity"] = ENTITY
    for key in ("zeroListPriceCount", "negativeListPriceCount", "enrichMissCount"):
        header.pop(key, None)
    return header


def _lines(rows: list[dict]) -> list[tuple[dict, dict]]:
    return [(rec, line) for rec in rows for line in rec.get("Details") or [] if line.get("ItemCode")]


# ================================================================== masters seeding


def _seed_masters(db) -> dict:
    """Every master the DO sample names, in the Sorento company - CI's database holds no
    data (the DO ingest test's own `_Env.__init__` recipe)."""
    from sqlalchemy import select

    from app.models.inventory import Warehouse
    from app.models.order import Customer, OrderStatus
    from app.models.product import Product, ProductCategory, UnitOfMeasure

    category = ProductCategory(category_code=unique_code(MARKER), category_name="c")
    uom = UnitOfMeasure(uom_code=unique_code(MARKER), uom_name="u")
    db.add_all([category, uom])
    db.flush()
    ids = {}
    for code in ("ZZAC-P1", "ZZAC-P2"):
        row = Product(product_code=code, product_name=code, category_id=category.id,
                      base_uom_id=uom.id, list_price=10, company_id=DEFAULT_COMPANY_ID)
        db.add(row)
        db.flush()
        ids[code] = str(row.id)
    wh = Warehouse(warehouse_code="ZZAC-WH1", warehouse_name="wh", company_id=DEFAULT_COMPANY_ID)
    db.add(wh)
    db.flush()
    ids["wh"] = str(wh.id)
    cust = Customer(customer_code="300-ZZAC01", customer_name="ZZAC Customer One",
                    company_id=DEFAULT_COMPANY_ID)
    db.add(cust)
    db.flush()
    ids["customer"] = str(cust.id)
    existing = db.execute(
        select(OrderStatus).where(OrderStatus.status_code.in_(["new", "NEW"]))
    ).scalars().first()
    if existing is None:
        existing = OrderStatus(status_code="new", status_name="New Order")
        db.add(existing)
        db.flush()
    ids["new_status"] = str(existing.id)
    db.commit()
    return ids


TRACKING_COLUMNS = (
    "actual_delivery_date", "pickup_time", "transporter", "driver_name", "lorry_plate",
    "checker", "trips", "delivery_days", "kpi_warning", "customer_ref", "salesman",
    "warehouse", "delivery_remarks", "delivery_remarks_cs", "remarks_cs", "order_status_id",
)


def _seed_tracking_row(db, ids: dict):
    """A DO the tracking upload created (no `doc_key`), numbered like the sample's first
    document, with two lines: P1 x 10 matches the incoming line and keeps its id, P2 x 99
    matches nothing and is the one line adoption deletes."""
    from app.models.order import Order, OrderLine

    row = Order(
        order_number="ZZDO-0001", company_id=DEFAULT_COMPANY_ID, order_date=date(2026, 9, 26),
        actual_delivery_date=date(2026, 9, 28), pickup_time="14:30", transporter="ZZ TRANS",
        driver_name="Ali", lorry_plate="WXX 1234", checker="Chong", trips=2, delivery_days=1,
        kpi_warning=False, customer_ref="iPad ref", salesman="SEAN", warehouse="BRW",
        delivery_remarks="dr", delivery_remarks_cs="drcs", remarks_cs="rcs",
        order_status_id=ids["new_status"], debtor_code="OLD", debtor_name="Old name",
    )
    db.add(row)
    db.flush()
    db.add(OrderLine(order_id=row.id, line_sequence=1, product_id=ids["ZZAC-P1"],
                     warehouse_id=ids["wh"], quantity=Decimal("10")))
    db.add(OrderLine(order_id=row.id, line_sequence=2, product_id=ids["ZZAC-P2"],
                     warehouse_id=ids["wh"], quantity=Decimal("99")))
    db.commit()
    return str(row.id)


def _tracking_snapshot(db, order_id: str) -> dict:
    row = db.execute(
        text(f"SELECT {', '.join(TRACKING_COLUMNS)}, doc_key FROM orders WHERE id = :id"),
        {"id": order_id},
    ).mappings().first()
    return dict(row)


def _orders(db, *doc_keys: int) -> list[dict]:
    rows = db.execute(
        text(
            "SELECT id, order_number, doc_key, source_book, source_record, updated_at, "
            "transporter, driver_name FROM orders WHERE company_id = :cid "
            "AND doc_key = ANY(:keys) ORDER BY doc_key"
        ),
        {"cid": DEFAULT_COMPANY_ID, "keys": list(doc_keys)},
    ).mappings().all()
    return [dict(r) for r in rows]


def _order_line_count(db, order_id) -> int:
    return db.execute(
        text("SELECT count(*) FROM order_lines WHERE order_id = :id"), {"id": str(order_id)}
    ).scalar()


# ================================================================== job seeding


def _prepare_do_preview(db, fake: _FakeFoundryX, *, rows: list[dict], header: dict | None = None,
                        snapshot_id: str | None = None, owner: str | None = None):
    """Seeds a `previewing` DO pull and points the fake at its snapshot."""
    snapshot_id = snapshot_id or f"{MARKER}-snap-{uuid.uuid4().hex[:8]}"
    valid = _do_header(rows)
    fake.status = (200, {**(header or valid), "snapshotId": snapshot_id})
    _set_rows_page(fake, snapshot_id, rows)
    return _seed_pull_job(
        db, job_type=PULL_JOB_TYPE, user_id=owner or str(uuid.uuid4()),
        company_id=DEFAULT_COMPANY_ID, entity=ENTITY, company_code="SRT",
        snapshot_id=snapshot_id, phase="previewing", header=_stored(header or valid),
    )


def _seed_do_apply_job(db, *, user_id: str, snapshot_id: str, pull_job_id: str | None = None):
    from app.models.job import ImportJob, JobStatus

    job = ImportJob(
        id=uuid.uuid4(), job_id=str(uuid.uuid4()), job_type=APPLY_JOB_TYPE,
        status=JobStatus.QUEUED.value, user_id=user_id, company_id=DEFAULT_COMPANY_ID,
        job_metadata={"autocount_apply": {
            "pull_job_id": pull_job_id or str(uuid.uuid4()),
            "snapshot_id": snapshot_id, "entity": ENTITY,
        }},
    )
    db.add(job)
    db.commit()
    return job.id


def _prepare_do_apply(db, fake: _FakeFoundryX, *, rows: list[dict], user_id: str | None = None,
                      snapshot_id: str | None = None):
    snapshot_id = snapshot_id or f"{MARKER}-apply-{uuid.uuid4().hex[:8]}"
    fake.status = (200, {**_do_header(rows), "snapshotId": snapshot_id})
    _set_rows_page(fake, snapshot_id, rows)
    return _seed_do_apply_job(db, user_id=user_id or str(uuid.uuid4()), snapshot_id=snapshot_id)


def _seed_do_review_job(env, *, owner, rows=None, phase="review"):
    rows = rows if rows is not None else _do_rows()
    snapshot_id = f"{MARKER}-rv-{uuid.uuid4().hex[:8]}"
    job_id = _seed_pull_job(
        env.db, job_type=PULL_JOB_TYPE, user_id=owner["id"], company_id=env.company_a,
        entity=ENTITY, company_code=env.company_a_code, snapshot_id=snapshot_id, phase=phase,
        header=_stored(_do_header(rows)),
    )
    _set_rows_page(env.fake, snapshot_id, rows)
    return job_id, rows


def _counts(db, job_id) -> dict:
    return _job_row(db, job_id)["metadata"]["autocount_pull"]["counts"]


def _progress(db, job_id) -> tuple[int | None, int | None]:
    row = db.execute(
        text("SELECT processed_rows, total_rows FROM import_jobs WHERE id = :id"),
        {"id": str(job_id)},
    ).first()
    return (row[0], row[1]) if row else (None, None)


# ======================================================================= AC-DP-01..04


class TestEntityMaps:
    def test_dp_01_maps_and_registry_slug(self):
        from app.rbac.permission_registry import PERMISSION_REGISTRY
        from app.services import autocount_pull_service as pull_service

        assert pull_service.ENTITY_PERMISSIONS[ENTITY] == SLUG
        assert pull_service.JOB_TYPES[ENTITY] == PULL_JOB_TYPE
        assert pull_service.APPLY_JOB_TYPES[ENTITY] == APPLY_JOB_TYPE
        assert SLUG in {p["slug"] for p in PERMISSION_REGISTRY}

    def test_dp_02_start_requires_slug_and_builds_do_entity(self, env):
        nobody = env.user("master_data.products.autocount_pull")
        env.as_user(nobody)
        before = _job_count(env.db)
        resp = env.post_pull(ENTITY)
        assert resp.status_code == 403, resp.text
        assert env.fake.calls == []
        assert _job_count(env.db) == before

        owner = env.user(SLUG)
        env.as_user(owner)
        resp = env.post_pull(ENTITY)
        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["entity"] == ENTITY and body["phase"] == "building"
        assert _job_row(env.db, body["job_id"])["job_type"] == PULL_JOB_TYPE
        build = [c for c in env.fake.calls if c["method"] == "POST"]
        assert len(build) == 1
        assert build[0]["json"] == {"companyCode": env.company_a_code, "entity": ENTITY}

    def test_dp_02b_scope_travels_flat_and_is_stored(self, env):
        owner = env.user(SLUG)
        env.as_user(owner)
        scope = {"fromDay": "2026-09-01", "toDay": "2026-09-30"}
        resp = env.client.post(PULLS_URL, json={"entity": ENTITY, "scope": scope})
        assert resp.status_code == 200, resp.text
        assert resp.json()["scope"] == scope
        build = [c for c in env.fake.calls if c["method"] == "POST"][0]
        assert build["json"] == {"companyCode": env.company_a_code, "entity": ENTITY, **scope}
        stored = _job_row(env.db, resp.json()["job_id"])["metadata"]["autocount_pull"]
        assert stored["scope"] == scope
        # The same open pull is reused on a second click (AC-PL-4, entity-agnostic).
        again = env.client.post(PULLS_URL, json={"entity": ENTITY, "scope": {"docNo": "X"}})
        assert again.status_code == 200
        assert again.json()["job_id"] == resp.json()["job_id"]

    def test_dp_02c_build_in_flight_409_leaves_no_job(self, env):
        owner = env.user(SLUG)
        env.as_user(owner)
        env.fake.build = (
            409,
            {"code": "BUILD_IN_FLIGHT", "message": "A build with a different scope is in flight."},
        )
        before = _job_count(env.db)
        resp = env.client.post(PULLS_URL, json={"entity": ENTITY, "scope": {"docNo": "ZZDO-0001"}})
        assert resp.status_code == 409, resp.text
        assert resp.json().get("code") == "BUILD_IN_FLIGHT"
        assert _job_count(env.db) == before

    def test_dp_03_wrong_entity_permission_403_and_current_finds_do_pull(self, env):
        owner = env.user(SLUG, "master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_do_review_job(env, owner=owner)
        current = env.get_current(ENTITY)
        assert current.status_code == 200, current.text
        assert current.json()["job_id"] == str(job_id)
        assert current.json()["entity"] == ENTITY

        products_only = env.user("master_data.products.autocount_pull")
        env.as_user(products_only)
        own_do, _ = _seed_do_review_job(env, owner=products_only)
        assert env.get_pull(own_do).status_code == 403
        assert env.client.get(f"{PULLS_URL}/{own_do}/rows").status_code == 403


class TestMigration:
    def _load(self):
        alembic_dir = str(VERSIONS.parent)
        if alembic_dir not in sys.path:
            sys.path.insert(0, alembic_dir)
        spec = importlib.util.spec_from_file_location(f"m_{MIGRATION}", VERSIONS / f"{MIGRATION}.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        return module

    def test_dp_04a_revision_fits_and_sits_on_the_single_head(self):
        """Not a literal `down_revision`: `scripts/alembic-reparent.sh` rewrites it every time
        main's head moves before merge (review blocker 1), so the pin is the property that
        matters - the chain has ONE head and this revision is on its ancestry. Not "this
        revision IS the head": the join migration main adds after a fork (merge_30sep_batch9
        after #1383 and #1386) sits on top of it and would turn a by-name pin red."""
        from alembic.config import Config
        from alembic.script import ScriptDirectory

        module = self._load()
        assert module.revision == MIGRATION
        assert len(module.revision) <= 32
        script = ScriptDirectory.from_config(Config(str(VERSIONS.parent.parent / "alembic.ini")))
        heads = list(script.get_heads())
        assert len(heads) == 1, heads
        assert MIGRATION in {r.revision for r in script.walk_revisions(base="base", head=heads[0])}

    def test_dp_04b_grants_to_import_holders_and_admin_not_integrations(self):
        """Runs `upgrade()` on the real tables inside a rolled-back transaction (the
        `test_migration_*` pattern): a role holding `order_management.orders.import`
        gains the slug, an `integration_*` role holding it does not, `admin` gains it."""
        from app.database import engine

        module = self._load()
        stem = uuid.uuid4().hex[:8]
        with engine.connect() as conn:
            outer = conn.begin()
            try:
                def q(sql, **params):
                    return conn.execute(text(sql), params)

                q("INSERT INTO user_permissions (id, slug, name, description, created_at) "
                  "SELECT gen_random_uuid()::text, 'order_management.orders.import', 'i', '', now() "
                  "WHERE NOT EXISTS (SELECT 1 FROM user_permissions WHERE slug = 'order_management.orders.import')")
                import_perm = q("SELECT id FROM user_permissions WHERE slug = 'order_management.orders.import'").scalar()
                roles = {
                    "human": f"{MARKER.lower()}-human-{stem}",
                    "integration": f"integration_{MARKER.lower()}_{stem}",
                    "bystander": f"{MARKER.lower()}-bystander-{stem}",
                }
                role_ids = {}
                for key, slug in roles.items():
                    role_ids[key] = q(
                        "INSERT INTO user_roles (id, slug, name, description, is_trashed, "
                        "is_protected, is_default) "
                        "VALUES (gen_random_uuid()::text, :slug, :slug, '', false, false, false) "
                        "RETURNING id",
                        slug=slug,
                    ).scalar()
                for key in ("human", "integration"):
                    q("INSERT INTO user_role_permissions (id, role_id, permission_id, assigned_at) "
                      "VALUES (gen_random_uuid()::text, :r, :p, now()) ON CONFLICT DO NOTHING",
                      r=role_ids[key], p=import_perm)
                q("INSERT INTO user_roles (id, slug, name, description, is_trashed, is_protected, "
                  "is_default) "
                  "SELECT gen_random_uuid()::text, 'admin', 'admin', '', false, true, false "
                  "WHERE NOT EXISTS (SELECT 1 FROM user_roles WHERE slug = 'admin')")

                with Operations.context(MigrationContext.configure(conn)):
                    module.upgrade()

                holders = {
                    r[0] for r in q(
                        "SELECT r.slug FROM user_role_permissions rp "
                        "JOIN user_permissions p ON p.id = rp.permission_id "
                        "JOIN user_roles r ON r.id = rp.role_id WHERE p.slug = :slug", slug=SLUG,
                    )
                }
                assert roles["human"] in holders
                assert "admin" in holders
                assert roles["integration"] not in holders
                assert roles["bystander"] not in holders

                with Operations.context(MigrationContext.configure(conn)):
                    module.downgrade()
                assert q("SELECT count(*) FROM user_permissions WHERE slug = :slug", slug=SLUG).scalar() == 0
            finally:
                outer.rollback()


# ======================================================================= AC-DP-10..14


class TestPreview:
    def test_dp_10_sample_all_created_and_nothing_written(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        _seed_masters(db)
        rows = _do_rows()
        job_id = _prepare_do_preview(db, fake, rows=rows)

        _run_preview(monkeypatch, factory, job_id)

        row = _job_row(db, job_id)
        assert row["status"] == "finished", row["error"]
        pull = row["metadata"]["autocount_pull"]
        assert pull["phase"] == "review"
        assert pull["confirm_blocked_reason"] is None
        assert pull["counts"] == {
            "received": 2, "created": 2, "updated": 0, "adopted": 0, "unchanged": 0,
            "lines_to_delete": 0, "failed": 0, "retryable": 0, "with_warnings": 0,
        }
        written = _job_rows(db, job_id)
        assert sorted(r["value"] for r in written) == ["ZZDO-0001", "ZZDO-0002"]
        assert {r["outcome"] for r in written} == {"created"}
        first = next(r for r in written if r["value"] == "ZZDO-0001")
        assert first["identity"]["doc_key"] == 900001
        assert first["identity"]["lines_created"] == 2
        # Dry run: the ingest rolled itself back, nothing landed.
        assert _orders(db, 900001, 900002) == []

    def test_dp_10b_unresolved_sales_order_is_a_warning_not_a_blocker(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        _seed_masters(db)
        rows = _do_rows()
        rows[0]["RefDocNo"] = f"{MARKER}-SO-NOWHERE"
        job_id = _prepare_do_preview(db, fake, rows=rows)

        _run_preview(monkeypatch, factory, job_id)

        pull = _job_row(db, job_id)["metadata"]["autocount_pull"]
        assert pull["phase"] == "review"
        assert pull["confirm_blocked_reason"] is None
        assert pull["counts"]["created"] == 2
        assert pull["counts"]["with_warnings"] == 1
        first = next(r for r in _job_rows(db, job_id) if r["value"] == "ZZDO-0001")
        assert first["outcome"] == "created"
        # Words, never the slug (review S4): the identity is printed as is on the rows card.
        assert first["identity"]["warnings"] == "sales order not found"
        assert "sales_order_unresolved" not in str(first["identity"])
        assert "sales order not found" in (first["message"] or "")

    def test_dp_11_adopts_tracking_row_by_number_and_writes_nothing(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        ids = _seed_masters(db)
        order_id = _seed_tracking_row(db, ids)
        before = _tracking_snapshot(db, order_id)
        job_id = _prepare_do_preview(db, fake, rows=_do_rows())

        _run_preview(monkeypatch, factory, job_id)

        row = _job_row(db, job_id)
        assert row["status"] == "finished", row["error"]
        counts = row["metadata"]["autocount_pull"]["counts"]
        assert counts["created"] == 1
        assert counts["adopted"] == 1
        assert counts["updated"] == 0
        assert counts["lines_to_delete"] == 1
        adopted = next(r for r in _job_rows(db, job_id) if r["value"] == "ZZDO-0001")
        assert adopted["outcome"] == "updated"
        assert "adopted" in (adopted["message"] or "").lower()
        # Adoption is the message itself, never a warning (review N5): no "(warnings" suffix,
        # no slug anywhere, and it does not count toward `with_warnings`.
        assert "(warnings" not in (adopted["message"] or "")
        assert "adopted_by_doc_no" not in str(adopted)
        assert counts["with_warnings"] == 0
        # Flat keys, non-zero only (review S4): `_json_safe_identity` prints identity in the
        # UI and stringifies any nested value, so the per-line counters are their own keys.
        assert {k: v for k, v in adopted["identity"].items() if k.startswith("lines_")} == {
            "lines_created": 1, "lines_deleted": 1, "lines_adopted": 1, "lines_unlinked": 2,
        }
        assert "source_ref" not in adopted["identity"]
        # Nothing landed: the tracking row still has no doc_key, its two old lines stand.
        assert _tracking_snapshot(db, order_id) == before
        assert before["doc_key"] is None
        assert _order_line_count(db, order_id) == 2

    def test_dp_12_doc_key_clash_failed_and_unknown_product_retryable(self, task_db, monkeypatch):
        from app.models.order import Order
        from app.services import import_outcome_codes as codes

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        ids = _seed_masters(db)
        # ZZDO-0001 is already another AutoCount document's number.
        db.add(Order(order_number="ZZDO-0001", company_id=DEFAULT_COMPANY_ID,
                     order_date=date(2026, 9, 1), source_book=BOOK, doc_key=777,
                     order_status_id=ids["new_status"]))
        db.commit()
        rows = _do_rows()
        rows[1]["Details"][0]["ItemCode"] = "ZZAC-NOPE"
        job_id = _prepare_do_preview(db, fake, rows=rows)

        _run_preview(monkeypatch, factory, job_id)

        row = _job_row(db, job_id)
        assert row["status"] == "finished", row["error"]
        counts = row["metadata"]["autocount_pull"]["counts"]
        assert counts["failed"] == 1 and counts["retryable"] == 1 and counts["created"] == 0
        written = {r["value"]: r for r in _job_rows(db, job_id)}
        clash = written["ZZDO-0001"]
        assert clash["outcome"] == "failed" and clash["code"] == "DocNo"
        assert "777" in (clash["message"] or "")
        retry = written["ZZDO-0002"]
        assert retry["outcome"] == "failed" and retry["code"] == codes.AUTOCOUNT_RETRYABLE
        assert "ZZAC-NOPE" in (retry["message"] or "")
        assert _orders(db, 900001, 900002) == []

    def test_dp_13_no_book_fails_and_two_books_fail(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        _seed_masters(db)

        no_ref = _do_rows(book=None)
        job_id = _prepare_do_preview(db, fake, rows=no_ref)
        _run_preview(monkeypatch, factory, job_id)
        row = _job_row(db, job_id)
        assert row["status"] == "failed"
        assert "names no book" in (row["error"] or "").lower()
        assert _job_rows(db, job_id) == []

        two = _do_rows()
        two[1]["source_ref"] = f"db2:DO:{two[1]['DocKey']}"
        job_id = _prepare_do_preview(db, fake, rows=two)
        _run_preview(monkeypatch, factory, job_id)
        row = _job_row(db, job_id)
        assert row["status"] == "failed"
        assert "book" in (row["error"] or "").lower()
        assert _orders(db, 900001, 900002) == []

    def test_dp_13b_header_book_is_accepted_when_rows_carry_none(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        _seed_masters(db)
        rows = _do_rows(book=None)
        header = _do_header(rows)
        header["book"] = BOOK
        job_id = _prepare_do_preview(db, fake, rows=rows, header=header)

        _run_preview(monkeypatch, factory, job_id)

        row = _job_row(db, job_id)
        assert row["status"] == "finished", row["error"]
        assert row["metadata"]["autocount_pull"]["counts"]["created"] == 2

    def test_dp_14_preview_progress_reaches_total(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        _seed_masters(db)
        job_id = _prepare_do_preview(db, fake, rows=_do_rows())

        _run_preview(monkeypatch, factory, job_id)

        assert _progress(db, job_id) == (2, 2)


# ======================================================================= AC-DP-20..23


class TestConfirmAndApply:
    def test_dp_20a_confirm_creates_do_apply_job(self, env, monkeypatch):
        owner = env.user(SLUG)
        env.as_user(owner)
        job_id, _ = _seed_do_review_job(env, owner=owner)
        captured = []
        monkeypatch.setattr(
            "app.services.autocount_pull_service.enqueue_job",
            lambda func, *a, **k: captured.append(func) or MagicMock(id=str(uuid.uuid4())),
        )

        resp = env.client.post(f"{PULLS_URL}/{job_id}/confirm")

        assert resp.status_code == 200, resp.text
        assert resp.json()["phase"] == "confirmed"
        apply_row = _job_row(env.db, resp.json()["apply_job_id"])
        assert apply_row["job_type"] == APPLY_JOB_TYPE
        assert apply_row["metadata"]["autocount_apply"]["entity"] == ENTITY
        assert len(captured) == 1 and captured[0].__name__ == "apply_autocount_pull"

    def test_dp_20b_apply_writes_through_the_ingest(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        _seed_masters(db)
        rows = _do_rows()
        job_id = _prepare_do_apply(db, fake, rows=rows)

        _run_apply(monkeypatch, factory, job_id)

        row = _job_row(db, job_id)
        assert row["status"] == "finished", row["error"]
        assert row["metadata"]["autocount_apply"]["counts"] == {
            "total": 2, "created": 2, "updated": 0, "adopted": 0, "unchanged": 0,
            "failed": 0, "retryable": 0, "lines_deleted": 0, "with_warnings": 0,
        }
        orders = _orders(db, 900001, 900002)
        assert [o["order_number"] for o in orders] == ["ZZDO-0001", "ZZDO-0002"]
        assert {o["source_book"] for o in orders} == {BOOK}
        assert orders[0]["source_record"]["DocKey"] == 900001
        assert _order_line_count(db, orders[0]["id"]) == 2
        assert _order_line_count(db, orders[1]["id"]) == 1
        written = _job_rows(db, job_id)
        assert sorted(r["value"] for r in written) == ["ZZDO-0001", "ZZDO-0002"]
        assert {r["outcome"] for r in written} == {"created"}

    def test_dp_21_apply_adopts_and_keeps_tracking_columns(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        ids = _seed_masters(db)
        order_id = _seed_tracking_row(db, ids)
        before = _tracking_snapshot(db, order_id)
        job_id = _prepare_do_apply(db, fake, rows=_do_rows())

        _run_apply(monkeypatch, factory, job_id)

        row = _job_row(db, job_id)
        assert row["status"] == "finished", row["error"]
        counts = row["metadata"]["autocount_apply"]["counts"]
        assert counts["adopted"] == 1 and counts["created"] == 1 and counts["lines_deleted"] == 1
        after = _tracking_snapshot(db, order_id)
        assert after["doc_key"] == 900001
        after.pop("doc_key")
        before.pop("doc_key")
        assert after == before
        adopted = _orders(db, 900001)[0]
        assert str(adopted["id"]) == order_id
        assert _order_line_count(db, order_id) == 2

    def test_dp_22_second_apply_of_the_same_snapshot_is_unchanged(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        _seed_masters(db)
        rows = _do_rows()
        first = _prepare_do_apply(db, fake, rows=rows)
        _run_apply(monkeypatch, factory, first)
        assert _job_row(db, first)["status"] == "finished"
        stamps = {o["order_number"]: o["updated_at"] for o in _orders(db, 900001, 900002)}

        second = _prepare_do_apply(db, fake, rows=rows)
        _run_apply(monkeypatch, factory, second)

        row = _job_row(db, second)
        assert row["status"] == "finished", row["error"]
        counts = row["metadata"]["autocount_apply"]["counts"]
        assert counts["unchanged"] == 2 and counts["created"] == 0 and counts["updated"] == 0
        assert {o["order_number"]: o["updated_at"] for o in _orders(db, 900001, 900002)} == stamps
        assert _job_rows(db, second) == []

    def test_dp_23_expired_snapshot_fails_and_writes_nothing(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        _seed_masters(db)
        snapshot_id = f"{MARKER}-expired"
        fake.status = (410, {"code": "SNAPSHOT_EXPIRED", "message": "expired"})
        job_id = _seed_do_apply_job(db, user_id=str(uuid.uuid4()), snapshot_id=snapshot_id)

        _run_apply(monkeypatch, factory, job_id)

        row = _job_row(db, job_id)
        assert row["status"] == "failed"
        assert "pull again" in (row["error"] or "").lower()
        assert _orders(db, 900001, 900002) == []


# ======================================================================= AC-DP-30..32

ROW_KEYS = [
    "doc_no", "doc_date", "debtor_code", "debtor_name", "item_code", "description",
    "location", "qty", "uom", "unit_price", "sub_total",
]
XLSX_HEADER = [
    "Doc No", "Doc Date", "Debtor Code", "Debtor Name", "Item Code", "Description",
    "Location", "Qty", "UOM", "Unit Price", "Sub Total",
]


class TestRowsDownloadCompare:
    def test_dp_30_rows_one_per_line_in_the_import_sheet_shape(self, env):
        owner = env.user(SLUG)
        env.as_user(owner)
        job_id, rows = _seed_do_review_job(env, owner=owner)

        resp = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"limit": 50})

        assert resp.status_code == 200, resp.text
        body = resp.json()
        expected = _lines(rows)
        assert body["pagination"]["total"] == len(expected) == 3
        for got, (rec, line) in zip(body["data"], expected):
            assert list(got.keys()) == ROW_KEYS, got
            assert got["doc_no"] == rec["DocNo"]
            assert got["doc_date"] == "2026-09-27"
            assert got["debtor_code"] == rec["DebtorCode"]
            assert got["debtor_name"] == rec["DebtorName"]
            assert got["item_code"] == line["ItemCode"]
            assert got["location"] == line["Location"]
            assert Decimal(str(got["qty"])) == Decimal(str(line["Qty"]))
            assert got["uom"] == line["UOM"]
            assert Decimal(str(got["unit_price"])) == Decimal(str(line["UnitPrice"]))
            assert Decimal(str(got["sub_total"])) == Decimal(str(line["SubTotal"]))

        by_doc = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"query": "zzdo-0002"})
        assert [r["doc_no"] for r in by_doc.json()["data"]] == ["ZZDO-0002"]
        by_item = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"query": "ZZAC-P1"})
        assert [r["item_code"] for r in by_item.json()["data"]] == ["ZZAC-P1", "ZZAC-P1"]

    def test_dp_30b_description_only_row_is_not_a_line(self, env):
        owner = env.user(SLUG)
        env.as_user(owner)
        rows = _do_rows()
        rows[0]["Details"].append({"DtlKey": 910099, "Seq": 64, "ItemCode": "",
                                   "Description": "Handle with care", "Qty": 0})
        job_id, _ = _seed_do_review_job(env, owner=owner, rows=rows)

        resp = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"limit": 50})
        assert resp.status_code == 200, resp.text
        assert resp.json()["pagination"]["total"] == 3

    def test_dp_31_download_xlsx(self, env):
        owner = env.user(SLUG)
        env.as_user(owner)
        job_id, rows = _seed_do_review_job(env, owner=owner)

        resp = env.client.get(f"{PULLS_URL}/{job_id}/download.xlsx")

        assert resp.status_code == 200, resp.text
        assert "autocount-delivery_orders-pull.xlsx" in resp.headers.get("Content-Disposition", "")
        sheet = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        data = list(sheet.iter_rows(values_only=True))
        assert list(data[0]) == XLSX_HEADER
        assert len(data) == 1 + len(_lines(rows))
        assert data[1][0] == "ZZDO-0001" and data[1][4] == "ZZAC-P1"
        assert Decimal(str(data[1][7])) == Decimal("10")

    def test_dp_32_compare_keyed_by_doc_item_location(self, env):
        owner = env.user(SLUG)
        env.as_user(owner)
        job_id, rows = _seed_do_review_job(env, owner=owner)
        excel = [
            {"Doc No": rec["DocNo"], "Item Code": line["ItemCode"], "Location": line["Location"],
             "Qty": line["Qty"]}
            for rec, line in _lines(rows)
        ]
        excel[0]["Qty"] = 11  # ZZDO-0001 / ZZAC-P1 / ZZAC-WH1: 10 in the pull
        excel.pop(2)  # ZZDO-0002 / ZZAC-P1 only in the pull
        excel.append({"Doc No": "zzdo-0009 ", "Item Code": "ZZAC-P2", "Location": "ZZAC-WH1", "Qty": 1})

        resp = env.client.post(
            f"{PULLS_URL}/{job_id}/compare", json={"filename": "do-lines.xlsx", "rows": excel},
        )

        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["summary"]["filename"] == "do-lines.xlsx"
        assert body["summary"]["total"] == 2
        assert body["summary"]["matched"] == 1
        assert body["summary"]["different"] == 1
        assert body["summary"]["only_in_excel"] == 1
        assert body["summary"]["only_in_pull"] == 1
        assert body["differences"] == [{
            "item_code": "ZZAC-P1", "doc_no": "ZZDO-0001", "location": "ZZAC-WH1",
            "field": "qty", "excel": 11, "pull": 10,
        }]
        assert body["only_in_excel"] == ["zzdo-0009|ZZAC-P2|ZZAC-WH1"]
        assert body["only_in_pull"] == ["ZZDO-0002|ZZAC-P1|ZZAC-WH1"]
        stored = _job_row(env.db, job_id)["metadata"]["autocount_pull"]["compare"]
        assert stored["matched"] == 1 and stored["different"] == 1
        assert "differences" not in stored


def test_compare_delivery_orders_is_pure_and_case_insensitive():
    from app.services.autocount_pull_compare import compare_delivery_orders

    pull_rows = _do_rows()
    excel = [{"doc number": "zzdo-0001", "product code": " zzac-p1 ", "warehouse": "zzac-wh1",
              "quantity": "10.0"}]
    # AC-CMM-7: no alias guessing any more - the sheet's columns are named by a mapping; the
    # header match is still trimmed and case-insensitive ("Doc Number" matches "doc number").
    mapping = {"sheet_name": "Master", "columns": [
        {"excel_header": "Doc Number", "transform": "text", "field": "doc_no"},
        {"excel_header": " PRODUCT CODE", "transform": "text", "field": "item_code"},
        {"excel_header": "Warehouse", "transform": "text", "field": "location"},
        {"excel_header": "Quantity", "transform": "number", "field": "qty"},
    ]}
    result = compare_delivery_orders(excel, pull_rows, mapping)
    assert result["summary"] == {"total": 1, "matched": 1, "different": 0}
    assert result["differences"] == []
    assert result["only_in_excel"] == []
    assert sorted(result["only_in_pull"]) == ["ZZDO-0001|ZZAC-P2|ZZAC-WH1", "ZZDO-0002|ZZAC-P1|ZZAC-WH1"]


# ================================================== security review fix round (30 Sep)


class TestSecurityFixRound:
    """Security-reviewer findings on 69c1f263: B1 (unbounded quantity), S1 (module gate
    per entity), S2 (CSV formula guard), S3 (book validation), N1 (scope values), N2 (odd
    vendor cells), N3 (document cap), N5 (test gaps)."""

    def test_b1_compare_quantities_are_bounded_and_finite(self):
        import json
        import time

        from app.services.autocount_pull_compare import compare_delivery_orders

        pull_rows = _do_rows()
        hostile = [
            {"Doc No": "ZZDO-0001", "Item Code": "ZZAC-P1", "Location": "ZZAC-WH1", "Qty": q}
            for q in ("1e3000000", "NaN", "sNaN", "Infinity", "-Infinity", "1e-3000000")
        ]
        started = time.monotonic()
        result = compare_delivery_orders(hostile, pull_rows)
        assert time.monotonic() - started < 2.0
        json.dumps(result)  # serialisable: no NaN, no Infinity
        # Every hostile value reads as 0 (the "not a quantity" rule), against the pull's 10.
        assert result["differences"][-1]["excel"] == 0
        assert result["differences"][-1]["pull"] == 10

    def test_b1_compare_route_answers_200_on_hostile_quantity(self, env):
        owner = env.user(SLUG)
        env.as_user(owner)
        job_id, _ = _seed_do_review_job(env, owner=owner)
        resp = env.client.post(
            f"{PULLS_URL}/{job_id}/compare",
            json={"filename": "x.xlsx", "rows": [
                {"Doc No": "ZZDO-0001", "Item Code": "ZZAC-P1", "Location": "ZZAC-WH1",
                 "Qty": "1e3000000"},
            ]},
        )
        assert resp.status_code == 200, resp.text
        assert resp.json()["summary"]["different"] == 1

    def test_s1_disabled_order_module_blocks_the_do_pull_under_strict_guard(self, env, monkeypatch):
        from app.config import settings
        from app.models.app_modules import AppModuleCatalog, TenantModule
        from app.modules.runtime.installer import DEFAULT_TENANT_ID

        monkeypatch.setattr(settings, "module_guard_strict", True)
        for key, enabled in (("product", True), ("inventory", True), ("order", False)):
            if not env.db.query(AppModuleCatalog).filter_by(module_key=key).first():
                env.db.add(AppModuleCatalog(module_key=key, display_name=key, dependencies=[]))
                env.db.flush()
            env.db.add(TenantModule(tenant_id=DEFAULT_TENANT_ID, module_key=key, enabled=enabled))
        env.db.commit()

        owner = env.user(SLUG)
        env.as_user(owner)
        before = _job_count(env.db)
        resp = env.post_pull(ENTITY)
        assert resp.status_code == 403, resp.text
        assert resp.json().get("code") == "MODULE_DISABLED"
        assert env.fake.calls == [] and _job_count(env.db) == before

        row = env.db.query(TenantModule).filter_by(tenant_id=DEFAULT_TENANT_ID, module_key="order").one()
        row.enabled = True
        env.db.commit()
        assert env.post_pull(ENTITY).status_code == 200

    def test_s2_csv_export_guards_formula_prefixes(self):
        from app.api.v1.system.jobs import _csv_safe

        assert _csv_safe("=cmd|' /C calc'!A0") == "'=cmd|' /C calc'!A0"
        assert _csv_safe("+1+cmd") == "'+1+cmd"
        assert _csv_safe("@SUM(1)") == "'@SUM(1)"
        assert _csv_safe("\tx") == "'\tx"
        # Any leading dash (OWASP; review N4): `-1+cmd|' /C calc'!A0` is a formula too.
        assert _csv_safe("-1+cmd|' /C calc'!A0") == "'-1+cmd|' /C calc'!A0"
        assert _csv_safe("-5") == "'-5"
        assert _csv_safe("-cmd") == "'-cmd"
        assert _csv_safe("DO-2609/0077") == "DO-2609/0077"
        assert _csv_safe(None) == ""

    def test_s3_book_must_match_the_ingest_pattern_and_the_rows(self, task_db, monkeypatch):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        _seed_masters(db)

        rows = _do_rows()
        header = _do_header(rows)
        header["book"] = "db2"  # rows say db1
        job_id = _prepare_do_preview(db, fake, rows=rows, header=header)
        _run_preview(monkeypatch, factory, job_id)
        row = _job_row(db, job_id)
        assert row["status"] == "failed"
        assert "more than one book" in (row["error"] or "")

        bad = _do_rows(book="a-book-name-far-longer-than-twenty")
        job_id = _prepare_do_preview(db, fake, rows=bad)
        _run_preview(monkeypatch, factory, job_id)
        row = _job_row(db, job_id)
        assert row["status"] == "failed"
        assert "cannot store" in (row["error"] or "")
        assert _orders(db, 900001, 900002) == []

    def test_n1_scope_values_are_validated(self, env):
        owner = env.user(SLUG, "master_data.products.autocount_pull")
        env.as_user(owner)
        bad_day = env.client.post(PULLS_URL, json={"entity": ENTITY, "scope": {"fromDay": "01/09/2026"}})
        assert bad_day.status_code == 422, bad_day.text
        long_doc = env.client.post(PULLS_URL, json={"entity": ENTITY, "scope": {"docNo": "X" * 51}})
        assert long_doc.status_code == 422, long_doc.text
        on_products = env.client.post(
            PULLS_URL, json={"entity": "products", "scope": {"fromDay": "2026-09-01"}}
        )
        assert on_products.status_code == 422, on_products.text
        assert env.fake.calls == []

    def test_n2_rows_and_download_tolerate_odd_vendor_cells(self, env):
        owner = env.user(SLUG)
        env.as_user(owner)
        rows = _do_rows()
        rows[0]["DocNo"] = 12345  # not a string
        rows[0]["Details"][0]["Qty"] = "nan"
        rows[0]["Details"][0]["Description"] = {"nested": "object"}
        rows[0]["Details"][1]["UnitPrice"] = "not a number"
        job_id, _ = _seed_do_review_job(env, owner=owner, rows=rows)

        listed = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"query": "12345"})
        assert listed.status_code == 200, listed.text
        data = listed.json()["data"]
        assert [r["doc_no"] for r in data] == ["12345", "12345"]
        assert data[0]["qty"] is None and data[0]["description"] is None
        assert data[1]["unit_price"] is None
        assert env.client.get(f"{PULLS_URL}/{job_id}/download.xlsx").status_code == 200

    def test_n3_snapshot_past_the_document_cap_is_refused(self, task_db, monkeypatch):
        from app.services.autocount_pull_service import MAX_DELIVERY_ORDER_DOCS

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        rows = [
            {"source_ref": f"{BOOK}:DO:{i}", "DocKey": i, "DocNo": f"ZZCAP-{i}",
             "DocDate": "2026-09-27", "Details": []}
            for i in range(1, MAX_DELIVERY_ORDER_DOCS + 2)
        ]
        job_id = _prepare_do_preview(db, fake, rows=rows)
        _run_preview(monkeypatch, factory, job_id)
        row = _job_row(db, job_id)
        assert row["status"] == "failed"
        assert "narrower window" in (row["error"] or "")
        assert _job_rows(db, job_id) == []

    def test_n5_company_b_row_with_the_same_number_is_never_adopted(self, task_db, monkeypatch):
        from app.models.company import Company
        from app.models.order import Order

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        ids = _seed_masters(db)
        other = Company(id=str(uuid.uuid4()), name=f"{MARKER} B", code=f"ZB{uuid.uuid4().hex[:6]}")
        db.add(other)
        db.flush()
        foreign = Order(order_number="ZZDO-0001", company_id=str(other.id),
                        order_date=date(2026, 9, 26), transporter="B TRANS",
                        order_status_id=ids["new_status"])
        db.add(foreign)
        db.commit()
        foreign_id = str(foreign.id)
        job_id = _prepare_do_apply(db, fake, rows=_do_rows())

        _run_apply(monkeypatch, factory, job_id)

        row = _job_row(db, job_id)
        assert row["status"] == "finished", row["error"]
        assert row["metadata"]["autocount_apply"]["counts"]["created"] == 2
        assert row["metadata"]["autocount_apply"]["counts"]["adopted"] == 0
        untouched = db.execute(
            text("SELECT doc_key, transporter, company_id FROM orders WHERE id = :id"),
            {"id": foreign_id},
        ).mappings().one()
        assert untouched["doc_key"] is None and untouched["transporter"] == "B TRANS"
        assert str(untouched["company_id"]) == str(other.id)

    def test_n5_xlsx_neutralises_a_formula_shaped_document_number(self, env):
        owner = env.user(SLUG)
        env.as_user(owner)
        rows = _do_rows()
        rows[0]["DocNo"] = "=cmd|' /C calc'!A0"
        job_id, _ = _seed_do_review_job(env, owner=owner, rows=rows)

        resp = env.client.get(f"{PULLS_URL}/{job_id}/download.xlsx")
        assert resp.status_code == 200, resp.text
        sheet = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        cell = sheet["A2"]
        assert cell.value == "=cmd|' /C calc'!A0"
        assert cell.data_type == "s"


# ============================================== correctness review fix round (30 Sep)


class TestCorrectnessFixRound:
    """Reviewer findings on 69c1f263 / 40884ec3: S1 (a push-written DO must preview as
    unchanged), S2 (duplicate lines summed), S3 (`with_warnings` counts shown rows only)."""

    def test_s1_a_document_the_feed_pushed_previews_as_unchanged(self, task_db, monkeypatch):
        """The feed's push carries the vendor record WITHOUT `source_ref`; the pull's rows
        carry it. The pull must strip it before the ingest, or the stored `source_record`
        differs and every document reads `updated` with nothing to show."""
        from app.services.autocount_doc_ingest_service import AutocountDocIngestService

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        _seed_masters(db)
        pushed = AutocountDocIngestService(db, None, company_id=DEFAULT_COMPANY_ID, book=BOOK).ingest(
            ENTITY, _do_rows(book=None)
        )
        db.commit()
        assert {r.outcome.value for r in pushed.records} == {"created"}
        stamps = {o["order_number"]: o["updated_at"] for o in _orders(db, 900001, 900002)}

        job_id = _prepare_do_preview(db, fake, rows=_do_rows())
        _run_preview(monkeypatch, factory, job_id)
        counts = _counts(db, job_id)
        assert counts["unchanged"] == 2 and counts["updated"] == 0, counts
        assert _job_rows(db, job_id) == []

        apply_id = _prepare_do_apply(db, fake, rows=_do_rows())
        _run_apply(monkeypatch, factory, apply_id)
        assert _job_row(db, apply_id)["metadata"]["autocount_apply"]["counts"]["unchanged"] == 2
        assert {o["order_number"]: o["updated_at"] for o in _orders(db, 900001, 900002)} == stamps
        assert all("source_ref" not in o["source_record"] for o in _orders(db, 900001, 900002))

    def test_s2_compare_sums_the_same_item_listed_twice_on_one_document(self):
        from app.services.autocount_pull_compare import compare_delivery_orders

        pull_rows = _do_rows()
        twice = copy.deepcopy(pull_rows[0]["Details"][0])
        twice.update({"DtlKey": 910077, "Seq": 48, "Qty": 3.0})
        pull_rows[0]["Details"].append(twice)  # ZZDO-0001 / ZZAC-P1 / ZZAC-WH1: 10 + 3

        split = [
            {"Doc No": "ZZDO-0001", "Item Code": "ZZAC-P1", "Location": "ZZAC-WH1", "Qty": 5},
            {"Doc No": "ZZDO-0001", "Item Code": "ZZAC-P1", "Location": "ZZAC-WH1", "Qty": 8},
        ]
        result = compare_delivery_orders(split, pull_rows)
        assert result["summary"]["matched"] == 1 and result["differences"] == []

        one_row = [{"Doc No": "ZZDO-0001", "Item Code": "ZZAC-P1", "Location": "ZZAC-WH1", "Qty": 12}]
        result = compare_delivery_orders(one_row, pull_rows)
        assert result["differences"] == [{
            "item_code": "ZZAC-P1", "doc_no": "ZZDO-0001", "location": "ZZAC-WH1",
            "field": "qty", "excel": 12, "pull": 13,
        }]

    def test_s3_with_warnings_counts_shown_rows_only(self, task_db, monkeypatch):
        """An unchanged document writes no row, so it must not count toward the tile; a
        second pull of a window whose DOs all carry an unresolved sales order reads
        `with_warnings 0`, not the whole window."""
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake, db)
        _seed_masters(db)
        rows = _do_rows()
        for row in rows:
            row["RefDocNo"] = f"{MARKER}-SO-NOWHERE"
        first = _prepare_do_apply(db, fake, rows=rows)
        _run_apply(monkeypatch, factory, first)
        assert _job_row(db, first)["metadata"]["autocount_apply"]["counts"]["with_warnings"] == 2

        job_id = _prepare_do_preview(db, fake, rows=rows)
        _run_preview(monkeypatch, factory, job_id)
        counts = _counts(db, job_id)
        assert counts["unchanged"] == 2
        assert counts["with_warnings"] == 0
        assert _job_rows(db, job_id) == []


# ========================================= two-file compare + Q4 switch (owner, 30 Sep)


def _lines_sheet(rows: list[dict], doc_date="27/09/2026") -> list[dict]:
    """The Order Listing macro's `Master` sheet, one row per DO line, as `sheet_to_json`
    hands it over (a date as text here; the serial-number form is pinned separately)."""
    out = []
    for rec, line in _lines(rows):
        out.append({
            "Doc No": rec["DocNo"], "Doc Date": doc_date, "Created Time": "09:12",
            "Cancelled": "F", "Item Code": line["ItemCode"], "Qty": line["Qty"],
            "Location": line["Location"], "Unit Price": line["UnitPrice"], "Discount": "",
            "Total (Ex)": line["SubTotal"], "Total (Inc)": line["SubTotal"],
        })
    return out


def _headers_sheet(rows: list[dict], date="27/09/2026") -> list[dict]:
    """The Order Tracking macro's `Master` sheet, one row per DO."""
    return [
        {"Doc. No.": rec["DocNo"], "Date": date, "Created Time": "09:12",
         "Debtor Code": rec["DebtorCode"], "Debtor Name": rec["DebtorName"],
         "Agent": rec["SalesAgent"], "Cancel": "F", "Remarks CS": "", "Type": "Normal"}
        for rec in rows
    ]


class TestTwoFileCompare:
    def test_excel_day_reads_serials_and_text(self):
        from datetime import date as _date

        from app.services.autocount_pull_compare import _excel_day

        assert _excel_day(46292) == _date(2026, 9, 27)  # sheet_to_json's serial number
        assert _excel_day(46292.0) == _date(2026, 9, 27)
        assert _excel_day("27/09/2026") == _date(2026, 9, 27)
        assert _excel_day("2026-09-27T00:00:00") == _date(2026, 9, 27)
        assert _excel_day("20260927") == _date(2026, 9, 27)
        assert _excel_day("") is None and _excel_day(None) is None
        assert _excel_day("not a date") is None and _excel_day(5) is None

    def test_lines_source_is_windowed_and_compares_the_q3_fields(self, env):
        owner = env.user(SLUG)
        env.as_user(owner)
        job_id, rows = _seed_do_review_job(env, owner=owner)
        env.db.execute(
            text("UPDATE import_jobs SET metadata = jsonb_set(metadata, '{autocount_pull,scope}', "
                 "'{\"fromDay\": \"2026-09-01\", \"toDay\": \"2026-09-30\"}') WHERE id = :id"),
            {"id": str(job_id)},
        )
        env.db.commit()
        sheet = _lines_sheet(rows)
        sheet[0]["Unit Price"] = 13  # 12.5 in the pull
        sheet[0]["Discount"] = "5%"  # blank in the pull
        sheet[0]["Total (Ex)"] = 130  # 125 in the pull
        sheet.append({**sheet[1], "Doc No": "ZZDO-AUG", "Doc Date": "15/08/2026"})  # outside
        sheet.append({**sheet[1], "Doc No": "ZZDO-OCT", "Doc Date": 46296})  # 1 Oct, outside

        resp = env.client.post(
            f"{PULLS_URL}/{job_id}/compare",
            json={"filename": "Order Listing 2026-09.xlsm", "rows": sheet, "source": "lines"},
        )

        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["source"] == "lines"
        assert body["window"] == {"fromDay": "2026-09-01", "toDay": "2026-09-30"}
        assert body["ignored_outside_window"] == 2 and body["rows_in_window"] == 3
        assert body["only_in_excel"] == [] and body["only_in_pull"] == []
        fields = {(d["doc_no"], d["field"]): (d["excel"], d["pull"]) for d in body["differences"]}
        assert fields == {
            ("ZZDO-0001", "unit_price"): (13, 12.5),
            ("ZZDO-0001", "discount"): ("5", None),
            ("ZZDO-0001", "total_ex"): (130, 125),
        }
        assert body["source_summary"]["matched"] == 2 and body["source_summary"]["different"] == 1
        # The stored combined summary is the lines file alone until the headers file lands.
        assert body["summary"]["filename"] == "Order Listing 2026-09.xlsm"
        assert body["summary"]["total"] == 3

    def test_headers_source_compares_date_debtor_and_cancel(self, env):
        owner = env.user(SLUG)
        env.as_user(owner)
        job_id, rows = _seed_do_review_job(env, owner=owner)
        sheet = _headers_sheet(rows)
        sheet[0]["Cancel"] = "Y"
        sheet[1]["Debtor Code"] = "300-OTHER"
        sheet[1]["Date"] = "28/09/2026"
        sheet.append({**sheet[0], "Doc. No.": "ZZDO-0142", "Cancel": "F"})  # not in the pull

        resp = env.client.post(
            f"{PULLS_URL}/{job_id}/compare",
            json={"filename": "Order Tracking 2026-09.xlsm", "rows": sheet, "source": "headers"},
        )

        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["source"] == "headers"
        fields = {(d["doc_no"], d["field"]): (d["excel"], d["pull"]) for d in body["differences"]}
        assert fields == {
            ("ZZDO-0001", "cancel"): (True, False),
            ("ZZDO-0002", "doc_date"): ("2026-09-28", "2026-09-27"),
            ("ZZDO-0002", "debtor_code"): ("300-OTHER", "300-ZZAC01"),
        }
        assert all(d["item_code"] == "" for d in body["differences"])
        assert body["only_in_excel"] == ["ZZDO-0142"]
        assert body["only_in_pull"] == []
        assert body["source_summary"]["different"] == 2

    def test_both_sources_add_up_in_the_stored_summary(self, env):
        owner = env.user(SLUG)
        env.as_user(owner)
        job_id, rows = _seed_do_review_job(env, owner=owner)
        env.client.post(f"{PULLS_URL}/{job_id}/compare",
                        json={"filename": "lines.xlsm", "rows": _lines_sheet(rows), "source": "lines"})
        resp = env.client.post(f"{PULLS_URL}/{job_id}/compare",
                               json={"filename": "tracking.xlsm", "rows": _headers_sheet(rows), "source": "headers"})
        assert resp.status_code == 200, resp.text
        summary = resp.json()["summary"]
        assert summary["filename"] == "lines.xlsm and tracking.xlsm"
        assert summary["total"] == 5 and summary["matched"] == 5 and summary["different"] == 0
        pull = env.client.get(f"{PULLS_URL}/{job_id}").json()
        assert set(pull["compare_sources"]) == {"lines", "headers"}
        assert pull["compare_sources"]["headers"]["filename"] == "tracking.xlsm"
        assert pull["window"] == {"fromDay": None, "toDay": None} or set(pull["window"]) == {"fromDay", "toDay"}

    def test_source_is_refused_on_products(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        from tests.test_autocount_pull_sr3 import _seed_review_job

        job_id, _ = _seed_review_job(env, owner=owner)
        resp = env.client.post(f"{PULLS_URL}/{job_id}/compare",
                               json={"filename": "x.xlsx", "rows": [], "source": "lines"})
        assert resp.status_code == 422

    def test_default_window_is_the_31_days_ending_on_the_snapshot_day(self):
        from app.services.autocount_pull_service import pull_window

        pull = {"entity": ENTITY, "scope": None, "header": {"extractedAt": "2026-09-30T01:12:00Z"}}
        assert pull_window(pull) == ("2026-08-31", "2026-09-30")
        pull["header"]["fromDay"], pull["header"]["toDay"] = "2026-09-01", "2026-09-30"
        assert pull_window(pull) == ("2026-09-01", "2026-09-30")
        pull["scope"] = {"docNo": "ZZDO-0001"}
        assert pull_window(pull) == (None, None)
        pull["scope"] = {"fromDay": "2026-09-10", "toDay": "2026-09-20"}
        assert pull_window(pull) == ("2026-09-10", "2026-09-20")


class TestConfirmRequiresMatchSwitch:
    """Owner Q4: one cheap switch, default advisory. On: Confirm is held until BOTH files
    compare clean; off: the compare never gates Confirm (the Products / Stock rule)."""

    def test_default_is_advisory(self, env, monkeypatch):
        owner = env.user(SLUG)
        env.as_user(owner)
        job_id, rows = _seed_do_review_job(env, owner=owner)
        monkeypatch.setattr("app.services.autocount_pull_service.enqueue_job",
                            lambda *a, **k: MagicMock(id=str(uuid.uuid4())))
        assert env.get_pull(job_id).json()["confirm_blocked_reason"] is None
        assert env.get_pull(job_id).json()["confirm_requires_match"] is False
        assert env.client.post(f"{PULLS_URL}/{job_id}/confirm").status_code == 200

    def test_switch_on_holds_confirm_until_both_files_compare_clean(self, env, monkeypatch):
        from app.config import settings

        monkeypatch.setattr(settings, "autocount_do_pull_confirm_requires_match", True)
        owner = env.user(SLUG)
        env.as_user(owner)
        job_id, rows = _seed_do_review_job(env, owner=owner)
        monkeypatch.setattr("app.services.autocount_pull_service.enqueue_job",
                            lambda *a, **k: MagicMock(id=str(uuid.uuid4())))

        before = env.get_pull(job_id).json()
        assert before["confirm_requires_match"] is True
        assert "both Excel files" in before["confirm_blocked_reason"]
        refused = env.client.post(f"{PULLS_URL}/{job_id}/confirm")
        assert refused.status_code == 409 and refused.json()["code"] == "NOT_READY"

        lines = _lines_sheet(rows)
        lines[0]["Qty"] = 99
        env.client.post(f"{PULLS_URL}/{job_id}/compare",
                        json={"filename": "lines.xlsm", "rows": lines, "source": "lines"})
        env.client.post(f"{PULLS_URL}/{job_id}/compare",
                        json={"filename": "tracking.xlsm", "rows": _headers_sheet(rows), "source": "headers"})
        assert env.client.post(f"{PULLS_URL}/{job_id}/confirm").status_code == 409

        clean = env.client.post(f"{PULLS_URL}/{job_id}/compare",
                                json={"filename": "lines.xlsm", "rows": _lines_sheet(rows), "source": "lines"})
        assert clean.json()["confirm_blocked_reason"] is None
        assert env.get_pull(job_id).json()["confirm_blocked_reason"] is None
        assert env.client.post(f"{PULLS_URL}/{job_id}/confirm").status_code == 200

    def test_switch_on_never_touches_a_products_pull(self, env, monkeypatch):
        from app.config import settings
        from tests.test_autocount_pull_sr3 import _seed_review_job

        monkeypatch.setattr(settings, "autocount_do_pull_confirm_requires_match", True)
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_review_job(env, owner=owner)
        assert env.get_pull(job_id).json()["confirm_blocked_reason"] is None
