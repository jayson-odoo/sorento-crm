"""Dealer targets import, third real-file gate (lane SALES-TARGET-DEALER, crew 16 Oct).

Round 2 on the real files: Sorento 41/109 mapped, Mocha 1/135. Shapes found, masked:
1. Mocha: "QUARTERLY (POINT)" merged over JAN-MAR..OCT-DEC only; the INCENTIVE column right after
   the block sits outside the merge.
2. Several CRM customers per name + account: a bracketed word other than the account (a branch)
   tells them apart.
3. Per-account tiers beside dealer tiers in a QUARTERLY block; account targets with no window of
   their own beside a Dec-Nov dealer; a rate naming an account with no row in the file; one
   threshold repeated across tiers that name different accounts.
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from .test_sales_targets_s1 import BASE, _code, _uid, api, world  # noqa: F401  (fixtures)

D = Decimal


def _wb(fill):
    from openpyxl import Workbook

    wb = Workbook()
    out = fill(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), out


def _parse(fill):
    from app.services.sales.dealer_target_import import parse_workbook

    data, _ = _wb(fill)
    return parse_workbook(data, year=2026)


def _upload(client, db, data, *, dry_run):
    files = {"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    db.autoflush = False
    try:
        return client.post(f"{BASE}/import", params={"dry_run": str(dry_run).lower(), "year": 2026}, files=files).json()
    finally:
        db.autoflush = True


def _customer(db, company_id, name, level=None):
    from app.models.order import Customer

    row = Customer(id=_uid(), company_id=company_id, customer_code=_code(), customer_name=name, account_level=level)
    db.add(row)
    db.flush()
    return row


# 1. the INCENTIVE right after a block whose merge stops short of it --------------------------


def _mocha(ws):
    ws["A2"] = "CUSTOMER NAME"
    ws.merge_cells("A2:A3")
    ws["B2"] = "QUARTERLY (POINT)"
    ws.merge_cells("B2:E2")
    for col, text in zip("BCDEF", ("JAN-MAR", "APR-JUN", "JUL-SEP", "OCT-DEC", "INCENTIVE")):
        ws[f"{col}3"] = text
    ws["A4"] = "KILO LIVING A/C I"
    for i, col in enumerate("BCDE"):
        ws[f"{col}4"] = 100000 + 10000 * i
    ws["F4"] = 0.04


def test_an_incentive_right_after_a_quarterly_block_pays_every_quarter():
    parsed = _parse(_mocha)
    assert parsed.rows[0].status == "mapped", parsed.rows[0].reason
    (q,) = [s for s in parsed.schemes if s.kind == "quarterly"]
    assert q.threshold_unit == "points"
    assert [(t.period_start, t.threshold, t.rate) for t in q.tiers] == [
        (date(2026, 1, 1), D("100000"), D("4")), (date(2026, 4, 1), D("110000"), D("4")),
        (date(2026, 7, 1), D("120000"), D("4")), (date(2026, 10, 1), D("130000"), D("4")),
    ]


# 2. several customers per name and account -------------------------------------------------


def _one_row(name):
    def fill(ws):
        ws["A1"], ws["B1"] = "DEALER NAME", "Q1"
        ws["A2"], ws["B2"] = name, "100,000 - 1%"

    return fill


def test_a_bracketed_branch_word_breaks_a_tie(api):
    client, db, company_id = api
    jb = _customer(db, company_id, "ALPHA HARDWARE (JB) SDN BHD [A/C I]", 1)
    kl = _customer(db, company_id, "ALPHA HARDWARE (KL) SDN BHD [A/C I]", 1)
    data, _ = _wb(_one_row("Alpha Hardware (KL) Sdn Bhd A/C I"))
    row = _upload(client, db, data, dry_run=True)["rows"][0]
    assert row["match"] == "auto" and row["customer"] == kl.customer_name
    data, _ = _wb(_one_row("ALPHA HARDWARE SDN BHD A/C I"))
    row = _upload(client, db, data, dry_run=True)["rows"][0]
    assert row["status"] == "needs_pick"
    assert {s["name"] for s in row["suggestions"]} == {jb.customer_name, kl.customer_name}


def test_the_masked_dry_run_says_what_tells_several_candidates_apart():
    from scripts.dealer_targets_dry_run import _differ

    assert _differ(["ALPHA (JB) SDN BHD [A/C I]", "ALPHA (KL) SDN BHD [A/C I]"]) == "differ by: xx | xx"
    assert _differ(["ALPHA SDN BHD [A/C I]", "ALPHA SDN BHD [A/C I]"]) == "identical names"
    assert _differ(["ALPHA SDN BHD [A/C I]", "ALPHA SDN. BHD. - [A/C I]"]) == "differ by: punctuation or spacing only"


# 3. per-account tiers beside dealer tiers in a quarterly block ----------------------------


def _quarterly_split(ws):
    ws["A1"], ws["B1"], ws["D1"] = "DEALER NAME", "Q1", "Q2"
    ws.merge_cells("B1:C1")
    ws.merge_cells("D1:E1")
    for col, text in zip("BCDE", ("Tier 1", "Tier 1 Incentive", "Tier 1", "Tier 1 Incentive")):
        ws[f"{col}2"] = text
    ws["A3"], ws["A4"] = "X TRADING A/C I", "X TRADING A/C III"
    ws["B3"], ws["C3"] = 500000, 0.02
    ws.merge_cells("B3:B4")
    ws.merge_cells("C3:C4")
    ws["D3"], ws["E3"] = 200000, 0.01
    ws["D4"], ws["E4"] = 150000, 0.01


def test_per_account_quarter_tiers_beside_dealer_quarter_tiers_split():
    parsed = _parse(_quarterly_split)
    assert all(r.status == "mapped" for r in parsed.rows), [(r.row, r.reason) for r in parsed.rows]
    keys = [r.key for r in parsed.rows]
    dealer = next(s for s in parsed.schemes if s.codes == keys)
    assert [(t.period_start, t.threshold) for t in dealer.tiers] == [(date(2026, 1, 1), D("500000"))]
    owns = [s for s in parsed.schemes if len(s.codes) == 1]
    assert sorted((s.codes[0], [(t.period_start, t.threshold) for t in s.tiers]) for s in owns) == [
        (keys[0], [(date(2026, 4, 1), D("200000"))]),
        (keys[1], [(date(2026, 4, 1), D("150000"))]),
    ]
    assert all(s.replaced_by == dealer.key for s in owns)


def _yearly_dec_nov(ws):
    ws["A1"], ws["B1"], ws["F1"] = "DEALER NAME", "Yearly", "Yearly DEC'25-NOV'26"
    ws.merge_cells("B1:E1")
    ws.merge_cells("F1:G1")
    for col, text in zip("BCDEFG", ("TIER 1", "INCENTIVE", "TIER 2", "INCENTIVE", "TIER 4", "INCENTIVE")):
        ws[f"{col}2"] = text
    ws["A3"], ws["A4"] = "X TRADING A/C I", "X TRADING A/C III"
    for r, base in ((3, 100000), (4, 120000)):
        ws[f"B{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"E{r}"], ws[f"G{r}"] = base, 0.01, base * 2, 0.015, 0.02
    ws["F3"] = 2000000
    ws.merge_cells("F3:F4")


def test_account_targets_take_the_dealer_window_when_they_state_none():
    parsed = _parse(_yearly_dec_nov)
    assert all(r.status == "mapped" for r in parsed.rows), [(r.row, r.reason) for r in parsed.rows]
    assert {(s.start, s.end) for s in parsed.schemes} == {(date(2025, 12, 1), date(2026, 11, 30))}


def _pays_an_account_with_no_row(ws):
    ws["A1"], ws["B1"] = "DEALER NAME", "Q1"
    ws.merge_cells("B1:C1")
    ws["B2"], ws["C2"] = "Tier 1", "Tier 1 Incentive"
    ws["A3"], ws["B3"], ws["C3"] = "DELTA BUILDERS SDN BHD A/C III", 80000, "A/C I- 2%"


def test_a_rate_naming_an_account_with_no_row_pays_it_by_name(api):
    client, db, company_id = api
    a1 = _customer(db, company_id, "DELTA BUILDERS SDN BHD - [A/C I]", 1)
    a3 = _customer(db, company_id, "DELTA BUILDERS SDN BHD - [A/C III]", 3)
    data, _ = _wb(_pays_an_account_with_no_row)
    body = _upload(client, db, data, dry_run=False)
    row = body["rows"][0]
    assert row["status"] == "mapped", row
    rows = client.get(BASE, params={"subject": "dealer", "all": "true"}).json()["rows"]
    target = next(r for r in rows if r["target_no"] in row["target_nos"])
    detail = client.get(f"{BASE}/{target['target_id']}").json()
    assert [a["code"] for a in detail["accounts"]] == [a3.customer_code]
    assert detail["payout_account"]["code"] == a1.customer_code


def test_a_rate_naming_an_account_with_no_row_and_no_customer_fails(api):
    client, db, company_id = api
    _customer(db, company_id, "DELTA BUILDERS SDN BHD - [A/C III]", 3)
    data, _ = _wb(_pays_an_account_with_no_row)
    row = _upload(client, db, data, dry_run=True)["rows"][0]
    assert row["status"] == "not_representable"
    assert "A/C I" in row["reason"]


def _repeated_threshold_per_account(ws):
    ws["A1"], ws["B1"] = "DEALER NAME", "Q1"
    ws.merge_cells("B1:E1")
    for col, text in zip("BCDE", ("Tier 1", "Tier 1 Incentive", "Tier 2", "Tier 2 Incentive")):
        ws[f"{col}2"] = text
    ws["A3"], ws["A4"] = "X TRADING A/C I", "X TRADING A/C III"
    for col, value in zip("BCDE", (300000, "A/C I- 3%", 300000, "A/C III- 2%")):
        ws[f"{col}3"] = value
        ws.merge_cells(f"{col}3:{col}4")


def test_one_threshold_repeated_by_tiers_naming_different_accounts_is_one_tier():
    parsed = _parse(_repeated_threshold_per_account)
    assert all(r.status == "mapped" for r in parsed.rows), [(r.row, r.reason) for r in parsed.rows]
    keys = [r.key for r in parsed.rows]
    (q,) = parsed.schemes
    assert [(t.threshold, t.rate, t.account_rates) for t in q.tiers] == [
        (D("300000"), None, {keys[0]: D("3"), keys[1]: D("2")}),
    ]


# review (gate 3) ----------------------------------------------------------------------


def test_two_branches_in_one_upload_each_map_to_their_own_ledger(api):
    """Review B1: the tie-break is not cached across branches."""
    client, db, company_id = api
    jb = _customer(db, company_id, "ALPHA HARDWARE (JB) SDN BHD [A/C I]", 1)
    kl = _customer(db, company_id, "ALPHA HARDWARE (KL) SDN BHD [A/C I]", 1)

    def fill(ws):
        ws["A1"], ws["B1"] = "DEALER NAME", "Q1"
        ws["A2"], ws["B2"] = "Alpha Hardware (KL) Sdn Bhd A/C I", "100,000 - 1%"
        ws["A3"], ws["B3"] = "Alpha Hardware (JB) Sdn Bhd A/C I", "200,000 - 2%"

    data, _ = _wb(fill)
    rows = _upload(client, db, data, dry_run=True)["rows"]
    assert [r["customer"] for r in rows] == [kl.customer_name, jb.customer_name]


def test_an_at_sign_in_a_sheet_name_or_pay_to_never_crashes(api):
    """Review B2: "@" is file text, never a key marker."""
    client, db, company_id = api
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "@HQ"
    ws["A1"], ws["B1"], ws["C1"] = "DEALER NAME", "Q1", "PAY TO"
    ws["A2"], ws["B2"] = "ZZT NOBODY A/C I", "100,000 - 1%"
    ws["A3"], ws["B3"], ws["C3"] = "ZZT NOBODY TWO A/C I", "100,000 - 1%", "@ZH01"
    buf = io.BytesIO()
    wb.save(buf)
    files = {"file": ("t.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res = client.post(f"{BASE}/import", params={"dry_run": "true", "year": 2026}, files=files)
    assert res.status_code == 200, res.text


def test_an_unpaired_incentive_in_a_split_block_fails_the_rows():
    """Review S1: the split never drops an incentive it cannot pair."""

    def fill(ws):
        ws["A1"], ws["B1"] = "DEALER NAME", "Q1"
        ws.merge_cells("B1:F1")
        for col, text in zip("BCDEF", ("Tier 1", "Tier 1 Incentive", "Tier 4", "Tier 4 Incentive", "Tier 5 Incentive")):
            ws[f"{col}2"] = text
        ws["A3"], ws["A4"] = "X TRADING A/C I", "X TRADING A/C III"
        ws["B3"], ws["B4"], ws["C3"], ws["C4"] = 100000, 120000, 0.01, 0.01
        ws["D3"], ws["E3"], ws["E4"], ws["F3"], ws["F4"] = 500000, 0.02, 0.02, 0.03, 0.03
        ws.merge_cells("D3:D4")

    parsed = _parse(fill)
    assert all(r.status == "not_representable" and "no threshold" in r.reason for r in parsed.rows), [
        (r.row, r.status, r.reason) for r in parsed.rows
    ]


def test_only_a_plain_incentive_after_a_quarterly_block_is_adopted():
    """Review S2: any other rate column the reader cannot place still fails loudly."""

    def fill(ws):
        ws["A1"], ws["B1"], ws["C1"] = "DEALER NAME", "YEARLY TARGET", "TOTAL INCENTIVE"
        ws["A2"], ws["B2"], ws["C2"] = "X TRADING A/C I", 500000, 0.05

    row = _parse(fill).rows[0]
    assert row.status == "not_representable" and "cannot place" in row.reason


def test_a_shared_quarterly_incentive_reaches_both_halves_of_a_split():
    """Review nit: the dealer's quarters and the accounts' quarters both read the one INCENTIVE."""

    def fill(ws):
        ws["A2"] = "CUSTOMER NAME"
        ws.merge_cells("A2:A3")
        ws["B2"] = "QUARTERLY (POINT)"
        ws.merge_cells("B2:C2")
        ws["B3"], ws["C3"], ws["D3"] = "JAN-MAR", "APR-JUN", "INCENTIVE"
        ws["A4"], ws["A5"] = "KILO LIVING A/C I", "KILO LIVING A/C II"
        ws["B4"], ws["B5"] = 100000, 90000             # Q1 per account
        ws["C4"] = 300000                               # Q2 merged: the dealer's
        ws.merge_cells("C4:C5")
        ws["D4"], ws["D5"] = 0.04, 0.03

    parsed = _parse(fill)
    assert all(r.status == "mapped" for r in parsed.rows), [(r.row, r.reason) for r in parsed.rows]
    keys = [r.key for r in parsed.rows]
    dealer = next(s for s in parsed.schemes if s.codes == keys)
    assert [(t.period_start, t.threshold, t.account_rates) for t in dealer.tiers] == [
        (date(2026, 4, 1), D("300000"), {keys[0]: D("4"), keys[1]: D("3")}),
    ]
    owns = sorted((s.codes[0], [(t.threshold, t.rate) for t in s.tiers]) for s in parsed.schemes if len(s.codes) == 1)
    assert owns == [(keys[0], [(D("100000"), D("4"))]), (keys[1], [(D("90000"), D("3"))])]


def test_a_rate_column_after_a_quarterly_block_not_named_incentive_still_fails():
    """Review nit: the header check itself, beside a quarterly block."""

    def fill(ws):
        ws["A1"], ws["B1"] = "DEALER NAME", "QUARTERLY (POINT)"
        ws["B2"], ws["C2"] = "OCT-DEC", "TOTAL INCENTIVE"
        ws["A3"], ws["B3"], ws["C3"] = "X TRADING A/C I", 100000, 0.04

    row = _parse(fill).rows[0]
    assert row.status == "not_representable" and "cannot place" in row.reason
