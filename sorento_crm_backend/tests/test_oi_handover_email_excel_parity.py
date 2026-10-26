"""OI-EMAIL-PURCHASING-FOCUS (owner, 26 Oct 2026): "the rows in OI must be truthful to the rows in
the email". The OI Excel a handover attaches is written from the email's own lines, so the two can
never differ: a new order, an amendment and a cancel-all each read the same rows in both.

Plan: `documentation/plans/scm/PLAN-oi-email-purchasing-focus-26oct.md` (rules 1 and 4).

TEST-FIRST. Driven through the board's own Confirm route; the dispatch is captured (nothing is
sent) and storage is kept in memory (nothing leaves the sandbox).
"""
from __future__ import annotations

import io
import uuid
from decimal import Decimal

import pytest

from app.models.order import SalesOrderLine

from .test_oi_saved_line_email import TRIGGER, _order, _post
from .test_planning_change_apply_on_board import BASE, api, world  # noqa: F401 - fixtures
from .test_planning_changes import _line_payload

__all__ = ["api", "world"]

#: Rule 4: the email's own line columns plus the S/O facts per row, change columns only when a
#: line in this email has one. No reserve figures, no TOTAL QTY.
NEW_ORDER_HEADINGS = [
    "SO DATE",
    "S/O NO",
    "ITEM CODE",
    "QTY",
    "DELIVERY DATE",
    "PROJECT/CUSTOMER",
    "LOCATION",
    "REMARK",
]


@pytest.fixture()
def sent(api, monkeypatch):
    """`(handover context, stored uploads)`: the drain bound to the test's own connection so it
    reads the uncommitted write, an enabled handover automation so the workbook is stored, the
    dispatch captured, storage in memory."""
    from sqlalchemy.orm import Session

    from app.models.automation import Automation
    from app.models.email_template import EmailTemplate
    from app.services.project_order_inquiry_service import (
        register_order_inquiry_post_commit_dispatch,
    )

    _client, world = api
    register_order_inquiry_post_commit_dispatch()
    connection = world.db.connection()
    monkeypatch.setattr(
        "app.database.SessionLocal",
        lambda: Session(bind=connection, join_transaction_mode="create_savepoint"),
    )
    uploads: dict = {}

    class _Backend:
        def upload_file(self, *, file_content, file_path, content_type=None, **_kw):
            uploads[file_path] = file_content
            return file_path, None

    monkeypatch.setattr(
        "app.services.storage_router.get_backend", lambda provider=None: _Backend()
    )
    contexts: list = []

    def _fake_dispatch(self, trigger_type, *, context, source_kind, source_id):
        if trigger_type == TRIGGER:
            contexts.append(context)
        return {"trigger_type": trigger_type, "fired": 0, "results": []}

    monkeypatch.setattr(
        "app.services.automation_service.AutomationService.dispatch_event", _fake_dispatch
    )
    template = EmailTemplate(
        id=str(uuid.uuid4()), code=f"zzt-oip-{uuid.uuid4().hex[:6]}", name="ZZT OIP",
        subject="OI: {{ handover.subject_scope }}", body_html="<p>x</p>", is_active=True,
    )
    world.db.add(template)
    world.db.flush()
    world.db.add(
        Automation(
            id=str(uuid.uuid4()), name="ZZT OIP handover", enabled=True,
            trigger_type=TRIGGER, trigger_config={}, action_type="send_email",
            email_template_id=str(template.id), recipient_config={"include_actor": True},
            group_matches=False, schedule_type="manual", timezone="Asia/Kuala_Lumpur",
        )
    )
    world.db.commit()
    return contexts, uploads


def _sheet(context, uploads):
    """`(headings, rows as dicts)` of the one OI Excel this email attaches."""
    import openpyxl

    excel = [
        a for a in context.get("email_attachments") or [] if a["filename"].endswith(".xlsx")
    ]
    assert len(excel) == 1, context.get("email_attachments")
    workbook = openpyxl.load_workbook(io.BytesIO(uploads[excel[0]["storage_key"]]))
    assert len(workbook.sheetnames) == 1, workbook.sheetnames
    rows = list(workbook.active.iter_rows(values_only=True))
    headings = [h for h in rows[1] if h is not None]
    body = [dict(zip(headings, row)) for row in rows[2:] if any(v is not None for v in row)]
    return headings, body


def _num(value) -> str:
    """A sheet number the way the email prints one (`10`, `2.5`)."""
    if value in (None, ""):
        return ""
    text = f"{float(value):f}".rstrip("0").rstrip(".")
    return text


def _excel_rows(body):
    return sorted(
        (
            row.get("ITEM CODE") or "",
            _num(row.get("QTY")),
            _num(row.get("QTY CHANGE TO")),
            row.get("LOCATION") or "",
            row.get("REMARK") or "",
        )
        for row in body
    )


def _email_rows(context):
    out = []
    for line in context["handover"]["lines"]:
        was = line.get("was") or {}
        if was.get("qty") is not None:
            qty, change_to = was["qty"], line.get("qty") or ""
        else:
            qty, change_to = line.get("qty") or "", ""
        out.append(
            (line.get("item_code") or "", qty, change_to, line.get("location") or "",
             line.get("remark") or "")
        )
    return sorted(out)


def test_a_new_order_excel_has_exactly_the_email_rows_and_the_short_column_list(api, sent):
    contexts, uploads = sent
    order, line_a, line_b = _order(api, qty_a="10", qty_b="6")

    _post(api, order, [
        _line_payload(line_a.id, buy_qty="10"),
        _line_payload(line_b.id, buy_qty="6"),
    ])

    [context] = contexts
    headings, body = _sheet(context, uploads)
    assert headings == NEW_ORDER_HEADINGS
    assert _excel_rows(body) == _email_rows(context)
    assert len(body) == context["handover"]["line_count"] == 2
    assert {row["S/O NO"] for row in body} == {order.autocount_doc_no}


def test_an_amendment_excel_is_the_one_changed_line_with_its_change_column(api, sent):
    _client, world = api
    contexts, uploads = sent
    order, line_a, line_b = _order(api, qty_a="10", qty_b="6")
    _post(api, order, [
        _line_payload(line_a.id, buy_qty="10"),
        _line_payload(line_b.id, buy_qty="6"),
    ])
    contexts.clear()
    # AutoCount cut line A from 10 to 8.
    core_a = world.db.get(SalesOrderLine, line_a.core_sales_order_line_id)
    core_a.qty_ordered = Decimal("8")
    line_a.qty = Decimal("8")
    world.db.commit()

    _post(api, order, [_line_payload(line_a.id, buy_qty="8", amend_reason="Site cut.")])

    [context] = contexts
    headings, body = _sheet(context, uploads)
    assert "QTY CHANGE TO" in headings
    for gone in ("REQUESTED RESERVE", "RESERVED", "TOTAL QTY"):
        assert gone not in headings
    assert _excel_rows(body) == _email_rows(context)
    assert len(body) == 1, body


def test_a_cancel_all_excel_lists_the_cancelled_lines_the_email_lists(api, sent):
    """SO394881's shape (masked): every line of an order with raised OI rows transferred to Cancel
    SO. The email reads one CANCEL line per line; the Excel used to print the OI's rows that were
    still open (none, once the cancel had closed them) instead of those lines."""
    from .test_oi_cancelled_line_used_confirm import _inquiry, _oi_row

    client, world = api
    contexts, uploads = sent
    order, line_a, line_b = _order(api, qty_a="10", qty_b="6")
    inquiry = _inquiry(world.db, world.company_id, order.id)
    code = world.product.product_code
    _oi_row(world.db, world.company_id, inquiry.id, line_a.id, code, "10")
    _oi_row(world.db, world.company_id, inquiry.id, line_b.id, code, "6")
    world.db.commit()
    contexts.clear()

    response = client.post(f"{BASE}/sales-orders/{order.id}/confirm", json={
        "lines": [],
        "cancelled_lines": [
            {"project_line_id": str(line_a.id), "reason": "Customer dropped it."},
            {"project_line_id": str(line_b.id), "reason": "Customer dropped it."},
        ],
    })
    assert response.status_code == 200, response.text
    world.db.commit()

    [context] = contexts
    assert {line["remark"] for line in context["handover"]["lines"]} == {"CANCEL"}
    _headings, body = _sheet(context, uploads)
    assert _excel_rows(body) == _email_rows(context)
    assert len(body) == 2, body


# ------------------------------------------- the hand-test script reads both the same way


def _render_current(lines):
    """The handover template as a real database holds it now (`oihr_0006` seeded, then every
    later layout migration), rendered for these lines."""
    from app.services.email_template_service import EmailTemplateService

    from ._pg_fixture import blank_session
    from .test_oi_saved_line_email import _context, _migration
    from .test_order_inquiry_handover_automation import _run_upgrade
    from app.models.email_template import EmailTemplate

    with blank_session() as db:
        for name in (
            "oihr_0006_no_remark_column", "oihr_0007_excel_remark", "oihr_0008_so_table_on_top",
            "oihr_0009_split_rows", "oihr_0010_so_remark_column",
        ):
            _run_upgrade(_migration(name), db)
        template = (
            db.query(EmailTemplate)
            .filter(EmailTemplate.code == "order_inquiry_handover_default")
            .one()
        )
        return EmailTemplateService(db).render(template, _context(lines))


def test_the_compare_script_reads_the_rendered_email_and_the_excel_as_the_same_rows():
    from scripts.compare_oi_handover_email_excel import email_rows, excel_rows

    from app.services.order_inquiry_worklist_service import handover_xlsx

    from .test_oi_saved_line_email import _line

    lines = [
        _line(),
        _line(item_code="ZZT-B", qty="8", was={"qty": "10", "delivery_date": "01/11/2026"},
              remark="DELAY, CANCEL 2"),
        _line(item_code="ZZT-NEW", was={"item_code": "ZZT-OLD"}, remark="NEW ITEM CODE"),
        _line(item_code="ZZT-C", qty="0", was={"qty": "6"}, remark="CANCEL"),
        # A split line: its second row prints only what that split changes (`oihr_0009`).
        _line(item_code="ZZT-S", qty="4", was={"qty": "10"}, remark="ORDER 4"),
        _line(item_code="ZZT-S", qty="6", was={"qty": "10", "delivery_date": "01/12/2026"},
              remark="ORDER 6", continued=True),
    ]
    rendered = _render_current(lines)

    head, emailed = email_rows(rendered["body_html"])
    sheet_head, sheet = excel_rows(handover_xlsx(lines), head)

    assert {"NEW ITEM CODE", "NEW QTY", "NEW DATE"} <= set(head), head
    assert {"NEW ITEM CODE", "QTY CHANGE TO", "DELIVERY DATE CHANGE TO"} <= set(sheet_head), sheet_head
    assert len(emailed) == 6, emailed
    assert emailed == sheet
