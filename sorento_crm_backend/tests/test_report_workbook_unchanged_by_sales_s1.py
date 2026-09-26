"""S1 (#1267), AC-R2-2 / AC-R4-7: the kernel extensions leave every existing workbook as it
was. The fixture is the cell-for-cell dump (value, bold, number format, merged ranges) of
the synthetic report and the sponsorship report rendered by origin/main's renderer before
S1 (reviewer B2: the TOTAL row's amounts had turned bold). Regenerate it only when a change
to an existing workbook is intended, and say so in the PR.
"""
from __future__ import annotations

import io
import json
from pathlib import Path

from openpyxl import load_workbook

from app.main import app  # noqa: F401
from tests import _report_fixture as fixture
from tests._pg_fixture import blank_session

_FIXTURE = Path(__file__).parent / "fixtures" / "reports" / "workbook_before_sales_s1.json"


def _dump(content: bytes) -> dict:
    book = load_workbook(io.BytesIO(content))
    out = {}
    for sheet in book.worksheets:
        cells = []
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None and not cell.font.b:
                    continue
                value = str(cell.value) if cell.value is not None else None
                cells.append([cell.coordinate, value, bool(cell.font.b), cell.number_format])
        out[sheet.title] = {"cells": cells, "merged": sorted(str(m) for m in sheet.merged_cells.ranges)}
    return out


def test_existing_workbooks_are_unchanged_cell_for_cell():
    from app.services.reports import engine, registry as reg
    from app.services.reports.xlsx_renderer import render_workbook

    expected = json.loads(_FIXTURE.read_text())
    with blank_session() as db:
        fixture.create_table(db)
        definition = fixture.definition()
        data = engine.run_workbook(
            db, definition, {"date_basis": "booked_on", "period": {"kind": "year", "year": 2026}}
        )
        synthetic = json.loads(json.dumps(_dump(render_workbook(definition, data))))
        sponsorship = reg.get("sponsorship")
        data = engine.run_workbook(db, sponsorship, {"period": {"kind": "year", "year": 2025}})
        empty = json.loads(json.dumps(_dump(render_workbook(sponsorship, data))))
    assert synthetic == expected["synthetic"]
    assert empty == expected["sponsorship_empty"]
