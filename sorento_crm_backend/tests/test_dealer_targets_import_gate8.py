"""Dealer targets import, eighth real-file gate (lane SALES-TARGET-DEALER, crew 17 Oct, round 7).

Sorento's yearly block ends in a "Yearly / Special Remark" column, merged over a dealer's
account rows and holding the yearly period ("DEC'25-NOV'26"). The reader took it for a yearly
THRESHOLD merged over the accounts, so the account / dealer split made it the dealer side; that
side had no tiers, and `if not dealer: return []` dropped every account's own yearly tiers. 13
blocks, 20 rows with yearly figures and no yearly target, 16 of them counted as mapped.

1. A yearly remark column is the block's window, never a threshold.
2. A dealer side with no tiers ("NIL" merged over the accounts) never drops the accounts' own
   tiers: they are built as the accounts' own targets.
Parse-only (no database), masked names.
"""
from __future__ import annotations

import io
from datetime import date


def _parse(remark=None, tier4=None, with_remark=True):
    from openpyxl import Workbook

    from app.services.sales.dealer_target_import import parse_workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "DEALER NAME"
    ws.merge_cells("A1:A2")
    ws["B1"] = "ACCOUNT"
    ws.merge_cells("B1:B2")
    ws["C1"] = "YEARLY"
    ws.merge_cells("C1:F1")
    fourth = ("Tier 4", "Tier 4 Incentive") if tier4 is not None else ("Tier 2", "Tier 2 Incentive")
    for col, text in zip("CDEF", ("Tier 1", "Tier 1 Incentive") + fourth):
        ws[f"{col}2"] = text
    if with_remark:
        ws["G1"] = "Yearly / Special Remark"
        ws.merge_cells("G1:G2")
    ws["A3"] = "ECHO HOME"
    ws.merge_cells("A3:A4")
    ws["B3"], ws["B4"] = "A/C I", "A/C II"
    ws["C3"], ws["D3"], ws["C4"], ws["D4"] = 100000, 0.01, 50000, 0.01
    if tier4 is not None:
        ws["E3"] = tier4
        ws.merge_cells("E3:E4")
        ws.merge_cells("F3:F4")
    else:
        ws["E3"], ws["F3"], ws["E4"], ws["F4"] = 200000, 0.02, 80000, 0.02
    if with_remark:
        ws["G3"] = remark
        ws.merge_cells("G3:G4")
    buf = io.BytesIO()
    wb.save(buf)
    return parse_workbook(buf.getvalue(), year=2026)


def _yearly(parsed):
    return sorted(
        ((s.codes, s.start, s.end, [(t.threshold, t.rate) for t in s.tiers], s.replaced_by)
         for s in parsed.schemes if s.kind == "yearly"),
        key=lambda x: x[0],
    )


def test_a_yearly_remark_holding_the_period_is_the_window_of_each_account_s_target():
    parsed = _parse(remark="DEC'25-NOV'26")
    assert [(r.status, r.warnings) for r in parsed.rows] == [("mapped", []), ("mapped", [])], [
        (r.reason, r.warnings) for r in parsed.rows
    ]
    got = _yearly(parsed)
    assert len(got) == 2 and all(len(codes) == 1 for codes, *_ in got), got
    assert {(start, end) for _c, start, end, _t, _r in got} == {(date(2025, 12, 1), date(2026, 11, 30))}
    assert [len(tiers) for _c, _s, _e, tiers, _r in got] == [2, 2]


def test_a_blank_yearly_remark_merged_over_the_accounts_does_not_join_them():
    parsed = _parse(remark=None)
    assert [r.status for r in parsed.rows] == ["mapped", "mapped"], [r.reason for r in parsed.rows]
    got = _yearly(parsed)
    assert len(got) == 2 and {(s, e) for _c, s, e, _t, _r in got} == {(date(2026, 1, 1), date(2026, 12, 31))}


def test_a_yearly_remark_that_is_no_period_is_told_never_read_as_a_figure():
    parsed = _parse(remark="INCLUDE DELUXE")
    assert [r.status for r in parsed.rows] == ["mapped", "mapped"], [r.reason for r in parsed.rows]
    assert len(_yearly(parsed)) == 2
    assert all(any("not a period" in w for w in r.warnings) for r in parsed.rows), [r.warnings for r in parsed.rows]


def test_a_dealer_side_with_no_tiers_keeps_the_accounts_own_tiers():
    """"NIL" Tier 4 merged over the accounts: no dealer tier, each account its own Tier 1."""
    parsed = _parse(tier4="NIL", with_remark=False)
    assert [r.status for r in parsed.rows] == ["mapped", "mapped"], [r.reason for r in parsed.rows]
    got = _yearly(parsed)
    assert len(got) == 2, got
    assert all(len(codes) == 1 and replaced is None and len(tiers) == 1 for codes, _s, _e, tiers, replaced in got), got
