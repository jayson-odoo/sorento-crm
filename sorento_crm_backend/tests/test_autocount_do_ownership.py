"""Column ownership between the AutoCount DO ingest and the uploads (#1354 S2, plan section 3).

UAC: documentation/plans/autocount/autocount-grn-do-ingest-29sep-acceptance-criteria.md.

  AC-AG043  the tracking upload's Master sheet on an AutoCount-owned DO keeps AutoCount's
            columns and still writes its own (`remarks_cs`)
  AC-AG046  the DO detail import skips a row whose DO is AutoCount-owned (`autocount_owned`)

A row is AutoCount-owned when `doc_key IS NOT NULL`.
"""
from __future__ import annotations

import uuid
from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.models.order import Order, OrderStatus
from app.services.order_service import OrderService
from tests._pg_fixture import blank_session

# The DO detail import's own per-schema fixtures, reused as they are.
from tests.test_import_outcome_attribution import (  # noqa: F401 - pytest fixtures
    _breakdown_counts,
    _build_workbook_bytes,
    _make_job,
    _run_import,
    seeded,
    session_factory,
)


@pytest.fixture()
def db():
    with blank_session() as session:
        yield session


def _master_workbook(rows: list[list]) -> bytes:
    wb = Workbook()
    default = wb.active
    if default is not None:
        wb.remove(default)
    master = wb.create_sheet("Master")
    master.append(["Doc. No.", "Date", "Debtor Code", "Debtor Name", "Agent", "Cancel", "Remarks CS"])
    for r in rows:
        master.append(r)
    tracking = wb.create_sheet("Overall Tracking")
    tracking.append(["Doc Number", "Date", "Transporter", "Driver Name", "Lorry Plate"])
    tracking.append([rows[0][0], date(2026, 9, 29), "ZZ TRANS", "Ali", "WXX 1"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_master_sheet_keeps_autocount_columns(db):
    """AC-AG043."""
    for code, name in [("NEW", "New Order"), ("DELIVERED", "Delivered")]:
        db.add(OrderStatus(status_code=code, status_name=name))
    order = Order(
        order_number="ZZDO-OWN-1", order_date=date(2026, 9, 27), debtor_code="300-AC",
        debtor_name="AutoCount Name", agent="AC-AGENT", is_cancelled=False,
        doc_key=990001, source_book="db1",
    )
    plain = Order(order_number="ZZDO-OWN-2", order_date=date(2026, 9, 27), debtor_code="OLD")
    db.add_all([order, plain])
    db.commit()

    data = _master_workbook([
        ["ZZDO-OWN-1", date(2026, 9, 1), "999-SHEET", "Sheet Name", "SHEET-AGENT", "Y", "cs note"],
        ["ZZDO-OWN-2", date(2026, 9, 1), "999-SHEET", "Sheet Name", "SHEET-AGENT", "N", "cs 2"],
    ])
    result = OrderService(db).import_excel_tracking(data, "11111111-1111-1111-1111-111111111111")
    assert not result.get("errors"), result.get("errors")

    db.refresh(order)
    assert order.order_date == date(2026, 9, 27)
    assert (order.debtor_code, order.debtor_name, order.agent) == ("300-AC", "AutoCount Name", "AC-AGENT")
    assert order.is_cancelled is False
    assert order.customer_id is None
    # Upload-owned columns still land on the AutoCount-owned row.
    assert order.remarks_cs == "cs note"
    assert order.transporter == "ZZ TRANS" and order.lorry_plate == "WXX 1"
    # A row AutoCount does not own is written as before.
    db.refresh(plain)
    assert plain.debtor_code == "999-SHEET" and plain.order_date == date(2026, 9, 1)


def test_do_detail_import_skips_autocount_owned_do(seeded):  # noqa: F811
    """AC-AG046."""
    factory, fx = seeded
    s = factory()
    try:
        # A Core UPDATE on the Table: no ORM company filter, still schema-translated.
        table = Order.__table__
        s.execute(
            table.update()
            .where(table.c.order_number == fx.doc_no)
            .values(doc_key=990002, source_book="db1")
        )
        s.commit()
    finally:
        s.close()
    user_id = str(uuid.uuid4())
    job = _run_import(factory, _make_job(factory, user_id), _build_workbook_bytes(fx), user_id)
    assert job.successful_rows == 0
    counts = _breakdown_counts(job)
    assert counts["skipped"].get("autocount_owned") == 3


def test_grn_excel_import_leaves_autocount_grn_untouched(db):
    """AC-AG047: the GRN Excel import writes neither the header nor a line of an
    AutoCount-owned GRN."""
    from app.models.procurement import PickingHeader, PickingLine
    from app.models.product import Product, ProductCategory, UnitOfMeasure
    from app.services.procurement_service import PickingHeaderService

    category = ProductCategory(category_code=f"ZZAC-{uuid.uuid4().hex[:6]}", category_name="c")
    uom = UnitOfMeasure(uom_code=f"ZZAC-{uuid.uuid4().hex[:6]}", uom_name="u")
    db.add_all([category, uom])
    db.flush()
    product = Product(product_code=f"ZZAC-{uuid.uuid4().hex[:6]}", product_name="p",
                      category_id=category.id, base_uom_id=uom.id, list_price=1)
    header = PickingHeader(picking_number="ZZGRN-OWN-1", picking_type="goods_received",
                           picking_date=date(2026, 7, 27), picking_status="approved",
                           source_system="autocount", doc_key=880001, source_book="db1",
                           spo_number=None)
    db.add_all([product, header])
    db.commit()

    svc = PickingHeaderService(db)
    got, created = svc.upsert_grn_header_for_import("ZZGRN-OWN-1", "SPO-SHEET", date(2026, 1, 1))
    assert created is False and str(got.id) == str(header.id)
    db.refresh(header)
    assert header.spo_number is None and header.picking_date == date(2026, 7, 27)

    svc.upsert_grn_line_for_import(
        picking_header_id=header.id, product_id=product.id, source_warehouse_id=None,
        quantity=5,
    )
    db.expire_all()
    assert db.query(PickingLine).filter(PickingLine.picking_header_id == header.id).count() == 0
