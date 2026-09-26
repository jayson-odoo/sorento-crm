"""S1 of the retail sales reports (#1267): the sales dataset, the Yearly comparison and the
report kernel extensions it needs.

`documentation/plans/sales/PLAN-retail-sales-reports-26sep.md` (5.1 to 5.3, S1) and its UAC
(AC-S1-1 to AC-S1-17, AC-R2-2, AC-R2-4, AC-R2-17, AC-R3-3, AC-R4-7). Engine level: the
routes are in `test_sales_yearly_routes.py`, the chatbot seam in
`test_sales_analysis_route.py`.

Postgres only, every row seeded here (CI's database holds none). Sorento is seeded into
every blank schema by conftest; Mocha is added by `seed_mocha`.
"""
from __future__ import annotations

import io
import uuid
from datetime import date
from decimal import Decimal

import pytest

from app.main import app  # noqa: F401  (import order: resolves the guards cycle)
from app.models.base import set_company_scope
from app.models.order import SalesOrder, SalesOrderLine
from app.services.company_scope import DEFAULT_COMPANY_ID
from tests._mc_lookup_seed import MOCHA_ID, customer, product, seed_mocha
from tests._pg_fixture import blank_session, unique_code

KEY = "sales_yearly"


@pytest.fixture
def db():
    with blank_session() as s:
        set_company_scope(s, frozenset({DEFAULT_COMPANY_ID}))
        yield s


@pytest.fixture
def definition():
    from app.services.reports import registry as reg

    found = reg.get(KEY)
    assert found is not None, "sales_yearly is not registered"
    return found


def _line(
    db,
    *,
    company_id=DEFAULT_COMPANY_ID,
    order_date,
    ordered,
    delivered,
    line_total,
    demand_class="retail",
    status="open",
    line_status="open",
    required_date=None,
    product_id=None,
    customer_id=None,
):
    so = SalesOrder(
        id=str(uuid.uuid4()),
        so_number=unique_code("SO"),
        order_date=order_date,
        status=status,
        demand_class=demand_class,
        customer_id=customer_id,
        company_id=company_id,
    )
    db.add(so)
    db.flush()
    if product_id is None:
        product_id = product(db, company_id=company_id).id
    db.add(
        SalesOrderLine(
            id=str(uuid.uuid4()),
            sales_order_id=so.id,
            product_id=product_id,
            qty_ordered=ordered,
            qty_delivered=delivered,
            line_total=line_total,
            line_status=line_status,
            required_date=required_date,
            company_id=company_id,
        )
    )
    db.flush()
    return so


def _params(**over):
    params = {
        "date_basis": "order_date",
        "period": {"kind": "custom", "from": "2024-01-01", "to": "2026-09-26"},
        "company": [DEFAULT_COMPANY_ID],
        "channel": ["dealer", "project"],
        "basis": ["delivered"],
    }
    params.update(over)
    return params


def _view(rows="year", cols="month_of_year", measures=("sales_value",), **params):
    from app.schemas.report import ReportViewConfig

    return ReportViewConfig.model_validate(
        {
            "params": _params(**params),
            "detail": {"columns": [], "order": []},
            "pivot": {"rows": rows, "cols": cols, "measures": list(measures)},
        }
    )


def _run(db, definition, grants=frozenset({DEFAULT_COMPANY_ID}), **kw):
    from app.services.reports import engine

    view = _view(**kw)
    return engine.run(db, definition, view.params, view, company_grants=grants)


# ------------------------------------------------------------------------ the grid


def test_ac_s1_1_both_bases_equal_sales_report_to_the_sen(db, definition):
    """AC-S1-1: the dataset's ordered and delivered totals equal `sales_report`'s."""
    from app.services.sales_report_service import sales_report

    cust = customer(db, company_id=DEFAULT_COMPANY_ID)
    # 3 of 7 delivered on 100.00: confirmed = round(100 * 3 / 7, 2) = 42.86, per line.
    _line(db, order_date=date(2026, 2, 10), ordered=7, delivered=3, line_total=Decimal("100.00"), customer_id=cust.id)
    _line(db, order_date=date(2026, 3, 1), ordered=2, delivered=2, line_total=Decimal("55.55"), customer_id=cust.id)
    # Closed and under-delivered: ordered basis counts only what was confirmed.
    _line(db, order_date=date(2026, 4, 1), ordered=4, delivered=1, line_total=Decimal("40.00"),
          status="closed", line_status="closed", customer_id=cust.id)

    ref = sales_report(db, customer_ids=[cust.id], date_from=date(2026, 1, 1), date_to=date(2026, 12, 31))
    ref_ordered = sum(Decimal(str(m["ordered_value"])) for m in ref["months"])
    ref_confirmed = sum(Decimal(str(m["confirmed_value"])) for m in ref["months"])

    ordered = _run(db, definition, rows="channel", cols="year", basis=["ordered"],
                   period={"kind": "year", "year": 2026})
    delivered = _run(db, definition, rows="channel", cols="year", basis=["delivered"],
                     period={"kind": "year", "year": 2026})
    assert Decimal(ordered.layouts.summary.grand_total["sales_value"]) == ref_ordered
    assert Decimal(delivered.layouts.summary.grand_total["sales_value"]) == ref_confirmed
    assert ref_confirmed == Decimal("42.86") + Decimal("55.55") + Decimal("10.00")


def test_ac_s1_2_cancelled_order_and_cancelled_line_are_nowhere(db, definition):
    _line(db, order_date=date(2026, 1, 5), ordered=1, delivered=1, line_total=Decimal("10.00"))
    _line(db, order_date=date(2026, 1, 6), ordered=1, delivered=1, line_total=Decimal("999.00"), status="cancelled")
    _line(db, order_date=date(2026, 1, 7), ordered=1, delivered=1, line_total=Decimal("888.00"), line_status="cancelled")
    result = _run(db, definition)
    assert Decimal(result.layouts.summary.grand_total["sales_value"]) == Decimal("10.00")
    assert result.row_count == 1


def test_ac_s1_3_a_march_line_on_a_february_order_counts_in_february(db, definition):
    _line(db, order_date=date(2026, 2, 20), ordered=1, delivered=1, line_total=Decimal("30.00"),
          required_date=date(2026, 3, 15))
    summary = _run(db, definition).layouts.summary
    assert summary.cells["2026"]["02"]["sales_value"] == "30.00"
    assert "03" not in summary.cells["2026"]


def test_ac_s1_4_a_null_demand_class_is_a_blank_row_last_and_in_the_total(db, definition):
    _line(db, order_date=date(2026, 1, 5), ordered=1, delivered=1, line_total=Decimal("10.00"), demand_class="retail")
    _line(db, order_date=date(2026, 1, 5), ordered=1, delivered=1, line_total=Decimal("20.00"), demand_class="project")
    _line(db, order_date=date(2026, 1, 5), ordered=1, delivered=1, line_total=Decimal("5.00"), demand_class=None)
    summary = _run(db, definition, rows="channel", cols="year", channel=[]).layouts.summary
    assert summary.row_values == ["Dealer", "Project team", "(blank)"]
    assert summary.grand_total["sales_value"] == "35.00"


def test_ac_s1_5_ac_r2_4_company_isolation_and_a_company_outside_the_grant_is_403(db, definition):
    from app.services.error_handler import AppException

    seed_mocha(db)
    _line(db, order_date=date(2026, 1, 5), ordered=1, delivered=1, line_total=Decimal("10.00"))
    _line(db, company_id=MOCHA_ID, order_date=date(2026, 1, 5), ordered=1, delivered=1,
          line_total=Decimal("700.00"))

    sorento = _run(db, definition)
    assert sorento.layouts.summary.grand_total["sales_value"] == "10.00"

    with pytest.raises(AppException) as refused:
        _run(db, definition, company=[MOCHA_ID])
    assert refused.value.status_code == 403

    both = frozenset({DEFAULT_COMPANY_ID, MOCHA_ID})
    mocha = _run(db, definition, grants=both, company=[MOCHA_ID])
    assert mocha.layouts.summary.grand_total["sales_value"] == "700.00"


def test_ac_r2_4_company_arm_is_fail_closed_with_no_scope(db, definition):
    """No resolved scope = no rows, never every company's."""
    from app.models.base import UNSET

    _line(db, order_date=date(2026, 1, 5), ordered=1, delivered=1, line_total=Decimal("10.00"))
    for grants in (UNSET, frozenset()):
        result = _run(db, definition, grants=grants, company=[])
        assert result.row_count == 0
        assert result.layouts.summary.grand_total == {}


def test_company_defaults_to_the_session_company_when_none_is_named(db, definition):
    from app.services.reports import engine

    _line(db, order_date=date(2026, 1, 5), ordered=1, delivered=1, line_total=Decimal("12.00"))
    params = _params()
    params.pop("company")
    ctx = engine.resolve(db, definition, params, company_grants=frozenset({DEFAULT_COMPANY_ID}))
    assert ctx.values["company"] == [DEFAULT_COMPANY_ID]


# ---------------------------------------------------------------------- the report


def test_ac_s1_9_rows_are_every_year_and_months_after_the_as_at_date_are_blank(db, definition):
    _line(db, order_date=date(2025, 11, 3), ordered=1, delivered=1, line_total=Decimal("100.00"))
    _line(db, order_date=date(2026, 9, 1), ordered=1, delivered=1, line_total=Decimal("50.00"))
    summary = _run(db, definition).layouts.summary
    assert summary.row_values == ["2024", "2025", "2026"]
    assert summary.col_dim.values == [f"{m:02d}" for m in range(1, 13)]
    assert summary.col_dim.value_labels["01"] == "JAN"
    assert summary.col_dim.value_labels["12"] == "DEC"
    # blank, not 0: no cell at all for October to December 2026
    for month in ("10", "11", "12"):
        assert month not in summary.cells.get("2026", {})
    assert "2024" not in summary.cells  # an empty year is a row with no cells


def test_ac_s1_10_variance_is_last_year_minus_the_one_before_over_this_years_months(db, definition):
    _line(db, order_date=date(2025, 1, 10), ordered=1, delivered=1, line_total=Decimal("100.00"))
    _line(db, order_date=date(2025, 2, 10), ordered=1, delivered=1, line_total=Decimal("80.00"))
    _line(db, order_date=date(2025, 12, 10), ordered=1, delivered=1, line_total=Decimal("500.00"))
    _line(db, order_date=date(2026, 1, 10), ordered=1, delivered=1, line_total=Decimal("60.00"))
    _line(db, order_date=date(2026, 2, 10), ordered=1, delivered=1, line_total=Decimal("90.00"))
    summary = _run(db, definition).layouts.summary
    assert summary.variance_row["01"]["sales_value"] == "-40.00"
    assert summary.variance_row["02"]["sales_value"] == "10.00"
    assert "12" not in summary.variance_row  # after the as-at month: blank
    # G5 (a): 2026 YTD minus 2025 over the same months, not minus the whole of 2025
    assert summary.variance_total["sales_value"] == "-30.00"
    assert summary.variance_label == "VARIANCE"


def test_ac_r4_7_ac_s1_11_one_block_per_ticked_channel_titled_by_company(db, definition):
    _line(db, order_date=date(2026, 1, 10), ordered=1, delivered=1, line_total=Decimal("10.00"), demand_class="retail")
    _line(db, order_date=date(2026, 1, 10), ordered=1, delivered=1, line_total=Decimal("7.00"), demand_class="project")
    result = _run(db, definition)
    blocks = result.layouts.blocks
    assert [b.title for b in blocks] == ["SORENTO - DEALER", "SORENTO - PROJECT TEAM"]
    assert blocks[0].summary.grand_total["sales_value"] == "10.00"
    assert blocks[1].summary.grand_total["sales_value"] == "7.00"
    assert blocks[0].summary.chart == "line"
    assert blocks[0].summary.row_values == ["2024", "2025", "2026"]

    one = _run(db, definition, channel=["project"]).layouts.blocks
    assert [b.title for b in one] == ["SORENTO - PROJECT TEAM"]


def test_ac_r3_3_basis_is_required_single_and_defaults_to_delivered(db, definition):
    from app.services.error_handler import AppException
    from app.services.reports import engine

    basis = next(p for p in definition.params if p.key == "basis")
    assert basis.default == ("delivered",)
    assert basis.multi is False and basis.clearable is False
    with pytest.raises(AppException) as bad:
        _run(db, definition, basis=["invoiced"])
    assert bad.value.status_code == 422
    view = engine.view_config(definition)
    assert view.params["basis"] == ["delivered"]


def test_the_basis_line_is_on_the_result(db, definition):
    delivered = _run(db, definition)
    assert delivered.note == (
        "Basis: Delivered (transferred to DO), by sales order date. Sales orders, not invoices."
    )
    ordered = _run(db, definition, basis=["ordered"])
    assert ordered.note.startswith("Basis: Ordered,")


def test_ac_r2_17_screen_prints_whole_ringgit_and_opens_on_the_summary(db, definition):
    summary = _run(db, definition).layouts.summary
    assert summary.whole_units is True
    assert summary.show_column_totals is False
    assert definition.opens_on == "summary"


def test_a_long_detail_is_truncated_not_refused(db, definition, monkeypatch):
    """Three years of lines exceed the 5,000 row cap; the summary must still answer."""
    from app.services.reports import engine

    monkeypatch.setattr(engine, "DETAIL_ROW_CAP", 2)
    for day in (1, 2, 3):
        _line(db, order_date=date(2026, 1, day), ordered=1, delivered=1, line_total=Decimal("1.00"))
    result = _run(db, definition)
    assert result.layouts.detail.truncated is True
    assert len(result.layouts.detail.rows) == 2
    assert result.layouts.summary.grand_total["sales_value"] == "3.00"


# -------------------------------------------------------------------- the workbook


def _workbook(db, definition, **kw):
    from openpyxl import load_workbook

    from app.services.reports import engine
    from app.services.reports.xlsx_renderer import render_workbook

    view = _view(**kw)
    data = engine.run_workbook(
        db, definition, view.params, view, company_grants=frozenset({DEFAULT_COMPANY_ID})
    )
    return load_workbook(io.BytesIO(render_workbook(definition, data)))


def _cells(sheet):
    return [c.value for row in sheet.iter_rows() for c in row if c.value is not None]


def test_ac_s1_17_ac_r4_7_one_sheet_both_blocks_variance_chart_and_values(db, definition):
    _line(db, order_date=date(2025, 1, 10), ordered=1, delivered=1, line_total=Decimal("100.00"))
    _line(db, order_date=date(2026, 1, 10), ordered=1, delivered=1, line_total=Decimal("60.00"))
    _line(db, order_date=date(2026, 1, 10), ordered=1, delivered=1, line_total=Decimal("9.00"), demand_class="project")
    book = _workbook(db, definition)
    assert book.sheetnames == ["SUMMARY"]  # month_sheets off
    sheet = book["SUMMARY"]
    values = _cells(sheet)
    texts = [v for v in values if isinstance(v, str)]
    assert "YEARLY SALES COMPARISON" in texts
    assert "AS AT 26/09/2026" in texts
    assert any(t.startswith("Basis: Delivered") for t in texts)
    dealer = texts.index("SORENTO - DEALER")
    project = texts.index("SORENTO - PROJECT TEAM")
    assert dealer < project
    assert texts.count("VARIANCE") == 2
    assert len(sheet._charts) == 2
    # totals are values, never formulas
    assert not any(isinstance(v, str) and v.startswith("=") for v in values)
    money = [c for row in sheet.iter_rows() for c in row if isinstance(c.value, (int, float, Decimal))]
    assert money and all(c.number_format == "#,##0.00;(#,##0.00)" for c in money)
    assert Decimal("-40.00") in [Decimal(str(c.value)) for c in money]


def test_ac_s1_17_a_formula_looking_text_cell_is_escaped(db, definition, monkeypatch):
    from app.services.reports import xlsx_renderer

    assert xlsx_renderer.safe_text("=HYPERLINK(1)") == "'=HYPERLINK(1)"
    assert xlsx_renderer.safe_text("+1") == "'+1"
    assert xlsx_renderer.safe_text("-1") == "'-1"
    assert xlsx_renderer.safe_text("@x") == "'@x"
    assert xlsx_renderer.safe_text("SORENTO") == "SORENTO"


def test_the_sponsorship_workbook_keeps_its_month_sheets(db):
    from app.services.reports import registry as reg

    sponsorship = reg.get("sponsorship")
    assert sponsorship.workbook.month_sheets is True
    assert sponsorship.workbook.sheet_per is None
    assert sponsorship.module_key == "procurement"
    assert sponsorship.pivot.variance is None and sponsorship.pivot.chart is None
