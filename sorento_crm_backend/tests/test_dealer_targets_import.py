"""Dealer targets Excel import (lane SALES-TARGET-DEALER, UAC D3).

The owner's acceptance gate (16 Oct 00:19): every row of both real files must be representable.
The real files never enter the repo; `_dealer_target_fixture.build` writes a workbook with every
pattern the crew found (P1 to P11, listed in that module), and this file proves each one is read,
imported and computed. Layout: documentation/plans/sales/dealer-targets-import-layout.md.

Two halves: `parse_workbook` (pure, no database: the sheet reader and the cell grammar) and the
`POST /sales/targets/import` route (codes resolved against the CRM, dry run, apply, re-import).
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal

from . import _dealer_target_fixture as fx
from .test_sales_targets_s1 import (  # noqa: F401  (fixtures)
    BASE,
    _category,
    _code,
    _product,
    _so_line,
    _uid,
    api,
    world,
)

IMPORT = f"{BASE}/import"
D = Decimal


# --------------------------------------------------------------------------------------- #
# pure parse
# --------------------------------------------------------------------------------------- #


def _literal_codes():
    return {key: f"C-{key}" for key in fx.KEYS}


def _parse():
    from app.services.sales.dealer_target_import import parse_workbook

    return parse_workbook(fx.build(_literal_codes()), year=2026)


def _scheme(parsed, first_code, kind):
    return next(s for s in parsed.schemes if s.codes[0] == first_code and s.kind == kind)


def _tiers(scheme, period_start):
    return [
        (t.threshold, t.rate, t.is_additional, t.account_rates)
        for t in scheme.tiers
        if t.period_start == period_start
    ]


def test_parse_reports_every_row():
    parsed = _parse()
    by_code = {r.code: r for r in parsed.rows}
    # 9 Sorento rows, 3 monthly, 4 Mocha.
    assert len(parsed.rows) == 16
    assert by_code["C-E1"].status == "not_representable"
    assert "see note" in by_code["C-E1"].reason
    # The parser does not know the CRM: an unknown code is the route's call.
    assert by_code[fx.UNKNOWN_CODE].status == "mapped"
    assert all(r.status == "mapped" for r in parsed.rows if r.code != "C-E1")
    assert {r.sheet for r in parsed.rows} == {"Sorento 2026", "Monthly", "Mocha"}


def test_parse_p1_p2_p3_merged_accounts_named_rate_and_stacking_additionals():
    parsed = _parse()
    q = _scheme(parsed, "C-A1", "quarterly")
    assert q.codes == ["C-A1", "C-A3"]                      # P1: one target over both accounts
    assert q.payout_code == "C-A1"
    assert (q.start, q.end, q.split_every, q.split_unit) == (date(2026, 1, 1), date(2026, 12, 31), 3, "month")
    assert q.threshold_unit == "amount"
    assert _tiers(q, date(2026, 1, 1)) == [                 # P3
        (D("150000"), D("2"), False, None),
        (D("200000"), D("3"), False, None),
        (D("220000"), D("1"), True, None),
        (D("250000"), D("0.5"), True, None),
    ]
    assert _tiers(q, date(2026, 4, 1)) == [                 # P2: only A/C I earns
        (D("200000"), None, False, {"C-A1": D("2")}),
        (D("250000"), None, False, {"C-A1": D("3")}),
    ]
    assert _tiers(q, date(2026, 7, 1)) == [(D("240000"), D("2.5"), False, None)]   # P9 "RM"
    assert _tiers(q, date(2026, 10, 1)) == [(D("260000"), D("3"), False, None)]    # P9 "k"
    y = _scheme(parsed, "C-A1", "yearly")                   # P5: three tiers
    assert [t.threshold for t in y.tiers] == [D("850000"), D("1000000"), D("1200000")]
    assert (y.start, y.end, y.split_every) == (date(2026, 1, 1), date(2026, 12, 31), None)


def test_parse_p4_rate_per_account_deluxe_mastile_wply():
    parsed = _parse()
    q = _scheme(parsed, "C-B_DELUXE", "quarterly")
    assert q.codes == ["C-B_DELUXE", "C-B_MASTILE", "C-B_WPLY"]
    assert _tiers(q, date(2026, 1, 1)) == [
        (D("300000"), None, False, {"C-B_DELUXE": D("2"), "C-B_MASTILE": D("3"), "C-B_WPLY": D("1.5")}),
    ]
    # "-" is no scheme for Q2 to Q4.
    assert _tiers(q, date(2026, 4, 1)) == []
    assert q.payout_code == "C-B_DELUXE"


def test_parse_p6_dec_to_nov_window():
    y = _scheme(_parse(), "C-C1", "yearly")
    assert (y.start, y.end) == (date(2025, 12, 1), date(2026, 11, 30))
    assert [(t.threshold, t.rate) for t in y.tiers] == [(D("1200000"), D("1")), (D("1500000"), D("1.25"))]
    assert all(t.period_start == date(2025, 12, 1) for t in y.tiers)


def test_parse_p7_monthly_pairs_mixed_cells():
    parsed = _parse()
    for key in ("M1", "M2", "M3"):
        m = _scheme(parsed, f"C-{key}", "monthly")
        assert (m.split_every, m.split_unit) == (1, "month")
        assert len({t.period_start for t in m.tiers}) == 12
        assert _tiers(m, date(2026, 1, 1)) == [(D("50000"), D("1"), False, None)]
        assert _tiers(m, date(2026, 2, 1)) == [(D("51000"), D("1"), False, None)]


def test_parse_p8_mocha_points_header_rate_and_yearly_tiers():
    parsed = _parse()
    k = _scheme(parsed, "C-K1", "quarterly")
    assert k.codes == ["C-K1", "C-K2", "C-K3"]
    assert k.threshold_unit == "points"
    assert _tiers(k, date(2026, 4, 1)) == [(D("100000"), D("4"), False, None)]   # text "100,000"
    assert _tiers(k, date(2026, 10, 1)) == [(D("120000"), D("4"), False, None)]
    ky = _scheme(parsed, "C-K1", "yearly")
    assert ky.threshold_unit == "points"
    assert [(t.threshold, t.rate) for t in ky.tiers] == [
        (D("450000"), D("1")), (D("500000"), D("1.5")), (D("550000"), D("2")),
    ]
    ly = _scheme(parsed, "C-L1", "yearly")
    assert [(t.threshold, t.rate) for t in ly.tiers] == [(D("200000"), D("1")), (D("250000"), D("2"))]


def test_parse_cell_grammar():
    from app.services.sales.dealer_target_import import parse_tier_cell

    assert [(t.threshold, t.rate) for t in parse_tier_cell("300,000 = 2%")] == [(D("300000"), D("2"))]
    assert [(t.threshold, t.rate) for t in parse_tier_cell("300k (2%); 400K (3%)")] == [
        (D("300000"), D("2")), (D("400000"), D("3")),
    ]
    extra = parse_tier_cell("350,000 + Additional 1%")
    assert [(t.threshold, t.rate, t.is_additional) for t in extra] == [(D("350000"), D("1"), True)]
    same_line = parse_tier_cell("300,000 - 3% + Additional 1%")
    assert [(t.threshold, t.rate, t.is_additional) for t in same_line] == [
        (D("300000"), D("3"), False), (D("300000"), D("1"), True),
    ]
    assert parse_tier_cell("-") == [] and parse_tier_cell("N/A") == [] and parse_tier_cell(None) == []
    bare = parse_tier_cell(300000, default_rate=D("4"))
    assert [(t.threshold, t.rate) for t in bare] == [(D("300000"), D("4"))]


# --------------------------------------------------------------------------------------- #
# the route
# --------------------------------------------------------------------------------------- #


NAMES = {
    "A1": "Dealer A Sdn Bhd", "A3": "Dealer A Sdn Bhd", "B_DELUXE": "Dealer B Trading",
    "B_MASTILE": "Dealer B Trading", "B_WPLY": "Dealer B Trading", "C1": "Dealer C Enterprise",
    "D1": "Dealer D Sdn Bhd", "E1": "Dealer E", "M1": "Monthly dealer M1", "M2": "Monthly dealer M2",
    "M3": "Monthly dealer M3", "K1": "Mocha K", "K2": "Mocha K [CERAMIC]", "K3": "Mocha K [IBORN]",
    "L1": "Mocha L",
}


def _seed(db, company_id):
    from app.models.order import Customer

    customers = {}
    for key in fx.KEYS:
        row = Customer(id=_uid(), company_id=company_id, customer_code=_code(), customer_name=NAMES[key])
        db.add(row)
        customers[key] = row
    db.flush()
    return customers


def _upload(client, data, *, dry_run, db=None):
    files = {"file": ("dealer targets.xlsx", data,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    # `SessionLocal` runs with autoflush off; so does the import here, as in production.
    if db is not None:
        db.autoflush = False
    try:
        return client.post(IMPORT, params={"dry_run": str(dry_run).lower(), "year": 2026}, files=files)
    finally:
        if db is not None:
            db.autoflush = True


def _dealer_target_count(db):
    from app.models.sales import SalesTarget

    return db.query(SalesTarget).filter(SalesTarget.subject_kind == "dealer").count()


def test_import_dry_run_writes_nothing_and_reports_every_row(api):
    client, db, company_id = api
    customers = _seed(db, company_id)
    data = fx.build({k: c.customer_code for k, c in customers.items()})
    before = _dealer_target_count(db)
    res = _upload(client, data, dry_run=True, db=db)
    assert res.status_code == 200, res.text
    body = res.json()
    assert _dealer_target_count(db) == before
    rows = body["rows"]
    assert len(rows) == 16
    by_code = {r["code"]: r for r in rows}
    unknown = by_code[fx.UNKNOWN_CODE]
    assert unknown["status"] == "not_representable"
    assert fx.UNKNOWN_CODE in unknown["reason"]
    assert by_code[customers["E1"].customer_code]["status"] == "not_representable"
    d1 = by_code[customers["D1"].customer_code]
    assert d1["status"] == "mapped"
    assert d1["customer"] == "Dealer D Sdn Bhd"
    assert any("name differs" in w for w in d1["warnings"])
    assert body["summary"]["mapped"] == 14
    assert body["summary"]["not_representable"] == 2
    assert body["summary"]["targets_new"] == 12
    assert body["summary"]["targets_updated"] == 0
    assert all(r["target_keys"] for r in rows if r["status"] == "mapped")


def test_import_every_pattern_applies_and_computes(api):
    client, db, company_id = api
    customers = _seed(db, company_id)
    pid = _product(db, company_id, _category(db, company_id).id).id

    def sale(key, amount, on):
        _so_line(db, company_id, customer_id=customers[key].id, order_date=on,
                 line_total=D(str(amount)), product_id=pid)

    sale("A1", "253450.50", date(2026, 2, 10))        # A Q1: P3
    sale("A1", 180000, date(2026, 5, 3))              # A Q2: P1/P2
    sale("A3", 70000, date(2026, 5, 3))
    sale("B_DELUXE", 150000, date(2026, 1, 15))        # B Q1: P4
    sale("B_MASTILE", 100000, date(2026, 1, 15))
    sale("B_WPLY", 60000, date(2026, 1, 15))
    sale("C1", 1300000, date(2025, 12, 15))           # C yearly, Dec'25 window: P6
    sale("M1", 60000, date(2026, 1, 20))              # M1 January: P7
    sale("K2", 80000, date(2026, 8, 1))               # K Q3 points: P8
    sale("K3", 50000, date(2026, 8, 2))
    ratio = client.post(f"{BASE}/point-ratios", json={"effective_from": "2026-01-01", "rm_per_point": 1.2})
    assert ratio.status_code in (200, 201), ratio.text

    data = fx.build({k: c.customer_code for k, c in customers.items()})
    res = _upload(client, data, dry_run=False, db=db)
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["summary"]["targets_new"] == 12
    by_code = {r["code"]: r for r in body["rows"]}
    assert all(r["target_nos"] for r in body["rows"] if r["status"] == "mapped")

    def detail(key, kind):
        rows = client.get(BASE, params={"subject": "dealer", "all": "true"}).json()["rows"]
        code = customers[key].customer_code
        target_no = next(
            no for no in by_code[code]["target_nos"]
            if next(r for r in rows if r["target_no"] == no)["name"].endswith(kind)
        )
        row = next(r for r in rows if r["target_no"] == target_no)
        return client.get(f"{BASE}/{row['target_id']}").json()

    def period(body, start):
        return next(p for p in body["periods"] if p["period_start"] == start)

    a_q = detail("A1", "quarterly")
    assert [a["code"] for a in a_q["accounts"]] == [customers["A1"].customer_code, customers["A3"].customer_code]
    assert a_q["commission_method"] == "retroactive"
    assert period(a_q, "2026-01-01")["commission_earned"] == 11405.28
    assert period(a_q, "2026-04-01")["commission_earned"] == 5400
    assert period(detail("B_DELUXE", "quarterly"), "2026-01-01")["commission_earned"] == 6900
    c_y = detail("C1", "yearly")
    assert (c_y["start_date"], c_y["end_date"]) == ("2025-12-01", "2026-11-30")
    assert c_y["periods"][0]["commission_earned"] == 13000
    assert period(detail("M1", "monthly"), "2026-01-01")["commission_earned"] == 600
    k_q = detail("K1", "quarterly")
    assert k_q["threshold_unit"] == "points"
    # 100,000 pts x 1.20 = RM 120,000 reached by 130,000; 4% on each account.
    assert period(k_q, "2026-07-01")["commission_earned"] == 5200


def test_import_reimport_updates_never_duplicates(api):
    client, db, company_id = api
    customers = _seed(db, company_id)
    data = fx.build({k: c.customer_code for k, c in customers.items()})
    first = _upload(client, data, dry_run=False, db=db)
    assert first.status_code == 200, first.text
    count = _dealer_target_count(db)
    again = _upload(client, data, dry_run=True, db=db)
    assert again.json()["summary"]["targets_new"] == 0
    assert again.json()["summary"]["targets_updated"] == 12
    second = _upload(client, data, dry_run=False, db=db)
    assert second.status_code == 200, second.text
    assert _dealer_target_count(db) == count


def test_import_rejects_a_file_that_is_not_a_workbook(api):
    client, db, company_id = api
    res = _upload(client, b"not a workbook", dry_run=True)
    assert res.status_code == 422


# --------------------------------------------------------------------------------------- #
# bounds (security review B1): a crafted file cannot make the reader walk a billion cells
# --------------------------------------------------------------------------------------- #


def _workbook(fill) -> bytes:
    import io

    from openpyxl import Workbook

    wb = Workbook()
    fill(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_far_cell_finishes_fast():
    import time

    from app.services.sales.dealer_target_import import parse_workbook

    def fill(ws):
        ws["A1"], ws["B1"], ws["C1"] = "CODE", "Q1", "NAME"
        ws["A2"], ws["B2"] = "C-1", "100,000 - 1%"
        ws["XFD1048576"] = "far away"

    started = time.monotonic()
    parsed = parse_workbook(_workbook(fill), year=2026)
    assert time.monotonic() - started < 5
    assert [r.code for r in parsed.rows] == ["C-1", None]  # the far cell is a row with no code
    assert parsed.rows[1].status == "not_representable"


def test_parse_refuses_an_overlong_cell_and_a_sheet_past_the_bounds():
    from app.services.sales import dealer_target_import as imp

    def long_cell(ws):
        ws["A1"], ws["B1"] = "CODE", "Q1"
        ws["A2"], ws["B2"] = "C-1", "1," * 2000 + "1%"

    parsed = imp.parse_workbook(_workbook(long_cell), year=2026)
    assert parsed.rows[0].status == "not_representable"
    assert "characters" in parsed.rows[0].reason

    def wide(ws):
        ws["A1"] = "CODE"
        for c in range(2, imp.MAX_SHEET_COLS + 10):
            ws.cell(row=1, column=c, value="Q1")

    parsed = imp.parse_workbook(_workbook(wide), year=2026)
    assert parsed.rows == []
    assert any("columns" in note for note in parsed.notes)


def test_import_code_of_another_company_is_not_representable(api):
    client, db, company_id = api
    from sqlalchemy import text

    from app.models.base import company_scope
    from app.models.order import Customer

    other = db.execute(
        text("select id from companies where id <> :c order by code limit 1"), {"c": company_id}
    ).scalar()
    with company_scope(db, frozenset({other})):
        theirs = Customer(id=_uid(), company_id=other, customer_code=_code(), customer_name="ZZT THEIRS")
        db.add(theirs)
        db.flush()

    def fill(ws):
        ws["A1"], ws["B1"] = "CODE", "Q1"
        ws["A2"], ws["B2"] = theirs.customer_code, "100,000 - 1%"

    body = _upload(client, _workbook(fill), dry_run=True, db=db).json()
    assert body["rows"][0]["status"] == "not_representable"
    assert "no customer with code" in body["rows"][0]["reason"]


# --------------------------------------------------------------------------------------- #
# review round (continuation): rates on one row, groups per scheme kind, code matching
# --------------------------------------------------------------------------------------- #


def test_parse_rate_on_one_row_of_a_merged_target_is_that_accounts_rate():
    """Ruling 2: the threshold merged over A/C III and A/C I, the rate typed on A/C I's row
    only, pays A/C I only; with A/C III listed first, the payout defaults to A/C I, the first
    account given a rate (not the first code)."""
    from app.services.sales.dealer_target_import import parse_workbook

    def fill(ws):
        # A threshold / % pair under one Q1 header (the monthly layout, P7).
        ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "CODE", "ACCOUNT", "Q1", "Q1"
        ws["C2"], ws["D2"] = "TARGET", "%"
        ws["A3"], ws["B3"] = "C-A3", "A/C III"
        ws["A4"], ws["B4"] = "C-A1", "A/C I"
        ws["C3"] = 200000
        ws.merge_cells("C3:C4")
        ws["D4"] = 3

    parsed = parse_workbook(_workbook(fill), year=2026)
    assert all(r.status == "mapped" for r in parsed.rows), [(r.code, r.reason) for r in parsed.rows]
    q = _scheme(parsed, "C-A3", "quarterly")
    assert q.codes == ["C-A3", "C-A1"]
    assert _tiers(q, date(2026, 1, 1)) == [(D("200000"), None, False, {"C-A1": D("3")})]
    assert q.payout_code == "C-A1"


def test_parse_merge_in_one_scheme_kind_does_not_join_the_other_kinds():
    """A yearly threshold merged over two accounts with a quarterly threshold per account is
    one yearly target over both and one quarterly target per account (rule 2)."""
    from app.services.sales.dealer_target_import import parse_workbook

    def fill(ws):
        for col, text in enumerate(["CODE", "ACCOUNT", "Q1", "JAN-DEC'26"], start=1):
            ws.cell(row=1, column=col, value=text)
        ws["A2"], ws["B2"], ws["C2"] = "C-X1", "A/C I", "100,000 - 2%"
        ws["A3"], ws["B3"], ws["C3"] = "C-X3", "A/C III", "80,000 - 2%"
        ws["D2"] = "500,000 - 1%"
        ws.merge_cells("D2:D3")

    parsed = parse_workbook(_workbook(fill), year=2026)
    assert all(r.status == "mapped" for r in parsed.rows), [(r.code, r.reason) for r in parsed.rows]
    quarterly = sorted((s.codes, _tiers(s, date(2026, 1, 1))) for s in parsed.schemes if s.kind == "quarterly")
    assert quarterly == [
        (["C-X1"], [(D("100000"), D("2"), False, None)]),
        (["C-X3"], [(D("80000"), D("2"), False, None)]),
    ]
    yearly = [s for s in parsed.schemes if s.kind == "yearly"]
    assert [s.codes for s in yearly] == [["C-X1", "C-X3"]]


def test_import_matches_codes_without_regard_to_case_or_spaces(api):
    client, db, company_id = api
    from app.models.order import Customer

    customer = Customer(id=_uid(), company_id=company_id, customer_code=_code(), customer_name="ZZT CASE")
    db.add(customer)
    db.flush()

    def fill(ws):
        ws["A1"], ws["B1"] = "CODE", "Q1"
        ws["A2"], ws["B2"] = f"  {customer.customer_code.lower()} ", "100,000 - 1%"

    body = _upload(client, _workbook(fill), dry_run=True, db=db).json()
    assert body["rows"][0]["status"] == "mapped", body["rows"][0]
    assert body["rows"][0]["customer"] == "ZZT CASE"


def test_import_pay_to_code_without_a_row_of_its_own_resolves(api):
    """E6: a target counting A/C III pays to A/C I, and A/C I has no row in the file."""
    client, db, company_id = api
    from app.models.order import Customer

    counted = Customer(id=_uid(), company_id=company_id, customer_code=_code(), customer_name="ZZT A3")
    payee = Customer(id=_uid(), company_id=company_id, customer_code=_code(), customer_name="ZZT A1")
    db.add_all([counted, payee])
    db.flush()

    def fill(ws):
        ws["A1"], ws["B1"], ws["C1"] = "CODE", "Q1", "PAY TO"
        ws["A2"], ws["B2"], ws["C2"] = counted.customer_code, "100,000 - 1%", payee.customer_code

    res = _upload(client, _workbook(fill), dry_run=False, db=db)
    assert res.status_code == 200, res.text
    row = res.json()["rows"][0]
    assert row["status"] == "mapped", row
    rows = client.get(BASE, params={"subject": "dealer", "all": "true"}).json()["rows"]
    mine = next(r for r in rows if r["target_no"] in row["target_nos"])
    assert mine["payout_label"] and "ZZT A1" in mine["payout_label"]
