"""Dealer targets import, second real-file gate (lane SALES-TARGET-DEALER, crew 16 Oct).

Crew ran the dry run on the real Sorento and Mocha files at a24581de: 0 rows mapped. The shapes
it found, copied here with masked names and shape-only figures:

1. No ACCOUNT column: the account is a SUFFIX of the dealer name cell ("<dealer> A/C I").
2. Mocha: "TIER n" next to a plain "INCENTIVE" column (no tier number on it).
3. Sorento yearly: Tier 1-3 per account, Tier 4-5 merged over the dealer's accounts (crew ruling
   (b), 16 Oct: the dealer target replaces each account's own target once it reaches a tier).
4. Columns must never be dropped silently: a "Monthly" block with tiers but no months, and
   "QUARTERLY (POINT)" over JAN-MAR .. OCT-DEC plus one INCENTIVE column.
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from .test_sales_targets_s1 import BASE, _code, _uid, api, world  # noqa: F401  (fixtures)

D = Decimal
S, M = "SRT 2026", "MCH 2026"


def _head(ws, row, col, text, across=1, down=1):
    ws.cell(row=row, column=col, value=text)
    if across > 1 or down > 1:
        ws.merge_cells(start_row=row, start_column=col, end_row=row + down - 1, end_column=col + across - 1)


def _put(ws, row, col, value, down=1):
    _head(ws, row, col, value, down=down)


def sorento(ws):
    """NO | DEALER NAME (account as suffix) | Q1/JAN-MAR/Tier 1, Tier 1 Incentive |
    Yearly/TIER 1..5 + INCENTIVE each | Monthly/Tier 1, Tier 1 Incentive, Tier 2, Tier 2 Incentive."""
    ws.title = S
    _head(ws, 1, 1, "NO", down=3)
    _head(ws, 1, 2, "DEALER NAME", down=3)
    _head(ws, 1, 3, "Q1", across=2)
    _head(ws, 2, 3, "JAN-MAR", across=2)
    ws.cell(row=3, column=3, value="Tier 1")
    ws.cell(row=3, column=4, value="Tier 1 Incentive")
    _head(ws, 1, 5, "Yearly", across=10)
    for n in range(1, 6):
        _head(ws, 2, 3 + 2 * n, f"TIER {n}", down=2)
        _head(ws, 2, 4 + 2 * n, "INCENTIVE", down=2)
    _head(ws, 1, 15, "Monthly", across=4)
    for i, text in enumerate(("Tier 1", "Tier 1 Incentive", "Tier 2", "Tier 2 Incentive")):
        _head(ws, 2, 15 + i, text, down=2)
    rows = {}
    # ALPHA: A/C I + A/C III share Q1 (merged); "A/C I- 3%" pays only A/C I.
    rows["ALPHA:A/C I"], rows["ALPHA:A/C III"] = 4, 5
    ws["A4"], ws["B4"] = 1, "ALPHA HARDWARE SDN BHD A/C I"
    ws["A5"], ws["B5"] = 2, "ALPHA HARDWARE SDN BHD A/C III"
    _put(ws, 4, 3, 400000, down=2)
    _put(ws, 4, 4, "A/C I- 3%", down=2)
    # BRAVO (one account): yearly tiers 1-2 in TIER n / INCENTIVE pairs; monthly tiers with no
    # months (every month the same).
    rows["BRAVO:A/C I"] = 6
    ws["A6"], ws["B6"] = 3, "BRAVO TRADING (A/C I)"
    ws.cell(row=6, column=5, value=1000000)
    ws.cell(row=6, column=6, value=0.01)
    ws.cell(row=6, column=7, value=1200000)
    ws.cell(row=6, column=8, value="A/C I- 1.5%")
    ws.cell(row=6, column=15, value=80000)
    ws.cell(row=6, column=16, value=0.01)
    ws.cell(row=6, column=17, value=100000)
    ws.cell(row=6, column=18, value=0.015)
    # CHARLIE: Tier 1-3 per account, Tier 4-5 merged over the 3 accounts (item 3).
    for i, label in enumerate(("A/C I", "A/C III", "A/C IV")):
        r = 7 + i
        rows[f"CHARLIE:{label}"] = r
        ws.cell(row=r, column=1, value=4 + i)
        ws.cell(row=r, column=2, value=f"CHARLIE ENTERPRISE {label}")
        for n in range(1, 4):
            ws.cell(row=r, column=3 + 2 * n, value=100000 * n + 10000 * i)
            ws.cell(row=r, column=4 + 2 * n, value=0.01)
        ws.cell(row=r, column=12, value=0.02 + 0.005 * i)
        ws.cell(row=r, column=14, value=0.03 + 0.005 * i)
    _put(ws, 7, 11, 2000000, down=3)
    _put(ws, 7, 13, 2500000, down=3)
    return rows


def mocha(ws):
    """NO | CUSTOMER NAME | QUARTERLY (POINT)/JAN-MAR..OCT-DEC + INCENTIVE |
    <title> RM (YEARLY)/TIER 1, INCENTIVE, TIER 2, INCENTIVE."""
    ws.title = M
    _head(ws, 1, 1, "NO", down=2)
    _head(ws, 1, 2, "CUSTOMER NAME", down=2)
    _head(ws, 1, 3, "QUARTERLY (POINT)", across=5)
    for i, text in enumerate(("JAN-MAR", "APR-JUN", "JUL-SEP", "OCT-DEC", "INCENTIVE")):
        ws.cell(row=2, column=3 + i, value=text)
    _head(ws, 1, 8, "MOCHA TARGET RM (YEARLY)", across=4)
    for i, text in enumerate(("TIER 1", "INCENTIVE", "TIER 2", "INCENTIVE")):
        ws.cell(row=2, column=8 + i, value=text)
    rows = {"KILO:A/C I": 3, "KILO:A/C II": 4}
    ws["A3"], ws["B3"] = 1, "KILO LIVING A/C I"
    ws["A4"], ws["B4"] = 2, "KILO LIVING A/C II"
    for i in range(4):
        _put(ws, 3, 3 + i, 100000 + 10000 * i, down=2)
    _put(ws, 3, 7, 0.04, down=2)
    _put(ws, 3, 8, 600000, down=2)
    _put(ws, 3, 9, 0.01, down=2)
    _put(ws, 3, 10, 750000, down=2)
    _put(ws, 3, 11, 0.02, down=2)
    return rows


def _book(*fills):
    from openpyxl import Workbook

    wb = Workbook()
    out = []
    for i, fill in enumerate(fills):
        ws = wb.active if i == 0 else wb.create_sheet()
        out.append(fill(ws))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), out


def _parse():
    from app.services.sales.dealer_target_import import parse_workbook

    data, (s_rows, m_rows) = _book(sorento, mocha)
    return parse_workbook(data, year=2026), s_rows, m_rows


def _k(sheet, row):
    return f"{sheet}!{row}"


def _row(parsed, sheet, row):
    return next(r for r in parsed.rows if (r.sheet, r.row) == (sheet, row))


def _schemes(parsed, key, kind):
    return [s for s in parsed.schemes if key in s.codes and s.kind == kind]


def _tiers(scheme, start):
    return [(t.threshold, t.rate, t.is_additional, t.account_rates) for t in scheme.tiers if t.period_start == start]


# 1. the account is a suffix of the name cell --------------------------------------------


def test_account_suffix_of_the_name_cell_is_the_account():
    parsed, s_rows, m_rows = _parse()
    a3 = _row(parsed, S, s_rows["ALPHA:A/C III"])
    assert (a3.name_in_file, a3.account_in_file) == ("ALPHA HARDWARE SDN BHD", "A/C III")
    b = _row(parsed, S, s_rows["BRAVO:A/C I"])
    assert (b.name_in_file, b.account_in_file) == ("BRAVO TRADING", "A/C I")
    k2 = _row(parsed, M, m_rows["KILO:A/C II"])
    assert (k2.name_in_file, k2.account_in_file) == ("KILO LIVING", "A/C II")


def test_a_rate_naming_a_suffix_account_pays_that_account():
    parsed, s_rows, _ = _parse()
    a1, a3 = _k(S, s_rows["ALPHA:A/C I"]), _k(S, s_rows["ALPHA:A/C III"])
    (q,) = _schemes(parsed, a1, "quarterly")
    assert q.codes == [a1, a3]
    assert _tiers(q, date(2026, 1, 1)) == [(D("400000"), None, False, {a1: D("3")})]
    (y,) = _schemes(parsed, _k(S, s_rows["BRAVO:A/C I"]), "yearly")
    b = _k(S, s_rows["BRAVO:A/C I"])
    assert [(t.threshold, t.rate, t.account_rates) for t in y.tiers] == [
        (D("1000000"), D("1"), None), (D("1200000"), None, {b: D("1.5")}),
    ]


def test_name_match_reads_the_suffix_on_both_sides(api):
    """The CRM may carry the account bracketed or as a bare suffix; both match."""
    from app.models.order import Customer
    from app.services.sales.dealer_target_import import customer_account_keys, dealer_key

    assert dealer_key("ALPHA HARDWARE SDN BHD A/C I") == dealer_key("Alpha Hardware Sdn Bhd - [A/C III]")
    assert customer_account_keys("ALPHA HARDWARE SDN BHD A/C III", None) == {"L3"}

    client, db, company_id = api
    for name, level in (("ALPHA HARDWARE SDN BHD - [A/C I]", 1), ("ALPHA HARDWARE SDN BHD A/C III", None)):
        db.add(Customer(id=_uid(), company_id=company_id, customer_code=_code(), customer_name=name, account_level=level))
    db.flush()
    data, _ = _book(sorento)
    files = {"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    db.autoflush = False
    try:
        body = client.post(f"{BASE}/import", params={"dry_run": "true", "year": 2026}, files=files).json()
    finally:
        db.autoflush = True
    by_row = {r["row"]: r for r in body["rows"]}
    assert by_row[4]["match"] == "auto" and by_row[4]["customer"] == "ALPHA HARDWARE SDN BHD - [A/C I]"
    assert by_row[5]["match"] == "auto" and by_row[5]["customer"] == "ALPHA HARDWARE SDN BHD A/C III"


# 2. TIER n next to a plain INCENTIVE --------------------------------------------------


def test_tier_n_pairs_with_the_plain_incentive_column_after_it():
    parsed, _, m_rows = _parse()
    k1 = _k(M, m_rows["KILO:A/C I"])
    (y,) = _schemes(parsed, k1, "yearly")
    assert y.threshold_unit == "amount"
    assert [(t.threshold, t.rate) for t in y.tiers] == [(D("600000"), D("1")), (D("750000"), D("2"))]


# 4. nothing dropped: Monthly with no months, QUARTERLY (POINT) over month ranges --------


def test_monthly_tiers_with_no_months_apply_to_every_month():
    parsed, s_rows, _ = _parse()
    (m,) = _schemes(parsed, _k(S, s_rows["BRAVO:A/C I"]), "monthly")
    assert (m.split_every, m.split_unit) == (1, "month")
    assert _tiers(m, None) == [(D("80000"), D("1"), False, None), (D("100000"), D("1.5"), False, None)]


def test_quarterly_point_over_month_ranges_with_one_incentive():
    parsed, _, m_rows = _parse()
    k1, k2 = _k(M, m_rows["KILO:A/C I"]), _k(M, m_rows["KILO:A/C II"])
    (q,) = _schemes(parsed, k1, "quarterly")
    assert q.codes == [k1, k2]
    assert q.threshold_unit == "points"
    for i, start in enumerate((date(2026, 1, 1), date(2026, 4, 1), date(2026, 7, 1), date(2026, 10, 1))):
        assert _tiers(q, start) == [(D(100000 + 10000 * i), D("4"), False, None)]


def test_a_period_looking_column_the_reader_cannot_place_fails_the_rows_holding_data():
    """Item 4: a column under period words that is neither Q, month nor yearly is reported on
    every row with a figure in it, never skipped."""
    from openpyxl import Workbook

    from app.services.sales.dealer_target_import import parse_workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"], ws["D1"] = "DEALER NAME", "Q1", "SPECIAL", "SPECIAL"
    ws["C2"], ws["D2"] = "TIER 1", "INCENTIVE"
    ws["B2"] = "Tier 1"
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "X A/C I", "100,000 - 1%", 50000, 0.01
    ws["A4"], ws["B4"] = "Y A/C I", "100,000 - 1%"
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_workbook(buf.getvalue(), year=2026)
    x, y = parsed.rows
    assert x.status == "not_representable" and "SPECIAL" in x.reason
    assert y.status == "mapped"




# 3. ruling (b): per-account Tier 1-3 targets replaced by the dealer's Tier 4-5 target --------


def test_per_account_tiers_and_merged_dealer_tiers_become_linked_targets():
    parsed, s_rows, _ = _parse()
    keys = [_k(S, s_rows[f"CHARLIE:{label}"]) for label in ("A/C I", "A/C III", "A/C IV")]
    for key in keys:
        assert next(r for r in parsed.rows if r.key == key).status == "mapped"
    yearly = [s for s in parsed.schemes if s.kind == "yearly" and set(s.codes) & set(keys)]
    dealer = next(s for s in yearly if s.codes == keys)
    assert [(t.threshold, t.rate, t.account_rates) for t in dealer.tiers] == [
        (D("2000000"), None, dict(zip(keys, (D("2"), D("2.5"), D("3"))))),
        (D("2500000"), None, dict(zip(keys, (D("3"), D("3.5"), D("4"))))),
    ]
    assert dealer.replaced_by is None
    for i, key in enumerate(keys):
        own = next(s for s in yearly if s.codes == [key])
        assert [(t.threshold, t.rate) for t in own.tiers] == [
            (D(100000 * n + 10000 * i), D("1")) for n in (1, 2, 3)
        ]
        assert own.replaced_by == dealer.key


def test_apply_links_each_account_target_to_its_dealer_target(api):
    from app.models.order import Customer

    client, db, company_id = api
    for label, level in (("A/C I", 1), ("A/C III", 3), ("A/C IV", 4)):
        db.add(Customer(id=_uid(), company_id=company_id, customer_code=_code(),
                        customer_name=f"CHARLIE ENTERPRISE - [{label}]", account_level=level))
    db.flush()
    data, _ = _book(sorento)
    files = {"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    db.autoflush = False
    try:
        body = client.post(f"{BASE}/import", params={"dry_run": "false", "year": 2026}, files=files).json()
    finally:
        db.autoflush = True
    charlie = [r for r in body["rows"] if r["row"] in (7, 8, 9)]
    assert all(r["status"] == "mapped" for r in charlie), charlie
    rows = client.get(BASE, params={"subject": "dealer", "all": "true"}).json()["rows"]
    yearly = [r for r in rows if r["name"].startswith("CHARLIE") and r["name"].endswith("yearly")]
    details = [client.get(f"{BASE}/{r['target_id']}").json() for r in yearly]
    dealer = next(d for d in details if len(d["accounts"]) == 3)
    owns = [d for d in details if len(d["accounts"]) == 1]
    assert len(owns) == 3
    assert all(d["replaced_by"]["target_no"] == dealer["target_no"] for d in owns)


def test_the_account_suffix_reader_is_linear():
    """Security review (gate 2) B1: the suffix regex never backtracks on a long run of spaces."""
    import time

    from app.services.sales.dealer_target_import import customer_account_keys, dealer_key, split_account_suffix

    bad = " " * 167 + "A/C 1" + " " * 328 + "x"
    started = time.monotonic()
    for _ in range(50):
        split_account_suffix(bad)
        dealer_key(bad[:255])
        customer_account_keys(bad[:255], None)
    assert time.monotonic() - started < 0.5
    assert split_account_suffix("X - A/C III ") == ("X", "A/C III")
    assert split_account_suffix("Y (A/C 2)") == ("Y", "A/C 2")



# review (gate 2): the split goes by MERGE, never by equal values -------------------------


def _charlie_equal(tiers):
    """Sorento with CHARLIE's per-account `tiers` (1-3) typed the same on every account row."""

    def fill(ws):
        rows = sorento(ws)
        for i in range(3):
            for n in tiers:
                ws.cell(row=7 + i, column=3 + 2 * n, value=100000 * n)
        return rows

    return fill


def _split(parsed, s_rows):
    keys = [_k(S, s_rows[f"CHARLIE:{label}"]) for label in ("A/C I", "A/C III", "A/C IV")]
    yearly = [s for s in parsed.schemes if s.kind == "yearly" and set(s.codes) & set(keys)]
    dealer = next(s for s in yearly if s.codes == keys)
    owns = [s for s in yearly if len(s.codes) == 1]
    return dealer, owns


def test_an_account_tier_equal_on_every_row_by_chance_stays_the_account_s():
    from app.services.sales.dealer_target_import import parse_workbook

    for tiers in ((1,), (1, 2, 3)):
        data, (s_rows,) = _book(_charlie_equal(tiers))
        dealer, owns = _split(parse_workbook(data, year=2026), s_rows)
        assert [t.threshold for t in dealer.tiers] == [D("2000000"), D("2500000")], tiers
        assert len(owns) == 3
        assert all(len(o.tiers) == 3 for o in owns)


def test_account_targets_with_another_window_than_their_dealer_fail():
    """Review should-fix 4: the replacement needs the dealer's period to cover the account's."""
    from openpyxl import Workbook

    from app.services.sales.dealer_target_import import parse_workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "DEALER NAME", "Yearly"
    ws.merge_cells("B1:F1")
    for col, text in zip("BCDEF", ("PERIOD", "TIER 1", "INCENTIVE", "TIER 4", "INCENTIVE")):
        ws[f"{col}2"] = text
    ws["A3"], ws["B3"], ws["C3"], ws["D3"], ws["F3"] = "X A/C I", "JAN-DEC'26", 100000, 0.01, 0.02
    ws["A4"], ws["B4"], ws["C4"], ws["D4"], ws["F4"] = "X A/C III", "DEC'25-NOV'26", 200000, 0.01, 0.02
    ws["E3"] = 1000000
    ws.merge_cells("E3:E4")
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_workbook(buf.getvalue(), year=2026)
    assert all(r.status == "not_representable" and "window" in r.reason for r in parsed.rows), [
        (r.status, r.reason) for r in parsed.rows
    ]


def test_apply_writes_a_monthly_block_with_no_months(api):
    """Review should-fix 1: tiers with no period of their own are every month's tiers."""
    from app.models.order import Customer

    client, db, company_id = api
    db.add(Customer(id=_uid(), company_id=company_id, customer_code=_code(),
                    customer_name="BRAVO TRADING - [A/C I]", account_level=1))
    db.flush()
    data, _ = _book(sorento)
    files = {"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    db.autoflush = False
    try:
        body = client.post(f"{BASE}/import", params={"dry_run": "false", "year": 2026}, files=files).json()
    finally:
        db.autoflush = True
    bravo = next(r for r in body["rows"] if r["row"] == 6)
    assert bravo["status"] == "mapped", bravo
    rows = client.get(BASE, params={"subject": "dealer", "all": "true"}).json()["rows"]
    monthly = next(r for r in rows if r["target_no"] in bravo["target_nos"] and r["name"].endswith("monthly"))
    detail = client.get(f"{BASE}/{monthly['target_id']}").json()
    assert detail["tiers_per_period"] is False
    assert [t["threshold"] for t in detail["tiers"]] == [80000, 100000]
