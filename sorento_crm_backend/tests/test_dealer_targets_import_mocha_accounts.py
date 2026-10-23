"""Dealer targets import: the Mocha account rule (lane SALES-TARGET-DEALER, owner 17 Oct ~21:00).

Owner, asked "Is A/C II CERAMIC and A/C III IBORN?": "yes for mocha". In Mocha's file A/C I is
the dealer's "- [A/C I]" ledger, A/C II its "- [CERAMIC]" ledger and A/C III its "- [IBORN]"
ledger. An explicit rule for that one company, keyed on its code; no name guessing. Every other
company keeps reading A/C II / A/C III as account levels only.
"""
from __future__ import annotations

import io

from .test_sales_targets_s1 import BASE, _code, _uid, api, world  # noqa: F401  (fixtures)


def _bytes(fill):
    from openpyxl import Workbook

    wb = Workbook()
    fill(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _customer(db, company_id, name, level):
    from app.models.order import Customer

    row = Customer(id=_uid(), company_id=company_id, customer_code=_code(), customer_name=name, account_level=level)
    db.add(row)
    db.flush()
    return row


def _post(client, db, data):
    files = {"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    db.autoflush = False
    try:
        return client.post(f"{BASE}/import", params={"dry_run": "true", "year": 2026}, files=files).json()
    finally:
        db.autoflush = True


def _as_mocha(db, company_id, monkeypatch):
    """The fixture company stands in for Mocha: the rule is keyed on the company's code."""
    from app.models.company import Company
    from app.services.sales import dealer_target_import

    code = db.query(Company.code).filter(Company.id == company_id).scalar()
    monkeypatch.setitem(dealer_target_import.ACCOUNT_LABEL_LEDGERS, code, dealer_target_import.ACCOUNT_LABEL_LEDGERS["MCH"])


def _trio(db, company_id):
    return (
        _customer(db, company_id, "MIKE TILES SDN BHD - [A/C I]", 1),
        _customer(db, company_id, "MIKE TILES SDN BHD - [CERAMIC]", None),
        _customer(db, company_id, "MIKE TILES SDN BHD - [IBORN]", None),
    )


def _dealer_file(ws):
    ws["A1"], ws["B1"], ws["C1"] = "DEALER NAME", "ACCOUNT", "Q1"
    ws["A2"] = "MIKE TILES SDN BHD"
    ws.merge_cells("A2:A4")
    ws["C2"] = "100,000 - 1%"
    ws.merge_cells("C2:C4")
    ws["B2"], ws["B3"], ws["B4"] = "A/C I", "A/C II", "A/C III"


def test_the_rule_is_the_owner_s_and_mocha_s_only():
    from app.services.sales.dealer_target_import import ACCOUNT_LABEL_LEDGERS

    assert ACCOUNT_LABEL_LEDGERS == {"MCH": {"A/C II": "CERAMIC", "A/C III": "IBORN"}}


def test_mocha_a_c_ii_and_iii_are_the_ceramic_and_iborn_ledgers(api, monkeypatch):
    client, db, company_id = api
    _as_mocha(db, company_id, monkeypatch)
    one, ceramic, iborn = _trio(db, company_id)
    result = _post(client, db, _bytes(_dealer_file))
    assert [(r["status"], r["match"], r["customer"]) for r in result["rows"]] == [
        ("mapped", "auto", one.customer_name),
        ("mapped", "auto", ceramic.customer_name),
        ("mapped", "auto", iborn.customer_name),
    ], [r["reason"] for r in result["rows"]]
    assert result["summary"]["targets_new"] == 1


def test_another_company_still_waits_on_a_pick_for_a_c_ii(api):
    client, db, company_id = api
    _trio(db, company_id)
    rows = _post(client, db, _bytes(_dealer_file))["rows"]
    assert rows[1]["status"] == "needs_pick"
    assert rows[1]["reason"].startswith('no ledger of this name has the account "A/C II"')


def test_the_rule_never_names_a_ledger_of_another_dealer(api, monkeypatch):
    """No name guessing: A/C II is the CERAMIC ledger OF THAT DEALER, or a pick."""
    client, db, company_id = api
    _as_mocha(db, company_id, monkeypatch)
    _customer(db, company_id, "MIKE TILES SDN BHD - [A/C I]", 1)
    _customer(db, company_id, "OSCAR STONE SDN BHD - [CERAMIC]", None)
    rows = _post(client, db, _bytes(_dealer_file))["rows"]
    assert rows[1]["status"] == "needs_pick", rows[1]
