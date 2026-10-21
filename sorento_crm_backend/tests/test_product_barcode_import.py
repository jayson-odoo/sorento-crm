"""BARCODE-IMPORT: bulk-set `products.barcode` from an uploaded sheet (preview, then apply).

UAC: documentation/plans/master-data/barcode-import-acceptance-criteria.md

Postgres only (tests/_pg_fixture.py); every test seeds its own products.
"""
from __future__ import annotations

import io
import uuid

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

from app.models.audit import AuditLog
from app.models.product import Product, ProductCategory, UnitOfMeasure
from tests._pg_fixture import blank_session, unique_code

BASE = "/api/v1/master-data/products/barcode-import"

# Checksum-valid EAN-13s (GS1 mod-10), and one 13-digit value that is not.
EAN_A = "9555123400011"
EAN_B = "9555123400028"
EAN_C = "9555123400035"
EAN_BAD = "9555123400014"


@pytest.fixture()
def db():
    with blank_session() as session:
        yield session


@pytest.fixture()
def refs(db):
    cat, uom = str(uuid.uuid4()), str(uuid.uuid4())
    db.add(ProductCategory(id=cat, category_code=unique_code("BC")[:50], category_name="ZZT"))
    db.add(UnitOfMeasure(id=uom, uom_code=unique_code("BU")[:20], uom_name="Each"))
    db.flush()
    return cat, uom


def _mk(db, refs, code, barcode=None) -> str:
    pid = str(uuid.uuid4())
    db.add(
        Product(
            id=pid, product_code=code, product_name=f"N {code}", category_id=refs[0],
            base_uom_id=refs[1], list_price=0, is_active=True, barcode=barcode,
        )
    )
    db.flush()
    return pid


def _barcode(db, pid):
    db.expire_all()
    return db.execute(text("SELECT barcode FROM products WHERE id = :i"), {"i": pid}).scalar()


def _xlsx(rows: list[list]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(rows: list[list]) -> bytes:
    return "\n".join(",".join(str(c) for c in r) for r in rows).encode()


@pytest.fixture()
def perm_client(db, monkeypatch):
    """A signed-in user whose permission set is the returned `allow` set."""
    from app.main import app
    from app.database import get_db
    from app.dependencies import get_current_user, get_current_user_or_api_key
    from app.services.company_scope_resolver import apply_company_scope
    from app.services.user_service import UserPermissionService

    allow: set = set()
    caller = {"id": str(uuid.uuid4()), "email": "barcode@example.com"}
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: caller
    app.dependency_overrides[get_current_user_or_api_key] = lambda: caller
    app.dependency_overrides[apply_company_scope] = lambda: None
    monkeypatch.setattr(
        UserPermissionService, "check_user_has_permission", lambda self, uid, slug: slug in allow
    )
    try:
        yield TestClient(app), allow
    finally:
        app.dependency_overrides.clear()


@pytest.fixture()
def client(perm_client):
    c, allow = perm_client
    allow.add("master_data.products.edit")
    return c


def _post(client, path, data: bytes, name="barcodes.xlsx"):
    return client.post(f"{BASE}/{path}", files={"file": (name, data, "application/octet-stream")})


def _apply(client, data: bytes, name="barcodes.xlsx"):
    """Preview, then apply with the token the preview returned (the only way apply runs)."""
    token = _post(client, "preview", data, name).json()["preview_token"]
    return client.post(
        f"{BASE}/apply",
        files={"file": (name, data, "application/octet-stream")},
        data={"preview_token": token},
    )


def _by_code(body):
    return {r["product_code"]: r for r in body["rows"]}


# --------------------------------------------------------------------------- service unit
@pytest.mark.parametrize("value, ok", [
    (EAN_A, True), (EAN_BAD, False), ("123", False), ("95551234000AB", False), ("", False),
    ("4006381333931", True),
])
def test_ean13_checksum(value, ok):
    from app.services.product_barcode_import_service import is_valid_ean13

    assert is_valid_ean13(value) is ok


# --------------------------------------------------------------------------- preview
def test_preview_every_status_and_writes_nothing(client, db, refs):
    stem = unique_code("BCI", alpha=True)
    new = _mk(db, refs, f"{stem}-NEW")
    chg = _mk(db, refs, f"{stem}-CHG", barcode=EAN_A)
    same = _mk(db, refs, f"{stem}-SAME", barcode=EAN_B)
    blank = _mk(db, refs, f"{stem}-BLANK", barcode=EAN_C)
    longp = _mk(db, refs, f"{stem}-LONG")
    bad = _mk(db, refs, f"{stem}-BAD")
    rows = [
        ["Product Code", "Barcode"],
        [f"{stem}-NEW", "9555123400042"],
        [f"{stem}-CHG", "9555123400059"],
        [f"{stem}-SAME", EAN_B],
        [f"{stem}-BLANK", ""],
        [f"{stem}-NOPE", "9555123400066"],
        [f"{stem}-LONG", "X" * 101],
        [f"{stem}-BAD", EAN_BAD],
    ]
    res = _post(client, "preview", _xlsx(rows))
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["readable"] is True
    assert body["code_column"] == "Product Code"
    assert body["barcode_column"] == "Barcode"
    by = _by_code(body)
    assert by[f"{stem}-NEW"]["status"] == "new"
    assert by[f"{stem}-CHG"]["status"] == "change"
    assert by[f"{stem}-CHG"]["old_barcode"] == EAN_A
    assert by[f"{stem}-CHG"]["new_barcode"] == "9555123400059"
    assert by[f"{stem}-SAME"]["status"] == "same"
    assert by[f"{stem}-BLANK"]["status"] == "blank_barcode"
    assert by[f"{stem}-NOPE"]["status"] == "unknown_code"
    assert by[f"{stem}-LONG"]["status"] == "too_long"
    assert by[f"{stem}-BAD"]["status"] == "new"
    assert "not_ean13" in by[f"{stem}-BAD"]["warnings"]
    assert by[f"{stem}-NEW"]["warnings"] == []
    c = body["counts"]
    assert (c["new"], c["change"], c["same"], c["blank_barcode"], c["unknown_code"],
            c["too_long"], c["not_ean13"]) == (2, 1, 1, 1, 1, 1, 1)
    assert body["total_rows"] == 7
    # Nothing written.
    assert _barcode(db, new) is None
    assert _barcode(db, chg) == EAN_A
    assert _barcode(db, blank) == EAN_C
    assert _barcode(db, longp) is None
    assert _barcode(db, bad) is None
    assert _barcode(db, same) == EAN_B


def test_preview_duplicates_warn_but_stay_writable(client, db, refs):
    stem = unique_code("BCD", alpha=True)
    _mk(db, refs, f"{stem}-HOLDER", barcode=EAN_A)
    _mk(db, refs, f"{stem}-P1")
    _mk(db, refs, f"{stem}-P2")
    _mk(db, refs, f"{stem}-P3")
    rows = [
        ["item code", "EAN"],
        [f"{stem}-P1", EAN_A],  # already on HOLDER
        [f"{stem}-P2", EAN_B],  # twice in the file
        [f"{stem}-P3", EAN_B],
    ]
    body = _post(client, "preview", _xlsx(rows)).json()
    by = _by_code(body)
    for code in ("P1", "P2", "P3"):
        row = by[f"{stem}-{code}"]
        assert row["status"] == "new"
        assert "duplicate" in row["warnings"], row
    assert body["counts"]["duplicate"] == 3
    assert body["code_column"] == "item code" and body["barcode_column"] == "EAN"


def test_preview_csv_headers_case_insensitive_numeric_and_code_case(client, db, refs):
    stem = unique_code("BCC", alpha=True)
    pid = _mk(db, refs, f"{stem}-ONE")
    data = _csv([["PRODUCT_CODE", "BAR CODE"], [f"  {stem.lower()}-one ", EAN_A]])
    body = _post(client, "preview", data, name="b.csv").json()
    assert body["readable"] is True
    assert body["rows"][0]["status"] == "new"
    assert body["rows"][0]["product_code"] == f"{stem}-ONE"
    assert _barcode(db, pid) is None


def test_preview_numeric_cell_is_not_float_text(client, db, refs):
    stem = unique_code("BCN", alpha=True)
    _mk(db, refs, f"{stem}-NUM")
    body = _post(client, "preview", _xlsx([["Code", "Barcode"], [f"{stem}-NUM", int(EAN_A)]])).json()
    assert body["rows"][0]["new_barcode"] == EAN_A


def test_preview_missing_column_is_unreadable(client):
    body = _post(client, "preview", _xlsx([["Product Code", "Name"], ["X", "Y"]])).json()
    assert body["readable"] is False
    assert "barcode" in body["missing_columns"]


def test_preview_refuses_other_file_types(client):
    res = _post(client, "preview", b"hello", name="b.txt")
    assert res.status_code == 400


# --------------------------------------------------------------------------- apply
def test_apply_writes_only_new_and_change(client, db, refs):
    stem = unique_code("BCA", alpha=True)
    new = _mk(db, refs, f"{stem}-NEW")
    chg = _mk(db, refs, f"{stem}-CHG", barcode=EAN_A)
    same = _mk(db, refs, f"{stem}-SAME", barcode=EAN_B)
    blank = _mk(db, refs, f"{stem}-BLANK", barcode=EAN_C)
    longp = _mk(db, refs, f"{stem}-LONG", barcode="KEEP")
    rows = [
        ["Product Code", "Barcode"],
        [f"{stem}-NEW", "9555123400042"],
        [f"{stem}-CHG", "9555123400059"],
        [f"{stem}-SAME", EAN_B],
        [f"{stem}-BLANK", ""],
        [f"{stem}-NOPE", "9555123400066"],
        [f"{stem}-LONG", "X" * 101],
    ]
    res = _apply(client, _xlsx(rows))
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["applied"] == 2
    assert body["counts"]["new"] == 1 and body["counts"]["change"] == 1
    assert _barcode(db, new) == "9555123400042"
    assert _barcode(db, chg) == "9555123400059"
    assert _barcode(db, same) == EAN_B
    assert _barcode(db, blank) == EAN_C  # blank never clears
    assert _barcode(db, longp) == "KEEP"


def test_apply_is_audited_per_product(client, db, refs):
    stem = unique_code("BCU", alpha=True)
    chg = _mk(db, refs, f"{stem}-CHG", barcode=EAN_A)
    same = _mk(db, refs, f"{stem}-SAME", barcode=EAN_B)
    res = _apply(client, _xlsx([
        ["Product Code", "Barcode"], [f"{stem}-CHG", EAN_C], [f"{stem}-SAME", EAN_B],
    ]))
    assert res.status_code == 200, res.text
    db.expire_all()
    logs = db.query(AuditLog).filter(AuditLog.entity_type == "product",
                                     AuditLog.entity_id == chg).all()
    hit = [r for r in logs if (r.new_values or {}).get("barcode") == EAN_C]
    assert hit, [(r.action, r.old_values, r.new_values) for r in logs]
    assert (hit[0].old_values or {}).get("barcode") == EAN_A
    assert hit[0].event == "product.barcode_import"
    # The Same row is not written, so it gains no import row (its seed INSERT row stays).
    assert not db.query(AuditLog).filter(AuditLog.entity_type == "product",
                                         AuditLog.entity_id == same,
                                         AuditLog.event == "product.barcode_import").all()


# --------------------------------------------------------------------------- permission
def test_both_endpoints_need_products_edit(perm_client, db, refs):
    c, allow = perm_client
    stem = unique_code("BCP", alpha=True)
    pid = _mk(db, refs, f"{stem}-X")
    data = _xlsx([["Product Code", "Barcode"], [f"{stem}-X", EAN_A]])
    allow.add("master_data.products.view")
    assert _post(c, "preview", data).status_code == 403
    assert _post(c, "apply", data).status_code == 403
    assert _barcode(db, pid) is None
    allow.add("master_data.products.edit")
    assert _post(c, "preview", data).status_code == 200


def test_repeated_code_keeps_first_row_only(client, db, refs):
    stem = unique_code("BCR", alpha=True)
    pid = _mk(db, refs, f"{stem}-R")
    body = _apply(client, _xlsx([
        ["Product Code", "Barcode"], [f"{stem}-R", EAN_A], [f"{stem}-r", EAN_B],
    ])).json()
    assert [r["status"] for r in body["rows"]] == ["new", "repeated_code"]
    assert body["rows"][1]["repeated_of_row"] == 2
    assert body["applied"] == 1
    assert _barcode(db, pid) == EAN_A


def test_title_lines_above_header_are_skipped(client, db, refs):
    stem = unique_code("BCT", alpha=True)
    _mk(db, refs, f"{stem}-T")
    body = _post(client, "preview", _xlsx([
        ["Barcode list 21 Oct"], [], ["Item Code", "Barcode"], [f"{stem}-T", EAN_A],
    ])).json()
    assert body["readable"] is True
    assert body["rows"][0]["row"] == 4 and body["rows"][0]["status"] == "new"


def test_blank_first_row_does_not_claim_the_code(client, db, refs):
    stem = unique_code("BCK", alpha=True)
    pid = _mk(db, refs, f"{stem}-K")
    body = _apply(client, _xlsx([
        ["Product Code", "Barcode"], [f"{stem}-K", ""], [f"{stem}-K", EAN_A],
    ])).json()
    assert [r["status"] for r in body["rows"]] == ["blank_barcode", "new"]
    assert _barcode(db, pid) == EAN_A


def test_code_matching_two_products_is_skipped(client, db, refs):
    # The unique index is (company_id, product_code), case-sensitive; the import matches
    # trimmed and case-insensitive, so `abc` and `ABC` in one company are both hits.
    stem = unique_code("BCM", alpha=True)
    upper = _mk(db, refs, f"{stem}-AMB")
    lower = _mk(db, refs, f"{stem}-amb")
    body = _apply(client, _xlsx([["Product Code", "Barcode"], [f"{stem}-AMB", EAN_A]])).json()
    assert body["rows"][0]["status"] == "ambiguous_code"
    assert body["applied"] == 0
    assert _barcode(db, upper) is None and _barcode(db, lower) is None


# --------------------------------------------------------------------------- crew audit round
def test_owner_sheet_headers(client, db, refs):
    """The owner's real header row: 'Barcode (EAN-13)' is the barcode, never 'Barcode Image'
    (pictures, empty cells), and 'Note' is carried onto the row."""
    stem = unique_code("BCH", alpha=True)
    _mk(db, refs, f"{stem}-OLD")
    _mk(db, refs, f"{stem}-NEW")
    body = _post(client, "preview", _xlsx([
        ["No.", "Product Code", "Barcode (EAN-13)", "Barcode Image", "Note"],
        [1, f"{stem}-OLD", EAN_A, None, None],
        [2, f"{stem}-NEW", EAN_A, None, f"Same barcode as {stem}-OLD"],
    ])).json()
    assert body["readable"] is True
    assert body["code_column"] == "Product Code"
    assert body["barcode_column"] == "Barcode (EAN-13)"
    assert body["note_column"] == "Note"
    by = _by_code(body)
    assert [by[f"{stem}-{c}"]["status"] for c in ("OLD", "NEW")] == ["new", "new"]
    assert by[f"{stem}-NEW"]["note"] == f"Same barcode as {stem}-OLD"
    assert by[f"{stem}-OLD"]["note"] is None
    # The intentional share is explained: each row names its partner.
    assert by[f"{stem}-OLD"]["duplicate_of"] == [f"{stem}-NEW"]
    assert by[f"{stem}-NEW"]["duplicate_of"] == [f"{stem}-OLD"]
    assert body["counts"]["duplicate"] == 2


def test_barcode_image_column_alone_is_not_a_barcode(client):
    body = _post(client, "preview", _xlsx([["Product Code", "Barcode Image"], ["X", None]])).json()
    assert body["readable"] is False
    assert body["missing_columns"] == ["barcode"]


def test_duplicate_names_the_product_already_holding_it(client, db, refs):
    stem = unique_code("BCG", alpha=True)
    _mk(db, refs, f"{stem}-HOLD", barcode=EAN_B)
    _mk(db, refs, f"{stem}-P")
    body = _post(client, "preview", _xlsx([["Product Code", "Barcode"], [f"{stem}-P", EAN_B]])).json()
    assert body["rows"][0]["duplicate_of"] == [f"{stem}-HOLD"]


def test_scientific_notation_text_is_skipped(client, db, refs):
    stem = unique_code("BCS", alpha=True)
    pid = _mk(db, refs, f"{stem}-S")
    data = _csv([["Product Code", "Barcode"], [f"{stem}-S", "9.55103E+12"]])
    body = _post(client, "preview", data, name="b.csv").json()
    assert body["rows"][0]["status"] == "scientific"
    assert body["counts"]["scientific"] == 1
    out = _apply(client, data, name="b.csv").json()
    assert out["applied"] == 0
    assert _barcode(db, pid) is None


def test_header_on_a_later_sheet(client, db, refs):
    import openpyxl

    stem = unique_code("BCW", alpha=True)
    _mk(db, refs, f"{stem}-W")
    wb = openpyxl.Workbook()
    wb.active.append(["Read me first"])
    ws = wb.create_sheet("Barcodes")
    ws.append(["Product Code", "Barcode"])
    ws.append([f"{stem}-W", EAN_A])
    buf = io.BytesIO()
    wb.save(buf)
    body = _post(client, "preview", buf.getvalue()).json()
    assert body["readable"] is True and body["rows"][0]["status"] == "new"


def test_apply_refuses_without_a_preview_of_the_same_file(client, db, refs):
    stem = unique_code("BCZ", alpha=True)
    pid = _mk(db, refs, f"{stem}-Z")
    data = _xlsx([["Product Code", "Barcode"], [f"{stem}-Z", EAN_A]])
    other = _xlsx([["Product Code", "Barcode"], [f"{stem}-Z", EAN_B]])
    files = lambda d: {"file": ("b.xlsx", d, "application/octet-stream")}  # noqa: E731

    assert client.post(f"{BASE}/apply", files=files(data)).status_code == 409
    token_for_other = _post(client, "preview", other).json()["preview_token"]
    res = client.post(f"{BASE}/apply", files=files(data), data={"preview_token": token_for_other})
    assert res.status_code == 409
    assert _barcode(db, pid) is None

    token = _post(client, "preview", data).json()["preview_token"]
    res = client.post(f"{BASE}/apply", files=files(data), data={"preview_token": token})
    assert res.status_code == 200 and res.json()["applied"] == 1
    assert _barcode(db, pid) == EAN_A


def test_preview_names_the_active_company(client, db, refs):
    """Owner ruling 21 Oct: the file is run once per company, so the preview names it."""
    from app.models.base import set_company_scope
    from app.models.company import Company

    cid = str(uuid.uuid4())
    db.add(Company(id=cid, name="ZZT Mocha Test Sdn Bhd", code=unique_code("C")[:50]))
    db.flush()
    set_company_scope(db, frozenset({cid}))
    body = _post(client, "preview", _xlsx([["Product Code", "Barcode"], ["NOPE", EAN_A]])).json()
    assert body["companies"] == ["ZZT Mocha Test Sdn Bhd"]


# --------------------------------------------------------------------------- caps, scope, token
def _xlsx_with_row(row_no: int, rows: list[list]) -> bytes:
    """A normal workbook whose last row/cell refs are rewritten to `row_no` (past Excel's max)."""
    import re
    import zipfile

    src = zipfile.ZipFile(io.BytesIO(_xlsx(rows)))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for item in src.infolist():
            blob = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                xml = blob.decode()
                last = len(rows)
                xml = re.sub(rf'<row r="{last}"', f'<row r="{row_no}"', xml)
                xml = re.sub(rf'<c r="([A-Z]+){last}"', rf'<c r="\g<1>{row_no}"', xml)
                xml = re.sub(r'<dimension ref="A1:([A-Z]+)\d+"', rf'<dimension ref="A1:\g<1>{row_no}"', xml)
                blob = xml.encode()
            zf.writestr(item, blob)
    return out.getvalue()


def test_sheet_row_past_excel_max_is_refused_quickly(client):
    import time

    data = _xlsx_with_row(1048577, [["Product Code", "Barcode"], ["X", EAN_A]])
    start = time.monotonic()
    res = _post(client, "preview", data)
    assert res.status_code == 400, res.text
    assert time.monotonic() - start < 10


def test_csv_over_50000_rows_is_refused(client):
    rows = [["Product Code", "Barcode"]] + [[f"C{i}", EAN_A] for i in range(50001)]
    res = _post(client, "preview", _csv(rows), name="b.csv")
    assert res.status_code == 400, res.text
    assert "50,000" in res.json()["detail"]


def test_uncompressed_size_cap_refuses_a_workbook(client, monkeypatch):
    from app.services import product_barcode_import_service as svc

    monkeypatch.setattr(svc, "_MAX_UNCOMPRESSED_BYTES", 10)
    res = _post(client, "preview", _xlsx([["Product Code", "Barcode"], ["X", EAN_A]]))
    assert res.status_code == 400, res.text


def _company(db, name):
    from app.models.company import Company

    cid = str(uuid.uuid4())
    db.add(Company(id=cid, name=name, code=unique_code("C")[:50]))
    db.flush()
    return cid


def test_token_is_bound_to_the_company_scope(client, db, refs):
    from app.models.base import set_company_scope

    a, b = _company(db, "ZZT A"), _company(db, "ZZT B")
    stem = unique_code("BCQ", alpha=True)
    _mk(db, refs, f"{stem}-Q")
    data = _xlsx([["Product Code", "Barcode"], [f"{stem}-Q", EAN_A]])
    set_company_scope(db, frozenset({a}))
    token = _post(client, "preview", data).json()["preview_token"]
    set_company_scope(db, frozenset({b}))
    res = client.post(f"{BASE}/apply", files={"file": ("b.xlsx", data, "application/octet-stream")},
                      data={"preview_token": token})
    assert res.status_code == 409, res.text


def test_token_from_another_user_is_refused(client, db, refs):
    from app.services import product_barcode_import_service as svc

    stem = unique_code("BCO", alpha=True)
    pid = _mk(db, refs, f"{stem}-O")
    data = _xlsx([["Product Code", "Barcode"], [f"{stem}-O", EAN_A]])
    token = svc.preview_token(data, str(uuid.uuid4()), db)
    res = client.post(f"{BASE}/apply", files={"file": ("b.xlsx", data, "application/octet-stream")},
                      data={"preview_token": token})
    assert res.status_code == 409, res.text
    assert _barcode(db, pid) is None


def test_same_code_in_two_companies_writes_only_the_scoped_one(client, db, refs):
    from app.models.base import set_company_scope

    a, b = _company(db, "ZZT A"), _company(db, "ZZT B")
    code = unique_code("BCX", alpha=True)
    ids = {}
    for key, cid in (("a", a), ("b", b)):
        pid = str(uuid.uuid4())
        db.add(Product(id=pid, product_code=code, product_name=f"N {key}", category_id=refs[0],
                       base_uom_id=refs[1], list_price=0, is_active=True, company_id=cid,
                       barcode=EAN_B if key == "b" else None))
        ids[key] = pid
    db.flush()
    set_company_scope(db, frozenset({a}))
    data = _xlsx([["Product Code", "Barcode"], [code, EAN_A]])
    body = _post(client, "preview", data).json()
    assert body["rows"][0]["status"] == "new"
    assert body["rows"][0]["warnings"] == []
    res = _apply(client, data)
    assert res.status_code == 200 and res.json()["applied"] == 1, res.text
    assert _barcode(db, ids["a"]) == EAN_A
    assert _barcode(db, ids["b"]) == EAN_B


def test_non_ascii_preview_token_is_409_not_500(client, db, refs):
    stem = unique_code("BCY", alpha=True)
    _mk(db, refs, f"{stem}-Y")
    data = _xlsx([["Product Code", "Barcode"], [f"{stem}-Y", EAN_A]])
    res = client.post(f"{BASE}/apply", files={"file": ("b.xlsx", data, "application/octet-stream")},
                      data={"preview_token": "toké"})
    assert res.status_code == 409, res.text


def test_csv_field_over_the_csv_limit_is_400(client):
    data = b"Product Code,Barcode\n" + b"A," + b"9" * 140000 + b"\n"
    res = _post(client, "preview", data, name="b.csv")
    assert res.status_code == 400, res.text


def test_sheet_that_fails_while_rows_are_read_is_400(client, monkeypatch):
    import openpyxl.worksheet._read_only as ro

    def boom(self, *a, **k):
        raise ValueError("broken sheet")
        yield  # pragma: no cover

    monkeypatch.setattr(ro.ReadOnlyWorksheet, "_cells_by_row", boom)
    res = _post(client, "preview", _xlsx([["Product Code", "Barcode"], ["X", EAN_A]]))
    assert res.status_code == 400, res.text


# --------------------------------------------------------------------------- round 3 caps
def _rewrite_zip(data: bytes, edit, extra: dict | None = None) -> bytes:
    import zipfile

    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for item in src.infolist():
            zf.writestr(item, edit(item.filename, src.read(item.filename)))
        for name, blob in (extra or {}).items():
            zf.writestr(name, blob)
    return out.getvalue()


def test_a_huge_cell_is_cut_to_256_chars_in_the_response(client, db, refs):
    stem = unique_code("BCL", alpha=True)
    _mk(db, refs, f"{stem}-L")
    big = "9" * 100_000
    data = _xlsx([
        ["Product Code", "Barcode", "Note"],
        [f"{stem}-L", big, big],
        [big, big, big],
    ])
    res = _post(client, "preview", data)
    assert res.status_code == 200, res.text
    rows = res.json()["rows"]
    assert rows[0]["status"] == "too_long"
    for r in rows:
        for key in ("product_code", "new_barcode", "note"):
            assert len(r[key] or "") <= 256, (key, len(r[key]))


def test_media_members_do_not_count_toward_the_xml_cap(client, monkeypatch):
    from app.services import product_barcode_import_service as svc

    monkeypatch.setattr(svc, "_MAX_UNCOMPRESSED_BYTES", 1024 * 1024)
    data = _rewrite_zip(_xlsx([["Product Code", "Barcode"], ["X", EAN_A]]), lambda n, b: b,
                        extra={"xl/media/image1.png": b"\x89PNG\r\n\x1a\n" + b"\0" * (2 * 1024 * 1024)})
    assert _post(client, "preview", data).status_code == 200


def test_non_media_members_over_the_cap_are_refused(client, monkeypatch):
    from app.services import product_barcode_import_service as svc

    monkeypatch.setattr(svc, "_MAX_UNCOMPRESSED_BYTES", 1024 * 1024)
    data = _rewrite_zip(_xlsx([["Product Code", "Barcode"], ["X", EAN_A]]), lambda n, b: b,
                        extra={"xl/customXml/big.xml": b"\0" * (2 * 1024 * 1024)})
    assert _post(client, "preview", data).status_code == 400


def test_header_past_column_64_is_not_found(client):
    row = [None] * 64 + ["Product Code", "Barcode"]
    body = _post(client, "preview", _xlsx([row, [None] * 64 + ["X", EAN_A]])).json()
    assert body["readable"] is False
    csv_row = ",".join([""] * 64 + ["Product Code", "Barcode"]) + "\n" + ",".join([""] * 64 + ["X", EAN_A])
    body = _post(client, "preview", csv_row.encode(), name="b.csv").json()
    assert body["readable"] is False


def _sheets_with_header_on(n: int) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    for i in range(1, 12):
        ws = wb.active if i == 1 else wb.create_sheet(f"S{i}")
        ws.append(["Product Code", "Barcode"] if i == n else ["Read me"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_header_is_searched_on_ten_sheets_only(client):
    assert _post(client, "preview", _sheets_with_header_on(10)).json()["readable"] is True
    assert _post(client, "preview", _sheets_with_header_on(11)).json()["readable"] is False


def test_wide_dimension_with_empty_rows_returns_quickly(client):
    import re
    import time

    def edit(name, blob):
        if name != "xl/worksheets/sheet1.xml":
            return blob
        xml = blob.decode()
        xml = re.sub(r'<dimension ref="[^"]*"', '<dimension ref="A1:ZZZ1048576"', xml)
        return xml.replace("</sheetData>", "<row/>" * 100000 + "</sheetData>").encode()

    data = _rewrite_zip(_xlsx([["Product Code", "Barcode"], ["X", EAN_A]]), edit)
    start = time.monotonic()
    _post(client, "preview", data)
    assert time.monotonic() - start < 5


def test_other_company_holding_the_same_barcode_is_not_a_duplicate(client, db, refs):
    from app.models.base import set_company_scope

    a, b = _company(db, "ZZT A"), _company(db, "ZZT B")
    stem = unique_code("BCJ", alpha=True)
    for code, cid, bc in ((f"{stem}-A", a, None), (f"{stem}-B", b, EAN_A)):
        db.add(Product(id=str(uuid.uuid4()), product_code=code, product_name=code,
                       category_id=refs[0], base_uom_id=refs[1], list_price=0, is_active=True,
                       company_id=cid, barcode=bc))
    db.flush()
    set_company_scope(db, frozenset({a}))
    body = _post(client, "preview", _xlsx([["Product Code", "Barcode"], [f"{stem}-A", EAN_A]])).json()
    assert body["rows"][0]["status"] == "new"
    assert body["rows"][0]["warnings"] == []
    assert body["rows"][0]["duplicate_of"] == []


def test_a_lying_small_dimension_does_not_drop_data_rows(client, db, refs):
    import re

    stem = unique_code("BCV", alpha=True)
    _mk(db, refs, f"{stem}-V")

    def edit(name, blob):
        if name != "xl/worksheets/sheet1.xml":
            return blob
        xml = blob.decode()
        xml = re.sub(r'<dimension ref="[^"]*"', '<dimension ref="A1:B3"', xml)
        return xml.encode()

    rows = [["Product Code", "Barcode"]] + [[] for _ in range(8)] + [[f"{stem}-V", EAN_A]]
    data = _rewrite_zip(_xlsx(rows), edit)
    body = _post(client, "preview", data).json()
    assert [r["status"] for r in body["rows"]] == ["new"]


def test_a_non_picture_member_under_xl_media_counts(client, monkeypatch):
    from app.services import product_barcode_import_service as svc

    monkeypatch.setattr(svc, "_MAX_UNCOMPRESSED_BYTES", 1024 * 1024)
    data = _rewrite_zip(_xlsx([["Product Code", "Barcode"], ["X", EAN_A]]), lambda n, b: b,
                        extra={"xl/media/image1.png": b"<xml>" + b"\0" * (2 * 1024 * 1024)})
    assert _post(client, "preview", data).status_code == 400


def test_a_sheet_disguised_under_xl_media_counts(client, monkeypatch):
    from app.services import product_barcode_import_service as svc

    monkeypatch.setattr(svc, "_MAX_UNCOMPRESSED_BYTES", 1024 * 1024)

    def edit(name, blob):
        if name == "xl/worksheets/sheet1.xml":
            return blob.replace(b"<sheetData>", b"<!--" + b"x" * (2 * 1024 * 1024) + b"--><sheetData>")
        return blob.replace(b"worksheets/sheet1.xml", b"media/sheet1.xml")

    import zipfile

    src = zipfile.ZipFile(io.BytesIO(_xlsx([["Product Code", "Barcode"], ["X", EAN_A]])))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
        for item in src.infolist():
            new_name = "xl/media/sheet1.xml" if item.filename == "xl/worksheets/sheet1.xml" else item.filename
            zf.writestr(new_name, edit(item.filename, src.read(item.filename)))
    assert _post(client, "preview", out.getvalue()).status_code == 400


def test_leading_spaces_are_trimmed_before_the_cell_is_cut(client, db, refs):
    stem = unique_code("BCW2", alpha=True)
    _mk(db, refs, f"{stem}-W")
    body = _post(client, "preview", _xlsx([
        ["Product Code", "Barcode"], [" " * 300 + f"{stem}-W", EAN_A],
    ])).json()
    assert body["rows"][0]["status"] == "new"
