"""S2 / AC-2.1 to AC-2.4 - the order inquiry search box splits on whitespace and ANDs.

TEST-FIRST (`PLAN-scm-ui-feedback-14sep.md`, J2). At the time this file is written
`order_inquiry_worklist_service._base` wraps the WHOLE query string in one `%phrase%` and
ORs it across eleven columns, so `SO366990 SRTWT6801` is looked for as a single literal and
matches nothing. Every test asserting the AND is therefore expected to be RED until the
token split lands.

The journey is CS's, and it is the only reason the box exists in two words rather than one:
an order has forty lines and a product sits on twenty orders, so either word alone is the
wrong answer and only their intersection is the row being asked about. What each word may
hit is unchanged - any of the eleven searchable columns (OR within a token), all of the
tokens (AND across them).

Seeded as ADOPTED orders (no project registration), which is the cheaper of the two shapes
and the one that carries a real `so_number` off the core sales order. Every FK target is
this file's own: CI's database is empty and the local one is a prod copy, so borrowing a row
would pass here and fail there.
"""
from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import openpyxl
import pytest

from app.models.company import Company  # noqa: F401  (FK target, imported for clarity)
from app.models.order import SalesOrder
from app.models.project_so import (
    INQUIRY_RAISED,
    ProjectSalesOrder,
    ProjectSalesOrderLine,
)

from ._pg_fixture import blank_session
from .test_order_inquiry_worklist import (
    LIST,
    MARKER,
    _client,
    _customer,
    _inquiry_for,
    _product,
    _restore,
    _row,
    _sorento,
    _uid,
    _user,
)

READ_ONLY = ["projects.projects.view"]

SUMMARY = f"{LIST}/summary"
EXPORT = f"{LIST}/export"

#: The two item codes, one shared between both orders and one on the first order alone.
#: Deliberately NOT a prefix of each other, so a match can only be the token itself.
ALPHA = "ZZTTOK-SRTWT6801"
BETA = "ZZTTOK-SRTWC9042"


def _adopted_order(db, company_id: str, so_number: str, customer_name: str):
    """One AutoCount-book order: a core sales order, its customer, and the adopted
    `ProjectSalesOrder` the worklist reads `so_number` off."""
    customer = _customer(db, company_id, customer_name)
    core = SalesOrder(
        id=_uid(),
        company_id=company_id,
        so_number=so_number,
        customer_id=customer.id,
        order_date=date(2026, 1, 8),
    )
    db.add(core)
    db.flush()
    adopted = ProjectSalesOrder(
        id=_uid(),
        company_id=company_id,
        project_id=None,
        so_id=core.id,
        provisional_ref=so_number,
        autocount_doc_no=so_number,
        status="adopted",
    )
    db.add(adopted)
    db.flush()
    return adopted


def _line(db, company_id: str, order, product, line_no: int):
    line = ProjectSalesOrderLine(
        id=_uid(),
        company_id=company_id,
        project_sales_order_id=order.id,
        line_no=line_no,
        product_id=product.id,
        description=f"{MARKER} {product.product_code}",
        qty=Decimal("10"),
        uom="UNIT",
        unit_price=Decimal("10.00"),
        amount=Decimal("100.00"),
        delivery_date=date(2026, 3, 2),
    )
    db.add(line)
    db.flush()
    return line


def _seed_two_orders(db, company_id: str, user_id: str) -> dict:
    """Two orders over an OVERLAPPING catalogue.

    * SO-A carries ALPHA and BETA;
    * SO-B carries ALPHA.

    So `SO-A` alone answers two rows, `ALPHA` alone answers two rows, and only the pair of
    them answers the one row the question was about.
    """
    alpha = _product(db, ALPHA, f"{MARKER} Wall hung WC")
    beta = _product(db, BETA, f"{MARKER} Close coupled WC")

    so_a_number = "ZZTTOKSOA366990"
    so_b_number = "ZZTTOKSOB411221"
    order_a = _adopted_order(db, company_id, so_a_number, f"{MARKER} Optad Sdn Bhd")
    order_b = _adopted_order(db, company_id, so_b_number, f"{MARKER} Bintang Sdn Bhd")

    inquiry_a = _inquiry_for(db, company_id, order_a)
    inquiry_b = _inquiry_for(db, company_id, order_b)
    # The raiser, so the email column has a real principal to keep its PREFIX match against.
    inquiry_a.raised_by = user_id
    inquiry_b.raised_by = user_id
    db.flush()

    a_alpha = _row(
        db,
        company_id,
        inquiry_a,
        so_line_id=_line(db, company_id, order_a, alpha, 1).id,
        item_code=ALPHA,
        qty="11",
        delivery_date=date(2026, 3, 2),
        state=INQUIRY_RAISED,
    )
    a_beta = _row(
        db,
        company_id,
        inquiry_a,
        so_line_id=_line(db, company_id, order_a, beta, 2).id,
        item_code=BETA,
        qty="12",
        delivery_date=date(2026, 3, 2),
        state=INQUIRY_RAISED,
    )
    b_alpha = _row(
        db,
        company_id,
        inquiry_b,
        so_line_id=_line(db, company_id, order_b, alpha, 1).id,
        item_code=ALPHA,
        qty="13",
        delivery_date=date(2026, 3, 2),
        state=INQUIRY_RAISED,
    )
    db.commit()
    return {
        "so_a": so_a_number,
        "so_b": so_b_number,
        "a_alpha": a_alpha,
        "a_beta": a_beta,
        "b_alpha": b_alpha,
    }


@pytest.fixture()
def api():
    from app.models.base import company_scope

    with blank_session() as db:
        company_id = _sorento(db)
        user_id = _user(db, f"{MARKER} Cindy Tan")
        seeded = _seed_two_orders(db, company_id, user_id)
        client, originals = _client(db, user_id, READ_ONLY)
        try:
            with company_scope(db, frozenset({company_id})):
                yield client, db, user_id, seeded
        finally:
            _restore(originals)


def _ids(client, query: str | None) -> set[str]:
    params = {"limit": 100}
    if query is not None:
        params["query"] = query
    body = client.get(LIST, params=params).json()
    return {row["id"] for row in body["data"]}


# ------------------------------------------------------------------ AC-2.1


def test_two_tokens_answer_only_their_intersection(api):
    """`SO366990 SRTWT6801` is the order AND the product, not the literal string.

    Today the phrase is looked for verbatim across every column, so this returns nothing at
    all - the failure the buyer reports as "the search is broken".
    """
    client, _db, _user_id, seeded = api

    found = _ids(client, f"{seeded['so_a']} {ALPHA}")

    assert found == {seeded["a_alpha"].id}


def test_the_order_of_the_tokens_does_not_matter(api):
    """AND is commutative, and a person types whichever word they remember first."""
    client, _db, _user_id, seeded = api

    assert _ids(client, f"{ALPHA} {seeded['so_a']}") == {seeded["a_alpha"].id}
    assert _ids(client, f"{seeded['so_a']} {ALPHA}") == {seeded["a_alpha"].id}


def test_a_token_may_hit_any_column_so_customer_plus_item_narrows_too(api):
    """The OR inside a token is the whole eleven-column set, unchanged: `Bintang` is a
    customer and `ALPHA` is an item code, and together they name one row."""
    client, _db, _user_id, seeded = api

    assert _ids(client, f"Bintang {ALPHA}") == {seeded["b_alpha"].id}


# ------------------------------------------------------------------ AC-2.2


def test_a_single_token_behaves_exactly_as_it_does_today(api):
    """One word is one filter - the split must not change what a one-word query answers."""
    client, _db, _user_id, seeded = api

    assert _ids(client, ALPHA) == {seeded["a_alpha"].id, seeded["b_alpha"].id}
    assert _ids(client, seeded["so_a"]) == {seeded["a_alpha"].id, seeded["a_beta"].id}
    assert _ids(client, "Wall hung WC".split()[0]) >= {seeded["a_alpha"].id}


def test_the_raiser_email_keeps_its_prefix_match_per_token(api):
    """The email column matches the FRONT of the address and nothing else, or a buyer
    typing a common word gets every row raised by anyone at that domain. That rule belongs
    to the TOKEN now, not to the whole box.
    """
    client, _db, user_id, seeded = api

    # `_user` builds the address as `<uuid>@zzt.test`, so the id is its prefix.
    assert _ids(client, user_id[:8]) == {
        seeded["a_alpha"].id,
        seeded["a_beta"].id,
        seeded["b_alpha"].id,
    }
    # Inside the address, never the front: "zzt.test" matches nothing through the email.
    assert _ids(client, "zzt.test") == set()
    # And the prefix still narrows when it is one token of two.
    assert _ids(client, f"{user_id[:8]} {ALPHA}") == {
        seeded["a_alpha"].id,
        seeded["b_alpha"].id,
    }


# ------------------------------------------------------------------ AC-2.3


def test_padding_and_repeated_spaces_are_ignored(api):
    """A pasted order number arrives with whitespace around it; an empty token between two
    real ones would AND against nothing and empty the list."""
    client, _db, _user_id, seeded = api

    assert _ids(client, f"   {seeded['so_a']}    {ALPHA}  ") == {seeded["a_alpha"].id}


def test_an_all_space_query_filters_nothing(api):
    """No tokens is no filter, the same answer an empty box gives - never zero rows."""
    client, _db, _user_id, seeded = api

    everything = _ids(client, None)
    assert len(everything) == 3

    assert _ids(client, "   ") == everything
    assert _ids(client, "") == everything


# ------------------------------------------------------------------ the box has limits
# (security review, round 1): a filter per token is a JOIN-heavy OR per token, and the box
# is a public-facing string. Both halves are pinned here.


def test_only_the_first_ten_tokens_are_applied(api):
    """A pasted paragraph must not turn into a hundred ILIKE filters over eleven columns.

    The tenth token still narrows; the eleventh is dropped rather than refused, because a
    search box that answers 422 to a clumsy paste is a worse answer than a search result.
    """
    client, _db, _user_id, seeded = api

    # Nine tokens that match everything (the marker every seeded row carries), then the
    # order number: ten in all, and the tenth still does its work.
    nine = " ".join([MARKER] * 9)
    assert _ids(client, f"{nine} {seeded['so_a']}") == {
        seeded["a_alpha"].id,
        seeded["a_beta"].id,
    }

    # The same query with the item code as an ELEVENTH token: it falls off the end, so the
    # answer is still the whole order rather than the one line.
    assert _ids(client, f"{nine} {seeded['so_a']} {ALPHA}") == {
        seeded["a_alpha"].id,
        seeded["a_beta"].id,
    }


def test_an_overlong_query_is_refused(api):
    """The cap above is on tokens; this is the cap on the string itself, at the route."""
    client, _db, _user_id, _seeded = api

    assert client.get(LIST, params={"query": "x" * 201}).status_code == 422
    assert client.get(SUMMARY, params={"query": "x" * 201}).status_code == 422
    assert client.get(EXPORT, params={"query": "x" * 201}).status_code == 422
    assert client.get(LIST, params={"query": "x" * 200}).status_code == 200


def test_a_wildcard_in_the_query_is_a_literal_character(api):
    """`%` and `_` are SQL, not search syntax. Unescaped, `50%` matched every row with 50 in
    it and `_` matched any character at all - so the box quietly answered a different
    question from the one that was typed."""
    client, _db, _user_id, seeded = api

    # The item codes here are ZZTTOK-SRTWT6801 / ZZTTOK-SRTWC9042: `ZZTTOK-%` is a literal
    # nobody carries, and `ZZTTOK_S` would match `ZZTTOK-S` if the underscore were a
    # wildcard.
    assert _ids(client, "ZZTTOK-%") == set()
    assert _ids(client, "ZZTTOK_S") == set()
    # And the plain prefix still answers, so the escaping has not broken matching itself.
    assert _ids(client, "ZZTTOK-S") == {
        seeded["a_alpha"].id,
        seeded["a_beta"].id,
        seeded["b_alpha"].id,
    }


# ------------------------------------------------------------------ AC-2.4


def test_the_summary_tiles_count_the_same_rows_the_list_shows(api):
    """`_worklist_filters` is shared, so the tiles cannot be allowed to answer the phrase
    while the list answers the tokens - the strip would read 0 above one visible row."""
    client, _db, _user_id, seeded = api

    body = client.get(SUMMARY, params={"query": f"{seeded['so_a']} {ALPHA}"}).json()

    assert body["total_rows"] == 1


def test_the_export_carries_the_same_rows_the_list_shows(api):
    """The workbook is what purchasing sends on, so it has to be the rows on screen."""
    client, _db, _user_id, seeded = api

    response = client.get(EXPORT, params={"query": f"{seeded['so_a']} {ALPHA}"})

    assert response.status_code == 200, response.text
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    # One sheet per delivery MONTH, two heading rows, then the rows themselves.
    body_rows = [
        sheet.max_row - 2 for sheet in book.worksheets if sheet.title != "ORDER INQUIRY"
    ]
    assert sum(body_rows) == 1, book.sheetnames
