"""Dealer targets import on the files AS DELIVERED (lane SALES-TARGET-DEALER, gate-fail round).

Crew ran the owner's gate on the real Sorento and Mocha files (16 Oct): the engine could hold
every scheme, the importer could not read the files. This suite pins items 1 to 8 of that round
and the owner's 09:12 ruling (match dealers by NAME for this seeding; an optional CUSTOMER CODE
column still wins; anything not an exact normalised match is listed with up to 3 suggestions and
the user picks or skips it before Apply). The workbook is `_dealer_target_real_layout.build`, a
copy of the real LAYOUT with masked names and shape-only figures.
"""
from __future__ import annotations

import io
import json
from datetime import date
from decimal import Decimal

from . import _dealer_target_real_layout as rl
from .test_sales_targets_s1 import (  # noqa: F401  (fixtures)
    BASE,
    _category,
    _code,
    _product,
    _so_line,
    _uid,
    api,
    world,
)

IMPORT = f"{BASE}/import"
D = Decimal
S, M = "SORENTO 2026", "MOCHA 2026"


def _parse(**kwargs):
    from app.services.sales.dealer_target_import import parse_workbook

    data, s_rows, m_rows = rl.build(**kwargs)
    return parse_workbook(data, year=2026), s_rows, m_rows


def _k(sheet, row):
    """The key a row without a customer code goes by in a parsed scheme."""
    return f"{sheet}!{row}"


def _schemes(parsed, key, kind, unit=None):
    return [
        s for s in parsed.schemes
        if key in s.codes and s.kind == kind and (unit is None or s.threshold_unit == unit)
    ]


def _one(parsed, key, kind, unit=None):
    found = _schemes(parsed, key, kind, unit)
    assert len(found) == 1, [(s.kind, s.threshold_unit, s.codes) for s in parsed.schemes]
    return found[0]


def _tiers(scheme, start):
    return [
        (t.threshold, t.rate, t.is_additional, t.account_rates)
        for t in scheme.tiers
        if t.period_start == start
    ]


def _row(parsed, sheet, row):
    return next(r for r in parsed.rows if (r.sheet, r.row) == (sheet, row))


# --------------------------------------------------------------------------------------- #
# 1. the files as delivered: stacked headers, merged account rows
# --------------------------------------------------------------------------------------- #


def test_reads_both_sheets_with_stacked_headers_and_every_data_row():
    parsed, s_rows, m_rows = _parse()
    assert {r.sheet for r in parsed.rows} == {S, M}
    assert sorted(r.row for r in parsed.rows if r.sheet == S) == sorted(s_rows.values())
    assert sorted(r.row for r in parsed.rows if r.sheet == M) == sorted(m_rows.values())
    # A name merged down over the account rows is every account row's name.
    a3 = _row(parsed, S, s_rows["ALPHA:A/C III"])
    assert a3.name_in_file == "Alpha Hardware Sdn. Bhd."
    assert a3.account_in_file == "A/C III"
    assert parsed.notes == []


def test_q1_over_jan_mar_over_tier_and_tier_incentive_with_merged_accounts():
    parsed, s_rows, _ = _parse()
    a1, a3 = _k(S, s_rows["ALPHA:A/C I"]), _k(S, s_rows["ALPHA:A/C III"])
    q = _one(parsed, a1, "quarterly")
    assert q.codes == [a1, a3]
    assert (q.start, q.end, q.split_every, q.split_unit) == (date(2026, 1, 1), date(2026, 12, 31), 3, "month")
    assert q.threshold_unit == "amount"
    # 6. "+ Additional 1%" as a SECOND tier at the same threshold is that tier's additional; a
    # second additional line at its own threshold stacks too (ruling 5).
    assert _tiers(q, date(2026, 1, 1)) == [
        (D("500000"), D("2"), False, None),
        (D("500000"), D("1"), True, None),
        (D("600000"), D("0.5"), True, None),
    ]


def test_monthly_header_over_jan_to_dec():
    parsed, s_rows, _ = _parse()
    m = _one(parsed, _k(S, s_rows["ECHO:A/C I"]), "monthly")
    assert (m.split_every, m.split_unit) == (1, "month")
    assert _tiers(m, date(2026, 1, 1)) == [(D("50000"), D("1"), False, None)]
    assert _tiers(m, date(2026, 2, 1)) == [(D("52000"), D("1"), False, None)]


def test_mocha_quarterly_point_and_point_yearly_headers():
    parsed, _, m_rows = _parse()
    k1 = _k(M, m_rows["KILO:A/C I"])
    q = _one(parsed, k1, "quarterly")
    assert q.codes == [k1, _k(M, m_rows["KILO:[CERAMIC]"]), _k(M, m_rows["KILO:[IBORN]"])]
    assert q.threshold_unit == "points"
    for start in (date(2026, 1, 1), date(2026, 4, 1), date(2026, 7, 1), date(2026, 10, 1)):
        assert _tiers(q, start) == [(D("100000"), D("4"), False, None)]
    y = _one(parsed, k1, "yearly")
    assert y.threshold_unit == "points"
    assert [(t.threshold, t.rate) for t in y.tiers] == [(D("450000"), D("1")), (D("500000"), D("1.5"))]


# --------------------------------------------------------------------------------------- #
# 3. rates: "A/C I- 3%" names the account it pays; a fraction is a percent
# --------------------------------------------------------------------------------------- #


def test_fraction_is_a_percent_and_rate_per_account():
    parsed, s_rows, _ = _parse()
    keys = [_k(S, s_rows[f"BRAVO:{label}"]) for label in ("Deluxe", "Mastile", "W-Ply")]
    q = _one(parsed, keys[0], "quarterly")
    assert q.codes == keys
    assert _tiers(q, date(2026, 1, 1)) == [
        (D("300000"), None, False, dict(zip(keys, (D("2"), D("3"), D("1.5"))))),
    ]
    y = _one(parsed, _k(S, s_rows["ALPHA:A/C I"]), "yearly")
    assert [(t.threshold, t.rate) for t in y.tiers] == [(D("1500000"), D("1")), (D("2000000"), D("1.5"))]


def test_rate_naming_one_account_of_the_target_pays_only_that_account():
    parsed, s_rows, _ = _parse()
    a1 = _k(S, s_rows["ALPHA:A/C I"])
    q = _one(parsed, a1, "quarterly")
    assert _tiers(q, date(2026, 4, 1)) == [(D("400000"), None, False, {a1: D("3")})]
    assert q.payout_code == a1


def test_rate_naming_another_account_of_the_dealer_is_the_pay_to_case():
    parsed, s_rows, _ = _parse()
    d1, d3 = _k(S, s_rows["DELTA:A/C I"]), _k(S, s_rows["DELTA:A/C III"])
    own = _one(parsed, d1, "quarterly")
    assert own.codes == [d1] and own.payout_code == d1
    paid = _one(parsed, d3, "quarterly")
    assert paid.codes == [d3]                     # counts A/C III
    assert paid.payout_code == d1                 # pays A/C I
    assert _tiers(paid, date(2026, 1, 1)) == [(D("80000"), D("2"), False, None)]


# --------------------------------------------------------------------------------------- #
# 4. units per column; 5. a Dec-Nov yearly window leaves the quarters alone
# --------------------------------------------------------------------------------------- #


def test_units_are_read_per_column_not_per_sheet():
    parsed, _, m_rows = _parse()
    lima = _k(M, m_rows["LIMA:A/C I"])
    y = _one(parsed, lima, "yearly")
    assert y.threshold_unit == "amount"           # "YEARLY (RM)" on the Mocha sheet
    assert [(t.threshold, t.rate) for t in y.tiers] == [(D("600000"), D("1")), (D("750000"), D("2"))]
    assert _schemes(parsed, lima, "yearly", "points") == []


def test_dec_to_nov_yearly_window_does_not_shift_the_quarters():
    parsed, s_rows, _ = _parse()
    c = _k(S, s_rows["CHARLIE:A/C I"])
    y = _one(parsed, c, "yearly")
    assert (y.start, y.end) == (date(2025, 12, 1), date(2026, 11, 30))
    assert [(t.threshold, t.rate) for t in y.tiers] == [(D("1500000"), D("1"))]   # 7. "≥1500000"
    q = _one(parsed, c, "quarterly")
    assert (q.start, q.end) == (date(2026, 1, 1), date(2026, 12, 31))
    assert _tiers(q, date(2026, 1, 1)) == [(D("100000"), D("2"), False, None)]


# --------------------------------------------------------------------------------------- #
# 7. typos are reported on their row, never mapped
# --------------------------------------------------------------------------------------- #


def test_typos_fail_their_row_with_a_reason():
    parsed, s_rows, _ = _parse()
    foxtrot = _row(parsed, S, s_rows["FOXTROT:A/C I"])
    # "A/C 1" is never read as A/C I on its own: it waits on an account pick (round 10).
    assert foxtrot.status == "needs_pick"
    assert "A/C 1" in foxtrot.reason and "not an account" in foxtrot.reason
    assert foxtrot.account_pick["label"] == "A/C 1"
    golf = _row(parsed, S, s_rows["GOLF:A/C I"])
    # Owner 20 Oct 10:11: a tier cell that is not a rate skips that tier, with a warning on the row.
    assert golf.status == "mapped"
    assert any("skipped" in w and "not a rate" in w for w in golf.warnings), golf.warnings
    hotel = _row(parsed, S, s_rows["HOTEL:A/C I"])
    assert hotel.status == "not_representable"
    assert "NIL" in hotel.reason
    for key in ("FOXTROT", "GOLF", "HOTEL"):
        assert _schemes(parsed, _k(S, s_rows[f"{key}:A/C I"]), "quarterly") == []


def test_text_threshold_forms_parse():
    from app.services.sales.dealer_target_import import _number

    assert _number("≥1500000") == D("1500000")
    assert _number(">= 1,500,000") == D("1500000")
    assert _number(">1500000") == D("1500000")


# --------------------------------------------------------------------------------------- #
# 8. nothing is dropped silently
# --------------------------------------------------------------------------------------- #


def test_a_sheet_the_reader_cannot_read_fails_its_rows():
    def notes(ws):
        ws["A1"], ws["B1"] = "Dealer X", "call back in Nov"
        ws["A2"], ws["B2"] = "Dealer Y", 300000

    from app.services.sales.dealer_target_import import gate_passes, parse_workbook

    data, _, _ = rl.build(extra_sheet=notes)
    parsed = parse_workbook(data, year=2026)
    failed = [r for r in parsed.rows if r.sheet == "NOTES"]
    assert [r.row for r in failed] == [1, 2]
    assert all(r.status == "not_representable" and "NOTES" in r.reason for r in failed)
    assert gate_passes({"summary": {"not_representable": 0, "needs_pick": 0}, "notes": ["x"]}) is False
    assert gate_passes({"summary": {"not_representable": 1, "needs_pick": 0}, "notes": []}) is False
    assert gate_passes({"summary": {"not_representable": 0, "needs_pick": 1}, "notes": []}) is False
    assert gate_passes({"summary": {"not_representable": 0, "needs_pick": 0}, "notes": []}) is True


def test_a_filled_period_cell_that_yields_no_tier_fails_its_row():
    """A cell with a figure in a period column is never dropped: a threshold with no
    incentive, or text that reads as nothing, fails the row."""
    from openpyxl import Workbook

    from app.services.sales.dealer_target_import import parse_workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"], ws["E1"] = "DEALER NAME", "ACCOUNT", "Q1", "Q2"
    ws.merge_cells("C1:D1")
    ws.merge_cells("E1:F1")
    ws["C2"], ws["D2"], ws["E2"], ws["F2"] = "Tier 1", "Tier 1 Incentive", "Tier 1", "Tier 1 Incentive"
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "X", "A/C I", 100000, 0.02
    ws["E3"] = 200000                              # Q2 threshold, no incentive
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_workbook(buf.getvalue(), year=2026)
    assert parsed.rows[0].status == "not_representable"
    assert "Q2" in parsed.rows[0].reason


# --------------------------------------------------------------------------------------- #
# the route: name matching (owner ruling 16 Oct 09:12), picks, apply
# --------------------------------------------------------------------------------------- #


def _upload(client, data, *, dry_run, db=None, picks=None):
    files = {"file": ("dealer targets.xlsx", data,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    form = {"picks": json.dumps(picks)} if picks is not None else None
    if db is not None:
        db.autoflush = False
    try:
        return client.post(IMPORT, params={"dry_run": str(dry_run).lower(), "year": 2026}, files=files, data=form)
    finally:
        if db is not None:
            db.autoflush = True


def _by(body, sheet, row):
    return next(r for r in body["rows"] if (r["sheet"], r["row"]) == (sheet, row))


def test_dry_run_matches_by_name_and_account_and_lists_the_rest(api):
    client, db, company_id = api
    crm = rl.seed(db, company_id, _uid, _code)
    data, s_rows, m_rows = rl.build()
    res = _upload(client, data, dry_run=True, db=db)
    assert res.status_code == 200, res.text
    body = res.json()

    def row(key, sheet=S, rows=s_rows):
        return _by(body, sheet, rows[key])

    # Exact normalised name + account label: auto-mapped.
    for key, crm_key in (
        ("ALPHA:A/C I", ("ALPHA", "A/C I")),            # "Sdn. Bhd." vs "SDN BHD", level 1
        ("ALPHA:A/C III", ("ALPHA", "A/C III")),        # level 3
        ("BRAVO:Deluxe", ("BRAVO", "Deluxe")),          # the [DELUXE] marker, no level
        ("BRAVO:W-Ply", ("BRAVO", "W-Ply")),            # [W-PLY]
        ("CHARLIE:A/C I", ("CHARLIE", "A/C I")),        # Enterprises vs ENTERPRISE
        ("ECHO:A/C I", ("ECHO", "A/C I")),              # file without SDN BHD
    ):
        r = row(key)
        assert r["status"] == "mapped", (key, r)
        assert r["match"] == "auto"
        assert r["customer"] == crm[crm_key].customer_name
    kilo = row("KILO:[CERAMIC]", M, m_rows)
    assert kilo["status"] == "mapped" and kilo["customer"] == crm[("KILO", "[CERAMIC]")].customer_name

    # No exact match: listed with up to 3 suggestions, never auto-linked.
    india = row("INDIA:A/C I")
    assert india["status"] == "needs_pick"
    assert india["customer"] is None
    names = [s["name"] for s in india["suggestions"]]
    assert 1 <= len(names) <= 3
    assert "INDIA MARBLE SDN BHD - [A/C I]" in names
    assert all(s["customer_id"] and s["code"] for s in india["suggestions"])

    summary = body["summary"]
    assert summary["auto_mapped"] == 14                     # GOLF too: its typo tier is skipped, not the row
    assert summary["needs_pick"] == 2                       # INDIA, and FOXTROT's "A/C 1" account pick
    assert summary["not_representable"] == 1                # the HOTEL typo row
    assert summary["rows"] == len(s_rows) + len(m_rows)


def test_two_candidates_for_one_name_and_account_are_never_auto_linked(api):
    client, db, company_id = api
    from app.models.order import Customer

    rl.seed(db, company_id, _uid, _code)
    db.add(Customer(id=_uid(), company_id=company_id, customer_code=_code(),
                    customer_name="Alpha Hardware (A/C I)", account_level=1))
    db.flush()
    data, s_rows, _ = rl.build()
    body = _upload(client, data, dry_run=True, db=db).json()
    a1 = _by(body, S, s_rows["ALPHA:A/C I"])
    assert a1["status"] == "needs_pick"
    assert len(a1["suggestions"]) == 2
    # Its merged partner cannot become a target without it.
    a3 = _by(body, S, s_rows["ALPHA:A/C III"])
    assert a3["status"] != "mapped" or a3["target_keys"] == []


def test_customer_code_column_wins_over_the_name(api):
    client, db, company_id = api
    crm = rl.seed(db, company_id, _uid, _code)
    codes = {
        ("ALPHA", "A/C I"): crm[("ALPHA", "A/C I")].customer_code.lower(),   # case ignored
        ("ECHO", "A/C I"): crm[("BRAVO", "Deluxe")].customer_code,          # code beats name
        ("CHARLIE", "A/C I"): "ZZT-NO-SUCH-CODE",                            # unknown: fails
    }
    data, s_rows, _ = rl.build(with_code_column=True, codes=codes)
    body = _upload(client, data, dry_run=True, db=db).json()
    assert _by(body, S, s_rows["ALPHA:A/C I"])["customer"] == crm[("ALPHA", "A/C I")].customer_name
    echo = _by(body, S, s_rows["ECHO:A/C I"])
    assert echo["customer"] == crm[("BRAVO", "Deluxe")].customer_name
    assert any("name differs" in w for w in echo["warnings"])
    charlie = _by(body, S, s_rows["CHARLIE:A/C I"])
    assert charlie["status"] == "not_representable"
    assert "ZZT-NO-SUCH-CODE" in charlie["reason"]
    # A row with an empty code cell still matches by name.
    assert _by(body, S, s_rows["BRAVO:Deluxe"])["match"] == "auto"


def test_apply_uses_picks_and_skips_and_computes(api):
    client, db, company_id = api
    from app.models.order import Customer

    crm = rl.seed(db, company_id, _uid, _code)
    near = db.query(Customer).filter(Customer.customer_name == "INDIA MARBLE SDN BHD - [A/C I]").one()
    pid = _product(db, company_id, _category(db, company_id).id).id

    def sale(key, amount, on):
        _so_line(db, company_id, customer_id=crm[key].id, order_date=on, line_total=D(str(amount)), product_id=pid)

    sale(("ALPHA", "A/C I"), 400000, date(2026, 2, 1))     # Q1 550,000: 2% + 1% additional
    sale(("ALPHA", "A/C III"), 150000, date(2026, 2, 1))
    sale(("ALPHA", "A/C I"), 300000, date(2026, 5, 1))     # Q2 450,000 >= 400,000: 3% on A/C I only
    sale(("ALPHA", "A/C III"), 150000, date(2026, 5, 1))
    sale(("DELTA", "A/C III"), 90000, date(2026, 2, 1))    # paid to A/C I

    data, s_rows, m_rows = rl.build()
    picks = {
        f"{S}!{s_rows['INDIA:A/C I']}": near.id,
        f"{S}!{s_rows['HOTEL:A/C I']}": "skip",
    }
    res = _upload(client, data, dry_run=False, db=db, picks=picks)
    assert res.status_code == 200, res.text
    body = res.json()
    india = _by(body, S, s_rows["INDIA:A/C I"])
    assert india["status"] == "mapped" and india["match"] == "picked"
    assert india["customer"] == "INDIA MARBLE SDN BHD - [A/C I]"
    assert india["target_nos"]
    assert _by(body, S, s_rows["HOTEL:A/C I"])["status"] == "skipped"
    assert body["summary"]["skipped"] == 1

    rows = client.get(BASE, params={"subject": "dealer", "all": "true"}).json()["rows"]

    def detail(key, sheet=S, which=s_rows, kind="quarterly"):
        nos = _by(body, sheet, which[key])["target_nos"]
        target = next(r for r in rows if r["target_no"] in nos and r["name"].endswith(kind))
        return client.get(f"{BASE}/{target['target_id']}").json()

    def period(d, start):
        return next(p for p in d["periods"] if p["period_start"] == start)

    alpha = detail("ALPHA:A/C I")
    assert period(alpha, "2026-01-01")["commission_earned"] == 16500      # 11,000 + 5,500
    assert period(alpha, "2026-04-01")["commission_earned"] == 9000       # 3% of A/C I's 300,000
    delta = detail("DELTA:A/C III")
    assert [a["code"] for a in delta["accounts"]] == [crm[("DELTA", "A/C III")].customer_code]
    assert delta["payout_account"]["code"] == crm[("DELTA", "A/C I")].customer_code
    assert period(delta, "2026-01-01")["commission_earned"] == 1800
    charlie_q = detail("CHARLIE:A/C I")
    assert (charlie_q["start_date"], charlie_q["end_date"]) == ("2026-01-01", "2026-12-31")
    lima = detail("LIMA:A/C I", M, m_rows, "yearly")
    assert lima["threshold_unit"] == "amount"
    kilo = detail("KILO:A/C I", M, m_rows, "yearly")
    assert kilo["threshold_unit"] == "points"


def test_a_pick_must_be_a_customer_of_this_company(api):
    client, db, company_id = api
    rl.seed(db, company_id, _uid, _code)
    data, s_rows, _ = rl.build()
    picks = {f"{S}!{s_rows['INDIA:A/C I']}": _uid()}
    body = _upload(client, data, dry_run=True, db=db, picks=picks).json()
    india = _by(body, S, s_rows["INDIA:A/C I"])
    assert india["status"] == "not_representable"
    assert "picked customer" in india["reason"]


def test_picks_that_are_not_a_json_object_are_refused(api):
    client, db, company_id = api
    data, _, _ = rl.build()
    files = {"file": ("t.xlsx", data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res = client.post(IMPORT, params={"dry_run": "true", "year": 2026}, files=files, data={"picks": "[1, 2]"})
    assert res.status_code == 422


# --------------------------------------------------------------------------------------- #
# 2. the template: the real layout plus a CUSTOMER CODE column
# --------------------------------------------------------------------------------------- #


def test_template_downloads_and_reads_back(api):
    client, db, company_id = api
    from openpyxl import load_workbook

    from app.services.sales.dealer_target_import import parse_workbook

    res = client.get(f"{IMPORT}/template", params={"year": 2026})
    assert res.status_code == 200, res.text
    assert "spreadsheetml" in res.headers["content-type"]
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.worksheets[0]
    headers = {str(c.value).upper() for row in ws.iter_rows(max_row=6) for c in row if c.value}
    for text in ("DEALER NAME", "CUSTOMER CODE", "ACCOUNT", "Q1", "JAN-MAR", "TIER 1", "TIER 1 INCENTIVE",
                 "MONTHLY", "JAN"):
        assert text in headers, text
    # The office fills one row; it reads back as one quarterly target.
    first = max(r for r in range(1, 10) if any(c.value for c in ws[r])) + 1
    cols = {}
    for c in ws[first - 1]:
        if c.value:
            cols.setdefault(str(c.value).upper(), c.column)
    name_col = next(c.column for row in ws.iter_rows(max_row=first - 1) for c in row if str(c.value).upper() == "DEALER NAME")
    code_col = next(c.column for row in ws.iter_rows(max_row=first - 1) for c in row if str(c.value).upper() == "CUSTOMER CODE")
    acc_col = next(c.column for row in ws.iter_rows(max_row=first - 1) for c in row if str(c.value).upper() == "ACCOUNT")
    ws.cell(row=first, column=name_col, value="ZZT DEALER")
    ws.cell(row=first, column=code_col, value="ZZT-1")
    ws.cell(row=first, column=acc_col, value="A/C I")
    tier = cols["TIER 1"]
    ws.cell(row=first, column=tier, value=100000)
    ws.cell(row=first, column=tier + 1, value=0.02)
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_workbook(buf.getvalue(), year=2026)
    assert parsed.notes == []
    assert [r.status for r in parsed.rows] == ["mapped"]
    assert [s.kind for s in parsed.schemes] == ["quarterly"]
    assert parsed.schemes[0].codes == ["ZZT-1"]


def test_a_plain_second_tier_at_the_same_threshold_is_reported_not_guessed():
    """Item 6 covers "+ Additional x%" at a repeated threshold; a repeated threshold with a
    plain rate has no reading, so its row fails with the reason instead of a guess."""
    from openpyxl import Workbook

    from app.services.sales.dealer_target_import import parse_workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"] = "DEALER NAME", "ACCOUNT", "Q1"
    ws.merge_cells("C1:F1")
    for col, text in zip("CDEF", ("Tier 1", "Tier 1 Incentive", "Tier 2", "Tier 2 Incentive")):
        ws[f"{col}2"] = text
    ws["A3"], ws["B3"], ws["C3"], ws["D3"], ws["E3"], ws["F3"] = "X", "A/C I", 500000, 0.02, 500000, 0.01
    buf = io.BytesIO()
    wb.save(buf)
    parsed = parse_workbook(buf.getvalue(), year=2026)
    assert parsed.rows[0].status == "not_representable"
    assert "repeats" in parsed.rows[0].reason


# --------------------------------------------------------------------------------------- #
# security review (gate-fail round): the name match cannot be made to eat the server
# --------------------------------------------------------------------------------------- #


def _one_row_workbook(rows):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["C1"] = "DEALER NAME", "ACCOUNT", "Q1"
    for i, (name, account) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=name)
        ws.cell(row=i, column=2, value=account)
        ws.cell(row=i, column=3, value="100,000 - 1%")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_an_overlong_name_fails_its_row_before_any_match():
    """B1: a dealer name of unbounded length never reaches the matcher (difflib indexes it)."""
    from app.services.sales import dealer_target_import as imp

    parsed = imp.parse_workbook(_one_row_workbook([("AB" * 5000, "A/C I")]), year=2026)
    assert parsed.rows[0].status == "not_representable"
    assert "characters" in parsed.rows[0].reason
    assert parsed.schemes == []


def test_the_fuzzy_match_runs_once_per_name_and_is_capped(api, monkeypatch):
    """B2: the same unmatched name on many rows is looked up once; past the cap a row still
    waits on a pick but gets no suggestions."""
    import difflib

    from app.services.sales import dealer_target_import as imp

    client, db, company_id = api
    calls = []
    real = difflib.get_close_matches
    monkeypatch.setattr(difflib, "get_close_matches", lambda *a, **k: calls.append(a[0]) or real(*a, **k))
    monkeypatch.setattr(imp, "MAX_FUZZY_LOOKUPS", 2)
    rows = [("ZZT NOWHERE ONE", "A/C I")] * 20 + [("ZZT NOWHERE TWO", "A/C I"), ("ZZT NOWHERE THREE", "A/C I")]
    body = _upload(client, _one_row_workbook(rows), dry_run=True, db=db).json()
    assert body["summary"]["needs_pick"] == 22
    assert len(calls) == 2
    last = body["rows"][-1]
    assert last["status"] == "needs_pick" and last["suggestions"] == []
    assert "CUSTOMER CODE" in last["reason"]


def test_deeply_nested_picks_are_refused_with_422(api):
    client, db, company_id = api
    files = {"file": ("t.xlsx", _one_row_workbook([("X", "A/C I")]),
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    res = client.post(IMPORT, params={"dry_run": "true", "year": 2026}, files=files, data={"picks": "[" * 100000})
    assert res.status_code == 422


# --------------------------------------------------------------------------------------- #
# code review (gate-fail round)
# --------------------------------------------------------------------------------------- #


def _wb(fill):
    from openpyxl import Workbook

    wb = Workbook()
    fill(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_a_scheme_whose_pay_to_row_is_skipped_is_not_created(api):
    """Review B1: DELTA A/C III pays A/C I; skipping A/C I fails A/C III, no target pays it."""
    client, db, company_id = api
    rl.seed(db, company_id, _uid, _code)
    data, s_rows, _ = rl.build()
    picks = {f"{S}!{s_rows['DELTA:A/C I']}": "skip"}
    body = _upload(client, data, dry_run=True, db=db, picks=picks).json()
    d3 = _by(body, S, s_rows["DELTA:A/C III"])
    assert d3["status"] == "not_representable"
    assert d3["reason"] == f"shares a target with row {s_rows['DELTA:A/C I']}, which is skipped"
    assert "!" not in d3["reason"]


def test_a_text_only_first_data_row_is_never_read_as_a_header():
    """Review S1: "PTS TRADING" with tier text under the header is data, not a sub-header."""
    from app.services.sales.dealer_target_import import parse_workbook

    def fill(ws):
        ws["A1"], ws["B1"] = "DEALER NAME", "Q1"
        ws["A2"], ws["B2"] = "PTS TRADING", "300,000 - 2%"
        ws["A3"], ws["B3"] = "MAY HARDWARE", "200,000 - 1%"

    parsed = parse_workbook(_wb(fill), year=2026)
    assert [(r.row, r.name_in_file, r.status) for r in parsed.rows] == [
        (2, "PTS TRADING", "mapped"), (3, "MAY HARDWARE", "mapped"),
    ]


def test_one_cell_paying_a_counted_and_an_uncounted_account_fails():
    """Review S2: the paired-column path already refuses it; the one-cell path does too."""
    from app.services.sales.dealer_target_import import parse_workbook

    def fill(ws):
        ws["A1"], ws["B1"], ws["C1"] = "DEALER NAME", "ACCOUNT", "Q1"
        ws["A2"], ws["B2"], ws["C2"] = "X", "A/C I", "300,000 A/C I 2% / A/C III 3%"
        ws["A3"], ws["B3"], ws["C3"] = "X", "A/C III", "100,000 - 1%"

    parsed = parse_workbook(_wb(fill), year=2026)
    row2 = next(r for r in parsed.rows if r.row == 2)
    assert row2.status == "not_representable"
    assert "both" in row2.reason


def test_a_rate_above_twenty_percent_fails_and_a_percent_is_no_threshold():
    """Review S3, crew ruling (b): 0.5 reads as 50% under the fraction rule; above
    RATE_LIMIT_PCT the row fails with a reason, never a possible 100x overpay. A "2%" typed in a
    Tier threshold column is not a threshold of RM 2."""
    from app.services.sales.dealer_target_import import parse_workbook

    def fill(ws):
        ws["A1"], ws["B1"], ws["C1"] = "DEALER NAME", "ACCOUNT", "Q1"
        ws.merge_cells("C1:D1")
        ws["C2"], ws["D2"] = "Tier 1", "Tier 1 Incentive"
        ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "X", "A/C I", 100000, 0.5
        ws["A4"], ws["B4"], ws["C4"], ws["D4"] = "Y", "A/C I", "2%", 0.02

    parsed = parse_workbook(_wb(fill), year=2026)
    x, y = parsed.rows
    assert x.status == "not_representable"
    assert "rate 50% looks wrong: Excel 0.5 means 50%, type 0.005 or 0.5% for half a percent" in x.reason
    assert y.status == "not_representable"
    assert "threshold" in y.reason


def test_two_unlabelled_rates_in_one_incentive_cell_fail():
    from app.services.sales.dealer_target_import import CellError, parse_rate_cell

    import pytest

    with pytest.raises(CellError):
        parse_rate_cell("2% / 3%")


def test_a_typo_in_the_account_cell_says_so(api):
    client, db, company_id = api
    rl.seed(db, company_id, _uid, _code)
    body = _upload(client, _one_row_workbook([("ECHO HOME", "A/C 1")]), dry_run=True, db=db).json()
    row = body["rows"][0]
    assert row["status"] == "needs_pick"
    assert "A/C 1" in row["reason"]


def test_two_rows_of_one_file_with_the_same_accounts_and_dates_fail_the_second(api):
    client, db, company_id = api
    rl.seed(db, company_id, _uid, _code)
    rows = [("ECHO HOME", "A/C I"), ("ECHO HOME", "A/C I")]

    def fill(ws):
        ws["A1"], ws["B1"], ws["C1"] = "DEALER NAME", "ACCOUNT", "Q1"
        for i, (name, account) in enumerate(rows, start=2):
            ws.cell(row=i, column=1, value=name)
            ws.cell(row=i, column=2, value=account)
            ws.cell(row=i, column=3, value=f"{i}00,000 - 1%")
        # Two dealers' blocks: a blank NO-less gap keeps them apart as separate rows.

    body = _upload(client, _wb(fill), dry_run=True, db=db).json()
    statuses = [r["status"] for r in body["rows"]]
    assert statuses.count("mapped") == 1
    second = body["rows"][1]
    assert second["status"] == "not_representable"
    assert "row 2" in second["reason"]


def test_a_refused_scheme_rolls_back_its_dealer_s_other_schemes(api, monkeypatch):
    """Review S4: no partial dealer on apply. ALPHA's yearly target is refused by the target
    service; its quarterly target must not be created either."""
    from app.services.error_handler import AppException
    from app.services.sales import target_service

    client, db, company_id = api
    rl.seed(db, company_id, _uid, _code)
    real = target_service.create_target

    def refuse_alpha_yearly(db_, payload, **kwargs):
        if payload.name.startswith("Alpha") and payload.name.endswith("yearly"):
            raise AppException(status_code=422, message="refused for the test", code="X")
        return real(db_, payload, **kwargs)

    monkeypatch.setattr(target_service, "create_target", refuse_alpha_yearly)
    data, s_rows, _ = rl.build()
    body = _upload(client, data, dry_run=False, db=db).json()
    a1 = _by(body, S, s_rows["ALPHA:A/C I"])
    assert a1["status"] == "not_representable" and "refused for the test" in a1["reason"]
    rows = client.get(BASE, params={"subject": "dealer", "all": "true"}).json()["rows"]
    assert not [r for r in rows if r["name"].startswith("Alpha")]
    assert body["summary"]["targets_new"] == len(rows)


def test_a_long_digit_run_in_a_header_cell_parses_fast():
    """Security re-check B3: `_PCT` / `_AMOUNT` are quadratic on a long run of digits; every
    cell's text is cut once, so a few-KB upload cannot hold the CPU for hours."""
    import time

    from app.services.sales.dealer_target_import import _number, parse_workbook

    def fill(ws):
        ws["A1"], ws["B1"] = "DEALER NAME", "Q1"
        ws["B2"] = "1" * 100_000
        ws["A3"], ws["B3"] = "ACME", "100,000 - 1%"

    started = time.monotonic()
    parse_workbook(_wb(fill), year=2026)
    assert _number("9" * 100_000 + "x") is None
    assert time.monotonic() - started < 5


def test_every_string_is_cut_when_the_grid_is_built():
    """Security re-check B4: one long shared string can sit in every cell; the grid keeps at
    most MAX_CELL_TEXT + 1 characters of it, so no later read (dash, NIL, %) walks it whole."""
    from openpyxl import load_workbook

    from app.services.sales import dealer_target_import as imp

    def fill(ws):
        ws["A1"], ws["B1"] = "DEALER NAME", "Q1"
        ws["A2"], ws["B2"] = "ACME", "1 " * 10_000

    ws = load_workbook(io.BytesIO(_wb(fill))).active
    grid = imp._Grid(ws)
    assert len(grid.value(2, 2)) == imp.MAX_CELL_TEXT + 1


def test_a_sheet_title_past_excels_limit_is_refused():
    """Security re-check B5: a title is copied into every row's key and reason; Excel itself
    stops at 31 characters, so a longer one is not a real file."""
    import pytest

    from app.services.error_handler import AppException
    from app.services.sales.dealer_target_import import parse_workbook

    def fill(ws):
        ws.title = "X" * 40
        ws["A1"], ws["B1"] = "DEALER NAME", "Q1"

    with pytest.raises(AppException) as exc:
        parse_workbook(_wb(fill), year=2026)
    assert exc.value.status_code == 422


def test_a_number_format_past_excels_limit_is_never_scanned():
    """Security re-check B6: one long shared number format would be scanned per cell; a format
    longer than Excel's 255 characters is treated as no format."""
    from types import SimpleNamespace

    from app.services.sales import dealer_target_import as imp

    long_fmt = "0" * 10_000 + "%"
    ws = SimpleNamespace(
        _cells={(1, 1): SimpleNamespace(value=0.02, number_format=long_fmt),
                (1, 2): SimpleNamespace(value=0.02, number_format="0%")},
        merged_cells=SimpleNamespace(ranges=[]),
    )
    grid = imp._Grid(ws)
    assert grid.fmt(1, 1) is None
    assert grid.fmt(1, 2) == "0%"
