"""RED tests for SR4 of the AutoCount pull + review lane (issue #1048).

Covers: AC-SP-1..5, AC-SC-1..4, AC-CM-3 (stock compare), plus the stock halves
of AC-RV-3 and AC-RV-5.

Plan: documentation/plans/_archive/autocount/PLAN-autocount-pull-review.md
UAC:  documentation/plans/_archive/autocount/autocount-pull-review-acceptance-criteria.md

Nothing under test here exists yet - `classify_stock_rows` (autocount_pull_
service.py), `compare_stock` (autocount_pull_compare.py), the stock branches of
`preview_autocount_pull` / `apply_autocount_pull` (autocount_pull_tasks.py),
`app.services.stock_list_archive_service` (new module) and the entity dispatch
that currently answers `_require_products_entity` -> 501 for every stock route
are ALL the coder's SR4 deliverable. Every one of those is imported (or
invoked) INSIDE a test/fixture body, never at module import time, so a missing
piece reds one test, not collection.

**Valid red reasons today**, all pinned by name in the captain's brief: a bare
`ModuleNotFoundError`/`ImportError` on `app.services.stock_list_archive_service`
or on a not-yet-defined `import_outcome_codes` constant; `AC-BD/RV/CM/PC`-style
routes answering `501 NOT_IMPLEMENTED` because `_require_products_entity` still
gates every `/{job_id}/...` route by entity; and `app.tasks.autocount_pull_
tasks.UnsupportedPullEntity` because `preview_autocount_pull` / `apply_
autocount_pull` still dispatch products-only. A test that would ALSO pass today
purely because "the job ends up failed either way" is not a real red - every
guard test below additionally pins something today's `UnsupportedPullEntity`
message does not contain (a guard-specific phrase, a phase of `review` instead
of `failed`, a route call that never happens), so it fails for the pinned
reason and not by coincidence.

**Stock header trap** (captain's brief): the committed fixture pair
`stock-header-ready.json` / `stock-rows-page1.json` describes FoundryX's real
~12,133-row snapshot (`recordCount`, `zeroPairs`, `negativePairs`, ...) while
the rows fixture holds only the first 10-row page - the header's own counts do
NOT describe those 10 rows. The header fixture is used ONLY where a test wants
those exact ten rows unmodified (the review-page tests below, which never lean
on `recordCount`/`negativePairs`/etc, only on the rows themselves being
served). Every other test builds its own header with `_stock_header(rows,
...)`, sized to the handful of rows it actually serves.

Substrate reused from SR1 (`tests/test_autocount_pull_sr1.py`) and SR3
(`tests/test_autocount_pull_sr3.py`) by import, not copy - same technique SR3
used to reuse SR1's `env`/`task_db`: importing a `@pytest.fixture` function's
NAME into this module's globals is enough for pytest to discover it here too.
`_FakeFoundryX` defaults to the PRODUCTS fixtures; every stock test below
points `fake.status` / `fake.rows` at stock shapes explicitly (never relies on
the constructor default).
"""
from __future__ import annotations

import io
import uuid
from decimal import Decimal
from unittest.mock import MagicMock

import openpyxl
import pytest
from sqlalchemy import text

# MUST be the first app import - resolves the circular import in
# app.modules.runtime.guards (repo convention, see test_ingest_deletions.py).
from app.main import app  # noqa: E402,F401

from app.services.company_scope import DEFAULT_COMPANY_ID

from tests._pg_fixture import unique_code
from tests.test_autocount_pull_sr1 import (  # noqa: F401 - env/task_db are fixtures
    API_KEY,
    BASE_URL,
    FIXTURE_DIR,
    PULLS_URL,
    _FakeFoundryX,
    _content_hash,
    _fixture,
    _header,
    _job_count,
    _job_row,
    _job_rows,
    _patch_foundryx,
    _run_preview,
    _seed_pull_job,
    _stored,
    env,
    task_db,
)
from tests.test_autocount_pull_sr3 import (  # noqa: F401 - reused helpers
    _run_apply,
    _seed_apply_job,
)
from tests.test_stock_list_xlsm_upload import (  # noqa: F401 - `client` is a fixture
    REPLACE_URL,
    _USER_ID,
    _FakeStorageBackend,
    _xlsx_bytes,
    client,
)

MARKER = "ZZTAP4"

_FIXTURE_ACTIVE_LOCATIONS = ("BRW-BB", "BRW", "MWH", "WH3")
_FIXTURE_INACTIVE_LOCATIONS = ("MKT-D",)
_STOCK_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# =========================================================== row / header shaping


def _stock_row(item_code: str, location_code: str, qty: int, **overrides) -> dict:
    """One raw FoundryX stock snapshot row (the `stock-rows-page1.json` shape)."""
    row = {
        "source_ref": f"{MARKER}:{item_code}|{location_code}",
        "item_code": item_code,
        "item_description": f"{MARKER} DESC {item_code}",
        "location_code": location_code,
        "uom_code": "UNIT",
        "qty": qty,
    }
    row.update(overrides)
    return row


def _raw_stock_to_template(row: dict) -> dict:
    """The manual-template shape (`Item Code`/`Item Description`/`Location`/`On
    Hand Qty`), built FROM a raw snapshot row - what a checker's own Excel file,
    or `classify_stock_rows`'s `fed` bucket, is expected to carry."""
    return {
        "Item Code": row["item_code"],
        "Item Description": row.get("item_description") or "",
        "Location": row["location_code"],
        "On Hand Qty": row["qty"],
    }


def _expected_stock_view_row(row: dict) -> dict:
    """The `/rows` shape AC-RV-3's stock half pins."""
    return {
        "item_code": row["item_code"],
        "item_description": row.get("item_description") or "",
        "location": row["location_code"],
        "on_hand_qty": row["qty"],
    }


def _stock_header(
    rows: list[dict],
    *,
    record_count: int | None = None,
    company_code: str = "SRT",
    complete: bool = True,
    content_hash: str | None = None,
    zero_pairs: int = 0,
    negative_pairs: int = 0,
    negative_pair_list: list | None = None,
    excluded_rows: list | None = None,
    excluded_nonzero_count: int = 0,
) -> dict:
    """A ready STOCK header, built from SR1's `_header` (which hardcodes
    `entity: "products"` and the products-only price counters - harmless,
    unread by the stock path) then re-shaped with the stock-only keys the real
    fixture (`stock-header-ready.json`) carries: `zeroPairs`, `negativePairs`,
    `negativePairList`, `excludedNonzeroCount`. Sized to `rows` - see the
    module docstring's "stock header trap"."""
    header = _header(
        rows, record_count=record_count, company_code=company_code, complete=complete,
        content_hash=content_hash, excluded_rows=excluded_rows,
    )
    header["entity"] = "stock_balances"
    header["zeroPairs"] = zero_pairs
    header["negativePairs"] = negative_pairs
    header["fractionalPairs"] = 0
    header["excludedNonzeroCount"] = excluded_nonzero_count
    if negative_pair_list is not None:
        header["negativePairList"] = negative_pair_list
    return header


def _prepare_stock_preview(
    db, fake: _FakeFoundryX, *, rows: list[dict], fetched_header: dict | None = None,
    stored_header: dict | None = None, pages: dict | None = None,
    snapshot_id: str | None = None, owner: str | None = None,
    company_id: str = DEFAULT_COMPANY_ID, company_code: str = "SRT",
):
    """Seeds a `previewing` STOCK pull and points the fake at its snapshot -
    the stock analogue of SR1's `_prepare_preview`."""
    snapshot_id = snapshot_id or f"{MARKER}-stock-snap-{uuid.uuid4().hex[:8]}"
    valid = _stock_header(rows, company_code=company_code)
    fake.status = (200, {**(fetched_header or valid), "snapshotId": snapshot_id})
    fake.rows = pages or {
        1: (200, {"snapshotId": snapshot_id, "page": 1, "pageSize": 1000, "totalPages": 1,
                  "recordCount": len(rows), "rows": rows})
    }
    return _seed_pull_job(
        db, job_type="autocount_stock_pull", user_id=owner or str(uuid.uuid4()),
        company_id=company_id, entity="stock_balances", company_code=company_code,
        snapshot_id=snapshot_id, phase="previewing", header=_stored(stored_header or valid),
    )


def _seed_stock_review_job(env, *, owner, phase="review", rows=None, extra=None,
                            fetched_header=None):
    """Seeds a stock pull job directly at `phase` (bypassing the preview task) and
    points the fake FoundryX's `/rows` at `rows` (default: the committed 10-row
    stock fixture - see the module docstring's header trap)."""
    rows = rows if rows is not None else _fixture("stock-rows-page1.json")["rows"]
    snapshot_id = f"{MARKER}-snap-{uuid.uuid4().hex[:8]}"
    header = fetched_header or _stock_header(rows)
    job_id = _seed_pull_job(
        env.db, job_type="autocount_stock_pull", user_id=owner["id"], company_id=env.company_a,
        entity="stock_balances", company_code=env.company_a_code, snapshot_id=snapshot_id,
        phase=phase, header=_stored(header), extra=extra,
    )
    env.fake.rows = {
        1: (200, {"snapshotId": snapshot_id, "page": 1, "pageSize": 1000, "totalPages": 1,
                  "recordCount": len(rows), "rows": rows})
    }
    return job_id, rows


def _seed_stock_active_warehouses(db, company_id) -> None:
    """The four active + one inactive warehouse codes the committed stock
    fixture's ten rows resolve against (see module docstring): FED = the rows
    at BRW-BB/BRW/MWH/WH3 (4 of 10), `not applied, inactive` = the MKT-D row,
    `not applied, unknown` = the remaining five (BRW-DISP, PARTS, BRW-IR,
    PJ-SR, MAINTANC - deliberately unseeded)."""
    from app.models.inventory import Warehouse

    for code in _FIXTURE_ACTIVE_LOCATIONS:
        db.add(Warehouse(id=str(uuid.uuid4()), warehouse_code=code, warehouse_name=code,
                          is_active=True, company_id=company_id))
    for code in _FIXTURE_INACTIVE_LOCATIONS:
        db.add(Warehouse(id=str(uuid.uuid4()), warehouse_code=code, warehouse_name=code,
                          is_active=False, company_id=company_id))
    db.commit()


def _seed_product_with_stock(db, company_id, warehouse_id, *, code: str, qty: int) -> str:
    from app.models.inventory import Stock
    from app.models.product import Product, ProductCategory, UnitOfMeasure

    cat = ProductCategory(category_code=unique_code(MARKER), category_name="cat")
    uom = UnitOfMeasure(uom_code=unique_code(MARKER)[:20], uom_name="unit")
    db.add_all([cat, uom])
    db.flush()
    product = Product(
        product_code=code, product_name=code, category_id=cat.id, base_uom_id=uom.id,
        list_price=Decimal("1.00"), company_id=company_id,
    )
    db.add(product)
    db.flush()
    db.add(Stock(id=str(uuid.uuid4()), product_id=product.id, warehouse_id=warehouse_id,
                  quantity_on_hand=qty, quantity_reserved=0, quantity_damaged=0))
    db.commit()
    return product.id


def _seed_product_only(db, company_id, *, code: str) -> str:
    from app.models.product import Product, ProductCategory, UnitOfMeasure

    cat = ProductCategory(category_code=unique_code(MARKER), category_name="cat")
    uom = UnitOfMeasure(uom_code=unique_code(MARKER)[:20], uom_name="unit")
    db.add_all([cat, uom])
    db.flush()
    product = Product(
        product_code=code, product_name=code, category_id=cat.id, base_uom_id=uom.id,
        list_price=Decimal("1.00"), company_id=company_id,
    )
    db.add(product)
    db.commit()
    return product.id


def _stock_qty(db, product_code: str) -> int | None:
    return db.execute(
        text(
            "SELECT s.quantity_on_hand FROM stock s "
            "JOIN products p ON p.id = s.product_id "
            "WHERE p.product_code = :c"
        ),
        {"c": product_code},
    ).scalar()


# ======================================================================= SP-1


class TestStockPreviewGuards:
    """AC-SP-1: "Guards as AC-PP-1" - the SAME `fetch_verified_snapshot` guard
    function SR1 already proved for products, reused unchanged for stock. Each
    arm pins a guard-specific phrase in the stored error - today's actual
    failure reason (`UnsupportedPullEntity`, "Stock preview is not implemented
    yet") contains none of them, so this is red for the guard, not merely
    red because nothing runs."""

    @pytest.mark.parametrize(
        "overrides,expected_snippet",
        [
            ({"complete": False}, "incomplete"),
            ({"record_count": 999}, "does not match"),
            ({"company_code": "MCH"}, "does not match"),
        ],
        ids=["incomplete", "record_count_mismatch", "company_code_mismatch"],
    )
    def test_sp_1a_bad_fetched_header_fails_the_stock_job(
        self, task_db, monkeypatch, overrides, expected_snippet
    ):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        rows = [_stock_row(f"{MARKER}-SP1A", f"{MARKER}-SP1A-WH", 5)]
        job_id = _prepare_stock_preview(
            db, fake, rows=rows, fetched_header=_stock_header(rows, **overrides)
        )
        before = db.execute(text("SELECT count(*) FROM stock")).scalar()

        _run_preview(monkeypatch, factory, job_id)

        row_after = _job_row(db, job_id)
        assert row_after["status"] == "failed", row_after
        assert row_after["metadata"]["autocount_pull"]["phase"] == "failed"
        error = (row_after["error"] or "").lower()
        assert expected_snippet in error, error
        assert _job_rows(db, job_id) == []
        after = db.execute(text("SELECT count(*) FROM stock")).scalar()
        assert after == before


class TestStockPreviewConfirmBlocked:
    def test_sp_1b_excluded_nonzero_finishes_in_review_with_a_blocked_reason(
        self, task_db, monkeypatch
    ):
        from app.models.inventory import Warehouse

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)

        wh_code = f"{MARKER}-SP1B-WH"
        db.add(Warehouse(id=str(uuid.uuid4()), warehouse_code=wh_code, warehouse_name="Fed",
                          is_active=True, company_id=DEFAULT_COMPANY_ID))
        db.commit()

        rows = [_stock_row(f"{MARKER}-SP1B", wh_code, 4)]
        header = _stock_header(rows, excluded_nonzero_count=3)
        job_id = _prepare_stock_preview(
            db, fake, rows=rows, fetched_header=header, stored_header=header
        )

        _run_preview(monkeypatch, factory, job_id)

        row_after = _job_row(db, job_id)
        assert row_after["status"] == "finished", row_after["error"]
        pull = row_after["metadata"]["autocount_pull"]
        assert pull["phase"] == "review"
        reason = pull.get("confirm_blocked_reason")
        assert isinstance(reason, str) and reason.strip(), pull


# ======================================================================= SP-2


class TestClassifyStockRows:
    def test_sp_2a_active_inactive_unknown_trim_case_and_cross_company(self, env):
        from app.models.base import set_company_scope
        from app.models.inventory import Warehouse

        active_code = f"{MARKER}-FED"
        inactive_code = f"{MARKER}-OFF"
        shared_code = f"{MARKER}-SHARED"  # exists, but only for company_b

        env.db.add_all([
            Warehouse(id=str(uuid.uuid4()), warehouse_code=active_code, warehouse_name="Fed",
                      is_active=True, company_id=env.company_a),
            Warehouse(id=str(uuid.uuid4()), warehouse_code=inactive_code, warehouse_name="Off",
                      is_active=False, company_id=env.company_a),
            Warehouse(id=str(uuid.uuid4()), warehouse_code=shared_code, warehouse_name="OtherCo",
                      is_active=True, company_id=env.company_b),
        ])
        env.db.commit()

        # Ambient scope covers BOTH companies, deliberately - if `classify_stock_
        # rows` ever forgot its own `company_id` filter and leaned on ambient scope
        # instead, the cross-company row would wrongly classify as FED here.
        set_company_scope(env.db, frozenset({env.company_a, env.company_b}))

        rows = [
            _stock_row("A1", f" {active_code.lower()} ", 5),   # trimmed + lowercased -> FED
            _stock_row("A2", inactive_code, 3),                 # -> inactive
            _stock_row("A3", f"{MARKER}-GHOST", 1),              # -> unknown, no such warehouse
            _stock_row("A4", shared_code, 2),                    # exists, but for company_b
        ]

        from app.services.autocount_pull_service import classify_stock_rows

        result = classify_stock_rows(env.db, env.company_a, rows)

        assert {r["Item Code"] for r in result["fed"]} == {"A1"}
        fed_row = result["fed"][0]
        assert fed_row == {
            "Item Code": "A1", "Item Description": rows[0]["item_description"],
            "Location": rows[0]["location_code"], "On Hand Qty": 5,
        }

        def _codes(bucket):
            return {r.get("item_code") or r.get("Item Code") for r in bucket}

        assert _codes(result["inactive"]) == {"A2"}
        assert _codes(result["unknown"]) == {"A3", "A4"}


class TestStockPreviewFeedsFedOnly:
    def test_sp_2b_only_fed_rows_reach_bulk_import_stock(self, task_db, monkeypatch):
        from app.models.inventory import Warehouse
        from app.services.inventory_service import StockService

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)

        active_code = f"{MARKER}-FED2B"
        inactive_code = f"{MARKER}-OFF2B"
        db.add_all([
            Warehouse(id=str(uuid.uuid4()), warehouse_code=active_code, warehouse_name="Fed",
                      is_active=True, company_id=DEFAULT_COMPANY_ID),
            Warehouse(id=str(uuid.uuid4()), warehouse_code=inactive_code, warehouse_name="Off",
                      is_active=False, company_id=DEFAULT_COMPANY_ID),
        ])
        db.commit()

        rows = [
            _stock_row("FEDONE", active_code, 4),
            _stock_row("OFFONE", inactive_code, 9),
            _stock_row("GHOSTONE", f"{MARKER}-NOWHERE", 1),
        ]
        job_id = _prepare_stock_preview(db, fake, rows=rows)

        captured: dict = {}
        real_bulk_import = StockService.bulk_import_stock

        def _spy(self, stock_data, user_id, validate_only=False, outcome=None):
            captured["stock_data"] = stock_data
            captured["validate_only"] = validate_only
            return real_bulk_import(
                self, stock_data, user_id, validate_only=validate_only, outcome=outcome
            )

        monkeypatch.setattr(StockService, "bulk_import_stock", _spy)

        _run_preview(monkeypatch, factory, job_id)

        assert captured, "bulk_import_stock was never called by the stock preview"
        assert captured["validate_only"] is True
        codes_sent = {r["Item Code"] for r in captured["stock_data"]}
        assert codes_sent == {"FEDONE"}


# =============================================================== SP-3 / SP-4


def _seed_sp_scenario(db, fake: _FakeFoundryX) -> str:
    """One preview scenario exercising every SP-3/SP-4 counter at once: a qty
    change, an unchanged pair, a new pair, a product-not-found row, a not-
    applied-inactive row, a not-applied-unknown row, a header-only negative
    pair, and an existing active-warehouse pair the payload never mentions
    (so it should be swept to zero)."""
    from app.models.inventory import Warehouse

    active_wh = Warehouse(id=str(uuid.uuid4()), warehouse_code=f"{MARKER}-SP-FED",
                           warehouse_name="Fed", is_active=True, company_id=DEFAULT_COMPANY_ID)
    inactive_wh = Warehouse(id=str(uuid.uuid4()), warehouse_code=f"{MARKER}-SP-OFF",
                             warehouse_name="Off", is_active=False, company_id=DEFAULT_COMPANY_ID)
    db.add_all([active_wh, inactive_wh])
    db.flush()

    _seed_product_with_stock(db, DEFAULT_COMPANY_ID, active_wh.id, code=f"{MARKER}-SP1", qty=5)
    _seed_product_with_stock(db, DEFAULT_COMPANY_ID, active_wh.id, code=f"{MARKER}-SP2", qty=5)
    _seed_product_with_stock(db, DEFAULT_COMPANY_ID, active_wh.id, code=f"{MARKER}-SP5", qty=7)
    _seed_product_only(db, DEFAULT_COMPANY_ID, code=f"{MARKER}-SP3")

    rows = [
        _stock_row(f"{MARKER}-SP1", active_wh.warehouse_code, 9),   # existing 5 -> 9: qty change
        _stock_row(f"{MARKER}-SP2", active_wh.warehouse_code, 5),   # existing 5 -> 5: unchanged
        _stock_row(f"{MARKER}-SP3", active_wh.warehouse_code, 3),   # no existing stock: new pair
        _stock_row(f"{MARKER}-SP4", active_wh.warehouse_code, 2),   # no such product
        _stock_row(f"{MARKER}-SP6", inactive_wh.warehouse_code, 1),  # not applied, inactive
        _stock_row(f"{MARKER}-SP7", f"{MARKER}-SP-GHOST", 1),        # not applied, unknown
    ]
    negative_entry = {
        "item_code": f"{MARKER}-SP8", "location_code": active_wh.warehouse_code, "qty": -3,
    }
    header = _stock_header(rows, negative_pair_list=[negative_entry])
    return _prepare_stock_preview(
        db, fake, rows=rows, fetched_header=header, stored_header=header
    )
    # SP5 (qty 7, active warehouse) is deliberately NOT in `rows` - it is the
    # "existing pair AutoCount no longer sends" case (AC-SP-3's "would set to 0").


class TestStockPreviewValidateOnly:
    def test_sp_3_validate_only_leaves_stock_unchanged_and_stores_a_summary(
        self, task_db, monkeypatch
    ):
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        job_id = _seed_sp_scenario(db, fake)

        def _snapshot():
            rows = db.execute(
                text(
                    "SELECT p.product_code, w.warehouse_code, s.quantity_on_hand "
                    "FROM stock s JOIN products p ON p.id = s.product_id "
                    "JOIN warehouses w ON w.id = s.warehouse_id "
                    "WHERE p.product_code LIKE :prefix"
                ),
                {"prefix": f"{MARKER}-SP%"},
            ).mappings().all()
            return {(r["product_code"], r["warehouse_code"]): r["quantity_on_hand"] for r in rows}

        before = _snapshot()
        assert before, "seed produced no stock rows to compare against"

        _run_preview(monkeypatch, factory, job_id)

        row_after = _job_row(db, job_id)
        assert row_after["status"] == "finished", row_after["error"]
        assert row_after["metadata"]["autocount_pull"]["phase"] == "review"

        after = _snapshot()
        assert after == before, "a validate_only preview must not write to stock"

        counts = row_after["metadata"]["autocount_pull"]["counts"]
        assert counts, "no summary was stored on the pull"


class TestStockPreviewJobRowsAndCounts:
    def test_sp_4_import_job_rows_and_all_eight_counters(self, task_db, monkeypatch):
        from app.services import import_outcome_codes as codes

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        job_id = _seed_sp_scenario(db, fake)

        _run_preview(monkeypatch, factory, job_id)

        row_after = _job_row(db, job_id)
        assert row_after["status"] == "finished", row_after["error"]
        counts = row_after["metadata"]["autocount_pull"]["counts"]
        assert counts["received"] == 6, counts
        assert counts["fed"] == 4, counts
        assert counts["not_applied_inactive"] == 1, counts
        assert counts["not_applied_unknown"] == 1, counts
        assert counts["qty_changes"] == 1, counts
        assert counts["set_to_zero"] == 1, counts
        assert counts["skipped_product_not_found"] == 1, counts
        assert counts["negative_in_autocount"] == 1, counts

        rows_written = _job_rows(db, job_id)

        def _blob(r) -> str:
            return " ".join(str(v) for v in (r["message"], r["identity"], r["value"]) if v)

        not_applied_inactive = [
            r for r in rows_written if r["code"] == codes.AUTOCOUNT_NOT_APPLIED_INACTIVE
        ]
        not_applied_unknown = [
            r for r in rows_written if r["code"] == codes.AUTOCOUNT_NOT_APPLIED_UNKNOWN
        ]
        negative_rows = [r for r in rows_written if r["code"] == codes.AUTOCOUNT_NEGATIVE]
        product_not_found = [r for r in rows_written if r["code"] == codes.PRODUCT_NOT_FOUND]

        assert len(not_applied_inactive) == 1, rows_written
        assert f"{MARKER}-SP6" in _blob(not_applied_inactive[0])
        assert len(not_applied_unknown) == 1, rows_written
        assert f"{MARKER}-SP7" in _blob(not_applied_unknown[0])
        assert len(negative_rows) == 1, rows_written
        assert f"{MARKER}-SP8" in _blob(negative_rows[0])
        assert len(product_not_found) == 1, rows_written
        assert f"{MARKER}-SP4" in _blob(product_not_found[0])

        qty_change_rows = [r for r in rows_written if f"{MARKER}-SP1" in _blob(r)]
        assert len(qty_change_rows) == 1, rows_written
        blob = _blob(qty_change_rows[0])
        assert "5" in blob and "9" in blob, blob

        assert len(rows_written) == 5, rows_written

    def test_b3_progress_is_set_to_the_total_once_the_preview_finishes(self, task_db, monkeypatch):
        """B3 (small-fix track): stock has no per-record hook to report progress through
        mid-run (unlike `MasterIngestService.ingest`'s `on_progress`) - `_preview_stock`
        publishes `processed_rows`/`total_rows` once, at the end, both equal to the
        received row count."""
        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        job_id = _seed_sp_scenario(db, fake)

        _run_preview(monkeypatch, factory, job_id)

        progress = db.execute(
            text("SELECT processed_rows, total_rows FROM import_jobs WHERE id = :id"),
            {"id": str(job_id)},
        ).mappings().first()
        received = _job_row(db, job_id)["metadata"]["autocount_pull"]["counts"]["received"]
        assert progress["total_rows"] == received == 6
        assert progress["processed_rows"] == received


# ======================================================================= SP-5


class TestStockConfirmNeverBlockedBySkip:
    def test_sp_5_would_skip_rows_never_block_confirm(self, env, monkeypatch):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_stock_review_job(
            env, owner=owner,
            extra={
                "counts": {
                    "received": 1, "fed": 1, "not_applied_inactive": 0,
                    "not_applied_unknown": 0, "qty_changes": 0, "set_to_zero": 0,
                    "skipped_product_not_found": 1, "negative_in_autocount": 0,
                },
                "confirm_blocked_reason": None,
            },
        )
        monkeypatch.setattr(
            "app.services.autocount_pull_service.enqueue_job",
            lambda *a, **k: MagicMock(id=str(uuid.uuid4())),
        )

        resp = env.client.post(f"{PULLS_URL}/{job_id}/confirm")

        assert resp.status_code == 200, resp.text
        assert resp.json()["phase"] == "confirmed"


# ======================================================================= RV-3s


class TestStockRowsRoute:
    def test_rv_3s_rows_route_returns_fed_only_in_key_order(self, env):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        _seed_stock_active_warehouses(env.db, env.company_a)
        job_id, rows = _seed_stock_review_job(env, owner=owner)

        resp = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"limit": 50})

        assert resp.status_code == 200, resp.text
        data = resp.json()["data"]
        assert len(data) == len(_FIXTURE_ACTIVE_LOCATIONS)

        expected = {
            r["item_code"]: _expected_stock_view_row(r)
            for r in rows if r["location_code"] in _FIXTURE_ACTIVE_LOCATIONS
        }
        assert {r["item_code"] for r in data} == set(expected)
        for got in data:
            assert list(got.keys()) == ["item_code", "item_description", "location", "on_hand_qty"]
            want = expected[got["item_code"]]
            assert got["item_description"] == want["item_description"]
            assert got["location"] == want["location"]
            assert int(got["on_hand_qty"]) == int(want["on_hand_qty"])

    def test_rv_3s_query_filters_by_item_code(self, env):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        _seed_stock_active_warehouses(env.db, env.company_a)
        job_id, rows = _seed_stock_review_job(env, owner=owner)

        resp = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"query": "acc-cb", "limit": 50})

        assert resp.status_code == 200, resp.text
        codes = {r["item_code"] for r in resp.json()["data"]}
        assert codes == {"ACC-CB8001"}

    def test_rv_3s_non_owner_404(self, env):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        _seed_stock_active_warehouses(env.db, env.company_a)
        job_id, _ = _seed_stock_review_job(env, owner=owner)

        other = env.user("inventory.stock.autocount_pull")
        env.as_user(other)
        resp = env.client.get(f"{PULLS_URL}/{job_id}/rows")
        assert resp.status_code == 404
        assert resp.json().get("code") == "NOT_FOUND", resp.json()


# ======================================================================= RV-5s


class TestStockDownloadRoute:
    def test_rv_5s_header_and_rows_match_fed_only(self, env):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        _seed_stock_active_warehouses(env.db, env.company_a)
        job_id, rows = _seed_stock_review_job(env, owner=owner)

        rows_resp = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"limit": 50})
        assert rows_resp.status_code == 200, rows_resp.text
        api_rows = rows_resp.json()["data"]
        assert len(api_rows) == len(_FIXTURE_ACTIVE_LOCATIONS)

        resp = env.client.get(f"{PULLS_URL}/{job_id}/download.xlsx")
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"].startswith(_STOCK_XLSX_MEDIA_TYPE)
        disposition = resp.headers.get("content-disposition", "")
        assert "filename" in disposition
        # D3 (small-fix track): the stock DOWNLOAD uses the same stock-list name the
        # apply's archive step gets - lowercase company code, no more the generic
        # `autocount-stock_balances-pull.xlsx` (nor the apply-side `autocount-pull-pull.xlsx`).
        from datetime import datetime
        from zoneinfo import ZoneInfo

        today = datetime.now(ZoneInfo("Asia/Kuala_Lumpur")).strftime("%Y%m%d")
        assert f"autocount-stock-list-srt-{today}.xlsx" in disposition, disposition

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert header == ["Item Code", "Item Description", "Location", "On Hand Qty"]

        body_rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(body_rows) == len(api_rows)
        by_code = {r["item_code"]: r for r in api_rows}
        for sheet_row in body_rows:
            api_row = by_code[sheet_row[0]]
            assert (sheet_row[1] or "") == (api_row["item_description"] or "")
            assert sheet_row[2] == api_row["location"]
            assert int(sheet_row[3]) == int(api_row["on_hand_qty"])

    def test_empty_company_code_falls_back_the_same_way_the_download_does(self):
        """N2 (opus review, fix round 3): `_company_code` (both `autocount_pull_service.py`
        and `autocount_pull_tasks.py`'s own copy) returns "" for a job whose `company_id`
        does not resolve - `stock_list_archive_filename` must fall back exactly like
        `download_filename` already did, not print a bare double hyphen."""
        from app.services.autocount_pull_service import stock_list_archive_filename

        assert "autocount-stock-list--" not in stock_list_archive_filename("")


# ======================================================================= CM-3a


class TestCompareStockPureFunction:
    """`app.services.autocount_pull_compare.compare_stock(excel_rows, fed_rows)`
    - both sides in the manual-template shape (`Item Code`/`Item Description`/
    `Location`/`On Hand Qty`), keyed by (Item Code, Location) trimmed and
    case-insensitive, On Hand Qty compared as integers. Same return shape as
    `compare_products`, plus `summary.qty_total_excel`/`qty_total_pull`
    (AC-CM-3)."""

    def test_cm_3a_key_trim_case_insensitive_and_integer_qty_forms_match(self):
        from app.services.autocount_pull_compare import compare_stock

        fed_row = {"Item Code": "ABC-1", "Item Description": "Widget", "Location": "LOC1",
                   "On Hand Qty": 5}
        excel_row = {"Item Code": " abc-1 ", "Item Description": "Widget", "Location": " loc1 ",
                     "On Hand Qty": "5.0"}

        result = compare_stock([excel_row], [fed_row])

        assert result["differences"] == []
        assert result["only_in_excel"] == []
        assert result["only_in_pull"] == []
        assert result["summary"]["matched"] == 1
        assert result["summary"]["total"] == 1

    def test_cm_3b_differing_qty_is_one_difference_entry(self):
        from app.services.autocount_pull_compare import compare_stock

        fed_row = {"Item Code": "ABC-2", "Item Description": "Widget", "Location": "LOC1",
                   "On Hand Qty": 5}
        excel_row = {"Item Code": "ABC-2", "Item Description": "Widget", "Location": "LOC1",
                     "On Hand Qty": 9}

        result = compare_stock([excel_row], [fed_row])

        assert len(result["differences"]) == 1, result["differences"]
        diff = result["differences"][0]
        assert diff["item_code"] == "ABC-2"
        assert diff["location"] == "LOC1"
        assert diff["field"] == "on_hand_qty"
        assert int(diff["excel"]) == 9
        assert int(diff["pull"]) == 5

    def test_cm_3c_same_item_code_different_location_are_distinct_pairs(self):
        from app.services.autocount_pull_compare import compare_stock

        fed_rows = [
            {"Item Code": "ABC-3", "Item Description": "W", "Location": "LOC1", "On Hand Qty": 5},
            {"Item Code": "ABC-3", "Item Description": "W", "Location": "LOC2", "On Hand Qty": 5},
        ]
        excel_rows = [
            {"Item Code": "ABC-3", "Item Description": "W", "Location": "LOC1", "On Hand Qty": 5},
        ]

        result = compare_stock(excel_rows, fed_rows)

        assert result["only_in_pull"] == [("ABC-3", "LOC2")] or result["only_in_pull"] == ["ABC-3|LOC2"] \
            or any("LOC2" in str(x) for x in result["only_in_pull"]), result["only_in_pull"]
        assert result["summary"]["total"] == 1
        assert result["summary"]["matched"] == 1

    def test_cm_3d_only_in_excel_only_in_pull_and_totals(self):
        from app.services.autocount_pull_compare import compare_stock

        fed_rows = [{"Item Code": "A", "Item Description": "x", "Location": "L1", "On Hand Qty": 3}]
        excel_rows = [{"Item Code": "B", "Item Description": "y", "Location": "L2", "On Hand Qty": 4}]

        result = compare_stock(excel_rows, fed_rows)

        assert len(result["only_in_excel"]) == 1
        assert len(result["only_in_pull"]) == 1
        assert result["summary"]["qty_total_excel"] == 4
        assert result["summary"]["qty_total_pull"] == 3


class TestCompareStockRoute:
    def test_cm_3e_file_built_from_fed_rows_is_a_perfect_match(self, env):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        _seed_stock_active_warehouses(env.db, env.company_a)
        job_id, rows = _seed_stock_review_job(env, owner=owner)

        fed_raw = [r for r in rows if r["location_code"] in _FIXTURE_ACTIVE_LOCATIONS]
        template_rows = [_raw_stock_to_template(r) for r in fed_raw]

        resp = env.client.post(
            f"{PULLS_URL}/{job_id}/compare",
            json={"filename": "stock-check.xlsx", "rows": template_rows},
        )

        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["summary"]["matched"] == len(fed_raw), body
        assert body["summary"]["total"] == len(fed_raw), body
        assert body["differences"] == []

    def test_cm_3f_wrong_entity_permission_403(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_stock_review_job(env, owner=owner)

        resp = env.client.post(
            f"{PULLS_URL}/{job_id}/compare", json={"filename": "x.xlsx", "rows": []},
        )
        assert resp.status_code == 403


# ======================================================================= SC-1


class TestStockConfirmRoute:
    def test_sc_1a_confirm_creates_stock_apply_job_enqueues_once(self, env, monkeypatch):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_stock_review_job(env, owner=owner)

        captured = []

        def _enqueue(func, *a, **k):
            captured.append({"func": func, "kwargs": k})
            return MagicMock(id=str(uuid.uuid4()))

        monkeypatch.setattr("app.services.autocount_pull_service.enqueue_job", _enqueue)

        resp = env.client.post(f"{PULLS_URL}/{job_id}/confirm")

        assert resp.status_code == 200, resp.text
        body = resp.json()
        assert body["phase"] == "confirmed"
        apply_job_id = body["apply_job_id"]
        assert apply_job_id

        apply_row = _job_row(env.db, apply_job_id)
        assert apply_row is not None
        assert apply_row["job_type"] == "autocount_stock_apply"
        assert apply_row["metadata"]["autocount_apply"]["entity"] == "stock_balances"
        assert apply_row["metadata"]["autocount_apply"]["pull_job_id"] == str(job_id)

        assert len(captured) == 1
        assert getattr(captured[0]["func"], "__name__", "") == "apply_autocount_pull"

    def test_sc_1a_second_confirm_is_idempotent(self, env, monkeypatch):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_stock_review_job(env, owner=owner)

        captured = []
        monkeypatch.setattr(
            "app.services.autocount_pull_service.enqueue_job",
            lambda *a, **k: captured.append(1) or MagicMock(id=str(uuid.uuid4())),
        )

        first = env.client.post(f"{PULLS_URL}/{job_id}/confirm")
        second = env.client.post(f"{PULLS_URL}/{job_id}/confirm")

        assert first.status_code == 200, first.text
        assert second.status_code == 200, second.text
        assert second.json()["apply_job_id"] == first.json()["apply_job_id"]
        assert len(captured) == 1

    @pytest.mark.parametrize("phase", ["building", "previewing", "failed", "expired"])
    def test_sc_1b_non_review_phase_is_409(self, env, phase):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_stock_review_job(env, owner=owner, phase=phase)

        resp = env.client.post(f"{PULLS_URL}/{job_id}/confirm")
        assert resp.status_code == 409, resp.text

    def test_sc_1b_excluded_nonzero_blocked_reason_is_409(self, env):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_stock_review_job(
            env, owner=owner,
            extra={
                "confirm_blocked_reason": "AutoCount reports non-zero excluded pairs; pull again.",
            },
        )

        resp = env.client.post(f"{PULLS_URL}/{job_id}/confirm")
        assert resp.status_code == 409, resp.text

    def test_sc_1c_non_owner_404(self, env):
        owner = env.user("inventory.stock.autocount_pull")
        env.as_user(owner)
        job_id, _ = _seed_stock_review_job(env, owner=owner)

        other = env.user("inventory.stock.autocount_pull")
        env.as_user(other)
        resp = env.client.post(f"{PULLS_URL}/{job_id}/confirm")
        assert resp.status_code == 404
        assert resp.json().get("code") == "NOT_FOUND", resp.json()


# ======================================================================= SC-2


class TestStockApplyTask:
    def test_sc_2_fed_applied_active_absent_zeroed_inactive_always_untouched(
        self, task_db, monkeypatch
    ):
        from app.models.inventory import Warehouse

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)

        active_wh = Warehouse(id=str(uuid.uuid4()), warehouse_code=f"{MARKER}-SC2-FED",
                               warehouse_name="Fed", is_active=True, company_id=DEFAULT_COMPANY_ID)
        inactive_wh = Warehouse(id=str(uuid.uuid4()), warehouse_code=f"{MARKER}-SC2-OFF",
                                 warehouse_name="Off", is_active=False,
                                 company_id=DEFAULT_COMPANY_ID)
        db.add_all([active_wh, inactive_wh])
        db.flush()

        _seed_product_with_stock(db, DEFAULT_COMPANY_ID, active_wh.id, code=f"{MARKER}-SC2A", qty=5)
        _seed_product_with_stock(db, DEFAULT_COMPANY_ID, active_wh.id, code=f"{MARKER}-SC2B", qty=7)
        _seed_product_with_stock(db, DEFAULT_COMPANY_ID, inactive_wh.id, code=f"{MARKER}-SC2C",
                                  qty=11)
        _seed_product_with_stock(db, DEFAULT_COMPANY_ID, inactive_wh.id, code=f"{MARKER}-SC2D",
                                  qty=13)

        rows = [
            _stock_row(f"{MARKER}-SC2A", active_wh.warehouse_code, 9),
            _stock_row(f"{MARKER}-SC2D", inactive_wh.warehouse_code, 99),
            _stock_row(f"{MARKER}-SC2E", active_wh.warehouse_code, 1),  # no such product
        ]
        snapshot_id = f"{MARKER}-snap-sc2"
        header = _stock_header(rows)
        fake.status = (200, {**header, "snapshotId": snapshot_id})
        fake.rows = {1: (200, {"snapshotId": snapshot_id, "page": 1, "pageSize": 1000,
                                "totalPages": 1, "recordCount": len(rows), "rows": rows})}

        job_id = _seed_apply_job(
            db, user_id=str(uuid.uuid4()), company_id=DEFAULT_COMPANY_ID, entity="stock_balances",
            snapshot_id=snapshot_id,
        )

        _run_apply(monkeypatch, factory, job_id)

        row_after = _job_row(db, job_id)
        assert row_after["status"] == "finished", row_after["error"]

        assert _stock_qty(db, f"{MARKER}-SC2A") == 9   # updated to the fed value
        assert _stock_qty(db, f"{MARKER}-SC2B") == 0   # active, absent from FED -> zeroed
        assert _stock_qty(db, f"{MARKER}-SC2C") == 11  # inactive, never sent -> untouched
        assert _stock_qty(db, f"{MARKER}-SC2D") == 13  # inactive, WAS sent -> still untouched

        no_stock = db.execute(
            text(
                "SELECT count(*) FROM stock s JOIN products p ON p.id = s.product_id "
                "WHERE p.product_code = :c"
            ),
            {"c": f"{MARKER}-SC2E"},
        ).scalar()
        assert no_stock == 0


# ======================================================================= SC-3


class TestStockListArchiveServiceUsedByRoute:
    """AC-SC-3: the archive logic is ONE service function used by both the
    existing `replace-latest-stock-list` route and the pull Confirm. Pinned by
    patching `app.services.stock_list_archive_service.replace_latest_stock_list`
    (the module the route is expected to import it from, late, per this repo's
    own inline-import convention for every other storage/webhook helper in
    `app/api/v1/resources/attachments.py`) and observing the route call it with
    the file bytes, filename and confirming user id.

    The coder must keep every test in `tests/test_stock_list_xlsm_upload.py`
    green: `test_replace_latest_xlsm_is_converted_to_template_only_xlsx`,
    `test_replace_latest_xlsm_without_template_sheet_is_422`,
    `test_replace_latest_xlsx_passes_through_unchanged`,
    `test_replace_latest_requires_auth`,
    `test_generic_create_attachment_converts_stock_list_xlsm`,
    `test_generic_create_attachment_xlsm_without_template_is_422`,
    `test_generic_create_attachment_plain_file_still_uploads`.
    """

    def test_sc_3_replace_latest_stock_list_route_delegates_to_the_shared_service(
        self, client, monkeypatch
    ):
        import app.services.stock_list_archive_service as archive_mod  # not built yet

        c, db, backend = client
        captured: dict = {}

        def _spy(db_arg, **kwargs):
            captured["called"] = True
            captured["kwargs"] = kwargs
            raise RuntimeError(f"{MARKER}-stop-after-spy")

        monkeypatch.setattr(archive_mod, "replace_latest_stock_list", _spy)

        raw = _xlsx_bytes()
        c.post(
            REPLACE_URL,
            files={"file": ("stock.xlsx", raw, _STOCK_XLSX_MEDIA_TYPE)},
        )

        assert captured.get("called") is True, "route did not delegate to the shared archive service"
        kwargs = captured["kwargs"]
        assert kwargs.get("file_bytes") == raw
        assert kwargs.get("filename") == "stock.xlsx"
        assert kwargs.get("user_id") == _USER_ID


# ======================================================================= SC-4


class TestStockApplyArchivesStockList:
    def _seed_attachment_type(self, db) -> str:
        from app.models.resources import AttachmentType

        type_id = str(uuid.uuid4())
        db.add(AttachmentType(id=type_id, type_name="Stock_List", allowed_extensions="xls,xlsx,xlsm",
                               max_file_size_mb=10))
        db.commit()
        return type_id

    def _seed_previous_attachment(self, db, type_id: str) -> str:
        from app.models.resources import Attachment

        att_id = str(uuid.uuid4())
        db.add(Attachment(
            id=att_id, attachment_type_id=type_id, original_filename="old.xlsx",
            stored_filename="old.xlsx", file_path="https://cdn.test/old.xlsx", is_deleted=False,
        ))
        db.commit()
        return att_id

    def _patch_storage(self, monkeypatch) -> _FakeStorageBackend:
        import app.services.storage_router as storage_router

        backend = _FakeStorageBackend()
        monkeypatch.setattr(storage_router, "default_provider", lambda: "s3")
        monkeypatch.setattr(storage_router, "get_backend", lambda provider: backend)
        monkeypatch.setattr(
            storage_router, "cdn_base_url", lambda provider, key: f"https://cdn.test/{key}"
        )
        return backend

    def test_sc_4a_successful_apply_archives_the_new_stock_list(self, task_db, monkeypatch):
        from app.models.company import Company
        from app.models.inventory import Warehouse
        from app.models.resources import Attachment

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        backend = self._patch_storage(monkeypatch)
        type_id = self._seed_attachment_type(db)
        previous_id = self._seed_previous_attachment(db, type_id)
        if not db.query(Company).filter(Company.id == DEFAULT_COMPANY_ID).first():
            db.add(Company(id=DEFAULT_COMPANY_ID, name="Sorento", code="SRT"))
            db.commit()

        active_wh = Warehouse(id=str(uuid.uuid4()), warehouse_code=f"{MARKER}-SC4-FED",
                               warehouse_name="Fed", is_active=True, company_id=DEFAULT_COMPANY_ID)
        db.add(active_wh)
        db.flush()
        _seed_product_with_stock(db, DEFAULT_COMPANY_ID, active_wh.id, code=f"{MARKER}-SC4A", qty=1)

        rows = [
            _stock_row(f"{MARKER}-SC4A", active_wh.warehouse_code, 8),
            _stock_row(f"{MARKER}-SC4B", active_wh.warehouse_code, 2),  # no such product
        ]
        snapshot_id = f"{MARKER}-snap-sc4a"
        header = _stock_header(rows)
        fake.status = (200, {**header, "snapshotId": snapshot_id})
        fake.rows = {1: (200, {"snapshotId": snapshot_id, "page": 1, "pageSize": 1000,
                                "totalPages": 1, "recordCount": len(rows), "rows": rows})}

        job_id = _seed_apply_job(
            db, user_id=str(uuid.uuid4()), company_id=DEFAULT_COMPANY_ID, entity="stock_balances",
            snapshot_id=snapshot_id,
        )

        _run_apply(monkeypatch, factory, job_id)

        row_after = _job_row(db, job_id)
        assert row_after["status"] == "finished", row_after["error"]

        assert len(backend.uploads) == 1, backend.uploads
        uploaded_path, stored_bytes, _ = backend.uploads[0]
        # D3 (small-fix track): a real name, not the old `autocount-pull-pull.xlsx` -
        # lowercase company code + the apply's own local (Asia/Kuala_Lumpur) date.
        from datetime import datetime
        from zoneinfo import ZoneInfo

        today = datetime.now(ZoneInfo("Asia/Kuala_Lumpur")).strftime("%Y%m%d")
        assert uploaded_path == f"stock_list/autocount-stock-list-srt-{today}.xlsx", uploaded_path
        wb = openpyxl.load_workbook(io.BytesIO(stored_bytes))
        ws = wb.active
        header_row = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        assert header_row == ["Item Code", "Item Description", "Location", "On Hand Qty"]
        body = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(body) == 2, body
        codes_in_file = {r[0] for r in body}
        # Every FED row, INCLUDING the one the import skipped as product-not-found.
        assert codes_in_file == {f"{MARKER}-SC4A", f"{MARKER}-SC4B"}

        db.expire_all()
        previous = db.query(Attachment).filter(Attachment.id == previous_id).first()
        assert previous is not None and previous.is_deleted is True

        current = (
            db.query(Attachment)
            .filter(Attachment.attachment_type_id == type_id, Attachment.is_deleted == False)  # noqa: E712
            .all()
        )
        assert len(current) == 1, current
        assert current[0].id != previous_id

    def test_sc_4b_failed_apply_replaces_nothing(self, task_db, monkeypatch):
        from app.models.resources import Attachment

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        backend = self._patch_storage(monkeypatch)
        type_id = self._seed_attachment_type(db)
        previous_id = self._seed_previous_attachment(db, type_id)

        rows = [_stock_row(f"{MARKER}-SC4C", f"{MARKER}-SC4C-WH", 4)]
        snapshot_id = f"{MARKER}-snap-sc4b"
        # AC-SC-2/AC-PC-2 style re-check: the apply must refuse when the FETCHED
        # header carries excludedNonzeroCount > 0, the same guard AC-SP-1 gates
        # Confirm on.
        header = _stock_header(rows, excluded_nonzero_count=2)
        fake.status = (200, {**header, "snapshotId": snapshot_id})
        fake.rows = {1: (200, {"snapshotId": snapshot_id, "page": 1, "pageSize": 1000,
                                "totalPages": 1, "recordCount": len(rows), "rows": rows})}

        job_id = _seed_apply_job(
            db, user_id=str(uuid.uuid4()), company_id=DEFAULT_COMPANY_ID, entity="stock_balances",
            snapshot_id=snapshot_id,
        )

        _run_apply(monkeypatch, factory, job_id)

        row_after = _job_row(db, job_id)
        assert row_after["status"] == "failed", row_after
        # Pins the GUARD reason specifically - "nothing written because nothing
        # runs" (today's `UnsupportedPullEntity`) would ALSO leave `backend.
        # uploads` empty and the previous attachment untouched, so those two
        # assertions alone would already be green before SR4 exists. This is
        # what makes the test red for the guard and not by coincidence.
        assert "exclud" in (row_after["error"] or "").lower(), row_after["error"]

        assert backend.uploads == []
        db.expire_all()
        previous = db.query(Attachment).filter(Attachment.id == previous_id).first()
        assert previous is not None and previous.is_deleted is False

    def test_sc_4c_archive_happens_after_the_import_not_before(self, task_db, monkeypatch):
        from app.models.inventory import Warehouse
        from app.services.inventory_service import StockService

        import app.services.stock_list_archive_service as archive_mod  # not built yet

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        self._patch_storage(monkeypatch)
        self._seed_attachment_type(db)

        active_wh = Warehouse(id=str(uuid.uuid4()), warehouse_code=f"{MARKER}-SC4C-FED",
                               warehouse_name="Fed", is_active=True, company_id=DEFAULT_COMPANY_ID)
        db.add(active_wh)
        db.flush()
        _seed_product_with_stock(db, DEFAULT_COMPANY_ID, active_wh.id, code=f"{MARKER}-SC4CA", qty=1)

        rows = [_stock_row(f"{MARKER}-SC4CA", active_wh.warehouse_code, 5)]
        snapshot_id = f"{MARKER}-snap-sc4c"
        header = _stock_header(rows)
        fake.status = (200, {**header, "snapshotId": snapshot_id})
        fake.rows = {1: (200, {"snapshotId": snapshot_id, "page": 1, "pageSize": 1000,
                                "totalPages": 1, "recordCount": len(rows), "rows": rows})}

        job_id = _seed_apply_job(
            db, user_id=str(uuid.uuid4()), company_id=DEFAULT_COMPANY_ID, entity="stock_balances",
            snapshot_id=snapshot_id,
        )

        order: list[str] = []
        real_bulk_import = StockService.bulk_import_stock

        def _spy_import(self, stock_data, user_id, validate_only=False, outcome=None):
            order.append("import")
            return real_bulk_import(
                self, stock_data, user_id, validate_only=validate_only, outcome=outcome
            )

        real_archive = archive_mod.replace_latest_stock_list

        def _spy_archive(db_arg, **kwargs):
            order.append("archive")
            return real_archive(db_arg, **kwargs)

        monkeypatch.setattr(StockService, "bulk_import_stock", _spy_import)
        monkeypatch.setattr(archive_mod, "replace_latest_stock_list", _spy_archive)

        _run_apply(monkeypatch, factory, job_id)

        assert order == ["import", "archive"], order


# ================================================================= SR4 fix round


class TestStockApplyArchiveFailureIsAVisibleWarning:
    """Captain ruling (SR4 fix round): an archive failure must not fail an apply whose
    stock import already committed (sc_2 pins that), but it must not be silent either -
    the chatbot/n8n would otherwise keep answering from a stale Stock List file with no
    visible sign anything is wrong. `stock_list_not_archived` lands on the PULL job's own
    `metadata.autocount_pull.warnings` (the row `serialize()`/the review page reads),
    never the apply job's."""

    def test_sc_4d_archive_failure_appends_a_warning_to_the_pull_job(self, task_db, monkeypatch):
        from app.models.inventory import Warehouse

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)
        # Deliberately NO Stock_List AttachmentType seeded - reproduces sc_2's own
        # "the import succeeds, the archive step has nothing to write to" case.

        active_wh = Warehouse(id=str(uuid.uuid4()), warehouse_code=f"{MARKER}-SC4D-FED",
                               warehouse_name="Fed", is_active=True, company_id=DEFAULT_COMPANY_ID)
        db.add(active_wh)
        db.flush()
        _seed_product_with_stock(db, DEFAULT_COMPANY_ID, active_wh.id, code=f"{MARKER}-SC4D", qty=1)

        rows = [_stock_row(f"{MARKER}-SC4D", active_wh.warehouse_code, 8)]
        snapshot_id = f"{MARKER}-snap-sc4d"
        header = _stock_header(rows)
        fake.status = (200, {**header, "snapshotId": snapshot_id})
        fake.rows = {1: (200, {"snapshotId": snapshot_id, "page": 1, "pageSize": 1000,
                                "totalPages": 1, "recordCount": len(rows), "rows": rows})}

        pull_job_id = _seed_pull_job(
            db, job_type="autocount_stock_pull", user_id=str(uuid.uuid4()),
            company_id=DEFAULT_COMPANY_ID, entity="stock_balances", company_code="SRT",
            snapshot_id=snapshot_id, phase="confirmed", header=_stored(header),
        )
        job_id = _seed_apply_job(
            db, user_id=str(uuid.uuid4()), company_id=DEFAULT_COMPANY_ID, entity="stock_balances",
            snapshot_id=snapshot_id, pull_job_id=str(pull_job_id),
        )

        _run_apply(monkeypatch, factory, job_id)

        row_after = _job_row(db, job_id)
        assert row_after["status"] == "finished", row_after["error"]
        assert _stock_qty(db, f"{MARKER}-SC4D") == 8, "the stock import itself is real"

        pull_row = _job_row(db, pull_job_id)
        warnings = pull_row["metadata"]["autocount_pull"]["warnings"]
        assert "stock_list_not_archived" in (warnings or []), warnings
