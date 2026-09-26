"""The synthetic TAIYANG price-list workbook, for the cost-price-from-supplier lane
(#1288, `PLAN-cost-price-supplier-26sep.md` section 5.3).

The owner's real file (`TO SORENTO-19&12&22&28&25 series price list 20260917.xlsx`) was
never attached to #1288 or #1291, so this is built from the plan's DESCRIPTION of the
file's shape, not from the file itself: one sheet per series (19, 12, 22, 28, 25), each
with 5 letterhead rows, a header at row 6 (`序号 型号 产品配置 价格`), and 40, 24, 18, 58 and
118 body rows in that order (AC-S1-01). Every code and price is synthetic except the two
the owner quoted verbatim in #1288 (`CB2500SS-BL（彩盒）`, `SRTWT1900-BL-DIY`) and the one
row deliberately shaped for the shared matching engine's separator rung (`CB2500SS GY`).

``taiyang_workbook()`` is the golden input AC-S1-01 to AC-S1-05 parse against, and the
same bytes are written to ``taiyang_price_list.xlsx`` (committed, this directory) for the
browser verification run (AC-S1-24). The smaller builders below it are NOT part of that
golden shape - they build a single minimal sheet each, for the header-detection and
merge edge cases (AC-S1-02, AC-S1-03) that the golden file does not itself exercise
beyond its one 产品配置 merge.
"""
from __future__ import annotations

from io import BytesIO

import openpyxl

LETTERHEAD_TEXT = "XIAMEN TAIYANG TECHNOLOGY CO., LTD"
HEADER_ROW = ("序号", "型号", "产品配置", "价格")

#: Sheet name -> body row count, IN THE ORDER the plan and AC-S1-01 name them.
SHEET_ROW_COUNTS: dict[str, int] = {
    "19 series": 40,
    "12 series": 24,
    "22 series": 18,
    "28 series": 58,
    "25 series": 118,
}

#: The owner's two quoted codes (#1288), placed in the sheets the plan names.
SRTWT_CODE_RAW = " SRTWT1900-BL-DIY "
SRTWT_SHEET = "19 series"
SRTWT_ROW_NO = 10  # 序号 value, not the spreadsheet row

CB2500_CODE_RAW = "CB2500SS-BL（彩盒）"
CB2500_SHEET = "25 series"
CB2500_ROW_NO = 5

#: A code that only binds through the shared engine's separator rung (AC-S1-08):
#: `CB2500SS GY` (a space, the supplier's own separator) for our `CB2500SS-GY`.
SEPARATOR_CODE_RAW = "CB2500SS GY"
SEPARATOR_SHEET = "12 series"
SEPARATOR_ROW_NO = 3

#: 产品配置 merged over rows 7 to 9 of "19 series" (AC-S1-03): the configuration typed once
#: on the anchor row and left blank on the two rows under it.
MERGED_CONFIG_SHEET = "19 series"
MERGED_CONFIG_ANCHOR_SPREADSHEET_ROW = 7  # header is row 6, so this is the first body row
MERGED_CONFIG_SPAN_SPREADSHEET_ROWS = (7, 9)
MERGED_CONFIG_TEXT = "304不锈钢 单把 冷热"


def _letterhead_rows() -> list[list]:
    rows = [[None, None, None, None] for _ in range(5)]
    rows[0][0] = LETTERHEAD_TEXT
    return rows


def _body_row(line_no: int, code: str, configuration: str, price) -> list:
    return [line_no, code, configuration, price]


def _sheet_rows(sheet_name: str, count: int) -> list[list]:
    """The letterhead, the header, and `count` synthetic body rows for one sheet."""
    rows = _letterhead_rows()
    rows.append(list(HEADER_ROW))
    for i in range(1, count + 1):
        code = f"ZZCPC-{sheet_name.split()[0]}-{i:03d}"
        configuration = f"{sheet_name} configuration {i}"
        price = 100 + i
        line_no = i

        if sheet_name == MERGED_CONFIG_SHEET and i == 1:
            configuration = MERGED_CONFIG_TEXT
        if sheet_name == MERGED_CONFIG_SHEET and i in (2, 3):
            # Rows 8 and 9: left blank in the cell, filled by the merge below.
            configuration = None

        # 序号 in this fixture equals its position among the body rows (i), so overriding
        # `line_no` here is redundant with `i` in every case below - stated anyway, so the
        # intent (this specific 序号 carries this specific code) reads without cross-checking.
        if sheet_name == SRTWT_SHEET and i == SRTWT_ROW_NO:
            code = SRTWT_CODE_RAW
            line_no = SRTWT_ROW_NO
        if sheet_name == CB2500_SHEET and i == CB2500_ROW_NO:
            code = CB2500_CODE_RAW
            line_no = CB2500_ROW_NO
        if sheet_name == SEPARATOR_SHEET and i == SEPARATOR_ROW_NO:
            code = SEPARATOR_CODE_RAW
            line_no = SEPARATOR_ROW_NO

        rows.append(_body_row(line_no, code, configuration, price))
    return rows


def taiyang_workbook() -> bytes:
    """The 5-sheet synthetic TAIYANG price list (AC-S1-01 to AC-S1-05 golden input)."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name, count in SHEET_ROW_COUNTS.items():
        ws = wb.create_sheet(sheet_name)
        for row in _sheet_rows(sheet_name, count):
            ws.append(row)
        if sheet_name == MERGED_CONFIG_SHEET:
            start, end = MERGED_CONFIG_SPAN_SPREADSHEET_ROWS
            ws.merge_cells(start_row=start, start_column=3, end_row=end, end_column=3)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def header_at_row(row_number: int, *, sheet_name: str = "Sheet") -> bytes:
    """One sheet whose header (序号 型号 产品配置 价格) sits at `row_number` instead of 6
    (AC-S1-02): `row_number - 1` blank/letterhead rows above it, 3 body rows below."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for r in range(1, row_number):
        ws.append([None, None, None, None])
    ws.append(list(HEADER_ROW))
    for i in range(1, 4):
        ws.append([i, f"ZZCPC-HDR{row_number}-{i:03d}", f"config {i}", 100 + i])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sheet_with_no_header(*, sheet_name: str = "no header") -> bytes:
    """A sheet with none of the four fields resolvable in rows 1 to 20 (AC-S1-02): the
    reader must record it as skipped by name, not raise."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for r in range(1, 15):
        ws.append([f"random text {r}", None, None, None])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def merged_price_workbook(*, sheet_name: str = "merged price") -> bytes:
    """One sheet where 价格 is merged over two body rows (AC-S1-03): both rows must carry
    the price with `price_from_merge`, and `型号` (item_code) is never filled down."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in _letterhead_rows():
        ws.append(row)
    ws.append(list(HEADER_ROW))
    ws.append([1, "ZZCPC-MP-001", "shared configuration", 250])
    ws.append([2, "ZZCPC-MP-002", "shared configuration", None])
    ws.merge_cells(start_row=7, start_column=4, end_row=8, end_column=4)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def price_cell_workbook(prices: list) -> bytes:
    """One row per entry in `prices` (AC-S1-05): a price cell holding each of the given
    raw values, so the reader's cleaning can be asserted per value without a full sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "prices"
    for row in _letterhead_rows():
        ws.append(row)
    ws.append(list(HEADER_ROW))
    for i, value in enumerate(prices, start=1):
        ws.append([i, f"ZZCPC-PRICE-{i:03d}", "configuration", value])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def simple_price_list_workbook(
    rows: list[tuple],
    *,
    sheet_name: str = "Sheet",
    letterhead: str | None = LETTERHEAD_TEXT,
) -> bytes:
    """A minimal one-sheet workbook of `(supplier_code, configuration, price)` rows, for
    tests that need a specific code/price shape (matching, duplicates, apply) without the
    full 258-row golden fixture."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for r in _letterhead_rows():
        ws.append(r)
    if letterhead is not None:
        ws.cell(row=1, column=1, value=letterhead)
    ws.append(list(HEADER_ROW))
    for i, (code, configuration, price) in enumerate(rows, start=1):
        ws.append([i, code, configuration, price])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    # Regenerates the committed golden fixture. Run from `sorento_crm_backend/`:
    #   venv/bin/python -m tests.fixtures.cost_price.taiyang_shapes
    import pathlib

    out = pathlib.Path(__file__).parent / "taiyang_price_list.xlsx"
    out.write_bytes(taiyang_workbook())
    print(f"wrote {out}")
