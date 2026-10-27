"""SALES-UNIFIED S3: the reports (AC-18, AC-20 to AC-23).

Workbook, Yearly comparison on the invoiced basis, Top selling items, server-side sorting and
access, over the report kernel (contract sections 4 and 5). Postgres only, every row seeded here.
"""
from __future__ import annotations

import io
import uuid
from datetime import date

import pytest

from app.main import app  # noqa: F401  (import order)
from app.models.base import set_company_scope
from app.services.company_scope import DEFAULT_COMPANY_ID
from tests._mc_lookup_seed import MOCHA_ID, product, seed_mocha
from tests._pg_fixture import blank_session, unique_code
from tests._sales_unified_seed import (
    ALL_TYPES, D, JAN, SALES_TYPE_LABELS, agent, cust, definition_of, doc, grand, line, run, six, so, view,
)
from tests.test_sales_yearly_routes import SLUG, _client, _user

JAN15 = date(2026, 1, 15)
RUN = "/api/v1/reports/sales_by_debtor_type/run"


@pytest.fixture
def db():
    with blank_session() as s:
        set_company_scope(s, frozenset({DEFAULT_COMPANY_ID}))
        seed_mocha(s)
        yield s


def _dealer_invoice(db, net="1000.00", when=JAN15, customer=None, agent_row=None):
    _h, (so_line,) = so(db, demand_class="retail", customer=customer)
    return doc(db, "invoice", net, when, customer=customer, agent_row=agent_row,
               lines=[line(so_line=so_line)])


# ------------------------------------------------------------------------------ AC-18


def _workbook(db, period):
    from openpyxl import load_workbook

    from app.services.reports import engine
    from app.services.reports.xlsx_renderer import render_workbook

    definition = definition_of("sales_by_debtor_type")
    v = view(period=period)
    data = engine.run_workbook(db, definition, v.params, v,
                               company_grants=frozenset({DEFAULT_COMPANY_ID}))
    return load_workbook(io.BytesIO(render_workbook(definition, data)))


def _texts(sheet):
    return [c.value for row in sheet.iter_rows() for c in row if isinstance(c.value, str)]


def _numbers(sheet):
    return [D(str(c.value)) for row in sheet.iter_rows() for c in row
            if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)]


def test_ac_18_the_excel_has_the_blocks_on_the_summary_sheet_and_one_sheet_per_month(db):
    a = agent(db, unique_code("ZZTA"))
    jan = _dealer_invoice(db, net="123.45", when=JAN15, agent_row=a)
    feb = doc(db, "cash_sale", "67.89", date(2026, 2, 3), agent_row=a)
    book = _workbook(db, {"kind": "custom", "from": "2026-01-01", "to": "2026-02-28"})
    assert book.sheetnames == ["SUMMARY", "JAN'26", "FEB'26"]

    summary = book["SUMMARY"]
    texts = _texts(summary)
    assert "WEEKLY SALES BY DEBTOR TYPE BY SALES AGENT" in texts
    titles = [f"SORENTO - {label.upper()}" for label in SALES_TYPE_LABELS.values()]
    positions = [texts.index(t) for t in titles]
    assert positions == sorted(positions), "the six blocks, in the sales type order"
    assert D("123.45") in _numbers(summary) and D("67.89") in _numbers(summary)
    assert a.sales_agent in texts

    assert jan.doc_no in _texts(book["JAN'26"]) and feb.doc_no not in _texts(book["JAN'26"])
    assert feb.doc_no in _texts(book["FEB'26"]) and jan.doc_no not in _texts(book["FEB'26"])


# ------------------------------------------------------------------------------ AC-20


def _yearly(db, **over):
    from app.schemas.report import ReportViewConfig
    from app.services.reports import engine

    params = {"date_basis": "order_date", "period": JAN, "company": [DEFAULT_COMPANY_ID],
              "basis": ["invoiced"], "channel": ALL_TYPES}
    params.update(over)
    v = ReportViewConfig.model_validate({
        "params": params, "detail": {"columns": [], "order": []},
        "pivot": {"rows": "year", "cols": "month_of_year", "measures": ["sales_value"]},
    })
    return engine.run(db, definition_of("sales_yearly"), v.params, v,
                      company_grants=frozenset({DEFAULT_COMPANY_ID}))


def test_ac_20_basis_defaults_to_invoiced_and_the_channel_options_are_the_sales_types(db):
    from app.services.reports import engine

    definition = definition_of("sales_yearly")
    basis = next(p for p in definition.params if p.key == "basis")
    channel = next(p for p in definition.params if p.key == "channel")
    assert basis.default == ("invoiced",)
    assert six(list(channel.options(db))) == list(SALES_TYPE_LABELS.items())
    assert six(list(channel.default)) == ALL_TYPES
    assert engine.view_config(definition).params["basis"] == ["invoiced"]


def test_ac_20_on_invoiced_the_blocks_are_the_sales_types(db):
    _dealer_invoice(db, net="10.00")
    doc(db, "cash_sale", "4.00", JAN15)
    result = _yearly(db)
    assert [b.title for b in result.layouts.blocks if b.key != "other"] == [
        f"SORENTO - {label.upper()}" for label in SALES_TYPE_LABELS.values()]
    by_key = {b.key: b.summary.grand_total.get("sales_value") for b in result.layouts.blocks}
    assert by_key["dealer"] == "10.00" and by_key["no_sales_order"] == "4.00"
    assert result.note.startswith("Basis: Invoiced")


def _order_line(db, total, demand_class):
    from app.models.order import SalesOrder, SalesOrderLine

    header = SalesOrder(id=str(uuid.uuid4()), so_number=unique_code("SO"), order_date=JAN15,
                        status="open", demand_class=demand_class, company_id=DEFAULT_COMPANY_ID)
    db.add(header)
    db.flush()
    db.add(SalesOrderLine(id=str(uuid.uuid4()), sales_order_id=header.id,
                          product_id=product(db, company_id=DEFAULT_COMPANY_ID).id, qty_ordered=1,
                          qty_delivered=1, line_total=D(total), line_status="open",
                          company_id=DEFAULT_COMPANY_ID))
    db.flush()


@pytest.mark.parametrize("basis, word", [("delivered", "Delivered"), ("ordered", "Ordered")])
def test_ac_20_the_order_bases_still_run_and_say_so_in_the_basis_line(db, basis, word):
    _order_line(db, "10.00", "retail")
    _order_line(db, "20.00", "project")
    _order_line(db, "5.00", None)
    result = _yearly(db, basis=[basis], channel=["dealer", "project"])
    assert result.note.startswith(f"Basis: {word}")
    assert result.layouts.summary.grand_total["sales_value"] == "30.00"


def test_ac_20_on_the_order_bases_type_not_set_is_a_null_class_and_the_other_three_match_nothing(db):
    _order_line(db, "10.00", "retail")
    _order_line(db, "5.00", None)
    assert _yearly(db, basis=["delivered"], channel=["type_not_set"]).layouts.summary.grand_total == {
        "sales_value": "5.00"}
    for nothing in ("set_apart", "mixed", "no_sales_order"):
        assert _yearly(db, basis=["delivered"], channel=[nothing]).row_count == 0


# ------------------------------------------------------------------------------ AC-21


def _items(result):
    """{item code: (quantity, sales_value, documents)} from the item pivot's row totals."""
    out = {}
    for label, totals in result.layouts.summary.row_totals.items():
        out[label.split(" - ")[0]] = (
            D(totals["quantity"]), D(totals["sales_value"]), D(totals["documents"]))
    return out


def _top(db, **kw):
    return run(db, "top_selling_items", rows="item", cols="year",
               measures=("quantity", "sales_value", "documents"), **kw)


def test_ac_21_a_credit_note_line_nets_against_an_invoice_line_of_the_same_item(db):
    p = product(db, company_id=DEFAULT_COMPANY_ID)
    doc(db, "invoice", "250.00", JAN15, lines=[line(qty="5", net="250.00", prod=p)])
    doc(db, "credit_note", "100.00", JAN15, lines=[line(qty="2", net="100.00", prod=p)])
    doc(db, "invoice", "900.00", JAN15, status="cancelled", lines=[line(qty="9", net="900.00", prod=p)])
    assert _items(_top(db)) == {p.product_code: (D("3"), D("150.00"), D("2"))}


def test_ac_21_an_item_code_not_in_the_crm_still_lists_by_its_code(db):
    p = product(db, company_id=DEFAULT_COMPANY_ID)
    code = unique_code("ZZTGONE")
    doc(db, "invoice", "40.00", JAN15, lines=[line(qty="1", net="40.00", item_code=code, description="Old tile")])
    doc(db, "invoice", "10.00", JAN15, lines=[line(qty="1", net="10.00", prod=p)])
    items = _items(_top(db))
    assert items[code][:2] == (D("1"), D("40.00"))
    assert items[p.product_code][:2] == (D("1"), D("10.00"))


def test_ac_21_the_item_dimension_is_code_dash_description_or_the_code_alone(db):
    p = product(db, company_id=DEFAULT_COMPANY_ID)
    q = product(db, company_id=DEFAULT_COMPANY_ID)
    doc(db, "invoice", "5.00", JAN15, lines=[line(qty="1", net="5.00", prod=p, description="Basin")])
    doc(db, "invoice", "6.00", JAN15, lines=[line(qty="1", net="6.00", prod=q)])
    labels = set(_top(db).layouts.summary.row_values)
    assert labels == {f"{p.product_code} - Basin", q.product_code}


def test_ac_21_a_line_amount_is_net_amount_times_the_documents_currency_rate(db):
    p = product(db, company_id=DEFAULT_COMPANY_ID)
    doc(db, "invoice", "100.00", JAN15, rate="2", lines=[line(qty="1", net="50.00", prod=p)])
    assert _items(_top(db))[p.product_code][1] == D("100.00")


def test_ac_21_the_report_ranks_by_rm_and_hides_customer_and_category_from_the_meta(db):
    definition = definition_of("top_selling_items")
    keys = [x.key for x in definition.params]
    assert {"company", "channel", "agent", "brand", "debtor_type", "period", "customer", "category"} <= set(keys)
    assert definition.default_view["pivot"]["sort"] == {"key": "__total__", "dir": "desc"}
    assert definition.default_view["pivot"]["rows"] == "item"


# ------------------------------------------------------------------------------ AC-22


def _agent_rows(db):
    """Three agents on no-sales-order cash sales, two debtor types:
    A: T1 10        total 10
    B: T1 40, T2 5  total 45
    C: T2 30        total 30"""
    a, b, c = (agent(db, unique_code(f"ZZTSORT{n}")) for n in "ABC")
    t1, t2 = cust(db, debtor_type="ZZTT1"), cust(db, debtor_type="ZZTT2")
    doc(db, "cash_sale", "10.00", JAN15, customer=t1, agent_row=a)
    doc(db, "cash_sale", "40.00", JAN15, customer=t1, agent_row=b)
    doc(db, "cash_sale", "5.00", JAN15, customer=t2, agent_row=b)
    doc(db, "cash_sale", "30.00", JAN15, customer=t2, agent_row=c)
    doc(db, "cash_sale", "1.00", JAN15, customer=t2)  # no agent: sorts last either way
    return a, b, c


def _rows_in_order(db, sort):
    result = run(db, "sales_by_debtor_type", channel=["no_sales_order"], pivot_sort=sort)
    block = result.layouts.blocks[0].summary
    assert list(result.layouts.summary.row_values) == list(block.row_values)
    return block.row_values


@pytest.mark.parametrize("direction", ["desc", "asc"])
def test_ac_22_the_pivot_sorts_by_the_row_total_and_the_blank_row_stays_last(db, direction):
    a, b, c = _agent_rows(db)
    got = _rows_in_order(db, {"key": "__total__", "dir": direction})
    named = [b.sales_agent, c.sales_agent, a.sales_agent]
    assert got == (named if direction == "desc" else named[::-1]) + ["(blank)"]


def test_ac_22_the_pivot_sorts_by_the_row_label(db):
    a, b, c = _agent_rows(db)
    assert _rows_in_order(db, {"key": "__row__", "dir": "asc"}) == [
        a.sales_agent, b.sales_agent, c.sales_agent, "(blank)"]
    assert _rows_in_order(db, {"key": "__row__", "dir": "desc"}) == [
        c.sales_agent, b.sales_agent, a.sales_agent, "(blank)"]


def test_ac_22_the_pivot_sorts_by_a_column_value_and_a_missing_cell_sorts_last(db):
    a, b, c = _agent_rows(db)
    # T1: B 40, A 10, C has no T1 cell, the unnamed row neither.
    assert _rows_in_order(db, {"key": "ZZTT1", "dir": "desc"})[:2] == [b.sales_agent, a.sales_agent]
    assert _rows_in_order(db, {"key": "ZZTT1", "dir": "asc"})[:2] == [a.sales_agent, b.sales_agent]
    assert set(_rows_in_order(db, {"key": "ZZTT1", "dir": "asc"})[-2:]) == {c.sales_agent, "(blank)"}
    assert set(_rows_in_order(db, {"key": "ZZTT1", "dir": "desc"})[-2:]) == {c.sales_agent, "(blank)"}


def test_ac_22_the_detail_sorts_by_a_catalogue_column(db):
    from app.services.reports import engine

    early = doc(db, "invoice", "30.00", date(2026, 1, 3))
    late = doc(db, "invoice", "10.00", date(2026, 1, 20))
    mid = doc(db, "invoice", "20.00", date(2026, 1, 10))

    def order(key, direction):
        v = view(detail_sort={"key": key, "dir": direction})
        result = engine.run(db, definition_of("sales_by_debtor_type"), v.params, v,
                            company_grants=frozenset({DEFAULT_COMPANY_ID}))
        return [r["document_no"] for r in result.layouts.detail.rows]

    assert order("document_date", "asc") == [early.doc_no, mid.doc_no, late.doc_no]
    assert order("document_date", "desc") == [late.doc_no, mid.doc_no, early.doc_no]
    assert order("sales_value", "asc") == [late.doc_no, mid.doc_no, early.doc_no]
    assert order("document_no", "asc") == sorted([early.doc_no, mid.doc_no, late.doc_no])


def _post(client, v):
    return client.post(RUN, json={"params": v.params, "view": v.model_dump(mode="json")})


def test_ac_22_an_unknown_sort_key_is_422_through_the_run_route(db):
    principal = _user(db, DEFAULT_COMPANY_ID)
    with _client(db, principal) as client:
        ok = _post(client, view(pivot_sort={"key": "__total__", "dir": "desc"}))
        assert ok.status_code == 200, ok.text
        pivot = _post(client, view(pivot_sort={"key": "no_such_column", "dir": "asc"}))
        detail = _post(client, view(detail_sort={"key": "no_such_column", "dir": "asc"}))
    assert pivot.status_code == 422, pivot.text
    assert detail.status_code == 422, detail.text


# ------------------------------------------------------------------------------ AC-23


def _body(company=DEFAULT_COMPANY_ID):
    return {"params": {"company": [company], "period": JAN}, "view": None}


def test_ac_23_without_sales_reports_view_the_run_is_403(db):
    principal = _user(db, DEFAULT_COMPANY_ID)
    with _client(db, principal, allow=()) as client:
        assert client.post(RUN, json=_body()).status_code == 403
        keys = [r["key"] for r in client.get("/api/v1/reports").json()["reports"]]
    assert "sales_by_debtor_type" not in keys


def test_ac_23_a_holder_sees_the_report_in_the_catalogue_and_can_run_it(db):
    principal = _user(db, DEFAULT_COMPANY_ID)
    with _client(db, principal, allow=(SLUG,)) as client:
        keys = [r["key"] for r in client.get("/api/v1/reports").json()["reports"]]
        assert "sales_by_debtor_type" in keys and "top_selling_items" in keys
        assert client.post(RUN, json=_body()).status_code == 200


def test_ac_23_a_company_outside_the_grant_is_403_and_a_sorento_run_holds_no_mocha_row(db):
    seed = _dealer_invoice(db, net="10.00")
    assert seed is not None
    c = cust(db, company_id=MOCHA_ID)
    doc(db, "cash_sale", "700.00", JAN15, customer=c, company_id=MOCHA_ID)
    principal = _user(db, DEFAULT_COMPANY_ID)
    with _client(db, principal) as client:
        assert client.post(RUN, json=_body(company=MOCHA_ID)).status_code == 403
        ok = client.post(RUN, json=_body())
    assert ok.status_code == 200, ok.text
    assert ok.json()["layouts"]["summary"]["grand_total"]["sales_value"] == "10.00"
