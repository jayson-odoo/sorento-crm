"""One house style for exported workbooks: the LOW STOCK REPORT's (OI-DOWNLOAD, owner 21 Oct
2026: "our excel format needs to be nicer ... like the LOW STOCK REPORT, we should unify
this"). `app.services.excel_style` holds it; the low stock report and the order summary keep
their exact look, and the Order Inquiry workbook (the email attachment and both downloads)
now wears it, with its columns and values unchanged.
"""
from __future__ import annotations

import io
from datetime import date

import openpyxl
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.services import excel_style
from app.services.order_inquiry_worklist_service import (
    _EXPORT_DEFAULT_WIDTH,
    EXPORT_HEADINGS,
    EXPORT_TITLE,
    WORKSHEET_HEADINGS,
    OrderInquiryWorklistService,
    _style_sheet,
)


def _assert_house_header(cell) -> None:
    assert cell.font.bold and cell.font.color.rgb == "FFFFFFFF", cell.coordinate
    assert cell.fill.fgColor.rgb == "FF404040", cell.coordinate
    _assert_house_body(cell)


def _assert_house_body(cell) -> None:
    assert cell.border.left.style == "thin" and cell.border.bottom.style == "thin", cell.coordinate
    assert cell.alignment.wrap_text and cell.alignment.vertical == "top", cell.coordinate


def test_low_stock_workbook_keeps_its_look():
    from app.services.scm import low_stock_report_service as ls

    view = {
        "columns": ls.LOW_STOCK_COLUMNS,
        "rows": [tuple(range(16))],
        "sheets": [{"title": "Low stock", "row_indexes": [0]}],
    }
    wb = openpyxl.load_workbook(io.BytesIO(ls._write_workbook(view, list(ls._LOW_STOCK_WIDTHS))))
    ws = wb["Low stock"]

    _assert_house_header(ws["A1"])
    _assert_house_body(ws["B2"])
    assert not ws["B2"].font.bold
    assert ws.freeze_panes == "A2"
    assert ws.column_dimensions["B"].width == 42


def test_order_summary_sheet_keeps_its_look():
    from app.services.scm.summary_order_service import write_sheet

    wb = Workbook()
    write_sheet(wb.active, ("A", "B"), [(1, "x")], {"A": 10, "B": 20})
    ws = wb.active

    _assert_house_header(ws["B1"])
    _assert_house_body(ws["A2"])
    assert ws.freeze_panes == "A2"
    assert ws.column_dimensions["B"].width == 20


def _oi_workbook(rows, columns=EXPORT_HEADINGS):
    workbook = Workbook()
    workbook.remove(workbook.active)
    OrderInquiryWorklistService(db=None)._write_sheet(workbook, EXPORT_TITLE, rows, columns=columns)
    return workbook[EXPORT_TITLE]


_ROW = {
    "so_date": date(2025, 12, 1),
    "so_number": "SO380096",
    "item_code": "CWCY605",
    "qty": "5",
    "delivery_date": date(2026, 8, 3),
    "project_customer": "SHING SOON HENG HARDWARE SDN BHD / LOT 1940 / MUAR",
    "supplier": None,
    "location": "BRW-BB",
}


def _col(heading, columns=EXPORT_HEADINGS):
    """The 1-based column a heading prints in, so these tests follow the owner-approved order
    (REQUESTED RESERVE / RESERVED beside QTY, OI-RESERVE-PILL-QTY 21 Oct) rather than pin it."""
    return columns.index(heading) + 1


def _letter(heading, columns=EXPORT_HEADINGS):
    return get_column_letter(_col(heading, columns))


def test_oi_workbook_wears_the_house_style_with_its_columns_unchanged():
    first = {**_ROW, "asked_reserve_qty": "70", "reserved_qty": "0"}
    second = {**_ROW, "so_number": "SO394881", "change_date": date(2026, 11, 2)}
    ws = _oi_workbook([first, second])

    # Same title row and headings as before, in the same places.
    assert ws["A1"].value == EXPORT_TITLE and ws["A1"].font.bold
    assert [cell.value for cell in ws[2]] == list(EXPORT_HEADINGS)
    headings = ("S/O NO", "ITEM CODE", "QTY", "REQUESTED RESERVE", "RESERVED", "LOCATION")
    assert [ws.cell(row=3, column=_col(h)).value for h in headings] == [
        "SO380096", "CWCY605", 5.0, 70.0, None, "BRW-BB",
    ]
    # The reserve figures sit right beside QTY (owner ruling, 21 Oct 2026).
    assert EXPORT_HEADINGS[_col("QTY"):_col("QTY") + 2] == ("REQUESTED RESERVE", "RESERVED")
    # The low stock look: dark headings, every table cell bordered + wrapped, frozen
    # under the headings, explicit widths, dates dd/mm/yyyy.
    for cell in ws[2]:
        _assert_house_header(cell)
    for row in (3, 4):
        for cell in ws[row]:
            _assert_house_body(cell)
    assert ws.freeze_panes == "A3"
    assert ws.column_dimensions[_letter("PROJECT/CUSTOMER")].width == 40
    assert ws[f"{_letter('SO DATE')}3"].number_format == excel_style.DATE_FORMAT
    assert ws[f"{_letter('DELIVERY DATE')}3"].number_format == excel_style.DATE_FORMAT
    assert ws[f"{_letter('QTY')}3"].number_format == "General"  # QTY stays a plain number
    change_date = ws.cell(row=4, column=_col("CHANGE DATE"))
    assert change_date.value is not None
    assert change_date.number_format == excel_style.DATE_FORMAT


def test_oi_workbook_widths_follow_the_heading_not_the_position():
    """A cut column list keeps each heading's width; a heading the map does not know (a
    column added later) gets the default rather than its neighbour's."""
    # The planning-board worksheet's cut list puts PROJECT/CUSTOMER and LOCATION in other
    # columns than the full sheet does; each keeps its own width there.
    ws = _oi_workbook([dict(_ROW)], columns=WORKSHEET_HEADINGS)

    assert ws.column_dimensions[_letter("PROJECT/CUSTOMER", WORKSHEET_HEADINGS)].width == 40
    assert ws.column_dimensions[_letter("SUPPLIER", WORKSHEET_HEADINGS)].width == 24
    assert ws.cell(row=3, column=_col("LOCATION", WORKSHEET_HEADINGS)).value == "BRW-BB"

    # Values are printed by heading, so an unknown heading has no value to print; the style
    # alone decides its width. It sits after REMARK (width 30) and must not inherit that.
    columns = EXPORT_HEADINGS + ("NEW COLUMN",)
    workbook = Workbook()
    sheet = workbook.active
    sheet.append([EXPORT_TITLE])
    sheet.append(list(columns))
    _style_sheet(sheet, columns)

    new = _letter("NEW COLUMN", columns)
    assert sheet.column_dimensions[new].width == _EXPORT_DEFAULT_WIDTH
    assert sheet[f"{new}2"].fill.fgColor.rgb == "FF404040"
