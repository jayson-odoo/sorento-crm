"""DO-OWNERSHIP-GUARD: one two-way field-ownership list for AutoCount DOs vs Order Tracking.

Plan: documentation/plans/autocount/PLAN-do-ownership-guard.md
UAC:  documentation/plans/autocount/do-ownership-guard-acceptance-criteria.md

A row is AutoCount-owned when `doc_key IS NOT NULL`. The ingest-side ACs (AC-OG04, AC-OG06,
AC-OG07) live in tests/test_ingest_autocount_do_grn.py, next to the ingest harness.
"""
from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi import HTTPException
from openpyxl import Workbook

from app.models.order import Order, OrderLine, OrderStatus
from app.models.inventory import Warehouse
from app.models.product import Product, ProductCategory, UnitOfMeasure
from app.schemas.order import OrderLineCreate, OrderLineUpdate, OrderResponse, OrderUpdate
from app.services.order_field_ownership import (
    AUTOCOUNT_OWNED_ORDER_COLUMNS,
    ORDER_TRACKING_OWNED_ORDER_COLUMNS,
    SYSTEM_ORDER_COLUMNS,
)
from app.services.order_service import OrderService
from tests._pg_fixture import blank_session, unique_code

_USER = "11111111-1111-1111-1111-111111111111"

AUTOCOUNT_VALUES = dict(
    order_date=date(2026, 9, 27), created_time=datetime(2026, 9, 27, 8, 30),
    debtor_code="300-AC", debtor_name="AutoCount Name", agent="AC-AGENT", is_cancelled=False,
    remarks="ac remarks", subtotal_amount=Decimal("100.00"), discount_amount=Decimal("0.00"),
    tax_amount=Decimal("6.00"), total_amount=Decimal("106.00"),
)


@pytest.fixture()
def db():
    with blank_session() as session:
        for code, name in [("NEW", "New Order"), ("DELIVERED", "Delivered")]:
            session.add(OrderStatus(status_code=code, status_name=name))
        session.commit()
        yield session


def _autocount_order(db, number="ZZDO-OG-1") -> Order:
    order = Order(order_number=number, doc_key=990101, source_book="db1", **AUTOCOUNT_VALUES)
    db.add(order)
    db.commit()
    return order


def _plain_order(db, number="ZZDO-OG-2") -> Order:
    order = Order(order_number=number, **AUTOCOUNT_VALUES)
    db.add(order)
    db.commit()
    return order


def _line_fixture(db, order: Order) -> tuple[str, str, OrderLine]:
    category = ProductCategory(category_code=unique_code("OG"), category_name="c")
    uom = UnitOfMeasure(uom_code=unique_code("OG"), uom_name="u")
    db.add_all([category, uom])
    db.flush()
    product = Product(product_code=unique_code("OG"), product_name="p", category_id=category.id,
                      base_uom_id=uom.id, list_price=1)
    wh = Warehouse(warehouse_code=unique_code("OG"), warehouse_name="wh")
    db.add_all([product, wh])
    db.flush()
    line = OrderLine(order_id=order.id, line_sequence=1, product_id=product.id,
                     warehouse_id=wh.id, quantity=Decimal("10"))
    db.add(line)
    db.commit()
    return str(product.id), str(wh.id), line


def _snapshot(order: Order, columns) -> dict:
    return {c: getattr(order, c) for c in columns}


# ======================================================================= the registry
def test_registry_partitions_every_orders_column():
    """AC-OG01: every `orders` column has exactly one owner."""
    sets = [AUTOCOUNT_OWNED_ORDER_COLUMNS, ORDER_TRACKING_OWNED_ORDER_COLUMNS, SYSTEM_ORDER_COLUMNS]
    model_columns = {c.name for c in Order.__table__.columns}
    assert set().union(*sets) == model_columns
    assert sum(len(s) for s in sets) == len(model_columns)


def test_sheet_mappings_follow_registry():
    """AC-OG02: the Master columns AutoCount owns are exactly the old hard-coded skip list, and
    no Overall Tracking column is AutoCount-owned."""
    master_owned = {"order_date", "created_time", "debtor_code", "debtor_name", "agent",
                    "is_cancelled", "customer_id"}
    assert master_owned <= AUTOCOUNT_OWNED_ORDER_COLUMNS
    for col in ("remarks_cs", "order_type", "estimated_delivery_date", "actual_delivery_date",
                "pickup_time", "checker", "transporter", "transporter_id", "driver_name",
                "lorry_plate", "customer_ref", "delivery_remarks_cs", "delivery_remarks",
                "salesman", "trips", "warehouse", "delivery_days", "kpi_warning",
                "order_status_id"):
        assert col in ORDER_TRACKING_OWNED_ORDER_COLUMNS, col


# ======================================================================= Master sheet
def _master_workbook(header: list, rows: list[list]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    master = wb.create_sheet("Master")
    master.append(header)
    for r in rows:
        master.append(r)
    tracking = wb.create_sheet("Overall Tracking")
    tracking.append(["Doc Number", "Date", "Time", "Checker", "Transporter", "Driver Name",
                     "Lorry Plate", "Customer", "Remarks CS", "Remarks", "Salesman", "Trips",
                     "W/H"])
    for r in rows:
        tracking.append([r[0], date(2026, 9, 29), "2:30 PM", "Chong", "ZZ TRANS", "Ali", "WXX 1",
                         "cust ref", "dcs", "dr", "SEAN", 2, "BRW"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


MASTER_HEADER = ["Doc. No.", "Date", "Created Time", "Debtor Code", "Debtor Name", "Agent",
                 "Cancel", "Remarks CS", "Type"]


def test_master_keeps_created_time_and_cancel_on_autocount_row(db):
    """AC-OG03: Master Created Time and Cancel=Y are ignored on an AutoCount row."""
    order = _autocount_order(db)
    data = _master_workbook(MASTER_HEADER, [
        ["ZZDO-OG-1", date(2026, 9, 1), "01/09/2026 10:00", "999", "Sheet", "S-AG", "Y", "cs",
         "RMA"],
    ])
    result = OrderService(db).import_excel_tracking(data, _USER)
    assert not result.get("errors"), result.get("errors")
    db.refresh(order)
    assert order.created_time == datetime(2026, 9, 27, 8, 30)
    assert order.is_cancelled is False
    assert (order.remarks_cs, order.order_type) == ("cs", "RMA")


def test_unowned_row_gets_every_master_and_tracking_column(db):
    """AC-OG05: an RMA / non-AutoCount row (doc_key NULL) is written in full by both sheets."""
    order = _plain_order(db, "ZZRMA-OG-1")
    data = _master_workbook(MASTER_HEADER, [
        ["ZZRMA-OG-1", date(2026, 9, 1), "01/09/2026 10:00", "999-SHEET", "Sheet Debtor",
         "S-AG", "Y", "cs rma", "RMA"],
    ])
    result = OrderService(db).import_excel_tracking(data, _USER)
    assert not result.get("errors"), result.get("errors")
    db.refresh(order)
    assert order.order_date == date(2026, 9, 1)
    assert order.created_time == datetime(2026, 9, 1, 10, 0)
    assert (order.debtor_code, order.debtor_name, order.agent) == ("999-SHEET", "Sheet Debtor", "S-AG")
    assert order.is_cancelled is True
    assert order.customer_id is not None
    assert (order.remarks_cs, order.order_type) == ("cs rma", "RMA")
    assert order.actual_delivery_date == date(2026, 9, 29)
    assert (order.checker, order.transporter, order.driver_name, order.lorry_plate) == (
        "Chong", "ZZ TRANS", "Ali", "WXX 1")
    assert order.transporter_id is not None
    assert (order.customer_ref, order.delivery_remarks_cs, order.delivery_remarks) == (
        "cust ref", "dcs", "dr")
    assert (order.salesman, order.trips, order.warehouse) == ("SEAN", 2, "BRW")


# ======================================================================= JSON bulk import
def test_bulk_import_skips_autocount_owned_keys(db):
    """AC-OG08: on an AutoCount row the JSON bulk import writes only Order Tracking keys and
    reports what it kept; a plain row is written in full."""
    ac = _autocount_order(db)
    plain = _plain_order(db)
    row = {
        "Order Date": "2026-01-05", "Actual Delivery Date": "2026-09-30",
        "Subtotal Amount": "1", "Tax Amount": "0", "Remarks": "sheet remarks",
        "Customer ID": None,
    }
    result = OrderService(db).bulk_import_orders(
        [{"Order Number": "ZZDO-OG-1", **row}, {"Order Number": "ZZDO-OG-2", **row}], _USER)
    assert result["errors"] == []
    assert result["updated"] == 2
    db.refresh(ac)
    db.refresh(plain)
    assert _snapshot(ac, AUTOCOUNT_VALUES) == AUTOCOUNT_VALUES
    assert ac.actual_delivery_date == date(2026, 9, 30)
    assert plain.order_date == date(2026, 1, 5) and plain.total_amount == Decimal("1.00")
    assert plain.remarks == "sheet remarks"
    warnings = result["warnings"]
    assert len(warnings) == 1
    assert "ZZDO-OG-1" in warnings[0] and "owned by AutoCount" in warnings[0]
    named = set(warnings[0].split(" - ", 1)[1].split(" owned by")[0].split(", "))
    # The row sent no total: the import computes one and keeps AutoCount's, but does not claim
    # the row asked to change it.
    assert named == {"order_date", "subtotal_amount", "tax_amount", "remarks", "customer_id"}


def test_bulk_import_response_carries_warnings():
    """AC-OG08: the route's response_model keeps the warnings (an undeclared field is dropped)."""
    from app.schemas.order import BulkImportResponse

    body = BulkImportResponse(created=0, updated=1, errors=[], warnings=["w"]).model_dump()
    assert body["warnings"] == ["w"]


# ======================================================================= manual edits
def test_update_order_rejects_autocount_owned_field(db):
    """AC-OG09: a manual edit naming an AutoCount-owned field on an AutoCount row is a 409
    naming the field, and nothing is written."""
    order = _autocount_order(db)
    with pytest.raises(HTTPException) as exc:
        OrderService(db).update_order(
            str(order.id),
            OrderUpdate(debtor_name="Edited", driver_name="Bala", total_amount=Decimal("1")),
            _USER,
        )
    assert exc.value.status_code == 409
    assert exc.value.detail["code"] == "AUTOCOUNT_OWNED"
    assert "debtor_name" in exc.value.detail["message"]
    assert "total_amount" in exc.value.detail["message"]
    assert "owned by AutoCount" in exc.value.detail["message"]
    assert "driver_name" not in exc.value.detail["message"]
    db.expire_all()
    order = db.get(Order, order.id)
    assert order.debtor_name == "AutoCount Name" and order.driver_name is None


def test_update_order_tracking_fields_on_autocount_row(db):
    """AC-OG10: Order Tracking fields stay editable on an AutoCount row, and the edit does not
    re-point AutoCount's customer from the debtor text."""
    order = _autocount_order(db)
    before = _snapshot(order, AUTOCOUNT_VALUES)
    updated = OrderService(db).update_order(
        str(order.id),
        OrderUpdate(driver_name="Bala", transporter="ZZ TRANS", remarks_cs="cs",
                    order_type="RMA", actual_delivery_date=datetime(2026, 9, 30)),
        _USER,
    )
    assert updated.driver_name == "Bala" and updated.order_type == "RMA"
    assert updated.transporter_id is not None
    assert updated.customer_id is None
    assert _snapshot(updated, AUTOCOUNT_VALUES) == before


def test_cancel_autocount_row_is_rejected(db):
    """AC-OG09: the narrow cancel path (is_cancelled + remarks) goes through the same guard."""
    order = _autocount_order(db)
    with pytest.raises(HTTPException) as exc:
        OrderService(db).update_order(str(order.id), OrderUpdate(is_cancelled=True), _USER)
    assert exc.value.status_code == 409 and "is_cancelled" in exc.value.detail["message"]


def test_update_order_plain_row_unchanged(db):
    """AC-OG11: a row AutoCount does not own takes every field as before."""
    order = _plain_order(db)
    updated = OrderService(db).update_order(
        str(order.id), OrderUpdate(debtor_name="Edited", agent="X", is_cancelled=True), _USER)
    assert (updated.debtor_name, updated.agent, updated.is_cancelled) == ("Edited", "X", True)


def test_order_response_names_autocount_owned_fields(db):
    """AC-OG12: the single-order read tells the edit form which fields are AutoCount's."""
    ac = OrderService(db).get_order(str(_autocount_order(db).id))
    plain = OrderService(db).get_order(str(_plain_order(db).id))
    ac_body = OrderResponse.model_validate(ac).model_dump()
    plain_body = OrderResponse.model_validate(plain).model_dump()
    assert set(ac_body["autocount_owned_fields"]) == set(AUTOCOUNT_OWNED_ORDER_COLUMNS)
    assert plain_body["autocount_owned_fields"] == []


# ======================================================================= order lines
def test_order_line_crud_rejected_on_autocount_row(db):
    """AC-OG13: lines of an AutoCount row cannot be added, edited or deleted by hand."""
    order = _autocount_order(db)
    product_id, wh_id, line = _line_fixture(db, order)
    svc = OrderService(db)
    calls = [
        lambda: svc.create_order_line(str(order.id), OrderLineCreate(
            product_id=product_id, warehouse_id=wh_id, quantity=Decimal("1"))),
        lambda: svc.update_order_line(str(order.id), str(line.id),
                                      OrderLineUpdate(quantity=Decimal("5"))),
        lambda: svc.delete_order_line(str(order.id), str(line.id)),
        lambda: svc.bulk_delete_order_lines(str(order.id), [str(line.id)]),
    ]
    for call in calls:
        with pytest.raises(HTTPException) as exc:
            call()
        assert exc.value.status_code == 409
        assert exc.value.detail["code"] == "AUTOCOUNT_OWNED"
        assert "owned by AutoCount" in exc.value.detail["message"]
    db.expire_all()
    lines = db.query(OrderLine).filter(OrderLine.order_id == order.id).all()
    assert [(l.quantity, str(l.id)) for l in lines] == [(Decimal("10.0000"), str(line.id))]


def test_order_line_crud_on_plain_row_unchanged(db):
    """AC-OG14: lines of a row AutoCount does not own stay editable."""
    order = _plain_order(db)
    product_id, wh_id, line = _line_fixture(db, order)
    svc = OrderService(db)
    added = svc.create_order_line(str(order.id), OrderLineCreate(
        product_id=product_id, warehouse_id=wh_id, quantity=Decimal("1")))
    assert svc.update_order_line(str(order.id), str(line.id),
                                 OrderLineUpdate(quantity=Decimal("5"))).quantity == Decimal("5")
    svc.delete_order_line(str(order.id), str(added.id))
    assert svc.bulk_delete_order_lines(str(order.id), [str(line.id)])["deleted_count"] == 1
