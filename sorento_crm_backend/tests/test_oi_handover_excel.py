"""OI-RESEND-PURCHASING, owner ruling 16 Oct 00:10: the handover email carries the OI's own
Excel export (the workbook the detail page's Export Excel makes), on the automatic confirm
send AND on "Send to purchasing", through ONE shared build.

TEST-FIRST. Storage is stubbed (`stored`): nothing leaves the sandbox. Dispatch is captured.
"""
from __future__ import annotations

import pytest

from app.models.project_so import OrderInquiry

from .test_order_inquiry_handover_automation import (
    _captured_dispatches,
    _handover_calls,
    _register,
)
from .test_order_inquiry_handshake import _raise_one_row, api, world  # noqa: F401 - fixtures

__all__ = ["api", "world"]


def _drain_on_test_connection(monkeypatch, world) -> None:
    """The post-commit drain opens its own `SessionLocal()`, which in production sees the
    committed write. Here the write lives in the test's never-committed transaction, so the
    drain's session is bound to that same connection (a savepoint) to read it."""
    from sqlalchemy.orm import Session

    connection = world.db.connection()
    monkeypatch.setattr(
        "app.database.SessionLocal",
        lambda: Session(bind=connection, join_transaction_mode="create_savepoint"),
    )


def _enabled_handover(world) -> None:
    """An enabled handover automation: the workbook is stored only when one will send it."""
    import uuid

    from app.models.automation import Automation
    from app.models.email_template import EmailTemplate

    template = EmailTemplate(
        id=str(uuid.uuid4()), code=f"zzt-oix-{uuid.uuid4().hex[:6]}", name="ZZT OIX",
        subject="OI: {{ handover.subject_scope }}", body_html="<p>x</p>", is_active=True,
    )
    world.db.add(template)
    world.db.flush()
    world.db.add(
        Automation(
            id=str(uuid.uuid4()), name="ZZT OIX handover", enabled=True,
            trigger_type="order_inquiry_handover", trigger_config={}, action_type="send_email",
            email_template_id=str(template.id), recipient_config={"include_actor": True},
            group_matches=False, schedule_type="manual", timezone="Asia/Kuala_Lumpur",
        )
    )
    world.db.commit()


def _cells(data: bytes) -> set:
    import io

    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data))
    return {str(v) for ws in workbook for row in ws.iter_rows(values_only=True) for v in row if v}


@pytest.fixture(autouse=True)
def stored(monkeypatch):
    """Every upload the handover makes, kept in memory instead of S3/R2 (autouse: no test
    here may reach a real bucket, whatever `.env` it runs against)."""
    uploads: list[dict] = []

    class _Backend:
        def upload_file(self, *, file_content, file_path, content_type=None, **_kw):
            uploads.append({"bytes": file_content, "key": file_path, "type": content_type})
            return file_path, None

    monkeypatch.setattr("app.services.storage_router.get_backend", lambda provider=None: _Backend())
    return uploads


def test_confirm_handover_attaches_the_oi_excel(api, monkeypatch, stored):
    _client, world = api
    _register(world)
    calls = _captured_dispatches(monkeypatch)
    _drain_on_test_connection(monkeypatch, world)
    _enabled_handover(world)

    fixture = _raise_one_row(api, qty="10")

    inquiry = world.db.get(OrderInquiry, fixture["row"].order_inquiry_id)
    [call] = _handover_calls(calls)
    excel = [
        a for a in call["context"].get("email_attachments", [])
        if a["filename"] == f"{inquiry.inquiry_no}.xlsx"
    ]
    assert len(excel) == 1, call["context"].get("email_attachments")
    assert excel[0]["optional"] is True
    [upload] = [u for u in stored if u["key"] == excel[0]["storage_key"]]
    assert fixture["row"].item_code in _cells(upload["bytes"])  # this OI's own rows (B2)
    import io

    import openpyxl

    from app.services.order_inquiry_worklist_service import EXPORT_TITLE

    # Owner, 16 Oct 2026: the OI Excel is never split by month, the Confirm attachment included.
    assert openpyxl.load_workbook(io.BytesIO(upload["bytes"])).sheetnames == [EXPORT_TITLE]


def test_confirm_handover_still_sends_when_the_excel_cannot_be_stored(api, monkeypatch):
    _client, world = api
    _register(world)
    calls = _captured_dispatches(monkeypatch)

    class _Down:
        def upload_file(self, **_kw):
            raise RuntimeError("storage down")

    monkeypatch.setattr("app.services.storage_router.get_backend", lambda provider=None: _Down())
    _drain_on_test_connection(monkeypatch, world)
    _enabled_handover(world)

    _raise_one_row(api, qty="10")

    [call] = _handover_calls(calls)
    assert not any(
        a["filename"].endswith(".xlsx") for a in call["context"].get("email_attachments", [])
    )


def test_a_database_error_in_the_excel_still_sends_the_email(api, monkeypatch):
    """Review B1: the render runs in a savepoint, so an SQL error inside it leaves the drain
    session usable and the email still dispatches (and can still query)."""
    import sqlalchemy as sa

    _client, world = api
    _register(world)
    _drain_on_test_connection(monkeypatch, world)
    _enabled_handover(world)
    dispatched: list = []

    def _export_fails(self, **_kw):
        self.db.execute(sa.text("SELECT 1/0"))

    def _dispatch_queries(self, trigger_type, *, context, source_kind, source_id):
        self.db.execute(sa.text("SELECT 1"))  # the real dispatch reads the automations
        dispatched.append(context)
        return {"trigger_type": trigger_type, "fired": 0, "results": []}

    monkeypatch.setattr(
        "app.services.order_inquiry_worklist_service.OrderInquiryWorklistService.export_xlsx",
        _export_fails,
    )
    monkeypatch.setattr(
        "app.services.automation_service.AutomationService.dispatch_event", _dispatch_queries
    )

    _raise_one_row(api, qty="10")

    assert len(dispatched) == 1
    assert not any(
        a["filename"].endswith(".xlsx") for a in dispatched[0].get("email_attachments", [])
    )


def test_confirm_with_no_enabled_automation_stores_no_workbook(api, monkeypatch, stored):
    """Review SF1 / S2: nothing will send it, so nothing is rendered or stored."""
    _client, world = api
    _register(world)
    _captured_dispatches(monkeypatch)
    _drain_on_test_connection(monkeypatch, world)

    _raise_one_row(api, qty="10")

    assert stored == []
