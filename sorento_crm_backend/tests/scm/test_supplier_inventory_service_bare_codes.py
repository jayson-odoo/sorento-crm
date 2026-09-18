"""Stock list bare model codes - the SERVICE half (S3, `PLAN-stock-list-bare-model-codes.md`).

TEST-FIRST (Phase 2): `supplier_inventory_service` does not yet build a `WordList` for the
chosen supplier, so AC-S1/AC-S2 are expected to be RED - either a `ModuleNotFoundError`
seeding word rows, or an assertion that a bare code composed and bound where today it is
kept verbatim and unmatched. AC-S3 is a guard-rail rather than a red test: it is expected to
PASS now (the matcher is untouched) and to start failing the moment a slice edits it.

Every product/supplier/word row is seeded by this file under its own marker-prefixed codes,
per `tests/scm/test_supplier_inventory_service.py`'s own rule - the CI database starts empty.
"""
from __future__ import annotations

import subprocess
import uuid
from datetime import date
from io import BytesIO
from pathlib import Path

import pytest

from app.services.scm import supplier_inventory_service as svc
from tests._pg_fixture import pg_session
from tests.scm._outstanding_workbooks import MARKER, require_aliases

HEADER = ["型号", "品名", "商标", "规格", "包装好库存", "空瓷", "体积(cbm)", "备注"]


def workbook(rows, header=None) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(header or HEADER))
    for r in rows:
        ws.append(list(r))
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def merged_workbook(rows, merges: list[str], header=None) -> bytes:
    """Same shape as `workbook`, plus merged ranges named the way the real sheet does
    (`"A2:A5"`) - AC-R7's four `8613` siblings share one merged 型号/品名/商标 family."""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(header or HEADER))
    for r in rows:
        ws.append(list(r))
    for rng in merges:
        ws.merge_cells(rng)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class Codes:
    def __init__(self):
        tag = uuid.uuid4().hex[:8].upper()
        self.known = f"{MARKER}-SIBC-{tag}"
        self.supplier = f"{MARKER}-CRSIBC-{tag}"


def seed(db, codes: Codes, *, product_code: str | None = None) -> str:
    """The one catalogue product and the one supplier this file names. Returns supplier id."""
    from app.models.procurement import Supplier
    from app.models.product import Product, ProductCategory, UnitOfMeasure

    require_aliases(db, "supplier_inventory")

    cat = ProductCategory(
        id=str(uuid.uuid4()),
        category_code=f"{MARKER}-CATBC-{uuid.uuid4().hex[:8]}".upper(),
        category_name=f"{MARKER} bare-code category",
    )
    uom = UnitOfMeasure(
        id=str(uuid.uuid4()),
        uom_code=f"{MARKER}-UBC-{uuid.uuid4().hex[:6]}".upper(),
        uom_name="pcs",
    )
    db.add_all([cat, uom])
    db.flush()
    db.add(
        Product(
            id=str(uuid.uuid4()),
            product_code=product_code or codes.known,
            product_name=product_code or codes.known,
            category_id=cat.id,
            base_uom_id=uom.id,
            list_price=0,
            is_active=True,
            is_discontinued=False,
        )
    )
    supplier = Supplier(
        id=str(uuid.uuid4()),
        supplier_code=codes.supplier,
        supplier_name=f"{MARKER} bare-code supplier",
        is_active=True,
    )
    db.add(supplier)
    db.flush()
    return str(supplier.id)


_D7_PAIRS = (("SRT", "SORENTO"), ("WB", "盆"))


def _seed_shared_words(db, pairs=_D7_PAIRS) -> None:
    """The D7 seed's shape, restated by hand (not read off the migration) - the exact rows
    `SRTWB7055` needs: SORENTO -> SRT, 盆 -> WB. Real D7 spellings on purpose (fix round 1,
    item c): the point of AC-S1/AC-S2 is that the REAL spellings compose, not a marker-
    prefixed stand-in. `pairs` is overridable for a test that needs a different subset of
    the real D7 vocabulary (AC-R7's `连体马桶`/`横排`).

    Idempotent (fix round 1, item c): the migration's own seed already ships these exact
    rows as SHARED (`supplier_id IS NULL`), and the new partial unique index
    (`uq_import_field_alias_shared`) rightly rejects inserting them a second time on a
    migrated database. Checking first keeps this pass on a migrated DB (rows already exist)
    and a bare `create_all` DB (rows do not exist yet) alike.
    """
    from app.models.import_alias import ImportFieldAlias
    from app.services.scm.supplier_code_composer import WORD_DOC_TYPE

    for field, alias in pairs:
        exists = (
            db.query(ImportFieldAlias.id)
            .filter(
                ImportFieldAlias.doc_type == WORD_DOC_TYPE,
                ImportFieldAlias.field == field,
                ImportFieldAlias.alias == alias,
                ImportFieldAlias.supplier_id.is_(None),
            )
            .first()
        )
        if exists is None:
            db.add(
                ImportFieldAlias(
                    id=str(uuid.uuid4()), doc_type=WORD_DOC_TYPE, field=field,
                    alias=alias, supplier_id=None,
                )
            )
    db.flush()


def held(db, supplier_id: str):
    from app.models.scm import SupplierInventory

    return (
        db.query(SupplierInventory)
        .filter(SupplierInventory.supplier_id == supplier_id)
        .order_by(SupplierInventory.item_code)
        .all()
    )


def test_ac_s1_a_bare_code_file_binds_via_the_exact_rung_and_writes_no_alias():
    with pg_session() as db:
        from app.models.scm import SupplierProductCodeAlias

        codes = Codes()
        supplier_id = seed(db, codes, product_code="SRTWB7055")
        _seed_shared_words(db)
        data = workbook([["-7055", "盆", "SORENTO", None, 5, 0, 0.2, ""]])

        out = svc.apply(db, data, supplier_id=supplier_id, as_of=date(2026, 9, 14))

        assert out["rows_written"] == 1
        row = held(db, supplier_id)[0]
        assert row.item_code == "SRTWB7055"
        assert row.product_id is not None

        remembered = (
            db.query(SupplierProductCodeAlias)
            .filter(
                SupplierProductCodeAlias.supplier_id == supplier_id,
                SupplierProductCodeAlias.supplier_code == "SRTWB7055",
            )
            .count()
        )
        assert remembered == 0


def test_ac_r7_four_bare_rows_sharing_a_merged_model_become_four_rows_with_their_own_quantities():
    """The owner's own `8613` example (plan D2): 品名/商标/型号 merged over four rows that
    each state their own 规格 (150/200/250mm, 横排180mm) compose to FOUR different codes,
    so `apply`'s dedup-by-`item_code` never collapses them and each keeps its own quantity -
    unlike the merged-cells suite's quantity fields, which are never merge-fill fields at
    all and always stayed on their own row."""
    with pg_session() as db:
        codes = Codes()
        supplier_id = seed(db, codes)
        _seed_shared_words(db, pairs=(("SRT", "SORENTO"), ("WC", "连体马桶"), ("P", "横排")))
        data = merged_workbook(
            [
                ["8613", "连体马桶", "SORENTO", "150mm", 10, 0, 0.2, ""],
                [None, None, None, "200mm", 20, 0, 0.2, ""],
                [None, None, None, "250mm", 30, 0, 0.2, ""],
                [None, None, None, "横排180mm", 40, 0, 0.2, ""],
            ],
            merges=["A2:A5", "B2:B5", "C2:C5"],
        )

        out = svc.apply(db, data, supplier_id=supplier_id, as_of=date(2026, 9, 14))

        assert out["rows_written"] == 4
        assert out["duplicate_models_merged"] == 0
        rows = held(db, supplier_id)
        expected_codes = {"SRTWC8613-150", "SRTWC8613-200", "SRTWC8613-250", "SRTWC8613-P-180"}
        assert {r.item_code for r in rows} == expected_codes
        by_code = {r.item_code: r for r in rows}
        assert float(by_code["SRTWC8613-150"].qty_packed) == 10
        assert float(by_code["SRTWC8613-200"].qty_packed) == 20
        assert float(by_code["SRTWC8613-250"].qty_packed) == 30
        assert float(by_code["SRTWC8613-P-180"].qty_packed) == 40


def test_ac_s2_a_letter_led_fixture_is_identical_with_or_without_a_word_list():
    codes = Codes()

    def run(seed_words_flag: bool):
        with pg_session() as db:
            supplier_id = seed(db, codes)
            if seed_words_flag:
                _seed_shared_words(db)
            data = workbook([[codes.known, "座厕", "SORENTO", None, 120, 340, 0.21, ""]])

            out = svc.apply(db, data, supplier_id=supplier_id, as_of=date(2026, 7, 31))

            row = held(db, supplier_id)[0]
            # `product_id` is a fresh uuid seeded fresh in EACH `run()` call, so it can
            # never be equal across the two arms for any implementation - the bound
            # PRODUCT's own code is the stable, comparable fact (both arms seed the same
            # `codes.known` as the product's `product_code`).
            from app.models.product import Product

            bound_product_code = (
                db.query(Product.product_code).filter(Product.id == row.product_id).scalar()
                if row.product_id
                else None
            )
            return (
                out["rows_written"],
                out["summary"]["items_matched"],
                row.item_code,
                bound_product_code,
                float(row.qty_packed),
                float(row.qty_unfinished),
            )

    without_words = run(False)
    with_words = run(True)

    assert without_words == with_words


def test_ac_s3_supplier_code_matcher_has_no_diff_in_the_pr():
    repo_root = Path(__file__).resolve().parents[2]
    try:
        subprocess.run(["git", "--version"], capture_output=True, check=True)
    except (OSError, subprocess.CalledProcessError):
        pytest.skip("git is not available in this environment")

    result = subprocess.run(
        ["git", "diff", "origin/main", "--", "app/services/scm/supplier_code_matcher.py"],
        cwd=repo_root,
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        pytest.skip(f"git diff against origin/main failed: {result.stderr.strip()}")

    assert result.stdout.strip() == "", (
        "supplier_code_matcher.py has a diff against origin/main - the plan's out-of-scope "
        "line says this file is never touched:\n" + result.stdout
    )
