"""GET /api/v1/inventory/stock-position - the supplier stock sheet, every supplier, one row per item.

Plan: documentation/plans/inventory/PLAN-stock-position-report-27oct.md (column map, Q1-Q5).
Every figure is the Stock Status panel's own (`stock_status_service`), batched: the seeds reuse
that suite's helpers so a document the panel counts is one this report counts.
Each test seeds its own chain under a fresh company (CI's database is empty).
"""
from __future__ import annotations

import io
import uuid
from datetime import date

from fastapi.testclient import TestClient

from app.models.finance import BillingDocument, BillingDocumentLine
from app.models.inventory import Stock, Warehouse
from app.models.procurement import ProductSupplier, PurchaseOrder, PurchaseOrderLine, Supplier
from app.models.product import Brand
from app.services.scm.trajectory import month_shift
from tests.scm.conftest import as_user, grant_permission, requires_pg, seed_user
from tests.scm.test_loading_plan import World
from tests.scm.test_outstanding_import_routes import as_company_user
from tests.scm.test_stock_status import (
    MARKER,
    _on_hand,
    _pl_line,
    _po,
    _shipment,
    _so,
    _spo,
)

pytestmark = requires_pg

URL = "/api/v1/inventory/stock-position"


def _wh(db, code: str, *, segment=None, active: bool = True) -> Warehouse:
    wh = Warehouse(
        id=str(uuid.uuid4()),
        warehouse_code=code,
        warehouse_name=f"{MARKER} {code}",
        is_active=active,
        segment=segment,
    )
    db.add(wh)
    db.flush()
    return wh


def _setup(scm_app):
    app, db, gcu, gcuk = scm_app
    as_company_user(app, db, gcu, gcuk)  # purchasing: holds scm.reorder.run
    return app, db, World(db)


def _get(app, **params):
    r = TestClient(app).get(URL, params={"limit": 100, **params})
    assert r.status_code == 200, r.text
    return r.json()


def _row(body, product):
    return next((r for r in body["data"] if r["product_id"] == str(product.id)), None)


def _months_ago(n: int) -> date:
    """A day inside the full month `n` months before this one (n=1 is last month)."""
    return month_shift(month_shift(date.today(), 0), -n).replace(day=15)


def _billing(db, product, qty, doc_date, *, doc_type="invoice", status="posted") -> None:
    doc = BillingDocument(
        id=str(uuid.uuid4()),
        document_type=doc_type,
        doc_no=f"{MARKER}-IV-{uuid.uuid4().hex[:8]}",
        doc_date=doc_date,
        status=status,
        local_net_total=0,
        source_ref=uuid.uuid4().hex,
    )
    db.add(doc)
    db.flush()
    db.add(BillingDocumentLine(
        id=str(uuid.uuid4()), document_id=doc.id, product_id=product.id, quantity=qty,
        source_ref=uuid.uuid4().hex,
    ))
    db.flush()


def _sheet_world(db, w: World):
    """Warehouses BRW, MWH, WH3, DC1, RSW (site pool) + one project bin, and three items shaped
    like the owner's sheet rows (A: CWCX605-RL, B: SRTWC286-SH-NEW, C: CWC7606-SH)."""
    whs = {code: _wh(db, code) for code in ("BRW", "MWH", "WH3", "DC1", "RSW")}
    bin_ = _wh(db, "BRW-BB", segment="project")
    a, b, c = w.product("A"), w.product("B"), w.product("C")

    # A: BRW 40, DC1 244; SPO 352 to BRW (60 of it on a packing list of 100 still to land, so the
    # packing list counts 40), 880 to the project bin; dealer SO owing 70; PO 2,000 with 180 on SPO.
    _on_hand(db, a, whs["BRW"], 40)
    _on_hand(db, a, whs["DC1"], 244)
    ship = _shipment(db, w)
    _pl_line(db, ship, a, 100)
    _spo(db, w, a, whs["BRW"], 292)
    _spo(db, w, a, whs["BRW"], 60, shipment=ship)
    _spo(db, w, a, bin_, 880)
    _so(db, a, whs["BRW"], 100, delivered=30)
    _po(db, w, a, whs["BRW"], 2000, received=180)

    # B: BRW 42, MWH 846, WH3 81; SPO 84 to BRW, 51 to the bin; dealer SO owing 59; PO 75.
    _on_hand(db, b, whs["BRW"], 42)
    _on_hand(db, b, whs["MWH"], 846)
    _on_hand(db, b, whs["WH3"], 81)
    _spo(db, w, b, whs["BRW"], 84)
    _spo(db, w, b, bin_, 51)
    _so(db, b, whs["MWH"], 59)
    _po(db, w, b, whs["BRW"], 75)

    # C: BRW 98, MWH 85, DC1 594; nothing incoming, nothing owed.
    _on_hand(db, c, whs["BRW"], 98)
    _on_hand(db, c, whs["MWH"], 85)
    _on_hand(db, c, whs["DC1"], 594)
    return whs, bin_, (a, b, c)


# ------------------------------------------------------------------ the sheet's own rows
def test_the_sheet_rows_foot_column_by_column(scm_app):
    app, db, w = _setup(scm_app)
    whs, bin_, (a, b, c) = _sheet_world(db, w)

    body = _get(app)
    assert [x["code"] for x in body["warehouses"]] == ["BRW", "DC1", "MWH", "WH3"]

    ra = _row(body, a)
    assert ra["stock"] == {"BRW": 40, "DC1": 244}
    assert ra["incoming_pool"] == 352
    assert ra["total_stock"] == 636
    assert ra["incoming_project"] == 880
    assert ra["total_incoming"] == 1232
    assert ra["packing_list_open"] == 40
    assert ra["outstanding_dealer"] == -70
    assert ra["outstanding_po"] == 1820

    rb = _row(body, b)
    assert rb["stock"] == {"BRW": 42, "MWH": 846, "WH3": 81}
    assert (rb["incoming_pool"], rb["total_stock"], rb["incoming_project"], rb["total_incoming"]) == (
        84, 1053, 51, 135)
    assert (rb["outstanding_dealer"], rb["outstanding_po"], rb["packing_list_open"]) == (-59, 75, 0)

    rc = _row(body, c)
    assert rc["stock"] == {"BRW": 98, "MWH": 85, "DC1": 594}
    assert (rc["total_stock"], rc["total_incoming"], rc["outstanding_dealer"], rc["outstanding_po"]) == (
        777, 0, 0, 0)
    assert rb["item_code"] == b.product_code


# ------------------------------------------------------------------ warehouse columns (Q5)
def test_columns_are_active_site_pool_warehouses_holding_stock_project_bins_and_inactive_are_not(scm_app):
    app, db, w = _setup(scm_app)
    p = w.product("A")
    brw = _wh(db, "BRW")
    _wh(db, "EMPTY")  # holds nothing: no column
    old = _wh(db, "OLD", active=False)
    bin_ = _wh(db, "BRW-BB", segment="project")
    _on_hand(db, p, brw, 5)
    _on_hand(db, p, old, 7)
    _on_hand(db, p, bin_, 9)

    body = _get(app)
    assert [x["code"] for x in body["warehouses"]] == ["BRW"]
    r = _row(body, p)
    assert r["stock"] == {"BRW": 5}
    assert r["total_stock"] == 5  # neither the inactive location nor the project bin


def test_another_companys_warehouse_is_no_column_and_its_stock_no_figure(scm_app):
    app, db, w = _setup(scm_app)
    p = w.product("A")
    _on_hand(db, p, _wh(db, "MINE"), 3)
    theirs = str(uuid.uuid4())
    from app.models.company import Company

    db.add(Company(id=theirs, name=f"{MARKER} other {theirs[:6]}",
                   code=f"{MARKER}-{uuid.uuid4().hex[:6]}".upper(), is_active=True))
    db.flush()
    other = Warehouse(id=str(uuid.uuid4()), warehouse_code="THEIRS", warehouse_name="other",
                      is_active=True, company_id=theirs)
    db.add(other)
    db.flush()
    db.add(Stock(id=str(uuid.uuid4()), product_id=p.id, warehouse_id=other.id, quantity_on_hand=50,
                 company_id=theirs))
    db.flush()

    # Documents of OUR company bound to THEIR warehouse: the warehouse check is the only guard,
    # since the company scope lets these rows through.
    _so(db, p, other, 40)
    _po(db, w, p, other, 30)
    _spo(db, w, p, other, 20)

    body = _get(app)
    assert [x["code"] for x in body["warehouses"]] == ["MINE"]
    r = _row(body, p)
    assert r["stock"] == {"MINE": 3}
    assert (r["outstanding_dealer"], r["outstanding_po"], r["incoming_pool"], r["incoming_project"]) == (0, 0, 0, 0)


# ------------------------------------------------------------------ incoming
def test_incoming_counts_only_open_spo_still_to_come_split_pool_and_project(scm_app):
    app, db, w = _setup(scm_app)
    p = w.product("A")
    brw = _wh(db, "BRW")
    bin_ = _wh(db, "PRJ", segment="project")
    _spo(db, w, p, brw, 10, received=4)          # 6 still to come
    _spo(db, w, p, brw, 5, line_status="closed")  # left the book
    _spo(db, w, p, bin_, 8)

    r = _row(_get(app), p)
    assert (r["incoming_pool"], r["incoming_project"], r["total_incoming"]) == (6, 8, 14)
    assert r["total_stock"] == 6  # no warehouse stock, the pool's incoming only


def test_packing_list_open_is_what_no_spo_carries_and_stays_out_of_total_incoming(scm_app):
    app, db, w = _setup(scm_app)
    p = w.product("A")
    brw = _wh(db, "BRW")
    linked = _shipment(db, w)
    _pl_line(db, linked, p, 30)
    _spo(db, w, p, brw, 12, shipment=linked)
    by_container = _shipment(db, w, container="TGHU1234567 (PORT)")
    _pl_line(db, by_container, p, 20)
    _spo(db, w, p, brw, 5, container="TGHU1234567")

    r = _row(_get(app), p)
    assert r["packing_list_open"] == (30 - 12) + (20 - 5)
    assert r["total_incoming"] == 17  # the two SPO rows; the packing lists are not added


def test_packing_list_netting_is_per_item_on_a_shipment_carrying_two(scm_app):
    app, db, w = _setup(scm_app)
    a, b = w.product("A"), w.product("B")
    brw = _wh(db, "BRW")
    ship = _shipment(db, w)
    _pl_line(db, ship, a, 30)
    _pl_line(db, ship, b, 20)
    _spo(db, w, a, brw, 25, shipment=ship)  # on A only: B's 20 is still on no SPO

    body = _get(app)
    assert _row(body, a)["packing_list_open"] == 5
    assert _row(body, b)["packing_list_open"] == 20


# ------------------------------------------------------------------ outstanding
def test_outstanding_dealer_is_dealer_so_still_owed_printed_negative(scm_app):
    app, db, w = _setup(scm_app)
    p = w.product("A")
    brw = _wh(db, "BRW")
    _so(db, p, brw, 10, delivered=3)          # owes 7
    _so(db, p, brw, 4, transferable=False)    # not approved for the queue: out
    _so(db, p, brw, 6, line_status="closed")  # closed line: out
    project = _so(db, p, brw, 50)             # project demand is not dealer demand
    project.demand_class = "project"
    db.flush()

    assert _row(_get(app), p)["outstanding_dealer"] == -7


def test_outstanding_po_is_ordered_less_transferred_on_open_lines(scm_app):
    app, db, w = _setup(scm_app)
    p = w.product("A")
    brw = _wh(db, "BRW")
    _po(db, w, p, brw, 100, received=40)
    _po(db, w, p, brw, 9, line_status="closed")
    _po(db, w, p, brw, 9, status="draft")

    assert _row(_get(app), p)["outstanding_po"] == 60


# ------------------------------------------------------------------ monthly sales (Q3)
# Owner ruling 27 Oct: one sales calculation everywhere. The two columns come from SALES-UNIFIED's
# shared function (`invoiced_sales.quantity_by_product_month`, #1916): IV + CS + DN - CN line
# quantity by document date, cancelled excluded. This report adds only the 12-month window,
# the mean over all 12 months and the peak.
def test_monthly_sales_is_the_shared_invoiced_quantity_12_full_month_mean_and_peak(scm_app):
    app, db, w = _setup(scm_app)
    p = w.product("A")
    _wh(db, "BRW")
    _billing(db, p, 120, _months_ago(2))
    _billing(db, p, 30, _months_ago(2), doc_type="cash_sale")
    _billing(db, p, 6, _months_ago(3), doc_type="debit_note")
    _billing(db, p, 10, _months_ago(2), doc_type="credit_note")
    _billing(db, p, 500, _months_ago(4), status="cancelled")
    _billing(db, p, 40, _months_ago(12))
    _billing(db, p, 999, _months_ago(13))      # older than the window
    _billing(db, p, 777, date.today())         # this month is not a full month

    r = _row(_get(app), p)
    assert r["monthly_sales"] == (120 + 30 - 10 + 6 + 40) / 12
    assert r["highest_month"] == 140


def test_monthly_sales_reads_the_one_shared_sales_function(scm_app, monkeypatch):
    """No second calculation: the report asks `invoiced_sales` and sums nothing itself."""
    from app.services.sales import invoiced_sales

    app, db, w = _setup(scm_app)
    p = w.product("A")
    _wh(db, "BRW")
    asked = []

    def shared(session, since, until, product_ids=None):
        asked.append((since, until, product_ids))
        return [(str(p.id), _months_ago(1).strftime("%Y-%m"), 24)]

    monkeypatch.setattr(invoiced_sales, "quantity_by_product_month", shared)
    r = _row(_get(app), p)
    assert (r["monthly_sales"], r["highest_month"]) == (2, 24)
    (since, until, product_ids), = asked
    assert until == date.today().replace(day=1)
    assert (until.year - since.year) * 12 + until.month - since.month == 12
    assert product_ids is None


def test_sales_basis_is_no_longer_a_choice(scm_app):
    app, db, w = _setup(scm_app)
    body = _get(app)
    assert "sales_basis" not in body


# ------------------------------------------------------------------ supplier (Q4), filters, sort
def _supplier(db, code) -> Supplier:
    s = Supplier(id=str(uuid.uuid4()), supplier_code=f"{MARKER}-{code}-{uuid.uuid4().hex[:4]}",
                 supplier_name=f"{MARKER} {code}", is_active=True)
    db.add(s)
    db.flush()
    return s


def _po_from(db, product, supplier, wh, issued, *, status="active") -> None:
    po = PurchaseOrder(id=str(uuid.uuid4()), po_number=f"{MARKER}-PO-{uuid.uuid4().hex[:8]}",
                       supplier_id=supplier.id, issue_date=issued, status=status)
    db.add(po)
    db.flush()
    db.add(PurchaseOrderLine(id=str(uuid.uuid4()), purchase_order_id=po.id, product_id=product.id,
                             warehouse_id=wh.id, qty_ordered=5, qty_received=5, line_status="closed"))
    db.flush()


def test_supplier_is_the_last_po_supplier_and_filters_on_it(scm_app):
    """Owner ruling 27 Oct (Q4): the LAST PO's supplier, the shared rule Stock Debt and the low
    stock report read (`scm/last_po_supplier.py`). The product links play no part."""
    app, db, w = _setup(scm_app)
    brw = _wh(db, "BRW")
    s1, s2, s3 = _supplier(db, "S1"), _supplier(db, "S2"), _supplier(db, "S3")
    a, b, c = (w.product(k) for k in "ABC")
    for p in (a, b, c):
        _on_hand(db, p, brw, 1)
    _po_from(db, a, s1, brw, date(2025, 1, 1))
    _po_from(db, a, s2, brw, date(2026, 3, 1))                       # newest: S2
    _po_from(db, a, s3, brw, date(2026, 6, 1), status="cancelled")   # cancelled: skipped
    _po_from(db, b, s1, brw, date(2026, 2, 1))
    db.add(ProductSupplier(id=str(uuid.uuid4()), product_id=c.id, supplier_id=s3.id, is_primary_supplier=True))
    db.flush()                                                        # c: a link, no PO: none

    body = _get(app)
    assert _row(body, a)["supplier_name"] == s2.supplier_name
    assert _row(body, b)["supplier_name"] == s1.supplier_name
    assert _row(body, c)["supplier_name"] is None

    only = _get(app, supplier_id=s1.id)
    assert {r["product_id"] for r in only["data"]} == {str(b.id)}


def test_supplier_reads_the_shared_last_po_rule(scm_app, monkeypatch):
    from app.services.scm import last_po_supplier

    app, db, w = _setup(scm_app)
    p = w.product("A")
    s = _supplier(db, "spy")
    _on_hand(db, p, _wh(db, "BRW"), 1)
    monkeypatch.setattr(last_po_supplier, "last_po_supplier_ids", lambda session, ids: {str(p.id): str(s.id)})
    assert _row(_get(app), p)["supplier_name"] == s.supplier_name


def test_brand_item_code_and_warehouse_filters(scm_app):
    app, db, w = _setup(scm_app)
    brw, mwh = _wh(db, "BRW"), _wh(db, "MWH")
    brand = Brand(id=str(uuid.uuid4()), brand_code=f"{MARKER}-{uuid.uuid4().hex[:6]}", brand_name="ZZ brand")
    db.add(brand)
    db.flush()
    a, b = w.product("A"), w.product("B")
    a.brand_id = brand.id
    db.flush()
    _on_hand(db, a, brw, 1)
    _on_hand(db, b, mwh, 1)

    assert [r["product_id"] for r in _get(app, brand_id=brand.id)["data"]] == [str(a.id)]
    assert [r["product_id"] for r in _get(app, query=b.product_code.lower())["data"]] == [str(b.id)]
    # A warehouse filter keeps the items holding stock there; the columns stay whole.
    filtered = _get(app, warehouse_id=mwh.id)
    assert [r["product_id"] for r in filtered["data"]] == [str(b.id)]
    assert [x["code"] for x in filtered["warehouses"]] == ["BRW", "MWH"]
    assert [x["id"] for x in filtered["warehouses"]] == [str(brw.id), str(mwh.id)]
    # A value that is no id matches nothing rather than failing.
    assert _get(app, warehouse_id="nope")["data"] == []


def test_every_column_sorts_on_the_server_and_pages(scm_app):
    app, db, w = _setup(scm_app)
    brw, mwh = _wh(db, "BRW"), _wh(db, "MWH")
    a, b, c = w.product("A"), w.product("B"), w.product("C")
    _on_hand(db, a, brw, 5)
    _on_hand(db, b, brw, 50)
    _on_hand(db, c, mwh, 20)

    def order(**params):
        return [r["product_id"] for r in _get(app, **params)["data"]]

    assert order(sort="total_stock", dir="desc") == [str(b.id), str(c.id), str(a.id)]
    # A warehouse column sorts by its code; an item with nothing there reads 0.
    assert order(sort="wh:BRW", dir="asc") == [str(c.id), str(a.id), str(b.id)]
    assert order(sort="item_code", dir="asc") == [str(a.id), str(b.id), str(c.id)]

    page = _get(app, sort="total_stock", dir="desc", page=2, limit=2)
    assert [r["product_id"] for r in page["data"]] == [str(a.id)]
    assert page["pagination"] == {"page": 2, "limit": 2, "total": 3, "total_pages": 2}
    assert TestClient(app).get(URL, params={"sort": "nope"}).status_code == 422
    assert TestClient(app).get(URL, params={"sort": "wh:NOPE"}).status_code == 422


def test_items_with_nothing_to_show_are_not_rows(scm_app):
    app, db, w = _setup(scm_app)
    _wh(db, "BRW")
    idle = w.product("IDLE")
    assert _row(_get(app), idle) is None


# ------------------------------------------------------------------ access, export
def test_read_needs_one_of_the_stock_status_permissions(scm_app):
    app, db, gcu, gcuk = scm_app
    as_user(app, gcu, gcuk, seed_user(db, "guest"))
    assert TestClient(app).get(URL).status_code == 403
    grant_permission(db, "salesperson", "projects.projects.view")
    as_user(app, gcu, gcuk, seed_user(db, "salesperson"))
    assert TestClient(app).get(URL).status_code == 200


def test_export_writes_every_row_with_the_screen_headers(scm_app):
    from openpyxl import load_workbook

    app, db, w = _setup(scm_app)
    _sheet_world(db, w)
    a, b, c = (w.product(k) for k in "ABC")
    r = TestClient(app).get(f"{URL}/export", params={"sort": "item_code", "dir": "asc"})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    ws = load_workbook(io.BytesIO(r.content)).active
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    assert header[:3] == ("Item code", "Supplier", "Brand")
    assert header[3:7] == ("BRW", "DC1", "MWH", "WH3")
    assert "Incoming BRW" in header and "Outstanding dealer" in header
    codes = [row[0] for row in rows[1:]]
    assert codes == sorted([a.product_code, b.product_code, c.product_code])
    c_row = next(row for row in rows[1:] if row[0] == c.product_code)
    assert c_row[header.index("Total stock")] == 777


def test_export_writes_only_the_warehouse_columns_the_user_shows_totals_unchanged(scm_app):
    """Owner ruling 27 Oct (Q5): the user picks which warehouse columns show (the grid's Columns
    chooser); the export writes those, and Total stock still counts every warehouse."""
    from openpyxl import load_workbook

    app, db, w = _setup(scm_app)
    _sheet_world(db, w)
    c = w.product("C")
    r = TestClient(app).get(f"{URL}/export", params={"sort": "item_code", "wh": ["DC1", "BRW"]})
    assert r.status_code == 200, r.text
    rows = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    header = rows[0]
    assert header[3:5] == ("BRW", "DC1")  # in the report's own order, not the request's
    assert "MWH" not in header and "WH3" not in header
    c_row = next(row for row in rows[1:] if row[0] == c.product_code)
    assert c_row[header.index("Total stock")] == 777  # MWH's 85 still counts
    # A code the report does not have is ignored rather than refused.
    assert TestClient(app).get(f"{URL}/export", params={"wh": ["NOPE"]}).status_code == 200


def test_export_with_every_warehouse_hidden_writes_no_warehouse_column(scm_app):
    from openpyxl import load_workbook

    app, db, w = _setup(scm_app)
    _sheet_world(db, w)
    r = TestClient(app).get(f"{URL}/export", params={"wh": [""]})
    assert r.status_code == 200, r.text
    header = next(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert header[:4] == ("Item code", "Supplier", "Brand", "Incoming BRW")
