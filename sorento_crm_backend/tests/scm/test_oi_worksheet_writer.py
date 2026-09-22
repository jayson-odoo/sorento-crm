"""Lane C, PLAN-order-sheet-oi-reports-22sep.md (AC-C3/AC-C5): the OI worksheet is printed
through the SAME writer the OI worklist list page already uses -
`OrderInquiryWorklistService.export_xlsx` - over a different row set, not a new writer
(the plan's own "Rejected" section). Two new keyword arguments make that possible:

* `row_ids` - restricts `_all_rows` to exactly the named `OrderInquiryRow` ids (fed by
  `run_scope_oi_rows`, tested separately in `test_oi_worksheet_row_scope.py`). An empty
  list produces the headings-only one-sheet workbook (AC-C5).
* `columns` - the heading tuple to print, defaulting to today's full `EXPORT_HEADINGS` so
  the list-page export is unchanged; the worksheet passes a 10-column slice that stops at
  LOCATION (AC-C3: "no ACKNOWLEDGED column").

Neither kwarg exists on `export_xlsx`/`_base` today - `export_xlsx(self, **filters)`
forwards everything straight to `_base(**filters)`, and `_base`'s own signature (read
above) has no `row_ids` parameter at all - so every case below is red for a TypeError
("unexpected keyword argument"), not a fixture bug, until C3 lands.

Seeding reused from Lane A's own new helper (`tests.scm.test_order_summary_sheet._scope_row`
+ its `chain`/`db` fixtures) - the exact OI-row shape this writer already serializes
(a project SO line + a Buy-verb `OrderInquiryRow`, no supply decision).
"""
from __future__ import annotations

import io
from datetime import date

import openpyxl

from app.services.order_inquiry_worklist_service import (
    EXPORT_HEADINGS,
    EXPORT_TITLE,
    EXPORT_UNDATED_SHEET,
    OrderInquiryWorklistService,
)
from tests.scm.conftest import requires_pg
from tests.scm.test_order_summary_sheet import MARKER, _scope_row
from tests.scm.test_summary_order_service import chain, db  # noqa: F401

pytestmark = requires_pg

#: AC-C3's own 10-column list ("SO DATE ... LOCATION", no ACKNOWLEDGED / TAKEN /
#: REMAINING). Measured against the file, not assumed: `EXPORT_HEADINGS` carries 13
#: headings today (ACKNOWLEDGED was appended, then TAKEN/REMAINING in a later fix round),
#: not the 11 the plan text alone would suggest - a slice of the real constant is what
#: cannot drift out of sync with whichever of those additions lands first.
WORKSHEET_HEADINGS = EXPORT_HEADINGS[:10]


def _data_rows(workbook):
    """Every non-header cell-row across every sheet, as a flat list of row tuples."""
    out = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=3):
            values = [c.value for c in row]
            if any(v is not None for v in values):
                out.append(values)
    return out


def test_c3_row_ids_filter_restricts_all_rows_to_the_named_ones(db, chain):
    f = chain
    kept = _scope_row(db, product=f["product"], qty=5, delivery=date(2026, 10, 1))
    _scope_row(db, product=f["product"], qty=9, delivery=date(2026, 10, 1))
    db.flush()

    svc = OrderInquiryWorklistService(db)
    _filename, content = svc.export_xlsx(row_ids=[kept["row"].id])

    workbook = openpyxl.load_workbook(io.BytesIO(content))
    rows = _data_rows(workbook)
    assert len(rows) == 1, rows
    assert rows[0][3] == 5.0, rows  # QTY column


def test_c5_empty_row_ids_produces_a_headings_only_one_sheet_workbook():
    from tests._pg_fixture import pg_session

    with pg_session() as db:
        svc = OrderInquiryWorklistService(db)
        _filename, content = svc.export_xlsx(row_ids=[])

        workbook = openpyxl.load_workbook(io.BytesIO(content))
        assert len(workbook.worksheets) == 1, workbook.sheetnames
        sheet = workbook.worksheets[0]
        assert sheet.cell(row=1, column=1).value == EXPORT_TITLE
        headings = [c.value for c in sheet[2]]
        assert headings == list(EXPORT_HEADINGS), headings
        assert _data_rows(workbook) == []


def test_c3_columns_kwarg_prints_the_worksheets_own_headings_no_acknowledged(db, chain):
    f = chain
    leg = _scope_row(db, product=f["product"], qty=3, delivery=date(2026, 10, 1))
    db.flush()

    svc = OrderInquiryWorklistService(db)
    _filename, content = svc.export_xlsx(row_ids=[leg["row"].id], columns=WORKSHEET_HEADINGS)

    workbook = openpyxl.load_workbook(io.BytesIO(content))
    sheet = workbook.worksheets[0]
    assert sheet.cell(row=1, column=1).value == EXPORT_TITLE
    headings = [c.value for c in sheet[2]]
    assert headings == list(WORKSHEET_HEADINGS), headings
    assert "ACKNOWLEDGED" not in headings, headings
    assert len(headings) == 10, headings


def test_c3_default_columns_still_carries_every_heading_for_the_list_page_export(db, chain):
    """The list-page export (no `columns` kwarg) is UNCHANGED - it still prints every
    heading `EXPORT_HEADINGS` carries today (13: the original ten plus ACKNOWLEDGED,
    TAKEN, REMAINING), so a caller that never passes `columns` sees no regression."""
    f = chain
    leg = _scope_row(db, product=f["product"], qty=3, delivery=date(2026, 10, 1))
    db.flush()

    svc = OrderInquiryWorklistService(db)
    _filename, content = svc.export_xlsx(row_ids=[leg["row"].id])

    workbook = openpyxl.load_workbook(io.BytesIO(content))
    sheet = workbook.worksheets[0]
    headings = [c.value for c in sheet[2]]
    assert headings == list(EXPORT_HEADINGS), headings
    assert len(headings) == len(EXPORT_HEADINGS), headings


def test_c3_row_ids_scoped_undated_row_lands_on_the_no_date_sheet(db, chain):
    f = chain
    leg = _scope_row(db, product=f["product"], qty=4, delivery=None)
    db.flush()

    svc = OrderInquiryWorklistService(db)
    _filename, content = svc.export_xlsx(row_ids=[leg["row"].id])

    workbook = openpyxl.load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == [EXPORT_UNDATED_SHEET], workbook.sheetnames
