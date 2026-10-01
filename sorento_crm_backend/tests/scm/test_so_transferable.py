"""SO-TRANSFERABLE: an AutoCount sales order marked Transferable = F is not Stock Debt demand.

Owner, 1 Oct 2026: F means "not confirmed yet for the queue", so Stock Debt must not count it.
Scope ruling (b), same day: F leaves the SHARED assignment (`StockDebtService._demand`, R21's
one reader), so the Stock Debt page and the fulfilment board's ladder both skip it - no stock
is reserved for an F order anywhere until AutoCount flips it to T. NULL (AutoCount never
stated it) counts, exactly like T.

  AC-TR-6  list: an F order's line is not demand; T and unknown are
  AC-TR-7  cell drill: F lines are not listed and not in `demand_total_qty`
  AC-TR-8  the board's shared assignment (the ladder) carries no F line either
  AC-TR-9  GET /scm/sales-orders/{id} answers `is_transferable`, the list filters on it, and
           PUT cannot change it (AutoCount-owned)
"""
from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient

from app.services.scm.supply_assignment import month_key
from tests.scm.conftest import as_user, requires_pg, seed_user
from tests.scm.test_stock_debt_routes import (
    BASE,
    _client,
    _demand,
    _months_ahead,
    _product,
    _row_of,
    _u,
    _warehouse,
)

pytestmark = requires_pg

SO_BASE = "/api/v1/scm/sales-orders"


def _book(db):
    """One product at one BB bin, nothing on hand, three orders due next month:
    F x 50, T x 30, unknown x 20. Only 30 + 20 are owed on the Stock Debt page."""
    marker = f"ZZTTR{_u()[:6]}".upper()
    warehouse = _warehouse(db, f"ZZTTRW{_u()[:4]}-BB")
    product = _product(db, f"{marker}-A")
    due = _months_ahead(1)
    orders = {}
    for key, qty, flag in (("F", 50, False), ("T", 30, True), ("U", 20, None)):
        order, line = _demand(
            db, product, warehouse, qty=qty, required_date=due,
            so_number=f"{marker}-SO{key}",
        )
        order.is_transferable = flag
        orders[key] = (order, line)
    db.flush()
    return {
        "marker": marker, "warehouse": warehouse, "product": product,
        "month": month_key(due), "orders": orders,
    }


def test_the_list_does_not_count_a_non_transferable_order(scm_app):
    """AC-TR-6. -50 next month: T's 30 and the unknown 20; F's 50 is not demand."""
    app, db = _client(scm_app)
    book = _book(db)

    with TestClient(app) as c:
        got = c.get(BASE, params={"query": book["marker"], "only_debt": False, "limit": 25})

    assert got.status_code == 200, got.text
    row = _row_of(got.json(), book["product"].product_code)
    balances = {m["key"]: m["balance"] for m in row["months"]}
    assert balances[book["month"]] == -50, balances


def test_a_product_whose_only_demand_is_non_transferable_gets_no_row(scm_app):
    """AC-TR-6. With no stock and no supply, an F-only product owes nothing, so the candidate
    read must not list it either (the two reads share one rule)."""
    app, db = _client(scm_app)
    marker = f"ZZTTR{_u()[:6]}".upper()
    warehouse = _warehouse(db, f"ZZTTRW{_u()[:4]}-BB")
    product = _product(db, f"{marker}-ONLYF")
    order, _line = _demand(
        db, product, warehouse, qty=40, required_date=_months_ahead(1),
        so_number=f"{marker}-SOF",
    )
    order.is_transferable = False
    db.flush()

    with TestClient(app) as c:
        body = c.get(
            BASE, params={"query": marker, "only_debt": False, "limit": 25}
        ).json()

    assert body["pagination"]["total"] == 0, body["data"]


def test_the_cell_drill_does_not_list_a_non_transferable_line(scm_app):
    """AC-TR-7. The drill foots with the cell, so it lists exactly what the cell counted."""
    app, db = _client(scm_app)
    book = _book(db)

    with TestClient(app) as c:
        got = c.get(
            f"{BASE}/{book['product'].id}/cell", params={"month": book["month"]}
        )

    assert got.status_code == 200, got.text
    cell = got.json()
    listed = sorted(row["so_number"] for row in cell["demand"])
    assert listed == sorted(
        [f"{book['marker']}-SOT", f"{book['marker']}-SOU"]
    ), listed
    assert cell["demand_total_qty"] == 50


def test_the_export_carries_the_same_figure(scm_app):
    """AC-TR-6. The workbook is built off `list()`."""
    from openpyxl import load_workbook

    from app.services.scm.stock_debt_service import StockDebtService

    _app, db = _client(scm_app)
    book = _book(db)
    blob, _type, _name, counts = StockDebtService(db).export(
        query=book["marker"], only_debt=False, split="none",
    )
    assert counts["rows"] == 1
    ws = load_workbook(BytesIO(blob))["Stock debt"]
    header = [cell.value for cell in ws[1]]
    data = [cell.value for cell in ws[2]]
    assert data[header.index("Total")] == -50


def test_the_shared_assignment_drops_the_non_transferable_line(scm_app):
    """AC-TR-8. Owner ruling 1 Oct 2026 (b): F is excluded from the SHARED assignment
    (`assignments_for`, R21's one reader), so the board ladder agrees with the page."""
    from app.services.scm.stock_debt_service import StockDebtService

    _app, db = _client(scm_app)
    book = _book(db)
    warehouse = book["warehouse"]

    result = StockDebtService(db).assignments_for(
        [str(book["product"].id)], {str(warehouse.id): warehouse}
    )[str(book["product"].id)]

    keys = {line.line.key for line in result.lines}
    assert str(book["orders"]["F"][1].id) not in keys
    assert {str(book["orders"][k][1].id) for k in ("T", "U")} <= keys


def test_the_board_ladder_reserves_nothing_for_a_non_transferable_line(scm_app):
    """AC-TR-8, read through the board's own entry point (`planning_assignments`, the
    ladder). On hand 60 at the bin: with F counted the F line (50, the oldest SO number)
    would draw from it; with F out the T and unknown lines (30 + 20) take 50 and 10 stays
    free. The F line has no assignment at all, so no stock is reserved for it anywhere
    until AutoCount flips it to T."""
    from app.services.project_supply_service import ProjectSupplyService
    from tests.scm.test_stock_debt_routes import _stock

    _app, db = _client(scm_app)
    book = _book(db)
    _stock(db, book["product"], book["warehouse"], 60)

    result = ProjectSupplyService(db).planning_assignments([str(book["product"].id)])[
        str(book["product"].id)
    ]

    by_key = {line.line.key: line for line in result.lines}
    assert str(book["orders"]["F"][1].id) not in by_key
    for key in ("T", "U"):
        line = by_key[str(book["orders"][key][1].id)]
        assert round(sum(item.qty for item in line.assigned), 4) == line.line.open_qty


def test_the_board_pile_queue_does_not_rank_a_non_transferable_line(scm_app):
    """AC-TR-8, the board's own pile reads (`_pile_book`, `_pile_read`, `_group_pile_members`,
    `_check_group_borrow`). They rank the lines competing for one pile; an F line competing
    there would claim stock the shared assignment has already said it gets none of, which is
    the R21 disagreement ruling (b) exists to prevent."""
    from app.services.project_supply_service import ProjectSupplyService

    _app, db = _client(scm_app)
    book = _book(db)

    queue = ProjectSupplyService(db).pile_book(
        str(book["product"].id), str(book["warehouse"].id)
    )

    ids = {row["line_id"] for row in queue}
    assert str(book["orders"]["F"][1].id) not in ids
    assert {str(book["orders"][k][1].id) for k in ("T", "U")} <= ids


# ------------------------------------------------------------------ the SO screens (AC-TR-9)


def _so_client(scm_app):
    app, db, gcu, gcuak = scm_app
    from tests.scm.test_stock_debt_routes import ensure_reference_data

    ensure_reference_data(db)
    as_user(app, gcu, gcuak, seed_user(db, "purchasing"))
    return app, db


def test_the_detail_answers_is_transferable(scm_app):
    app, db = _so_client(scm_app)
    book = _book(db)

    with TestClient(app) as c:
        got = {
            key: c.get(f"{SO_BASE}/{book['orders'][key][0].id}")
            for key in ("F", "T", "U")
        }

    for res in got.values():
        assert res.status_code == 200, res.text
    assert got["F"].json()["is_transferable"] is False
    assert got["T"].json()["is_transferable"] is True
    assert got["U"].json()["is_transferable"] is None


def test_the_list_carries_and_filters_on_it(scm_app):
    app, db = _so_client(scm_app)
    book = _book(db)
    marker = book["marker"]

    def numbers(c, **params):
        res = c.get(SO_BASE, params={"query": marker, "limit": 50, **params})
        assert res.status_code == 200, res.text
        return {row["so_number"]: row["is_transferable"] for row in res.json()["data"]}

    with TestClient(app) as c:
        every = numbers(c)
        no = numbers(c, transferable="no")
        yes = numbers(c, transferable="yes")
        unknown = numbers(c, transferable="unknown")
        nonsense = numbers(c, transferable="maybe")

    assert every == {
        f"{marker}-SOF": False, f"{marker}-SOT": True, f"{marker}-SOU": None,
    }
    assert no == {f"{marker}-SOF": False}
    assert yes == {f"{marker}-SOT": True}
    assert unknown == {f"{marker}-SOU": None}
    # Same rule as `source` / `demand_class`: an unknown word matches nothing rather than
    # being ignored, so a heading never claims a narrowing it did not apply.
    assert nonsense == {}


def test_a_header_edit_cannot_change_it(scm_app):
    """AutoCount owns the field (plan D3): the SO edit PUT has no such input."""
    app, db = _so_client(scm_app)
    book = _book(db)
    order = book["orders"]["F"][0]

    with TestClient(app) as c:
        res = c.put(f"{SO_BASE}/{order.id}", json={"is_transferable": True})
        after = c.get(f"{SO_BASE}/{order.id}").json()

    assert res.status_code in (200, 422), res.text
    assert after["is_transferable"] is False
