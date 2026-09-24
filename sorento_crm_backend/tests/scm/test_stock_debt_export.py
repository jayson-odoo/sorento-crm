"""Stock Debt export: the download route, the worker task, and the workbook (24 Sep slice,
PLAN-stock-debt-filters-totals-export-24sep.md, AC-12..AC-18).

WRITTEN BEFORE THE IMPLEMENTATION EXISTS (Phase 2 is test-first): `POST .../stock-debt/
export`, `app.tasks.export_tasks.generate_stock_debt_xlsx` and `StockDebtService.export()`
are all named by the plan and none of them exist yet. Every not-yet-existing import is made
INSIDE the test that needs it (`_task()`, `from app.services.scm.stock_debt_service import
StockDebtService` inside each test body), so the file COLLECTS and each test fails on its
own missing behaviour rather than one ImportError taking the whole file down.

Postgres only (`requires_pg`), marker-prefixed seeding, nothing borrowed with LIMIT 1 beyond
what `tests.scm.conftest.ensure_reference_data` already guarantees.
"""
from __future__ import annotations

import re
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

from app.services.error_handler import AppException
from tests.scm.conftest import SORENTO_COMPANY_ID, requires_pg
from tests.scm.test_order_sheet_export_downloads import _NoCloseSession
from tests.scm.test_stock_debt_routes import (
    BASE,
    TODAY,
    _client,
    _category,
    _demand,
    _product,
    _product_with_category,
    _row_of,
    _supplier,
    _u,
    _warehouse,
)

pytestmark = requires_pg


def _sheets(blob: bytes):
    from openpyxl import load_workbook

    return load_workbook(BytesIO(blob))


def _task():
    """`app.tasks.export_tasks.generate_stock_debt_xlsx` - fetched by name so its absence
    is an explicit, readable failure rather than an ImportError at collection (AC-12b)."""
    from app.tasks import export_tasks

    fn = getattr(export_tasks, "generate_stock_debt_xlsx", None)
    assert fn is not None, (
        "app.tasks.export_tasks.generate_stock_debt_xlsx does not exist yet (AC-12b)"
    )
    return export_tasks, fn


# =========================================================================== #
# AC-12: the route
# =========================================================================== #


def test_export_route_creates_download_and_enqueues(scm_app, monkeypatch):
    """AC-12: `POST /stock-debt/export` creates a `user_downloads` row
    (`kind=stock_debt_xlsx`, `filename=stock-debt-<ddmmyyyy>.xlsx`, owned by the caller),
    enqueues `generate_stock_debt_xlsx` on `imports`, and answers 201 with the download
    row - behind `projects.stock_debt.view` (403, no row, without it)."""
    from app.services import queue_service

    calls: list[dict] = []

    def _fake_enqueue(func, *args, **kwargs):
        calls.append({"func": func, "args": args, "kwargs": kwargs})
        return type("J", (), {"id": "fake-job-id"})()

    monkeypatch.setattr(queue_service, "enqueue_job", _fake_enqueue)

    app, db = _client(scm_app)
    payload = {
        "split": "supplier", "cutoff": "2026-11-30", "supplier_id": None,
        "book": "all", "only_debt": True,
    }
    with TestClient(app) as c:
        resp = c.post(f"{BASE}/export", json=payload)

    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["kind"] == "stock_debt_xlsx", body

    row = db.execute(
        text("SELECT kind, filename, user_id FROM user_downloads WHERE id = :id"),
        {"id": body["id"]},
    ).mappings().first()
    assert row is not None, "no user_downloads row was created"
    assert row["kind"] == "stock_debt_xlsx"
    assert row["filename"] == f"stock-debt-{date.today().strftime('%d%m%Y')}.xlsx", row

    assert len(calls) == 1, f"expected exactly one enqueue, got {calls}"
    from app.tasks.export_tasks import generate_stock_debt_xlsx

    call = calls[0]
    assert call["func"] is generate_stock_debt_xlsx, call["func"]
    assert call["args"][0] == body["id"], call["args"]
    assert call["kwargs"].get("queue_name") == "imports", call["kwargs"]

    # Without the permission: 403, and no row left behind for the drawer.
    no_perm_app, no_perm_db = _client(scm_app, permission=None)
    before = no_perm_db.execute(text("SELECT count(*) FROM user_downloads")).scalar()
    with TestClient(no_perm_app) as c:
        denied = c.post(f"{BASE}/export", json={"split": "none"})
    assert denied.status_code == 403, denied.text
    after = no_perm_db.execute(text("SELECT count(*) FROM user_downloads")).scalar()
    assert after == before, "a denied export left a download row behind"


# =========================================================================== #
# AC-12b: the task
# =========================================================================== #


def test_generate_stock_debt_xlsx_marks_ready_and_failed(scm_app, monkeypatch):
    """AC-12b: the task mirrors `generate_low_stock_report` - mark_processing, build the
    workbook through `StockDebtService.export()`, store it, then `mark_ready` with
    `row_count` and `sheet_count`; on any exception it marks the row `failed` with the
    message and never raises into RQ."""
    from app.services.download_service import DownloadService

    export_tasks, task_fn = _task()
    app, db = _client(scm_app)
    marker = f"ZZTSD{_u()[:6]}".upper()
    warehouse = _warehouse(db, f"ZZTBRW{_u()[:4]}-BB")
    product = _product(db, f"{marker}-A")
    _demand(
        db, product, warehouse, qty=5, required_date=date(2026, 11, 10),
        so_number=f"{marker}-SO1",
    )
    user_id = db.execute(text("SELECT id FROM users LIMIT 1")).scalar()
    assert user_id, "seed_user (via _client) must have left a user row"
    download = DownloadService(db).create(
        user_id=str(user_id), kind="stock_debt_xlsx", filename="zzt-stock-debt-test.xlsx",
    )
    db.flush()

    # The task's own `SessionLocal()` must reuse THIS rolled-back session, or it opens a
    # separate connection that cannot see the uncommitted seed above.
    monkeypatch.setattr(export_tasks, "SessionLocal", lambda: _NoCloseSession(db))

    calls: list[dict] = []
    original_mark_ready = DownloadService.mark_ready

    def _fake_mark_ready(self, download_id, **kwargs):
        calls.append(kwargs)
        return original_mark_ready(self, download_id, **kwargs)

    monkeypatch.setattr(DownloadService, "mark_ready", _fake_mark_ready)

    class _StubBackend:
        @staticmethod
        def upload_file(*, file_content, file_path, content_type):
            return file_path, None

    monkeypatch.setattr(export_tasks, "get_backend", lambda provider: _StubBackend())
    monkeypatch.setattr(export_tasks, "default_provider", lambda: "s3")

    task_fn(str(download.id), str(user_id), {"query": marker, "split": "none"})

    row = DownloadService(db).get(str(download.id))
    assert row is not None
    assert row.status == "ready", getattr(row, "error", None)
    assert calls, "mark_ready was never called"
    assert calls[-1].get("row_count") == 1, calls[-1]
    assert calls[-1].get("sheet_count") == 1, calls[-1]

    # The failure path: the service raising marks the row failed and swallows the error.
    download2 = DownloadService(db).create(
        user_id=str(user_id), kind="stock_debt_xlsx", filename="zzt-stock-debt-test2.xlsx",
    )
    db.flush()
    monkeypatch.setattr(
        "app.services.scm.stock_debt_service.StockDebtService.export",
        lambda self, **kwargs: (_ for _ in ()).throw(RuntimeError("boom")),
    )
    task_fn(str(download2.id), str(user_id), {"query": marker, "split": "none"})
    row2 = DownloadService(db).get(str(download2.id))
    assert row2 is not None
    assert row2.status == "failed", getattr(row2, "error", None)
    assert "boom" in (row2.error or "")


# =========================================================================== #
# AC-13..AC-18: the workbook (`StockDebtService.export()`, called directly)
# =========================================================================== #


def test_export_split_none_one_sheet_with_total_row(scm_app):
    """AC-13: `split=none` writes ONE sheet, "Stock debt" - Product, Name, Category,
    Supplier, one column per axis month (`Sep 26` style), TBA, No date, No location,
    Total - with a `Total` footer row summing every column. A name equal to its code
    prints blank."""
    app, db = _client(scm_app)
    from app.services.scm.stock_debt_service import StockDebtService

    marker = f"ZZTSD{_u()[:6]}".upper()
    warehouse = _warehouse(db, f"ZZTBRW{_u()[:4]}-BB")
    equal_code = f"{marker}-EQ"
    product = _product(db, equal_code)
    product.product_name = equal_code
    _demand(
        db, product, warehouse, qty=10, required_date=date(2026, 11, 10),
        so_number=f"{marker}-SO1",
    )
    db.flush()

    svc = StockDebtService(db)
    listing = svc.list(query=marker, only_debt=False, limit=50)
    row = _row_of(listing, equal_code)

    blob, content_type, filename, counts = svc.export(
        query=marker, only_debt=False, split="none",
    )
    assert content_type == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert filename == f"stock-debt-{date.today().strftime('%d%m%Y')}.xlsx", filename
    assert counts == {"rows": 1, "sheets": 1}, counts

    wb = _sheets(blob)
    assert wb.sheetnames == ["Stock debt"], wb.sheetnames
    ws = wb["Stock debt"]
    header = [cell.value for cell in ws[1]]
    assert header[:4] == ["Product", "Name", "Category", "Supplier"], header
    assert header[-4:] == ["TBA", "No date", "No location", "Total"], header
    assert len(header) == 8 + len(listing["months"]), header
    month_labels = header[4:-4]
    assert len(month_labels) == len(listing["months"])
    for label in month_labels:
        assert re.match(r"^[A-Za-z]{3} \d{2}$", str(label)), label

    data_row = [cell.value for cell in ws[2]]
    assert data_row[0] == equal_code
    assert data_row[1] in (None, ""), "a name equal to its code must print blank"

    total_row = [cell.value for cell in ws[ws.max_row]]
    assert total_row[0] == "Total", total_row
    assert total_row[-1] == row["total"], (total_row, row["total"])


def test_export_split_supplier_one_sheet_each_plus_no_supplier(scm_app):
    """AC-14: `split=supplier` writes one sheet per LAST supplier name, plus "No
    supplier" for a product with none, sorted by title, each with its own Total row."""
    app, db = _client(scm_app)
    from app.models.procurement import ProductSupplier
    from app.services.scm.stock_debt_service import StockDebtService

    marker = f"ZZTSD{_u()[:6]}".upper()
    warehouse = _warehouse(db, f"ZZTBRW{_u()[:4]}-BB")
    supplier_z = _supplier(db, f"ZZTSZ{_u()[:4]}".upper())
    supplier_z.supplier_name = "Zeta supplier"
    supplier_a = _supplier(db, f"ZZTSA{_u()[:4]}".upper())
    supplier_a.supplier_name = "Alpha supplier"
    db.flush()

    p_z = _product(db, f"{marker}-Z")
    db.add(
        ProductSupplier(
            id=_u(), product_id=p_z.id, supplier_id=supplier_z.id,
            standard_lead_time_days=7, is_primary_supplier=True,
            company_id=SORENTO_COMPANY_ID,
        )
    )
    _demand(
        db, p_z, warehouse, qty=5, required_date=date(2026, 11, 10),
        so_number=f"{marker}-SOZ",
    )

    p_a = _product(db, f"{marker}-Y")
    db.add(
        ProductSupplier(
            id=_u(), product_id=p_a.id, supplier_id=supplier_a.id,
            standard_lead_time_days=7, is_primary_supplier=True,
            company_id=SORENTO_COMPANY_ID,
        )
    )
    _demand(
        db, p_a, warehouse, qty=5, required_date=date(2026, 11, 10),
        so_number=f"{marker}-SOY",
    )

    p_none = _product(db, f"{marker}-N")
    _demand(
        db, p_none, warehouse, qty=5, required_date=date(2026, 11, 10),
        so_number=f"{marker}-SON",
    )
    db.flush()

    svc = StockDebtService(db)
    blob, _ct, _fn, counts = svc.export(query=marker, only_debt=False, split="supplier")
    assert counts["sheets"] == 3, counts

    wb = _sheets(blob)
    expected_order = sorted(["Alpha supplier", "Zeta supplier", "No supplier"])
    assert wb.sheetnames == expected_order, wb.sheetnames
    for name in wb.sheetnames:
        ws = wb[name]
        last_row = [cell.value for cell in ws[ws.max_row]]
        assert last_row[0] == "Total", (name, last_row)


def test_export_split_category(scm_app):
    """AC-15: `split=category` writes one sheet per `category_code`, "No category" for
    blanks."""
    app, db = _client(scm_app)
    from app.services.scm.stock_debt_service import StockDebtService

    marker = f"ZZTSD{_u()[:6]}".upper()
    warehouse = _warehouse(db, f"ZZTBRW{_u()[:4]}-BB")
    cat_x = _category(db, f"{marker}-CATX")
    cat_y = _category(db, f"{marker}-CATY")

    p_x = _product_with_category(db, f"{marker}-X", cat_x.id)
    _demand(
        db, p_x, warehouse, qty=5, required_date=date(2026, 11, 10),
        so_number=f"{marker}-SOX",
    )
    p_y = _product_with_category(db, f"{marker}-Y", cat_y.id)
    _demand(
        db, p_y, warehouse, qty=5, required_date=date(2026, 11, 10),
        so_number=f"{marker}-SOY",
    )
    # `products.category_id` is NOT NULL, so "no category" is a BLANK `category_code`
    # (the same convention `low_stock_report_service._master_map` already uses), never a
    # null FK.
    no_category = _category(db, "")
    p_none = _product_with_category(db, f"{marker}-N", no_category.id)
    _demand(
        db, p_none, warehouse, qty=5, required_date=date(2026, 11, 10),
        so_number=f"{marker}-SON",
    )
    db.flush()

    svc = StockDebtService(db)
    blob, _ct, _fn, counts = svc.export(query=marker, only_debt=False, split="category")
    assert counts["sheets"] == 3, counts

    wb = _sheets(blob)
    expected_order = sorted([cat_x.category_code, cat_y.category_code, "No category"])
    assert wb.sheetnames == expected_order, wb.sheetnames


def test_export_split_pair_titles_sanitised(scm_app):
    """AC-16: `split=supplier_category` writes one sheet per `<supplier> - <category>`
    pair, title cut to 31 chars with `[]:*?/\\` removed; two pairs that collide after the
    cut get a `(2)` suffix."""
    app, db = _client(scm_app)
    from app.models.procurement import ProductSupplier
    from app.services.scm.stock_debt_service import StockDebtService

    marker = f"ZZTSD{_u()[:6]}".upper()
    warehouse = _warehouse(db, f"ZZTBRW{_u()[:4]}-BB")
    # Both names share their first 40 characters, so titling either against the SAME
    # category and cutting to 31 chars collapses them onto the identical string.
    base = "A" * 40
    supplier1 = _supplier(db, f"ZZTSA{_u()[:4]}".upper())
    supplier1.supplier_name = f"{base}ONE"
    supplier2 = _supplier(db, f"ZZTSB{_u()[:4]}".upper())
    supplier2.supplier_name = f"{base}TWO"
    category = _category(db, f"{marker}-SHARED")
    db.flush()

    p1 = _product_with_category(db, f"{marker}-P1", category.id)
    db.add(
        ProductSupplier(
            id=_u(), product_id=p1.id, supplier_id=supplier1.id,
            standard_lead_time_days=7, is_primary_supplier=True,
            company_id=SORENTO_COMPANY_ID,
        )
    )
    _demand(
        db, p1, warehouse, qty=5, required_date=date(2026, 11, 10),
        so_number=f"{marker}-SO1",
    )

    p2 = _product_with_category(db, f"{marker}-P2", category.id)
    db.add(
        ProductSupplier(
            id=_u(), product_id=p2.id, supplier_id=supplier2.id,
            standard_lead_time_days=7, is_primary_supplier=True,
            company_id=SORENTO_COMPANY_ID,
        )
    )
    _demand(
        db, p2, warehouse, qty=5, required_date=date(2026, 11, 10),
        so_number=f"{marker}-SO2",
    )
    db.flush()

    svc = StockDebtService(db)
    blob, _ct, _fn, counts = svc.export(
        query=marker, only_debt=False, split="supplier_category",
    )
    assert counts["sheets"] == 2, counts

    wb = _sheets(blob)
    assert len(wb.sheetnames) == 2, wb.sheetnames
    for title in wb.sheetnames:
        assert len(title) <= 31, title
        for forbidden in "[]:*?/\\":
            assert forbidden not in title, title
    suffixed = [t for t in wb.sheetnames if t.endswith("(2)")]
    assert len(suffixed) == 1, wb.sheetnames
    plain = next(t for t in wb.sheetnames if t not in suffixed)
    assert suffixed[0].startswith(plain[: min(len(plain), 20)]), (suffixed, plain)


def test_export_honours_list_filters(scm_app):
    """AC-17: the export's rows are exactly the list's rows for the same filters,
    unpaged - a cutoff that drops a product from the list drops it from the workbook
    too."""
    app, db = _client(scm_app)
    from app.services.scm.stock_debt_service import StockDebtService

    marker = f"ZZTSD{_u()[:6]}".upper()
    warehouse = _warehouse(db, f"ZZTBRW{_u()[:4]}-BB")
    kept = _product(db, f"{marker}-KEEP")
    dropped = _product(db, f"{marker}-DROP")
    _demand(
        db, kept, warehouse, qty=5, required_date=date(2026, 11, 10),
        so_number=f"{marker}-SO1",
    )
    _demand(
        db, dropped, warehouse, qty=5, required_date=date(2026, 12, 5),
        so_number=f"{marker}-SO2",
    )
    db.flush()

    svc = StockDebtService(db)
    filters = dict(query=marker, only_debt=True, cutoff=date(2026, 11, 30))
    listing = svc.list(**filters, limit=50)
    blob, _ct, _fn, counts = svc.export(**filters, split="none")

    assert counts["rows"] == listing["pagination"]["total"], counts
    wb = _sheets(blob)
    ws = wb["Stock debt"]
    exported_codes = {
        row[0]
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row - 1, values_only=True)
    }
    listed_codes = {row["product_code"] for row in listing["data"]}
    assert exported_codes == listed_codes, (exported_codes, listed_codes)
    assert kept.product_code in exported_codes
    assert dropped.product_code not in exported_codes


def test_export_refuses_above_cap(scm_app, monkeypatch):
    """AC-18: above the low stock report's own `MAX_LOW_STOCK_ROWS` the export raises a
    422 "Narrow the plan first" and writes nothing."""
    app, db = _client(scm_app)
    from app.services.scm import low_stock_report_service
    from app.services.scm import stock_debt_service as svc_mod
    from app.services.scm.stock_debt_service import StockDebtService

    marker = f"ZZTSD{_u()[:6]}".upper()
    warehouse = _warehouse(db, f"ZZTBRW{_u()[:4]}-BB")
    for stem in ("A", "B"):
        product = _product(db, f"{marker}-{stem}")
        _demand(
            db, product, warehouse, qty=5, required_date=date(2026, 11, 10),
            so_number=f"{marker}-SO-{stem}",
        )
    db.flush()

    # Patched on both modules: whichever one the coder reads the cap off at call time.
    monkeypatch.setattr(low_stock_report_service, "MAX_LOW_STOCK_ROWS", 1)
    monkeypatch.setattr(svc_mod, "MAX_LOW_STOCK_ROWS", 1, raising=False)

    with pytest.raises(AppException) as excinfo:
        StockDebtService(db).export(query=marker, only_debt=False, split="none")
    assert excinfo.value.status_code == 422
    assert "Narrow the plan first" in str(excinfo.value.detail)
