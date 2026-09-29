"""R44 (owner, 29 Sep 2026, #1359): on the stock debt page the overdue rule does not apply.

"the supply here doesn't care about anything received, same like SO, nothing delivered also
is fine, I just need to know what's my sold quantity (demand) and purchased quantity
(supply), so I don't really care about the fulfilment"

The issue's shape, `_production`'s seed: SO419208 x CSK14A-NL, 1,309 outstanding on two
lines pinned to PO 202609-S0029 lines of 4 and 1,305, and a third, unpinned line of 41 on the
same PO. Every PO date has passed and nothing is received. Before R44 the 41 was "overdue,
not counted" and the month cell read 0; it reads +41 whatever the policy's grace and dead
numbers say, because this page counts every outstanding PO line as supply in its arrival
month, the way an SO line counts as demand whether delivered or not.

The board path (`assignments_for`, the ladder / coverage / front planning) keeps the rule:
`test_stock_debt_po_supply_routes.test_the_board_path_is_unchanged_and_differs_from_the_view_only_by_r42`
and `test_overdue_grace_ladder.py` guard it.
"""
from __future__ import annotations

from io import BytesIO

import pytest
from fastapi.testclient import TestClient

from app.services.scm.supply_assignment import month_key
from tests.scm.conftest import requires_pg
from tests.scm.test_stock_debt_pinned_po_fulfils_routes import (
    _assert_fulfilled,
    _cell,
    _production,
)
from tests.scm.test_stock_debt_routes import BASE, TODAY, _client, _row_of

pytestmark = requires_pg


def _board_row(c, seed):
    got = c.get(BASE, params={"query": seed["marker"], "only_debt": False, "limit": 25})
    assert got.status_code == 200, got.text
    return _row_of(got.json(), seed["product"].product_code)


@pytest.mark.parametrize("grace,dead", [(0, 0), (14, 90), (45, 45)])
def test_every_outstanding_po_line_is_supply_whatever_the_overdue_policy(
    scm_app, grace, dead
):
    """Cell +41, row total +41, the 41 line Free 41 and never "overdue, not counted", the two
    pinned lines still fulfil their SO lines (Short 0), under any grace/dead policy."""
    app, db = _client(scm_app)
    seed = _production(db, grace=grace, dead=dead)
    month = month_key(TODAY)

    with TestClient(app) as c:
        cell = _cell(c, seed, month)
        row = _board_row(c, seed)

    _assert_fulfilled(cell, seed)
    balances = {m["key"]: m["balance"] for m in row["months"]}
    assert balances[month] == 41, balances
    assert all(v == 0 for k, v in balances.items() if k != month), balances
    assert row["total"] == 41

    po_rows = {r["po_line_id"]: r for r in cell["supply"] if r["kind"] == "po"}
    assert set(po_rows) == {
        str(seed[k].id) for k in ("free", "line_2", "line_3")
    }, "all three PO lines are listed in the month they arrive in"
    assert all(r["overdue"] is False for r in po_rows.values()), po_rows
    free = po_rows[str(seed["free"].id)]
    assert (free["free_qty"], free["outstanding_qty"], free["assigned_to"]) == (41, 41, [])
    # Late is information, never exclusion.
    assert free["days_late"] == 18
    assert po_rows[str(seed["line_2"].id)]["free_qty"] == 0
    assert po_rows[str(seed["line_3"].id)]["free_qty"] == 0
    assert cell["supply_total_qty"] == 1350
    assert cell["demand_total_qty"] == 1309
    # The drill foots with the cell: Free less Short over the month.
    assert sum(r["free_qty"] for r in cell["supply"]) - sum(
        r["short_qty"] for r in cell["demand"]
    ) == 41


def test_the_export_reads_the_same_plus_41(scm_app):
    """The workbook is built off `list()`, so it carries the same +41 in the month and Total."""
    from openpyxl import load_workbook

    from app.services.scm.stock_debt_service import StockDebtService

    _app, db = _client(scm_app)
    seed = _production(db, grace=0, dead=0)
    blob, _type, _name, counts = StockDebtService(db).export(
        query=seed["marker"], only_debt=False, split="none",
    )
    assert counts["rows"] == 1
    ws = load_workbook(BytesIO(blob))["Stock debt"]
    header = [cell.value for cell in ws[1]]
    data = [cell.value for cell in ws[2]]
    assert data[0] == seed["product"].product_code
    assert data[header.index("Total")] == 41
    assert data[header.index(TODAY.strftime("%b %y"))] == 41
