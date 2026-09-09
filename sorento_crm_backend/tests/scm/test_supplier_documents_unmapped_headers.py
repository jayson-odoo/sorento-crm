"""S5 - AC-E2: every preview response carries `unmapped_headers` per file.

TEST-FIRST (Phase 2): `AliasResolver.unmapped_headers()` already exists and both readers
already populate it on their own `ProformaReadResult`/`PackingReadResult` (confirmed by
reading `proforma_invoice_reader.py:484` and `packing_list_reader.py:483` directly) - what
is missing is `supplier_document_service._file_preview` surfacing it on the dict `preview()`
returns. So this file is red today with a plain `KeyError`/missing-key assertion, not an
`ImportError` - the resolver-level mechanism this AC reuses is already built.

The E2/E5 variant fixture (Jinbaichuan's `箱数` header renamed to the traditional `箱數`, so
a header that resolves TODAY stops resolving) is built here with openpyxl, from the real
committed `Jinbaichuan_Invoice.xlsx`, rather than committed as a second binary - the rename
is the whole point and reads clearly as one `ws.cell(...).value = ...` line.
"""
from __future__ import annotations

import importlib.util
import uuid
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from app.config import settings
from app.models.procurement import Supplier
from app.models.product import Product, ProductCategory, UnitOfMeasure
from app.services.scm import supplier_document_service as svc
from tests._pg_fixture import blank_session

pytestmark = pytest.mark.usefixtures("no_live_llm")

MARKER = "ZZUNMAP"
_VERSIONS = Path(__file__).resolve().parents[2] / "alembic" / "versions"
_LANE_FIXTURES = Path(__file__).resolve().parents[3] / "documentation" / "plans" / "scm" / "fixtures"

_JINBAICHUAN_CODES = [
    "SRTWC8366-RL-300", "SRTWC8366-RL-250", "CWB242", "CWB242海关样品",
    "SRTWC8152-SH-250-UF", "MWB243", "MWB243海关样品", "SRTWC286-SH-250-NEW",
]


@pytest.fixture(autouse=True)
def _no_ai_translation(monkeypatch):
    monkeypatch.setattr(settings, "openai_api_key", None, raising=False)


def _load(name: str):
    spec = importlib.util.spec_from_file_location(name, _VERSIONS / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def _jinbaichuan_with_renamed_carton_header() -> bytes:
    """The real Jinbaichuan file, header row 8: `箱数` (simplified, resolves via the
    migration seed) renamed `箱數` (traditional) - the exact E2/E5 scenario."""
    wb = openpyxl.load_workbook(_LANE_FIXTURES / "Jinbaichuan_Invoice.xlsx")
    ws = wb.active
    found = False
    for row in ws.iter_rows():
        for cell in row:
            if cell.value == "箱数":
                cell.value = "箱數"
                found = True
    assert found, "fixture no longer states 箱数 - update this test's rename target"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_e2_the_renamed_carton_header_is_reported_as_unmapped():
    with blank_session() as db:
        conn = db.connection()
        _load("311_scm_purchasing_base").seed_import_field_aliases(conn)
        _load("375_scm_proforma_invoice").seed(conn)
        _load("428_scm_pi_cbm_adjust_revision").seed(conn)
        _load("483_supplier_doc_aliases").seed(conn)
        db.commit()

        tag = uuid.uuid4().hex[:8].upper()
        cat = ProductCategory(id=str(uuid.uuid4()), category_code=f"{MARKER}-CAT-{tag}", category_name="c")
        uom = UnitOfMeasure(id=str(uuid.uuid4()), uom_code=f"{MARKER}U"[:20], uom_name="pcs")
        db.add_all([cat, uom])
        db.flush()
        supplier = Supplier(
            id=str(uuid.uuid4()), supplier_code=f"{MARKER}-{tag}", supplier_name="Jinbaichuan",
            is_active=True,
        )
        db.add(supplier)
        db.flush()
        for code in _JINBAICHUAN_CODES:
            db.add(
                Product(
                    id=str(uuid.uuid4()), product_code=code, product_name=code,
                    category_id=cat.id, base_uom_id=uom.id, list_price=0,
                    is_active=True, is_discontinued=False,
                )
            )
        db.flush()

        out = svc.preview(
            db, [("Jinbaichuan_Invoice.xlsx", _jinbaichuan_with_renamed_carton_header())],
            supplier_id=str(supplier.id),
        )

        assert len(out["files"]) == 1
        file_preview = out["files"][0]
        assert "unmapped_headers" in file_preview, (
            "AC-E2: the preview response must carry unmapped_headers per file"
        )
        assert "箱數" in file_preview["unmapped_headers"]
        # The pre-existing, always-unmapped headers (the plan's own measured fact) are
        # still reported too - this AC does not narrow that set, only adds the renamed one.
        assert "尺寸（mm）" in file_preview["unmapped_headers"]
