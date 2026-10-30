"""DO-IMPORT-DATE: an undated DO is imported and flagged, never refused, and a stored DO date
is never cleared.

Owner, 30 Oct 2026: no existing import or integration may fail because of a DO date (manual
DOs can come without one); unreadable dates stay out of reports. The import used to write NULL
over a stored date and gave no word about a blank or unreadable cell.
Plan: documentation/plans/order-management/PLAN-do-import-date.md.
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.models.order import Order, OrderStatus
from app.schemas.order import OrderCreate, OrderResponse, OrderUpdate
from app.services import import_outcome_codes as oc
from app.services.import_outcome import ImportOutcome
from app.services.order_service import OrderService
from tests._pg_fixture import blank_session, unique_code

USER = "11111111-1111-1111-1111-111111111111"


@pytest.fixture()
def db():
    with blank_session() as session:
        for code, name in [("NEW", "New Order"), ("DELIVERED", "Delivered")]:
            session.add(OrderStatus(status_code=code, status_name=name))
        session.commit()
        yield session


def _workbook(master_rows: list[list]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    master = wb.create_sheet("Master")
    master.append(["Doc. No.", "Date", "Type"])
    for r in master_rows:
        master.append(r)
    tracking = wb.create_sheet("Overall Tracking")
    tracking.append(["Doc Number", "Date"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _order(db, number: str) -> Order | None:
    return db.query(Order).filter(Order.order_number == number).first()


def _row_errors(result: dict) -> dict[int, str]:
    return {e["row"]: e["error"] for e in result.get("errors") or [] if isinstance(e, dict)}


def _successful_codes(outcome: ImportOutcome) -> set[str]:
    return {e["code"] for e in outcome.breakdown()["successful"]}


def _warnings(result: dict) -> list[str]:
    return [w.get("warning") if isinstance(w, dict) else w for w in result.get("warnings") or []]


UNDATED = "imported without a DO date (left out of reports until it has one)"


# --- Excel tracking import ------------------------------------------------------------


def test_import_creates_master_row_with_blank_date_flagged(db):
    good, blank = unique_code("DO"), unique_code("RMA")
    outcome = ImportOutcome(None)
    result = OrderService(db).import_excel_tracking(
        _workbook([[good, date(2026, 8, 28), "DO"], [blank, None, "RMA"]]), USER, outcome=outcome
    )

    assert not _row_errors(result), result.get("errors")
    assert f"{blank}: no Date, {UNDATED}" in _warnings(result)
    created = _order(db, blank)
    assert created is not None and created.order_date is None
    assert _order(db, good).order_date == date(2026, 8, 28)
    assert oc.MISSING_DOC_DATE in _successful_codes(outcome)
    assert outcome.successful == 2


def test_import_creates_master_row_with_unparseable_date_flagged(db):
    bad = unique_code("RMA")
    outcome = ImportOutcome(None)
    result = OrderService(db).import_excel_tracking(
        _workbook([[bad, "not a date", "RMA"]]), USER, outcome=outcome
    )

    assert not _row_errors(result), result.get("errors")
    assert (
        f"{bad}: Date 'not a date' is not a valid date (use DD/MM/YYYY), {UNDATED}"
    ) in _warnings(result)
    assert _order(db, bad).order_date is None
    assert oc.INVALID_DOC_DATE in _successful_codes(outcome)


def test_import_blank_date_does_not_wipe_existing_do_date(db):
    number = unique_code("DO")
    db.add(Order(order_number=number, order_date=date(2026, 7, 1)))
    db.commit()

    result = OrderService(db).import_excel_tracking(_workbook([[number, None, "RMA"]]), USER)

    assert not _row_errors(result), result.get("errors")
    assert not _warnings(result), "nothing lost, nothing to say"
    order = _order(db, number)
    db.refresh(order)
    assert order.order_date == date(2026, 7, 1)
    assert order.order_type == "RMA", "the row's other columns still apply"


def test_import_unreadable_date_keeps_existing_do_date_and_says_so(db):
    number = unique_code("DO")
    db.add(Order(order_number=number, order_date=date(2026, 7, 1)))
    db.commit()

    result = OrderService(db).import_excel_tracking(_workbook([[number, "31/13/2026", "DO"]]), USER)

    assert (
        f"{number}: Date '31/13/2026' is not a valid date (use DD/MM/YYYY), kept the stored DO date"
    ) in _warnings(result)
    order = _order(db, number)
    db.refresh(order)
    assert order.order_date == date(2026, 7, 1)


def test_import_validate_only_warns_about_undated_row(db):
    bad = unique_code("RMA")
    result = OrderService(db).import_excel_tracking(
        _workbook([[bad, None, "RMA"]]), USER, validate_only=True
    )
    # The preview lists bare messages, so the doc number rides in the text.
    assert result["valid"] is True
    assert result["errors"] == []
    assert f"{bad}: no Date, {UNDATED}" in result["warnings"]
    assert result["summary"]["would_create"] == 1


# --- API schemas ----------------------------------------------------------------------


def test_order_create_accepts_a_missing_order_date():
    assert OrderCreate(order_number="DO-1").order_date is None
    assert OrderCreate(order_number="DO-1", order_date=None).order_date is None
    assert OrderCreate(order_number="DO-1", order_date="2026-10-29").order_date is not None


def test_order_update_null_order_date_keeps_the_stored_date():
    # A sent null is dropped, exactly like leaving the field out.
    assert "order_date" not in OrderUpdate(order_date=None).model_dump(exclude_unset=True)
    assert "order_date" not in OrderUpdate(remarks="x").model_dump(exclude_unset=True)
    assert OrderUpdate(order_date="2026-10-29").order_date is not None


def test_update_order_with_null_date_does_not_clear_it(db):
    order = Order(order_number=unique_code("DO"), order_date=date(2026, 7, 1))
    db.add(order)
    db.commit()
    updated = OrderService(db).update_order(
        order.id, OrderUpdate.model_validate({"order_date": None, "remarks": "kept"}), USER
    )
    assert updated.remarks == "kept" and updated.order_date == date(2026, 7, 1)


def test_legacy_undated_row_still_reads(db):
    number = unique_code("RMA")
    order = Order(order_number=number, order_date=None)
    db.add(order)
    db.commit()
    db.refresh(order)
    assert OrderResponse.model_validate(order).order_date is None


def test_update_without_date_on_legacy_undated_row_still_works(db):
    order = Order(order_number=unique_code("RMA"), order_date=None)
    db.add(order)
    db.commit()
    updated = OrderService(db).update_order(order.id, OrderUpdate(remarks="kept"), USER)
    assert updated.remarks == "kept" and updated.order_date is None


def test_create_order_route_accepts_missing_date(db):
    from fastapi.testclient import TestClient

    from app.database import get_db
    from app.dependencies import get_current_user, get_current_user_or_api_key
    from app.main import app
    from app.services.company_scope_resolver import apply_company_scope

    number = unique_code("DO")
    principal = {"id": USER, "roles": ["superadmin"]}
    overrides = {
        get_current_user: lambda: principal,
        get_current_user_or_api_key: lambda: principal,
        apply_company_scope: lambda: None,
        get_db: lambda: db,
    }
    app.dependency_overrides.update(overrides)
    try:
        resp = TestClient(app).post("/api/v1/order-management/orders/", json={"order_number": number})
    finally:
        for dep in overrides:
            app.dependency_overrides.pop(dep, None)
    assert resp.status_code in (200, 201), resp.text
    assert resp.json()["order_date"] is None
    assert _order(db, number).order_date is None


# --- JSON bulk import -----------------------------------------------------------------


def test_bulk_import_creates_new_order_without_date_flagged(db):
    number = unique_code("DO")
    result = OrderService(db).bulk_import_orders([{"Order Number": number}], USER)
    assert result["created"] == 1 and result["errors"] == [], result
    assert result["warnings"] == [f"Row 1: Order {number}: no Date, {UNDATED}"]
    assert _order(db, number).order_date is None


def test_bulk_import_blank_date_keeps_existing_date(db):
    number = unique_code("DO")
    db.add(Order(order_number=number, order_date=date(2026, 7, 1)))
    db.commit()
    result = OrderService(db).bulk_import_orders(
        [{"Order Number": number, "Order Date": ""}], USER
    )
    assert result["updated"] == 1 and result["errors"] == [], result
    order = _order(db, number)
    db.refresh(order)
    assert order.order_date == date(2026, 7, 1)


def test_bulk_import_update_without_date_column_keeps_date(db):
    number = unique_code("DO")
    db.add(Order(order_number=number, order_date=date(2026, 7, 1)))
    db.commit()
    result = OrderService(db).bulk_import_orders([{"Order Number": number, "Remarks": "x"}], USER)
    assert result["updated"] == 1, result
    order = _order(db, number)
    db.refresh(order)
    assert order.order_date == date(2026, 7, 1)


def test_bulk_import_creates_dated_order(db):
    number = unique_code("DO")
    result = OrderService(db).bulk_import_orders(
        [{"Order Number": number, "Order Date": "29/10/2026"}], USER
    )
    assert result["created"] == 1, result
    assert _order(db, number).order_date == date(2026, 10, 29)


# --- read-only report of existing undated DOs -----------------------------------------


def test_report_lists_live_undated_dos_only(db):
    from scripts.report_undated_dos import undated_dos

    undated = unique_code("RMA")
    dated = unique_code("DO")
    gone = unique_code("RMA")
    db.add_all([
        Order(order_number=undated, order_date=None, order_type="RMA"),
        Order(order_number=dated, order_date=date(2026, 7, 1)),
        Order(order_number=gone, order_date=None, deleted_at=datetime(2026, 9, 1)),
    ])
    db.flush()

    rows = {r["order_number"]: r for r in undated_dos(db.connection())}
    assert undated in rows and dated not in rows and gone not in rows
    assert rows[undated]["line_count"] == 0
    assert rows[undated]["from_autocount"] is False
    assert rows[undated]["order_type"] == "RMA"


# --- review round 1 -------------------------------------------------------------------


def _sheet(master_header: list, master_rows: list[list], tracking_rows: list[list] = ()) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    master = wb.create_sheet("Master")
    master.append(master_header)
    for r in master_rows:
        master.append(r)
    tracking = wb.create_sheet("Overall Tracking")
    tracking.append(["Doc Number", "Date"])
    for r in tracking_rows:
        tracking.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _excel_serial(d: date) -> int:
    return (d - date(1899, 12, 30)).days


def test_import_blank_date_on_autocount_do_still_applies_sheet_columns(db):
    """AutoCount owns order_date on its own DOs, so the sheet's Date cell is irrelevant
    there: a blank one must not block the columns the sheet does own."""
    number = unique_code("DO")
    db.add(Order(order_number=number, order_date=date(2026, 7, 1), doc_key=990001, remarks_cs="old"))
    db.commit()

    result = OrderService(db).import_excel_tracking(
        _sheet(["Doc. No.", "Date", "Remarks CS"], [[number, None, "new"]]), USER
    )

    assert not _row_errors(result), result.get("errors")
    order = _order(db, number)
    db.refresh(order)
    assert order.remarks_cs == "new"
    assert order.order_date == date(2026, 7, 1)


def test_import_reads_an_excel_date_serial(db):
    number = unique_code("DO")
    result = OrderService(db).import_excel_tracking(
        _sheet(["Doc. No.", "Date"], [[number, _excel_serial(date(2026, 10, 29))]]), USER
    )
    assert not _row_errors(result), result.get("errors")
    assert _order(db, number).order_date == date(2026, 10, 29)


def test_import_tracking_row_of_undated_master_row_still_applies(db):
    number = unique_code("RMA")
    result = OrderService(db).import_excel_tracking(
        _sheet(["Doc. No.", "Date"], [[number, None]], [[number, date(2026, 10, 30)]]), USER
    )
    assert not any("not found in Master sheet" in (w or "") for w in _warnings(result))
    order = _order(db, number)
    assert order.order_date is None
    assert order.actual_delivery_date == date(2026, 10, 30)


def test_import_master_sheet_without_date_column_imports_with_one_warning(db):
    number = unique_code("DO")
    result = OrderService(db).import_excel_tracking(
        _sheet(["Doc. No.", "Type"], [[number, "DO"]]), USER
    )
    assert not _row_errors(result), result.get("errors")
    assert (
        "Master sheet has no Date column: new delivery orders are imported without a DO date"
    ) in _warnings(result)
    assert _order(db, number).order_date is None


def test_bulk_import_reads_an_excel_date_serial(db):
    """The FE template upload reads cells with SheetJS, so a real date arrives as a serial."""
    number = unique_code("DO")
    result = OrderService(db).bulk_import_orders(
        [{"Order Number": number, "Order Date": _excel_serial(date(2026, 10, 29))}], USER
    )
    assert result["created"] == 1, result
    assert _order(db, number).order_date == date(2026, 10, 29)


def test_bulk_import_names_an_unreadable_date(db):
    number = unique_code("DO")
    result = OrderService(db).bulk_import_orders([{"Order Number": number, "Order Date": "soon"}], USER)
    assert result["created"] == 1 and result["errors"] == [], result
    assert result["warnings"] == [
        f"Row 1: Order {number}: Date 'soon' is not a valid date (use DD/MM/YYYY), {UNDATED}"
    ]
    assert _order(db, number).order_date is None


def test_bulk_import_blank_date_on_autocount_do_is_not_refused(db):
    number = unique_code("DO")
    db.add(Order(order_number=number, order_date=date(2026, 7, 1), doc_key=990002))
    db.commit()
    result = OrderService(db).bulk_import_orders(
        [{"Order Number": number, "Order Date": "", "Remarks": "x"}], USER
    )
    assert result["updated"] == 1, result
    order = _order(db, number)
    db.refresh(order)
    assert order.order_date == date(2026, 7, 1)


def test_import_number_outside_plausible_dates_is_flagged_not_dated(db):
    """A typed ddmmyy (280826) is not a serial for the year 2668: no date, said visibly."""
    number = unique_code("DO")
    result = OrderService(db).import_excel_tracking(
        _sheet(["Doc. No.", "Date"], [[number, 280826]]), USER
    )
    assert (
        f"{number}: Date '280826' is not a valid date (use DD/MM/YYYY), {UNDATED}"
    ) in _warnings(result)
    assert _order(db, number).order_date is None


def test_import_empty_master_tab_still_applies_tracking(db):
    number = unique_code("DO")
    db.add(Order(order_number=number, order_date=date(2026, 7, 1)))
    db.commit()
    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Master")
    tracking = wb.create_sheet("Overall Tracking")
    tracking.append(["Doc Number", "Date"])
    tracking.append([number, date(2026, 7, 5)])
    buf = BytesIO()
    wb.save(buf)

    OrderService(db).import_excel_tracking(buf.getvalue(), USER)

    order = _order(db, number)
    db.refresh(order)
    assert order.actual_delivery_date == date(2026, 7, 5)
