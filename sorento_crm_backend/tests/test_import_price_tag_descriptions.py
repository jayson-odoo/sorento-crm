"""PRICETAG-DESC-IMPORT: `scripts/import_price_tag_descriptions.py` (red tests, written first).

UAC: documentation/plans/products/pricetag-desc-import-15oct-acceptance-criteria.md
Plan: documentation/plans/products/PLAN-pricetag-desc-import-15oct.md (section 4 contract).

Everything runs through `script.main(argv, db=db, out=StringIO)` on a blank Postgres schema;
every test seeds its own products and specs. The CSVs are read from the paths the script
prints.
"""
from __future__ import annotations

import contextlib
import csv
import io
import os
import re
from pathlib import Path

import openpyxl
import pytest

from app.models.audit import AuditLog
from app.models.base import company_scope, set_company_scope
from app.models.company import Company
from app.models.product import Product
from app.models.product_spec import ProductSpecifications
from app.services.company_scope import DEFAULT_COMPANY_ID

from tests._mc_lookup_seed import MOCHA_ID, product, seed_mocha
from tests._pg_fixture import blank_session

CHANGES_COLUMNS = [
    "product_code", "sheet_code", "sheet_rows", "rule", "status",
    "active", "old_description", "new_description",
]
SKIPPED_COLUMNS = ["sheet_code", "sheet_rows", "product_code", "reason", "sheet_texts"]

SH_TEXT = (
    "Twister Flushing Technology\nD: L680xW365xH735mm\nS-Trap: Bottom Outlet\n"
    "150,200,250mm\nP-Trap: Horizontal\nOutlet 180mm\n*PP Seat Cover"
)
SH_P_EXPECTED = (
    "Twister Flushing Technology\nD: L680xW365xH735mm\nP-Trap: Horizontal\n"
    "Outlet 180mm\n*PP Seat Cover"
)
SH_S_EXPECTED = (
    "Twister Flushing Technology\nD: L680xW365xH735mm\nS-Trap: Bottom Outlet\n"
    "150,200,250mm\n*PP Seat Cover"
)


def _script():
    from scripts import import_price_tag_descriptions as script

    return script


@pytest.fixture()
def db():
    with blank_session() as session:
        set_company_scope(session, frozenset({DEFAULT_COMPANY_ID}))
        yield session


# --------------------------------------------------------------------------- helpers


def _prod(db, code, *, text=None, trap=None, finish=None, seat=None, company_id=DEFAULT_COMPANY_ID,
          active=True):
    row = product(db, company_id=company_id, code=code)
    row.price_tag_description = text
    row.is_active = active
    values = {}
    if trap:
        values["trap_type"] = {"value": trap}
    if finish:
        values["finish"] = {"value": finish}
    if seat:
        values["seat_material"] = {"value": seat}
    if values:
        db.add(ProductSpecifications(product_id=row.id, values=values))
    db.flush()
    db.commit()
    return row


def _xlsx(tmp_path, rows, *, headers=("No.", "Product Code", "Description", "Flyer Page(s)"),
          name="flyer.xlsx", extra_sheet=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(list(headers))
    for i, (code, desc) in enumerate(rows, start=1):
        ws.append([i, code, desc, "1"])
        for cell in ws[ws.max_row]:  # openpyxl would store "=..." as a formula; keep it text
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
    if extra_sheet:
        ws2 = wb.create_sheet("Other")
        ws2.append(list(headers))
        for i, (code, desc) in enumerate(extra_sheet, start=1):
            ws2.append([i, code, desc, "1"])
    path = tmp_path / name
    wb.save(path)
    return path


def _read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        return reader.fieldnames, list(reader)


class Run:
    def __init__(self, code, out, xlsx):
        self.code = code
        self.out = out
        self.xlsx = xlsx
        self.lines = [l.strip() for l in out.splitlines()]

    def path(self, label):
        for l in self.lines:
            if l.startswith(f"{label} report:"):
                return Path(l.split(":", 1)[1].strip())
        raise AssertionError(f"no '{label} report:' line in output:\n{self.out}")

    @property
    def changes(self):
        return _read_csv(self.path("changes"))[1]

    @property
    def skipped(self):
        return _read_csv(self.path("skipped"))[1]

    def change(self, product_code):
        found = [r for r in self.changes if r["product_code"] == product_code]
        assert len(found) == 1, (product_code, self.changes)
        return found[0]

    def skip_reasons(self, product_code=None, sheet_code=None):
        return {
            r["reason"]
            for r in self.skipped
            if (product_code is None or r["product_code"] == product_code)
            and (sheet_code is None or r["sheet_code"].upper() == sheet_code)
        }

    def count(self, label):
        for l in self.lines:
            m = re.fullmatch(rf"{re.escape(label)}: (\d+)", l)
            if m:
                return int(m.group(1))
        raise AssertionError(f"no '{label}: <n>' line in output:\n{self.out}")


def _run(db, tmp_path, rows, *extra, apply=False, **xlsx_kwargs):
    xlsx = _xlsx(tmp_path, rows, **xlsx_kwargs)
    argv = ["--file", str(xlsx), *extra] + (["--apply"] if apply else [])
    out = io.StringIO()
    code = _script().main(argv, db=db, out=out)
    return Run(code, out.getvalue(), xlsx)


def _desc(db, row):
    db.expire_all()
    with company_scope(db, None):
        return db.query(Product.price_tag_description).filter(Product.id == row.id).scalar()


def _audit_updates(db, row=None):
    q = db.query(AuditLog).filter(
        AuditLog.entity_type == "product",
        AuditLog.action == "UPDATE",
        AuditLog.job_id == "import_price_tag_descriptions",
    )
    if row is not None:
        q = q.filter(AuditLog.entity_id == str(row.id))
    return [a for a in q.all() if "price_tag_description" in (a.new_values or {})]


# --------------------------------------------------------------------------- AC-1 / AC-2 / AC-3 / AC-4 parse


def test_headers_found_by_name_case_insensitive_and_any_column_order(db, tmp_path):
    p = _prod(db, "SRTH1")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Flyer", "DESCRIPTION", "product code"])
    ws.append(["1", "Hello", "SRTH1"])
    path = tmp_path / "odd.xlsx"
    wb.save(path)
    out = io.StringIO()
    assert _script().main(["--file", str(path), "--apply"], db=db, out=out) == 0
    assert _desc(db, p) == "Hello"


def test_missing_header_exits_1_and_writes_nothing(db, tmp_path):
    p = _prod(db, "SRTH2")
    run = _run(db, tmp_path, [("SRTH2", "Hello")], apply=True, headers=("No.", "Code", "Text"))
    assert run.code == 1
    assert _desc(db, p) is None
    assert list(tmp_path.glob("*.csv")) == []


def test_unreadable_file_exits_1(db, tmp_path):
    out = io.StringIO()
    code = _script().main(["--file", str(tmp_path / "nope.xlsx")], db=db, out=out)
    assert code == 1


def test_only_the_first_worksheet_is_read(db, tmp_path):
    a = _prod(db, "SRTS1")
    b = _prod(db, "SRTS2")
    run = _run(db, tmp_path, [("SRTS1", "first")], apply=True, extra_sheet=[("SRTS2", "second")])
    assert run.code == 0
    assert _desc(db, a) == "first"
    assert _desc(db, b) is None


def test_row_numbers_are_excel_rows_header_is_row_1(db, tmp_path):
    _prod(db, "SRTR1")
    run = _run(db, tmp_path, [("SRTR1", "one"), (None, "orphan text"), ("SRTR1", "one")])
    assert run.change("SRTR1")["sheet_rows"].replace(" ", "") in ("2,4", "2;4", "2|4") or set(
        re.findall(r"\d+", run.change("SRTR1")["sheet_rows"])
    ) == {"2", "4"}
    no_code = [r for r in run.skipped if r["reason"] == "no code"]
    assert [re.findall(r"\d+", r["sheet_rows"]) for r in no_code] == [["3"]]


def test_cell_text_is_cleaned_crlf_trailing_spaces_and_outer_blank(db, tmp_path):
    p = _prod(db, "SRTC1")
    _run(db, tmp_path, [("SRTC1", "  \r\nLine one  \r\nLine two \t\nLine three   \n  ")], apply=True)
    assert _desc(db, p) == "Line one\nLine two\nLine three"


def test_rows_with_no_code_or_no_description_are_skipped_and_listed(db, tmp_path):
    p = _prod(db, "SRTB1", text="keep")
    run = _run(db, tmp_path, [(None, "text with no code"), ("SRTB1", None), ("SRTB1", "   ")],
               apply=True)
    reasons = {r["reason"] for r in run.skipped}
    assert {"no code", "no description"} <= reasons
    assert _desc(db, p) == "keep"
    assert run.changes == []


@pytest.mark.parametrize("fragment", ["-BL\nPillar Mounted\nKitchen Mixer Tap", "+SRTMR1\nMirror"])
def test_text_starting_with_part_of_a_code_is_skipped(db, tmp_path, fragment):
    p = _prod(db, "SRTKT71SS")
    run = _run(db, tmp_path, [("SRTKT71SS", fragment)], apply=True)
    assert "text starts with part of a code" in run.skip_reasons()
    assert _desc(db, p) is None


# --------------------------------------------------------------------------- AC-5 / AC-7 exact


def test_exact_match_after_code_normalisation(db, tmp_path):
    p = _prod(db, "SRTWC1-SH")
    run = _run(db, tmp_path, [("  srtwc1 - sh ", "Flushing")], apply=True)
    assert _desc(db, p) == "Flushing"
    row = run.change("SRTWC1-SH")
    assert row["rule"] == "exact" and row["status"] == "update"
    assert row["new_description"] == "Flushing" and row["old_description"] == ""


def test_code_not_found_is_skipped_and_listed(db, tmp_path):
    run = _run(db, tmp_path, [("SRTGHOST", "text")])
    assert "code not found in the CRM" in run.skip_reasons(sheet_code="SRTGHOST")
    assert run.changes == []


# --------------------------------------------------------------------------- AC-6 combo


def test_combo_puts_the_whole_text_on_the_first_code_only(db, tmp_path):
    text = ("SRTBF11831 (CABINET)\nL800xW500xH300mm\nSRTWB831 (BASIN)\nL800xW500xH180mm\n"
            "SRTMR502-BL (MIRROR)\nL800xH600mm")
    bf = _prod(db, "SRTBF11831")
    mr = _prod(db, "SRTMR502-BL")
    run = _run(db, tmp_path, [("SRTBF11831 + SRTMR502-BL", text)], apply=True)
    assert _desc(db, bf) == text
    assert _desc(db, mr) is None
    assert run.change("SRTBF11831")["rule"] == "combo"
    assert [r for r in run.changes if r["product_code"] == "SRTMR502-BL"] == []


def test_combo_of_three_codes_is_skipped(db, tmp_path):
    a = _prod(db, "SRTCA1")
    run = _run(db, tmp_path, [("SRTCA1 + SRTCB1 + SRTCC1", "text")], apply=True)
    assert "combo is not two codes" in run.skip_reasons()
    assert _desc(db, a) is None


def test_combo_whose_first_code_is_unknown_is_skipped(db, tmp_path):
    mr = _prod(db, "SRTCM1")
    run = _run(db, tmp_path, [("SRTNOBF + SRTCM1", "text")], apply=True)
    assert "code not found in the CRM" in run.skip_reasons()
    assert _desc(db, mr) is None


# --------------------------------------------------------------------------- AC-8 variants


def test_variant_with_same_finish_and_seat_gets_the_base_text(db, tmp_path):
    base = _prod(db, "SRTVA100", finish="black", seat="pp")
    var = _prod(db, "SRTVA100-X", finish="black", seat="pp")
    run = _run(db, tmp_path, [("SRTVA100", "Base text")], apply=True)
    assert _desc(db, base) == "Base text"
    assert _desc(db, var) == "Base text"
    row = run.change("SRTVA100-X")
    assert row["rule"] == "variant" and row["sheet_code"].upper() == "SRTVA100"


def test_variant_with_no_specs_at_all_counts_as_equal(db, tmp_path):
    _prod(db, "SRTVN1")
    var = _prod(db, "SRTVN1-X")
    _run(db, tmp_path, [("SRTVN1", "T")], apply=True)
    assert _desc(db, var) == "T"


@pytest.mark.parametrize(
    "var_kwargs, reason",
    [
        ({"finish": "white", "seat": "pp"}, "variant has a different finish"),
        ({"finish": "black", "seat": "ud"}, "variant has a different seat material"),
        ({"finish": "white", "seat": "ud"}, "variant has a different finish and seat material"),
        ({"seat": "pp"}, "variant has a different finish"),  # empty vs set is different
    ],
)
def test_variant_with_other_finish_or_seat_is_skipped(db, tmp_path, var_kwargs, reason):
    base = _prod(db, "SRTVF1", finish="black", seat="pp")
    var = _prod(db, "SRTVF1-X", **var_kwargs)
    run = _run(db, tmp_path, [("SRTVF1", "Base")], apply=True)
    assert _desc(db, base) == "Base"
    assert _desc(db, var) is None
    assert reason in run.skip_reasons(product_code="SRTVF1-X")


def test_variants_of_a_sheet_code_that_is_not_a_product_are_skipped(db, tmp_path):
    var = _prod(db, "SRTVG1-A")
    run = _run(db, tmp_path, [("SRTVG1", "T")], apply=True)
    assert _desc(db, var) is None
    assert "variant of a sheet code that is not a product" in run.skip_reasons(product_code="SRTVG1-A")
    assert "code not found in the CRM" in run.skip_reasons(sheet_code="SRTVG1")


def test_a_printed_code_is_never_a_variant(db, tmp_path):
    _prod(db, "SRTVB100")
    printed = _prod(db, "SRTVB100-X")
    run = _run(db, tmp_path, [("SRTVB100", "Base"), ("SRTVB100-X", "Own text")], apply=True)
    assert _desc(db, printed) == "Own text"
    assert run.change("SRTVB100-X")["rule"] == "exact"


def test_a_code_printed_on_a_combo_side_is_never_a_variant(db, tmp_path):
    bf = _prod(db, "SRTVD1")
    mr = _prod(db, "SRTVD1-M")
    run = _run(db, tmp_path, [("SRTVD1 + SRTVD1-M", "Cabinet and mirror")], apply=True)
    assert _desc(db, bf) == "Cabinet and mirror"
    assert _desc(db, mr) is None
    assert [r for r in run.changes if r["product_code"] == "SRTVD1-M"] == []


def test_the_longest_printed_base_wins(db, tmp_path):
    _prod(db, "SRTVC1")
    _prod(db, "SRTVC1-BL")
    deep = _prod(db, "SRTVC1-BL-X")
    run = _run(db, tmp_path, [("SRTVC1", "Short"), ("SRTVC1-BL", "Long")], apply=True)
    assert _desc(db, deep) == "Long"
    assert run.change("SRTVC1-BL-X")["sheet_code"].upper() == "SRTVC1-BL"


def test_variant_of_a_skipped_sheet_code_is_skipped(db, tmp_path):
    base = _prod(db, "SRTVE1")
    var = _prod(db, "SRTVE1-X")
    run = _run(db, tmp_path, [("SRTVE1", "Alpha text"), ("SRTVE1", "Beta words")], apply=True)
    assert _desc(db, base) is None and _desc(db, var) is None
    assert "variant of SRTVE1, which was skipped" in run.skip_reasons(product_code="SRTVE1-X")


# --------------------------------------------------------------------------- AC-9 company


def _prod_like_scope(db, monkeypatch):
    """Production conditions: a script never installs the app's company-scope listener, so the
    session is unscoped and the script's own `company_scope` does nothing. Only the explicit
    `Product.company_id == ...` filter can keep the companies apart."""
    set_company_scope(db, None)
    monkeypatch.setattr(_script(), "company_scope", lambda *a, **k: contextlib.nullcontext())


def test_other_companys_product_with_the_same_code_is_untouched(db, tmp_path, monkeypatch):
    seed_mocha(db)
    other = _prod(db, "SRTDUP1", company_id=MOCHA_ID)  # seeded first: it wins if nothing filters
    srt = _prod(db, "SRTDUP1")
    _prod_like_scope(db, monkeypatch)
    run = _run(db, tmp_path, [("SRTDUP1", "SRT text")], apply=True)
    assert run.code == 0
    assert _desc(db, srt) == "SRT text"
    assert _desc(db, other) is None
    assert [r["product_code"] for r in run.changes] == ["SRTDUP1"]


def test_company_flag_targets_that_companys_products(db, tmp_path, monkeypatch):
    mocha = seed_mocha(db)
    srt = _prod(db, "SRTDUP2")
    other = _prod(db, "SRTDUP2", company_id=MOCHA_ID)
    _prod_like_scope(db, monkeypatch)
    run = _run(db, tmp_path, [("SRTDUP2", "Mocha text")], "--company", mocha.code, apply=True)
    assert run.code == 0
    assert _desc(db, other) == "Mocha text"
    assert _desc(db, srt) is None
    assert [r["product_code"] for r in run.changes] == ["SRTDUP2"]


def test_unknown_company_exits_1_and_writes_nothing(db, tmp_path):
    p = _prod(db, "SRTUC1")
    run = _run(db, tmp_path, [("SRTUC1", "T")], "--company", "NOPE", apply=True)
    assert run.code == 1
    assert _desc(db, p) is None


# --------------------------------------------------------------------------- AC-10..12 duplicates


def test_format_equal_rows_are_one_and_the_first_rows_text_is_written(db, tmp_path):
    p = _prod(db, "SRTWD1")
    first = "Washdown With Rimless\nFlushing Technology\nD: L700xW400xH760mm"
    second = "Washdown With Rimless\nFlushing Technology\nD : L700xW400xH760mm"
    run = _run(db, tmp_path, [("SRTWD1", first), ("SRTWD1", second)], apply=True)
    assert _desc(db, p) == first
    assert run.skipped == []


def test_contained_shorter_text_loses_to_the_longest(db, tmp_path):
    p = _prod(db, "SRTMRL707")
    full = "Frameless LED Mirror\nL1000 x H700mm\nAnti-fog,\nLED Light"
    run = _run(db, tmp_path, [("SRTMRL707", "Frameless LED Mirror"), ("SRTMRL707", full)], apply=True)
    assert _desc(db, p) == full
    assert run.skipped == []


def test_conflicting_texts_skip_the_code_and_list_every_row_and_text(db, tmp_path):
    t1 = "SRTBF11415 (CABINET)\nL400xW345xH390mm"
    t2 = "SRTBF11415 (CABINET)\nD: L370xW360xH400mm"
    p = _prod(db, "SRTBF11415", text="old")
    run = _run(db, tmp_path, [("SRTBF11415", t1), ("SRTBF11415", t2)], apply=True)
    assert _desc(db, p) == "old"
    rows = [r for r in run.skipped if r["reason"] == "same code on several rows with different text"]
    assert len(rows) == 1
    r = rows[0]
    assert r["sheet_code"].upper() == "SRTBF11415"
    assert set(re.findall(r"\d+", r["sheet_rows"])) == {"2", "3"}
    assert "row 2:" in r["sheet_texts"] and "row 3:" in r["sheet_texts"]
    assert t1 in r["sheet_texts"] and t2 in r["sheet_texts"]
    assert "\n\n" in r["sheet_texts"]


def test_a_combo_row_and_a_plain_row_of_the_same_bf_count_as_the_same_code(db, tmp_path):
    bf = _prod(db, "SRTBF200")
    _prod(db, "SRTMR200")
    run = _run(db, tmp_path, [("SRTBF200 + SRTMR200", "Cabinet one"), ("SRTBF200", "Cabinet two")],
               apply=True)
    assert _desc(db, bf) is None
    assert "same code on several rows with different text" in run.skip_reasons()


def test_skipped_csv_is_sorted_by_first_sheet_row(db, tmp_path):
    run = _run(db, tmp_path, [("SRTZ2", "a"), ("SRTA1", "b"), (None, "c")])
    firsts = [int(re.findall(r"\d+", r["sheet_rows"])[0]) for r in run.skipped]
    assert firsts == sorted(firsts) and len(firsts) == 3


# --------------------------------------------------------------------------- AC-13..16 traps


def test_s_trap_product_keeps_the_s_block_and_p_trap_product_keeps_the_p_block(db, tmp_path):
    sh = _prod(db, "SRTWC286-SH", trap="s_trap")
    shp = _prod(db, "SRTWC286-SH-P", trap="p_trap")
    run = _run(db, tmp_path, [("SRTWC286-SH", SH_TEXT)], apply=True)
    assert _desc(db, sh) == SH_S_EXPECTED
    assert _desc(db, shp) == SH_P_EXPECTED
    assert run.change("SRTWC286-SH")["rule"] == "exact+trap"
    assert run.change("SRTWC286-SH-P")["rule"] == "variant+trap"


def test_trap_line_match_is_case_insensitive_with_optional_space_or_dash(db, tmp_path):
    text = "Head\ns trap: Bottom\n150mm\nP-Trap: Horizontal\n180mm\n*Seat"
    s = _prod(db, "SRTTL1", trap="s_trap")
    _prod(db, "SRTTL1-P", trap="p_trap")
    _run(db, tmp_path, [("SRTTL1", text)], apply=True)
    assert _desc(db, s) == "Head\ns trap: Bottom\n150mm\n*Seat"


def test_a_block_ends_at_the_next_trap_line_so_trailing_lines_after_p_are_kept(db, tmp_path):
    text = "Head\nP-Trap: Horizontal\n180mm\nS-Trap: Bottom\n200mm"
    p = _prod(db, "SRTTL2", trap="p_trap")
    _run(db, tmp_path, [("SRTTL2", text)], apply=True)
    assert _desc(db, p) == "Head\nP-Trap: Horizontal\n180mm"


def test_both_traps_printed_and_no_trap_type_is_skipped(db, tmp_path):
    p = _prod(db, "SRTWC286-NT")
    run = _run(db, tmp_path, [("SRTWC286-NT", SH_TEXT)], apply=True)
    assert _desc(db, p) is None
    assert "both traps printed, product has no trap type" in run.skip_reasons(product_code="SRTWC286-NT")


def test_one_trap_printed_and_product_has_the_other_trap_is_skipped(db, tmp_path):
    text = "Washdown Flushing Technology\nS-Trap: Bottom Outlet 200mm\n*PP Seat Cover"
    base = _prod(db, "SRTWC201", trap="s_trap")
    var = _prod(db, "SRTWC201-P", trap="p_trap")
    run = _run(db, tmp_path, [("SRTWC201", text)], apply=True)
    assert _desc(db, base) == text
    assert run.change("SRTWC201")["rule"] == "exact"  # nothing was cut, no +trap
    assert _desc(db, var) is None
    assert "product's trap type is not printed" in run.skip_reasons(product_code="SRTWC201-P")


def test_one_trap_printed_and_no_trap_type_takes_the_text_unchanged(db, tmp_path):
    text = "Washdown\nS-Trap: Bottom Outlet 200mm\n*PP Seat Cover"
    p = _prod(db, "SRTWC202")
    run = _run(db, tmp_path, [("SRTWC202", text)], apply=True)
    assert _desc(db, p) == text
    assert run.change("SRTWC202")["rule"] == "exact"


# --------------------------------------------------------------------------- AC-17..22 dry run, apply


def test_dry_run_writes_nothing_and_no_audit_row_but_reports(db, tmp_path):
    p = _prod(db, "SRTDR1", text="old text")
    run = _run(db, tmp_path, [("SRTDR1", "new\ntext")])
    assert run.code == 0
    assert _desc(db, p) == "old text"
    assert _audit_updates(db) == []
    assert any(l.startswith("DRY RUN: 1 products to update. Run with --apply to write them.")
               for l in run.lines)
    assert run.count("updates") == 1


def test_update_line_shows_rule_code_old_and_new_with_line_breaks_as_slashes(db, tmp_path):
    _prod(db, "SRTUL1", text="old\nlines")
    run = _run(db, tmp_path, [("SRTUL1", "new\ntext")])
    line = next(l for l in run.lines if "SRTUL1" in l and "exact" in l)
    assert "old / lines" in line and "new / text" in line
    assert "\n" not in line


def test_summary_counts_updates_unchanged_and_skip_reasons(db, tmp_path):
    _prod(db, "SRTSM1")
    _prod(db, "SRTSM2", text="same")
    run = _run(db, tmp_path, [("SRTSM1", "x"), ("SRTSM2", "same"), (None, "y"), ("SRTGONE", "z")])
    assert run.count("updates") == 1
    assert run.count("unchanged") == 1
    assert run.count("skipped") == 2
    assert run.count("no code") == 1
    assert run.count("code not found in the CRM") == 1


def test_apply_overwrites_existing_text_and_reports_old_and_new(db, tmp_path):
    p = _prod(db, "SRTOW1", text="old words")
    run = _run(db, tmp_path, [("SRTOW1", "new words")], apply=True)
    assert _desc(db, p) == "new words"
    row = run.change("SRTOW1")
    assert (row["old_description"], row["new_description"], row["status"]) == (
        "old words", "new words", "update")
    assert any(l == "1 products updated." for l in run.lines)


def test_equal_current_text_is_unchanged_never_updated(db, tmp_path):
    p = _prod(db, "SRTUN1", text="Same\ntext")
    run = _run(db, tmp_path, [("SRTUN1", "Same\ntext")], apply=True)
    assert run.change("SRTUN1")["status"] == "unchanged"
    assert run.count("updates") == 0
    assert _audit_updates(db, p) == []


def test_apply_commits_once(db, tmp_path, monkeypatch):
    _prod(db, "SRTTX1")
    _prod(db, "SRTTX2")
    calls = []
    real = db.commit
    monkeypatch.setattr(db, "commit", lambda: (calls.append(1), real())[1])
    _run(db, tmp_path, [("SRTTX1", "a"), ("SRTTX2", "b")], apply=True)
    assert len(calls) == 1


def test_a_failing_commit_leaves_every_product_unchanged(db, tmp_path, monkeypatch):
    a = _prod(db, "SRTRB1", text="a-old")
    b = _prod(db, "SRTRB2", text="b-old")

    def boom():
        raise RuntimeError("commit failed")

    monkeypatch.setattr(db, "commit", boom)
    xlsx = _xlsx(tmp_path, [("SRTRB1", "a-new"), ("SRTRB2", "b-new")])
    try:
        code = _script().main(["--file", str(xlsx), "--apply"], db=db, out=io.StringIO())
        assert code != 0
    except RuntimeError:
        pass
    monkeypatch.undo()
    db.rollback()
    assert _desc(db, a) == "a-old" and _desc(db, b) == "b-old"


def test_one_audit_update_row_per_changed_product_with_old_and_new(db, tmp_path):
    a = _prod(db, "SRTAU1", text="old a")
    b = _prod(db, "SRTAU2")
    c = _prod(db, "SRTAU3", text="same")
    _run(db, tmp_path, [("SRTAU1", "new a"), ("SRTAU2", "new b"), ("SRTAU3", "same")], apply=True)
    ra = _audit_updates(db, a)
    rb = _audit_updates(db, b)
    assert len(ra) == 1 and len(rb) == 1
    assert ra[0].entity_type == "product" and ra[0].action == "UPDATE"
    assert ra[0].old_values["price_tag_description"] == "old a"
    assert ra[0].new_values["price_tag_description"] == "new a"
    assert rb[0].old_values.get("price_tag_description") in (None, "")
    assert rb[0].new_values["price_tag_description"] == "new b"
    assert _audit_updates(db, c) == []


def test_second_apply_updates_nothing_and_writes_no_audit_rows(db, tmp_path):
    p = _prod(db, "SRTID1")
    rows = [("SRTID1", "text")]
    _run(db, tmp_path, rows, apply=True)
    before = len(_audit_updates(db))
    assert before == 1
    again = _run(db, tmp_path, rows, apply=True)
    assert again.count("updates") == 0
    assert any(l == "0 products updated." for l in again.lines)
    assert len(_audit_updates(db)) == before
    assert _desc(db, p) == "text"


def test_dry_run_after_apply_reports_zero_to_update(db, tmp_path):
    _prod(db, "SRTID2")
    rows = [("SRTID2", "text")]
    _run(db, tmp_path, rows, apply=True)
    dry = _run(db, tmp_path, rows)
    assert any(l.startswith("DRY RUN: 0 products to update.") for l in dry.lines)


def test_only_price_tag_description_is_written(db, tmp_path):
    p = _prod(db, "SRTOO1", finish="black")
    name_before, price_before = p.product_name, p.list_price
    _run(db, tmp_path, [("SRTOO1", "text")], apply=True)
    db.expire_all()
    with company_scope(db, None):
        row = db.query(Product).filter(Product.id == p.id).one()
    assert (row.product_name, row.list_price) == (name_before, price_before)
    spec = db.query(ProductSpecifications).filter_by(product_id=p.id).one()
    assert spec.values == {"finish": {"value": "black"}}


# --------------------------------------------------------------------------- AC-23..25 reports


@pytest.mark.parametrize("apply, mode", [(False, "dryrun"), (True, "apply")])
def test_csvs_land_next_to_the_xlsx_with_the_contract_names(db, tmp_path, apply, mode):
    _prod(db, "SRTCS1")
    run = _run(db, tmp_path, [("SRTCS1", "x"), (None, "y")], apply=apply, name="My Sheet.xlsx")
    changes, skipped = run.path("changes"), run.path("skipped")
    assert changes.parent == tmp_path and skipped.parent == tmp_path
    assert re.fullmatch(rf"My Sheet-price-tag-changes-{mode}-\d{{8}}-\d{{6}}\.csv", changes.name)
    assert re.fullmatch(rf"My Sheet-price-tag-skipped-{mode}-\d{{8}}-\d{{6}}\.csv", skipped.name)
    assert changes.exists() and skipped.exists()


def test_csv_columns_are_in_contract_order_and_multiline_text_is_whole(db, tmp_path):
    _prod(db, "SRTCS2", text="old\nmulti", active=False)
    run = _run(db, tmp_path, [("SRTCS2", "new\nmulti\nline"), (None, "y")])
    assert _read_csv(run.path("changes"))[0] == CHANGES_COLUMNS
    assert _read_csv(run.path("skipped"))[0] == SKIPPED_COLUMNS
    row = run.change("SRTCS2")
    assert row["old_description"] == "old\nmulti" and row["new_description"] == "new\nmulti\nline"
    assert row["active"].strip().lower() in ("false", "0", "no", "n", "inactive")
    assert row["sheet_code"].upper() == "SRTCS2" and "2" in row["sheet_rows"]


def test_changes_csv_lists_unchanged_rows_too_and_empty_reports_still_have_headers(db, tmp_path):
    _prod(db, "SRTCS3", text="same")
    run = _run(db, tmp_path, [("SRTCS3", "same")])
    assert run.change("SRTCS3")["status"] == "unchanged"
    assert _read_csv(run.path("skipped")) == (SKIPPED_COLUMNS, [])


# --------------------------------------------------------------------------- fix round 1: AC-10 format-equal pairs


@pytest.mark.parametrize(
    "first, second",
    [
        ("Washdown\nFlushing Technology", "Washdown Flushing Technology"),
        ("Washdown Flushing", "WASHDOWN flushing"),
        ("Head\nD: L700xW400xH760mm", "Head\nL700xW400xH760mm"),
        ("Head\nSize :  L700", "Head\nSize: L700"),
        ("STANLEY 30oz TUMBLER", ":\nSTANLEY 30oz TUMBLER"),
        ("STANLEY 30oz TUMBLER", "  STANLEY 30oz TUMBLER :"),
    ],
)
def test_format_equal_pairs_are_one_row_and_the_first_text_is_written(db, tmp_path, first, second):
    p = _prod(db, "SRTFE1")
    run = _run(db, tmp_path, [("SRTFE1", first), ("SRTFE1", second)], apply=True)
    assert _desc(db, p) == first.strip()
    assert run.skipped == []


# --------------------------------------------------------------------------- fix round 1: AC-8 printed anywhere / skipped bases


def test_a_code_printed_only_on_a_no_description_row_is_never_a_variant(db, tmp_path):
    _prod(db, "SRTWC8318-RL")
    bl = _prod(db, "SRTWC8318-RL-BL")
    run = _run(db, tmp_path, [("SRTWC8318-RL", "Some text"), ("SRTWC8318-RL-BL", None)], apply=True)
    assert _desc(db, bl) is None
    assert [r for r in run.changes if r["product_code"] == "SRTWC8318-RL-BL"] == []


def test_variant_of_a_base_printed_only_on_fragment_rows_is_listed_as_base_skipped(db, tmp_path):
    base = _prod(db, "SRTBF11827")
    var = _prod(db, "SRTBF11827-BL")
    run = _run(db, tmp_path, [("SRTBF11827", "+SRTMCB8082-BL\nCabinet")], apply=True)
    assert _desc(db, base) is None and _desc(db, var) is None
    assert "variant of SRTBF11827, which was skipped" in run.skip_reasons(product_code="SRTBF11827-BL")


def test_product_whose_longest_base_is_only_a_combo_second_code_is_left_alone_and_unlisted(db, tmp_path):
    _prod(db, "SRTBF9")
    _prod(db, "SRTMCB9")
    new = _prod(db, "SRTMCB9-NEW")
    run = _run(db, tmp_path, [("SRTBF9 + SRTMCB9", "Cabinet and mirror")], apply=True)
    assert _desc(db, new) is None
    assert [r for r in run.changes if r["product_code"] == "SRTMCB9-NEW"] == []
    assert [r for r in run.skipped if r["product_code"] == "SRTMCB9-NEW"] == []


# --------------------------------------------------------------------------- fix round 1: AC-4 formula / template text


@pytest.mark.parametrize(
    "text, reason",
    [
        ("=SUM(1+1)", "text starts with = or @"),
        ("@mention\nmore", "text starts with = or @"),
        ("Head\nuses {{price_tag_description}} inside", "text contains {{"),
    ],
)
def test_formula_or_template_text_is_skipped_and_listed(db, tmp_path, text, reason):
    p = _prod(db, "SRTFX1", text="keep")
    run = _run(db, tmp_path, [("SRTFX1", text)], apply=True)
    assert reason in run.skip_reasons()
    assert _desc(db, p) == "keep"
    assert run.changes == []


# --------------------------------------------------------------------------- fix round 1: AC-25b CSV escaping


@pytest.mark.parametrize("old", ["-old", "+old", "\told", "\rold", "=old", "@old"])
def test_changes_csv_escapes_a_leading_formula_char_but_the_database_keeps_it(db, tmp_path, old):
    p = _prod(db, "SRTEC1", text=old)
    run = _run(db, tmp_path, [("SRTEC1", "new words")], apply=True)
    assert run.change("SRTEC1")["old_description"] == "'" + old
    assert _desc(db, p) == "new words"
    # the escape never reaches the database: the old text was raw, not "'"-prefixed
    audit = _audit_updates(db, p)
    assert audit[0].old_values["price_tag_description"] == old


def test_exactly_max_rows_rows_is_accepted(db, tmp_path, monkeypatch):
    rows = [(f"SRTCAPOK{i}", "t") for i in range(3)]
    ps = [_prod(db, c) for c, _ in rows]
    monkeypatch.setattr(_script(), "MAX_ROWS", 3)
    run = _run(db, tmp_path, rows, apply=True)
    assert run.code == 0
    assert all(_desc(db, p) == "t" for p in ps)


@pytest.mark.parametrize("code", ["=SRTX", "@SRTX"])
def test_skipped_csv_escapes_a_leading_formula_char_in_the_code_cell(db, tmp_path, code):
    run = _run(db, tmp_path, [(code, "some text")])
    assert [r["sheet_code"] for r in run.skipped] == ["'" + code]
    assert run.skipped[0]["reason"] == "code not found in the CRM"


# --------------------------------------------------------------------------- fix round 1: AC-23 names, report-before-commit


def test_a_file_name_with_a_percent_sign_works(db, tmp_path):
    _prod(db, "SRTPC1")
    run = _run(db, tmp_path, [("SRTPC1", "text")], name="Flyer 50% off.xlsx")
    assert run.code == 0
    assert run.path("changes").name.startswith("Flyer 50% off-price-tag-changes-dryrun-")
    assert run.path("skipped").name.startswith("Flyer 50% off-price-tag-skipped-dryrun-")
    assert run.path("changes").exists() and run.path("skipped").exists()


@pytest.mark.skipif(os.geteuid() == 0, reason="root ignores directory permissions")
def test_a_report_write_failure_on_apply_exits_1_and_changes_no_product(db, tmp_path):
    p = _prod(db, "SRTRW1", text="old")
    xlsx = _xlsx(tmp_path, [("SRTRW1", "new")])
    tmp_path.chmod(0o500)
    try:
        out = io.StringIO()
        code = _script().main(["--file", str(xlsx), "--apply"], db=db, out=out)
    finally:
        tmp_path.chmod(0o700)
    assert code == 1, out.getvalue()
    assert _desc(db, p) == "old"


# --------------------------------------------------------------------------- fix round 1: AC-1 caps


def test_more_rows_than_the_cap_exits_1_and_writes_nothing(db, tmp_path, monkeypatch):
    rows = [(f"SRTCAP{i}", "t") for i in range(4)]
    ps = [_prod(db, c) for c, _ in rows]
    monkeypatch.setattr(_script(), "MAX_ROWS", 3)
    run = _run(db, tmp_path, rows, apply=True)
    assert run.code == 1
    assert all(_desc(db, p) is None for p in ps)
    assert list(tmp_path.glob("*.csv")) == []


def test_a_file_over_the_byte_cap_exits_1_and_writes_nothing(db, tmp_path, monkeypatch):
    p = _prod(db, "SRTCAP9")
    monkeypatch.setattr(_script(), "MAX_FILE_BYTES", 10)
    run = _run(db, tmp_path, [("SRTCAP9", "t")], apply=True)
    assert run.code == 1
    assert _desc(db, p) is None
    assert list(tmp_path.glob("*.csv")) == []
