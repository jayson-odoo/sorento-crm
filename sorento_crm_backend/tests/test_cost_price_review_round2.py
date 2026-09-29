"""Fix lane round 2 (#1305): one test per finding of the reviewer pass at 232e5706.

Every test here was run RED against 232e5706 before its fix landed; the test names carry
the reviewer's finding id (b = Blocking, s = Should fix, n = Nit).
"""
from __future__ import annotations

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import openpyxl
import pytest

from tests.fixtures.cost_price.taiyang_shapes import (
    HEADER_ROW,
    LETTERHEAD_TEXT,
    simple_price_list_workbook,
)
from tests.support.cost_price_env import (
    PS_EDIT_PERM,
    PS_VIEW_PERM,
    UPLOAD_PERM,
    VERIFY_PERM,
    VIEW_PERM,
    cost_price_env,
)
from tests._pg_fixture import blank_session, unique_code

# 16:05 UTC on 30 Sep is 00:05 Malaysia time on 1 Oct: the moment the daily tick fires.
# Far in the future so the real clock can never make an unfixed build pass by accident.
_TICK_UTC = datetime(2031, 9, 30, 16, 5, tzinfo=timezone.utc)
_MY_TODAY = date(2031, 10, 1)


def _freeze_utc_host(monkeypatch, instant: datetime = _TICK_UTC) -> None:
    """A UTC host at `instant`: `date.today()` answers the UTC day, while a clock asked
    for Malaysia time answers the Malaysia day. Both halves are patched so the test is
    red on a build that reads `date.today()` and green on one that asks for Malaysia."""
    import app.services.pdf_render as pdf_render
    import app.services.procurement.cost_price_change_service as change_service
    import app.services.procurement.supplier_cost_service as cost_service

    class _UtcHostDate(date):
        @classmethod
        def today(cls):
            return instant.date()

    class _FrozenDatetime(datetime):
        @classmethod
        def now(cls, tz=None):
            return instant.astimezone(tz) if tz else instant.replace(tzinfo=None)

    monkeypatch.setattr(pdf_render, "datetime", _FrozenDatetime)
    for module in (cost_service, change_service):
        if hasattr(module, "date"):
            monkeypatch.setattr(module, "date", _UtcHostDate)


def _cost_row(e, link, unit_cost, *, start=None, end=None, currency="CNY"):
    from app.models.cost_price import ProductSupplierCost

    row = ProductSupplierCost(
        product_supplier_id=link.id, unit_cost=unit_cost, currency=currency,
        start_date=start, end_date=end,
    )
    e.db.add(row)
    e.db.flush()
    return row


# ------------------------------------------------------------------------------ Blocking 1


def test_b1_daily_tick_at_0005_myt_on_a_utc_host_takes_todays_start(cost_price_env, monkeypatch):
    from app.models.procurement import ProductSupplier
    from app.services.scheduled_task_service import TASK_HANDLERS

    e = cost_price_env
    supplier = e.supplier()
    product = e.product()
    link = e.link(product, supplier, unit_cost=50, currency="CNY")
    _cost_row(e, link, 50)
    _cost_row(e, link, 60, start=_MY_TODAY)
    e.db.commit()

    _freeze_utc_host(monkeypatch)
    TASK_HANDLERS["cost_price_daily_tick"](e.db, None)

    e.db.expire_all()
    assert e.db.query(ProductSupplier).filter_by(id=link.id).one().unit_cost == Decimal("60.00")


def test_b1_hand_edit_and_prices_tab_use_the_malaysia_day(cost_price_env, monkeypatch):
    from app.models.procurement import ProductSupplier

    e = cost_price_env
    editor = e.user(PS_EDIT_PERM, PS_VIEW_PERM)
    e.as_user(editor)
    supplier = e.supplier()
    product = e.product()
    link = e.link(product, supplier, unit_cost=50, currency="CNY")
    _cost_row(e, link, 50)
    e.db.commit()

    _freeze_utc_host(monkeypatch)
    r = e.post_cost(link.id, {"unit_cost": 70, "currency": "CNY", "start_date": _MY_TODAY.isoformat()})
    assert r.status_code == 201, r.text
    assert r.json()["status"] == "in_force"

    e.db.expire_all()
    assert e.db.query(ProductSupplier).filter_by(id=link.id).one().unit_cost == Decimal("70.00")

    tab = e.cost_lists(str(supplier.id))
    assert tab.status_code == 200, tab.text
    assert tab.json()["today"] == _MY_TODAY.isoformat()


# ------------------------------------------------------------------------------ Blocking 2


def test_b2_future_dated_first_hand_row_keeps_the_directly_written_price(cost_price_env):
    from app.models.procurement import ProductSupplier

    e = cost_price_env
    editor = e.user(PS_EDIT_PERM, PS_VIEW_PERM)
    e.as_user(editor)
    supplier = e.supplier()
    product = e.product()
    link = e.link(product, supplier, unit_cost=100, currency="CNY")
    e.db.commit()

    r = e.post_cost(link.id, {"unit_cost": 120, "currency": "CNY", "start_date": "2099-01-01"})
    assert r.status_code == 201, r.text

    e.db.expire_all()
    got = e.db.query(ProductSupplier).filter_by(id=link.id).one()
    assert (got.unit_cost, got.currency) == (Decimal("100.00"), "CNY")


def test_b2_daily_tick_never_nulls_a_price_when_nothing_is_in_force(cost_price_env):
    from app.models.procurement import ProductSupplier
    from app.services.procurement.supplier_cost_service import refresh_prices_in_force

    e = cost_price_env
    supplier = e.supplier()
    future = e.link(e.product(), supplier, unit_cost=100, currency="CNY")
    _cost_row(e, future, 120, start=date(2099, 1, 1))
    ended = e.link(e.product(), e.supplier(), unit_cost=80, currency="USD")
    _cost_row(e, ended, 80, currency="USD", end=date(2020, 1, 1))
    e.db.commit()

    refresh_prices_in_force(e.db, date(2030, 1, 1))

    e.db.expire_all()
    for link, expected in ((future, (Decimal("100.00"), "CNY")), (ended, (Decimal("80.00"), "USD"))):
        got = e.db.query(ProductSupplier).filter_by(id=link.id).one()
        assert (got.unit_cost, got.currency) == expected


# ------------------------------------------------------------------------------ Blocking 3


def _migration(name: str):
    import importlib.util
    from pathlib import Path

    path = Path(__file__).resolve().parents[1] / "alembic" / "versions" / f"{name}.py"
    spec = importlib.util.spec_from_file_location(f"zzt_r2_{name}", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _run(db, module, direction: str) -> None:
    from alembic.migration import MigrationContext
    from alembic.operations import Operations

    context = MigrationContext.configure(connection=db.connection())
    with Operations.context(context):
        getattr(module, direction)()


def _slugs_of(db, role_id: str) -> set[str]:
    from sqlalchemy import text

    return {
        r[0] for r in db.execute(
            text(
                "SELECT p.slug FROM user_role_permissions rp "
                "JOIN user_permissions p ON p.id = rp.permission_id WHERE rp.role_id = :r"
            ),
            {"r": role_id},
        ).all()
    }


def test_b3_cpc1_downgrade_runs_and_removes_only_what_it_added():
    import uuid

    from sqlalchemy import text

    from app.models.user import UserPermission, UserRole, UserRolePermission

    def _perm(db, slug):
        row = db.query(UserPermission).filter_by(slug=slug).one_or_none()
        if row is None:
            row = UserPermission(id=str(uuid.uuid4()), slug=slug, name=slug, description="")
            db.add(row)
            db.flush()
        return row

    with blank_session() as db:
        # A custom role granted the pre-existing product-supplier slugs BEFORE cpc1.
        role = UserRole(
            id=str(uuid.uuid4()), slug=unique_code("cpcr2custom"), name="custom",
            description="", is_protected=False, is_default=False,
        )
        db.add(role)
        db.flush()
        for slug in ("procurement.product_suppliers.view", "procurement.product_suppliers.edit"):
            db.add(UserRolePermission(id=str(uuid.uuid4()), role_id=role.id, permission_id=_perm(db, slug).id))
        db.commit()

        cpc1 = _migration("cpc1_supplier_cost_lists")
        _run(db, cpc1, "upgrade")
        _run(db, cpc1, "downgrade")

        def _exists(table):
            # The scratch schema is `current_schema()`; `public` (later on the search_path)
            # holds the migrated CI database's own copy of every table.
            return db.execute(
                text("SELECT to_regclass(quote_ident(current_schema()) || '.' || :t)"), {"t": table}
            ).scalar() is not None

        for table in ("cost_price_change_lines", "product_supplier_costs", "cost_price_change_sets", "supplier_price_links"):
            assert not _exists(table), table
        remaining = {
            r[0] for r in db.execute(
                text("SELECT slug FROM user_permissions WHERE slug LIKE 'procurement.%'")
            ).all()
        }
        assert not remaining & {
            "procurement.cost_price_changes.upload", "procurement.cost_price_changes.view",
            "procurement.cost_price_changes.verify", "procurement.suppliers.price_link",
        }
        assert {
            "procurement.product_suppliers.view", "procurement.product_suppliers.add",
            "procurement.product_suppliers.edit", "procurement.product_suppliers.delete",
        } <= remaining
        # The grant that existed before cpc1 survives its downgrade.
        assert {"procurement.product_suppliers.view", "procurement.product_suppliers.edit"} <= _slugs_of(db, role.id)
        # Nit 8: the numbering rule cpc1 seeded goes with it.
        assert not db.execute(
            text("SELECT count(*) FROM document_numbering_rules WHERE doc_type = 'cost_price_change_set'")
        ).scalar()

        # Upgrade again: the slugs come back. (The tables are proven by the real
        # `alembic upgrade head / downgrade / upgrade head` round trip on a fresh database;
        # inside this scratch schema the migration's `has_table` sees `public`'s copy.)
        _run(db, cpc1, "upgrade")
        assert db.execute(
            text("SELECT count(*) FROM user_permissions WHERE slug = 'procurement.cost_price_changes.upload'")
        ).scalar() == 1


# --------------------------------------------------------------------------- Should fix 1


def test_s1_migration_sweeps_product_supplier_slugs_onto_product_roles():
    import uuid

    from app.models.user import UserPermission, UserRole, UserRolePermission

    def _perm(db, slug):
        row = db.query(UserPermission).filter_by(slug=slug).one_or_none()
        if row is None:
            row = UserPermission(id=str(uuid.uuid4()), slug=slug, name=slug, description="")
            db.add(row)
            db.flush()
        return row

    def _role(db, tag, *slugs):
        role = UserRole(
            id=str(uuid.uuid4()), slug=unique_code(tag), name=tag, description="",
            is_protected=False, is_default=False,
        )
        db.add(role)
        db.flush()
        for slug in slugs:
            db.add(UserRolePermission(id=str(uuid.uuid4()), role_id=role.id, permission_id=_perm(db, slug).id))
        return role

    with blank_session() as db:
        viewer = _role(db, "cpcr2pview", "master_data.products.view")
        editor = _role(db, "cpcr2pedit", "master_data.products.view", "master_data.products.edit")
        db.commit()

        cpc1 = _migration("cpc1_supplier_cost_lists")
        _run(db, cpc1, "upgrade")
        _run(db, cpc1, "upgrade")

        viewer_slugs = _slugs_of(db, viewer.id)
        assert "procurement.product_suppliers.view" in viewer_slugs
        assert not viewer_slugs & {
            "procurement.product_suppliers.add", "procurement.product_suppliers.edit",
            "procurement.product_suppliers.delete",
        }
        assert {
            "procurement.product_suppliers.view", "procurement.product_suppliers.add",
            "procurement.product_suppliers.edit", "procurement.product_suppliers.delete",
        } <= _slugs_of(db, editor.id)


# --------------------------------------------------------------------------- Should fix 2


@pytest.mark.parametrize(
    "body",
    [
        {"unit_cost": 12},
        {"unit_cost": "abc", "currency": "CNY"},
        {"currency": "CNY"},
        {"unit_cost": 12, "currency": "CNY", "start_date": "not-a-date"},
    ],
)
def test_s2_cost_row_body_is_validated_as_422(cost_price_env, body):
    e = cost_price_env
    e.as_user(e.user(PS_EDIT_PERM))
    link = e.link(e.product(), e.supplier())
    e.db.commit()

    r = e.post_cost(link.id, body)
    assert r.status_code == 422, r.text


def test_s2_put_cost_row_text_price_is_422(cost_price_env):
    e = cost_price_env
    e.as_user(e.user(PS_EDIT_PERM))
    link = e.link(e.product(), e.supplier())
    row = _cost_row(e, link, 10)
    e.db.commit()

    r = e.put_cost(link.id, row.id, {"unit_cost": "abc"})
    assert r.status_code == 422, r.text


# --------------------------------------------------------------------------- Should fix 3


def _wide_letterhead_workbook(rows: int, *, stray_column: int | None = None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([LETTERHEAD_TEXT, None, None, None, None, None, None, None, None, "Tel: 0592"])
    ws.append(list(HEADER_ROW))
    for i in range(1, rows + 1):
        ws.append([i, f"ZZCPC-W-{i:05d}", "cfg", 10])
    if stray_column:
        ws.cell(row=1, column=stray_column).number_format = "0.00"
        ws.cell(row=1, column=stray_column, value=" ")
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_s3_a_2100_row_sheet_with_a_ten_column_letterhead_is_read():
    from app.services.procurement.supplier_price_list_reader import read_supplier_price_list

    parsed = read_supplier_price_list(_wide_letterhead_workbook(2100), "wide.xlsx")
    assert parsed.total_rows == 2100


def test_s3_stray_formatting_far_to_the_right_does_not_refuse_the_file():
    from app.services.procurement.supplier_price_list_reader import read_supplier_price_list

    parsed = read_supplier_price_list(_wide_letterhead_workbook(300, stray_column=16000), "stray.xlsx")
    assert parsed.total_rows == 300


def test_s3_an_oversized_sheet_has_its_own_code():
    from app.services.error_handler import AppException
    from app.services.procurement.supplier_price_list_reader import read_supplier_price_list

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(HEADER_ROW))
    for i in range(1, 10_002):
        ws.append([None, None, f"note {i}", None])  # filled rows that are never price rows
    buf = BytesIO()
    wb.save(buf)

    with pytest.raises(AppException) as exc:
        read_supplier_price_list(buf.getvalue(), "notes.xlsx")
    assert exc.value.status_code == 422
    assert exc.value.code == "sheet_too_large"


# --------------------------------------------------------------------------- Should fix 4


def test_s4_prices_tab_search_finds_the_suppliers_own_code(cost_price_env):
    from app.models.scm import SupplierProductCodeAlias

    e = cost_price_env
    e.as_user(e.user(PS_VIEW_PERM, VIEW_PERM))
    supplier = e.supplier()
    product = e.product(description="Nothing like the code")
    e.link(product, supplier, unit_cost=5, currency="CNY")
    e.link(e.product(), supplier, unit_cost=6, currency="CNY")
    e.db.add(SupplierProductCodeAlias(
        supplier_id=supplier.id, supplier_code="THEIRS-777", product_id=product.id,
        source="manual", matched_by="test",
    ))
    e.db.commit()

    r = e.cost_lists(str(supplier.id), query="THEIRS-777")
    assert r.status_code == 200, r.text
    assert [row["product"]["id"] for row in r.json()["data"]] == [str(product.id)]


# --------------------------------------------------------------------------- Should fix 5


def test_s5_explicit_form_currency_beats_the_suppliers_links(cost_price_env):
    e = cost_price_env
    e.as_user(e.user(UPLOAD_PERM, VIEW_PERM))
    supplier = e.supplier(name=LETTERHEAD_TEXT)
    e.link(e.product(), supplier, unit_cost=50, currency="CNY")
    e.db.commit()

    data = simple_price_list_workbook([("ZZCPC-USD-1", "cfg", 10)], letterhead=LETTERHEAD_TEXT)
    r = e.upload(data, supplier_id=str(supplier.id), currency="USD")
    assert r.status_code == 201, r.text
    assert r.json()["currency"] == "USD"


def test_s5_form_currency_conflicting_with_the_header_token_is_422(cost_price_env):
    e = cost_price_env
    e.as_user(e.user(UPLOAD_PERM, VIEW_PERM))
    supplier = e.supplier(name=LETTERHEAD_TEXT)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([LETTERHEAD_TEXT])
    ws.append(["序号", "型号", "产品配置", "价格(RMB)"])
    ws.append([1, "ZZCPC-HDR-1", "cfg", 10])
    buf = BytesIO()
    wb.save(buf)

    r = e.upload(buf.getvalue(), supplier_id=str(supplier.id), currency="USD")
    assert r.status_code == 422, r.text
    assert r.json().get("detail", {}).get("code") == "currency_conflict" or "currency_conflict" in r.text


# --------------------------------------------------------------------------- Should fix 6


def test_s6_a_draft_set_can_recapture_live_prices_after_going_stale(cost_price_env):
    from app.models.procurement import ProductSupplier

    e = cost_price_env
    e.as_user(e.user(UPLOAD_PERM, VIEW_PERM))
    e.seed_settings(cost_price_verification_enabled=False)
    supplier = e.supplier(name=LETTERHEAD_TEXT)
    product = e.product(code="ZZCPC-STALE-1")
    link = e.link(product, supplier, unit_cost=100, currency="CNY")
    e.db.commit()

    data = simple_price_list_workbook([("ZZCPC-STALE-1", "cfg", 110)], letterhead=LETTERHEAD_TEXT)
    set_id = e.upload(data, supplier_id=str(supplier.id), currency="CNY").json()["id"]

    # The tick or a hand edit moves the live price after the upload.
    e.db.query(ProductSupplier).filter_by(id=link.id).update({"unit_cost": Decimal("105")})
    e.db.commit()

    first = e.apply(set_id)
    assert first.status_code == 409, first.text
    line = e.lines(set_id).json()["data"][0]
    assert line["stale"] == {"live_unit_cost": 105.0, "live_currency": "CNY"}
    assert e.detail(set_id).json()["actions"]["can_refresh_prices"] is True

    refreshed = e.client.post(f"/api/v1/procurement/cost-price-changes/{set_id}/refresh-prices")
    assert refreshed.status_code == 200, refreshed.text
    assert refreshed.json()["actions"]["can_refresh_prices"] is False
    line = e.lines(set_id).json()["data"][0]
    assert line["stale"] is None
    assert line["current_unit_cost"] == 105.0

    applied = e.apply(set_id)
    assert applied.status_code == 200, applied.text


# --------------------------------------------------------------------------- Should fix 8


def _pending_set(e, *, codes=("ZZCPC-PEND-1", "ZZCPC-PEND-2")):
    e.seed_settings(cost_price_verification_enabled=True)
    uploader = e.user(UPLOAD_PERM, VIEW_PERM, VERIFY_PERM, name="Uploader")
    e.as_user(uploader)
    supplier = e.supplier(name=LETTERHEAD_TEXT)
    for code in codes:
        e.link(e.product(code=code), supplier, unit_cost=100, currency="CNY")
    e.db.commit()
    data = simple_price_list_workbook([(c, "cfg", 120) for c in codes], letterhead=LETTERHEAD_TEXT)
    set_id = e.upload(data, supplier_id=str(supplier.id), currency="CNY").json()["id"]
    assert e.submit(set_id).status_code == 200
    return uploader, set_id


def test_s8_pending_set_says_why_apply_is_disabled(cost_price_env):
    e = cost_price_env
    uploader, set_id = _pending_set(e)

    own = e.detail(set_id).json()["actions"]
    assert own["can_apply"] is False
    assert own["apply_blocked_reason"] == "You uploaded, submitted or mapped this set"

    e.as_user(e.user(VIEW_PERM, VERIFY_PERM, name="Kelvin"))
    other = e.detail(set_id).json()["actions"]
    assert other["can_apply"] is False
    assert other["apply_blocked_reason"] == "2 lines undecided"


# ----------------------------------------------------------------------------------- Nits


def test_n1_verifier_notice_skips_inactive_users_and_the_submitter(cost_price_env):
    from app.models.notification import Notification
    from app.models.user import User

    e = cost_price_env
    e.seed_settings(cost_price_verification_enabled=True)
    active = e.user(VIEW_PERM, VERIFY_PERM, name="Active")
    inactive = e.user(VIEW_PERM, VERIFY_PERM, name="Gone")
    uploader = e.user(UPLOAD_PERM, VIEW_PERM, VERIFY_PERM, name="Uploader")
    for u in (active, inactive, uploader):
        e.grant_company(u, e.company_a)
    e.db.query(User).filter_by(id=inactive["id"]).update({"status": "INACTIVE"})
    e.db.commit()

    e.as_user(uploader)
    supplier = e.supplier(name=LETTERHEAD_TEXT)
    e.link(e.product(code="ZZCPC-N1-1"), supplier, unit_cost=100, currency="CNY")
    e.db.commit()
    data = simple_price_list_workbook([("ZZCPC-N1-1", "cfg", 120)], letterhead=LETTERHEAD_TEXT)
    set_id = e.upload(data, supplier_id=str(supplier.id), currency="CNY").json()["id"]
    assert e.submit(set_id).status_code == 200

    notified = {
        str(n.user_id)
        for n in e.db.query(Notification).filter(Notification.source_entity_id == set_id).all()
    }
    assert active["id"] in notified
    assert inactive["id"] not in notified
    assert uploader["id"] not in notified


def test_n2_apply_audit_names_the_product_and_the_dates(cost_price_env):
    from app.models.audit import AuditLog

    e = cost_price_env
    e.as_user(e.user(UPLOAD_PERM, VIEW_PERM))
    e.seed_settings(cost_price_verification_enabled=False)
    supplier = e.supplier(name=LETTERHEAD_TEXT)
    product = e.product(code="ZZCPC-AUD-1")
    e.link(product, supplier, unit_cost=100, currency="CNY")
    e.db.commit()
    data = simple_price_list_workbook([("ZZCPC-AUD-1", "cfg", 110)], letterhead=LETTERHEAD_TEXT)
    set_id = e.upload(
        data, supplier_id=str(supplier.id), currency="CNY", start_date="2099-01-01", end_date="2099-12-31",
    ).json()["id"]
    assert e.apply(set_id).status_code == 200

    row = e.db.query(AuditLog).filter(
        AuditLog.entity_id == set_id, AuditLog.new_values["event"].astext == "COST_SET_APPLY"
    ).one()
    change = row.new_values["changes"][0]
    assert change["product_code"] == "ZZCPC-AUD-1"
    assert (change["start_date"], change["end_date"]) == ("2099-01-01", "2099-12-31")


def test_n3_prices_tab_query_count_does_not_grow_with_links(cost_price_env):
    from sqlalchemy import event

    e = cost_price_env
    e.as_user(e.user(PS_VIEW_PERM, VIEW_PERM))
    supplier = e.supplier()
    for _ in range(12):
        link = e.link(e.product(), supplier, unit_cost=5, currency="CNY")
        _cost_row(e, link, 5)
    e.db.commit()

    statements: list[str] = []

    def _count(conn, cursor, statement, *args):
        if statement.lstrip().upper().startswith("SELECT"):
            statements.append(statement)

    engine = e.db.get_bind()
    event.listen(engine, "before_cursor_execute", _count)
    try:
        r = e.cost_lists(str(supplier.id))
    finally:
        event.remove(engine, "before_cursor_execute", _count)
    assert r.status_code == 200, r.text
    assert len(r.json()["data"]) == 12
    assert len(statements) < 12, len(statements)


def test_n6_apply_count_on_a_pending_set_counts_accepted_lines(cost_price_env):
    e = cost_price_env
    _uploader, set_id = _pending_set(e, codes=("ZZCPC-N6-1", "ZZCPC-N6-2", "ZZCPC-N6-3"))
    e.as_user(e.user(VIEW_PERM, VERIFY_PERM, name="Kelvin"))
    lines = e.lines(set_id).json()["data"]
    e.decide(set_id, lines[0]["id"], {"decision": "rejected"})
    e.decide(set_id, lines[1]["id"], {"decision": "rejected"})
    e.decide(set_id, lines[2]["id"], {"decision": "accepted"})

    assert e.detail(set_id).json()["actions"]["apply_count"] == 1


def test_n8_cpc1_company_id_columns_reference_companies_and_are_indexed():
    import ast
    from pathlib import Path

    source = (Path(__file__).resolve().parents[1] / "alembic" / "versions" / "cpc1_supplier_cost_lists.py").read_text()
    upgrade_src = source[source.index("def upgrade"):source.index("def downgrade")]
    tree = ast.parse(upgrade_src)
    company_cols = [
        node for node in ast.walk(tree)
        if isinstance(node, ast.Call) and getattr(node.func, "attr", "") == "Column"
        and node.args and isinstance(node.args[0], ast.Constant) and node.args[0].value == "company_id"
    ]
    assert len(company_cols) == 3
    for col in company_cols:
        assert "companies.id" in ast.unparse(col), ast.unparse(col)
    for table in ("cost_price_change_sets", "product_supplier_costs", "supplier_price_links"):
        assert f'op.create_index("ix_{table}_company_id", "{table}", ["company_id"])' in upgrade_src


def test_n9_a_racing_second_upload_gets_409_not_500(cost_price_env, monkeypatch):
    from app.models.cost_price import CostPriceChangeSet
    from app.services.procurement import cost_price_change_service as svc

    e = cost_price_env
    e.as_user(e.user(UPLOAD_PERM, VIEW_PERM))
    supplier = e.supplier(name=LETTERHEAD_TEXT)
    e.db.commit()

    real_reader = svc.read_supplier_price_list

    def _reader_losing_the_race(data, filename):
        # Another request commits an open set for the same supplier after this one's
        # open-set check has already passed.
        e.db.add(CostPriceChangeSet(
            code=unique_code("CPCRACE"), supplier_id=supplier.id, channel="staff_upload",
            status="draft", currency="CNY",
        ))
        e.db.commit()
        return real_reader(data, filename)

    monkeypatch.setattr(svc, "read_supplier_price_list", _reader_losing_the_race)
    data = simple_price_list_workbook([("ZZCPC-RACE-1", "cfg", 10)], letterhead=LETTERHEAD_TEXT)
    r = e.upload(data, supplier_id=str(supplier.id), currency="CNY")
    assert r.status_code == 409, r.text


def test_n10_non_integer_lead_time_is_422(cost_price_env):
    e = cost_price_env
    e.as_user(e.user(UPLOAD_PERM, VIEW_PERM))
    supplier = e.supplier(name=LETTERHEAD_TEXT)
    e.db.commit()
    data = simple_price_list_workbook([("ZZCPC-LT-1", "cfg", 10)], letterhead=LETTERHEAD_TEXT)
    set_id = e.upload(data, supplier_id=str(supplier.id), currency="CNY").json()["id"]
    line_id = e.lines(set_id).json()["data"][0]["id"]

    r = e.patch_line(set_id, line_id, {"new_link_lead_time_days": "abc"})
    assert r.status_code == 422, r.text


@pytest.mark.parametrize("action", ["decide", "decide_all", "return"])
def test_n11_decide_and_return_need_one_company(cost_price_env, action):
    e = cost_price_env
    _uploader, set_id = _pending_set(e, codes=("ZZCPC-N11-1",))
    verifier = e.user(VIEW_PERM, VERIFY_PERM, name="Kelvin")
    e.as_user(verifier, scope=frozenset({e.company_a, e.company_b}))
    line_id = e.lines(set_id).json()["data"][0]["id"]

    if action == "decide":
        r = e.decide(set_id, line_id, {"decision": "accepted"})
    elif action == "decide_all":
        r = e.decide_all(set_id, "accepted")
    else:
        r = e.return_set(set_id, "wrong list")
    assert r.status_code == 422, r.text
    assert "pick_one_company" in r.text
