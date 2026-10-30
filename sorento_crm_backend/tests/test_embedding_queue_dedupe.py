"""EMBED-QUEUE-DEDUPE: one pending embedding_queue row per source, cheap poll, prune, backlog dedupe.

Plan: documentation/plans/chatbot/PLAN-embed-queue-dedupe-27oct.md

Prod (27 Oct 2026) held ~1.17M pending rows, ~94% duplicates: every flush of a tracked model
queued a row, also when nothing embedded had changed, and the poll sorted the whole table.
Every test runs in a blank scratch schema (empty tables, writes discarded).
"""
from __future__ import annotations

import uuid
from datetime import date, datetime, timedelta
from io import BytesIO

import pytest
from sqlalchemy import inspect as sa_inspect, text

from app.database import engine
from app.models.embeddings import EmbeddingQueue
from app.models.order import Order, OrderStatus
from app.services import embedding_change_listener as ecl
from app.services.embedding_change_listener import (
    EMBEDDED_FIELDS,
    SUPERSEDED_ERROR,
    bulk_enqueue_embedding_events,
    enqueue_pending,
    register_embedding_change_listeners,
)
from tests._pg_fixture import _globally_required_tables, blank_session

register_embedding_change_listeners()

_USER = "11111111-1111-1111-1111-111111111111"


@pytest.fixture()
def db():
    with blank_session() as session:
        yield session


def _rows(db, source_type, source_id, status=None):
    q = db.query(EmbeddingQueue).filter(
        EmbeddingQueue.source_type == source_type, EmbeddingQueue.source_id == str(source_id)
    )
    if status:
        q = q.filter(EmbeddingQueue.status == status)
    return q.order_by(EmbeddingQueue.created_at).all()


def _queue_row(db, source_type="promotion", source_id=None, *, status="pending", age=timedelta(0), payload=None):
    now = datetime.utcnow() - age
    row = EmbeddingQueue(
        id=str(uuid.uuid4()), source_type=source_type, source_id=source_id or str(uuid.uuid4()),
        event_type=f"{source_type}.updated", event_version=1, payload=payload or {},
        status=status, retry_count=0, available_at=now, created_at=now,
    )
    db.add(row)
    db.flush()
    return row


def _order(db, number="ZZT-EQD-1"):
    """A keep-set source to queue against (named `order` from when this file used orders; orders no
    longer queue). `description` is the embedded field, `expiry_notified_at` one that is not."""
    from app.models.marketing import Promotion

    order = Promotion(description="r1", start_date=date(2026, 10, 1), end_date=date(2026, 12, 1), is_active=True)
    db.add(order)
    db.flush()
    return order


def _complete_all(db):
    db.query(EmbeddingQueue).update({"status": "completed"}, synchronize_session=False)
    db.flush()


# --------------------------------------------------------------------------- #
# (2) an update queues only when an embedded field changed
# --------------------------------------------------------------------------- #


def test_every_embedded_field_is_a_column_of_its_model():
    """A misspelt field would silently never count as a change."""
    from app.models.certificate import Certificate
    from app.models.marketing import Promotion, PromotionAttachment, PromotionProduct
    from app.models.order import Customer, OrderLine, Transporter
    from app.models.procurement import (
        InboundShipment, InboundShipmentLine, PickingHeader, PickingLine, SPOAllocation,
    )
    from app.models.product import Product, ProductAttachment

    models = {
        "product": Product, "promotion": Promotion, "promotion_product": PromotionProduct,
        "promotion_attachment": PromotionAttachment, "product_attachment": ProductAttachment,
        "inbound_shipment": InboundShipment, "inbound_shipment_line": InboundShipmentLine,
        "spo_allocation": SPOAllocation, "picking_header": PickingHeader, "picking_line": PickingLine,
        "certificate": Certificate, "order": Order, "order_status": OrderStatus,
        "order_line": OrderLine, "customer": Customer, "transporter": Transporter,
    }
    assert set(EMBEDDED_FIELDS) == set(models)
    for source_type, model in models.items():
        columns = {a.key for a in sa_inspect(model).column_attrs}
        missing = set(EMBEDDED_FIELDS[source_type]) - columns
        assert not missing, (source_type, missing)


def test_insert_queues_one_row(db):
    order = _order(db)
    assert len(_rows(db, "promotion", order.id, "pending")) == 1


def test_update_of_a_non_embedded_field_queues_nothing(db):
    order = _order(db)
    _complete_all(db)
    order.expiry_notified_at = datetime(2026, 10, 2)
    db.flush()
    assert _rows(db, "promotion", order.id, "pending") == []


def test_setting_an_embedded_field_to_its_current_value_queues_nothing(db):
    order = _order(db)
    _complete_all(db)
    order.description = "r1"
    db.flush()
    assert _rows(db, "promotion", order.id, "pending") == []


def test_update_of_an_embedded_field_queues_one_row(db):
    order = _order(db)
    _complete_all(db)
    order.description = "r2"
    db.flush()
    assert len(_rows(db, "promotion", order.id, "pending")) == 1


def test_spo_retirement_still_queues(db):
    """retired_at is not in the body but decides whether the worker hides the line (R9)."""
    assert "retired_at" in EMBEDDED_FIELDS["spo_allocation"]
    assert "quantity_received" in EMBEDDED_FIELDS["spo_allocation"]


# --------------------------------------------------------------------------- #
# (1) dedupe at enqueue
# --------------------------------------------------------------------------- #


def test_repeated_changes_keep_one_pending_row(db):
    order = _order(db)
    for i in range(5):
        order.description = f"change {i}"
        db.flush()
    assert len(_rows(db, "promotion", order.id, "pending")) == 1


def test_a_processing_row_does_not_block_a_new_pending_row(db):
    """A claimed row may already have read the source; a later change needs its own row."""
    order = _order(db)
    db.query(EmbeddingQueue).update({"status": "processing"}, synchronize_session=False)
    order.description = "after the claim"
    db.flush()
    assert len(_rows(db, "promotion", order.id, "pending")) == 1


def test_bulk_enqueue_skips_sources_already_pending(db):
    pending = _queue_row(db, "promotion")
    fresh = str(uuid.uuid4())
    inserted = bulk_enqueue_embedding_events(db, "promotion", [pending.source_id, fresh, fresh])
    assert inserted == 1
    assert len(_rows(db, "promotion", pending.source_id, "pending")) == 1
    assert len(_rows(db, "promotion", fresh, "pending")) == 1


def test_queue_event_reuses_the_pending_row_and_takes_the_new_payload(db, monkeypatch):
    """A queued row carries its payload; the newest payload must win (a keep-set `form` source stands in for the dropped mcp_tool)."""
    from app.services import embedding_service

    jobs = []

    class _Job:
        id = "job-1"

    monkeypatch.setattr(embedding_service, "enqueue_job", lambda *a, **k: jobs.append(a) or _Job())
    svc = embedding_service.EmbeddingEventService(db)
    first = svc.queue_event(
        source_type="form", source_id="zzt_tool", event_type="form.updated",
        payload={"capability": {"body_text": "old"}},
    )
    second = svc.queue_event(
        source_type="form", source_id="zzt_tool", event_type="form.updated",
        payload={"capability": {"body_text": "new"}},
    )
    assert second.id == first.id
    assert len(_rows(db, "form", "zzt_tool")) == 1
    assert second.payload["payload"]["capability"]["body_text"] == "new"
    assert len(jobs) == 1  # the pending row already has its RQ job


def test_a_failed_row_does_not_absorb_a_new_change(db):
    """A row in retry backoff is one failure nearer dead_letter; the new change gets its own."""
    order = _order(db)
    db.query(EmbeddingQueue).update({"retry_count": 2}, synchronize_session=False)
    order.description = "fresh change"
    db.flush()
    rows = _rows(db, "promotion", order.id, "pending")
    assert sorted(r.retry_count for r in rows) == [0, 2]


def test_queue_event_enqueues_a_job_for_a_pending_row_that_has_none(db, monkeypatch):
    """A listener row carries no RQ job; queue_event on the same source must give it one."""
    from app.services import embedding_service

    jobs = []

    class _Job:
        id = "job-2"

    monkeypatch.setattr(embedding_service, "enqueue_job", lambda *a, **k: jobs.append(a) or _Job())
    listener_row = _queue_row(db, "form", "zzt_tool_2")
    item = embedding_service.EmbeddingEventService(db).queue_event(
        source_type="form", source_id="zzt_tool_2", event_type="form.updated",
        payload={"capability": {"body_text": "x"}},
    )
    assert item.id == listener_row.id
    assert item.rq_job_id == "job-2" and len(jobs) == 1


# --------------------------------------------------------------------------- #
# Race safety: two real transactions on a scratch copy of embedding_queue
# --------------------------------------------------------------------------- #


@pytest.fixture()
def two_sessions():
    """Two sessions on separate connections, both seeing a scratch embedding_queue, so one can
    hold row locks while the other tries to claim. Rows here commit (into the scratch schema)."""
    import os

    from sqlalchemy.orm import Session

    name = f"zzs_eqdrace_{os.getpid()}_{uuid.uuid4().hex[:8]}"
    admin = engine.connect().execution_options(isolation_level="AUTOCOMMIT")
    admin.exec_driver_sql(f'CREATE SCHEMA "{name}"')
    admin.exec_driver_sql(f'CREATE TABLE "{name}".embedding_queue (LIKE public.embedding_queue INCLUDING ALL)')
    # The search_path below is this schema alone, so the tables app.main's global flush
    # listeners query (lookup_bindings on every ORM insert/update) must exist here too. Without
    # them the claim's flush raised UndefinedTable whenever an earlier file on the same
    # xdist worker had primed lookup_validator's per-engine "table exists" cache on public.
    for table in _globally_required_tables():
        admin.exec_driver_sql(
            f'CREATE TABLE "{name}".{table.name} (LIKE public.{table.name} INCLUDING ALL)'
        )
    conns, sessions = [], []
    for _ in range(2):
        conn = engine.connect()
        conn.exec_driver_sql(f'SET search_path TO "{name}"')
        # A broken lock fails the test instead of hanging it.
        conn.exec_driver_sql("SET statement_timeout = '5s'")
        conn.commit()
        conns.append(conn)
        sessions.append(Session(bind=conn))
    try:
        yield sessions
    finally:
        for sess in sessions:
            sess.rollback()
            sess.close()
        for conn in conns:
            conn.exec_driver_sql("RESET search_path")
            conn.exec_driver_sql("RESET statement_timeout")
            conn.close()
        admin.exec_driver_sql(f'DROP SCHEMA "{name}" CASCADE')
        admin.close()


def _row_dict(source_id, source_type="promotion"):
    now = datetime.utcnow()
    return {
        "id": str(uuid.uuid4()), "source_type": source_type, "source_id": source_id,
        "event_type": f"{source_type}.updated", "event_version": 1, "payload": {},
        "status": "pending", "retry_count": 0, "available_at": now, "created_at": now,
    }


def test_worker_cannot_claim_a_row_an_open_enqueue_relied_on(two_sessions, monkeypatch):
    """The lost update this guards: an enqueue skips its insert because a row is pending, the
    worker claims that row and reads the source before the enqueue's change commits."""
    from app.services import embedding_worker

    enqueuer, worker_db = two_sessions
    sid = str(uuid.uuid4())
    pending = _row_dict(sid)
    enqueue_pending(enqueuer, [pending])
    enqueuer.commit()

    assert enqueue_pending(enqueuer, [_row_dict(sid)]) == []  # skipped, transaction still open
    monkeypatch.setattr(embedding_worker, "SessionLocal", lambda: worker_db)
    monkeypatch.setattr(worker_db, "close", lambda: None)
    assert embedding_worker.process_embedding_queue_item(pending["id"]) == {
        "status": "busy", "queue_id": pending["id"],
    }
    assert worker_db.get(EmbeddingQueue, pending["id"]).status == "pending"

    enqueuer.commit()
    claimed = embedding_worker._claim(worker_db, pending["id"])
    assert claimed is not None and claimed.status == "pending"


def test_an_enqueue_after_an_uncommitted_claim_queues_its_own_row(two_sessions):
    """The claim may already have read the source, so the enqueue must not count on it."""
    import threading

    from app.services import embedding_worker

    enqueuer, worker_db = two_sessions
    sid = str(uuid.uuid4())
    pending = _row_dict(sid)
    enqueue_pending(worker_db, [pending])
    worker_db.commit()

    claimed = embedding_worker._claim(worker_db, pending["id"])
    claimed.status = "processing"
    worker_db.flush()  # claimed, not committed

    result = {}
    thread = threading.Thread(target=lambda: result.setdefault("rows", enqueue_pending(enqueuer, [_row_dict(sid)])))
    thread.start()
    thread.join(0.5)
    assert thread.is_alive(), "the enqueue should wait for the claim to commit"
    worker_db.commit()
    thread.join(5)
    assert len(result["rows"]) == 1


def test_claim_and_backlog_dedupe_never_wait_on_a_held_row(two_sessions, monkeypatch):
    """A duplicate an open transaction holds is skipped by both, not waited on (the 5s
    statement_timeout turns a wait into a failure) and stays pending."""
    from app.services import embedding_worker
    from app.services.embedding_service import dedupe_pending_embedding_queue

    holder, worker_db = two_sessions
    sid = str(uuid.uuid4())
    older, held = _row_dict(sid), _row_dict(sid)
    held["created_at"] = held["available_at"] = datetime.utcnow() - timedelta(minutes=5)
    older["created_at"] = older["available_at"] = datetime.utcnow() - timedelta(minutes=9)
    newest = _row_dict(sid)
    holder.execute(EmbeddingQueue.__table__.insert(), [older, held, newest])
    holder.commit()
    holder.execute(text("SELECT id FROM embedding_queue WHERE id = :i FOR SHARE"), {"i": held["id"]})

    claimed = embedding_worker._claim(worker_db, older["id"])
    assert embedding_worker._supersede_pending_duplicates(worker_db, claimed) == 1  # newest only
    worker_db.rollback()
    assert dedupe_pending_embedding_queue(worker_db, batch_size=10, pause_seconds=0)["rows_skipped"] == 1
    holder.commit()
    statuses = dict(worker_db.execute(text("SELECT id::text, status FROM embedding_queue")).all())
    assert statuses[held["id"]] == "pending"


def test_claiming_a_retrying_row_keeps_the_fresh_rows(db, worker):
    order = _order(db)
    _complete_all(db)
    retrying = _queue_row(db, "promotion", str(order.id), age=timedelta(minutes=9))
    retrying.retry_count = 2
    fresh = _queue_row(db, "promotion", str(order.id))
    db.flush()
    worker.process_embedding_queue_item(retrying.id)
    db.expire_all()
    assert db.get(EmbeddingQueue, fresh.id).status == "pending"


def test_backlog_dedupe_keeps_a_fresh_row_over_a_newer_retrying_one(db):
    from app.services.embedding_service import dedupe_pending_embedding_queue

    sid = str(uuid.uuid4())
    fresh = _queue_row(db, "promotion", sid, age=timedelta(minutes=5))
    retrying = _queue_row(db, "promotion", sid, age=timedelta(minutes=1))
    retrying.retry_count = 1
    db.flush()
    dedupe_pending_embedding_queue(db, batch_size=10, pause_seconds=0)
    db.expire_all()
    assert (db.get(EmbeddingQueue, fresh.id).status, db.get(EmbeddingQueue, retrying.id).status) == (
        "pending", "skipped",
    )


# --------------------------------------------------------------------------- #
# Worker claim: supersede, retry
# --------------------------------------------------------------------------- #


@pytest.fixture()
def worker(db, monkeypatch):
    from app.services import embedding_worker

    monkeypatch.setattr(embedding_worker, "SessionLocal", lambda: db)
    monkeypatch.setattr(db, "close", lambda: None)
    monkeypatch.setattr(embedding_worker, "_embed_text_chunks", lambda chunks, timeout=None: [[0.0] * 1536 for _ in chunks])
    return embedding_worker


def test_claim_supersedes_the_other_pending_rows_of_its_source(db, worker):
    order = _order(db)
    _complete_all(db)
    old = [_queue_row(db, "promotion", str(order.id), age=timedelta(minutes=m)) for m in (3, 2, 1)]
    other_source = _queue_row(db, "order")

    result = worker.process_embedding_queue_item(old[0].id)

    assert result["status"] == "completed"
    db.expire_all()
    statuses = {r.id: (r.status, r.last_error) for r in _rows(db, "promotion", order.id)}
    assert statuses[old[0].id][0] == "completed"
    assert statuses[old[1].id] == ("skipped", SUPERSEDED_ERROR)
    assert statuses[old[2].id] == ("skipped", SUPERSEDED_ERROR)
    assert db.get(EmbeddingQueue, other_source.id).status == "pending"


def test_claiming_an_older_mcp_tool_row_leaves_the_newer_payload_pending(db, worker):
    """mcp_tool rows carry their text in the payload, so an older one never covers a newer."""
    older = _queue_row(db, "mcp_tool", "zzt_tool_3", age=timedelta(minutes=5),
                       payload={"payload": {"capability": {"body_text": "old"}}})
    newer = _queue_row(db, "mcp_tool", "zzt_tool_3", age=timedelta(minutes=1),
                       payload={"payload": {"capability": {"body_text": "new"}}})
    worker.process_embedding_queue_item(older.id)
    db.expire_all()
    assert db.get(EmbeddingQueue, newer.id).status == "pending"


def test_failed_run_does_not_go_back_to_pending_beside_a_newer_row(db, worker, monkeypatch):
    order = _order(db)
    _complete_all(db)
    claimed = _queue_row(db, "promotion", str(order.id))

    def _fail(chunks, timeout=None):
        order.description = "changed while the worker ran"
        db.commit()  # another transaction's change: the listener queues a fresh pending row
        raise RuntimeError("provider down")

    monkeypatch.setattr(worker, "_embed_text_chunks", _fail)
    worker.process_embedding_queue_item(claimed.id)

    db.expire_all()
    assert db.get(EmbeddingQueue, claimed.id).status == "skipped"
    assert len(_rows(db, "promotion", order.id, "pending")) == 1


def test_failed_run_retries_when_nothing_newer_is_pending(db, worker, monkeypatch):
    order = _order(db)
    _complete_all(db)
    claimed = _queue_row(db, "promotion", str(order.id))
    monkeypatch.setattr(worker, "_embed_text_chunks", lambda chunks, timeout=None: (_ for _ in ()).throw(RuntimeError("x")))
    worker.process_embedding_queue_item(claimed.id)
    db.expire_all()
    row = db.get(EmbeddingQueue, claimed.id)
    assert (row.status, row.retry_count) == ("pending", 1)


# --------------------------------------------------------------------------- #
# (3) Order Tracking import queues only changed orders
# --------------------------------------------------------------------------- #

_MASTER_HEADER = ["Doc. No.", "Date", "Created Time", "Debtor Code", "Debtor Name", "Agent",
                  "Cancel", "Remarks CS", "Type"]


def _workbook(rows: list[list]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    master = wb.create_sheet("Master")
    master.append(_MASTER_HEADER)
    for r in rows:
        master.append(r)
    tracking = wb.create_sheet("Overall Tracking")
    tracking.append(["Doc Number", "Date", "Time", "Checker", "Transporter", "Driver Name",
                     "Lorry Plate", "Customer", "Remarks CS", "Remarks", "Salesman", "Trips", "W/H"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_tracking_import_queues_no_order_embeddings_at_all(db):
    from app.services.order_service import OrderService

    for code, name in [("NEW", "New Order"), ("DELIVERED", "Delivered")]:
        db.add(OrderStatus(status_code=code, status_name=name))
    db.commit()
    rows = [
        ["ZZT-EQD-A", date(2026, 9, 1), "01/09/2026 10:00", "300-A", "Debtor A", "AG", "N", "cs", "RMA"],
        ["ZZT-EQD-B", date(2026, 9, 1), "01/09/2026 10:00", "300-B", "Debtor B", "AG", "N", "cs", "RMA"],
    ]
    result = OrderService(db).import_excel_tracking(_workbook(rows), _USER)
    assert not result.get("errors"), result.get("errors")
    orders = {o.order_number: o for o in db.query(Order).filter(Order.order_number.like("ZZT-EQD-%"))}
    # Changed meaning (LOOKUP-EMBED-CENTRAL): orders left the keep set, so an import queues nothing for them,
    # new or changed. The per-field "queue only on a real change" rule is exercised on a keep-set source above.
    assert orders and all(_rows(db, "order", o.id) == [] for o in orders.values())

    _complete_all(db)
    db.commit()
    rows[1][4] = "Debtor B renamed"
    result = OrderService(db).import_excel_tracking(_workbook(rows), _USER)
    assert not result.get("errors"), result.get("errors")

    assert all(_rows(db, "order", o.id) == [] for o in orders.values())


# --------------------------------------------------------------------------- #
# (5) poll: index-backed pick, configurable batch, no count(*)
# --------------------------------------------------------------------------- #


def test_poll_takes_mcp_tool_first_then_oldest_up_to_the_batch(db, monkeypatch):
    from app.scheduler import task_scheduler
    from app.services import embedding_worker

    picked: list[str] = []
    monkeypatch.setattr(embedding_worker, "process_embedding_queue_item", picked.append)
    oldest = _queue_row(db, "promotion", age=timedelta(minutes=10))
    _queue_row(db, "promotion", age=timedelta(minutes=5))
    tool = _queue_row(db, "mcp_tool", age=timedelta(minutes=1))
    _queue_row(db, "promotion", status="completed", age=timedelta(hours=1))
    _queue_row(db, "promotion", age=timedelta(minutes=-10))  # backing off, not available yet

    result = task_scheduler._run_embedding_db_queue_fallback(db, max_jobs_per_run=2)

    assert picked == [tool.id, oldest.id]
    assert result == {"processed": 2}


def test_poll_pick_reads_the_pending_poll_index(db):
    from app.models.embeddings import EMBEDDING_QUEUE_POLL_PRIORITY

    db.execute(text("SET LOCAL enable_seqscan = off"))
    plan = "\n".join(
        r[0]
        for r in db.execute(
            text(
                "EXPLAIN SELECT id FROM embedding_queue WHERE status = 'pending' AND available_at <= now() "
                f"ORDER BY {EMBEDDING_QUEUE_POLL_PRIORITY} ASC, created_at ASC LIMIT 10"
            )
        )
    )
    assert "ix_embedding_queue_pending_poll_cov" in plan, plan
    assert "Sort" not in plan, plan


def test_processor_uses_the_configured_batch_size(db, monkeypatch):
    from app.scheduler import task_scheduler

    seen = {}
    monkeypatch.setattr(task_scheduler.settings, "embedding_queue_batch_size", 37)
    monkeypatch.setattr(
        task_scheduler, "_run_queue_jobs_impl",
        lambda name, max_jobs_per_run: seen.setdefault("rq", max_jobs_per_run) and {"processed": 0, "queued": 0},
    )
    monkeypatch.setattr(
        task_scheduler, "_run_embedding_db_queue_fallback",
        lambda db, max_jobs_per_run: seen.setdefault("db", max_jobs_per_run) and {"processed": 0},
    )
    task_scheduler._handler_embedding_job_processor(db, None)
    assert seen == {"rq": 37, "db": 37}


def test_batch_size_defaults_to_ten():
    from app.config import Settings

    assert Settings.model_fields["embedding_queue_batch_size"].default == 10


# --------------------------------------------------------------------------- #
# (6) prune of old completed/skipped rows
# --------------------------------------------------------------------------- #


def test_prune_deletes_only_old_terminal_rows_in_batches(db):
    from app.services.embedding_service import prune_embedding_queue

    old = timedelta(days=15)
    doomed = [_queue_row(db, status=s, age=old).id for s in ("completed", "completed", "skipped")]
    kept = [
        _queue_row(db, status="completed", age=timedelta(days=13)).id,
        _queue_row(db, status="pending", age=old).id,
        _queue_row(db, status="processing", age=old).id,
        _queue_row(db, status="dead_letter", age=old).id,
    ]

    result = prune_embedding_queue(db, retention_days=14, batch_size=2, max_batches=10)

    assert result["rows_deleted"] == 3 and result["batches"] == 2
    left = {r.id for r in db.query(EmbeddingQueue.id)}
    assert not left & set(doomed)
    assert set(kept) <= left


def test_prune_stops_after_max_batches_and_the_next_run_continues(db):
    from app.services.embedding_service import prune_embedding_queue

    for _ in range(5):
        _queue_row(db, status="completed", age=timedelta(days=20))
    assert prune_embedding_queue(db, retention_days=14, batch_size=2, max_batches=1)["rows_deleted"] == 2
    assert prune_embedding_queue(db, retention_days=14, batch_size=2, max_batches=10)["rows_deleted"] == 3


# --------------------------------------------------------------------------- #
# One-time backlog dedupe job
# --------------------------------------------------------------------------- #


def test_backlog_dedupe_keeps_the_newest_pending_row_per_source(db):
    from app.services.embedding_service import dedupe_pending_embedding_queue

    sid = str(uuid.uuid4())
    olds = [_queue_row(db, "promotion", sid, age=timedelta(minutes=m)).id for m in (5, 4, 3)]
    newest = _queue_row(db, "promotion", sid, age=timedelta(minutes=1)).id
    tool_old = _queue_row(db, "mcp_tool", "zzt_tool", age=timedelta(minutes=9), payload={"v": 1}).id
    tool_new = _queue_row(db, "mcp_tool", "zzt_tool", age=timedelta(minutes=2), payload={"v": 2}).id
    single = _queue_row(db, "order").id
    done = _queue_row(db, "promotion", sid, status="completed", age=timedelta(minutes=30)).id

    result = dedupe_pending_embedding_queue(db, batch_size=2, max_batches=10, pause_seconds=0)

    assert result == {"rows_skipped": 4, "batches": 2}
    db.expire_all()
    status = {r.id: (r.status, r.last_error) for r in db.query(EmbeddingQueue)}
    for row_id in olds + [tool_old]:
        assert status[row_id] == ("skipped", "dedup: superseded")
    for row_id in (newest, tool_new, single):
        assert status[row_id][0] == "pending"
    assert status[done][0] == "completed"


def test_backlog_dedupe_second_run_is_a_no_op(db):
    from app.services.embedding_service import dedupe_pending_embedding_queue

    sid = str(uuid.uuid4())
    for m in (3, 2, 1):
        _queue_row(db, "promotion", sid, age=timedelta(minutes=m))
    assert dedupe_pending_embedding_queue(db, batch_size=10, pause_seconds=0)["rows_skipped"] == 2
    assert dedupe_pending_embedding_queue(db, batch_size=10, pause_seconds=0) == {"rows_skipped": 0, "batches": 0}


def test_backlog_dedupe_resumes_after_a_capped_run(db):
    from app.services.embedding_service import dedupe_pending_embedding_queue

    sid = str(uuid.uuid4())
    for m in range(6, 0, -1):
        _queue_row(db, "promotion", sid, age=timedelta(minutes=m))
    first = dedupe_pending_embedding_queue(db, batch_size=2, max_batches=1, pause_seconds=0)
    second = dedupe_pending_embedding_queue(db, batch_size=2, max_batches=10, pause_seconds=0)
    assert (first["rows_skipped"], second["rows_skipped"]) == (2, 3)
    assert len(_rows(db, "promotion", sid, "pending")) == 1


def test_scheduled_tasks_are_registered():
    from app.scheduler.task_scheduler import register_task_handlers
    from app.services.scheduled_task_service import TASK_HANDLERS

    register_task_handlers()
    assert {"embedding_queue_prune", "embedding_queue_dedupe"} <= set(TASK_HANDLERS)


# --------------------------------------------------------------------------- #
# Index builds never stall a release: CONCURRENTLY, bounded, failure is a WARNING
# --------------------------------------------------------------------------- #


@pytest.fixture()
def queue_schema():
    """An autocommit connection whose search_path holds a bare copy of embedding_queue (no
    indexes), since CREATE INDEX CONCURRENTLY cannot run inside the rolled-back test session."""
    import os

    name = f"zzs_eqd_{os.getpid()}_{uuid.uuid4().hex[:8]}"
    conn = engine.connect().execution_options(isolation_level="AUTOCOMMIT")
    conn.exec_driver_sql(f'CREATE SCHEMA "{name}"')
    conn.exec_driver_sql(f'CREATE TABLE "{name}".embedding_queue (LIKE public.embedding_queue INCLUDING DEFAULTS)')
    conn.exec_driver_sql(f'SET search_path TO "{name}"')
    try:
        yield conn
    finally:
        conn.exec_driver_sql("RESET search_path")
        conn.exec_driver_sql(f'DROP SCHEMA "{name}" CASCADE')
        conn.close()


def _add_pending(conn, n=3):
    for _ in range(n):
        conn.execute(
            text(
                "INSERT INTO embedding_queue (id, source_type, source_id, event_type, event_version, payload, "
                "status, retry_count, available_at, created_at) VALUES (gen_random_uuid(), 'order', "
                "gen_random_uuid()::text, 'order.updated', 1, '{}', 'pending', 0, now(), now())"
            )
        )


def test_index_definitions_match_the_model():
    from app.models.embeddings import EmbeddingQueue as Model
    from app.services.embedding_queue_indexes import EMBEDDING_QUEUE_INDEXES

    model_indexes = {i.name for i in Model.__table__.indexes}
    assert set(EMBEDDING_QUEUE_INDEXES) <= model_indexes


def test_ensure_builds_missing_indexes_and_then_is_a_no_op(queue_schema):
    from app.services.embedding_queue_indexes import ensure_embedding_queue_indexes, index_valid

    _add_pending(queue_schema)
    assert ensure_embedding_queue_indexes(queue_schema) == {
        "ix_embedding_queue_pending_poll_cov": "built",
        "ix_embedding_queue_pending_source": "built",
    }
    assert index_valid(queue_schema, "ix_embedding_queue_pending_poll_cov") is True
    assert set(ensure_embedding_queue_indexes(queue_schema).values()) == {"present"}


def test_ensure_restores_the_session_timeouts(queue_schema):
    from app.services.embedding_queue_indexes import ensure_embedding_queue_indexes

    queue_schema.exec_driver_sql("SET lock_timeout = '7s'")
    queue_schema.exec_driver_sql("SET statement_timeout = '123s'")
    ensure_embedding_queue_indexes(queue_schema)
    assert queue_schema.execute(text("SELECT current_setting('lock_timeout')")).scalar() == "7s"
    assert queue_schema.execute(text("SELECT current_setting('statement_timeout')")).scalar() == "123s"


def test_a_build_that_waits_on_a_live_lock_gives_up_warns_and_does_not_raise(queue_schema, caplog):
    """Owner, 27 Oct: an index build must never queue behind (or hold up) live work."""
    from app.services.embedding_queue_indexes import ensure_embedding_queue_indexes, index_valid

    schema = queue_schema.execute(text("SELECT current_schema()")).scalar()
    with engine.connect() as live:
        live.exec_driver_sql(f'LOCK TABLE "{schema}".embedding_queue IN SHARE UPDATE EXCLUSIVE MODE')
        with caplog.at_level("WARNING"):
            outcome = ensure_embedding_queue_indexes(queue_schema, lock_timeout="200ms")
        live.rollback()
    assert set(outcome.values()) == {"failed"}
    assert "not built" in caplog.text
    assert index_valid(queue_schema, "ix_embedding_queue_pending_poll_cov") is None
    # The retry path (hourly prune task) builds it once the table is free.
    assert set(ensure_embedding_queue_indexes(queue_schema).values()) == {"built"}


def test_a_failed_build_drops_its_invalid_leftover(queue_schema, caplog):
    from app.services.embedding_queue_indexes import ensure_embedding_queue_indexes, index_valid

    _add_pending(queue_schema)
    # Division by zero on the existing rows fails the build after its catalog entry exists,
    # which is where a statement timeout mid-build leaves an INVALID index.
    broken = {"ix_embedding_queue_pending_source": "ON embedding_queue ((1 / retry_count)) WHERE status = 'pending'"}
    with caplog.at_level("WARNING"):
        assert ensure_embedding_queue_indexes(queue_schema, indexes=broken) == {
            "ix_embedding_queue_pending_source": "failed"
        }
    assert index_valid(queue_schema, "ix_embedding_queue_pending_source") is None


def test_an_invalid_index_of_the_same_name_is_rebuilt(queue_schema):
    """A hand-made prod index that ended INVALID is replaced; a valid one is reused."""
    from sqlalchemy.exc import DBAPIError

    from app.services.embedding_queue_indexes import ensure_embedding_queue_indexes, index_valid

    _add_pending(queue_schema)
    queue_schema.execute(text("UPDATE embedding_queue SET source_type = 'dup'"))
    with pytest.raises(DBAPIError):
        queue_schema.exec_driver_sql(
            "CREATE UNIQUE INDEX CONCURRENTLY ix_embedding_queue_pending_source ON embedding_queue (source_type)"
        )
    assert index_valid(queue_schema, "ix_embedding_queue_pending_source") is False
    outcome = ensure_embedding_queue_indexes(queue_schema)
    assert outcome["ix_embedding_queue_pending_source"] == "built"
    assert index_valid(queue_schema, "ix_embedding_queue_pending_source") is True


def test_prune_task_retries_a_missing_index_only_in_the_retry_hour(db, monkeypatch):
    from app.scheduler import task_scheduler
    from app.services import embedding_queue_indexes as eqi

    built = []
    monkeypatch.setattr(eqi, "missing_indexes", lambda conn: ["ix_embedding_queue_pending_source"])
    monkeypatch.setattr(eqi, "ensure_embedding_queue_indexes", lambda conn: built.append(1) or {"x": "built"})

    class _Clock:
        hour = 0

        @classmethod
        def utcnow(cls):
            return datetime(2026, 10, 27, cls.hour, 0)

    monkeypatch.setattr(task_scheduler, "datetime", _Clock)
    monkeypatch.setattr(task_scheduler.settings, "embedding_queue_index_retry_hour_utc", 19)

    _Clock.hour = 10
    result = task_scheduler._handler_embedding_queue_prune(db, None)
    assert built == [] and result["indexes"] == {"ix_embedding_queue_pending_source": "missing"}
    _Clock.hour = 19
    assert task_scheduler._handler_embedding_queue_prune(db, None)["indexes"] == {"x": "built"}
    assert built == [1]
