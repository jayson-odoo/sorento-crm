"""S1 (#1267): the Yearly comparison through the kernel's own routes, and its export.

UAC: AC-R2-2 (catalog to holders only), AC-R2-3 (the per-report module check, fail
closed), AC-R2-4 (company arm and the export under the enqueuer's company), AC-R2-8 and
AC-R4-6 (My Downloads, the kernel's file name), AC-S1-5 (403 outside the grant).
"""
from __future__ import annotations

import uuid
from contextlib import contextmanager
from datetime import date
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient

from app.main import app  # noqa: E402  (first app import)
from app.dependencies import get_current_user, get_current_user_or_api_key, get_db
from app.models.base import set_company_scope
from app.models.company import UserCompany
from app.models.order import SalesOrder, SalesOrderLine
from app.models.user import User
from app.services.company_scope import DEFAULT_COMPANY_ID
from app.services.company_scope_resolver import apply_company_scope
from tests._mc_lookup_seed import MOCHA_ID, product, seed_mocha
from tests._pg_fixture import blank_session, unique_code

BASE = "/api/v1/reports/sales_yearly"
SLUG = "sales.reports.view"


@pytest.fixture
def db():
    with blank_session() as s:
        set_company_scope(s, frozenset({DEFAULT_COMPANY_ID}))
        seed_mocha(s)
        yield s


def _user(db, *companies):
    user = User(id=str(uuid.uuid4()), email=f"{unique_code('U')}@t.test", name="Sales Manager", status="ACTIVE")
    db.add(user)
    db.flush()
    for company_id in companies:
        db.add(UserCompany(user_id=user.id, company_id=company_id))
    db.flush()
    return {"id": user.id, "email": user.email}


def _sale(db, company_id, amount, demand_class="retail", when=date(2026, 3, 3)):
    so = SalesOrder(id=str(uuid.uuid4()), so_number=unique_code("SO"), order_date=when,
                    status="open", demand_class=demand_class, company_id=company_id)
    db.add(so)
    db.flush()
    db.add(SalesOrderLine(id=str(uuid.uuid4()), sales_order_id=so.id,
                          product_id=product(db, company_id=company_id).id,
                          qty_ordered=1, qty_delivered=1, line_total=Decimal(amount),
                          line_status="open", company_id=company_id))
    db.flush()


@contextmanager
def _client(db, principal, allow=(SLUG,)):
    from app.services.user_service import UserPermissionService

    def _override_db():
        yield db

    async def _scope():
        scope = frozenset({DEFAULT_COMPANY_ID})
        set_company_scope(db, scope)
        return scope

    app.dependency_overrides[get_db] = _override_db
    app.dependency_overrides[get_current_user] = lambda: principal
    app.dependency_overrides[get_current_user_or_api_key] = lambda: principal
    app.dependency_overrides[apply_company_scope] = _scope
    original = UserPermissionService.check_user_has_permission
    UserPermissionService.check_user_has_permission = lambda self, uid, slug: slug in allow
    try:
        yield TestClient(app)
    finally:
        UserPermissionService.check_user_has_permission = original
        app.dependency_overrides.clear()


def _body(company=DEFAULT_COMPANY_ID, **over):
    params = {
        "date_basis": "order_date",
        "period": {"kind": "custom", "from": "2024-01-01", "to": "2026-09-26"},
        "company": [company],
        "channel": ["dealer", "project"],
        "basis": ["delivered"],
    }
    params.update(over)
    return {"params": params, "view": None}


def test_ac_r2_2_catalog_lists_the_yearly_comparison_to_holders_only(db):
    principal = _user(db, DEFAULT_COMPANY_ID)
    with _client(db, principal) as client:
        keys = [r["key"] for r in client.get("/api/v1/reports").json()["reports"]]
        assert "sales_yearly" in keys
    with _client(db, principal, allow=()) as client:
        keys = [r["key"] for r in client.get("/api/v1/reports").json()["reports"]]
        assert "sales_yearly" not in keys
        assert client.get(BASE).status_code == 403


def test_meta_offers_only_the_callers_companies_by_name_and_defaults_to_the_current_one(db):
    principal = _user(db, DEFAULT_COMPANY_ID)
    with _client(db, principal) as client:
        meta = client.get(BASE).json()
    company = next(p for p in meta["params"] if p["key"] == "company")
    assert company["options"] == [{"value": DEFAULT_COMPANY_ID, "label": "Sorento"}]
    assert company["multi"] is False and company["clearable"] is False
    assert meta["default_view"]["params"]["company"] == [DEFAULT_COMPANY_ID]
    assert meta["default_view"]["params"]["basis"] == ["delivered"]
    assert meta["opens_on"] == "summary"
    assert meta["title"] == "Yearly comparison"


def test_ac_s1_5_run_for_a_company_outside_the_grant_is_403(db):
    _sale(db, MOCHA_ID, "500.00")
    principal = _user(db, DEFAULT_COMPANY_ID)
    with _client(db, principal) as client:
        response = client.post(f"{BASE}/run", json=_body(company=MOCHA_ID))
    assert response.status_code == 403, response.text


def test_a_user_granted_both_companies_reads_mocha(db):
    _sale(db, DEFAULT_COMPANY_ID, "10.00")
    _sale(db, MOCHA_ID, "500.00")
    principal = _user(db, DEFAULT_COMPANY_ID, MOCHA_ID)
    with _client(db, principal) as client:
        response = client.post(f"{BASE}/run", json=_body(company=MOCHA_ID))
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["layouts"]["summary"]["grand_total"]["sales_value"] == "500.00"
    assert [b["title"] for b in body["layouts"]["blocks"]] == ["MOCHA - DEALER", "MOCHA - PROJECT TEAM"]
    assert body["note"].startswith("Basis: Delivered")
    summary = body["layouts"]["summary"]
    for field in ("variance_row", "variance_total", "variance_label", "chart", "whole_units"):
        assert field in summary, field


def test_ac_r2_9_mocha_with_no_orders_is_the_empty_state_not_an_error(db):
    principal = _user(db, DEFAULT_COMPANY_ID, MOCHA_ID)
    with _client(db, principal) as client:
        response = client.post(f"{BASE}/run", json=_body(company=MOCHA_ID))
    assert response.status_code == 200
    assert response.json()["row_count"] == 0


def _strict(monkeypatch, db, enabled):
    from app.config import settings
    from app.models.app_modules import AppModuleCatalog, TenantModule

    monkeypatch.setattr(settings, "module_guard_strict", True)
    for key in ("procurement", "sales"):
        if not db.query(AppModuleCatalog).filter_by(module_key=key).first():
            db.add(AppModuleCatalog(module_key=key, display_name=key, dependencies=[]))
        db.flush()
        db.add(TenantModule(tenant_id="__default__", module_key=key, enabled=key in enabled))
    db.flush()


def test_ac_r2_3_sales_disabled_blocks_the_sales_report_and_not_sponsorship(db, monkeypatch):
    _strict(monkeypatch, db, enabled={"procurement"})
    principal = _user(db, DEFAULT_COMPANY_ID)
    allow = (SLUG, "procurement.sponsorship_forms.report")
    with _client(db, principal, allow=allow) as client:
        assert client.get(BASE).status_code == 403
        assert client.post(f"{BASE}/run", json=_body()).status_code == 403
        assert client.get("/api/v1/reports/sponsorship").status_code == 200


def test_ac_r2_3_procurement_disabled_no_longer_blocks_the_sales_report(db, monkeypatch):
    _strict(monkeypatch, db, enabled={"sales"})
    principal = _user(db, DEFAULT_COMPANY_ID)
    allow = (SLUG, "procurement.sponsorship_forms.report")
    with _client(db, principal, allow=allow) as client:
        assert client.get(BASE).status_code == 200
        assert client.get("/api/v1/reports/sponsorship").status_code == 403


def test_ac_r2_3_a_definition_with_no_module_key_is_403(db):
    from app.services.reports import registry as reg
    from tests import _report_fixture as fixture

    fixture.create_table(db)
    definition = fixture.definition()
    assert definition.module_key is None
    reg.register(definition)
    principal = _user(db, DEFAULT_COMPANY_ID)
    try:
        with _client(db, principal, allow=("zzt.reports.orders",)) as client:
            assert client.get("/api/v1/reports/zzt_orders").status_code == 403
    finally:
        reg._REGISTRY.pop("zzt_orders", None)


def test_ac_r2_8_ac_r4_6_export_lands_in_my_downloads_with_the_company_snapshot(db, monkeypatch):
    from app.models.download import UserDownload
    from app.services import queue_service

    queued = []
    monkeypatch.setattr(queue_service, "enqueue_job", lambda fn, *a, **kw: queued.append((fn, a, kw)))
    principal = _user(db, DEFAULT_COMPANY_ID)
    with _client(db, principal) as client:
        response = client.post(f"{BASE}/export", json=_body())
    assert response.status_code == 200, response.text
    filename = response.json()["filename"]
    assert filename == "Yearly comparison-01-01-2024 TO 26-09-2026.xlsx"
    row = db.query(UserDownload).filter_by(id=response.json()["download_id"]).one()
    assert row.kind == "report_xlsx" and row.user_id == principal["id"]
    (fn, args, kwargs), = queued
    assert fn.__name__ == "generate_report_xlsx"
    assert kwargs["company_grants"] == [DEFAULT_COMPANY_ID]


class _FakeBackend:
    def __init__(self):
        self.uploads = []

    def upload_file(self, file_content, file_path, content_type=None, **_kw):
        self.uploads.append((file_path, file_content))
        return file_path, f"https://cdn.test/{file_path}"


def _run_task(db, download_id, grants, **kw):
    from unittest.mock import patch

    from app.tasks import report_export_tasks as tasks

    backend = _FakeBackend()
    with patch.object(tasks, "SessionLocal", lambda: db), patch.object(
        tasks, "get_backend", lambda _p: backend
    ), patch.object(tasks, "default_provider", lambda: "s3"), patch.object(db, "close", lambda: None):
        out = tasks.generate_report_xlsx(
            download_id, "sales_yearly", _body()["params"], None, "u", company_grants=grants, **kw
        )
    return out, backend


def test_ac_r2_4_the_export_job_runs_under_the_enqueuers_company(db):
    from openpyxl import load_workbook
    import io

    from app.services.download_service import DownloadService

    _sale(db, DEFAULT_COMPANY_ID, "10.00")
    _sale(db, MOCHA_ID, "500.00")
    principal = _user(db, DEFAULT_COMPANY_ID)
    download = DownloadService(db).create(user_id=principal["id"], kind="report_xlsx", filename="y.xlsx")
    out, backend = _run_task(db, str(download.id), [DEFAULT_COMPANY_ID])
    assert out["status"] == "ready", out
    book = load_workbook(io.BytesIO(backend.uploads[0][1]))
    numbers = [c.value for r in book["SUMMARY"].iter_rows() for c in r
               if isinstance(c.value, (int, float, Decimal))]
    assert Decimal("500.00") not in [Decimal(str(n)) for n in numbers]
    assert Decimal("10.00") in [Decimal(str(n)) for n in numbers]


def test_ac_r2_4_an_export_job_with_no_company_snapshot_writes_nothing_from_any_company(db):
    import io

    from openpyxl import load_workbook

    from app.services.download_service import DownloadService

    _sale(db, DEFAULT_COMPANY_ID, "10.00")
    principal = _user(db, DEFAULT_COMPANY_ID)
    download = DownloadService(db).create(user_id=principal["id"], kind="report_xlsx", filename="y.xlsx")
    from app.tasks import report_export_tasks as tasks

    out, backend = _run_task(db, str(download.id), tasks.NO_COMPANY_SNAPSHOT)
    if out["status"] == "ready":
        book = load_workbook(io.BytesIO(backend.uploads[0][1]))
        numbers = [c.value for r in book["SUMMARY"].iter_rows() for c in r
                   if isinstance(c.value, (int, float, Decimal))]
        assert Decimal("10.00") not in [Decimal(str(n)) for n in numbers]


def test_a_shared_default_saved_on_another_company_never_hands_over_its_id(db):
    """Security nit: a published default view carrying Mocha's id opens a Sorento-only
    user on Sorento, and the Mocha id never reaches their screen."""
    from app.schemas.report import ReportViewConfig
    from app.services.reports.views_service import ReportViewsService

    admin = _user(db, DEFAULT_COMPANY_ID, MOCHA_ID)
    view = ReportViewConfig.model_validate({
        "params": {"company": [MOCHA_ID], "channel": ["dealer"], "basis": ["delivered"],
                   "date_basis": "order_date"},
        "detail": {"columns": [], "order": []},
        "pivot": {"rows": "year", "cols": "month_of_year", "measures": ["sales_value"]},
    })
    svc = ReportViewsService(db)
    saved = svc.create("sales_yearly", admin["id"], "Mocha default", view)
    svc.publish("sales_yearly", saved.id, admin["id"], True)
    svc.set_default("sales_yearly", saved.id, admin["id"])
    principal = _user(db, DEFAULT_COMPANY_ID)
    with _client(db, principal) as client:
        meta = client.get(BASE).json()
    assert meta["default_view"]["params"]["company"] == [DEFAULT_COMPANY_ID]
    assert MOCHA_ID not in str(meta)
