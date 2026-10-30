"""SO-LINE-DESC: the OI handover ITEM CODE carries the SO line description for a ticked product
(AC-10..AC-15). Plan `documentation/plans/order-management/so-line-desc-acceptance-criteria.md`.

TEST-FIRST. Driven through the real flows (Send to purchasing preview and send, and the board's
Confirm), never through a stubbed line: the one string the email body, the plain text, the
preview and the Excel print is decided upstream of all four.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

import pytest
import sqlalchemy as sa

from app.models.project_so import INQUIRY_RAISED

from .test_email_layout_migrations import _run
from .test_oi_handover_email_excel_parity import _sheet, sent  # noqa: F401 - fixture
from .test_oi_handover_resend import (  # noqa: F401 - fixtures
    PREVIEW,
    SEND,
    _automation,
    _notifications,
    _ok,
    rs,
    stored,
)
from .test_oi_product_follow_apply import _apply, _swap_fixture
from .test_oi_saved_line_email import _migration, _order, _post, _rows
from .test_order_inquiry_handover_automation import (
    _captured_dispatches,
    _handover_calls,
    _register,
)
from .test_order_inquiry_worklist import _inquiry_for, _row, _uid
from .test_planning_change_apply_on_board import api, world  # noqa: F401 - fixtures
from .test_planning_changes import _core_line, _core_so, _line_payload, _project_line
from .test_so_supply_confirmation import _product, _warehouse

__all__ = ["api", "world", "rs", "stored", "sent"]

PUSH = "SORENTO ACCESSORIES SRTWC8320-S150-UF PUSH BUTTON ONLY"
TAP = "CABANA ACCESSORIES - TAP HEAD ONLY FOR MODEL KT886SS-NL"
RPACC_TEXT = f"RPACC - {PUSH}"

#: (product code, ticked, SO line description)
RPACC = ("RPACC", True, PUSH)
RPAC = ("RPAC", False, TAP)
EMPTY = ("RPJC", True, None)
SAME = ("RPBK", True, "RPBK")

_BASE_CHAIN = (
    "oihr_0004_wide_line_table",
    "oihr_0005_print_one_page",
    "oihr_0006_no_remark_column",
    "oihr_0007_excel_remark",
    "oihr_0008_so_table_on_top",
    "oihr_0009_split_rows",
    "oihr_0010_so_remark_column",
)


def _layout_chain() -> list[str]:
    """Every handover layout migration a real database has run, including any this lane adds
    after `oihr_0010` (the ITEM CODE wrap rule lives in the stored template)."""
    later = sorted(
        p.stem for p in Path("alembic/versions").glob("oihr_*.py") if p.stem > _BASE_CHAIN[-1]
    )
    return [*_BASE_CHAIN, *later]


def _automation_at_head(world):
    buyer = _automation(world)
    for name in _layout_chain():
        _run(_migration(name), world.db)
    world.db.commit()
    return buyer


def _tick(db, product_id, value=True):
    db.execute(
        sa.text("UPDATE products SET show_so_line_description = :v WHERE id = :id"),
        {"v": value, "id": product_id},
    )


def _describe(db, core_line_id, text):
    db.execute(
        sa.text("UPDATE sales_order_lines SET description = :d WHERE id = :id"),
        {"d": text, "id": core_line_id},
    )


def _oi_world(world, specs):
    """One SO with a core line per spec, its project mirror and one open OI row per line whose
    ITEM CODE is the line's product code."""
    from app.models.project_so import ProjectSalesOrder

    db = world.db
    wh = _warehouse(db, f"ZZT-{_uid()[:6]}".upper())
    core_so = _core_so(db, world.company_id)
    order = ProjectSalesOrder(
        id=_uid(), company_id=world.company_id, provisional_ref=f"ZZT-{_uid()[:8]}",
        so_id=core_so.id, autocount_doc_no=core_so.so_number, status="adopted",
    )
    db.add(order)
    db.flush()
    inquiry = _inquiry_for(db, world.company_id, order)
    inquiry.raised_by = world.requester
    for i, (code, ticked, description) in enumerate(specs):
        product = _product(db)
        product.product_code = code
        db.flush()
        core = _core_line(db, core_so, product, wh, qty_ordered="5")
        mirror = _project_line(db, order, line_no=10 * (i + 1), product=product, core_line=core)
        _tick(db, product.id, ticked)
        _describe(db, core.id, description)
        _row(db, world.company_id, inquiry, state=INQUIRY_RAISED, qty="5", item_code=code,
             stock_location="BRW-BB", so_line_id=mirror.id, delivery_date=None)
    db.commit()
    return inquiry


def _authored(html):
    if "<!--block:custom_text-->" in html:
        from tests._email_body import authored_html

        return authored_html(html)
    return html


def _email_item_cells(html) -> list[str]:
    """The ITEM CODE value of every line row of the email's line table."""
    for index, row in enumerate(rows := _rows(html)):
        if "ITEM CODE" in row:
            col, width = row.index("ITEM CODE"), len(row)
            return [r[col] for r in rows[index + 1:] if len(r) == width and r[col]]
    raise AssertionError(f"no ITEM CODE heading in {rows}")


def _excel_item_cells(book_bytes) -> list[str]:
    import openpyxl

    sheet = openpyxl.load_workbook(io.BytesIO(book_bytes)).active
    rows = list(sheet.iter_rows())
    head = next(r for r in rows if any(c.value == "ITEM CODE" for c in r))
    col = [c.value for c in head].index("ITEM CODE")
    return [r[col].value for r in rows[rows.index(head) + 1:] if r[col].value]


def _td_for(html, text):
    """The opening `<td ...>` of the cell whose content starts with `text`."""
    found = re.search(r"(<td[^>]*>)\s*(?:<[^>]+>\s*)*" + re.escape(text) + r"\s*<", _authored(html))
    assert found, f"no cell for {text!r}"
    return found.group(1)


def _preview_html(client, inquiry):
    body = _ok(client.post(PREVIEW, json={"inquiry_ids": [inquiry.id]}))
    assert len(body["emails"]) == 1, body
    return body["emails"][0]["body_html"]


# ------------------------------------------------------------------ AC-10, AC-11, AC-12


def test_ac10_ac11_preview_decorates_only_a_ticked_product_with_a_real_description(rs):
    client, world = rs
    _automation_at_head(world)
    inquiry = _oi_world(world, [RPACC, RPAC, EMPTY, SAME])

    html = _preview_html(client, inquiry)

    assert sorted(_email_item_cells(html)) == sorted([RPACC_TEXT, "RPAC", "RPJC", "RPBK"])
    assert "CABANA ACCESSORIES" not in html, "an unticked product never prints its description"
    assert not re.search(r"RPJC\s*-\s*<", html) and "RPBK - " not in html


def test_ac12_email_text_part_preview_and_excel_print_the_same_item_code(rs, stored):
    client, world = rs
    _automation_at_head(world)
    inquiry = _oi_world(world, [RPACC, RPAC, EMPTY, SAME])
    preview_cells = sorted(_email_item_cells(_preview_html(client, inquiry)))

    _ok(client.post(SEND, json={"inquiry_ids": [inquiry.id]}))

    [note] = _notifications(world, inquiry.id)
    body_cells = sorted(_email_item_cells(note.data["body_html"]))
    assert body_cells == preview_cells == sorted([RPACC_TEXT, "RPAC", "RPJC", "RPBK"])
    assert RPACC_TEXT in note.data["body_text"]
    assert "CABANA ACCESSORIES" not in note.data["body_text"]
    [upload] = stored
    assert sorted(_excel_item_cells(upload["bytes"])) == body_cells


def test_ac15_the_resend_sent_hint_still_finds_a_decorated_line_in_the_earlier_body(rs):
    """`oi_handover_resend._lines` matches `row.item_code in body_html`: the code stays the
    cell's prefix, so a line sent decorated is still marked sent."""
    client, world = rs
    _automation_at_head(world)
    inquiry = _oi_world(world, [RPACC])
    _ok(client.post(SEND, json={"inquiry_ids": [inquiry.id]}))

    body = _ok(client.post(PREVIEW, json={"inquiry_ids": [inquiry.id]}))

    assert [line["sent_at"] is not None for line in body["emails"][0]["lines"]] == [True]
    assert RPACC_TEXT in _email_item_cells(body["emails"][0]["body_html"])


# ------------------------------------------------------------------ AC-14


def test_ac14_a_decorated_item_code_cell_wraps_a_bare_one_keeps_nowrap(rs):
    client, world = rs
    _automation_at_head(world)
    inquiry = _oi_world(world, [RPACC, EMPTY])

    html = _preview_html(client, inquiry)

    assert "nowrap" not in _td_for(html, RPACC_TEXT)
    assert "nowrap" in _td_for(html, "RPJC")


def test_ac14_the_excel_decorated_cell_wraps_and_the_column_width_is_unchanged(rs, stored):
    import openpyxl

    client, world = rs
    _automation_at_head(world)
    inquiry = _oi_world(world, [RPACC, EMPTY])

    _ok(client.post(SEND, json={"inquiry_ids": [inquiry.id]}))

    [upload] = stored
    sheet = openpyxl.load_workbook(io.BytesIO(upload["bytes"])).active
    rows = list(sheet.iter_rows())
    head = next(r for r in rows if any(c.value == "ITEM CODE" for c in r))
    col = [c.value for c in head].index("ITEM CODE")
    cell = next(r[col] for r in rows if r[col].value == RPACC_TEXT)
    assert cell.alignment.wrap_text is True
    assert sheet.column_dimensions[head[col].column_letter].width == 18


# ------------------------------------------------------------------ Confirm path (board)


def test_ac10_ac12_confirm_email_lines_and_excel_carry_the_same_decorated_code(api, sent):
    client, world = api
    contexts, uploads = sent
    order, line_a, line_b = _order(api, qty_a="10", qty_b="6")
    _tick(world.db, world.product.id)
    _describe(world.db, line_a.core_sales_order_line_id, PUSH)
    world.db.commit()
    code = world.product.product_code

    _post(api, order, [
        _line_payload(line_a.id, buy_qty="10"),
        _line_payload(line_b.id, buy_qty="6"),
    ])

    [context] = contexts
    printed = sorted(line["item_code"] for line in context["handover"]["lines"])
    assert printed == sorted([f"{code} - {PUSH}", code])
    _headings, body = _sheet(context, uploads)
    assert sorted(row["ITEM CODE"] for row in body) == printed


def test_ac13_an_unchanged_line_never_shows_new_item_code(api, sent):
    contexts, _uploads = sent
    order, line_a, line_b = _order(api, qty_a="10", qty_b="6")
    _tick(api[1].db, api[1].product.id)
    _describe(api[1].db, line_a.core_sales_order_line_id, PUSH)
    api[1].db.commit()

    _post(api, order, [_line_payload(line_a.id, buy_qty="10")])

    [context] = contexts
    for line in context["handover"]["lines"]:
        assert not (line.get("was") or {}).get("item_code")
        assert "NEW ITEM CODE" not in line["remark"]
    assert [line["item_code"] for line in context["handover"]["lines"]] == [
        f"{api[1].product.product_code} - {PUSH}"
    ]


# ------------------------------------------------------------------ AC-13 code switch


@pytest.mark.parametrize(
    "old_ticked,new_ticked",
    [(True, False), (False, True), (True, True), (False, False)],
)
def test_ac13_each_side_of_a_code_switch_follows_its_own_products_tick(
    api, monkeypatch, old_ticked, new_ticked
):
    _client, world = api
    _register(world)
    calls = _captured_dispatches(monkeypatch)
    fx = _swap_fixture(api, acknowledged=True)
    db = fx["world"].db
    _tick(db, fx["old"].id, old_ticked)
    _tick(db, fx["new"].id, new_ticked)
    _describe(db, fx["line"].core_sales_order_line_id, PUSH)
    db.commit()
    calls.clear()

    response = _apply(fx)
    assert response.status_code == 200, response.text
    db.commit()

    lines = _handover_calls(calls)[-1]["context"]["handover"]["lines"]
    [line] = [entry for entry in lines if (entry.get("was") or {}).get("item_code")]
    old, new = fx["old"].product_code, fx["new"].product_code
    assert line["was"]["item_code"] == (f"{old} - {PUSH}" if old_ticked else old)
    assert line["item_code"] == (f"{new} - {PUSH}" if new_ticked else new)
