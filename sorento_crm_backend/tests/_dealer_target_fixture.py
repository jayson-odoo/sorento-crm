"""A generated workbook covering every pattern the crew found in the real dealer target files
(lane SALES-TARGET-DEALER, UAC D3-2). No real data: names and figures are invented, and the
real files never enter the repo.

Patterns, each named where it is written:
- P1 merged multi-account threshold (one target over A/C I + A/C III)
- P2 per-account rate named in the cell ("A/C I 3%"), only the named account earns
- P3 Q1 "+ Additional x%" with its own threshold, two additional lines stacking
- P4 rate per account Deluxe / Mastile / W-Ply on one shared threshold
- P5 yearly 2 to 3 tiers
- P6 Dec'25 to Nov'26 window stated in the cell
- P7 monthly targets (3 dealers), threshold / % sub-column pairs
- P8 Mocha points: quarterly figure with the 4% in the header, yearly 2 and 3 tier schemes
- P9 mixed number / text cells ("RM 300,000", 300000, "300k")
- P10 rows that must be reported, not mapped: unknown code, unreadable cell
- P11 name in the file differs from the CRM name (warning only)
"""
from __future__ import annotations

import io
from typing import Dict

from openpyxl import Workbook

#: Logical keys -> the customer codes the caller seeded (or literal codes in a pure test).
KEYS = (
    "A1", "A3",            # dealer A, P1/P2/P3
    "B_DELUXE", "B_MASTILE", "B_WPLY",  # dealer B, P4
    "C1",                  # dealer C, P6
    "D1",                  # dealer D, P11 (name differs)
    "E1",                  # dealer E, P10 unreadable cell
    "M1", "M2", "M3",      # monthly dealers, P7
    "K1", "K2", "K3",      # Mocha dealer K: A/C I, [CERAMIC], [IBORN], P8 merged
    "L1",                  # Mocha dealer L, 2 tier yearly
)
UNKNOWN_CODE = "ZZT-NO-SUCH"


def build(codes: Dict[str, str]) -> bytes:
    wb = Workbook()
    _sorento(wb.active, codes)
    _monthly(wb.create_sheet("Monthly"), codes)
    _mocha(wb.create_sheet("Mocha"), codes)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sorento(ws, c):
    ws.title = "Sorento 2026"
    ws["A1"] = "SORENTO DEALER TARGET 2026"
    header = ["CODE", "DEALER NAME", "ACCOUNT", "Q1", "Q2", "Q3", "Q4", "JAN-DEC'26", "PAY TO"]
    for col, text in enumerate(header, start=1):
        ws.cell(row=3, column=col, value=text)
    rows = [
        # P1 + P2 + P3: dealer A, A/C I and A/C III share every threshold (merged rows 4:5).
        [c["A1"], "Dealer A Sdn Bhd", "A/C I",
         "150,000 - 2%\n200,000 - 3%\n+ Additional 1% @ 220,000\n+ Additional 0.5% @ 250,000",
         "200,000 - A/C I 2%\n250,000 - A/C I 3%",
         "RM 240,000 (2.5%)",            # P9 text with RM
         "260k @ 3%",                     # P9 "k"
         "850,000 - 1%\n1,000,000 - 1.5%\n1,200,000 - 2%",  # P5 three tiers
         c["A1"]],
        [c["A3"], "Dealer A Sdn Bhd", "A/C III", None, None, None, None, None, None],
        # P4: dealer B, one threshold, a rate per account (merged rows 6:8).
        [c["B_DELUXE"], "Dealer B Trading", "Deluxe",
         "300,000 - Deluxe 2% / Mastile 3% / W-Ply 1.5%", "-", "-", "-", None, None],
        [c["B_MASTILE"], "Dealer B Trading", "Mastile", None, None, None, None, None, None],
        [c["B_WPLY"], "Dealer B Trading", "W-Ply", None, None, None, None, None, None],
        # P6: dealer C, yearly only, Dec'25 to Nov'26 stated in the cell.
        [c["C1"], "Dealer C Enterprise", "A/C I", "-", "-", "-", "-",
         "DEC'25-NOV'26: 1,200,000 - 1%\n1,500,000 - 1.25%", None],
        # P11: name differs from the CRM; only Q2 has a scheme.
        [c["D1"], "Dealer D Holdings (old name)", "A/C I", None, "300,000 - 2%", None, None, None, None],
        # P10: unknown code; and an unreadable cell.
        [UNKNOWN_CODE, "Dealer Z", "A/C I", "100,000 - 1%", None, None, None, None, None],
        [c["E1"], "Dealer E", "A/C I", "see note", None, None, None, None, None],
    ]
    for r, values in enumerate(rows, start=4):
        for col, value in enumerate(values, start=1):
            if value is not None:
                ws.cell(row=r, column=col, value=value)
    for col in "DEFGHI":
        ws.merge_cells(f"{col}4:{col}5")
    for col in "DEFG":
        ws.merge_cells(f"{col}6:{col}8")


def _monthly(ws, c):
    # P7: month headers merged over a threshold / % pair, sub-headers on row 4.
    ws["A3"], ws["B3"], ws["C3"] = "CODE", "NAME", "ACCOUNT"
    months = ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"]
    col = 4
    for month in months:
        ws.cell(row=3, column=col, value=month)
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + 1)
        ws.cell(row=4, column=col, value="TARGET")
        ws.cell(row=4, column=col + 1, value="%")
        col += 2
    for r, key in enumerate(("M1", "M2", "M3"), start=5):
        ws.cell(row=r, column=1, value=c[key])
        ws.cell(row=r, column=2, value=f"Monthly dealer {key}")
        ws.cell(row=r, column=3, value="A/C I")
        col = 4
        for i in range(12):
            # P9: numbers and text mixed across the months.
            ws.cell(row=r, column=col, value=50000 + 1000 * i if i % 2 == 0 else f"{50 + i},000")
            ws.cell(row=r, column=col + 1, value=1 if i % 3 else "1%")
            col += 2


def _mocha(ws, c):
    # P8: points, 4% in the header for the quarters; yearly tiers in the cell.
    header = ["CODE", "CUSTOMER", "A/C", "Q1 POINT (4%)", "Q2 POINT (4%)", "Q3 POINT (4%)",
              "Q4 POINT (4%)", "YEARLY POINT"]
    for col, text in enumerate(header, start=1):
        ws.cell(row=2, column=col, value=text)
    rows = [
        [c["K1"], "Mocha K", "A/C I", 100000, "100,000", 100000, 120000,
         "450,000 - 1%\n500,000 - 1.5%\n550,000 - 2%"],
        [c["K2"], "Mocha K [CERAMIC]", "[CERAMIC]", None, None, None, None, None],
        [c["K3"], "Mocha K [IBORN]", "[IBORN]", None, None, None, None, None],
        [c["L1"], "Mocha L", "A/C I", 50000, 50000, 50000, 50000, "200,000 - 1%\n250,000 - 2%"],
    ]
    for r, values in enumerate(rows, start=3):
        for col, value in enumerate(values, start=1):
            if value is not None:
                ws.cell(row=r, column=col, value=value)
    for col in "DEFGH":
        ws.merge_cells(f"{col}3:{col}5")
