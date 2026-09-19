"""RED tests for the Phase 3 fix round of the AutoCount pull + review lane
(issues #1045 #1046 #1047 #1048) - the reviewer's and security-reviewer's findings against
`4bbc3a431`.

Plan: documentation/plans/autocount/PLAN-autocount-pull-review.md
UAC:  documentation/plans/autocount/autocount-pull-review-acceptance-criteria.md

Every test below traces to one bullet of the captain's fix-round brief (F-1..F-12), never
new scope. Substrate is reused BY IMPORT from `tests/test_autocount_pull_sr1.py`
(`tests/test_autocount_pull_sr3.py` module as `sr3`, `tests/test_autocount_pull_sr4.py`
module as `sr4`) - the same technique SR3/SR4 already use to share `env`/`task_db` and their
own helpers: importing a name into this module's globals is enough for pytest to discover a
fixture, and a plain function is just as importable.

Postgres only; every test seeds its own data chain. A missing name (a constant, a kwarg, a
route behaviour the coder has not built yet) is imported/referenced INSIDE a test body so a
gap reds exactly one test, never collection.
"""
from __future__ import annotations

import io
import uuid
from unittest.mock import MagicMock

import httpx
import openpyxl
import pytest
from sqlalchemy import text

# MUST be the first app import - resolves the circular import in
# app.modules.runtime.guards (repo convention, see test_ingest_deletions.py).
from app.main import app  # noqa: E402,F401

from app.services.company_scope import DEFAULT_COMPANY_ID

import tests.test_autocount_pull_sr3 as sr3
import tests.test_autocount_pull_sr4 as sr4
from tests.test_autocount_pull_sr1 import (  # noqa: F401 - env/task_db are fixtures
    API_KEY,
    BASE_URL,
    PULLS_URL,
    _FakeFoundryX,
    _canonical_row,
    _header,
    _job_row,
    _job_rows,
    _patch_foundryx,
    _run_preview,
    _seed_pull_job,
    _stored,
    env,
    task_db,
)
from tests.test_autocount_pull_sr3 import _run_apply, _seed_apply_job  # noqa: F401
from tests.test_stock_list_xlsm_upload import client  # noqa: F401 - `client` is a fixture

MARKER = "ZZTAPFIX"


def _seed_permitted_user(db, *perm_slugs: str) -> str:
    """A user holding EXACTLY the given permission slugs - `_Env.user()`'s own recipe
    (`tests/test_autocount_pull_sr1.py`), reused here for the one test that drives a route
    through a hand-wired `TestClient` bound to `task_db`'s own connection rather than
    `env`'s independent `blank_session()`."""
    from app.models.user import User, UserPermission, UserRole, UserRoleAssignment, UserRolePermission

    uid = str(uuid.uuid4())
    role_id = str(uuid.uuid4())
    db.add(UserRole(
        id=role_id, slug=f"{MARKER.lower()}-role-{uid[:8]}", name=f"{MARKER} role {uid[:8]}",
        description="", is_protected=False, is_default=False,
    ))
    db.add(User(id=uid, email=f"{MARKER.lower()}-{uid[:8]}@test.com", name="U", status="ACTIVE"))
    db.flush()
    db.add(UserRoleAssignment(user_id=uid, role_id=role_id))
    for slug in perm_slugs:
        perm = db.query(UserPermission).filter_by(slug=slug).one_or_none()
        if perm is None:
            perm = UserPermission(id=str(uuid.uuid4()), slug=slug, name=slug, description="")
            db.add(perm)
            db.flush()
        db.add(UserRolePermission(id=str(uuid.uuid4()), role_id=role_id, permission_id=perm.id))
    db.commit()
    return uid


# ============================================================================ F-1


class TestF1EmptyFedGuard:
    """HIGH: an empty FED batch (nothing matched an active warehouse) must never reach
    `StockService.bulk_import_stock` for real - that call's own "active, absent from FED ->
    zeroed" sweep is company-wide, so an empty batch wipes every active warehouse's stock,
    not just the row this pull was ever about."""

    def test_f1a_preview_with_zero_fed_rows_blocks_confirm(self, task_db, monkeypatch):
        """a. No pulled row matches an active warehouse -> the preview still finishes in
        review, `confirm_blocked_reason` is set, `counts.fed == 0`, and Confirm itself then
        answers 409 - driven through the REAL preview task (`_run_preview`, task_db's own
        SessionLocal patch - `ImportOutcome.flush()` opens its own session and a plain
        `env`-style blank_session cannot see it, exactly the reason SR1's `task_db` fixture
        exists) and the real HTTP route (a hand-wired `TestClient` bound to the SAME
        connection, since `env`'s own blank_session is a second, independent transaction)."""
        from fastapi.testclient import TestClient

        from app.dependencies import get_current_user, get_current_user_or_api_key, get_db
        from app.models.base import set_company_scope
        from app.services.company_scope_resolver import apply_company_scope

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)

        owner_id = _seed_permitted_user(db, "inventory.stock.autocount_pull")
        rows = [sr4._stock_row(f"{MARKER}-F1A", f"{MARKER}-F1A-GHOST", 5)]  # no warehouse at all
        job_id = sr4._prepare_stock_preview(db, fake, rows=rows, owner=owner_id)

        _run_preview(monkeypatch, factory, job_id)

        row_after = _job_row(db, job_id)
        assert row_after["status"] == "finished", row_after["error"]
        pull = row_after["metadata"]["autocount_pull"]
        assert pull["phase"] == "review"
        assert pull["counts"]["fed"] == 0, pull["counts"]
        reason = pull.get("confirm_blocked_reason")
        assert isinstance(reason, str) and reason.strip(), pull

        def _override_get_db():
            yield db

        async def _override_scope():
            scope = frozenset({str(row_after["company_id"])})
            set_company_scope(db, scope)
            return scope

        app.dependency_overrides[get_db] = _override_get_db
        app.dependency_overrides[get_current_user] = lambda: {"id": owner_id}
        app.dependency_overrides[get_current_user_or_api_key] = lambda: {"id": owner_id}
        app.dependency_overrides[apply_company_scope] = _override_scope
        try:
            with TestClient(app) as c:
                resp = c.post(f"{PULLS_URL}/{job_id}/confirm")
        finally:
            app.dependency_overrides.clear()

        assert resp.status_code == 409, resp.text

    def test_f1b_apply_refuses_when_no_row_matches_an_active_warehouse(
        self, task_db, monkeypatch
    ):
        """b. The preview had FED rows; every warehouse those rows matched is deactivated
        before Apply runs -> the apply job fails, names the guard, touches NOT ONE stock
        row anywhere - including a product on an unrelated, still-active warehouse the pull
        never even mentioned, which is exactly what the unguarded zero-sweep would wipe."""
        from app.models.inventory import Warehouse
        import app.services.stock_list_archive_service as archive_mod

        db, factory = task_db
        fake = _FakeFoundryX()
        _patch_foundryx(monkeypatch, fake)

        matched_wh = Warehouse(
            id=str(uuid.uuid4()), warehouse_code=f"{MARKER}-F1B-MATCHED",
            warehouse_name="Matched", is_active=True, company_id=DEFAULT_COMPANY_ID,
        )
        bystander_wh = Warehouse(
            id=str(uuid.uuid4()), warehouse_code=f"{MARKER}-F1B-BYSTANDER",
            warehouse_name="Bystander", is_active=True, company_id=DEFAULT_COMPANY_ID,
        )
        db.add_all([matched_wh, bystander_wh])
        db.flush()

        sr4._seed_product_with_stock(
            db, DEFAULT_COMPANY_ID, matched_wh.id, code=f"{MARKER}-F1B-A", qty=5,
        )
        sr4._seed_product_with_stock(
            db, DEFAULT_COMPANY_ID, bystander_wh.id, code=f"{MARKER}-F1B-B", qty=7,
        )

        # The snapshot only ever named the matched pair - the bystander's stock was never
        # part of this pull's rows at all.
        rows = [sr4._stock_row(f"{MARKER}-F1B-A", matched_wh.warehouse_code, 9)]
        snapshot_id = f"{MARKER}-snap-f1b"
        header = sr4._stock_header(rows)
        fake.status = (200, {**header, "snapshotId": snapshot_id})
        fake.rows = {
            1: (200, {"snapshotId": snapshot_id, "page": 1, "pageSize": 1000,
                      "totalPages": 1, "recordCount": len(rows), "rows": rows}),
        }

        # Every warehouse the preview matched is deactivated before Apply runs.
        db.execute(
            text("UPDATE warehouses SET is_active = false WHERE id = :id"), {"id": matched_wh.id}
        )
        db.commit()

        archive_calls = []
        monkeypatch.setattr(
            archive_mod, "replace_latest_stock_list",
            lambda *a, **k: archive_calls.append(1),
        )

        job_id = _seed_apply_job(
            db, user_id=str(uuid.uuid4()), company_id=DEFAULT_COMPANY_ID, entity="stock_balances",
            snapshot_id=snapshot_id,
        )

        _run_apply(monkeypatch, factory, job_id)

        row_after = _job_row(db, job_id)
        assert row_after["status"] == "failed", row_after
        error = (row_after["error"] or "").lower()
        assert "no row matched an active warehouse" in error, error

        assert sr4._stock_qty(db, f"{MARKER}-F1B-A") == 5
        assert sr4._stock_qty(db, f"{MARKER}-F1B-B") == 7, (
            "an empty FED batch must never zero stock on an unrelated active warehouse"
        )
        assert archive_calls == [], "a refused apply must never reach the archive step"


# ============================================================================ F-2


class TestF2UnboundedPaging:
    def test_f2_all_rows_stops_after_max_pages_and_raises_row_limit(self, monkeypatch):
        """An upstream whose `totalPages` keeps growing must not be paged forever.
        `MAX_PAGES` is the coder's new module constant on `foundryx_autocount_client.py`,
        pinned here at 3 - `raising=False` so this test itself does not collection-error
        before the constant exists; the fake caps growth at page 10 so this run stays
        finite and fast even against TODAY'S unbounded loop."""
        from app.config import settings
        import app.services.foundryx_autocount_client as client_mod
        from app.services.foundryx_autocount_client import FoundryxAutocountClient, FoundryxPullError

        monkeypatch.setattr(settings, "foundryx_base_url", BASE_URL, raising=False)
        monkeypatch.setattr(settings, "foundryx_api_key", API_KEY, raising=False)

        snapshot_id = f"{MARKER}-snap-f2"
        calls: list[int] = []

        def _growing_handler(request: httpx.Request) -> httpx.Response:
            page = int(request.url.params.get("page", "1"))
            calls.append(page)
            total_pages = min(page + 1, 10)  # plateaus - keeps the fake finite either way
            body = {
                "snapshotId": snapshot_id, "page": page, "pageSize": 1,
                "totalPages": total_pages, "recordCount": total_pages,
                "rows": [_canonical_row(f"{MARKER}-P{page}")],
            }
            return httpx.Response(200, json=body)

        monkeypatch.setattr(
            client_mod, "TRANSPORT", httpx.MockTransport(_growing_handler), raising=False
        )
        monkeypatch.setattr(client_mod, "MAX_PAGES", 3, raising=False)

        client = FoundryxAutocountClient()
        with pytest.raises(FoundryxPullError) as exc_info:
            client.all_rows(snapshot_id)

        assert exc_info.value.code == "ROW_LIMIT", exc_info.value.code
        assert len(calls) <= 3, calls


# ============================================================================ F-3


class TestF3CompareBodyBounds:
    def test_f3a_filename_over_255_is_422_through_the_route(self, env):
        user = env.user("master_data.products.autocount_pull")
        env.as_user(user)
        job_id = str(uuid.uuid4())  # body validation must reject before the job is even looked up

        resp = env.client.post(
            f"{PULLS_URL}/{job_id}/compare",
            json={"filename": "x" * 256, "rows": []},
        )
        assert resp.status_code == 422, resp.text

    def test_f3b_rows_field_carries_a_200000_item_max_length(self):
        from app.api.v1.integrations.autocount_pull import ComparePostBody

        schema = ComparePostBody.model_json_schema()
        rows_schema = schema["properties"]["rows"]
        assert rows_schema.get("maxItems") == 200_000, rows_schema


# ============================================================================ F-4


class TestF4FormulaInjection:
    def test_f4a_products_download_stores_a_formula_leading_description_as_a_literal_string(
        self, env
    ):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        payload = '=HYPERLINK("http://x","y")'
        rows = [_canonical_row(f"{MARKER}-F4A", name=payload, description=payload)]
        job_id, _ = sr3._seed_review_job(env, owner=owner, rows=rows)

        resp = env.client.get(f"{PULLS_URL}/{job_id}/download.xlsx")
        assert resp.status_code == 200, resp.text

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        desc_col = header.index("Description")
        cell = next(ws.iter_rows(min_row=2, max_row=2))[desc_col]
        assert cell.data_type == "s", (cell.value, cell.data_type)
        assert cell.value == payload

    def test_f4b_stock_workbook_stores_a_formula_leading_item_description_as_a_literal_string(
        self,
    ):
        from app.services.autocount_pull_service import build_stock_workbook

        payload = "=1+1"
        template_rows = [
            {"Item Code": f"{MARKER}-F4B", "Item Description": payload, "Location": "L1",
             "On Hand Qty": 1},
        ]
        raw = build_stock_workbook(template_rows)

        wb = openpyxl.load_workbook(io.BytesIO(raw))
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        desc_col = header.index("Item Description")
        cell = next(ws.iter_rows(min_row=2, max_row=2))[desc_col]
        assert cell.data_type == "s", (cell.value, cell.data_type)
        assert cell.value == payload


# ============================================================================ F-5


class TestF5CompanyRecheck:
    def test_f5_owner_whose_current_scope_no_longer_covers_the_pull_gets_404(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner, scope=frozenset({env.company_a}))
        job_id, _ = sr3._seed_review_job(env, owner=owner)

        # The SAME owner, but the ambient company scope this request carries no longer
        # includes the pull's own company - a company switch, not a different user.
        env.as_user(owner, scope=frozenset({env.company_b}))

        responses = [
            env.get_pull(job_id),
            env.client.get(f"{PULLS_URL}/{job_id}/rows"),
            env.client.post(
                f"{PULLS_URL}/{job_id}/compare", json={"filename": "x.xlsx", "rows": []}
            ),
            env.client.post(f"{PULLS_URL}/{job_id}/confirm"),
        ]
        for resp in responses:
            assert resp.status_code == 404, resp.text
            assert resp.json().get("code") == "NOT_FOUND", resp.json()


# ============================================================================ F-6


class TestF6ExactlyOncePreview:
    def test_f6_claim_and_enqueue_preview_is_exactly_once_from_stale_state(
        self, task_db, monkeypatch
    ):
        """Two callers off the identical, stale pre-read `pull` dict - `_claim_and_enqueue_
        preview` is called twice with the SAME original state, the captain's own recipe.
        This depends entirely on the `ImportJob.status == PENDING` clause in the conditional
        UPDATE: with that clause removed the second call would also see rowcount==1 and
        enqueue a second time. Verified empirically before commit (see the tester report) by
        temporarily deleting that clause and re-running this one test, which then goes red;
        left in place, the test is a regression guard, not a red for a live bug."""
        from app.models.job import ImportJob
        from app.services.autocount_pull_service import _claim_and_enqueue_preview

        db, factory = task_db
        rows = [_canonical_row(f"{MARKER}-F6")]
        header = _header(rows)
        job_id = _seed_pull_job(
            db, job_type="autocount_products_pull", user_id=str(uuid.uuid4()),
            company_id=DEFAULT_COMPANY_ID, entity="products", company_code="SRT",
            snapshot_id=f"{MARKER}-snap-f6", phase="building",
        )

        captured = []
        monkeypatch.setattr(
            "app.services.autocount_pull_service.enqueue_job",
            lambda *a, **k: captured.append(1) or MagicMock(id=str(uuid.uuid4())),
        )

        job = db.query(ImportJob).filter(ImportJob.id == job_id).first()
        pull = dict(job.job_metadata["autocount_pull"])

        _claim_and_enqueue_preview(db, job, dict(pull), header)
        _claim_and_enqueue_preview(db, job, dict(pull), header)  # same stale pre-read pull

        assert len(captured) == 1, "a second caller off the same stale read must not enqueue again"
        row_after = _job_row(db, job_id)
        assert row_after["metadata"]["autocount_pull"]["phase"] == "previewing"


# ============================================================================ F-7


class TestF7ExactlyOnceConfirm:
    def test_f7_confirm_pull_is_exactly_once_no_orphan_apply_row(self, task_db, monkeypatch):
        """Two `confirm_pull` calls off two independent, stale `ImportJob` reads of the SAME
        row - the loser must enqueue nothing AND must leave no orphan `import_jobs` apply row
        behind. Today's `confirm_pull` inserts the loser's apply-job row and `db.commit()`s
        unconditionally BEFORE checking whether its own conditional UPDATE actually won the
        race, so the loser's insert survives even though nothing ever enqueues it - this is
        the red.

        One `task_db` session, not two (SR3/SR4's two-session recipe is for two REAL RQ task
        runs, each opening its own top-level `SessionLocal()`; `confirm_pull` is a plain
        function called directly here, and two live ORM Sessions doing interleaved
        begin/commit on ONE shared connection under `create_savepoint` corrupts the savepoint
        stack - proven empirically before this shape was chosen, see the tester report).
        Staleness is instead simulated by `expunge`-ing a first read (capturing its state as a
        genuinely detached copy) before a second, live read proceeds and confirms for real,
        then re-`add`-ing the detached copy to replay the loser's confirm against the SAME
        session once the identity slot is free again - both `confirm_pull` calls still run
        through the code path unmodified, including its own `db.refresh(job)`."""
        from app.models.job import ImportJob
        from app.services.autocount_pull_service import confirm_pull

        db, factory = task_db
        owner_id = str(uuid.uuid4())
        job_id = _seed_pull_job(
            db, job_type="autocount_products_pull", user_id=owner_id,
            company_id=DEFAULT_COMPANY_ID, entity="products", company_code="SRT",
            snapshot_id=f"{MARKER}-snap-f7", phase="review",
        )

        captured = []
        monkeypatch.setattr(
            "app.services.autocount_pull_service.enqueue_job",
            lambda *a, **k: captured.append(1) or MagicMock(id=str(uuid.uuid4())),
        )

        stale = db.query(ImportJob).filter(ImportJob.id == job_id).first()
        db.expunge(stale)  # a genuinely detached copy of the pre-confirm state

        live = db.query(ImportJob).filter(ImportJob.id == job_id).first()
        result1 = confirm_pull(db, live, user_id=owner_id)  # the winner
        db.expunge(live)  # free the identity slot so `stale` can be re-attached

        db.add(stale)  # the loser, replaying its own stale pre-read confirm
        result2 = confirm_pull(db, stale, user_id=owner_id)

        assert len(captured) == 1, "enqueue must fire exactly once across both racing callers"
        assert result1["apply_job_id"] == result2["apply_job_id"]

        apply_count = db.execute(
            text(
                "SELECT count(*) FROM import_jobs WHERE job_type = 'autocount_products_apply' "
                "AND metadata->'autocount_apply'->>'pull_job_id' = :pid"
            ),
            {"pid": str(job_id)},
        ).scalar()
        assert apply_count == 1, "the loser's own apply-job insert must not survive as an orphan"


# ============================================================================ F-8


class TestF8NonDictUpstreamBody:
    def test_f8_start_route_non_dict_upstream_body_is_unreachable_not_500(self, env):
        user = env.user("master_data.products.autocount_pull")
        env.as_user(user)
        env.fake.build = (200, [1, 2, 3])  # a JSON array, not the expected object

        resp = env.post_pull("products")

        assert resp.status_code == 502, resp.text
        assert resp.json().get("code") == "UNREACHABLE", resp.json()


# ============================================================================ F-9


class TestF9PriceIsNumeric:
    def test_f9_products_rows_price_is_numeric_and_download_cell_is_numeric(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        rows = [
            _canonical_row(f"{MARKER}-PRICE1", list_price="150.0"),
            _canonical_row(f"{MARKER}-PRICE2", list_price="0.0"),
            _canonical_row(f"{MARKER}-PRICE3", list_price="-1.0"),
        ]
        job_id, _ = sr3._seed_review_job(env, owner=owner, rows=rows)

        resp = env.client.get(f"{PULLS_URL}/{job_id}/rows", params={"limit": 50})
        assert resp.status_code == 200, resp.text
        by_code = {r["item_code"]: r for r in resp.json()["data"]}

        expected = {
            f"{MARKER}-PRICE1": 150.0, f"{MARKER}-PRICE2": 0.0, f"{MARKER}-PRICE3": -1.0,
        }
        for code, value in expected.items():
            price = by_code[code]["price"]
            assert isinstance(price, (int, float)), (code, price, type(price))
            assert float(price) == value

        dl = env.client.get(f"{PULLS_URL}/{job_id}/download.xlsx")
        assert dl.status_code == 200, dl.text
        wb = openpyxl.load_workbook(io.BytesIO(dl.content))
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        price_col = header.index("Price")
        for excel_row in ws.iter_rows(min_row=2):
            cell = excel_row[price_col]
            if cell.value in (None, ""):
                continue
            assert cell.data_type == "n", (cell.value, cell.data_type)


# ============================================================================ F-10


class TestF10CompareResponseSummary:
    def test_f10_compare_response_summary_equals_the_stored_get_compare(self, env):
        owner = env.user("master_data.products.autocount_pull")
        env.as_user(owner)
        job_id, rows = sr3._seed_review_job(env, owner=owner)
        template_rows = [sr3._template_row(r) for r in rows]

        # One row missing from the file (only_in_pull), one row edited to differ.
        template_rows.pop()
        template_rows[0]["Item Group"] = str(template_rows[0]["Item Group"]) + "-X"

        resp = env.client.post(
            f"{PULLS_URL}/{job_id}/compare", json={"filename": "chk.xlsx", "rows": template_rows},
        )
        assert resp.status_code == 200, resp.text
        body = resp.json()

        summary = body["summary"]
        for key in ("filename", "compared_at", "total", "matched", "different",
                    "only_in_excel", "only_in_pull"):
            assert key in summary, summary
        assert summary["filename"] == "chk.xlsx"
        assert isinstance(summary["only_in_excel"], int)
        assert isinstance(summary["only_in_pull"], int)
        assert summary["only_in_pull"] == 1

        for diff in body["differences"]:
            assert set(diff.keys()) >= {"item_code", "field", "excel", "pull"}

        # Top-level only_in_* stay LISTS of item codes - distinct from the summary's counts.
        assert isinstance(body["only_in_excel"], list)
        assert isinstance(body["only_in_pull"], list)
        assert len(body["only_in_pull"]) == summary["only_in_pull"]

        get_resp = env.client.get(f"{PULLS_URL}/{job_id}")
        assert get_resp.status_code == 200, get_resp.text
        stored_compare = get_resp.json()["compare"]
        assert stored_compare == summary


# ============================================================================ F-11


class TestF11DeadPreviewingPull:
    def test_f11_dead_previewing_pull_is_replaced_by_a_fresh_start(self, env):
        """What the orphan sweep leaves after a worker death: `import_jobs.status` is
        `failed` but the pull's own stored phase is still `previewing`."""
        user = env.user("master_data.products.autocount_pull")
        env.as_user(user)
        old_job_id = _seed_pull_job(
            env.db, job_type="autocount_products_pull", user_id=user["id"], company_id=env.company_a,
            entity="products", company_code=env.company_a_code, snapshot_id=f"{MARKER}-snap-dead",
            phase="previewing",
        )
        env.db.execute(text("UPDATE import_jobs SET status = 'failed' WHERE id = :id"), {"id": old_job_id})
        env.db.commit()

        resp = env.get_current("products")
        assert resp.status_code == 404, resp.text

        env.fake.calls.clear()
        start_resp = env.post_pull("products")
        assert start_resp.status_code in (200, 201, 202), start_resp.text
        new_job_id = start_resp.json()["job_id"]
        assert new_job_id != str(old_job_id)
        assert len(env.fake.calls) == 1
        assert env.fake.calls[0]["method"] == "POST"

        old_resp = env.get_pull(old_job_id)
        assert old_resp.status_code == 200, old_resp.text
        assert old_resp.json()["phase"] == "failed"


# ============================================================================ F-12


class TestF12ArchiveMime:
    def test_f12a_replace_route_stores_the_uploads_own_mime_for_a_non_xlsm_file(self, client):
        c, db, backend = client
        raw = sr4._xlsx_bytes()

        resp = c.post(
            sr4.REPLACE_URL,
            files={"file": ("stock.xls", raw, "application/vnd.ms-excel")},
        )
        assert resp.status_code == 201, resp.text
        assert len(backend.uploads) == 1, backend.uploads
        _, _, stored_mime = backend.uploads[0]
        assert stored_mime == "application/vnd.ms-excel", stored_mime

    def test_f12b_service_accepts_a_mime_type_keyword_task_path_defaults_to_xlsx(
        self, task_db, monkeypatch
    ):
        from app.models.resources import Attachment, AttachmentType
        from app.services.stock_list_archive_service import replace_latest_stock_list
        import app.services.storage_router as storage_router

        db, factory = task_db
        db.add(AttachmentType(
            id=str(uuid.uuid4()), type_name="Stock_List", allowed_extensions="xls,xlsx,xlsm",
            max_file_size_mb=10,
        ))
        db.commit()

        class _Backend:
            def __init__(self):
                self.uploads = []

            def upload_file(self, *, file_content, file_path, content_type=None):
                self.uploads.append((file_path, file_content, content_type))
                return file_path, f"https://cdn.test/{file_path}"

        backend = _Backend()
        monkeypatch.setattr(storage_router, "default_provider", lambda: "s3")
        monkeypatch.setattr(storage_router, "get_backend", lambda provider: backend)
        monkeypatch.setattr(
            storage_router, "cdn_base_url", lambda provider, key: f"https://cdn.test/{key}"
        )

        user_id = str(uuid.uuid4())
        raw = sr4._xlsx_bytes()

        # Task path: no mime_type passed at all - defaults to the xlsx spreadsheetml mime.
        attachment = replace_latest_stock_list(db, file_bytes=raw, filename="x.xlsx", user_id=user_id)
        db.expire_all()
        stored = db.query(Attachment).filter(Attachment.id == attachment.id).first()
        assert stored.mime_type.endswith("spreadsheetml.sheet"), stored.mime_type

        # The keyword itself must exist and be honoured - pins the coder's contract.
        attachment2 = replace_latest_stock_list(
            db, file_bytes=raw, filename="y.xls", user_id=user_id, mime_type="application/vnd.ms-excel",
        )
        db.expire_all()
        stored2 = db.query(Attachment).filter(Attachment.id == attachment2.id).first()
        assert stored2.mime_type == "application/vnd.ms-excel", stored2.mime_type

