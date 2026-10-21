"""Dealer targets import, fifth real-file gate (lane SALES-TARGET-DEALER, crew 16 Oct, round 4).

Round 4 on the owner's 15 Oct prod backup: Mocha parses (135 of 135), every failing row stops at
name matching. Shapes, masked:
1. Mocha A/C II / A/C III rows: the dealer's ledgers are "[A/C I]" (level 1), "[CERAMIC]" and
   "[IBORN]" (no level). None answers to A/C II, yet the reason said "more than one customer has
   this name and account". The row stays a pick (owner ruling 16 Oct 09:12: nothing ambiguous is
   linked; A/C II = [CERAMIC] is an open owner question), with a reason that says what is true.
2. Two A/C I ledgers whose names differ only in punctuation or spacing ("TIP-TOP TRADING" and
   "TIP TOP TRADING"): the one whose name the file writes letter for letter is the exact match. When
   the file writes neither, it stays a pick.
3. THRESHOLD_REPEATED "A/C I 9% | A/C I 9%": the same threshold written twice with the same rate
   in one period is one tier, with a warning on the row.
"""
from __future__ import annotations

import io
from decimal import Decimal

from .test_sales_targets_s1 import BASE, _code, _uid, api, world  # noqa: F401  (fixtures)

D = Decimal


def _bytes(fill):
    from openpyxl import Workbook

    wb = Workbook()
    fill(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _customer(db, company_id, name, level):
    from app.models.order import Customer

    row = Customer(id=_uid(), company_id=company_id, customer_code=_code(), customer_name=name, account_level=level)
    db.add(row)
    db.flush()
    return row


def _dry_run(client, db, name, account=None):
    def fill(ws):
        if account is None:
            ws["A1"], ws["B1"] = "DEALER NAME", "Q1"
            ws["A2"], ws["B2"] = name, "100,000 - 1%"
        else:
            ws["A1"], ws["B1"], ws["C1"] = "DEALER NAME", "ACCOUNT", "Q1"
            ws["A2"], ws["B2"], ws["C2"] = name, account, "100,000 - 1%"

    files = {"file": ("t.xlsx", _bytes(fill), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    db.autoflush = False
    try:
        return client.post(f"{BASE}/import", params={"dry_run": "true", "year": 2026}, files=files).json()["rows"][0]
    finally:
        db.autoflush = True


def _mocha_trio(db, company_id):
    return [
        _customer(db, company_id, "MIKE TILES SDN BHD [A/C I]", 1),
        _customer(db, company_id, "MIKE TILES SDN BHD [CERAMIC]", None),
        _customer(db, company_id, "MIKE TILES SDN BHD [IBORN]", None),
    ]


def test_an_account_no_ledger_of_the_name_has_is_not_called_several_matches(api):
    client, db, company_id = api
    trio = _mocha_trio(db, company_id)
    row = _dry_run(client, db, "MIKE TILES SDN BHD", "A/C II")
    assert row["status"] == "needs_pick"
    assert "more than one customer" not in row["reason"]
    assert row["reason"].startswith('no ledger of this name has the account "A/C II"'), row["reason"]
    # The dealer's own ledgers are the suggestions, so the pick is one click.
    assert {s["name"] for s in row["suggestions"]} == {c.customer_name for c in trio}


def test_a_suffix_account_no_ledger_has_reads_the_same(api):
    client, db, company_id = api
    _mocha_trio(db, company_id)
    row = _dry_run(client, db, "MIKE TILES SDN BHD A/C III")
    assert row["status"] == "needs_pick"
    assert row["reason"].startswith('no ledger of this name has the account "A/C III"'), row["reason"]


def test_the_file_s_account_still_auto_maps_beside_the_marker_ledgers(api):
    client, db, company_id = api
    one, _, _ = _mocha_trio(db, company_id)
    row = _dry_run(client, db, "MIKE TILES SDN BHD", "A/C I")
    assert row["match"] == "auto" and row["customer"] == one.customer_name


def test_a_name_no_ledger_has_keeps_its_own_reason(api):
    client, db, company_id = api
    _mocha_trio(db, company_id)
    row = _dry_run(client, db, "OSCAR STONE SDN BHD", "A/C I")
    assert row["reason"].startswith("no customer has this name and account"), row["reason"]


def test_the_ledger_the_file_writes_letter_for_letter_wins_a_punctuation_tie(api):
    client, db, company_id = api
    _customer(db, company_id, "TIP TOP TRADING SDN. BHD. [A/C I]", 1)
    hyphen = _customer(db, company_id, "TIP-TOP TRADING SDN BHD [A/C I]", 1)
    row = _dry_run(client, db, "Tip-Top Trading Sdn Bhd A/C I")
    assert row["match"] == "auto" and row["customer"] == hyphen.customer_name


def test_a_spacing_tie_breaks_the_same_way(api):
    client, db, company_id = api
    # Runs of spaces in the CRM name count as one (the reader already collapses the file's).
    spaced = _customer(db, company_id, "HOME  &  DECO  SDN BHD", 1)
    _customer(db, company_id, "HOME&DECO SDN BHD", 1)
    row = _dry_run(client, db, "home & deco sdn bhd", "A/C I")
    assert row["match"] == "auto" and row["customer"] == spaced.customer_name


def test_a_punctuation_tie_the_file_writes_neither_way_stays_a_pick(api):
    client, db, company_id = api
    _customer(db, company_id, "TIP TOP TRADING SDN. BHD. [A/C I]", 1)
    _customer(db, company_id, "TIP-TOP TRADING SDN BHD [A/C I]", 1)
    row = _dry_run(client, db, "TIP / TOP TRADING SDN BHD A/C I")
    assert row["status"] == "needs_pick"
    assert row["reason"].startswith("more than one customer has this name and account")
    assert len(row["suggestions"]) == 2


def _repeat(second_rate):
    from app.services.sales.dealer_target_import import parse_workbook

    def fill(ws):
        ws["A1"], ws["B1"] = "DEALER NAME", "Q1"
        ws.merge_cells("B1:E1")
        for col, text in zip("BCDE", ("Tier 1", "Tier 1 Incentive", "Tier 2", "Tier 2 Incentive")):
            ws[f"{col}2"] = text
        ws["A3"], ws["B3"], ws["C3"], ws["D3"], ws["E3"] = "X A/C I", 300000, "A/C I- 9%", 300000, second_rate

    return parse_workbook(_bytes(fill), year=2026)


def test_a_threshold_written_twice_with_the_same_rate_is_one_tier_with_a_warning():
    parsed = _repeat("A/C I- 9%")
    row = parsed.rows[0]
    assert row.status == "mapped", row.reason
    (q,) = [s for s in parsed.schemes if s.kind == "quarterly"]
    assert [(t.threshold, t.rate, dict(t.account_rates or {})) for t in q.tiers if not t.is_additional] == [
        (D("300000"), None, {next(iter(q.tiers[0].account_rates)): D("9")})
    ]
    assert any("written twice" in w for w in row.warnings), row.warnings


def test_a_threshold_written_twice_with_different_rates_still_fails():
    row = _repeat("A/C I- 7%").rows[0]
    assert row.status == "not_representable"
    assert "repeats in one period" in row.reason


# --- reviewer re-check of the gate-5 round --------------------------------------------------


def _dry_run_rows(client, db, rows):
    def fill(ws):
        ws["A1"], ws["B1"], ws["C1"] = "DEALER NAME", "ACCOUNT", "Q1"
        for i, (name, account, tier) in enumerate(rows, start=2):
            ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = name, account, tier

    files = {"file": ("t.xlsx", _bytes(fill), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    db.autoflush = False
    try:
        return client.post(f"{BASE}/import", params={"dry_run": "true", "year": 2026}, files=files).json()["rows"]
    finally:
        db.autoflush = True


def test_two_spellings_in_one_file_each_get_their_own_ledger(api):
    """Review B1: the match cache keyed on the punctuation-free name served row 2 row 1's ledger."""
    client, db, company_id = api
    dotted = _customer(db, company_id, "TIP TOP TRADING SDN. BHD. [A/C I]", 1)
    hyphen = _customer(db, company_id, "TIP-TOP TRADING SDN BHD [A/C I]", 1)
    first, second = _dry_run_rows(client, db, [
        ("Tip-Top Trading Sdn Bhd", "A/C I", "100,000 - 1%"),
        ("Tip Top Trading Sdn. Bhd.", "A/C I", "200,000 - 2%"),
    ])
    assert first["match"] == "auto" and first["customer"] == hyphen.customer_name
    assert second["match"] == "auto" and second["customer"] == dotted.customer_name, second


def test_a_spelling_tie_never_lands_on_a_ledger_naming_another_account(api):
    """Review B2: two ledgers name A/C I, a [CERAMIC] one answers only through a level."""
    client, db, company_id = api
    _customer(db, company_id, "TIP TOP SDN. BHD. [A/C I]", 1)
    _customer(db, company_id, "TIP-TOP SDN BHD [A/C I]", 1)
    ceramic = _customer(db, company_id, "TIP TOP SDN BHD [CERAMIC]", 1)
    row = _dry_run(client, db, "TIP TOP SDN BHD", "A/C I")
    assert row["customer"] != ceramic.customer_name
    assert row["status"] == "needs_pick"


def test_a_file_with_no_account_against_several_ledgers_says_so(api):
    """Review S1: not "no customer has this name"."""
    client, db, company_id = api
    a = _customer(db, company_id, "MIKE TILES SDN BHD [A/C I]", 1)
    b = _customer(db, company_id, "MIKE TILES SDN BHD [CERAMIC]", None)
    row = _dry_run(client, db, "MIKE TILES SDN BHD")
    assert row["status"] == "needs_pick"
    assert row["reason"].startswith("this name has several ledgers and the file names no account"), row["reason"]
    assert {s["name"] for s in row["suggestions"]} == {a.customer_name, b.customer_name}


def test_an_account_from_a_bracket_marker_is_quoted(api):
    """Review S2: the reason quoted an empty account."""
    client, db, company_id = api
    _mocha_trio(db, company_id)
    row = _dry_run(client, db, "MIKE TILES SDN BHD [DELUXE]")
    assert row["reason"].startswith('no ledger of this name has the account "DELUXE"'), row["reason"]


def test_a_branch_starting_with_l_is_not_an_account_marker():
    """Review S3: "(LABUAN)" read as an account marker vanished from both tie-breaks."""
    from app.services.sales.dealer_target_import import branch_key, literal_name

    assert literal_name("ABC (LABUAN) SDN BHD") == "ABC (LABUAN) SDN BHD"
    assert "LABUAN" in branch_key("ABC (LABUAN) SDN BHD")
    assert literal_name("ABC SDN BHD (A/C II)") == "ABC SDN BHD"
    assert "JB" in branch_key("ABC (JB) [A/C I]") and branch_key("ABC (A/C I)") == "ABC"


def test_a_monthly_repeat_warns_once_per_row():
    """Review nit: one warning per row, not one per repeated month."""
    from app.services.sales.dealer_target_import import parse_workbook

    def fill(ws):
        ws["A1"], ws["B1"] = "DEALER NAME", "Monthly"
        ws.merge_cells("B1:I1")
        ws["B2"], ws["F2"] = "JAN", "FEB"
        ws.merge_cells("B2:E2")
        ws.merge_cells("F2:I2")
        ws["A4"] = "X A/C I"
        for start in ("BCDE", "FGHI"):
            for col, text in zip(start, ("Tier 1", "Tier 1 Incentive", "Tier 2", "Tier 2 Incentive")):
                ws[f"{col}3"] = text
            for col, value in zip(start, (30000, "A/C I- 2%", 30000, "A/C I- 2%")):
                ws[f"{col}4"] = value

    row = parse_workbook(_bytes(fill), year=2026).rows[0]
    assert row.status == "mapped", row.reason
    assert sum("written twice" in w for w in row.warnings) == 1, row.warnings


def test_a_tie_where_every_ledger_names_another_account_stays_a_pick(api):
    """Re-check B2: [CERAMIC] and [IBORN] both at level 1, the file asks A/C I."""
    client, db, company_id = api
    _customer(db, company_id, "TIP TOP SDN BHD [CERAMIC]", 1)
    _customer(db, company_id, "TIP-TOP SDN BHD [IBORN]", 1)
    row = _dry_run(client, db, "TIP TOP SDN BHD", "A/C I")
    assert row["status"] == "needs_pick", row
    assert row["reason"].startswith("more than one customer has this name and account")


def test_an_account_marker_past_v_is_not_a_branch():
    """Re-check nit: "(A/C VI)" reads RAWVI, still an account marker."""
    from app.services.sales.dealer_target_import import branch_key

    assert branch_key("ABC (A/C VI) SDN BHD") == "ABC"


def test_a_branch_starting_with_raw_is_not_an_account_marker():
    """Re-check should-fix: "(RAWANG)" is a town, not A/C <n>."""
    from app.services.sales.dealer_target_import import branch_key

    assert branch_key("ABC (RAWANG) SDN BHD") == "ABC RAWANG"
