"""Order inquiry Project column and handover email read the sales order's project label.

PLAN-oi-project-label-from-so.md / oi-project-label-from-so-acceptance-criteria.md.

An order ADOPTED from the AutoCount book has `project_id` NULL by design
(`project_so_adoption_service.py`), so the worklist's PROJECT column and the handover
email used to print nothing for it even though the core sales order's own detail page
already shows a project label - `SalesOrder.project_label`, resolved by
`app.services.project_label_rules`. The rule under test is one coalesce, read by three
places: the worklist's `project_title` / `project_customer`, its free-text search, and
`ProjectOrderInquiryService._handover_order_facts` / `_project_customer_labels`.

AC-9..AC-12 (owner ruling, 21 Sep 2026, plan section 4): whether or not the order has a
project, purchasing is still handed the raise - a `ProjectTask` only when there is a
project (`tasks.project_id` is NOT NULL), the in-app notification either way, its
heading following the same fallback (`Project.title` > `SalesOrder.project_label` > the
SO reference).

AC-13..AC-18 (owner ask, 21 Sep 2026, plan section 5): the Project FILTER follows the
same text the column prints - `project`, exact match on `_PROJECT_TITLE`, on every
worklist entry point that reads `project_id` today (list, summary, matrix, export, and
the body-scoped bulk actions). `project_id` is unchanged, for an existing deep link.
"""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import text

from app.models.notification import Notification
from app.models.order import Customer, SalesOrder
from app.models.procurement import PurchaseOrder, PurchaseOrderLine, Supplier
from app.models.product import Product, ProductCategory, UnitOfMeasure
from app.models.project_so import (
    INQUIRY_RAISED,
    IV_ORDER,
    OrderInquiry,
    OrderInquiryLink,
    OrderInquiryRow,
    ProjectSalesOrder,
    ProjectSalesOrderLine,
    SOSupplyDecision,
)
from app.models.user import User, UserRole, UserRoleAssignment, UserStatus
from app.services import project_seed_service
from app.services.project_order_inquiry_service import ProjectOrderInquiryService

from ._pg_fixture import blank_session

MARKER = "zzt-oi-project-label"
BASE = "/api/v1/project-sales"
LIST = f"{BASE}/order-inquiries"

READ_ONLY = ["projects.projects.view"]

PROJECT_LABEL = f"{MARKER} KITACON / PHASE 6A"


def _uid() -> str:
    return str(uuid.uuid4())


def _sorento(db) -> str:
    return db.execute(text("select id from companies where code = 'SRT'")).scalar()


def _user(db, name: str) -> str:
    user_id = _uid()
    db.add(User(id=user_id, email=f"{user_id}@zzt.test", name=name))
    db.flush()
    return user_id


def _product(db, code: str, name: str) -> Product:
    uom = UnitOfMeasure(id=_uid(), uom_code=f"ZZT{_uid()[:6]}", uom_name="Unit")
    category = ProductCategory(
        id=_uid(), category_code=f"ZZT-{_uid()[:8]}", category_name=f"{MARKER} cat"
    )
    db.add_all([uom, category])
    db.flush()
    row = Product(
        id=_uid(),
        product_code=code,
        product_name=name,
        category_id=category.id,
        base_uom_id=uom.id,
        list_price=Decimal("100.00"),
    )
    db.add(row)
    db.flush()
    return row


def _customer(db, company_id: str, name: str) -> Customer:
    row = Customer(
        id=_uid(),
        company_id=company_id,
        customer_code=f"ZZT-{_uid()[:8]}",
        customer_name=name,
    )
    db.add(row)
    db.flush()
    return row


def _adopted_order(
    db,
    company_id: str,
    *,
    project_label: str | None,
    project_id: str | None = None,
    customer=None,
    row_id: str | None = None,
) -> dict:
    """One AutoCount-adopted `ProjectSalesOrder` (`project_id` NULL by design unless the
    caller passes a registered one), its core `SalesOrder` carrying `project_label`, and
    one raised `OrderInquiryRow` - the chain every AC needs.

    ``row_id`` lets a caller pin the row's own id (AC-4: two rows with equal
    `project_title` need to be told apart by something OTHER than the id the default
    tiebreak sorts by, or a coincidental pass reads as a fixed test)."""
    core = SalesOrder(
        id=_uid(),
        company_id=company_id,
        so_number=f"ZZTSO{_uid()[:8]}",
        customer_id=customer.id if customer else None,
        order_date=date(2026, 1, 8),
        project_label=project_label,
        project_label_source="inquiry" if project_label else None,
    )
    db.add(core)
    db.flush()
    order = ProjectSalesOrder(
        id=_uid(),
        company_id=company_id,
        project_id=project_id,
        so_id=core.id,
        provisional_ref=core.so_number,
        autocount_doc_no=core.so_number,
        status="adopted",
    )
    db.add(order)
    db.flush()
    product = _product(db, f"ZZT-P-{_uid()[:6]}", f"{MARKER} product")
    line = ProjectSalesOrderLine(
        id=_uid(),
        company_id=company_id,
        project_sales_order_id=order.id,
        line_no=1,
        product_id=product.id,
        description=f"{MARKER} line",
        qty=Decimal("10"),
        uom="UNIT",
        unit_price=Decimal("10.00"),
        amount=Decimal("100.00"),
        delivery_date=date(2026, 1, 19),
    )
    db.add(line)
    db.flush()
    inquiry = OrderInquiry(
        id=_uid(),
        company_id=company_id,
        project_sales_order_id=order.id,
        state=INQUIRY_RAISED,
    )
    db.add(inquiry)
    db.flush()
    row = OrderInquiryRow(
        id=row_id or _uid(),
        company_id=company_id,
        order_inquiry_id=inquiry.id,
        so_line_id=line.id,
        item_code=product.product_code,
        qty=Decimal("10"),
        delivery_date=date(2026, 1, 19),
        state=INQUIRY_RAISED,
        verb=IV_ORDER,
    )
    db.add(row)
    db.commit()
    return {"core": core, "pso": order, "inquiry": inquiry, "row": row, "product": product}


# ------------------------------------------------------------------ worklist (AC-1..AC-4)


#: S5 AC-17: "Select all N matching" (acknowledge's own `filter`) and the unplace-all
#: routes are write actions, gated on their own grants beside the plain view every
#: other test here reads with.
WRITE = READ_ONLY + [
    "projects.order_inquiry.action",
    "projects.order_inquiries.acknowledge",
]


def _client(db, user_id: str, permissions=READ_ONLY):
    from fastapi.testclient import TestClient

    from app.database import get_db
    from app.dependencies import get_current_user, get_current_user_or_api_key
    from app.main import app
    from app.services.company_scope_resolver import apply_company_scope
    from app.services.user_service import UserPermissionService

    actor = {"id": user_id, "email": f"{user_id}@zzt.test", "role": "user"}
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: dict(actor)
    app.dependency_overrides[get_current_user_or_api_key] = lambda: dict(actor)
    app.dependency_overrides[apply_company_scope] = lambda: None

    originals = (
        UserPermissionService.check_user_has_permission,
        UserPermissionService.get_user_permission_slugs,
    )
    granted = list(permissions)
    UserPermissionService.check_user_has_permission = (
        lambda self, uid, slug: slug in granted
    )
    UserPermissionService.get_user_permission_slugs = lambda self, uid: list(granted)
    return TestClient(app), originals


def _restore(originals) -> None:
    from app.main import app
    from app.services.user_service import UserPermissionService

    UserPermissionService.check_user_has_permission = originals[0]
    UserPermissionService.get_user_permission_slugs = originals[1]
    app.dependency_overrides.clear()


@pytest.fixture()
def api():
    from app.models.base import company_scope

    with blank_session() as db:
        company_id = _sorento(db)
        project_seed_service.run(db, company_id=company_id)
        user_id = _user(db, f"{MARKER} Eling")
        client, originals = _client(db, user_id)
        try:
            with company_scope(db, frozenset({company_id})):
                yield client, db, company_id
        finally:
            _restore(originals)


@pytest.fixture()
def write_api():
    from app.models.base import company_scope

    with blank_session() as db:
        company_id = _sorento(db)
        project_seed_service.run(db, company_id=company_id)
        user_id = _user(db, f"{MARKER} Buyer")
        client, originals = _client(db, user_id, WRITE)
        try:
            with company_scope(db, frozenset({company_id})):
                yield client, db, company_id
        finally:
            _restore(originals)


def test_worklist_project_cell_falls_back_to_the_sos_project_label(api):
    """AC-1: the adopted row's `project_title` / `project_customer` read the core
    sales order's free-text label rather than printing nothing."""
    client, db, company_id = api
    customer = _customer(db, company_id, f"{MARKER} Optad Sdn Bhd")
    db.commit()
    seeded = _adopted_order(db, company_id, project_label=PROJECT_LABEL, customer=customer)

    body = client.get(LIST).json()
    row = next(r for r in body["data"] if r["id"] == seeded["row"].id)

    assert row["project_title"] == PROJECT_LABEL
    assert row["project_customer"] == f"{customer.customer_name} / {PROJECT_LABEL}"


def test_worklist_project_cell_prints_the_label_alone_with_no_customer(api):
    """AC-1 (no resolvable customer): `project_customer` is the bare label."""
    client, db, company_id = api
    seeded = _adopted_order(db, company_id, project_label=PROJECT_LABEL, customer=None)

    body = client.get(LIST).json()
    row = next(r for r in body["data"] if r["id"] == seeded["row"].id)

    assert row["project_title"] == PROJECT_LABEL
    assert row["project_customer"] == PROJECT_LABEL


def test_a_registered_project_wins_over_the_sos_label(api):
    """AC-2: a registered `Project` outranks the SO's free-text label."""
    from app.services.project_service import register_project

    client, db, company_id = api
    user_id = _user(db, f"{MARKER} Registrar")
    project = register_project(
        db,
        company_id=company_id,
        actor_user_id=user_id,
        developer_party_id=None,
        title=f"{MARKER} TUJU RESIDENCE",
    )
    db.flush()
    seeded = _adopted_order(
        db, company_id, project_label=PROJECT_LABEL, project_id=project.id
    )

    body = client.get(LIST).json()
    row = next(r for r in body["data"] if r["id"] == seeded["row"].id)

    assert row["project_title"] == f"{MARKER} TUJU RESIDENCE"


def test_search_reaches_the_sos_project_label(api):
    """AC-3: the free-text search box hits an adopted row by its SO's project label."""
    client, db, company_id = api
    seeded = _adopted_order(db, company_id, project_label=PROJECT_LABEL, customer=None)

    hit = client.get(LIST, params={"query": "KITACON"}).json()
    miss = client.get(LIST, params={"query": "ZZZNOPE"}).json()

    assert seeded["row"].id in {r["id"] for r in hit["data"]}
    assert seeded["row"].id not in {r["id"] for r in miss["data"]}


def test_sort_by_project_title_orders_adopted_rows_by_the_sos_label(api):
    """AC-4: `sort=project_title` orders two adopted rows by the SO's own label.

    ALPHA is seeded with a HIGH row id and BETA a LOW one, so the default
    `OrderInquiryRow.id.asc()` tiebreak - what a pre-fix run falls back to, since both
    rows' `project_title` reads NULL off `Project.title` alone - orders them
    [beta, alpha], the WRONG way round. Only the real fix, sorting by the SO's own
    label, produces [alpha, beta]: this fails every time pre-fix rather than on a coin
    flip of freshly generated uuids (review round 1, reviewer measured 4/6 green
    pre-fix with random ids)."""
    client, db, company_id = api
    alpha = _adopted_order(
        db,
        company_id,
        project_label=f"{MARKER} ALPHA",
        customer=None,
        row_id="ffffffff-ffff-4fff-8fff-ffffffffffff",
    )
    beta = _adopted_order(
        db,
        company_id,
        project_label=f"{MARKER} BETA",
        customer=None,
        row_id="00000000-0000-4000-8000-000000000001",
    )
    ids = {alpha["row"].id, beta["row"].id}

    body = client.get(
        LIST, params={"sort": "project_title", "dir": "asc", "limit": 100}
    ).json()
    ordered = [row["id"] for row in body["data"] if row["id"] in ids]

    assert ordered == [alpha["row"].id, beta["row"].id]


def test_no_label_anywhere_prints_none_and_does_not_raise(api):
    """AC-7: no registered project and no SO project label - `project_title` is None,
    nothing raises."""
    client, db, company_id = api
    seeded = _adopted_order(db, company_id, project_label=None, customer=None)

    response = client.get(LIST)

    assert response.status_code == 200, response.text
    row = next(r for r in response.json()["data"] if r["id"] == seeded["row"].id)
    assert row["project_title"] is None


# ------------------------------------------------------------------ handover (AC-5, AC-6)


def _register_post_commit_dispatch() -> None:
    """Idempotent (module-global flag): a bare `TestClient(app)` never runs FastAPI's
    lifespan, so nothing else in this file's `seeded` fixture would register the
    `after_commit` listener `_hand_to_purchasing`'s notification depends on - same call
    site `tests/test_order_inquiry_handover_automation.py::_register` uses."""
    from app.services.project_order_inquiry_service import (
        register_order_inquiry_post_commit_dispatch,
    )

    register_order_inquiry_post_commit_dispatch()


@pytest.fixture()
def seeded():
    with blank_session() as db:
        company_id = _sorento(db)
        project_seed_service.run(db, company_id=company_id)
        _register_post_commit_dispatch()
        yield db, company_id


def test_handover_order_facts_read_the_sos_project_label(seeded):
    """AC-5: `_handover_order_facts` carries the SO's label for an adopted order."""
    db, company_id = seeded
    made = _adopted_order(db, company_id, project_label=PROJECT_LABEL, customer=None)

    service = ProjectOrderInquiryService(db)
    facts = service._handover_order_facts(made["pso"].id)

    assert facts["project"] == PROJECT_LABEL


def test_handover_order_facts_prefer_a_registered_project(seeded):
    """AC-5: a registered project still wins over the SO's label."""
    from app.services.project_service import register_project

    db, company_id = seeded
    user_id = _user(db, f"{MARKER} Registrar")
    project = register_project(
        db,
        company_id=company_id,
        actor_user_id=user_id,
        developer_party_id=None,
        title=f"{MARKER} TUJU RESIDENCE",
    )
    db.flush()
    made = _adopted_order(
        db, company_id, project_label=PROJECT_LABEL, project_id=project.id
    )

    service = ProjectOrderInquiryService(db)
    facts = service._handover_order_facts(made["pso"].id)

    assert facts["project"] == f"{MARKER} TUJU RESIDENCE"


def test_project_customer_labels_end_with_the_sos_project_label(seeded):
    """AC-6: `_project_customer_labels` ends with the SO's label for an adopted order."""
    db, company_id = seeded
    customer = _customer(db, company_id, f"{MARKER} Optad Sdn Bhd")
    db.commit()
    made = _adopted_order(db, company_id, project_label=PROJECT_LABEL, customer=customer)

    service = ProjectOrderInquiryService(db)
    labels = service._project_customer_labels({made["pso"].id})

    assert labels[made["pso"].id].endswith(PROJECT_LABEL)


# ----------------------------------- AC-9..AC-12: hand to purchasing either way


class _NoCloseSession:
    """Delegates to a real Session but ignores close(): `_fire_pending_purchasing_
    notifications`'s fresh `SessionLocal()` closes the session it believes it opened,
    while this wraps one bound to the fixture's own connection (same pattern as
    `test_conversation_ticket_send.py`, minus reusing the SAME Session object - the
    listener fires INSIDE `db.commit()` itself, and a second query on that same Session
    at that moment raises "This session is in 'committed' state")."""

    def __init__(self, inner):
        self._inner = inner

    def close(self):  # noqa: D401 - deliberate no-op
        return None

    def __getattr__(self, name):
        return getattr(self._inner, name)


def _patch_session_local(monkeypatch, db) -> None:
    """A genuinely separate `Session` bound to the SAME connection as the fixture's
    `db` - so it shares the scratch schema's `search_path` and can see what `db` just
    committed (a savepoint release, not a real commit, under
    `join_transaction_mode="create_savepoint"`) - rather than a second reference to
    `db` itself, which the listener cannot query from inside its own commit.

    A test using this then sees a `SAWarning: nested transaction already deassociated
    from connection` on its OWN later `db.commit()` - expected fixture bookkeeping (the
    fresh Session above released ITS savepoint on the shared connection first, ahead of
    `db`'s own), not a failure and nothing this fixture can silence without hiding a
    real one."""
    import app.database as database_module
    from sqlalchemy.orm import Session as SQLASession

    def _fresh():
        return _NoCloseSession(
            SQLASession(bind=db.get_bind(), join_transaction_mode="create_savepoint")
        )

    monkeypatch.setattr(database_module, "SessionLocal", _fresh)


def _purchasing_user(db, company_id: str) -> str:
    """One ACTIVE user holding a role whose slug starts with `purchasing` -
    `_purchasing_user_ids` matches by prefix, the same family SCM's own roles grow to."""
    user_id = _uid()
    db.add(
        User(
            id=user_id,
            email=f"{user_id}@zzt.test",
            name=f"{MARKER} buyer",
            status=UserStatus.ACTIVE.value,
            is_trashed=False,
        )
    )
    role = UserRole(
        id=_uid(),
        slug=f"purchasing-{_uid()[:8]}",
        name=f"{MARKER} Purchasing {_uid()[:8]}",
    )
    db.add(role)
    db.flush()
    db.add(UserRoleAssignment(id=_uid(), user_id=user_id, role_id=role.id))
    db.flush()
    return user_id


def _order_for_confirm(
    db, company_id: str, *, project_label: str | None, project_id: str | None = None
) -> dict:
    """A fresh order with ONE line and NO inquiry yet - what `_confirm` below raises
    against, so `_hand_to_purchasing` runs through the real confirm route rather than
    being called directly."""
    core = SalesOrder(
        id=_uid(),
        company_id=company_id,
        so_number=f"ZZTSO{_uid()[:8]}",
        order_date=date(2026, 1, 8),
        project_label=project_label,
        project_label_source="inquiry" if project_label else None,
    )
    db.add(core)
    db.flush()
    order = ProjectSalesOrder(
        id=_uid(),
        company_id=company_id,
        project_id=project_id,
        so_id=core.id,
        provisional_ref=core.so_number,
        autocount_doc_no=core.so_number,
        status="adopted",
    )
    db.add(order)
    db.flush()
    product = _product(db, f"ZZT-P-{_uid()[:6]}", f"{MARKER} confirm product")
    line = ProjectSalesOrderLine(
        id=_uid(),
        company_id=company_id,
        project_sales_order_id=order.id,
        line_no=1,
        product_id=product.id,
        description=f"{MARKER} confirm line",
        qty=Decimal("10"),
        uom="UNIT",
        unit_price=Decimal("10.00"),
        amount=Decimal("100.00"),
        delivery_date=date(2026, 1, 19),
    )
    db.add(line)
    db.commit()
    return {"core": core, "pso": order, "line": line, "product": product}


def _confirm(db, order, *, actor_user_id, buy=None) -> dict:
    """The Buy-only handoff `refresh_for_decision` writes - the same shape
    `tests/test_project_order_inquiry.py::_confirm` drives, so `_hand_to_purchasing`
    (called from inside `refresh_for_decision`) runs for real rather than being called
    directly."""
    lines = (
        db.query(ProjectSalesOrderLine)
        .filter(ProjectSalesOrderLine.project_sales_order_id == order.id)
        .order_by(ProjectSalesOrderLine.line_no.asc())
        .all()
    )
    revision = (
        db.query(SOSupplyDecision)
        .filter(SOSupplyDecision.project_sales_order_id == order.id)
        .count()
        + 1
    )
    decision = SOSupplyDecision(
        id=_uid(),
        company_id=order.company_id,
        project_sales_order_id=order.id,
        revision_no=revision,
        state="active" if revision == 1 else "superseded",
        line_snapshots=[{"line_no": line.line_no} for line in lines],
        confirmed_by=actor_user_id,
        confirmed_at=datetime.utcnow(),
    )
    db.add(decision)
    db.flush()
    service = ProjectOrderInquiryService(db)
    return service.refresh_for_decision(
        order,
        decision,
        [
            {
                "line": line,
                "line_no": line.line_no,
                "item_code": service._product_code(line.product_id),
                "buy_qty": Decimal(str(buy)) if buy is not None else Decimal(str(line.qty)),
                "required_date": line.delivery_date,
                "stock_location": line.stock_location,
            }
            for line in lines
        ],
        actor_user_id=actor_user_id,
    )


def test_adopted_order_with_no_project_still_hands_to_purchasing(seeded, monkeypatch):
    """AC-9: no registered project, but the SO carries a label - the notification
    fires, headed by that label, and no `ProjectTask` is created."""
    db, company_id = seeded
    buyer_id = _purchasing_user(db, company_id)
    owner = _user(db, f"{MARKER} CS")
    made = _order_for_confirm(db, company_id, project_label=PROJECT_LABEL, project_id=None)
    _patch_session_local(monkeypatch, db)

    result = _confirm(db, made["pso"], actor_user_id=owner)
    db.commit()

    inquiry = result["inquiry"]
    assert inquiry is not None
    notification = (
        db.query(Notification)
        .filter(
            Notification.user_id == buyer_id,
            Notification.event_type == "project_order_inquiry_raised",
            Notification.source_entity_id == inquiry.id,
        )
        .first()
    )
    assert notification is not None
    assert notification.body.startswith(f"{PROJECT_LABEL}:")
    assert notification.data.get("project_id") is None
    assert notification.data.get("sales_order_ref") == made["core"].so_number
    assert ProjectOrderInquiryService(db).task_for(inquiry.id) is None


def test_adopted_order_with_no_label_and_no_project_uses_the_reference(
    seeded, monkeypatch
):
    """AC-10: neither a project nor an SO label - the body falls back to the SO
    reference, and nothing raises."""
    db, company_id = seeded
    buyer_id = _purchasing_user(db, company_id)
    owner = _user(db, f"{MARKER} CS")
    made = _order_for_confirm(db, company_id, project_label=None, project_id=None)
    _patch_session_local(monkeypatch, db)

    result = _confirm(db, made["pso"], actor_user_id=owner)
    db.commit()

    inquiry = result["inquiry"]
    assert inquiry is not None
    notification = (
        db.query(Notification)
        .filter(
            Notification.user_id == buyer_id,
            Notification.event_type == "project_order_inquiry_raised",
            Notification.source_entity_id == inquiry.id,
        )
        .first()
    )
    assert notification is not None
    assert notification.body.startswith(f"{made['core'].so_number}:")


def test_a_registered_project_still_gets_a_task_and_the_project_title(
    seeded, monkeypatch
):
    """AC-11: a registered project - `ProjectTask` still created, and the
    notification's body still leads with the project's own title."""
    from app.services.project_service import register_project

    db, company_id = seeded
    buyer_id = _purchasing_user(db, company_id)
    owner = _user(db, f"{MARKER} CS")
    project = register_project(
        db,
        company_id=company_id,
        actor_user_id=owner,
        developer_party_id=None,
        title=f"{MARKER} TUJU RESIDENCE",
    )
    db.flush()
    made = _order_for_confirm(
        db, company_id, project_label=PROJECT_LABEL, project_id=project.id
    )
    _patch_session_local(monkeypatch, db)

    result = _confirm(db, made["pso"], actor_user_id=owner)
    db.commit()

    inquiry = result["inquiry"]
    assert inquiry is not None
    task = ProjectOrderInquiryService(db).task_for(inquiry.id)
    assert task is not None
    assert task.project_id == project.id
    notification = (
        db.query(Notification)
        .filter(
            Notification.user_id == buyer_id,
            Notification.event_type == "project_order_inquiry_raised",
            Notification.source_entity_id == inquiry.id,
        )
        .first()
    )
    assert notification is not None
    assert notification.body.startswith(f"{MARKER} TUJU RESIDENCE:")


def test_one_notification_per_header_across_reconfirms(seeded, monkeypatch):
    """AC-12: two GENUINE handoffs on the SAME adopted header (no project, so
    `task_for` never blocks a repeat call) still land exactly one notification row -
    the dedup key `{inquiry_id}:order_inquiry_raised`, enforced both by
    `NotificationService`'s own `existing` check and, as a backstop, the DB's
    `uq_notification_user_dedup_event`.

    `_hand_to_purchasing` is called a SECOND time directly (the S1 shape its own
    docstring describes - `_write`'s own raise and a reaction writing onto the SAME
    inquiry inside the SAME apply), rather than via a second `_confirm`: a plain
    reconfirm only ever calls it once per `refresh_for_decision`, which would make
    `len == 1` true trivially (review round 1, reviewer instrumented the call count)."""
    db, company_id = seeded
    buyer_id = _purchasing_user(db, company_id)
    owner = _user(db, f"{MARKER} CS")
    made = _order_for_confirm(db, company_id, project_label=PROJECT_LABEL, project_id=None)
    _patch_session_local(monkeypatch, db)

    first = _confirm(db, made["pso"], actor_user_id=owner, buy="10")
    inquiry = first["inquiry"]
    assert inquiry is not None
    # A second, genuine handoff on the SAME header before anything commits.
    ProjectOrderInquiryService(db)._hand_to_purchasing(made["pso"], inquiry, 1)
    db.commit()

    notifications = (
        db.query(Notification)
        .filter(
            Notification.user_id == buyer_id,
            Notification.event_type == "project_order_inquiry_raised",
            Notification.source_entity_id == inquiry.id,
        )
        .all()
    )
    assert len(notifications) == 1, (
        f"AC-12: expected exactly one notification for one header, got "
        f"{len(notifications)}"
    )


# ----------------------------------- AC-13..AC-18: the Project filter follows the column


def _link_row(db, company_id: str, row: OrderInquiryRow) -> None:
    """A real PO link on the row - the minimum `unplace_all_preview`'s `linked='any'`
    scope needs to see it at all. `_adopted_order`'s own rows carry none."""
    supplier = Supplier(
        id=_uid(),
        company_id=company_id,
        supplier_code=f"ZZT-{_uid()[:8]}",
        supplier_name=f"{MARKER} supplier",
    )
    db.add(supplier)
    db.flush()
    order = PurchaseOrder(
        id=_uid(),
        company_id=company_id,
        po_number=f"ZZT-PO-{_uid()[:8]}",
        supplier_id=supplier.id,
    )
    db.add(order)
    db.flush()
    product = _product(db, f"ZZT-LNK-{_uid()[:6]}", f"{MARKER} link product")
    line = PurchaseOrderLine(
        id=_uid(),
        company_id=company_id,
        purchase_order_id=order.id,
        product_id=product.id,
        qty_ordered=Decimal("10"),
    )
    db.add(line)
    db.flush()
    db.add(
        OrderInquiryLink(
            id=_uid(),
            company_id=company_id,
            row_id=row.id,
            po_line_id=line.id,
            document=order.po_number,
            qty=row.qty,
        )
    )
    db.commit()


def _project_filter_world(db, company_id: str) -> dict:
    """The scenario AC-13..AC-17 all read: `ALPHA` (2 rows, two separate adopted
    orders, both LINKED so `unplace_all_preview` has something to count), `BETA` (1
    row, also linked), a registered project `TUJU RESIDENCE` (1 row) and a row with
    NEITHER a project nor an SO label - the option `_projects()` must never offer
    (AC-13)."""
    from app.services.project_service import register_project

    alpha_1 = _adopted_order(db, company_id, project_label=f"{MARKER} ALPHA", customer=None)
    alpha_2 = _adopted_order(db, company_id, project_label=f"{MARKER} ALPHA", customer=None)
    beta = _adopted_order(db, company_id, project_label=f"{MARKER} BETA", customer=None)
    _link_row(db, company_id, alpha_1["row"])
    _link_row(db, company_id, alpha_2["row"])
    _link_row(db, company_id, beta["row"])
    owner = _user(db, f"{MARKER} Registrar")
    project = register_project(
        db,
        company_id=company_id,
        actor_user_id=owner,
        developer_party_id=None,
        title=f"{MARKER} TUJU RESIDENCE",
    )
    db.flush()
    registered = _adopted_order(db, company_id, project_label=None, project_id=project.id)
    neither = _adopted_order(db, company_id, project_label=None, customer=None)
    return {
        "alpha": [alpha_1, alpha_2],
        "beta": beta,
        "registered": registered,
        "project": project,
        "neither": neither,
    }


def test_summary_projects_facet_lists_the_column_labels(api):
    """AC-13: the facet groups by `_PROJECT_TITLE` - the SAME text the column prints -
    not `Project.title` alone. Two adopted orders sharing a label count as two rows of
    ONE option; a registered project is a second; a row with neither is not an option
    at all; ordered by label."""
    client, db, company_id = api
    _project_filter_world(db, company_id)

    body = client.get(f"{BASE}/order-inquiries/summary").json()
    options = {opt["id"]: opt for opt in body["projects"]}

    assert len(body["projects"]) == 3
    assert options[f"{MARKER} ALPHA"]["rows"] == 2
    assert options[f"{MARKER} BETA"]["rows"] == 1
    assert options[f"{MARKER} TUJU RESIDENCE"]["rows"] == 1
    labels = [opt["label"] for opt in body["projects"]]
    assert labels == sorted(labels)


def test_list_filters_by_project_text_exact_match(api):
    """AC-14: `project=ALPHA` returns exactly the two ALPHA rows; the registered
    project's own title reaches its row the same way; wrong case matches nothing."""
    client, db, company_id = api
    world = _project_filter_world(db, company_id)
    alpha_ids = {order["row"].id for order in world["alpha"]}

    alpha_hit = client.get(
        LIST, params={"project": f"{MARKER} ALPHA", "limit": 100}
    ).json()
    registered_hit = client.get(
        LIST, params={"project": f"{MARKER} TUJU RESIDENCE", "limit": 100}
    ).json()
    wrong_case = client.get(
        LIST, params={"project": f"{MARKER} alpha", "limit": 100}
    ).json()

    assert {row["id"] for row in alpha_hit["data"]} == alpha_ids
    assert {row["id"] for row in registered_hit["data"]} == {
        world["registered"]["row"].id
    }
    assert wrong_case["data"] == []


def test_projects_facet_ignores_its_own_filter(api):
    """AC-15: with `project=ALPHA` applied, the facet still lists all three options -
    a control that emptied itself the moment it was used could not be used again."""
    client, db, company_id = api
    _project_filter_world(db, company_id)

    body = client.get(
        f"{BASE}/order-inquiries/summary", params={"project": f"{MARKER} ALPHA"}
    ).json()

    assert len(body["projects"]) == 3


def test_summary_matrix_and_export_all_honour_the_project_filter(api):
    """AC-16: summary totals, the matrix and the export all narrow to ALPHA's two rows
    with `project=ALPHA` - not just the list."""
    client, db, company_id = api
    world = _project_filter_world(db, company_id)
    alpha_item_codes = {order["product"].product_code for order in world["alpha"]}

    summary = client.get(
        f"{BASE}/order-inquiries/summary", params={"project": f"{MARKER} ALPHA"}
    ).json()
    assert summary["total_rows"] == 2

    matrix = client.get(
        f"{BASE}/order-inquiries/matrix",
        params={"axis": "product", "by": "week", "project": f"{MARKER} ALPHA"},
    ).json()
    assert sum(cell["rows"] for cell in matrix["data"]) == 2

    export = client.get(
        f"{BASE}/order-inquiries/export", params={"project": f"{MARKER} ALPHA"}
    )
    assert export.status_code == 200, export.text
    book = openpyxl.load_workbook(io.BytesIO(export.content))
    sheet = book["JAN 26"]
    item_codes = {
        sheet.cell(row=r, column=3).value for r in range(3, sheet.max_row + 1)
    }
    assert item_codes == alpha_item_codes


def test_bulk_scopes_honour_the_project_filter(write_api):
    """AC-17: the acknowledge route's `filter` body and the unplace-all-preview route
    both narrow to ALPHA's two rows with `project: 'ALPHA'` - never BETA's, even though
    BETA is linked (and so in `unplace_all_preview`'s scope) and awaiting (and so in
    `acknowledge`'s) exactly the same as ALPHA."""
    client, db, company_id = write_api
    world = _project_filter_world(db, company_id)

    preview = client.get(
        f"{BASE}/order-inquiries/unplace-all-preview",
        params={"project": f"{MARKER} ALPHA"},
    ).json()
    assert preview["count"] == 2

    result = client.post(
        f"{BASE}/order-inquiries/acknowledge",
        json={"filter": {"project": f"{MARKER} ALPHA"}},
    )
    assert result.status_code == 200, result.text
    assert result.json()["acknowledged"] == 2

    #: The apply half of the SAME scope (`POST /order-inquiries/unplace-all`), not just
    #: its preview: a kill test that dropped `project=payload.project` from the apply
    #: route left this file green, because nothing else here calls the apply route with
    #: the project filter in play. ALPHA's two linked rows come back; BETA's own link -
    #: linked the same way, in scope for an unfiltered call - must survive untouched.
    beta_link = (
        db.query(OrderInquiryLink)
        .filter(OrderInquiryLink.row_id == world["beta"]["row"].id)
        .first()
    )
    assert beta_link is not None

    unplace_result = client.post(
        f"{BASE}/order-inquiries/unplace-all",
        json={"project": f"{MARKER} ALPHA"},
    )
    assert unplace_result.status_code == 200, unplace_result.text
    assert unplace_result.json()["unplaced"] == 2

    db.expire_all()
    beta_link_after = (
        db.query(OrderInquiryLink)
        .filter(OrderInquiryLink.row_id == world["beta"]["row"].id)
        .first()
    )
    assert beta_link_after is not None


def test_project_id_filter_still_works_by_uuid(api):
    """AC-18: `project_id=<uuid>` is unchanged - it still filters by the REGISTERED
    project, unaffected by `project` existing beside it."""
    client, db, company_id = api
    world = _project_filter_world(db, company_id)

    body = client.get(LIST, params={"project_id": world["project"].id}).json()

    assert {row["id"] for row in body["data"]} == {world["registered"]["row"].id}


def test_project_id_filter_still_rejects_a_non_uuid(api):
    """AC-18: a malformed `project_id` is still refused the way it is today -
    `validate_uuid_path` answers 404, never a 500 carrying a Postgres statement."""
    client, db, company_id = api
    _project_filter_world(db, company_id)

    response = client.get(LIST, params={"project_id": "not-a-uuid"})

    assert response.status_code == 404, response.text


def test_project_filter_carries_no_length_bound(api):
    """AC-21: `projects.title` is TEXT with no maximum, so a bounded `project` param
    would let the facet offer a title long enough for the filter itself to 422. A
    300-char value is still a well-formed request - it matches nothing and comes back
    200 with zero rows, never a validation error."""
    client, db, company_id = api
    _project_filter_world(db, company_id)

    response = client.get(LIST, params={"project": "A" * 300, "limit": 100})

    assert response.status_code == 200, response.text
    assert response.json()["data"] == []
