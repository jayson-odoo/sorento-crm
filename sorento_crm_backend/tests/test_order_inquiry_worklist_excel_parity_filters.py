"""S1 - the five new filters and the widened search box (R-K, AC-F1 to AC-F9).

`PLAN-scm-oi-worklist-excel-parity.md` S1, `scm-oi-worklist-excel-parity-acceptance-criteria.md`
S1. None of `location`, `agent`, `so_month`, `po_number`, `spo_number`, `delivery_from`,
`delivery_to` are declared on the route today - FastAPI drops an undeclared query param
rather than erroring on it, so every assertion here is "the filter actually narrows",
never a 422. That is also why every seed pairs a row that SHOULD survive a filter with one
that SHOULD NOT: a filter that is silently ignored returns both, which is the honest red
state before Phase 2 wiring.

Postgres only, via `tests/_pg_fixture.py::blank_session`. Every row and identifier carries
`MARKER` so nothing here can be confused with another test's fixture.
"""
from __future__ import annotations

import io
import uuid
from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest
from sqlalchemy import text

from app.models.inventory import Warehouse
from app.models.order import SalesOrder
from app.models.procurement import (
    InboundShipment,
    PurchaseOrder,
    PurchaseOrderLine,
    SPOAllocation,
    Supplier,
)
from app.models.product import Product, ProductCategory, UnitOfMeasure
from app.models.project_so import (
    INQUIRY_RAISED,
    IV_ORDER,
    OrderInquiry,
    OrderInquiryLink,
    OrderInquiryRow,
    ProjectSalesOrder,
    ProjectSalesOrderLine,
)
from app.models.sales_agent import SalesAgent
from app.models.user import User
from app.services import project_seed_service

from ._pg_fixture import blank_session

MARKER = "zzt-oi-parity-filters"
BASE = "/api/v1/project-sales"
LIST = f"{BASE}/order-inquiries"
SUMMARY = f"{LIST}/summary"
EXPORT = f"{LIST}/export"

READ_ONLY = ["projects.projects.view"]


def _uid() -> str:
    return str(uuid.uuid4())


def _sorento(db) -> str:
    return db.execute(text("select id from companies where code = 'SRT'")).scalar()


def _user(db, name: str) -> str:
    user_id = _uid()
    db.add(User(id=user_id, email=f"{user_id}@zzt.test", name=name))
    db.flush()
    return user_id


def _product(db, code: str, name: str) -> Product:
    uom = UnitOfMeasure(id=_uid(), uom_code=f"ZZT{_uid()[:6]}", uom_name="Unit")
    category = ProductCategory(
        id=_uid(), category_code=f"ZZT-{_uid()[:8]}", category_name=f"{MARKER} cat"
    )
    db.add_all([uom, category])
    db.flush()
    row = Product(
        id=_uid(),
        product_code=code,
        product_name=name,
        category_id=category.id,
        base_uom_id=uom.id,
        list_price=Decimal("100.00"),
    )
    db.add(row)
    db.flush()
    return row


def _supplier(db, company_id: str, name: str) -> Supplier:
    row = Supplier(
        id=_uid(), company_id=company_id, supplier_code=f"ZZT-{_uid()[:8]}",
        supplier_name=name,
    )
    db.add(row)
    db.flush()
    return row


def _agent(db, company_id: str, code: str, label: str) -> SalesAgent:
    row = SalesAgent(id=_uid(), company_id=company_id, sales_agent=code, person_label=label)
    db.add(row)
    db.flush()
    return row


def _client(db, user_id: str, permissions):
    from fastapi.testclient import TestClient

    from app.database import get_db
    from app.dependencies import get_current_user, get_current_user_or_api_key
    from app.main import app
    from app.services.company_scope_resolver import apply_company_scope
    from app.services.user_service import UserPermissionService

    actor = {"id": user_id, "email": f"{user_id}@zzt.test", "role": "user"}
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: dict(actor)
    app.dependency_overrides[get_current_user_or_api_key] = lambda: dict(actor)
    app.dependency_overrides[apply_company_scope] = lambda: None

    originals = (
        UserPermissionService.check_user_has_permission,
        UserPermissionService.get_user_permission_slugs,
    )
    granted = list(permissions)
    UserPermissionService.check_user_has_permission = (
        lambda self, uid, slug: slug in granted
    )
    UserPermissionService.get_user_permission_slugs = lambda self, uid: list(granted)
    return TestClient(app), originals


def _restore(originals) -> None:
    from app.main import app
    from app.services.user_service import UserPermissionService

    UserPermissionService.check_user_has_permission = originals[0]
    UserPermissionService.get_user_permission_slugs = originals[1]
    app.dependency_overrides.clear()


def _inquiry_for(db, company_id: str, order: ProjectSalesOrder) -> OrderInquiry:
    inquiry = OrderInquiry(
        id=_uid(), company_id=company_id, project_sales_order_id=order.id,
        state=INQUIRY_RAISED,
    )
    db.add(inquiry)
    db.flush()
    return inquiry


def _authored_pso(db, company_id: str, *, published_at: datetime) -> ProjectSalesOrder:
    """No `project_id`, no `so_id` - the minimum a `order_inquiries` row can hang off.
    `_base` INNER joins `ProjectSalesOrder`, never `Project`, so this is a legal shape."""
    order = ProjectSalesOrder(
        id=_uid(), company_id=company_id, project_id=None, so_id=None,
        provisional_ref=f"ZZT-PSO-{_uid()[:8]}", status="draft",
        published_at=published_at,
    )
    db.add(order)
    db.flush()
    return order


def _adopted_pso(db, company_id: str, core: SalesOrder) -> ProjectSalesOrder:
    order = ProjectSalesOrder(
        id=_uid(), company_id=company_id, project_id=None, so_id=core.id,
        provisional_ref=core.so_number, autocount_doc_no=core.so_number,
        status="adopted",
    )
    db.add(order)
    db.flush()
    return order


def _pso_line(db, company_id: str, pso: ProjectSalesOrder, product: Product, **overrides):
    line = ProjectSalesOrderLine(
        id=_uid(), company_id=company_id, project_sales_order_id=pso.id, line_no=1,
        product_id=product.id, description=f"{MARKER} line", qty=Decimal("10"),
        uom="UNIT", unit_price=Decimal("10.00"), amount=Decimal("100.00"),
        delivery_date=overrides.pop("delivery_date", date(2026, 6, 1)),
        **overrides,
    )
    db.add(line)
    db.flush()
    return line


def _row(db, company_id: str, inquiry: OrderInquiry, *, so_line_id=None, item_code: str,
          qty="10", delivery_date=None, stock_location=None) -> OrderInquiryRow:
    row = OrderInquiryRow(
        id=_uid(), company_id=company_id, order_inquiry_id=inquiry.id,
        so_line_id=so_line_id, item_code=item_code, qty=Decimal(str(qty)),
        delivery_date=delivery_date, verb=IV_ORDER, state=INQUIRY_RAISED,
        ack_state="acknowledged", stock_location=stock_location,
    )
    db.add(row)
    db.flush()
    return row


def _po_link(db, company_id: str, row: OrderInquiryRow, *, po_number: str,
              product: Product, document: str | None = None, qty="10") -> PurchaseOrder:
    supplier = _supplier(db, company_id, f"{MARKER} supplier {po_number}")
    po = PurchaseOrder(id=_uid(), company_id=company_id, po_number=po_number,
                        supplier_id=supplier.id)
    db.add(po)
    db.flush()
    line = PurchaseOrderLine(id=_uid(), company_id=company_id, purchase_order_id=po.id,
                              product_id=product.id, qty_ordered=Decimal("10"))
    db.add(line)
    db.flush()
    db.add(OrderInquiryLink(id=_uid(), company_id=company_id, row_id=row.id,
                             po_line_id=line.id, document=document or po_number,
                             qty=Decimal(str(qty))))
    db.flush()
    return po


def _spo_link(db, company_id: str, row: OrderInquiryRow, *, spo_number: str,
               product: Product, from_po_number: str | None = None, qty="10") -> SPOAllocation:
    allocation = SPOAllocation(
        id=_uid(), company_id=company_id, spo_number=spo_number,
        allocated_quantity=int(Decimal(str(qty))), product_id=product.id,
        from_po_number=from_po_number,
    )
    db.add(allocation)
    db.flush()
    db.add(OrderInquiryLink(id=_uid(), company_id=company_id, row_id=row.id,
                             spo_allocation_id=allocation.id, document=spo_number,
                             qty=Decimal(str(qty))))
    db.flush()
    return allocation


def _open_derived_spo(db, company_id: str, *, spo_number: str, po_number: str,
                        product: Product, landed: bool = False) -> SPOAllocation:
    """An SPOAllocation that answers to `spo_number`/`linked=spo` for a PO-linked row's
    OWN PO purely by derivation (S5, R-E) - `from_po_number == po_number AND
    product_id == product.id`, open per `spo_supply.open_incoming_clauses` - never
    itself linked to any row. `landed=True` books it on a shipment that has arrived, so
    the coordinator's addendum case (a landed shipment is not incoming) can be built
    with one flag."""
    shipment_id = None
    if landed:
        shipment = InboundShipment(
            id=_uid(), company_id=company_id, shipment_number=f"ZZT-{_uid()[:8]}",
            shipment_date=date.today(), actual_arrival_date=date.today(),
        )
        db.add(shipment)
        db.flush()
        shipment_id = shipment.id
    allocation = SPOAllocation(
        id=_uid(), company_id=company_id, spo_number=spo_number, allocated_quantity=5,
        quantity_received=0, product_id=product.id, from_po_number=po_number,
        line_status="open", receipt_status="pending", inbound_shipment_id=shipment_id,
    )
    db.add(allocation)
    db.flush()
    return allocation


def _core_order_with_agent(db, company_id: str, agent: SalesAgent | None) -> SalesOrder:
    order = SalesOrder(
        id=_uid(), company_id=company_id, so_number=f"ZZTSO{_uid()[:8]}",
        order_date=date(2026, 5, 1),
        sales_agent_id=agent.id if agent else None,
    )
    db.add(order)
    db.flush()
    return order


def _seed(db, company_id: str, user_id: str) -> dict:
    product = _product(db, f"ZZT-PROD-{_uid()[:6]}", f"{MARKER} product")

    # -- location + agent: two rows, each stamped with ITS OWN location and reached
    # through a core sales order carrying ITS OWN agent, so `location=`, `agent=` and
    # the facets can all be pinned off the same pair.
    agent_x = _agent(db, company_id, f"ZZT{_uid()[:6].upper()}", f"{MARKER} Zorro Agent")
    agent_y = _agent(db, company_id, f"ZZT{_uid()[:6].upper()}", f"{MARKER} Yusof Agent")
    core_x = _core_order_with_agent(db, company_id, agent_x)
    core_y = _core_order_with_agent(db, company_id, agent_y)
    pso_x = _adopted_pso(db, company_id, core_x)
    pso_y = _adopted_pso(db, company_id, core_y)
    line_x = _pso_line(db, company_id, pso_x, product)
    line_y = _pso_line(db, company_id, pso_y, product)
    inquiry_x = _inquiry_for(db, company_id, pso_x)
    inquiry_y = _inquiry_for(db, company_id, pso_y)
    row_agent_x_loc_a = _row(
        db, company_id, inquiry_x, so_line_id=line_x.id,
        item_code=f"{MARKER}-AGENTX", stock_location="ZZTLOCA",
    )
    row_agent_y_loc_b = _row(
        db, company_id, inquiry_y, so_line_id=line_y.id,
        item_code=f"{MARKER}-AGENTY", stock_location="ZZTLOCB",
    )

    # -- so_month: two authored rows, dated only through `published_at` (no core order,
    # no `SalesOrder.order_date` at all).
    pso_jan = _authored_pso(db, company_id, published_at=datetime(2026, 1, 10, 9, 0))
    pso_mar = _authored_pso(db, company_id, published_at=datetime(2026, 3, 10, 9, 0))
    inquiry_jan = _inquiry_for(db, company_id, pso_jan)
    inquiry_mar = _inquiry_for(db, company_id, pso_mar)
    row_so_jan = _row(db, company_id, inquiry_jan, item_code=f"{MARKER}-SOJAN")
    row_so_mar = _row(db, company_id, inquiry_mar, item_code=f"{MARKER}-SOMAR")

    # -- po_number / spo_number / search-hits-link-document.
    pso_docs = _authored_pso(db, company_id, published_at=datetime(2026, 6, 1, 9, 0))
    inquiry_docs = _inquiry_for(db, company_id, pso_docs)
    row_po_match = _row(db, company_id, inquiry_docs, item_code=f"{MARKER}-POMATCH")
    _po_link(db, company_id, row_po_match, po_number="202605-S0009", product=product)
    row_po_other = _row(db, company_id, inquiry_docs, item_code=f"{MARKER}-POOTHER")
    _po_link(db, company_id, row_po_other, po_number="999999-S0001", product=product)
    row_spo_via_po = _row(db, company_id, inquiry_docs, item_code=f"{MARKER}-SPOVIAPO")
    _spo_link(
        db, company_id, row_spo_via_po, spo_number=f"{MARKER}-SPO-VIAPO",
        product=product, from_po_number="202605-ABCD",
    )
    row_spo_match = _row(db, company_id, inquiry_docs, item_code=f"{MARKER}-SPOMATCH")
    _spo_link(
        db, company_id, row_spo_match, spo_number="SPO-2026/07-0005", product=product,
    )

    # -- delivery_from / delivery_to.
    pso_delivery = _authored_pso(db, company_id, published_at=datetime(2026, 6, 1, 9, 0))
    inquiry_delivery = _inquiry_for(db, company_id, pso_delivery)
    row_delivery_in = _row(
        db, company_id, inquiry_delivery, item_code=f"{MARKER}-DELIVIN",
        delivery_date=date(2026, 7, 10),
    )
    row_delivery_out = _row(
        db, company_id, inquiry_delivery, item_code=f"{MARKER}-DELIVOUT",
        delivery_date=date(2026, 9, 10),
    )

    db.commit()
    return {
        "product": product,
        "agent_x": agent_x,
        "agent_y": agent_y,
        "row_agent_x_loc_a": row_agent_x_loc_a,
        "row_agent_y_loc_b": row_agent_y_loc_b,
        "row_so_jan": row_so_jan,
        "row_so_mar": row_so_mar,
        "row_po_match": row_po_match,
        "row_po_other": row_po_other,
        "row_spo_via_po": row_spo_via_po,
        "row_spo_match": row_spo_match,
        "row_delivery_in": row_delivery_in,
        "row_delivery_out": row_delivery_out,
    }


@pytest.fixture()
def api():
    from app.models.base import company_scope

    with blank_session() as db:
        company_id = _sorento(db)
        project_seed_service.run(db, company_id=company_id)
        user_id = _user(db, f"{MARKER} tester")
        seeded = _seed(db, company_id, user_id)
        client, originals = _client(db, user_id, READ_ONLY)
        try:
            with company_scope(db, frozenset({company_id})):
                yield client, db, company_id, seeded
        finally:
            _restore(originals)


# ------------------------------------------------------------------ AC-F2 location


def test_location_filter_returns_only_that_locations_rows(api):
    client, _db, _company_id, seeded = api

    body = client.get(LIST, params={"location": "ZZTLOCA"}).json()

    ids = {row["id"] for row in body["data"]}
    assert seeded["row_agent_x_loc_a"].id in ids
    assert seeded["row_agent_y_loc_b"].id not in ids
    assert all(row["location"] == "ZZTLOCA" for row in body["data"])


# --------------------------------------------------------------------- AC-F3 agent


def test_agent_filter_returns_only_that_agents_rows(api):
    client, _db, _company_id, seeded = api

    body = client.get(LIST, params={"agent": seeded["agent_x"].id}).json()

    ids = {row["id"] for row in body["data"]}
    assert ids == {seeded["row_agent_x_loc_a"].id}


# ------------------------------------------------------------------- AC-F4 so_month


def test_so_month_filter_narrows_by_the_so_date(api):
    client, _db, _company_id, seeded = api

    march = client.get(LIST, params={"so_month": "2026-03"}).json()

    ids = {row["id"] for row in march["data"]}
    assert seeded["row_so_mar"].id in ids
    assert seeded["row_so_jan"].id not in ids


# ------------------------------------------------------------------ AC-F5 po_number


def test_po_number_prefix_matches_a_real_po_links_document(api):
    client, _db, _company_id, seeded = api

    body = client.get(LIST, params={"po_number": "202605"}).json()

    ids = {row["id"] for row in body["data"]}
    assert seeded["row_po_match"].id in ids
    assert seeded["row_po_other"].id not in ids


def test_po_number_prefix_also_matches_an_spo_links_source_po_number(api):
    """AC-F5b: the row's only link is an SPO allocation, but that allocation names
    the PO it was raised FROM (`from_po_number`), and `po_number` has to reach it."""
    client, _db, _company_id, seeded = api

    body = client.get(LIST, params={"po_number": "202605"}).json()

    ids = {row["id"] for row in body["data"]}
    assert seeded["row_spo_via_po"].id in ids
    assert seeded["row_po_other"].id not in ids


# ----------------------------------------------------------------- AC-F6 spo_number


def test_spo_number_prefix_matches_case_insensitively(api):
    client, _db, _company_id, seeded = api

    body = client.get(LIST, params={"spo_number": "spo-2026/07"}).json()

    ids = {row["id"] for row in body["data"]}
    assert ids == {seeded["row_spo_match"].id}


def test_spo_number_prefix_also_matches_a_derived_spo_on_the_rows_po(api):
    """AC-F6b: the row's only REAL link is a purchase order, but that PO carries an
    open SPO allocation for the same product (S5's derived SPO) - `spo_number` has to
    reach it, the same way `po_number` reaches an SPO link's `source_po_number`
    (AC-F5b). Not when the allocation's shipment has already landed."""
    client, db, company_id, seeded = api

    product = seeded["product"]
    inquiry_docs = db.get(OrderInquiryRow, seeded["row_po_match"].id).order_inquiry_id
    row_po_only = _row(
        db, company_id, db.get(OrderInquiry, inquiry_docs),
        item_code=f"{MARKER}-DERIVEDSPO",
    )
    po = _po_link(
        db, company_id, row_po_only, po_number=f"ZZT-DERIVED-PO-{_uid()[:6]}",
        product=product,
    )
    _open_derived_spo(
        db, company_id, spo_number="SPO-2026/09-0099", po_number=po.po_number,
        product=product,
    )
    row_landed = _row(
        db, company_id, db.get(OrderInquiry, inquiry_docs),
        item_code=f"{MARKER}-DERIVEDSPO-LANDED",
    )
    po_landed = _po_link(
        db, company_id, row_landed, po_number=f"ZZT-DERIVED-PO-LANDED-{_uid()[:6]}",
        product=product,
    )
    _open_derived_spo(
        db, company_id, spo_number="SPO-2026/09-0100", po_number=po_landed.po_number,
        product=product, landed=True,
    )
    db.commit()

    body = client.get(LIST, params={"spo_number": "SPO-2026/09"}).json()

    ids = {row["id"] for row in body["data"]}
    assert row_po_only.id in ids
    assert row_landed.id not in ids


# ------------------------------------------------------------------------ AC-F7 query


def test_the_search_box_matches_a_link_document(api):
    client, _db, _company_id, seeded = api

    body = client.get(LIST, params={"query": "202605-S0009"}).json()

    ids = {row["id"] for row in body["data"]}
    assert ids == {seeded["row_po_match"].id}


def test_the_search_box_matches_the_sales_agents_name(api):
    client, _db, _company_id, seeded = api

    body = client.get(LIST, params={"query": "Zorro"}).json()

    ids = {row["id"] for row in body["data"]}
    assert ids == {seeded["row_agent_x_loc_a"].id}


# --------------------------------------------------------------------- AC-F8 summary


def test_summary_carries_locations_and_agents_facets(api):
    client, _db, _company_id, seeded = api

    body = client.get(SUMMARY).json()

    assert "locations" in body, body.keys()
    assert "agents" in body, body.keys()
    location_ids = {entry["id"] for entry in body["locations"]}
    assert {"ZZTLOCA", "ZZTLOCB"} <= location_ids
    agent_ids = {entry["id"] for entry in body["agents"]}
    assert {seeded["agent_x"].id, seeded["agent_y"].id} <= agent_ids


def test_the_locations_facet_drops_its_own_filter(api):
    client, _db, _company_id, _seeded = api

    body = client.get(SUMMARY, params={"location": "ZZTLOCA"}).json()

    location_ids = {entry["id"] for entry in body["locations"]}
    assert {"ZZTLOCA", "ZZTLOCB"} <= location_ids


def test_the_locations_facet_honours_the_agent_filter(api):
    client, _db, _company_id, seeded = api

    body = client.get(SUMMARY, params={"agent": seeded["agent_x"].id}).json()

    location_ids = {entry["id"] for entry in body["locations"]}
    assert location_ids == {"ZZTLOCA"}


# ---------------------------------------------------------------------- AC-F9 export


def test_export_honours_the_location_filter(api):
    client, _db, _company_id, seeded = api

    response = client.get(EXPORT, params={"location": "ZZTLOCA"})

    book = openpyxl.load_workbook(io.BytesIO(response.content))
    item_codes = {
        cell[2].value
        for sheet in book.worksheets
        for cell in sheet.iter_rows(min_row=3)
    }
    assert seeded["row_agent_x_loc_a"].item_code in item_codes
    assert seeded["row_agent_y_loc_b"].item_code not in item_codes


def test_export_carries_row_level_taken_and_remaining_matching_the_grid(api):
    """Fix round (22 Sep, `PLAN-board-oi-mechanical-22sep.md`, AC-D15 parity): the export
    prints the SAME row-level Taken/Remaining the grid has shown since S3
    (`inquiryRowTaken`/`inquiryRowRemaining`), never the retired line-scoped
    `taken_from_po`/`remaining_open` pair. `row_po_match` (S1 seed above) is a plain
    `ORDER` row, qty 10, with one PO link for the full 10 - fully taken, nothing left."""
    client, _db, _company_id, seeded = api

    response = client.get(EXPORT)
    book = openpyxl.load_workbook(io.BytesIO(response.content))

    headers = None
    target_row = None
    for sheet in book.worksheets:
        sheet_headers = next(
            sheet.iter_rows(min_row=2, max_row=2, values_only=True), None
        )
        if sheet_headers is None:
            continue
        for values in sheet.iter_rows(min_row=3, values_only=True):
            if values[2] == seeded["row_po_match"].item_code:
                headers = sheet_headers
                target_row = values
                break
        if target_row is not None:
            break

    assert target_row is not None, "row_po_match not found in any export sheet"
    # Review round (22 Sep): APPENDED after ACKNOWLEDGED, never inserted before it - the
    # columns purchasing's own filters already point at keep the positions they have had.
    assert headers[10] == "ACKNOWLEDGED", headers
    assert headers[11] == "Taken", headers
    assert headers[12] == "Remaining", headers
    assert target_row[11] == "10", target_row
    assert target_row[12] == "0", target_row


# --------------------------------------------------------------- delivery range


def test_delivery_from_and_delivery_to_narrow_the_list_inclusively(api):
    client, _db, _company_id, seeded = api

    body = client.get(
        LIST, params={"delivery_from": "2026-07-01", "delivery_to": "2026-07-31"}
    ).json()

    ids = {row["id"] for row in body["data"]}
    assert ids == {seeded["row_delivery_in"].id}
    assert seeded["row_delivery_out"].id not in ids
