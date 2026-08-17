"""Reading an AutoCount sales order export into an ``IngestDocument`` (P8a, stage 1).

Stage 1's transport is a file: a CS exports the sales order from AutoCount and uploads it.
Its exact layout has not been seen yet - the open risk recorded in
`PLAN-project-so-divergence.md` - so nothing here depends on a column sitting in a
particular place.

**Everything is read by heading, with synonyms.** A column that moved, or that AutoCount
labels `Stock Code` where our own file says `Item`, must not silently shift every quantity
one place left, which is the failure mode of a positional reader and the one that produces
a divergence report full of differences nobody made.

**A row is a line only if it names a product and carries a number.** Section headers,
blank spacers, the `***TOWER***` area marker and the trailing `Total` row all live in the
same table on the real document.

The canonical JSON route does not come through here at all. When the ESB lands in stage 2
this file simply stops being the only way in.
"""
from __future__ import annotations

import csv
import io
import logging
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Sequence

from app.services.error_handler import AppException
from app.services.project_so_ingest_service import IngestDocument, IngestLine

logger = logging.getLogger(__name__)

# Column headings, normalised to lower case with punctuation stripped.
_LINE_HEADINGS: Dict[str, Sequence[str]] = {
    "product_code": ("item", "item code", "stock code", "product code", "code"),
    "description": ("description", "desc", "item description"),
    "qty": ("qty", "quantity", "order qty", "so qty"),
    "unit_price": ("uprice", "unit price", "price", "u price"),
    "delivery_date": ("delivery date", "delivery", "date", "req delivery date"),
    "uom": ("uom", "unit", "u o m"),
    "amount": ("total", "amount", "line total", "subtotal"),
}

# Header rows: `label, value` pairs above the line table.
_HEADER_LABELS: Dict[str, Sequence[str]] = {
    "doc_no": ("doc no", "document no", "so no", "s o no", "sales order no", "docno"),
    "customer_po_no": (
        "your ref no",
        "customer po",
        "cust po no",
        "po no",
        "customer po no",
        "your ref",
    ),
    "customer_code": ("debtor", "debtor code", "customer", "customer code"),
    "terms": ("terms", "term", "payment terms"),
    "area_group": ("area", "area group", "section"),
    "total_amount": ("total", "document total", "grand total", "net total"),
}

_AREA_MARKER = re.compile(r"^\*+\s*(?P<label>.+?)\s*\*+$")
_NUMERIC_NOISE = re.compile(r"[^0-9.\-]")

_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y", "%Y/%m/%d", "%d %b %Y")


def _norm(value: Any) -> str:
    """Lower case, punctuation dropped, so `U/Price` and `U Price` are one heading."""
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = re.sub(r"[^a-z0-9 ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _cell(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).strip()
    return text or None


def _decimal(value: Any) -> Optional[Decimal]:
    """`RM 12.50` and `1,810` are numbers a person typed, not text."""
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    cleaned = _NUMERIC_NOISE.sub("", str(value).strip())
    if not cleaned or cleaned in ("-", "."):
        return None
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return None


def _date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    logger.info("so ingest: unreadable delivery date %r, left empty", text)
    return None


def _rows_from_csv(payload: bytes) -> List[List[Any]]:
    text = payload.decode("utf-8-sig", errors="replace")
    return [list(row) for row in csv.reader(io.StringIO(text))]


def _rows_from_xlsx(payload: bytes) -> List[List[Any]]:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(payload), data_only=True)
    sheet = workbook.active
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def _heading_map(row: Sequence[Any]) -> Dict[int, str]:
    """Which column is which, or an empty map when this row is not the heading row."""
    found: Dict[int, str] = {}
    for index, cell in enumerate(row):
        label = _norm(cell)
        if not label:
            continue
        for field, synonyms in _LINE_HEADINGS.items():
            if field in found.values():
                continue
            if label in (_norm(s) for s in synonyms):
                found[index] = field
                break
    # A heading row has to name the product column and at least one number, or a
    # two-column `Terms, *Net 60 days` header row would qualify as a line table.
    fields = set(found.values())
    if "product_code" in fields and fields & {"qty", "unit_price", "amount"}:
        return found
    return {}


def parse_document(payload: bytes, *, filename: str) -> IngestDocument:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        rows = _rows_from_xlsx(payload)
    elif name.endswith((".csv", ".txt")):
        rows = _rows_from_csv(payload)
    else:
        raise AppException(
            status_code=422,
            message=(
                "Upload the AutoCount export as .csv or .xlsx. A PDF or a scan cannot be "
                "compared line by line."
            ),
            code="so_ingest_unsupported_file",
        )

    document = IngestDocument()
    heading: Dict[int, str] = {}
    header_values: Dict[str, Any] = {}
    line_total_candidates: List[Decimal] = []

    for row in rows:
        cells = [_cell(cell) for cell in row]
        if not any(cells):
            continue

        if not heading:
            found = _heading_map(row)
            if found:
                heading = found
                continue
            marker = _AREA_MARKER.match((cells[0] or "").strip()) if cells[0] else None
            if marker:
                document.area_group = marker.group("label")
                continue
            _read_header_pair(cells, header_values)
            continue

        marker = _AREA_MARKER.match((cells[0] or "").strip()) if cells[0] else None
        if marker:
            document.area_group = marker.group("label")
            continue

        line = _read_line(row, heading)
        if line is not None:
            document.lines.append(line)
            continue
        # Below the table, `Total` and any late header pair are still worth reading.
        _read_header_pair(cells, header_values)

    if not heading:
        raise AppException(
            status_code=422,
            message=(
                "No line table was found in that file: none of its rows carry a product "
                "heading (Item / Stock Code) beside a quantity or price heading."
            ),
            code="so_ingest_no_heading",
        )

    document.doc_no = header_values.get("doc_no")
    document.customer_po_no = header_values.get("customer_po_no")
    document.customer_code = header_values.get("customer_code")
    document.terms = header_values.get("terms")
    document.area_group = header_values.get("area_group") or document.area_group
    document.total_amount = _decimal(header_values.get("total_amount"))
    return document


def _read_header_pair(cells: Sequence[Optional[str]], into: Dict[str, Any]) -> None:
    """`Your Ref No.,PO-778` - the label in one cell, the value in the next non-empty one."""
    label = _norm(cells[0] if cells else None)
    if not label:
        # A `,Total,,,7500.00` row puts the label in the second column.
        label = _norm(cells[1]) if len(cells) > 1 else ""
        rest = [cell for cell in cells[2:] if cell]
    else:
        rest = [cell for cell in cells[1:] if cell]
    if not label or not rest:
        return
    for field, synonyms in _HEADER_LABELS.items():
        if label in (_norm(s) for s in synonyms):
            # First writer wins for everything except the total, where the LAST row is
            # the document total: a per-section subtotal prints above it.
            if field not in into or field == "total_amount":
                into[field] = rest[-1] if field == "total_amount" else rest[0]
            return


def _read_line(row: Sequence[Any], heading: Dict[int, str]) -> Optional[IngestLine]:
    values: Dict[str, Any] = {}
    for index, field in heading.items():
        if index < len(row):
            values[field] = row[index]

    code = _cell(values.get("product_code"))
    qty = _decimal(values.get("qty"))
    unit_price = _decimal(values.get("unit_price"))
    if not code:
        return None
    # `Total` sitting in the item column is the footer, not a product.
    if _norm(code) in ("total", "grand total", "sub total", "subtotal"):
        return None
    if qty is None and unit_price is None:
        return None

    return IngestLine(
        product_code=code,
        description=_cell(values.get("description")),
        qty=qty if qty is not None else Decimal("0"),
        unit_price=unit_price if unit_price is not None else Decimal("0"),
        uom=_cell(values.get("uom")),
        delivery_date=_date(values.get("delivery_date")),
    )
