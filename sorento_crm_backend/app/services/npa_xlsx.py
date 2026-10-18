"""Download Excel for a new product application (NEW-PRODUCT-APP v4, AC-NPA-77).

One sheet per product, laid out like the company's paper form: column A the label, column B
the value, rows in the paper's order. Staff and portal download the same workbook.
Values go through `_xlsx_safe_text` so free text never reads as a formula.
"""
from __future__ import annotations

import io
import re
from datetime import datetime, timedelta
from typing import List, Optional, Tuple

from sqlalchemy.orm import Session

from app.models.new_product_application import NewProductApplication, NewProductApplicationLine
from app.models.order import Customer
from app.models.product import Brand
from app.models.entity_attachment import EntityAttachmentLink
from app.models.resources import Attachment
from app.models.user import User
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from app.services.scm.proforma_invoice_service import _xlsx_safe_text

#: Larger images show their file name instead (memory, and Pillow's own bomb guard sits higher).
MAX_IMAGE_PIXELS = 40_000_000
_THUMB_PX = 200
_NAME_CAP = 200
_CUSTOMER_CAP = 60
_FORBIDDEN_IN_NAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
_MYT = timedelta(hours=8)


def _myt(value: Optional[datetime]) -> Optional[datetime]:
    """created_at is stored as naive UTC; Malaysia has no daylight saving."""
    return value + _MYT if value else None


def _tidy(text: Optional[str]) -> str:
    return " ".join(_FORBIDDEN_IN_NAME.sub("", text or "").split())


def _contact_name(db: Session, contact_id: Optional[str]) -> Optional[str]:
    from app.models.access import RespondContact

    row = db.query(RespondContact).filter(RespondContact.id == str(contact_id)).first() if contact_id else None
    if row is None:
        return None
    return (row.name or "").strip() or " ".join(p for p in [row.first_name, row.last_name] if p).strip() or None


def _sales_person(db: Session, app: NewProductApplication) -> Optional[str]:
    if app.source == "portal":
        name = _contact_name(db, app.contact_id)
        if name:
            return name
    return db.query(User.name).filter(User.id == app.applicant_id).scalar() if app.applicant_id else None


def _customer(db: Session, app: NewProductApplication) -> str:
    if app.customer_name:
        return app.customer_name
    if app.customer_id:
        return db.query(Customer.customer_name).filter(Customer.id == app.customer_id).scalar() or ""
    return ""


def _initials(db: Session, app: NewProductApplication) -> str:
    """Portal-raised: the application number (the contact is not a staff salesperson)."""
    if app.source != "portal" and app.applicant_id:
        name = db.query(User.name).filter(User.id == app.applicant_id).scalar() or ""
        letters = "".join(word[0] for word in _tidy(name).split()).upper()
        if letters:
            return letters
    return app.application_no


def file_stem(db: Session, app: NewProductApplication) -> str:
    stamp = (_myt(app.created_at) or datetime.utcnow()).strftime("%Y%m%d")
    parts = [f"NPA_{stamp}_{_tidy(_initials(db, app))}"]
    customer = _tidy(_customer(db, app))[:_CUSTOMER_CAP].strip()
    if customer:
        parts.append(customer)
    products = [
        f"({_tidy(' '.join(p for p in (line.proposed_code, line.description) if p))})"
        for line in app.lines
        if line.proposed_code or line.description
    ]
    stem = "_".join(parts) + (" " + " ".join(products) if products else "")
    return stem[:_NAME_CAP].strip()


def _thumbnail_or_original(backend, row) -> bytes:
    """The stored grid thumbnail when there is one (small, already decoded once); the original
    only when the upload never got one."""
    from app.services.image_thumbnailer import thumbnail_key_for

    try:
        return backend.download_file(thumbnail_key_for(row.file_path))
    except Exception:  # noqa: BLE001 - no thumbnail stored
        from app.services.new_product_application_service import MAX_IMAGE_BYTES

        if (row.file_size_bytes or 0) > MAX_IMAGE_BYTES:
            raise ValueError("original too large to embed")  # caller shows the file name
        return backend.download_file(row.file_path)


def _first_image(db: Session, line: NewProductApplicationLine) -> Optional[Tuple[str, Optional[bytes]]]:
    """(file name, thumbnail PNG bytes or None) of the product's first image."""
    from PIL import Image

    from app.services import storage_router

    row = (
        db.query(Attachment)
        .join(EntityAttachmentLink, EntityAttachmentLink.attachment_id == Attachment.id)
        .filter(
            EntityAttachmentLink.entity_type == "new_product_application_line",
            EntityAttachmentLink.entity_id == str(line.id),
        )
        .order_by(EntityAttachmentLink.sort_order.asc().nulls_last(), EntityAttachmentLink.created_at.asc())
        .first()
    )
    if row is None:
        return None
    try:
        provider = row.storage_provider or storage_router.default_provider()
        data = _thumbnail_or_original(storage_router.get_backend(provider), row)
        with Image.open(io.BytesIO(data), formats=("JPEG", "PNG", "WEBP")) as image:
            if image.width * image.height > MAX_IMAGE_PIXELS:
                return row.original_filename, None
            image.load()
            image.thumbnail((_THUMB_PX, _THUMB_PX))
            out = io.BytesIO()
            image.convert("RGB").save(out, format="PNG")
        return row.original_filename, out.getvalue()
    except Exception:  # noqa: BLE001 - an unreadable image shows its name, never fails the download
        return row.original_filename, None


def _cell(value):
    """Text with the characters a worksheet cannot hold removed, and formula-safe."""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    return _xlsx_safe_text(value) if value not in (None, "") else None


def to_xlsx(db: Session, app: NewProductApplication) -> Tuple[bytes, str]:
    import openpyxl
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.worksheet.properties import PageSetupProperties

    stem = file_stem(db, app)
    when = _myt(app.created_at)
    date_text = when.strftime("%d/%m/%Y") if when else ""
    needed = app.needed_by.strftime("%d/%m/%Y") if app.needed_by else ""
    person = _sales_person(db, app)
    customer = _customer(db, app)
    brands = {
        str(i): n for i, n in db.query(Brand.id, Brand.brand_name).filter(
            Brand.id.in_([l.brand_id for l in app.lines if l.brand_id])
        ).all()
    } if any(l.brand_id for l in app.lines) else {}

    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    lines: List[Optional[NewProductApplicationLine]] = list(app.lines) or [None]
    for number, line in enumerate(lines, start=1):
        ws = wb.create_sheet(f"Product {number}")
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 60
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1

        def heading(row: int, text: str) -> None:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
            cell = ws.cell(row=row, column=1, value=text)
            cell.font = bold
            cell.alignment = Alignment(horizontal="center", vertical="center")
            for col in (1, 2):
                ws.cell(row=row, column=col).border = box

        def field(row: int, label: str, value=None) -> None:
            head = ws.cell(row=row, column=1, value=label)
            head.font = bold
            head.alignment = Alignment(horizontal="right", vertical="top")
            body = ws.cell(row=row, column=2, value=_cell(value))
            body.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
            head.border = box
            body.border = box

        image = _first_image(db, line) if line is not None else None
        heading(1, "NEW PRODUCT APPLICATION")
        field(2, "REFF:", stem)
        field(3, "DATE:", date_text)
        field(4, "SALES PERSON:", person)
        field(5, "PRODUCT CODE:", line.proposed_code if line else None)
        field(6, "PRODUCT TYPE:", line.description if line else None)
        field(7, "DIMENSION:", line.dimension if line else None)
        field(8, "MATERIAL:", line.material if line else None)
        field(9, "THICKNESS:", line.thickness if line else None)
        field(10, "PACKING", line.packing if line else None)
        field(11, "BRAND:", brands.get(str(line.brand_id)) if line and line.brand_id else None)
        field(12, "IMAGE:", image[0] if image and image[1] is None else None)
        if image and image[1] is not None:
            picture = XlImage(io.BytesIO(image[1]))
            ws.add_image(picture, "B12")
            ws.row_dimensions[12].height = picture.height * 0.75 + 6
        field(13, "CUSTOMER:", customer)
        field(14, "PROJECT:", app.project_name)
        field(15, "QTY:", float(line.expected_qty) if line and line.expected_qty is not None else None)
        field(16, "TARGET DELIVERY:", needed)
        field(17, "REMARK:", app.remark)
        heading(18, "PURCHASING FEEDBACK")
        field(19, "COST:", float(line.cost_price) if line and line.cost_price is not None else None)
        field(20, "LEADTIME:", app.lead_time)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), f"{stem}.xlsx"
