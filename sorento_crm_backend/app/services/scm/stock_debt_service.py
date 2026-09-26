"""Stock Debt: every outstanding sales order without supply, as a month x product balance.

S2 of `PLAN-scm-borrow-ladder-v7-stock-debt.md` (section 3.4), rulings R6/R7, R15, R21, R23.

**The view SHOWS; the board DECIDES** (R23). Nothing here writes, proposes or reserves. It
reads the same book the ladder reads, hands it to `supply_assignment.assign()` - the one
piece of arithmetic both surfaces share - and prints the answer as a balance PER MONTH
(R37: what is debted in August stays in August) with the lines and documents behind each
cell. A cell and its drill are two readings of the same walk: free supply dated in the
month, less what the lines due in it went short of on their own dates.

**One read per input, never one per product.** The page is the whole flagged catalogue
(1,000-2,000 products on the live book), so every fact is fetched for the WHOLE set in one
query and then split by product in Python: on hand, open demand, SPO, PO, the decisions that
pin stock and the links that pin a document. A per-product query here is what would turn a
screen into a minute. The trigger for a persisted debt table is stated in plan 3.6 and it is
this cost - measured, not guessed.

**No debt table** (plan 3.6). Debt is computed from the book, and the parts of it that are
decisions (ORDER_BACK rows, placement links, confirmed allocations) are already persisted by
the board. A table would be a second copy of a derived figure, stale the moment anybody
confirms anything.
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Any, Dict, Iterable, List, Optional, Sequence, Set, Tuple

from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.models.base import get_company_scope
from app.models.inventory import Stock, Warehouse
from app.models.order import SalesOrder, SalesOrderLine
from app.models.procurement import (
    ProductSupplier,
    PurchaseOrder,
    PurchaseOrderLine,
    SPOAllocation,
    Supplier,
)
from app.models.product import Product, ProductCategory
from app.models.project_so import (
    INQUIRY_CANCELLED,
    OrderInquiry,
    OrderInquiryLink,
    OrderInquiryRow,
    ProjectSalesOrderLine,
    SOLineAllocation,
)
from app.models.sales_agent import SalesAgent
from app.services.company_scope import build_company_predicate
from app.services.error_handler import AppException
from app.services.project_supply_service import ProjectSupplyService, held_qty_expr
from app.services.scm import sales_agent_service, spo_supply
from app.services.scm.demand import demand_qty, is_open_demand, plan_qty
from app.services.scm.front_planning_engine import DEFAULT_LEAD_TIME_DAYS
# Reused, not reinvented (AC-18): the low stock report's own cap. A read-only export off a
# bounded catalogue does not need a cap of its own; it needs the SAME reason that one has -
# "narrow it first" past a size nobody opens a workbook to page through.
from app.services.scm.low_stock_report_service import MAX_LOW_STOCK_ROWS
from app.services.scm.planning_predicate import fulfilment_planning_predicate
from app.services.scm.workbook_split import split_rows, unique_sheet_title
from app.services.scm.supply_assignment import (
    BUCKET_TBA,
    BUCKET_UNDATED,
    BUCKET_UNLOCATED,
    KIND_ON_HAND,
    KIND_PO,
    KIND_SPO,
    Assignment,
    DemandLine,
    Hold,
    SupplyEvent,
    assign,
    effective_date,
    month_axis,
    month_key,
    tone_for,
)

_ZERO = Decimal("0")

#: The month keys that are not months. Addressable exactly like a `YYYY-MM` cell, because
#: the screen's TBA, No date and No location columns are cells a reader clicks like any
#: other (R28).
BUCKET_KEYS = (BUCKET_TBA, BUCKET_UNDATED, BUCKET_UNLOCATED)

EXPORT_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

_MONTH_NAMES = (
    "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
)

def _export_month_label(key: str) -> str:
    """`2026-09` -> `Sep 26` (AC-13). The export's own copy of the FE's `monthLabel` -
    the two are the same three lines twice, not a shared import, because one lives in
    Python and the other in TypeScript."""
    year, month = key.split("-")
    return f"{_MONTH_NAMES[int(month) - 1]} {year[2:]}"


def _float(value: Any) -> float:
    if value is None:
        return 0.0
    return float(value)


def _spo_ref(spo_number: Optional[str]) -> str:
    """`SPO-2026/08-0085` -> itself; `202607-S0105` -> `SPO 202607-S0105`.

    Every shipping order in the live book is already written `SPO-...`, and prefixing the
    word onto one of those reads "SPO SPO-2026/08-0085" - the exact sentence the captain
    flagged on SO418869 SRTWCX7405-RL-S-PJ (3 Sep 2026). `front_planning_engine.spo_reason`
    carries the same guard, for a document named in a WATER sentence rather than a hold.
    """
    if not spo_number:
        return "SPO"
    return spo_number if spo_number.upper().startswith("SPO") else f"SPO {spo_number}"


class StockDebtService:
    def __init__(self, db: Session):
        self.db = db
        #: The reader that already knows how to read SPO, PO and lead times, and the one the
        #: board uses. Sharing it is what stops the view and the board disagreeing about
        #: what is on the water.
        self.supply = ProjectSupplyService(db)

    # ------------------------------------------------------------------ the two answers

    def list(
        self,
        *,
        query: Optional[str] = None,
        group: Optional[str] = None,
        only_debt: bool = True,
        page: int = 1,
        limit: int = 50,
        date_from: Optional[date] = None,
        date_to: Optional[date] = None,
        supplier_ids: Optional[Sequence[str]] = None,
        book: str = "all",
    ) -> Dict[str, Any]:
        """The month x product board (AC-S2-6, extended AC-1 to AC-9; R14/R15 owner round).

        `date_from`/`date_to` (R14, replacing `cutoff` - REMOVED, not aliased) drop demand
        due before/after them, in `_demand()`; TBA reads 0 once the policy's
        `tba_date_from` sits after `date_to`, for free - every TBA line is dated on or
        after `tba_date_from`, so the same date filter drops the whole bucket without a
        second rule. `supplier_ids` (R15, replacing `supplier_id`) narrows to products
        whose LAST supplier (newest PO line, else the primary flag) is ANY of the values
        passed; `'none'` is one more value among the others, not a sentinel that excludes
        them. `book` (R1/A4) chooses the span `_warehouses` reads.

        `totals`, `suppliers` and `sheet_counts` travel on the ENVELOPE, over the WHOLE
        filtered set rather than the page, for the same reason the axis already does:
        derived per page, they would change under the reader as they page (AC-6/AC-7/AC-7b).
        """
        warehouses = self._warehouses(group, book)
        products = self._products(warehouses, query)
        product_ids = [product_id for product_id, _code, _name, _cat in products]
        assignments = self._assignments(
            [(pid, code, name) for pid, code, name, _cat in products],
            warehouses,
            date_from=date_from,
            date_to=date_to,
        )
        supplier_map = self._last_supplier_map(product_ids)

        pre_supplier = []
        for product_id, code, name, category_code in products:
            result = assignments[product_id]
            if only_debt and not self._in_debt(result):
                continue
            supplier = supplier_map.get(product_id) or {"id": None, "name": None}
            pre_supplier.append((product_id, code, name, category_code, supplier, result))

        # AC-7c: the supplier FACET is built from the set BEFORE the `supplier_ids` filter
        # narrows it - the select can then switch supplier without first clearing itself,
        # rather than a narrowed board silently dropping every option but the one chosen.
        suppliers = self._suppliers_list(row[4] for row in pre_supplier)

        wanted_suppliers = set(supplier_ids or [])
        filtered = []
        for entry in pre_supplier:
            supplier = entry[4]
            if wanted_suppliers:
                # R15: a product matches when its last supplier is ANY of the values
                # passed - `none` is one more value alongside real ids, never a sentinel
                # that excludes them.
                matches = ("none" in wanted_suppliers and supplier["id"] is None) or (
                    supplier["id"] is not None and supplier["id"] in wanted_suppliers
                )
                if not matches:
                    continue
            filtered.append(entry)

        axis = self._axis(
            (row[5] for row in filtered), date_from=date_from, date_to=date_to
        )
        filtered.sort(key=lambda row: self._sort_key(row[5], row[1]))

        data_rows = []
        for product_id, code, name, category_code, supplier, result in filtered:
            months = self._months_on_axis(result, axis, product_id)
            # R17: the row's `total` sums months + TBA ONLY - `undated`/`unlocated` are no
            # longer folded in (the "No date"/"No location" columns leave the screen and
            # the workbook both); the two still ride the row unchanged, just not here.
            total = sum(month["balance"] for month in months) + result.tba
            data_rows.append(
                {
                    "product_id": product_id,
                    "product_code": code,
                    # AC-9/R8: null when it equals the code, trimmed of surrounding
                    # whitespace, case-sensitively - done once here so the board and the
                    # export agree without each re-deriving it.
                    "product_name": (
                        None if name is not None and name.strip() == code else name
                    ),
                    "months": months,
                    "tba": result.tba,
                    "undated": result.undated,
                    "unlocated": result.unlocated,
                    "supplier_id": supplier["id"],
                    "supplier_name": supplier["name"],
                    "category_code": category_code,
                    "total": total,
                }
            )

        totals = self._totals(data_rows, axis)
        sheet_counts = self._sheet_counts(data_rows)

        start = max(page - 1, 0) * limit
        page_rows = data_rows[start : start + limit]
        return {
            "data": page_rows,
            "pagination": {"total": len(data_rows), "page": page, "limit": limit},
            "months": axis,
            "tba_month": month_key(self._tba_from()),
            "groups": self._groups(),
            "totals": totals,
            "suppliers": suppliers,
            "sheet_counts": sheet_counts,
        }

    def cell(
        self,
        product_id: str,
        month: str,
        group: Optional[str] = None,
        *,
        date_from: Optional[date] = None,
        date_to: Optional[date] = None,
        book: str = "all",
    ) -> Dict[str, Any]:
        """The demand and the supply behind one cell (AC-S2-7, R28; extended AC-11).

        The same reads as the board, narrowed to one product, so the two tables foot with the
        cell that opened them by construction rather than by agreement: the drill's
        `free_qty` less its `short_qty`, over the rows of one month, IS that month's balance
        (R37). `group` is the
        narrowing the BOARD was showing when the cell was pressed, and it is not optional
        detail: `group=BB` recomputes the balance from the BB span only, so a drill that read
        the whole book would answer a different question from the cell that opened it.
        `date_from`/`date_to` (R14, replacing `cutoff`) and `book` are the same narrowings
        the list route takes, threaded through for the same reason `group` already is
        (AC-11).
        """
        if month not in BUCKET_KEYS and not self._is_month_key(month):
            raise AppException(
                status_code=422,
                message="month must be YYYY-MM, 'tba', 'undated' or 'unlocated'.",
                code="stock_debt_bad_month",
            )
        warehouses = self._warehouses(group, book)
        product = (
            self.db.query(Product.id, Product.product_code, Product.product_name)
            .filter(Product.id == product_id)
            .first()
        )
        if product is None:
            raise AppException(
                status_code=404, message="Product not found.", code="NOT_FOUND"
            )
        products = [(str(product.id), product.product_code, product.product_name)]
        assignments = self._assignments(
            products, warehouses, keep_events=True, date_from=date_from, date_to=date_to,
        )
        result = assignments[str(product.id)]
        events = self._event_cache[str(product.id)]

        # R29: Supply's own "Assigned to" entries carry `line_no` (the CORE line's own
        # AutoCount `Seq`, `DemandLine.core_line_no`) beside `so_number`, so keyed on the
        # PAIR rather than `so_number` alone - two lines of the same SO could otherwise
        # merge into one entry.
        assigned_to: Dict[str, Dict[Tuple[str, Optional[int]], float]] = {}
        for line in result.lines:
            for item in line.assigned:
                key = (line.line.so_number, line.line.core_line_no)
                assigned_to.setdefault(item.event.key, {}).setdefault(key, 0.0)
                assigned_to[item.event.key][key] += item.qty

        demand = [
            {
                "so_number": line.line.so_number,
                "agent_code": line.line.agent_code,
                "warehouse_code": line.line.warehouse,
                "required_date": line.line.required_date,
                "open_qty": line.line.open_qty,
                # R22: Ordered/Delivered beside the existing Outstanding (`open_qty`,
                # unchanged) - `qty_ordered` is `plan_qty()` (CS's own `qty_required`
                # when stated, else the book's `qty_ordered`).
                "qty_ordered": line.line.qty_ordered,
                "qty_delivered": line.line.qty_delivered,
                "assigned_qty": round(sum(item.qty for item in line.assigned), 4),
                "assigned_source": self._source_text(line),
                "status": line.status,
                # What this line booked into the month it sits in (R37): what it was short
                # of ON ITS OWN DATE. A `late` line ends covered and still carries one,
                # which is why the drill states it rather than leaving the reader to
                # subtract Assigned from Open and get a different number from the cell.
                "short_qty": line.short_at_date,
                # R29: the Sales order cell's own link target.
                "sales_order_id": line.line.sales_order_id,
                # R29 + addendum: one LINKED entry per source, replacing `assigned_source`.
                "assigned_from": self._assigned_from(line),
            }
            for line in result.lines
            if line.bucket == month
        ]
        demand.sort(key=lambda row: (row["required_date"] or date.max, row["so_number"]))

        supply: List[Dict[str, Any]] = []
        if month not in BUCKET_KEYS:
            as_of = date.today()
            current = month_key(as_of)
            # The events AS THE WALK COUNTED THEM (R-O): a late-but-alive document was
            # admitted at an ASSUMED date, and the cell it is filed under has to be the one
            # its free quantity was credited to, or the drill and the cell that opened it
            # disagree about which month the goods are in.
            admitted = {event.key: event for event in result.supply}
            for event in events:
                # An uncounted document (dead, or with no date at all) is listed in the
                # CURRENT month: its own arrival month has gone, and the axis starts today.
                walked = admitted.get(event.key)
                counted = walked is not None
                arrival = walked.at if counted else event.at
                key = month_key(effective_date(arrival, as_of)) if counted else current
                if key != month:
                    continue
                stated = walked.stated_at if counted else None
                supply.append(
                    {
                        "kind": event.kind,
                        "ref": event.ref,
                        # R29: the Document cell's own link target.
                        "spo_number": event.spo_number,
                        "spo_line_number": event.spo_line_number,
                        "warehouse_code": event.warehouse,
                        # THE ASSUMED date where there is one (R-O), because that is what
                        # the walk planned against; the paperwork's own date travels beside
                        # it rather than instead of it.
                        "date": arrival,
                        "stated_date": stated,
                        "days_late": int(getattr(walked, "days_late", 0) or 0)
                        if counted
                        else 0,
                        "bought_for": event.bought_for,
                        # R26: an SPO's own Qty is the RAW ordered quantity - `event.qty`
                        # stays the netted OUTSTANDING balance the walk itself assigns
                        # against, stated here as its own column instead. On hand has no
                        # such split (it is not a document with a received/outstanding
                        # history), so both are `None` - blank, never a fabricated 0.
                        "qty": event.ordered_qty if event.kind == KIND_SPO else event.qty,
                        "received_qty": event.received_qty if event.kind == KIND_SPO else None,
                        "outstanding_qty": event.qty if event.kind == KIND_SPO else None,
                        # What nobody took, once the whole walk was over - the other half of
                        # the cell (R37). A DEAD document is free of nothing: it is not
                        # supply until somebody re-dates it (R31).
                        "free_qty": result.free.get(event.key, 0.0) if counted else 0.0,
                        "overdue": not counted and event.at is not None,
                        # R29 (owner ruling, 25 Sep): `line_no` is the CORE line's own
                        # AutoCount `Seq`, beside `so_number` - ALWAYS present, `None` when
                        # AutoCount has never numbered the line, the same "always present,
                        # null when unknown" discipline every other field on this contract
                        # already follows. Sorted on the pair, `line_no or 0` so a line
                        # with none sorts before a numbered one rather than raising on
                        # `None < int`.
                        "assigned_to": [
                            {
                                "so_number": so_number,
                                "qty": round(qty, 4),
                                "line_no": line_no,
                            }
                            for (so_number, line_no), qty in sorted(
                                assigned_to.get(event.key, {}).items(),
                                key=lambda item: (item[0][0], item[0][1] or 0),
                            )
                        ],
                    }
                )
            supply.sort(key=lambda row: (row["date"] or date.max, row["ref"] or ""))

        # R25: the envelope's own quantity totals, over the WHOLE tab (never the page - a
        # drill has no paging, but the same "sum here, not on the FE" reasoning the board's
        # own `totals` already applies). Demand sums Outstanding as it stood BEFORE this
        # walk's assignment (`open_qty`, the line's own full ask); Supply sums whatever
        # each row's own Qty column actually offers - Outstanding for an SPO, the on-hand
        # figure otherwise - so the two tab labels state what a reader would get by adding
        # the column up themselves.
        demand_total_qty = sum(row["open_qty"] for row in demand)
        supply_total_qty = sum(
            row["outstanding_qty"] if row["kind"] == KIND_SPO else row["qty"]
            for row in supply
        )
        return {
            "demand": demand,
            "supply": supply,
            "demand_total_qty": demand_total_qty,
            "supply_total_qty": supply_total_qty,
        }

    # ------------------------------------------------------------------ the reads

    def _warehouses(
        self, group: Optional[str], book: str = "all"
    ) -> Dict[str, Warehouse]:
        """The bins this read spans, narrowed to one ownership group and/or one `book`
        (R1/A4, AC-8).

        `group=BB` narrows the SPAN of every read below it rather than filtering finished
        rows (AC-S2-6): the balance asked for is the BB group's own, and a row filtered after
        the fact would still have let another group's stock cover a BB order. `group` only
        ever narrows the PROJECT half - a site pool is nobody's ownership group.

        `book`:
        * `project` - flagged bins only (the pre-24-Sep span), narrowed by `group`.
        * `retail` - site pools only; `group` is meaningless here and is ignored.
        * `all` (default) - both, in ONE span. `assign()` already seals a pool's own
          group (`POOL_GROUP`) off from every project group in both directions (A4), so
          adding the pools to this dict is the whole change - the ladder's
          `assignments_for` already relies on the same fact for its own pool step.
        """
        project_bins = self._project_bins(group)
        if book == "project":
            return project_bins
        pool_bins = dict(self.supply.site_pool_warehouses())
        if book == "retail":
            return pool_bins
        return {**project_bins, **pool_bins}

    def _project_bins(self, group: Optional[str]) -> Dict[str, Warehouse]:
        """The flagged bins alone, narrowed by `group` - `_warehouses`'s own project half,
        and the whole of `_groups`'s span (a site pool admits no ownership group)."""
        rows = self.db.query(Warehouse).filter(fulfilment_planning_predicate()).all()
        if group:
            wanted = group.strip().upper()
            rows = [
                row
                for row in rows
                if sales_agent_service.group_of_warehouse_code(row.warehouse_code)
                == wanted
            ]
        return {str(row.id): row for row in rows}

    def _groups(self) -> List[str]:
        """The ownership groups the flag currently admits, for the toolbar's select."""
        rows = (
            self.db.query(Warehouse.warehouse_code)
            .filter(fulfilment_planning_predicate())
            .all()
        )
        return sorted(
            {
                group
                for group in (
                    sales_agent_service.group_of_warehouse_code(row[0]) for row in rows
                )
                if group
            }
        )

    def _products(
        self, warehouses: Dict[str, Warehouse], query: Optional[str]
    ) -> List[Tuple[str, str, Optional[str], Optional[str]]]:
        """Every product with stock, demand or incoming at those bins: id, code, name and
        `category_code` (R4/AC-5) - one JOIN, not a second query, for the same reason the
        rest of this file reads once for the whole set.

        Four id reads and one product read, because a product with nothing at a flagged bin
        has no debt to state and no row to render. The demand read reaches one step further
        than the other three: an UNLOCATED sales-order line names no bin at all, and a screen
        that lists what is owed while silently dropping 2,312 open lines (30 Aug dev copy)
        is answering a narrower question than the one it is asked.
        """
        ids = set(warehouses)
        if not ids:
            return []
        candidates: set = set()
        candidates |= {
            str(row[0])
            for row in self.db.query(Stock.product_id)
            .filter(Stock.warehouse_id.in_(ids), Stock.quantity_on_hand > 0)
            .distinct()
            .all()
        }
        candidates |= {
            str(row[0])
            for row in self.db.query(SalesOrderLine.product_id)
            .join(SalesOrder, SalesOrder.id == SalesOrderLine.sales_order_id)
            .filter(
                self._demand_span(ids),
                SalesOrder.status == "open",
                is_open_demand(),
            )
            .distinct()
            .all()
        }
        candidates |= {
            str(row[0])
            for row in self.db.query(SPOAllocation.product_id)
            .filter(
                SPOAllocation.warehouse_id.in_(ids),
                SPOAllocation.line_status == "open",
                SPOAllocation.allocated_quantity > SPOAllocation.quantity_received,
            )
            .distinct()
            .all()
        }
        candidates |= {
            str(row[0])
            for row in self.db.query(PurchaseOrderLine.product_id)
            .filter(
                PurchaseOrderLine.warehouse_id.in_(ids),
                PurchaseOrderLine.line_status == "open",
                PurchaseOrderLine.qty_ordered > PurchaseOrderLine.qty_received,
            )
            .distinct()
            .all()
        }
        if not candidates:
            return []

        rows = (
            self.db.query(
                Product.id,
                Product.product_code,
                Product.product_name,
                ProductCategory.category_code,
            )
            .outerjoin(ProductCategory, ProductCategory.id == Product.category_id)
            .filter(Product.id.in_(candidates))
        )
        if query:
            needle = f"%{query.strip()}%"
            rows = rows.filter(
                or_(
                    Product.product_code.ilike(needle),
                    Product.product_name.ilike(needle),
                )
            )
        return [
            (str(row.id), row.product_code or "", row.product_name, row.category_code)
            for row in rows.all()
        ]

    @staticmethod
    def _demand_span(warehouse_ids):
        """"A line this view is answerable for": at one of these bins, or at NO bin.

        One expression, used by the candidate read and by `_demand`, so the products that
        get a row and the lines that fill it can never come from two different rules.
        `assign()` gives an unlocated line its own bucket - it is in no group's pile, so it
        draws nothing (AC-S2-1b's sibling case) - but it is COUNTED and it is listed.
        """
        return or_(
            SalesOrderLine.warehouse_id.in_(list(warehouse_ids)),
            SalesOrderLine.warehouse_id.is_(None),
        )

    def assignments_for(
        self,
        product_ids: Sequence[str],
        warehouses: Dict[str, Warehouse],
        *,
        as_of: Optional[date] = None,
        include_po: bool = True,
    ) -> Dict[str, Assignment]:
        """The same assignment, for a caller that holds product ids and a span of its own.

        PUBLIC for the LADDER (S3, R21): the board reads the assignment this view reads, or
        the two surfaces disagree about what is free. The ladder's span is this one plus the
        site pools, because it has a pool step and the view has not (see
        `ProjectSupplyService.planning_assignments`).

        `include_po` defaults `True` here: R23 ("got PO doesn't mean got supply") is the
        STOCK DEBT VIEW's own reading, not the shared assignment's - the board and the
        ladder still net a PO as supply (plan v7 R29), and `planning_assignments` is
        exactly this caller, so its own default has to keep doing that without having to
        say so at every call site.
        """
        return self._assignments(
            [(str(pid), "", None) for pid in product_ids],
            warehouses,
            as_of=as_of,
            include_po=include_po,
        )

    def _assignments(
        self,
        products: Sequence[Tuple[str, str, Optional[str]]],
        warehouses: Dict[str, Warehouse],
        *,
        keep_events: bool = False,
        as_of: Optional[date] = None,
        date_from: Optional[date] = None,
        date_to: Optional[date] = None,
        include_po: bool = False,
    ) -> Dict[str, Assignment]:
        """One `assign()` per product, off ONE read per input for the whole set.

        `date_from`/`date_to` (R14, replacing `cutoff`; AC-1/AC-1b/AC-3) drop demand due
        before/after them in `_demand()` below - they touch DEMAND only, never supply
        (AC-3: supply landing after a line's own due date, but on or before `date_to`,
        still covers it - the walk itself is unchanged).

        `include_po` defaults `False`: `list()`/`cell()` (the Stock Debt VIEW, R23) never
        pass it, so a PO is neither supply nor a hold's own document there. `assignments_for`
        (the board/ladder's own path) passes `True`.
        """
        self._event_cache: Dict[str, List[SupplyEvent]] = {}
        self._lead_cache: Dict[str, int] = {}
        product_ids = [product_id for product_id, _code, _name in products]
        # The CALLER's `as_of` when it pins one (the board's simulation does), else today.
        as_of = as_of or date.today()
        tba_from = self._tba_from()
        if not product_ids:
            return {}

        warehouse_ids = list(warehouses)
        codes = {
            warehouse_id: warehouse.warehouse_code or warehouse_id
            for warehouse_id, warehouse in warehouses.items()
        }
        pools = set(self.supply.site_pool_warehouses())

        # FIRST, and deliberately: `lead_times` fills the memo the per-line fallback in
        # `_po_rows` reads. Called after `_supply`, every PO line whose supplier agreement
        # states no lead paid its own round trip - ~1,900 extra queries per list request on
        # the dev copy.
        leads = self.supply.lead_times(product_ids)
        supply_rows = self._supply(
            product_ids, warehouse_ids, codes, pools, as_of=as_of, include_po=include_po,
        )
        demand_rows = self._demand(
            product_ids, warehouse_ids, codes, pools, date_from=date_from, date_to=date_to,
        )
        holds = self._holds(
            product_ids,
            {line.key for lines in demand_rows.values() for line in lines},
            include_po=include_po,
        )

        settings = self.supply._fulfilment_settings()
        grace = settings.get("overdue_grace_days")
        dead = settings.get("overdue_dead_days")

        out: Dict[str, Assignment] = {}
        for product_id in product_ids:
            # The product's own lead, or the ladder's default - the SAME source the reserve
            # window uses (plan risk 5), so the red horizon here and the window there agree.
            # `is None`, never `or` - the same test `_po_rows` and `reserve_window_end`
            # make. A supplier who states a 0-day lead is stating something (stock off the
            # shelf), and reading that as "nobody says" turned it into 90 days, which paints
            # three months of a product that can be bought today red.
            lead = leads.get(product_id)
            self._lead_cache[product_id] = (
                DEFAULT_LEAD_TIME_DAYS if lead is None else lead
            )
            events = supply_rows.get(product_id, [])
            lines = demand_rows.get(product_id, [])
            line_keys = {line.key for line in lines}
            out[product_id] = assign(
                product_id,
                as_of=as_of,
                tba_from=tba_from,
                lead_days=self._lead_cache[product_id],
                supply=events,
                demand=lines,
                pinned=[hold for hold in holds if hold.line_key in line_keys],
                # R-O (3 Sep 2026): the grace a late document is counted under, off the
                # SAME active policy row `tba_from` above is read from - the board, this
                # view and the ladder all walk one assignment, so they cannot come to two
                # views of how late is too late.
                overdue_grace_days=grace,
                overdue_dead_days=dead,
            )
            if keep_events:
                self._event_cache[product_id] = events
        return out

    def _supply(
        self,
        product_ids: Sequence[str],
        warehouse_ids: Sequence[str],
        codes: Dict[str, str],
        pools: set,
        *,
        as_of: Optional[date] = None,
        include_po: bool = False,
    ) -> Dict[str, List[SupplyEvent]]:
        """On hand and SPO for the whole page - two reads, neither of them per product.
        A THIRD, PO, joins them when `include_po` is set (the board/ladder's own path,
        `assignments_for` - plan v7 R29).

        On hand is `quantity_on_hand - quantity_reserved`, the same arithmetic
        `_free_stock` states: reserved stock is spoken for by a picking or despatch that is
        already under way, so offering it to a sales-order line here would promise the same
        units twice. The confirmed HOLDS are subtracted separately, by pinning them to the
        lines that hold them (`_holds`) - which is the more useful shape, because the drill
        can then say which order has them.

        `as_of` is the CALLER's day, not the clock: an on-hand event is stamped with the day
        the walk starts, and stamping it `date.today()` while the walk ran at a pinned
        earlier date put the stock after every line due between the two, so a board
        simulated at a past date read its own floor as arriving late.

        R23 (owner, 24 Sep, third red batch): "got PO doesn't mean got supply." Stock
        Debt's own reading (`include_po=False`, `list()`/`cell()`'s own default) counts on
        hand and SPO only - a PO is a plan to buy, not stock anybody has or a shipment
        already moving, and reading it as supply here let a line read `covered`/`pinned`
        off a document that could still fall through.

        Fix round (CI, 25 Sep): R23 is the VIEW's own reading, not the shared assignment's
        - the board and the ladder still net a PO as supply (plan v7 R29,
        `test_ladder_v7_po_never_supplies.py`/`test_ladder_v7_incoming_spo_only.py`'s own
        guards), which is `assignments_for`'s `include_po=True` default reaching here.
        """
        as_of = as_of or date.today()
        out: Dict[str, List[SupplyEvent]] = {}

        rows = (
            self.db.query(
                Stock.product_id,
                Stock.warehouse_id,
                func.sum(
                    func.coalesce(Stock.quantity_on_hand, 0)
                    - func.coalesce(Stock.quantity_reserved, 0)
                ).label("qty"),
            )
            .filter(
                Stock.product_id.in_(product_ids),
                Stock.warehouse_id.in_(warehouse_ids),
            )
            .group_by(Stock.product_id, Stock.warehouse_id)
            .all()
        )
        for row in rows:
            qty = _float(row.qty)
            if qty <= 0:
                continue
            warehouse_id = str(row.warehouse_id)
            out.setdefault(str(row.product_id), []).append(
                SupplyEvent(
                    key=f"on_hand:{warehouse_id}",
                    kind=KIND_ON_HAND,
                    warehouse=codes.get(warehouse_id),
                    at=as_of,
                    qty=qty,
                    is_pool=warehouse_id in pools,
                )
            )

        for (product_id, warehouse_id), refs in self.supply.incoming_by_location(
            product_ids, warehouse_ids
        ).items():
            for ref in refs:
                out.setdefault(product_id, []).append(
                    SupplyEvent(
                        key=f"spo:{ref.allocation_id}",
                        kind=KIND_SPO,
                        warehouse=codes.get(warehouse_id),
                        at=ref.arrival_date,
                        # The WALK's own figure stays the netted outstanding balance
                        # (R26) - `ordered_qty`/`received_qty` below are display-only.
                        qty=_float(ref.qty),
                        ref=_spo_ref(ref.spo_number),
                        is_pool=warehouse_id in pools,
                        ordered_qty=_float(ref.ordered_qty),
                        received_qty=_float(ref.received_qty),
                        # R29: the wire fields the Document cell links off.
                        spo_number=ref.spo_number,
                        spo_line_number=ref.spo_line_no,
                    )
                )

        # R23 (VIEW only, `include_po=False` by default) / plan v7 R29 (board and ladder,
        # `include_po=True`): see the docstring above.
        if include_po:
            for (product_id, warehouse_id), lines in self.supply.po_by_location(
                product_ids, warehouse_ids
            ).items():
                for line in lines:
                    out.setdefault(product_id, []).append(
                        SupplyEvent(
                            key=f"po:{line.line_id}",
                            kind=KIND_PO,
                            warehouse=codes.get(warehouse_id),
                            at=line.arrival_date,
                            qty=_float(line.qty),
                            ref=f"PO {line.po_number} line {line.po_line_no}",
                            bought_for=line.bought_for,
                            is_pool=warehouse_id in pools,
                        )
                    )
        return out

    def _demand(
        self,
        product_ids: Sequence[str],
        warehouse_ids: Sequence[str],
        codes: Dict[str, str],
        pools: set,
        *,
        date_from: Optional[date] = None,
        date_to: Optional[date] = None,
    ) -> Dict[str, List[DemandLine]]:
        """Every open sales-order line at those bins, plus the ones at NO bin - the same
        `is_open_demand()` rule the ladder and `scm.committed_v` share, so the debt and the
        plan count one book. `_demand_span` is why an unlocated line is here.

        `date_from`/`date_to` (R14, replacing `cutoff`; AC-1/AC-1b/AC-2) drop a line due
        before `date_from` or after `date_to`; an undated line has no date to test and
        always survives either bound. Every TBA line is dated on or after the policy's
        `tba_date_from`, so a `date_to` earlier than that date drops the whole TBA bucket
        for free, off this same clause - no second rule needed (AC-2).
        """
        extra_clauses = []
        if date_to is not None:
            extra_clauses.append(
                or_(
                    SalesOrderLine.required_date.is_(None),
                    SalesOrderLine.required_date <= date_to,
                )
            )
        if date_from is not None:
            extra_clauses.append(
                or_(
                    SalesOrderLine.required_date.is_(None),
                    SalesOrderLine.required_date >= date_from,
                )
            )
        rows = (
            self.db.query(
                SalesOrderLine.id,
                SalesOrderLine.product_id,
                SalesOrderLine.warehouse_id,
                SalesOrderLine.required_date,
                SalesOrderLine.sales_order_id,
                # R29: the CORE line's own `line_no` (AutoCount's `Seq`) - Supply's own
                # "Assigned to" entries name this, never the PROJECT mirror's `line_no`
                # selected below (a different number, the ladder's own).
                SalesOrderLine.line_no.label("core_line_no"),
                demand_qty().label("qty"),
                SalesOrder.so_number,
                SalesAgent.sales_agent,
                # The PROJECT mirror's own line number. A core line nobody has adopted has
                # none, and a missing number is treated as absent everywhere it is printed.
                # Carried because a v7 borrow names the donor's line ("SO414285 line 4") and
                # the ladder reads its donors out of this very list (AC-S3-2, AC-S3-11).
                ProjectSalesOrderLine.line_no,
                # R22: the drill's own Ordered/Delivered columns - `plan_qty()` is the
                # SAME coalesce the board plans against (`coalesce(qty_required,
                # qty_ordered)`), reused rather than re-derived so the two screens cannot
                # come to disagree about what "Ordered" means for a line CS has stated.
                plan_qty().label("qty_ordered"),
                func.coalesce(SalesOrderLine.qty_delivered, 0).label("qty_delivered"),
            )
            .join(SalesOrder, SalesOrder.id == SalesOrderLine.sales_order_id)
            .outerjoin(SalesAgent, SalesAgent.id == SalesOrder.sales_agent_id)
            .outerjoin(
                ProjectSalesOrderLine,
                ProjectSalesOrderLine.core_sales_order_line_id == SalesOrderLine.id,
            )
            .filter(
                SalesOrderLine.product_id.in_(product_ids),
                self._demand_span(warehouse_ids),
                SalesOrder.status == "open",
                is_open_demand(),
                *extra_clauses,
            )
            .all()
        )
        out: Dict[str, List[DemandLine]] = {}
        for row in rows:
            warehouse_id = str(row.warehouse_id) if row.warehouse_id else ""
            out.setdefault(str(row.product_id), []).append(
                DemandLine(
                    key=str(row.id),
                    so_number=row.so_number or "",
                    line_no=row.line_no,
                    # None for an unlocated line, which is what puts it in its own bucket.
                    warehouse=codes.get(warehouse_id),
                    agent_code=row.sales_agent,
                    required_date=row.required_date,
                    open_qty=_float(row.qty),
                    is_pool=warehouse_id in pools,
                    qty_ordered=_float(row.qty_ordered),
                    qty_delivered=_float(row.qty_delivered),
                    # R29: the Sales order cell's own link target.
                    sales_order_id=str(row.sales_order_id) if row.sales_order_id else None,
                    core_line_no=row.core_line_no,
                )
            )
        return out

    def _holds(
        self,
        product_ids: Sequence[str],
        line_keys: set,
        *,
        include_po: bool = False,
    ) -> List[Hold]:
        """What is already promised: confirmed allocations and placement links (R21).

        `include_po` (fix round, CI, 25 Sep): `False` is the Stock Debt VIEW's own reading
        (R23) - a placement link naming a PO line pins nothing here, the same way a PO is
        not a supply event in `_supply` above. `True` (the board/ladder's own path,
        `assignments_for`) restores the PO branch exactly as it read before R23.

        Two shapes, one meaning. A `so_line_allocations` row is a decision holding STOCK at a
        bin; an `order_inquiry_links` row is a placement holding a DOCUMENT. Both bind before
        anybody queues (AC-S2-2), and both are keyed to the CORE sales-order line, which is
        what the demand read above is keyed on.

        **The allocation half is `ProjectSupplyService._hold_query`, not a second spelling of
        it.** That is the one predicate for "this row is holding stock right now" - confirmed,
        located, not an ORDER source, and belonging to no decision or to an ACTIVE one - and
        it is shared with the free-stock arithmetic precisely so the two cannot come to
        disagree. Restated by hand here, it had already lost `confirmed_at IS NOT NULL`, so a
        decision saved but never confirmed held stock on this screen and nowhere else.

        **Each hold carries the bin or document it names.** `assign()` honours a hold whose
        supply is outside the read span (a site pool, an unflagged bin, another group under
        `group=`) off exactly these fields (AC-S2-1b), so the drill still says `On hand BRW`
        rather than leaving a pinned line looking unsourced.
        """
        if not line_keys:
            return []
        keys = list(line_keys)
        out: List[Hold] = []

        rows = (
            self.supply._hold_query(
                [str(pid) for pid in product_ids],
                exclude_line_ids=None,
                entities=(
                    ProjectSalesOrderLine.core_sales_order_line_id,
                    SOLineAllocation.warehouse_id,
                    held_qty_expr(),
                    Warehouse.warehouse_code,
                ),
            )
            .join(Warehouse, Warehouse.id == SOLineAllocation.warehouse_id)
            .filter(ProjectSalesOrderLine.core_sales_order_line_id.in_(keys))
            .all()
        )
        for core_line_id, warehouse_id, qty, warehouse_code in rows:
            if _float(qty) <= 0:
                continue
            out.append(
                Hold(
                    line_key=str(core_line_id),
                    supply_key=f"on_hand:{warehouse_id}",
                    qty=_float(qty),
                    kind=KIND_ON_HAND,
                    warehouse=warehouse_code,
                )
            )

        links = (
            self.db.query(
                ProjectSalesOrderLine.core_sales_order_line_id,
                OrderInquiryLink.spo_allocation_id,
                OrderInquiryLink.po_line_id,
                OrderInquiryLink.qty,
                SPOAllocation.spo_number,
                SPOAllocation.spo_line_number,
                PurchaseOrder.po_number,
                # R29 addendum: the order inquiry this placement came through - a
                # placement is part of an OI ROW's quantity on one document line, and
                # `_holds` already joins that row to get here.
                OrderInquiry.inquiry_no,
                OrderInquiry.id.label("order_inquiry_id"),
            )
            .join(OrderInquiryRow, OrderInquiryRow.id == OrderInquiryLink.row_id)
            .join(OrderInquiry, OrderInquiry.id == OrderInquiryRow.order_inquiry_id)
            .join(
                ProjectSalesOrderLine,
                ProjectSalesOrderLine.id == OrderInquiryRow.so_line_id,
            )
            .outerjoin(
                SPOAllocation, SPOAllocation.id == OrderInquiryLink.spo_allocation_id
            )
            .outerjoin(
                PurchaseOrderLine, PurchaseOrderLine.id == OrderInquiryLink.po_line_id
            )
            .outerjoin(
                PurchaseOrder,
                PurchaseOrder.id == PurchaseOrderLine.purchase_order_id,
            )
            .filter(
                ProjectSalesOrderLine.core_sales_order_line_id.in_(keys),
                # A cancelled inquiry row holds nothing - the same filter every other
                # consumer of these links applies (`planning_change_service`,
                # `project_order_inquiry_service`). Without it a withdrawn placement went
                # on pinning a document to a line nobody is waiting on.
                OrderInquiryRow.state != INQUIRY_CANCELLED,
                # R7/AC-E12: SPOAllocation is OUTER-joined, so this passes a plain PO
                # link (its columns come back NULL) untouched and only excludes a
                # hold whose SPO side names a retired line.
                *spo_supply.visible_line_clauses(),
            )
            .all()
        )
        for row in links:
            qty = _float(row.qty)
            if qty <= 0:
                continue
            if not row.spo_allocation_id:
                # R23 (VIEW only): a placement link to a PO line pins nothing in Stock
                # Debt's own reading - PO is not supply here at all, so there is no
                # PO-kind event left in this read's span for it to bind to (AC-S2-1b's
                # "stand an event up from the hold's own fields" branch would otherwise
                # manufacture one). `include_po` (board/ladder, plan v7 R29) restores
                # exactly the pre-R23 PO branch.
                if not include_po or not row.po_line_id:
                    continue
                out.append(
                    Hold(
                        line_key=str(row[0]),
                        supply_key=f"po:{row.po_line_id}",
                        qty=qty,
                        kind=KIND_PO,
                        ref=f"PO {row.po_number}" if row.po_number else "PO",
                    )
                )
                continue
            supply_key = f"spo:{row.spo_allocation_id}"
            kind, ref = KIND_SPO, _spo_ref(row.spo_number)
            out.append(
                Hold(
                    line_key=str(row[0]),
                    supply_key=supply_key,
                    qty=qty,
                    kind=kind,
                    ref=ref,
                    spo_number=row.spo_number,
                    spo_line_number=row.spo_line_number,
                    # R29 addendum: this hold IS a placement (an `order_inquiry_links`
                    # row), so it always names the order inquiry it came through.
                    oi_number=row.inquiry_no,
                    oi_id=str(row.order_inquiry_id),
                )
            )
        return out

    # ------------------------------------------------------------------ small helpers

    def _tba_from(self) -> date:
        """The policy's TBA line, read once per request off the row the admin screen edits."""
        from app.services.scm.priority import DEFAULT_TBA_DATE_FROM

        return (
            self.supply._fulfilment_settings().get("tba_date_from")
            or DEFAULT_TBA_DATE_FROM
        )

    @staticmethod
    def _is_month_key(value: str) -> bool:
        try:
            year, month = value.split("-")
            return len(value) == 7 and 1 <= int(month) <= 12 and len(year) == 4
        except (ValueError, AttributeError):
            return False

    @staticmethod
    def _in_debt(result: Assignment) -> bool:
        """Anything owed and unsupplied - a negative month, or bucketed demand.

        The three buckets count, and the AC's "no negative month" wording is read that way
        deliberately: AC-S2-3 calls a TBA line's whole quantity DEBT, so a product whose only
        debt is 2030 demand, or an order nobody has given a location, would otherwise be
        dropped from the one screen that exists to list what is owed.
        """
        return (
            any(month.balance < 0 for month in result.months)
            or result.tba < 0
            or result.undated < 0
            or result.unlocated < 0
        )

    @staticmethod
    def _sort_key(result: Assignment, code: str) -> tuple:
        """Earliest red month, then product code; no red month sorts after every red one.

        Red is "cannot be bought in time", so the order is the order the work is urgent in.
        """
        red = next(
            (month.key for month in result.months if month.tone == "red"), None
        )
        return (red is None, red or "", code or "")

    def _axis(
        self,
        results: Iterable[Assignment],
        *,
        date_from: Optional[date] = None,
        date_to: Optional[date] = None,
    ) -> List[str]:
        """The month columns of the whole filtered set: today (or `date_from`'s month, if
        later) to the last month anything is dated in (or `date_to`'s month, if earlier).
        One axis for every row, or the columns move as the reader pages.

        `date_from` (R14/AC-1b) raises the FIRST column to its own month, never below
        today's. `date_to` (A2/AC-1) caps the LAST column at its own month - demand past
        it is already dropped in `_demand()`, but a document arriving after `date_to` is
        still supply (AC-3) and could otherwise stretch the axis past a month nothing due
        survives in. Never below `first`: a `date_to` in the past still leaves the current
        month (or `date_from`'s) on screen.
        """
        first = month_key(date.today())
        if date_from is not None:
            first = max(first, month_key(date_from))
        last = first
        for result in results:
            for month in result.months:
                if month.key > last:
                    last = month.key
        if date_to is not None:
            cutoff_month = month_key(date_to)
            last = max(first, min(last, cutoff_month))
        return month_axis(first, last)

    def _months_on_axis(
        self, result: Assignment, axis: Sequence[str], product_id: str
    ) -> List[dict]:
        """One entry per axis key, in axis order. Past a row's own last event the months
        read 0, because nothing is due and nothing arrives in them (R37) - the balance does
        not carry, so a column states its own month or it states nothing owed."""
        as_of = date.today()
        lead = self._lead_cache.get(product_id, DEFAULT_LEAD_TIME_DAYS)
        by_key = {month.key: month for month in result.months}
        out: List[dict] = []
        for key in axis:
            month = by_key.get(key)
            balance = month.balance if month is not None else 0.0
            out.append(
                {
                    "key": key,
                    "balance": balance,
                    "tone": month.tone
                    if month is not None
                    else tone_for(balance, key, as_of=as_of, lead_days=lead),
                }
            )
        return out

    def _assigned_from(self, line) -> List[Dict[str, Any]]:
        """R29 + addendum: one LINKED entry per source behind `line`'s `assigned_qty` -
        a document (SPO/PO) or an on-hand bin, each with its OWN quantity, so the FE can
        render "SPO-... line 4 (100)" / "On hand BRW-BB (14)" instead of one merged
        sentence.

        Grouped by event key - the same grouping `_source_text` already applied to its own
        text label - so a line drawing from the same document across more than one walk
        step still names it once, with the two quantities summed, rather than twice.

        `oi_number`/`oi_id` (addendum) ride on the TAKE (`item.oi_number`/`oi_id`), not the
        event: the event object is SHARED across every line that draws from it, and only a
        PINNED placement names an order inquiry at all - a plain walk draw of the same
        document, by another line, names none.
        """
        entries: Dict[str, Dict[str, Any]] = {}
        order: List[str] = []
        for item in line.assigned:
            event = item.event
            if event.key not in entries:
                order.append(event.key)
                if event.kind == KIND_ON_HAND:
                    entries[event.key] = {
                        "kind": KIND_ON_HAND,
                        "ref": f"On hand {event.warehouse}",
                        "spo_number": None,
                        "spo_line_number": None,
                        "qty": 0.0,
                    }
                else:
                    entries[event.key] = {
                        "kind": event.kind,
                        "ref": event.ref or event.kind.upper(),
                        "spo_number": event.spo_number,
                        "spo_line_number": event.spo_line_number,
                        "qty": 0.0,
                        "oi_number": item.oi_number,
                        "oi_id": item.oi_id,
                    }
            entries[event.key]["qty"] += item.qty
        return [
            {**entries[key], "qty": round(entries[key]["qty"], 4)} for key in order
        ]

    def _source_text(self, line) -> Optional[str]:
        """What a demand row says it is covered FROM - `On hand BRW-BB`, `SPO ...`, `PO ...`.

        Human text, not an id: this is the sentence a planner reads before pressing Plan.
        """
        labels: List[str] = []
        for item in line.assigned:
            event = item.event
            label = (
                f"On hand {event.warehouse}"
                if event.kind == KIND_ON_HAND
                else (event.ref or event.kind.upper())
            )
            if label not in labels:
                labels.append(label)
        return ", ".join(labels) if labels else None

    # ------------------------------------------------------------------ last supplier (R3/A1)

    def _last_supplier_map(
        self, product_ids: Sequence[str]
    ) -> Dict[str, Dict[str, Optional[str]]]:
        """`{product_id: {"id": supplier_id | None, "name": supplier_name | None}}` -
        the LAST supplier (AC-4/AC-5): the supplier on the product's newest purchase-order
        LINE, falling back to the manually-flagged primary product supplier, else neither.

        The window (newest `PurchaseOrder.issue_date`, `PurchaseOrderLine.created_at`
        breaking a tie) is `po_last_cost_service.last_cost_rows`'s own pick key, reused
        rather than re-derived - the same "which PO line answers for this product" question,
        one PARTITION BY `product_id` instead of `(product_id, warehouse_id)` because the
        board asks per PRODUCT, not per bin. Cancelled POs are skipped (A1); a cancelled
        LINE is not, because A1 names only the PO's own status - an SPO's supplier is never
        consulted here at all (A1: "we should look at PO, not SPO").
        """
        out: Dict[str, Dict[str, Optional[str]]] = {
            str(pid): {"id": None, "name": None} for pid in product_ids
        }
        if not product_ids:
            return out

        rn = (
            func.row_number()
            .over(
                partition_by=PurchaseOrderLine.product_id,
                order_by=(
                    PurchaseOrder.issue_date.desc().nulls_last(),
                    PurchaseOrderLine.created_at.desc(),
                ),
            )
            .label("rn")
        )
        numbered = (
            self.db.query(
                PurchaseOrderLine.product_id,
                PurchaseOrder.supplier_id,
                rn,
            )
            .join(PurchaseOrder, PurchaseOrder.id == PurchaseOrderLine.purchase_order_id)
            .filter(
                PurchaseOrderLine.product_id.in_(product_ids),
                PurchaseOrder.status != "cancelled",
            )
        )
        # COMPANY SCOPE, EXPLICITLY, on BOTH tables named in this branch (belt-and-braces,
        # matching `spo_last_receipt_service`'s own windowed branch): `.subquery()` loses
        # the `with_loader_criteria` the session's `do_orm_execute` listener injects, and
        # the outer query below names only `Supplier`.
        for model in (PurchaseOrderLine, PurchaseOrder):
            predicate = build_company_predicate(model, get_company_scope(self.db))
            if predicate is not None:
                numbered = numbered.filter(predicate)
        numbered = numbered.subquery()
        po_supplier: Dict[str, Optional[str]] = {
            str(pid): (str(sid) if sid else None)
            for pid, sid in self.db.query(numbered.c.product_id, numbered.c.supplier_id)
            .filter(numbered.c.rn == 1)
            .all()
        }

        primary_rows = (
            self.db.query(ProductSupplier.product_id, ProductSupplier.supplier_id)
            .filter(
                ProductSupplier.product_id.in_(product_ids),
                ProductSupplier.is_primary_supplier.is_(True),
            )
            .all()
        )
        primary_supplier: Dict[str, str] = {
            str(pid): str(sid) for pid, sid in primary_rows
        }

        supplier_ids = {sid for sid in po_supplier.values() if sid} | set(
            primary_supplier.values()
        )
        names: Dict[str, str] = {}
        if supplier_ids:
            names = {
                str(row.id): row.supplier_name
                for row in self.db.query(Supplier.id, Supplier.supplier_name)
                .filter(Supplier.id.in_(supplier_ids))
                .all()
            }

        for pid in out:
            # The newest PO line's own answer wins WHENEVER one exists - even a PO line
            # with no supplier stated - and only a product with NO live PO at all falls
            # back to the primary flag (A1's "falling back to").
            supplier_id = po_supplier[pid] if pid in po_supplier else primary_supplier.get(pid)
            out[pid] = {
                "id": supplier_id,
                "name": names.get(supplier_id) if supplier_id else None,
            }
        return out

    # ------------------------------------------------------------------ envelope aggregates

    @staticmethod
    def _totals(data_rows: List[dict], axis: Sequence[str]) -> Dict[str, Any]:
        """The whole filtered set's totals (AC-6): summed here, over EVERY row already
        built for this request, before the page is sliced off - never the page's own rows,
        or the footer would change under the reader as they page."""
        months = {key: 0.0 for key in axis}
        tba = undated = unlocated = total = 0.0
        for row in data_rows:
            for month in row["months"]:
                months[month["key"]] += month["balance"]
            tba += row["tba"]
            undated += row["undated"]
            unlocated += row["unlocated"]
            total += row["total"]
        return {
            "months": months, "tba": tba, "undated": undated, "unlocated": unlocated,
            "total": total,
        }

    @staticmethod
    def _suppliers_list(
        suppliers: Iterable[Dict[str, Optional[str]]],
    ) -> List[Dict[str, str]]:
        """Distinct last suppliers, sorted by name (AC-7). Takes the raw `{id, name}`
        entries directly - AC-7c: the caller passes the set BEFORE the `supplier_id`
        filter narrows it, not a page of already-built rows, so the facet never loses an
        option the reader could still pick."""
        seen: Dict[str, str] = {}
        for supplier in suppliers:
            supplier_id = supplier["id"]
            if supplier_id and supplier_id not in seen:
                seen[supplier_id] = supplier["name"] or ""
        return [
            {"id": supplier_id, "name": name}
            for supplier_id, name in sorted(seen.items(), key=lambda pair: pair[1])
        ]

    @staticmethod
    def _sheet_counts(data_rows: List[dict]) -> Dict[str, int]:
        """The EXACT export sheet counts for the whole filtered set (AC-7b): distinct
        supplier keys, distinct category keys and distinct (supplier, category) PAIRS
        actually present - a pair with no row does not get a sheet, so this is never
        `len(suppliers) * len(categories)`. `None`/blank folds into one "none" bucket each,
        the same bucket the workbook titles "No supplier" / "No category" (AC-13..AC-16)."""
        supplier_keys = {row["supplier_id"] or "" for row in data_rows}
        category_keys = {row["category_code"] or "" for row in data_rows}
        pair_keys = {
            (row["supplier_id"] or "", row["category_code"] or "") for row in data_rows
        }
        return {
            "supplier": len(supplier_keys),
            "category": len(category_keys),
            "supplier_category": len(pair_keys),
        }
    # ------------------------------------------------------------------ export (R5/R10, AC-12..AC-18)

    def export(
        self,
        *,
        query: Optional[str] = None,
        group: Optional[str] = None,
        only_debt: bool = True,
        date_from: Optional[date] = None,
        date_to: Optional[date] = None,
        supplier_ids: Optional[Sequence[str]] = None,
        book: str = "all",
        split: str = "none",
    ) -> Tuple[bytes, str, str, Dict[str, int]]:
        """The workbook for the CURRENT filters (AC-17): `(bytes, content_type, filename,
        {"rows": n, "sheets": m})`, the same tuple shape `low_stock_report_service.
        export_low_stock` returns.

        Built off `list()` itself, unpaged (AC-17: the export's rows are exactly the
        list's, for the same filters) - `limit=MAX_LOW_STOCK_ROWS + 1` is enough to prove
        whether the filtered set fits under the cap (AC-18) without a second, differently-
        shaped read: `pagination["total"]` is always the WHOLE filtered set regardless of
        the limit used to slice `data`, so a set that fits is returned complete and a set
        that does not is refused before `data` is ever trusted to be complete.
        """
        listing = self.list(
            query=query, group=group, only_debt=only_debt, date_from=date_from,
            date_to=date_to, supplier_ids=supplier_ids, book=book, page=1,
            limit=MAX_LOW_STOCK_ROWS + 1,
        )
        if listing["pagination"]["total"] > MAX_LOW_STOCK_ROWS:
            raise AppException(422, "Narrow the filters first")
        return self._render_workbook(listing, split=split)

    def _render_workbook(
        self, listing: Dict[str, Any], *, split: str
    ) -> Tuple[bytes, str, str, Dict[str, int]]:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter

        from app.services.scm.summary_order_service import write_sheet

        axis: List[str] = listing["months"]
        rows: List[dict] = listing["data"]

        # R17: "No date" and "No location" leave the workbook (they still leave the screen
        # too - the row itself still carries `undated`/`unlocated`, just not as columns
        # here or in `Total`).
        columns = (
            ["Product", "Name", "Category", "Supplier"]
            + [_export_month_label(key) for key in axis]
            + ["TBA", "Total"]
        )
        widths = [16, 30, 14, 24] + [10] * len(axis) + [10, 12]
        width_map = {
            get_column_letter(index + 1): width for index, width in enumerate(widths)
        }

        # R5/AC-13: `split == "none"` is one sheet, one fixed title - never derived from a
        # row's own data. Any other split groups by the shared `workbook_split.split_rows`
        # (AC-2, PLAN-low-stock-export-split-25sep) - the SAME grouping and sanitised-title
        # sort the low stock report's own split uses, lifted here rather than reinvented.
        ordered: List[Tuple[str, List[dict]]] = (
            [("Stock debt", rows)]
            if split == "none"
            else split_rows(
                rows, split,
                supplier=lambda r: r["supplier_name"],
                category=lambda r: r["category_code"],
            )
        )

        wb = Workbook()
        used_titles: Set[str] = set()
        sheet_count = 0
        for index, (key, group_rows) in enumerate(ordered):
            ws = wb.active if index == 0 else wb.create_sheet()
            ws.title = key if split == "none" else unique_sheet_title(key, used_titles)
            sheet_count += 1
            data = [self._export_row(row, axis) for row in group_rows]
            data.append(self._export_total_row(group_rows, axis))
            write_sheet(ws, columns, data, width_map)

        buf = BytesIO()
        wb.save(buf)
        return (
            buf.getvalue(),
            EXPORT_CONTENT_TYPE,
            f"stock-debt-{date.today().strftime('%d%m%Y')}.xlsx",
            {"rows": len(rows), "sheets": sheet_count},
        )

    @staticmethod
    def _export_row(row: dict, axis: Sequence[str]) -> tuple:
        """One product, in the export's own column order (AC-13). `Name` prints blank when
        `list()` has already nulled it (AC-9); every other blank prints as `""`, never a
        bare 0 that would read as a fact somebody measured.

        Every text cell goes through `_xlsx_safe_text` (AC-13b, security review) - a
        supplier or a product code is free text off somebody else's document, and a
        leading `=`/`+`/`-`/`@` reaches openpyxl unescaped otherwise, which a spreadsheet
        reads as a formula the moment the file is opened. The same guard
        `low_stock_report_service._sheet_row` already applies to every one of ITS text
        cells.
        """
        from app.services.scm.proforma_invoice_service import _xlsx_safe_text

        balances = {month["key"]: month["balance"] for month in row["months"]}
        return (
            _xlsx_safe_text(row["product_code"]),
            _xlsx_safe_text(row["product_name"] or ""),
            _xlsx_safe_text(row["category_code"] or ""),
            _xlsx_safe_text(row["supplier_name"] or ""),
            *[balances.get(key, 0.0) for key in axis],
            row["tba"],
            row["total"],
        )

    @staticmethod
    def _export_total_row(rows: List[dict], axis: Sequence[str]) -> tuple:
        """The sheet's OWN Total footer (R5): summed over the rows THIS sheet carries, so
        a buyer reading one supplier's tab foots it without opening the others. R17: no
        "No date"/"No location" columns to foot any more."""
        month_sums = {key: 0.0 for key in axis}
        tba = total = 0.0
        for row in rows:
            for month in row["months"]:
                month_sums[month["key"]] += month["balance"]
            tba += row["tba"]
            total += row["total"]
        return (
            "Total", "", "", "",
            *[month_sums[key] for key in axis],
            tba, total,
        )
