"""Dealer targets Excel import (lane SALES-TARGET-DEALER, UAC D3).

The owner's gate (16 Oct 00:19): every row of the Sorento and Mocha files must be representable
on the one target engine, and the dry run says, per row, whether it is and why not. The layout
this reads is `documentation/plans/sales/dealer-targets-import-layout.md`; the real files never
enter the repo, so crew runs the dry run on them locally and sends back the patterns it rejects.

Two halves:

- `parse_workbook` (pure): sheets -> header by name -> merged cells resolved -> rows grouped into
  targets (a merged period cell over several rows is one target over those accounts) -> each
  period cell read by `parse_tier_cell` -> one `Scheme` per target and scheme kind (quarterly,
  monthly, yearly), with a `RowResult` per data row.
- `run` (database): codes resolved to the acting company's customers, then every scheme is
  created or updated through `target_service` (so the import obeys exactly the rules the screen
  does), each in its own savepoint; a scheme the service refuses marks its rows. A dry run is
  the same work in a savepoint the route rolls back, so its report is what an apply would do.
"""
from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass, field, replace
from datetime import date
from decimal import Decimal
from typing import Dict, List, Optional, Tuple

from sqlalchemy import func
from sqlalchemy.orm import Session

from app.services.error_handler import AppException
from app.services.sales.period_service import add_months

# --------------------------------------------------------------------------------------
# the cell grammar
# --------------------------------------------------------------------------------------

_BLANK = {"", "-", "--", "n/a", "na", "nil", "none", "x"}
_MONTHS = {
    m: i
    for i, names in enumerate(
        [
            ("JAN", "JANUARY"), ("FEB", "FEBRUARY"), ("MAR", "MARCH"), ("APR", "APRIL"),
            ("MAY",), ("JUN", "JUNE"), ("JUL", "JULY"), ("AUG", "AUGUST"),
            ("SEP", "SEPT", "SEPTEMBER"), ("OCT", "OCTOBER"), ("NOV", "NOVEMBER"),
            ("DEC", "DECEMBER"),
        ],
        start=1,
    )
    for m in names
}
_MONTH_RE = "|".join(sorted(_MONTHS, key=len, reverse=True))
#: `DEC'25-NOV'26`, `Jan-Dec 2026`, `JAN-DEC'26`, `Dec 25 - Nov 26`.
_WINDOW = re.compile(
    rf"\b({_MONTH_RE})\s*'?\s*(\d{{2}}|\d{{4}})?\s*(?:-|TO|\u2013)\s*({_MONTH_RE})\s*'?\s*(\d{{2}}|\d{{4}})\b",
    re.IGNORECASE,
)
_PCT = re.compile(r"(\d+(?:\.\d+)?)\s*%")
#: A money figure: `RM 300,000`, `300,000.50`, `300k`, `1.2m`, `300000`. Never followed by `%`.
_AMOUNT = re.compile(
    r"(?<![\w.])(?:RM\s*)?(\d{1,3}(?:,\d{3})+|\d+)(\.\d+)?\s*([kKmM])?(?![\w.,]*\s*%)(?![\d])"
)
_ADDITIONAL = re.compile(r"\+?\s*\badd(?:itional|'?l)?\b\.?", re.IGNORECASE)


class CellError(ValueError):
    """A cell this reader cannot turn into tiers; the message quotes it."""


class NotARate(CellError):
    """An incentive cell holding a figure where a rate belongs ("A/C I- 300,000", SRT 15-17).
    Its tier is skipped with a warning on the rows and the rest of the target loads (owner
    20 Oct 10:11: a typo, the corrected sheet adds the tier later)."""

    def __init__(self, cell: str, why: str = ""):
        super().__init__(f'"{cell}" is not a rate' + (f": {why}" if why else ""))
        self.cell, self.why = cell, why


class AccountPick(CellError):
    """A rate names an account no row of the dealer has ("A/C 9", SRT 57-59): the rows wait on
    the user picking which of the dealer's accounts it means, never on a guess (round 10)."""

    def __init__(self, message: str, pick_key: str, label: str, choices: List[str]):
        super().__init__(message)
        self.pick_key, self.label, self.choices = pick_key, label, choices


@dataclass
class CellTier:
    threshold: Optional[Decimal]
    rate: Optional[Decimal]
    is_additional: bool = False
    #: Account label as written ("A/C I", "Deluxe") -> rate.
    account_rates: Optional[Dict[str, Decimal]] = None


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", text.replace(" ", " ")).strip()


def _label_key(text: str) -> str:
    return re.sub(r"[^a-z0-9]", "", text.lower())


def _amount(match: re.Match) -> Decimal:
    whole = match.group(1).replace(",", "")
    value = Decimal(whole + (match.group(2) or ""))
    suffix = (match.group(3) or "").lower()
    if suffix == "k":
        value *= 1000
    elif suffix == "m":
        value *= 1000000
    return value


def _number(value) -> Optional[Decimal]:
    """A number cell, or a text cell that is only a number (`"100,000"`, `"RM 50,000"`, `"1%"`)."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    raw = str(value)
    if len(raw) > MAX_CELL_TEXT:
        return None
    text = re.sub(r"^(?:\u2265|>=?|=>)\s*", "", _norm(raw))
    m = re.fullmatch(r"(?:RM\s*)?(\d{1,3}(?:,\d{3})+|\d+)(\.\d+)?\s*([kKmM])?\s*%?", text, re.IGNORECASE)
    return _amount(m) if m else None


def _rates_in(text: str) -> Tuple[List[Tuple[str, Decimal]], Optional[Decimal]]:
    """`[(label, rate)]` for every `label n%` in `text`, and the rate written with no label."""
    named: List[Tuple[str, Decimal]] = []
    plain: Optional[Decimal] = None
    last = 0
    for m in _PCT.finditer(text):
        before = text[last:m.start()]
        label = re.sub(r"^[\s/,;|&+@=:()\-]+|[\s/,;|&+@=:()\-]+$", "", before)
        label = re.sub(r"^(?:and|at)\b\s*", "", label, flags=re.IGNORECASE).strip()
        rate = Decimal(m.group(1))
        if label and re.search(r"[A-Za-z\[]", label) and not re.fullmatch(r"(?i)rm|rate|@", label):
            named.append((_norm(label), rate))
        elif plain is None:
            plain = rate
        last = m.end()
    return named, plain


def _read_part(text: str) -> Tuple[Optional[Decimal], List[Tuple[str, Decimal]], Optional[Decimal]]:
    """(threshold, named rates, plain rate) of one part of a line."""
    amount = _AMOUNT.search(text)
    threshold = _amount(amount) if amount else None
    rest = text[: amount.start()] + " " + text[amount.end():] if amount else text
    named, plain = _rates_in(rest)
    return threshold, named, plain


def parse_tier_cell(value, *, default_rate: Optional[Decimal] = None) -> List[CellTier]:
    """The tiers one period cell holds (layout doc, "Cells"); `[]` for no scheme."""
    if value is None:
        return []
    number = None if isinstance(value, str) else _number(value)
    if number is not None:
        return [CellTier(number, default_rate)]
    if len(str(value)) > MAX_CELL_TEXT:
        raise CellError(f"a cell of more than {MAX_CELL_TEXT} characters")
    text = _norm(str(value))
    if text.lower() in _BLANK:
        return []
    if len(text) > MAX_CELL_TEXT:
        raise CellError(f"a cell of more than {MAX_CELL_TEXT} characters")
    tiers: List[CellTier] = []
    for raw_line in re.split(r"[\n;|]+", str(value)):
        line = _WINDOW.sub(" ", _norm(raw_line)).strip(" :")
        if not line or line.lower() in _BLANK:
            continue
        split = _ADDITIONAL.search(line)
        main, extra = (line[: split.start()], line[split.end():]) if split else (line, None)
        threshold, named, plain = _read_part(main)
        main_tier: Optional[CellTier] = None
        if named or plain is not None:
            main_tier = CellTier(threshold, None if named else plain, False, dict(named) if named else None)
        elif threshold is not None and extra is None:
            main_tier = CellTier(threshold, default_rate)
        if main_tier is not None:
            tiers.append(main_tier)
        if extra is not None:
            x_threshold, x_named, x_plain = _read_part(extra)
            if x_plain is None and not x_named:
                raise CellError(f'cannot read "{_norm(raw_line)}"')
            tiers.append(
                CellTier(
                    x_threshold if x_threshold is not None else threshold,
                    None if x_named else x_plain,
                    True,
                    dict(x_named) if x_named else None,
                )
            )
        if main_tier is None and extra is None:
            raise CellError(f'cannot read "{_norm(raw_line)}"')
    return tiers


def _window(text: str, year: int) -> Optional[Tuple[date, date]]:
    """The start and end a header or cell states (`DEC'25-NOV'26`, `JAN-DEC'26`)."""
    m = _WINDOW.search(text or "")
    if not m:
        return None

    def yr(token: Optional[str], fallback: int) -> int:
        if not token:
            return fallback
        value = int(token)
        return value + 2000 if value < 100 else value

    end_year = yr(m.group(4), year)
    start_month, end_month = _MONTHS[m.group(1).upper()], _MONTHS[m.group(3).upper()]
    start_year = yr(m.group(2), end_year if start_month <= end_month else end_year - 1)
    if not (2000 <= start_year <= 2100 and 2000 <= end_year <= 2100):
        raise CellError(f'the dates "{m.group(0)}" are out of range')
    start = date(start_year, start_month, 1)
    end = add_months(date(end_year, end_month, 1), 1).replace(day=1)
    if end <= start:
        raise CellError(f'the dates "{m.group(0)}" end before they start')
    return start, date.fromordinal(end.toordinal() - 1)


# --------------------------------------------------------------------------------------
# the sheet reader
# --------------------------------------------------------------------------------------


@dataclass
class SchemeTier:
    period_start: date
    threshold: Decimal
    rate: Optional[Decimal]
    is_additional: bool
    #: customer code -> rate
    account_rates: Optional[Dict[str, Decimal]]
    #: A fixed amount paid once per period when the tier is reached ("A/C I- 15,000" in an
    #: incentive column, round 10).
    bonus: Optional[Decimal] = None


@dataclass
class Scheme:
    key: str
    kind: str  # quarterly | monthly | yearly
    sheet: str
    rows: List[int]
    codes: List[str]
    names: List[str]
    payout_code: Optional[str]
    start: date
    end: date
    split_every: Optional[int]
    split_unit: Optional[str]
    threshold_unit: str
    rounding: str
    tiers: List[SchemeTier]
    #: "<dealer> <2026 | Dec 2025 to Nov 2026> <quarterly | monthly | yearly>", the kind last.
    name: str = ""
    #: Paid to an account of the dealer with no row in the file, found by name + account.
    payout_by_name: Optional["ByName"] = None
    #: The key of the dealer scheme that replaces this per-account one once it reaches a tier
    #: (crew ruling 16 Oct: Tier 1-3 per account, Tier 4-5 on the dealer, one scale).
    replaced_by: Optional[str] = None


@dataclass
class RowResult:
    sheet: str
    row: int
    code: Optional[str]
    name_in_file: Optional[str]
    account_in_file: Optional[str]
    status: str = "mapped"
    reason: Optional[str] = None
    warnings: List[str] = field(default_factory=list)
    target_keys: List[str] = field(default_factory=list)
    #: The CRM customer's name once the row matched (set by `run`).
    customer: Optional[str] = None
    #: How the row goes in a scheme: its code, or "<sheet>!<row>" (set by the reader).
    key: str = ""
    #: "<sheet>!<row>", the key a pick names (set by `run`).
    pick_key: str = ""
    #: auto | code | picked (set by `run`).
    match: Optional[str] = None
    #: Up to 3 customers to pick from, for a row with no exact name match (set by `run`).
    suggestions: List[dict] = field(default_factory=list)
    #: {pick_key, label, choices}: a rate's account label waiting on the user's pick (round 10).
    account_pick: Optional[dict] = None


@dataclass
class ParsedWorkbook:
    rows: List[RowResult]
    schemes: List[Scheme]
    notes: List[str]


@dataclass
class _Column:
    index: int
    role: str  # code | name | account | payto | round | window | period
    kind: Optional[str] = None  # Q | M | Y
    number: Optional[int] = None  # quarter or month
    window: Optional[Tuple[date, date]] = None
    points: bool = False
    default_rate: Optional[Decimal] = None
    is_rate: bool = False
    #: "Tier n": a threshold and its "Tier n Incentive" pair up by n inside one period.
    tier: Optional[int] = None
    #: The header texts from the top header row down, joined (" / ").
    header: str = ""


_KIND_NAME = {"Q": "quarterly", "M": "monthly", "Y": "yearly"}

#: A header cell that names a row's own field, matched on the WHOLE cell so a title such as
#: "SORENTO DEALER INCENTIVE 2026" is never read as the dealer column.
_ROLE_HEADERS = (
    ("payto", re.compile(r"PAY(?:OUT)?\s*TO(?:\s+(?:CODE|ACCOUNT))?")),
    ("code", re.compile(r"(?:(?:CUSTOMER|DEALER|DEBTOR|ACCOUNT)\s+)?CODE|CUST(?:OMER)?\.?\s*CODE")),
    ("round", re.compile(r"ROUND(?:ING)?")),
    ("account", re.compile(r"ACCOUNT(?:\s+(?:NO\.?|TYPE|LEVEL))?|A\s*/\s*C")),
    ("name", re.compile(r"(?:(?:DEALER|CUSTOMER|DEBTOR|COMPANY)\s+)?NAME|DEALERS?|CUSTOMERS?|DEBTORS?")),
)
#: Words a stacked period header is made of (below the dealer-name header row).
_PERIOD_WORDS = re.compile(
    rf"\bQ\s*[1-4]\b|QUARTER|YEAR|ANNUAL|MONTH|POINT|PTS|TIER|INCENTIVE|\b({_MONTH_RE})\b", re.IGNORECASE
)
#: A period header row stacked ABOVE the dealer-name row; narrower than `_PERIOD_WORDS` so a
#: sheet title ("SORENTO DEALER INCENTIVE 2026") above the header never joins it.
_GROUP_WORDS = re.compile(r"\bQ\s*[1-4]\b|QUARTERLY|YEARLY|ANNUAL|MONTHLY|\(POINTS?\)", re.IGNORECASE)


def _safe_window(text: str, year: int) -> Optional[Tuple[date, date]]:
    try:
        return _window(text, year)
    except CellError:
        return None


def _text(value) -> str:
    # Cut once, before any regex: `_PCT` / `_AMOUNT` are quadratic on a long digit run
    # (security re-check B3). One character over the cap keeps every length check firing.
    return "" if value is None else _norm(str(value)[: MAX_CELL_TEXT + 1])


def _role(text: str) -> Optional[str]:
    up = text.upper().strip(" :")
    for role, pattern in _ROLE_HEADERS:
        if pattern.fullmatch(up):
            return role
    return None


def _is_number(value) -> bool:
    return isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)


def _header_rows(grid: "_Grid") -> Optional[Tuple[int, int]]:
    """(first, last) header row: the row naming the dealer (or code) column, the period rows
    stacked above it, and the sub-header rows below it (`JAN-MAR`, `Tier 1`, `Tier 1
    Incentive`). None when no row in the first 20 names a dealer, customer or code column."""
    anchor = None
    for r in grid.rows:
        if r > 20:
            break
        roles = {_role(_text(grid.own(r, c))) for c in grid.row_cols.get(r, [])}
        if roles & {"name", "code"}:
            anchor = r
            break
    if anchor is None:
        return None
    # The columns the anchor names a row's own field in: a later row with its own value there
    # is a data row, whatever its words ("PTS TRADING", "MAY HARDWARE"; review S1).
    identity = {c for c in grid.row_cols.get(anchor, []) if _role(_text(grid.own(anchor, c))) in ("name", "code", "account")}
    first = anchor
    while first - 1 >= 1 and anchor - first < 3:
        texts = [_text(grid.own(first - 1, c)) for c in grid.row_cols.get(first - 1, [])]
        texts = [t for t in texts if t]
        if not texts or not any(_GROUP_WORDS.search(t) for t in texts):
            break
        first -= 1
    last = anchor
    while last + 1 in grid.row_cols and last - anchor < 4:
        r = last + 1
        own = [grid.own(r, c) for c in grid.row_cols[r]]
        own = [v for v in own if v is not None]
        if not own or any(_is_number(v) for v in own):
            break
        if any(grid.own(r, c) is not None for c in identity):
            break
        texts = [_text(v) for v in own]
        if any(_PCT.search(t) and _AMOUNT.search(t) for t in texts):
            break
        if not all(_PERIOD_WORDS.search(t) or _PCT.search(t) or _role(t) or _safe_window(t, 2000)
                   or re.fullmatch(r"(?i)%|RM|TARGET|RATE|THRESHOLD|PERIOD|WINDOW|SALES", t) for t in texts):
            break
        last = r
    return first, last


def _classify(chain: List[str], year: int) -> Optional[_Column]:
    """One column, from its header texts top to bottom (a merged header reads in every column
    it spans)."""
    texts: List[str] = []
    for t in chain:
        if t and (not texts or texts[-1] != t):
            texts.append(t)
    if not texts:
        return None
    header = " / ".join(texts)
    for t in texts:
        role = _role(t)
        if role:
            return _Column(0, role, header=header)
    up = [t.upper() for t in texts]
    joined = " / ".join(up)
    points = "POINT" in joined or "PTS" in joined
    is_rate = any(
        "INCENTIVE" in t or re.search(r"\bRATE\b|\bCOMM", t) or t.strip() == "%" for t in up
    )
    default_rate = None
    for t in up:
        if "INCENTIVE" in t or t.strip() == "%":
            continue
        m = _PCT.search(t)
        if m:
            default_rate = Decimal(m.group(1))
    tier = None
    for t in up:
        m = re.search(r"\bTIER\s*(\d+)", t)
        if m:
            tier = int(m.group(1))
    quarter = next((int(m.group(1)) for t in up for m in [re.search(r"\bQ\s*([1-4])\b", t)] if m), None)
    window = None
    for t in up:
        w = _safe_window(t, year)
        if w:
            window = w
    yearly = re.search(r"\b(YEAR|YEARLY|ANNUAL|ANNUALLY|FY)\b|\(YEARLY\)", joined) is not None or (
        window is not None and (window[1].year - window[0].year) * 12 + window[1].month - window[0].month >= 11
    )
    month = None
    for t in up:
        m = re.fullmatch(rf"({_MONTH_RE})\b[\s'0-9]*(?:\(.*\))?(?:\s*POINTS?)?", t)
        if m:
            month = _MONTHS[m.group(1)]
    quarterly = "QUARTER" in joined
    monthly = re.search(r"\bMONTHLY\b", joined) is not None
    if quarter is None and quarterly:
        for t in up:
            # "QUARTERLY (POINT)" over "JAN-MAR": the quarter its months are (gate 2, item 4).
            m = re.fullmatch(rf"({_MONTH_RE})\s*(?:-|TO|\u2013)\s*({_MONTH_RE})", t.strip())
            if m:
                first, last = _MONTHS[m.group(1)], _MONTHS[m.group(2)]
                if first % 3 == 1 and last == first + 2:
                    quarter = (first - 1) // 3 + 1
    col = _Column(0, "period", points=points, default_rate=default_rate, is_rate=is_rate, tier=tier, header=header)
    if quarter is not None:
        col.kind, col.number = "Q", quarter
    elif month is not None and not yearly:
        col.kind, col.number = "M", month
    elif yearly:
        col.kind, col.window = "Y", window
        # A yearly "PERIOD" or "Special Remark" column states the block's window, never a
        # threshold (gate 8: Sorento's remark, merged over the accounts, holds "DEC'25-NOV'26").
        if re.fullmatch(r"PERIOD|WINDOW|DATES?", up[-1]) or re.search(r"\bREMARKS?\s*$", up[-1]):
            col.role = "window"
    elif quarterly and is_rate:
        # One INCENTIVE under a QUARTERLY block: every quarter's rate (number 0 = all).
        col.kind, col.number = "Q", 0
    elif monthly:
        # A "Monthly" block with tiers and no months: the same tiers every month (number 0).
        col.kind, col.number = "M", 0
    elif _PERIOD_WORDS.search(joined):
        # Period words the reader cannot place (gate 2, item 4): never dropped silently, the
        # rows holding a figure in it fail (`_parse_sheet`).
        col.role = "unplaced"
    else:
        return None
    return col


def _columns(grid: "_Grid", first: int, last: int, year: int) -> List[_Column]:
    cols: List[_Column] = []
    for c in grid.cols:
        chain = [_text(grid.value(r, c)) for r in range(first, last + 1)]
        col = _classify(chain, year)
        if col is not None:
            col.index = c
            cols.append(col)
    # An incentive column right after a period block whose merged header stops short of it
    # (Mocha: "QUARTERLY (POINT)" merged over JAN-MAR..OCT-DEC, "INCENTIVE" beside it; gate 3)
    # belongs to that block: every quarter's rate on a quarterly block, else its neighbour's.
    by_index = {col.index: col for col in cols}
    for col in cols:
        left = by_index.get(col.index - 1)
        if (
            col.role == "period" and col.kind == "Q" and col.number == 0 and col.is_rate
            and left is not None and left.role == "period" and left.kind == "Q"
            and not left.is_rate and left.tier is None
        ):
            # Mocha F3 "INCENTIVE (<note with QUARTER>)": placed as every quarter's rate by its
            # own words, it still pays the block on its left, points included (gate 4).
            col.points = left.points
            continue
        if (
            col.role == "unplaced"
            and re.match(r"INCENTIVE\b", col.header.split(" / ")[-1].strip().upper())
            and left is not None
            and left.role == "period"
            and left.kind == "Q"
            and not left.is_rate
            and left.tier is None
        ):
            # Only the Mocha shape (review S2): any other rate column the reader cannot place
            # still fails its rows.
            col.role, col.kind, col.points, col.number = "period", "Q", left.points, 0
    return cols


#: Bounds on what one upload may make the reader do (security review B1): a sheet's declared
#: size is not trusted, only the cells that hold a value are visited, and a sheet past these is
#: skipped with a note. The real files are a few hundred rows wide of a few dozen columns.
MAX_SHEETS = 30
#: Excel's own limit on a sheet's name.
MAX_SHEET_TITLE = 31
#: Excel's own limit on a number format code.
MAX_NUMBER_FORMAT = 255
MAX_SHEET_CELLS = 100_000
MAX_SHEET_ROWS = 5_000
MAX_SHEET_COLS = 300
MAX_MERGE_AREA = 2_000
MAX_CELL_TEXT = 500
#: The highest incentive rate the import accepts (crew ruling 16 Oct, review S3): 0.5 reads as
#: 50% under the fraction rule where a 0.5% rebate may have been meant, so a rate above this
#: fails its row rather than risk a 100x overpay. The real files top out at 10%.
RATE_LIMIT_PCT = 20
#: A dealer name, code or account label; the real ones are under 100 characters.
MAX_ID_TEXT = 200
#: The workbook's XML once unzipped; a 10 MB upload that inflates past this is refused.
MAX_UNZIPPED_BYTES = 150 * 1024 * 1024


class SheetTooLarge(ValueError):
    pass


class _Grid:
    """The cells that hold a value, with merged ranges resolved: every cell of a merge reads the
    top-left value and knows the range's rows. Sparse: nothing is visited by the sheet's declared
    size, which one stray cell far away can make a billion cells (security review B1)."""

    def __init__(self, ws):
        self.ws = ws
        self.values: Dict[Tuple[int, int], object] = {}
        #: Number formats of numeric cells that show a percent (0.02 shown as 2%).
        self.formats: Dict[Tuple[int, int], str] = {}
        for (r, c), cell in getattr(ws, "_cells", {}).items():
            value = cell.value
            if isinstance(value, str):
                # Cut once, here: one long shared string can sit in every cell (security
                # re-check B4). One character over keeps every length check firing.
                value = value[: MAX_CELL_TEXT + 1]
            if value is None or (isinstance(value, str) and not value.strip()):
                continue
            self.values[(r, c)] = value
            fmt = cell.number_format
            # A format past Excel's own 255 characters is not a real one; never scanned, since one
            # shared format would be walked once per cell (security re-check B6).
            if _is_number(value) and fmt and len(fmt) <= MAX_NUMBER_FORMAT and "%" in fmt:
                self.formats[(r, c)] = fmt
            if len(self.values) > MAX_SHEET_CELLS:
                raise SheetTooLarge(f"more than {MAX_SHEET_CELLS:,} filled cells")
        self.merge: Dict[Tuple[int, int], Tuple[int, int, int, int]] = {}
        for rng in ws.merged_cells.ranges:
            area = (rng.max_row - rng.min_row + 1) * (rng.max_col - rng.min_col + 1)
            if area > MAX_MERGE_AREA:
                # A title banner merged across the sheet, never a threshold over a few rows.
                continue
            box = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
            if len(self.merge) + area > MAX_SHEET_CELLS:
                raise SheetTooLarge(f"more than {MAX_SHEET_CELLS:,} merged cells")
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.merge[(r, c)] = box
        cells = set(self.values) | {
            key for key, box in self.merge.items() if (box[0], box[1]) in self.values
        }
        self.rows = sorted({r for r, _ in cells})
        self.cols = sorted({c for _, c in cells})
        if len(self.rows) > MAX_SHEET_ROWS:
            raise SheetTooLarge(f"more than {MAX_SHEET_ROWS:,} rows")
        if len(self.cols) > MAX_SHEET_COLS:
            raise SheetTooLarge(f"more than {MAX_SHEET_COLS} columns")
        self.row_cols: Dict[int, List[int]] = {}
        for r, c in sorted(cells):
            self.row_cols.setdefault(r, []).append(c)

    def value(self, r: int, c: int):
        box = self.merge.get((r, c))
        if box:
            r, c = box[0], box[1]
        return self.values.get((r, c))

    def own(self, r: int, c: int):
        """The cell's own value, not its merge's."""
        return self.values.get((r, c))

    def box(self, r: int, c: int) -> Optional[Tuple[int, int, int, int]]:
        return self.merge.get((r, c))

    def fmt(self, r: int, c: int) -> Optional[str]:
        box = self.merge.get((r, c))
        if box:
            r, c = box[0], box[1]
        return self.formats.get((r, c))



def _union(groups: Dict[int, int], a: int, b: int) -> None:
    def find(x):
        while groups[x] != x:
            groups[x] = groups[groups[x]]
            x = groups[x]
        return x

    ra, rb = find(a), find(b)
    if ra != rb:
        groups[max(ra, rb)] = min(ra, rb)


def _root(groups: Dict[int, int], x: int) -> int:
    while groups[x] != x:
        x = groups[x]
    return x


def _rounding(text: str) -> Optional[str]:
    t = text.strip().lower()
    if not t:
        return None
    if t in {"sen", "0.01", "cent", "cents", "2", "2dp"}:
        return "sen"
    if t in {"ringgit", "rm", "1", "whole", "0", "0dp"}:
        return "ringgit"
    return "?"



# --------------------------------------------------------------------------------------
# the incentive grammar (a "Tier n Incentive" cell)
# --------------------------------------------------------------------------------------

#: Words that mean "no figure here" in an incentive or threshold cell. A cell holding one of
#: them BESIDE a figure is reported, never read as nothing (gate-fail round item 7).
_NIL = {"nil", "n/a", "na", "none", "x"}
_DASH = {"", "-", "--"}
_BARE_RATE = re.compile(r"^(?P<label>.*?)[\s\-:=@]*(?P<num>\d+(?:,\d{3})*(?:\.\d+)?)$")


@dataclass
class RateEntry:
    label: Optional[str]
    rate: Decimal
    is_additional: bool = False
    #: An account and an amount where a rate belongs, read as a fixed bonus (round 10).
    bonus: Optional[Decimal] = None


def _is_nil(value) -> bool:
    return isinstance(value, str) and _norm(value).lower() in _NIL


def _is_dash(value) -> bool:
    return value is None or (isinstance(value, str) and _norm(value).lower() in _DASH)


def _clean_label(text: str) -> Optional[str]:
    label = re.sub(r"^[\s/,;|&+@=:()\-]+|[\s/,;|&+@=:()\-]+$", "", text)
    label = re.sub(r"^(?:and|at)\b\s*", "", label, flags=re.IGNORECASE).strip()
    return _norm(label) if label and re.search(r"[A-Za-z\[]", label) else None


def _fraction(n: Decimal, percent_format: bool) -> Decimal:
    """A rate typed as a number: 0.02 is 2% (Excel keeps a % cell as its fraction, and the files
    type fractions too); 3 is 3%."""
    return n * 100 if percent_format or n < 1 else n


def parse_rate_cell(
    value, *, percent_format: bool = False, threshold: Optional[Decimal] = None
) -> List[RateEntry]:
    """The rates one incentive cell holds: `0.02`, `2%`, `3`, `A/C I- 3%`, `Deluxe 2% / Mastile
    3%`, `+ Additional 1%`, `3% + Additional 1%`. `[]` for an empty or dash cell; a NIL word
    raises (its caller decides whether the threshold beside it makes that an error).

    `A/C I- 15,000` (an account and an amount, SRT 15-17) is a fixed bonus on the tier when the
    amount is at most `RATE_LIMIT_PCT` of the tier's `threshold`; a larger amount, or an amount
    with no account, stays "not a rate" (round 10)."""
    if _is_dash(value):
        return []
    if _is_number(value):
        return [RateEntry(None, _fraction(Decimal(str(value)), percent_format))]
    raw = str(value)
    if len(raw) > MAX_CELL_TEXT:
        raise CellError(f"a cell of more than {MAX_CELL_TEXT} characters")
    text = _norm(raw)
    if text.lower() in _NIL:
        raise CellError(f'"{text}" is no rate')
    out: List[RateEntry] = []
    for line in re.split(r"[\n;|]+", raw):
        line = _norm(line)
        if not line or line.lower() in _DASH:
            continue
        split = _ADDITIONAL.search(line)
        parts = [(line[: split.start()], False), (line[split.end():], True)] if split else [(line, False)]
        for part, additional in parts:
            part = part.strip(" +:,")
            if not part:
                continue
            if _PCT.search(part):
                named, plain = _rates_in(part)
                if len(list(_PCT.finditer(part))) - len(named) > 1:
                    raise CellError(f'"{line}" holds two rates for one tier')
                tail = part[[m for m in _PCT.finditer(part)][-1].end():]
                if re.search(r"\d", tail):
                    raise NotARate(line)
                out.extend(RateEntry(label, rate, additional) for label, rate in named)
                if plain is not None:
                    out.append(RateEntry(None, plain, additional))
                continue
            m = _BARE_RATE.match(part)
            if not m:
                raise CellError(f'cannot read the rate "{line}"')
            n = Decimal(m.group("num").replace(",", ""))
            label = _clean_label(m.group("label"))
            if n > 100:
                if label and threshold and not additional and n <= threshold * RATE_LIMIT_PCT / 100:
                    out.append(RateEntry(label, Decimal(0), False, bonus=n))
                    continue
                if threshold:
                    share = f"{(n * 100 / threshold).quantize(Decimal('1')):f}"
                    raise NotARate(
                        line,
                        f"{n:,.0f} is {share}% of the threshold {threshold:,.0f}, "
                        f"more than a bonus of up to {RATE_LIMIT_PCT}%",
                    )
                raise NotARate(line)
            out.append(RateEntry(label, _fraction(n, False), additional))
    if not out:
        raise CellError(f'cannot read the rate "{text}"')
    return out




def parse_workbook(data: bytes, *, year: int, label_picks: Optional[Dict[str, str]] = None) -> ParsedWorkbook:
    import zipfile

    from openpyxl import load_workbook

    not_a_workbook = AppException(
        status_code=422, message="That file is not an Excel workbook (.xlsx).", code="NOT_A_WORKBOOK"
    )
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            unzipped = sum(info.file_size for info in archive.infolist())
    except zipfile.BadZipFile as exc:
        raise not_a_workbook from exc
    if unzipped > MAX_UNZIPPED_BYTES:
        raise AppException(status_code=422, message="That workbook is too large.", code="FILE_TOO_LARGE")
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # InvalidFileException, KeyError, ...
        raise not_a_workbook from exc
    if len(wb.worksheets) > MAX_SHEETS:
        raise AppException(
            status_code=422, message=f"That workbook has more than {MAX_SHEETS} sheets.", code="FILE_TOO_LARGE"
        )
    if any(len(ws.title) > MAX_SHEET_TITLE for ws in wb.worksheets):
        # A title is copied into every row's key and reason (security re-check B5); Excel
        # itself stops at 31 characters, so a longer one is not a real file.
        raise AppException(
            status_code=422,
            message=f"A sheet name is longer than {MAX_SHEET_TITLE} characters.",
            code="FILE_TOO_LARGE",
        )
    rows: List[RowResult] = []
    schemes: List[Scheme] = []
    notes: List[str] = []
    for ws in wb.worksheets:
        try:
            grid = _Grid(ws)
        except SheetTooLarge as exc:
            # Nothing of it is read, so nothing of it may pass for read (item 8).
            notes.append(f'Sheet "{ws.title}": {exc}; not read.')
            continue
        header = _header_rows(grid)
        if header is None:
            # Every filled row of a sheet the reader cannot read fails, so a dropped sheet is
            # never "all mapped" (item 8).
            reason = f'sheet "{ws.title}" not read: no DEALER NAME or CUSTOMER CODE header in its first 20 rows'
            for r in grid.rows:
                res = RowResult(sheet=ws.title, row=r, code=None, name_in_file=None, account_in_file=None)
                res.status, res.reason = "not_representable", reason
                rows.append(res)
            continue
        sheet_rows, sheet_schemes = _parse_sheet(grid, ws.title, header, year, notes, label_picks or {})
        rows.extend(sheet_rows)
        schemes.extend(sheet_schemes)
    return ParsedWorkbook(rows, schemes, notes)


def gate_passes(result: dict) -> bool:
    """The owner's gate on one file: every row mapped, none waiting on a pick, nothing skipped
    unread. The dry-run script's exit code is this."""
    summary = result["summary"]
    return not result.get("notes") and summary.get("not_representable", 0) == 0 and summary.get("needs_pick", 0) == 0


@dataclass(frozen=True)
class ByName:
    """An account paid by the dealer's name and its account label: it has no row in the file
    (gate 3). Its own type, never a string, so no file text can be read as one (review B2)."""

    dealer: str
    label: str


def _row_key(sheet: str, row: int, code: Optional[str]) -> str:
    """How a row goes in a scheme: its customer code when the file has one, else its place."""
    return code or f"{sheet}!{row}"


def _parse_sheet(
    grid: _Grid, sheet: str, header: Tuple[int, int], year: int, notes: List[str], label_picks: Dict[str, str]
):
    first_header, last_header = header
    cols = _columns(grid, first_header, last_header, year)
    by_role: Dict[str, Optional[_Column]] = {
        role: next((c for c in cols if c.role == role), None)
        for role in ("code", "name", "account", "payto", "round")
    }
    window_col = next((c for c in cols if c.role == "window"), None)
    periods = [c for c in cols if c.role == "period"]
    first = last_header + 1

    def cell(r: int, col: Optional[_Column]) -> Optional[str]:
        return (_text(grid.value(r, col.index)) or None) if col else None

    def too_long(*texts: Optional[str]) -> bool:
        return any(t and len(t) > MAX_ID_TEXT for t in texts)

    results: Dict[int, RowResult] = {}
    for r in grid.rows:
        if r < first:
            continue
        # A row only a merge from above reaches (no value of its own) is still a data row: it is
        # one of the accounts that merged threshold covers.
        if not any(_text(grid.value(r, c)) for c in grid.row_cols.get(r, [])):
            continue
        code = cell(r, by_role["code"])
        res = RowResult(
            sheet=sheet,
            row=r,
            code=code,
            name_in_file=cell(r, by_role["name"]),
            account_in_file=cell(r, by_role["account"]),
        )
        if res.name_in_file and not res.account_in_file and not too_long(res.name_in_file):
            # The real files carry the account as a SUFFIX of the dealer name cell ("<dealer>
            # A/C I"), with no ACCOUNT column (gate 2, item 1). After the length check.
            res.name_in_file, res.account_in_file = split_account_suffix(res.name_in_file)
        if too_long(code, res.name_in_file, res.account_in_file):
            # Security review B1: an unbounded name never reaches the matcher (difflib indexes
            # it) or the report; the row fails and only a cut of it is echoed back.
            res.code = (code or "")[:MAX_ID_TEXT] or None
            res.name_in_file = (res.name_in_file or "")[:MAX_ID_TEXT] or None
            res.account_in_file = (res.account_in_file or "")[:MAX_ID_TEXT] or None
            res.key = _row_key(sheet, r, None)
            res.status = "not_representable"
            res.reason = f"a name, code or account of more than {MAX_ID_TEXT} characters"
            results[r] = res
            continue
        res.key = _row_key(sheet, r, code)
        if not code and not res.name_in_file:
            res.status, res.reason = "not_representable", "no dealer name or customer code on this row"
        elif not periods:
            res.status, res.reason = "not_representable", f'sheet "{sheet}": no Q1-Q4, month or yearly column'
        results[r] = res

    # A dealer is the rows its name cell covers (merged down) or repeats on, one after another;
    # a rate naming an account resolves among the dealer's rows (ruling 2, pay-to).
    dealer_of: Dict[int, Tuple] = {}
    previous: Optional[Tuple[int, str]] = None
    for r in sorted(results):
        res = results[r]
        box = grid.box(r, by_role["name"].index) if by_role["name"] else None
        name_key = _label_key(res.name_in_file or res.key)
        if box:
            dealer_of[r] = ("merge", box[0], box[1])
        elif previous and previous[1] == name_key and previous[0] == r - 1:
            dealer_of[r] = dealer_of[previous[0]]
        else:
            dealer_of[r] = ("row", r)
        previous = (r, name_key)
    labels_of: Dict[Tuple, Dict[str, str]] = {}
    #: The dealer's first row (where an account pick is keyed) and its accounts as written.
    first_of: Dict[Tuple, int] = {}
    accounts_of: Dict[Tuple, Dict[str, str]] = {}
    for r, res in sorted(results.items()):
        labels = labels_of.setdefault(dealer_of[r], {})
        first_of.setdefault(dealer_of[r], r)
        if res.account_in_file:
            accounts_of.setdefault(dealer_of[r], {}).setdefault(_label_key(res.account_in_file), res.account_in_file)
        for text in (res.account_in_file, res.code):
            if text:
                labels.setdefault(_label_key(text), res.key)

    # A column under period words the reader cannot place is never dropped silently: every
    # row with a figure in it fails (gate 2, item 4).
    for col in (c for c in cols if c.role == "unplaced"):
        for r, res in results.items():
            if res.status == "mapped" and _text(grid.value(r, col.index)):
                res.status = "not_representable"
                res.reason = (
                    f'column "{col.header}" holds figures the reader cannot place '
                    "(not a Q1-Q4, month or yearly column)"
                )
    data_rows = [r for r, res in sorted(results.items()) if res.status == "mapped"]
    schemes: List[Scheme] = []
    # Rows group per scheme kind AND unit: a yearly threshold merged over two accounts is one
    # yearly target over both while their own quarterly thresholds stay two targets (rule 2);
    # an RM column and a POINT column never share a target (item 4).
    for kind in ("Q", "M", "Y"):
        for points in (False, True):
            kind_cols = [c for c in periods if c.kind == kind and c.points == points]
            if not kind_cols:
                continue
            groups = {r: r for r in data_rows}
            for col in kind_cols:
                for r in data_rows:
                    box = grid.box(r, col.index)
                    if box:
                        for other in data_rows:
                            if other != r and box[0] <= other <= box[2]:
                                _union(groups, r, other)
            members: Dict[int, List[int]] = {}
            for r in data_rows:
                members.setdefault(_root(groups, r), []).append(r)
            for _root_row, group_rows in sorted(members.items()):
                group_rows.sort()
                dealer = dealer_of[group_rows[0]]
                ctx = _Ctx(
                    grid, sheet, group_rows, results, labels_of[dealer], year,
                    dealer_row=first_of[dealer],
                    dealer_accounts=list(accounts_of.get(dealer, {}).values()),
                    label_picks=label_picks,
                )
                try:
                    schemes.extend(_group_schemes(ctx, kind, points, kind_cols, window_col, by_role))
                except AccountPick as exc:
                    for r in group_rows:
                        results[r].status, results[r].reason = "needs_pick", str(exc)
                        results[r].account_pick = {"pick_key": exc.pick_key, "label": exc.label, "choices": exc.choices}
                except CellError as exc:
                    for r in group_rows:
                        results[r].status, results[r].reason = "not_representable", str(exc)
    # A row whose own scheme failed fails every scheme it is in (no partial dealer).
    failed = {r for r, res in results.items() if res.status != "mapped"}
    kept = []
    for scheme in schemes:
        bad = [r for r in scheme.rows if r in failed]
        if bad:
            for r in scheme.rows:
                if r not in failed:
                    results[r].status = "not_representable"
                    results[r].reason = f"shares a target with row {bad[0]}: {results[bad[0]].reason}"
            continue
        kept.append(scheme)
    for scheme in kept:
        for r in scheme.rows:
            results[r].target_keys.append(scheme.key)
    for res in results.values():
        if res.status == "mapped" and not res.target_keys:
            res.warnings.append("no target figures on this row")
    return [results[r] for r in sorted(results)], kept


@dataclass
class _Ctx:
    grid: _Grid
    sheet: str
    rows: List[int]
    results: Dict[int, RowResult]
    #: label key -> row key, over the whole dealer (not only this target's rows).
    dealer_labels: Dict[str, str]
    year: int
    #: The dealer's first row, and its accounts as the file writes them (an account pick's choices).
    dealer_row: int = 0
    dealer_accounts: List[str] = field(default_factory=list)
    #: "<sheet>!<dealer row>~<label as written>" -> the dealer account the user picked for it.
    label_picks: Dict[str, str] = field(default_factory=dict)

    @property
    def keys(self) -> List[str]:
        return [self.results[r].key for r in self.rows]

    def per_row(self, col: _Column) -> Dict[int, object]:
        out = {}
        for r in self.rows:
            v = self.grid.value(r, col.index)
            if v is not None and _text(v):
                out[r] = v
        return out


def _account_and_dealer_columns(ctx: "_Ctx", cols: List[_Column]):
    """(per-account columns, dealer columns) when one period block holds thresholds that differ
    per account (Tier 1-3) beside thresholds merged over the accounts (Tier 4-5); else None."""
    shared_rates = [c for c in cols if c.kind != "Y" and c.number == 0 and c.is_rate]
    pairs, unpaired = _pairs([c for c in cols if c not in shared_rates], [])
    for col in unpaired:
        if ctx.per_row(col):
            # Never dropped by the split (review S1).
            raise CellError(f"{col.header}: an incentive with no threshold column before it")
    per_account: List[_Column] = list(shared_rates)
    dealer: List[_Column] = list(shared_rates)
    for threshold, rate in pairs:
        if not ctx.per_row(threshold):
            continue
        # By MERGE, never by equal values (review gate 2 B1): per-account thresholds that happen
        # to be the same number stay the accounts'.
        boxes = {ctx.grid.box(r, threshold.index) for r in ctx.rows}
        merged = len(boxes) == 1 and None not in boxes and all(
            next(iter(boxes))[0] <= r <= next(iter(boxes))[2] for r in ctx.rows
        )
        (dealer if merged else per_account).extend(c for c in (threshold, rate) if c is not None)
    if len(per_account) == len(shared_rates) or len(dealer) == len(shared_rates):
        return None
    return per_account, dealer


def _group_schemes(
    ctx: _Ctx, kind: str, points: bool, kind_cols, window_col, by_role, inherit_window=None
) -> List[Scheme]:
    """The one target of `kind` and unit over `ctx.rows`, or none. A block with Tier 1-3
    per account and Tier 4-5 merged over the accounts is one dealer target plus one target per
    account that the dealer target replaces once it reaches a tier (crew ruling 16 Oct)."""
    grid, results = ctx.grid, ctx.results
    if len(ctx.rows) > 1:
        split = _account_and_dealer_columns(ctx, list(kind_cols))
        if split is not None:
            per_account, dealer_cols = split
            dealer = _group_schemes(ctx, kind, points, dealer_cols, window_col, by_role)
            if not dealer:
                # No dealer tier ("NIL" merged over the accounts): each account keeps its own
                # tiers as its own target, never dropped (gate 8).
                out = []
                for r in ctx.rows:
                    single = replace(ctx, rows=[r])
                    out.extend(_group_schemes(single, kind, points, per_account, window_col, by_role))
                return out
            out = list(dealer)
            for r in ctx.rows:
                single = replace(ctx, rows=[r])
                own_schemes = _group_schemes(
                    single, kind, points, per_account, window_col, by_role,
                    inherit_window=(dealer[0].start, dealer[0].end),
                )
                for scheme in own_schemes:
                    if (scheme.start, scheme.end) != (dealer[0].start, dealer[0].end):
                        # The replacement needs the dealer's period to cover the account's.
                        raise CellError(
                            f"{_KIND_NAME[kind]}: an account's window ({scheme.start:%b %Y} to {scheme.end:%b %Y}) "
                            f"differs from its dealer's ({dealer[0].start:%b %Y} to {dealer[0].end:%b %Y})"
                        )
                    scheme.replaced_by = dealer[0].key
                    account = results[r].account_in_file
                    name = (results[r].name_in_file or "")[:80]
                    if account and name:
                        scheme.name = scheme.name.replace(name, f"{name} {account}", 1)
                    out.append(scheme)
            return out
    codes = ctx.keys
    first_name = next((results[r].name_in_file for r in ctx.rows if results[r].name_in_file), codes[0])

    rounding = "sen"
    if by_role["round"]:
        for r in ctx.rows:
            got = _rounding(_text(grid.value(r, by_role["round"].index)))
            if got == "?":
                raise CellError(f'ROUND: cannot read "{_text(grid.value(r, by_role["round"].index))}"')
            if got:
                rounding = got
    payto_text = None
    if by_role["payto"]:
        for r in ctx.rows:
            payto_text = payto_text or _text(grid.value(r, by_role["payto"].index)) or None

    # Only a YEARLY scheme reads a window stated on the row (a PERIOD cell, or the dates written
    # in a yearly cell); quarters and months keep the calendar year (item 5).
    if kind == "Y":
        stated = None
        sources = ([window_col] if window_col else []) + list(kind_cols)
        for col in sources:
            for r in ctx.rows:
                text = _text(grid.value(r, col.index))[:MAX_CELL_TEXT]
                win = _window(text, ctx.year)
                if win:
                    stated = win
                elif text and col is window_col:
                    # A remark that states no period is told, never read as a figure (gate 8).
                    note = f'{col.header}: "{text[:60]}" is not a period; read as a remark'
                    if note not in results[r].warnings:
                        results[r].warnings.append(note)
        header_window = next((c.window for c in kind_cols if c.window), None)
        # Account targets split off a dealer's block take the dealer's window when neither a row
        # cell nor their own headers state one (gate 3).
        start, end = stated or header_window or inherit_window or (date(ctx.year, 1, 1), date(ctx.year, 12, 31))
        split = (None, None)
    else:
        start, end = date(ctx.year, 1, 1), date(ctx.year, 12, 31)
        split = (3, "month") if kind == "Q" else (1, "month")

    tiers: List[SchemeTier] = []
    pay_to: List[str] = []
    # One incentive column for every period of the block ("QUARTERLY (POINT)" ... "INCENTIVE"),
    # unless the block's number-0 columns are tiers of their own (a "Monthly" block with no months).
    shared_rates = [c for c in kind_cols if kind != "Y" and c.number == 0 and c.is_rate]
    if any(c.number == 0 and not c.is_rate for c in kind_cols):
        shared_rates = []
    own = [c for c in kind_cols if c not in shared_rates]
    for number in sorted({c.number or 0 for c in own}):
        cols = [c for c in own if (c.number or 0) == number]
        # Number 0 on a monthly block: the same tiers every month (no period of their own).
        period_start = None if kind != "Y" and number == 0 else _period_start(kind, number, start)
        got, payees = _period_tiers(ctx, cols, period_start, kind, shared_rates)
        tiers.extend(got)
        pay_to.extend(payees)
    if not tiers:
        return []
    # One threshold repeated by tiers that each name other accounts ("A/C I- 3%" in Tier 1,
    # "A/C III- 2%" in Tier 2, both at 300,000) is one tier with a rate per account (gate 3).
    merged_tiers: List[SchemeTier] = []
    # Each tier as written, before the per-account merge below changes any of them in place, so
    # the repeat check does not depend on the order of the columns.
    written: set = set()
    for t in tiers:
        shape = (
            t.period_start, t.threshold, t.is_additional, t.rate, frozenset((t.account_rates or {}).items()), t.bonus
        )
        if shape in written:
            # The same threshold and rate written twice in one period is one tier (gate 5:
            # "A/C I 9% | A/C I 9%"), said on the row so it is never silent.
            note = f"{_KIND_NAME[kind]}: the threshold {t.threshold:,.0f} is written twice with the same rate; read once"
            for r in ctx.rows:
                if note not in results[r].warnings:
                    results[r].warnings.append(note)
            continue
        written.add(shape)
        twin = next(
            (
                m for m in merged_tiers
                if (m.period_start, m.threshold, m.is_additional) == (t.period_start, t.threshold, t.is_additional)
                and m.account_rates and t.account_rates and m.rate is None and t.rate is None
                and not set(m.account_rates) & set(t.account_rates)
            ),
            None,
        )
        if twin is not None:
            twin.account_rates.update(t.account_rates)
        else:
            merged_tiers.append(t)
    tiers = merged_tiers
    seen = set()
    labels = {res.key: res.account_in_file or "" for res in results.values()}

    def rate_text(t: SchemeTier) -> str:
        if t.account_rates:
            return " / ".join(f"{labels.get(k, '')} {v.normalize():f}%".strip() for k, v in t.account_rates.items())
        return f"{(t.rate or Decimal(0)).normalize():f}%"

    for t in tiers:
        if t.is_additional:
            continue
        key = (t.period_start, t.threshold)
        if key in seen:
            # The repeated tiers' rates, quoted, so the masked dry run shows their shape (gate 4).
            twins = [x for x in tiers if not x.is_additional and (x.period_start, x.threshold) == key]
            raise CellError(
                f"{_KIND_NAME[kind]}: the threshold {t.threshold:,.0f} repeats in one period "
                f'(rates "{" | ".join(rate_text(x) for x in twins)}")'
            )
        seen.add(key)
    if len(set(pay_to)) > 1:
        raise CellError(f"{_KIND_NAME[kind]}: the rates name two different accounts to pay")
    if payto_text:
        payout = ctx.dealer_labels.get(_label_key(payto_text), payto_text)
    elif pay_to:
        payout = pay_to[0]
    else:
        payout = next((next(iter(t.account_rates)) for t in tiers if t.account_rates), None) or codes[0]
    by_name = payout if isinstance(payout, ByName) else None
    if by_name is not None:
        payout = None
    span = (
        str(start.year)
        if (start.month, start.day, end.month, end.day) == (1, 1, 12, 31) and start.year == end.year
        else f"{start:%b %Y} to {end:%b %Y}"
    )
    unit = "points" if points else "amount"
    unit_word = " points" if points else ""
    return [
        Scheme(
            key=f"{_KIND_NAME[kind]}{unit_word} {start:%Y-%m} {'+'.join(codes)}",
            kind=_KIND_NAME[kind],
            sheet=ctx.sheet,
            rows=list(ctx.rows),
            codes=list(codes),
            names=[results[r].name_in_file or "" for r in ctx.rows],
            payout_code=payout,
            payout_by_name=by_name,
            start=start,
            end=end,
            split_every=split[0],
            split_unit=split[1],
            threshold_unit=unit,
            rounding=rounding,
            tiers=tiers,
            name=f"{first_name[:80]} {span}{unit_word} {_KIND_NAME[kind]}",
        )
    ]


def _period_start(kind: str, number: int, start: date) -> date:
    if kind == "Q":
        return add_months(start, 3 * (number - 1))
    if kind == "M":
        offset = (number - start.month) % 12
        return add_months(start.replace(day=1), offset)
    return start


def _pairs(
    cols: List[_Column], shared_rates: List[_Column]
) -> Tuple[List[Tuple[_Column, Optional[_Column]]], List[_Column]]:
    """(threshold column, its incentive column or None) per threshold, and the incentive
    columns nothing paired. An incentive pairs with the threshold column right before it (a
    plain "INCENTIVE" after "TIER n", gate 2 item 2), else with the same "Tier n", else with the
    one incentive shared by every period of the block ("QUARTERLY (POINT)" ... "INCENTIVE")."""
    rates = [c for c in cols if c.is_rate]
    used: set = set()
    out: List[Tuple[_Column, Optional[_Column]]] = []
    for col in (c for c in cols if not c.is_rate):
        rate = next(
            (r for r in rates if r.index == col.index + 1 and r.tier in (None, col.tier) and r.index not in used),
            None,
        )
        if rate is None and col.tier is not None:
            rate = next((r for r in rates if r.tier == col.tier and r.index not in used), None)
        if rate is not None:
            used.add(rate.index)
        elif len(shared_rates) == 1:
            rate = shared_rates[0]
        out.append((col, rate))
    return out, [r for r in rates if r.index not in used]


def _period_tiers(ctx: _Ctx, cols, period_start, kind, shared_rates=()) -> Tuple[List[SchemeTier], List[str]]:
    """One period's tiers for one target, and the accounts its rates name to be paid (pay-to)."""
    name = cols[0].header or _KIND_NAME[kind]
    out: List[SchemeTier] = []
    payees: List[str] = []
    in_target = {_label_key(ctx.results[r].account_in_file or ""): ctx.results[r].key for r in ctx.rows}
    in_target.update({_label_key(ctx.results[r].code): ctx.results[r].key for r in ctx.rows if ctx.results[r].code})
    in_target.pop("", None)

    def resolve(label: str) -> Tuple[str, bool]:
        """(row key, counted here) for an account a rate names."""
        key = _label_key(label)
        if key in in_target:
            return in_target[key], True
        if key in ctx.dealer_labels:
            return ctx.dealer_labels[key], False
        if (account_key(label) or "").startswith("L"):
            # An account of this dealer with no row in the file: paid to it, found by the
            # dealer's name and that account in the CRM when the import runs (gate 3).
            dealer = ctx.results[ctx.rows[0]].name_in_file or ""
            return ByName(dealer, split_account_suffix("x " + label)[1] or label), False
        message = f'{name}: the rate names "{label}", which is not an account of this dealer'
        pick_key = f"{ctx.sheet}!{ctx.dealer_row}~{label}"
        if not ctx.dealer_accounts or len(pick_key) > MAX_PICK_KEY:
            raise CellError(message)
        chosen = ctx.label_picks.get(pick_key)
        if chosen is not None and any(_label_key(chosen) == _label_key(a) for a in ctx.dealer_accounts):
            return resolve(chosen)
        raise AccountPick(f"{message}; pick the account it means", pick_key, label, list(ctx.dealer_accounts))

    def shared(col: _Column, *, every_row: bool) -> Tuple[object, Dict[int, object]]:
        """The value the target reads (a merge over the rows, or the same value on each), or the
        per-row values. A threshold typed on one row only is the target's; a RATE typed on some
        rows only is those accounts' rate (ruling 2)."""
        per_row = ctx.per_row(col)
        if not per_row:
            return None, {}
        distinct = {_text(v) for v in per_row.values()}
        if len(distinct) == 1 and (not every_row or len(per_row) == len(ctx.rows)):
            return next(iter(per_row.values())), {}
        return None, per_row

    def threshold_of(value) -> Optional[Decimal]:
        if _is_dash(value) or _is_nil(value):
            return None
        if isinstance(value, str) and "%" in value:
            raise CellError(f'{name}: cannot read the threshold "{_text(value)}" (a rate in a threshold column)')
        n = _number(value)
        if n is None:
            raise CellError(f'{name}: cannot read the threshold "{_text(value)}"')
        return n

    def pct_format(col: _Column, r: int) -> bool:
        return "%" in (ctx.grid.fmt(r, col.index) or "")

    pairs, unpaired = _pairs(cols, list(shared_rates))
    for col in unpaired:
        if ctx.per_row(col):
            raise CellError(f"{col.header}: an incentive with no threshold column before it")
    # Gate 2 item 3: in one period, thresholds per account beside thresholds merged over the
    # accounts (Tier 1-3 per account, Tier 4-5 for the dealer). Held for the owner's ruling.
    per_account = [t for t, _ in pairs if shared(t, every_row=False)[1]]
    merged = [t for t, _ in pairs if shared(t, every_row=False)[0] is not None]
    if per_account and merged:
        raise CellError(
            f"{name}: per-account {per_account[0].header} beside dealer {merged[0].header} is read "
            "only when the dealer tiers are merged over the accounts"
        )
    previous_threshold: Optional[Decimal] = None
    for thr_col, rate_col in pairs:
        if rate_col is None:
            # One cell per tier text ("150,000 - 2%"), or a bare figure paying the header's rate.
            value, per_row = shared(thr_col, every_row=False)
            if per_row:
                raise CellError(f"{name}: the accounts of one target give different tiers")
            if value is None or _is_dash(value):
                continue
            if _is_nil(value):
                continue
            for t in parse_tier_cell(value, default_rate=thr_col.default_rate):
                if t.threshold is None:
                    raise CellError(f"{name}: a rate with no threshold")
                rates = None
                if t.account_rates:
                    counted_rates: Dict[str, Decimal] = {}
                    paid: List[Tuple[str, Decimal]] = []
                    for label, rate in t.account_rates.items():
                        key, counted = resolve(label)
                        if counted:
                            counted_rates[key] = rate
                        else:
                            paid.append((key, rate))
                    if counted_rates and paid:
                        raise CellError(f"{name}: one tier pays both every account and named accounts")
                    if len(paid) > 1:
                        raise CellError(f"{_KIND_NAME[kind]}: the rates name two different accounts to pay")
                    if paid:
                        payees.append(paid[0][0])
                        t.rate = paid[0][1]
                    else:
                        rates = counted_rates
                if rates is None and t.rate is None:
                    raise CellError(f"{name}: the threshold {t.threshold:,.0f} has no rate")
                out.append(SchemeTier(period_start, t.threshold, None if rates else t.rate, t.is_additional, rates))
            continue

        thr_value, thr_rows = shared(thr_col, every_row=False)
        if thr_rows:
            raise CellError(f"{name}: the accounts of one target give different thresholds")
        rate_value, rate_rows = shared(rate_col, every_row=True)
        threshold = threshold_of(thr_value)
        has_rate = rate_value is not None or bool(rate_rows)
        if threshold is None and not has_rate:
            continue
        if threshold is None and not _is_nil(thr_value) and thr_value is not None and not _is_dash(thr_value):
            raise CellError(f'{name}: cannot read the threshold "{_text(thr_value)}"')
        # Rates: one shared cell, or one per row.
        entries: List[Tuple[Optional[int], RateEntry]] = []
        try:
            if rate_rows:
                for r, v in rate_rows.items():
                    entries.extend(
                        (r, e)
                        for e in parse_rate_cell(v, percent_format=pct_format(rate_col, r), threshold=threshold)
                    )
            elif rate_value is not None:
                entries.extend(
                    (None, e)
                    for e in parse_rate_cell(
                        rate_value, percent_format=pct_format(rate_col, ctx.rows[0]), threshold=threshold
                    )
                )
        except NotARate as exc:
            # Owner 20 Oct 10:11 (SRT 15-17): skip this tier, say so on the rows, load the rest.
            note = f'{thr_col.header} skipped: cell reads "{exc.cell}"' + (f" (not a rate: {exc.why})" if exc.why else "")
            for r in ctx.rows:
                if note not in ctx.results[r].warnings:
                    ctx.results[r].warnings.append(note)
            continue
        except CellError as exc:
            if threshold is not None and "threshold" not in str(exc):
                raise CellError(f"{rate_col.header}: {exc} for the threshold {threshold:,.0f}") from exc
            raise CellError(f"{rate_col.header}: {exc}") from exc
        if threshold is None:
            if entries and all(e.is_additional for _, e in entries) and previous_threshold is not None:
                threshold = previous_threshold   # "+ Additional 1%" on the tier above's threshold
            elif _is_nil(thr_value) and not entries:
                continue
            else:
                raise CellError(f"{name}: a rate with no threshold")
        if not entries:
            if thr_col.default_rate is not None:
                entries = [(None, RateEntry(None, thr_col.default_rate))]
            else:
                raise CellError(f"{name}: the threshold {threshold:,.0f} has no incentive")
        main: Dict[str, Decimal] = {}
        main_rate: Optional[Decimal] = None
        bonus: Optional[Decimal] = None
        extras: List[Tuple[Optional[Dict[str, Decimal]], Optional[Decimal]]] = []
        for r, e in entries:
            if e.bonus is not None:
                # Paid once the tier is reached, to the account the cell names.
                if bonus is not None:
                    raise CellError(f"{name}: two bonus amounts for the threshold {threshold:,.0f}")
                payees.append(resolve(e.label)[0])
                bonus = e.bonus
                continue
            additional = e.is_additional
            target_key = None
            rate_for_all = None
            if e.label:
                key, counted = resolve(e.label)
                if counted:
                    target_key = key
                else:
                    payees.append(key)
                    rate_for_all = e.rate
            elif r is not None:
                target_key = ctx.results[r].key
            else:
                rate_for_all = e.rate
            if additional:
                extras.append(({target_key: e.rate} if target_key else None, rate_for_all))
            elif target_key:
                main[target_key] = e.rate
            else:
                main_rate = rate_for_all
        if bonus is not None and not main and main_rate is None:
            main_rate = Decimal(0)
        if main or main_rate is not None:
            if main and main_rate is not None:
                raise CellError(f"{name}: one tier pays both every account and named accounts")
            out.append(SchemeTier(period_start, threshold, main_rate, False, main or None, bonus))
        for rates, rate in extras:
            out.append(SchemeTier(period_start, threshold, rate, True, rates))
        previous_threshold = threshold
    for t in out:
        for rate in [t.rate] + list((t.account_rates or {}).values()):
            if rate is not None and rate > RATE_LIMIT_PCT:
                shown = f"{rate.normalize():f}"
                raise CellError(
                    f"{name}: rate {shown}% looks wrong: Excel {(rate / 100).normalize():f} means "
                    f"{shown}%, type 0.005 or 0.5% for half a percent"
                )
    return out, payees


# --------------------------------------------------------------------------------------
# the database half: which customer each row is
# --------------------------------------------------------------------------------------

#: Words that name a company's legal form, never the business (owner ruling 16 Oct 09:12).
_LEGAL_FORM = {"SDN", "BHD", "BERHAD", "SB", "PLT", "LLP"}
_ROMAN = {"I": 1, "II": 2, "III": 3, "IV": 4, "V": 5}
_MAX_SUGGESTIONS = 3
#: Account labels a company's file writes for ledgers that carry a marker instead of a level, by
#: company code. Owner 17 Oct ("yes for mocha"): in Mocha's file A/C II is the dealer's "[CERAMIC]"
#: ledger and A/C III its "[IBORN]" ledger. Matched within the dealer's own name, never across.
ACCOUNT_LABEL_LEDGERS: Dict[str, Dict[str, str]] = {"MCH": {"A/C II": "CERAMIC", "A/C III": "IBORN"}}
#: Distinct names one upload may send to the near-match (security review B2): each one scans
#: every customer of the company. Past it a row still waits on a pick, with no suggestions.
MAX_FUZZY_LOOKUPS = 300
#: A pick names a customer per row; more than the bounded sheets can hold is not a real form.
MAX_PICKS = MAX_SHEETS * MAX_SHEET_ROWS
#: The longest pick key `_parse_picks` takes.
MAX_PICK_KEY = 200
#: "<sheet>!<dealer row>~<label>": an account pick's key, never a row's ("<sheet>!<row>").
_ACCOUNT_PICK_KEY = re.compile(r"!\d+~", re.DOTALL)


#: Possessive quantifiers and no leading class: linear on any text (security review gate 2 B1:
#: the earlier form backtracked quadratically on a long run of spaces).
_ACCOUNT_SUFFIX = re.compile(r"\bA\s*+/\s*+C\.?\s*+([IVX]++|\d++)\s*+[)\]]?\s*+\Z", re.IGNORECASE)


def _is_account_marker(text: str) -> bool:
    """A bracketed "A/C I" / "A/C 1", never a branch that merely starts with L ("(LABUAN)")."""
    return re.fullmatch(r"L\d+|RAW(?:\d+|[IVX]+)", account_key(text) or "") is not None


#: Parentheses of a company name that are never a branch: "(M)" (Malaysia) and a registration
#: number ("(123456-X)"), the bracket forms of the legal words `dealer_key` already drops.
_NOT_A_BRANCH = re.compile(r"M|.*\d{4,}.*")


def _branches(name: Optional[str], accounts: set) -> List[str]:
    """The branch words in a name's parentheses ("(KL)", "(JB)"), normalised: account markers,
    an account the ledger answers to ("(CERAMIC)"), "(M)" and registration numbers are not one."""
    out = []
    for text in re.findall(r"\(([^)]*)\)", name or ""):
        word = dealer_key(text)
        if not word or _is_account_marker(text) or account_key(text) in accounts or _NOT_A_BRANCH.fullmatch(word):
            continue
        out.append(word)
    return out


def missing_branches(name: Optional[str], ledger: "_Candidate") -> List[str]:
    """Branches the file name writes that the ledger's own name does not (gate 6)."""
    own = set(_branches(ledger.name, ledger.accounts))
    return [b for b in _branches(name, ledger.accounts) if b not in own]


def branch_key(name: Optional[str]) -> str:
    """`dealer_key` that keeps the words of a bracketed part which is not an account marker
    ("(JB)", "(KL)"): the tie-break between ledgers of one name and account."""
    text = re.sub(r"\[[^\]]*\]", " ", (name or "").upper())
    text = re.sub(r"\([^)]*\)", lambda m: " " if _is_account_marker(m.group(0)[1:-1]) else f" {m.group(0)[1:-1]} ", text)
    return dealer_key(text)


def literal_name(name: Optional[str]) -> str:
    """A name as written, letter for letter: case and runs of spaces ignored, punctuation kept,
    the account markers and suffix dropped. The tie-break between ledgers whose names differ only
    in punctuation or spacing (gate 5): the one the file writes the same way is the exact match."""
    text = re.sub(r"\[[^\]]*\]", " ", (name or "").upper())
    text = re.sub(r"\([^)]*\)", lambda m: " " if _is_account_marker(m.group(0)[1:-1]) else m.group(0), text)
    text = split_account_suffix(text)[0]
    return " ".join(text.split()).strip(" -\u2013")


def split_account_suffix(name: str) -> Tuple[str, Optional[str]]:
    """("<dealer>", "A/C <n>") from "<dealer> A/C <n>" (also "- A/C III", "(A/C I)"); the name
    whole and None when it carries no such suffix."""
    m = _ACCOUNT_SUFFIX.search(name or "")
    if not m:
        return name, None
    head = name[: m.start()].rstrip(" -\u2013([")
    if not head.strip():
        return name, None
    return head, f"A/C {m.group(1).upper()}"


def dealer_key(name: Optional[str]) -> str:
    """The normalised company name a file name and a CRM name are compared on: case, spaces and
    punctuation ignored, bracketed ledger markers (`[A/C I]`, `(SRT)`) dropped, the SDN BHD forms
    dropped, ENTERPRISES = ENTERPRISE (owner ruling 16 Oct 09:12), accents folded ("É" = "E",
    crew ruling 20 Oct: the same letter, not a looser match)."""
    text = unicodedata.normalize("NFKD", name or "")
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"\[[^\]]*\]|\([^)]*\)", " ", text.upper())
    text = split_account_suffix(text)[0]
    text = re.sub(r"\bS\s*/\s*B\b", " ", text)
    words = re.sub(r"[^A-Z0-9]+", " ", text).split()
    out = []
    for w in words:
        if w in _LEGAL_FORM:
            continue
        out.append("ENTERPRISE" if w == "ENTERPRISES" else w)
    # A trailing " - " before a marker leaves nothing behind; "-" alone is punctuation.
    return " ".join(out)


def account_key(label: Optional[str]) -> Optional[str]:
    """`A/C I`..`A/C V` (or `ACCOUNT 2`) -> `L1`..; any other label (`[CERAMIC]`, `Deluxe`,
    `W-Ply`) -> its letters and digits upper-cased; blank -> None."""
    text = (label or "").upper().strip()
    if not text:
        return None
    m = re.fullmatch(r"\[?\(?\s*(?:A\s*/\s*C|ACC(?:OUN)?T?\.?)\s*([IVX]+|\d+)\s*\)?\]?", text)
    if m:
        token = m.group(1)
        if token in _ROMAN:
            return f"L{_ROMAN[token]}"
        # "A/C 1" is a typo the office fixes, never read as A/C I (item 7).
        return f"RAW{token}"
    key = re.sub(r"[^A-Z0-9]", "", text)
    return key or None


def name_accounts(name: str) -> set:
    """The account keys a customer NAME itself states (markers and a bare suffix), level aside."""
    return customer_account_keys(name, None)


def customer_account_keys(name: str, level: Optional[int]) -> set:
    """The account labels a CRM customer answers to: its account level (`customers.account_level`,
    set by migration acct_ledger_0001 from the name's marker and edited by the office since) and
    every bracketed marker in its name (`[CERAMIC]`, `[IBORN]`, `[DELUXE]`), which is how the
    Mocha A/C II / III ledgers carry their account (they have no level)."""
    keys = set()
    if level:
        keys.add(f"L{level}")
    for marker in re.findall(r"\[([^\]]*)\]|\(([^)]*)\)", name or ""):
        key = account_key(marker[0] or marker[1])
        if key:
            keys.add(key)
    suffix = split_account_suffix(name or "")[1]
    if suffix:
        keys.add(account_key(suffix))
    return keys


@dataclass
class _Candidate:
    id: str
    code: str
    name: str
    family: str
    accounts: set


class _CustomerIndex:
    """The acting company's live customers, indexed once by dealer key."""

    def __init__(self, db: Session, company_id: str):
        from app.models.company import Company
        from app.models.order import Customer

        company_code = db.query(Company.code).filter(Company.id == company_id).scalar()
        aliases = {
            account_key(label): account_key(marker)
            for label, marker in ACCOUNT_LABEL_LEDGERS.get(company_code or "", {}).items()
        }

        self.by_family: Dict[str, List[_Candidate]] = {}
        self.all: List[_Candidate] = []
        self._cache: Dict[
            Tuple[str, str, Optional[str]], Tuple[Optional[_Candidate], List[_Candidate], bool, Optional[str]]
        ] = {}
        self.fuzzy_lookups = 0
        # The acting company by name, not only through the session's scope listener: a session
        # without the listeners (the dry-run script, gate 6) read every company's customers.
        q = db.query(Customer.id, Customer.customer_code, Customer.customer_name, Customer.account_level).filter(
            Customer.is_active.is_(True), Customer.company_id == company_id
        )
        for cid, code, name, level in q:
            accounts = customer_account_keys(name, level)
            accounts |= {label for label, marker in aliases.items() if marker in accounts}
            cand = _Candidate(str(cid), code, name, dealer_key(name), accounts)
            self.all.append(cand)
            self.by_family.setdefault(cand.family, []).append(cand)
        self.families = list(self.by_family)

    def match(
        self, name: Optional[str], account: Optional[str]
    ) -> Tuple[Optional[_Candidate], List[_Candidate], bool, Optional[str]]:
        """(the one exact match or None, up to 3 suggestions when there is none, whether the
        near-match was skipped for being past `MAX_FUZZY_LOOKUPS`, why there is no match:
        `several` ledgers answer to the name and account, the name is there but `no_account` of
        its ledgers has the account, or `no_name`). Cached per name and account, so a dealer's
        many rows cost one lookup."""
        # By `branch_key`, not `dealer_key`: two branches of one name are two lookups (review B1);
        # and by the spelling, which the letter-for-letter tie-break reads (gate 5 review B1).
        key = (branch_key(name), literal_name(name), account_key(account) or (name or "").upper())
        if key not in self._cache:
            self._cache[key] = self._match(name, account)
        return self._cache[key]

    def _match(self, name: Optional[str], account: Optional[str]):
        import difflib

        family = dealer_key(name)
        # A name carrying its own marker ("Mocha K [CERAMIC]") with no account cell says it there.
        acc = account_key(account)
        if acc is None:
            marked = customer_account_keys(name or "", None)
            acc = next(iter(marked)) if len(marked) == 1 else None
        same = self.by_family.get(family, [])
        if acc is not None:
            exact = [c for c in same if acc in c.accounts]
        else:
            exact = same if len(same) == 1 else [c for c in same if not c.accounts]
        if len(exact) > 1:
            # A bracketed word other than the account (a branch: "(JB)", "(KL)") breaks the tie
            # when the file writes it too (gate 3).
            strict = [c for c in exact if branch_key(c.name) == branch_key(name)]
            if len(strict) == 1:
                exact = strict
        if len(exact) > 1 and acc is not None:
            # In a tie, a ledger whose name states another account ("[CERAMIC]" at level 1) is
            # out before any tie-break can land on it (gate 5 review B2).
            plain = [c for c in exact if not name_accounts(c.name) or acc in name_accounts(c.name)]
            if not plain:
                # Every tied ledger names another account: no tie-break may pick one (re-check).
                return None, exact[:_MAX_SUGGESTIONS], False, "several"
            exact = plain
        if len(exact) > 1 and acc is not None:
            # Several ledgers answer to the account through `account_level`, but only one names
            # it in its own name ("[A/C I]", "A/C I"): that one (gate 4). A level never outranks
            # a name that says another account.
            named = [c for c in exact if name_accounts(c.name) == {acc}]
            if len(named) == 1:
                exact = named
        if len(exact) > 1:
            # Names that differ only in punctuation or spacing: the one the file writes the same
            # way (gate 5). When the file writes neither, the row stays a pick.
            written = [c for c in exact if literal_name(c.name) == literal_name(name)]
            if len(written) == 1:
                exact = written
        if len(exact) == 1:
            if missing_branches(name, exact[0]):
                # The file writes a branch ("<NAME> (KL)") the one ledger of the name lacks: not
                # an exact match, so a pick that suggests that ledger (gate 6).
                return None, exact, False, "no_branch"
            return exact[0], [], False, None
        if len(exact) > 1:
            return None, exact[:_MAX_SUGGESTIONS], False, "several"
        # The name is there, but none of its ledgers has the file's account (Mocha A/C II rows
        # beside "[A/C I]", "[CERAMIC]", "[IBORN]", gate 5): its ledgers are the suggestions.
        why = "no_name"
        if same:
            # Ledgers of the name, none answering: the file's account is not theirs, or the file
            # names no account and they all carry one (gate 5 review S1).
            why = "no_account" if acc is not None else "account_needed"
        if why != "no_name":
            ranked = sorted(same, key=lambda c: c.name)
            return None, ranked[:_MAX_SUGGESTIONS], False, why
        if not family:
            return None, [], False, why
        if self.fuzzy_lookups >= MAX_FUZZY_LOOKUPS:
            return None, [], True, why
        self.fuzzy_lookups += 1
        near = difflib.get_close_matches(family, self.families, n=10, cutoff=0.6)
        ranked = []
        for fam in near:
            score = difflib.SequenceMatcher(None, family, fam).ratio()
            for c in self.by_family[fam]:
                ranked.append((-(score + (0.1 if acc and acc in c.accounts else 0)), c.name, c))
        ranked.sort(key=lambda x: (x[0], x[1]))
        return None, [c for _, _, c in ranked[:_MAX_SUGGESTIONS]], False, why


def _same_name(a: Optional[str], b: Optional[str]) -> bool:
    return dealer_key(a) == dealer_key(b)


def _parse_picks(picks: Optional[dict]) -> Dict[str, str]:
    if picks is None:
        return {}
    if not isinstance(picks, dict) or len(picks) > MAX_PICKS or not all(
        isinstance(k, str) and isinstance(v, str) and len(k) <= MAX_PICK_KEY and len(v) <= 64 for k, v in picks.items()
    ):
        raise AppException(
            status_code=422,
            message='Picks are an object of "sheet!row" to a customer or "skip".',
            code="INVALID_PICKS",
        )
    return picks


def _acting_company(db: Session) -> str:
    """The one company the import reads customers of and writes targets for. A session scoped to
    every company (or to none) would match one company's file against another's ledgers."""
    from app.models.base import get_company_scope

    scope = get_company_scope(db)
    if not isinstance(scope, frozenset) or len(scope) != 1:
        raise AppException(
            status_code=422,
            message="Import dealer targets for one company at a time: choose the company first.",
            code="IMPORT_NEEDS_ONE_COMPANY",
        )
    return next(iter(scope))


def run(
    db: Session,
    data: bytes,
    *,
    year: int,
    dry_run: bool,
    user_id: Optional[str] = None,
    picks: Optional[dict] = None,
) -> dict:
    import uuid

    from app.models.order import Customer
    from app.models.sales import SalesTarget, SalesTargetAccount
    from app.schemas.sales import SalesTargetCreate, SalesTargetUpdate
    from app.services.sales import target_service

    company_id = _acting_company(db)
    picks = _parse_picks(picks)
    label_picks = {k: v for k, v in picks.items() if _ACCOUNT_PICK_KEY.search(k)}
    picks = {k: v for k, v in picks.items() if k not in label_picks}
    parsed = parse_workbook(data, year=year, label_picks=label_picks)
    rows_by_key = {(r.sheet, r.row): r for r in parsed.rows}

    # Codes the file names (a CUSTOMER CODE column, a PAY TO code), matched without regard to
    # case or surrounding spaces. A code wins over the name when the row has one.
    codes = {r.code for r in parsed.rows if r.code} | {
        s.payout_code for s in parsed.schemes
        if s.payout_code and "!" not in s.payout_code
    }
    wanted = sorted({code.strip().upper() for code in codes})
    by_code = {}
    if wanted:
        for c in db.query(Customer).filter(
            func.upper(func.trim(Customer.customer_code)).in_(wanted), Customer.company_id == company_id
        ):
            by_code[c.customer_code.strip().upper()] = c
    picked_ids = sorted({v for v in picks.values() if v != "skip"})
    valid_ids = []
    for v in picked_ids:
        try:
            valid_ids.append(str(uuid.UUID(v)))
        except ValueError:
            continue
    by_id = (
        {
            c.id: c
            for c in db.query(Customer).filter(
                Customer.id.in_(valid_ids), Customer.is_active.is_(True), Customer.company_id == company_id
            )
        }
        if valid_ids
        else {}
    )
    by_name_payouts = [s.payout_by_name for s in parsed.schemes if s.payout_by_name is not None]
    index = (
        _CustomerIndex(db, company_id)
        if by_name_payouts or any(r.status == "mapped" and not r.code for r in parsed.rows)
        else None
    )

    #: row key -> the CRM customer (code, picked or name-matched)
    customer_by_key: Dict[str, Customer] = {}
    candidates: Dict[str, _Candidate] = {}
    for res in parsed.rows:
        res.pick_key = f"{res.sheet}!{res.row}"
        pick = picks.get(res.pick_key)
        if pick == "skip":
            # Any row may be skipped, a row the reader could not read included: the user has
            # seen it and chosen to leave it out.
            res.status, res.reason = "skipped", "skipped in the preview"
            continue
        if res.status != "mapped":
            continue
        if pick is not None:
            chosen = by_id.get(pick)
            if chosen is None:
                res.status, res.reason = "not_representable", "the picked customer is not a customer of this company"
                continue
            customer_by_key[res.key] = chosen
            res.match, res.customer = "picked", chosen.customer_name
            continue
        if res.code:
            chosen = by_code.get(res.code.strip().upper())
            if chosen is None:
                res.status, res.reason = "not_representable", f"no customer with code {res.code}"
                continue
            customer_by_key[res.key] = chosen
            res.match, res.customer = "code", chosen.customer_name
            if res.name_in_file and not _same_name(res.name_in_file, chosen.customer_name):
                res.warnings.append(
                    f'name differs: the file says "{res.name_in_file}", the CRM "{chosen.customer_name}"'
                )
            continue
        exact, suggestions, capped, why = index.match(res.name_in_file, res.account_in_file)
        if exact is None:
            res.status = "needs_pick"
            if why == "several":
                res.reason = "more than one customer has this name and account"
            elif why == "no_branch":
                branch = (missing_branches(res.name_in_file, suggestions[0]) or [""])[0]
                res.reason = f'no ledger of this name has the branch "{branch}"; pick one of its ledgers or skip'
            elif why == "account_needed":
                res.reason = "this name has several ledgers and the file names no account; add an ACCOUNT column or pick one"
            elif why == "no_account":
                markers = re.findall(r"\[([^\]]*)\]|\(([^)]*)\)", res.name_in_file or "")
                shown = (
                    res.account_in_file
                    or split_account_suffix(res.name_in_file or "")[1]
                    or next(((a or b).strip() for a, b in markers if (a or b).strip()), "")
                )
                res.reason = f'no ledger of this name has the account "{shown}"; pick one of its ledgers or skip'
            else:
                res.reason = "no customer has this name and account"
            if capped:
                res.reason += "; too many unmatched names for suggestions, add a CUSTOMER CODE column"
            if (account_key(res.account_in_file) or "").startswith("RAW"):
                res.reason = f'the account "{res.account_in_file}" is not one of A/C I to A/C V; ' + res.reason
            res.suggestions = [{"customer_id": s.id, "code": s.code, "name": s.name} for s in suggestions]
            continue
        candidates[res.key] = exact
        res.match, res.customer = "auto", exact.name
    if candidates:
        for c in db.query(Customer).filter(Customer.id.in_([c.id for c in candidates.values()])):
            for key, cand in candidates.items():
                if cand.id == c.id:
                    customer_by_key[key] = c

    # Accounts paid by dealer name + account (no row of their own): one exact CRM match each.
    by_name_customer: Dict[ByName, Customer] = {}
    for payee in set(by_name_payouts):
        exact, _suggestions, _capped, _why = index.match(payee.dealer, payee.label)
        if exact is not None:
            found = db.query(Customer).filter(Customer.id == exact.id).first()
            if found is not None:
                by_name_customer[payee] = found

    def customer_of(key: Optional[str]):
        if not key:
            return None
        if key in customer_by_key:
            return customer_by_key[key]
        return by_code.get(key.strip().upper())

    # Existing dealer targets by (accounts, dates, split, unit): a re-import updates, never adds.
    existing: Dict[Tuple, SalesTarget] = {}
    dealer_targets = (
        db.query(SalesTarget)
        .filter(SalesTarget.subject_kind == "dealer", SalesTarget.company_id == company_id)
        .all()
    )
    if dealer_targets:
        links: Dict[str, set] = {}
        for link in db.query(SalesTargetAccount).filter(
            SalesTargetAccount.target_id.in_([t.id for t in dealer_targets]),
            SalesTargetAccount.company_id == company_id,
        ):
            links.setdefault(link.target_id, set()).add(link.customer_id)
        for t in dealer_targets:
            existing[
                (frozenset(links.get(t.id, ())), t.start_date, t.end_date, t.split_every, t.split_unit, t.threshold_unit)
            ] = t

    new = updated = 0
    target_nos: Dict[str, str] = {}
    key_to_row = {r.key: r for r in parsed.rows if r.key}

    # Schemes that share a row, or whose pay-to account is another scheme's row, are one
    # dealer: they are created together or not at all (review S4: no partial dealer).
    parent = list(range(len(parsed.schemes)))

    def find(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    owner: Dict[Tuple[str, int], int] = {}
    for i, scheme in enumerate(parsed.schemes):
        touched = [(scheme.sheet, r) for r in scheme.rows]
        payee = key_to_row.get(scheme.payout_code)
        if payee is not None:
            touched.append((payee.sheet, payee.row))
        for place in touched:
            if place in owner:
                parent[find(i)] = find(owner[place])
            else:
                owner[place] = i
    components: Dict[int, List[Scheme]] = {}
    for i, scheme in enumerate(parsed.schemes):
        components.setdefault(find(i), []).append(scheme)

    class _Refused(Exception):
        def __init__(self, scheme: Scheme, message: str):
            super().__init__(message)
            self.scheme, self.message = scheme, message

    #: match key -> the first row of the scheme this run already wrote it from.
    written_by: Dict[Tuple, int] = {}
    for _root, schemes in sorted(components.items()):
        all_rows = [rows_by_key[(s.sheet, r)] for s in schemes for r in s.rows]
        all_rows = list({id(r): r for r in all_rows}.values())

        def fail(reason: str) -> None:
            for r in all_rows:
                if r.status == "mapped":
                    r.status, r.reason = "not_representable", reason

        broken = [r for r in all_rows if r.status != "mapped"]
        for s in schemes:
            payee_row = key_to_row.get(s.payout_code)
            if payee_row is not None and payee_row.status != "mapped":
                broken.append(payee_row)
        if broken:
            first = broken[0]
            why = {
                "needs_pick": "which needs a customer picked",
                "skipped": "which is skipped",
            }.get(first.status)
            fail(f"shares a target with row {first.row}, {why}" if why else
                 f"shares a target with row {first.row}: {first.reason}")
            continue
        prepared = []
        problem = None
        for scheme in schemes:
            ids = [customer_of(code).id for code in scheme.codes]
            # account_rates per tier, by CRM customer: rows of one target on one ledger are one
            # account (owner 20 Oct, MCH 82-84: "it is all for A/C I").
            tier_rates: List[Optional[Dict[str, Decimal]]] = []
            if len(set(ids)) < len(ids):
                # Several rows of one target name the same ledger (three "A/C I" rows, gate 6).
                # By position: rows sharing a CUSTOMER CODE share one row key (review B1).
                twice = next(i for i in ids if ids.count(i) > 1)
                # Every row of the dealer on that ledger, across all its targets (gate 7: a
                # quarterly target on two rows beside a yearly one on three).
                numbers = sorted({
                    r for other in schemes for r, code in zip(other.rows, other.codes)
                    if customer_of(code) is not None and customer_of(code).id == twice
                })
                listed = ", ".join(str(n) for n in numbers[:-1]) + f" and {numbers[-1]}"
            for t in scheme.tiers:
                if not t.account_rates:
                    tier_rates.append(None)
                    continue
                by_customer: Dict[str, Decimal] = {}
                for code, rate in t.account_rates.items():
                    cid = customer_of(code).id
                    if cid in by_customer and by_customer[cid] != rate:
                        shown = " and ".join(f"{v.normalize():f}%" for v in sorted({by_customer[cid], rate}))
                        problem = f"rows {listed} are one CRM customer but give it different rates ({shown})"
                        break
                    by_customer[cid] = rate
                if problem:
                    break
                tier_rates.append(by_customer)
            if problem:
                break
            if len(set(ids)) < len(ids):
                note = f"rows {listed} are one CRM customer: read as one account"
                for r in all_rows:
                    if r.row in numbers and note not in r.warnings:
                        r.warnings.append(note)
                ids = list(dict.fromkeys(ids))
            if scheme.payout_by_name is not None:
                payout = by_name_customer.get(scheme.payout_by_name)
                if payout is None:
                    problem = (
                        f"pays to {scheme.payout_by_name.label}, which has no row in the file and no "
                        "single CRM customer of this dealer with that account"
                    )
                    break
            else:
                payout = customer_of(scheme.payout_code)
            if payout is None:
                problem = f"PAY TO: no customer with code {scheme.payout_code}"
                break
            match_key = (
                frozenset(ids), scheme.start, scheme.end, scheme.split_every, scheme.split_unit, scheme.threshold_unit
            )
            same_group = next((p[0].rows[0] for p in prepared if p[3] == match_key), None)
            if same_group is not None:
                problem = f"the same accounts and dates as row {same_group}"
                break
            if match_key in written_by:
                # Two rows of one file for the same accounts and dates: the second would
                # silently overwrite the first (review nit).
                problem = f"the same accounts and dates as row {written_by[match_key]}"
                break
            tiers = [
                {
                    "period_start": t.period_start if scheme.split_every else None,
                    "threshold": float(t.threshold),
                    "rate": float(t.rate or 0),
                    "is_additional": t.is_additional,
                    "account_rates": (
                        [{"customer_id": cid, "rate": float(rate)} for cid, rate in rates.items()]
                        if rates
                        else None
                    ),
                    "bonus_amount": float(t.bonus) if t.bonus is not None else None,
                }
                for t, rates in zip(scheme.tiers, tier_rates)
            ]
            fields = {
                "commission_method": "retroactive",
                "threshold_unit": scheme.threshold_unit,
                "rounding": scheme.rounding,
                # Tiers with no period of their own (a "Monthly" block with no months) are every
                # period's tiers.
                "tiers_per_period": bool(scheme.split_every) and any(t.period_start for t in scheme.tiers),
                "tiers": tiers,
            }
            prepared.append((scheme, ids, payout, match_key, fields))
        if problem:
            fail(problem)
            continue
        done = []
        try:
            with db.begin_nested():
                # A dealer target first, so the account targets it replaces can name it.
                written: Dict[str, str] = {}
                for scheme, ids, payout, match_key, fields in sorted(prepared, key=lambda p: p[0].replaced_by is not None):
                    match = existing.get(match_key)
                    if scheme.replaced_by and scheme.replaced_by not in written:
                        raise _Refused(scheme, "its dealer target was not written")
                    fields = {**fields, "replaced_by_target_id": written.get(scheme.replaced_by)}
                    try:
                        if match is not None:
                            target_service.update_target(
                                db,
                                match,
                                SalesTargetUpdate(customer_ids=ids, payout_customer_id=payout.id, **fields),
                                user_id=user_id,
                            )
                            target = match
                        else:
                            target = target_service.create_target(
                                db,
                                SalesTargetCreate(
                                    subject_kind="dealer",
                                    customer_ids=ids,
                                    payout_customer_id=payout.id,
                                    name=scheme.name[:120],
                                    metric="amount",
                                    basis="ordered",
                                    product_scope="all",
                                    start_date=scheme.start,
                                    end_date=scheme.end,
                                    split_every=scheme.split_every,
                                    split_unit=scheme.split_unit,
                                    **fields,
                                ),
                                user_id=user_id,
                            )
                    except AppException as exc:
                        raise _Refused(scheme, exc.message) from exc
                    except ValueError as exc:  # a payload the schema refuses (pydantic)
                        errors = getattr(exc, "errors", None)
                        message = errors()[0].get("msg", "not valid") if callable(errors) and errors() else str(exc)
                        raise _Refused(scheme, message) from exc
                    done.append((scheme, match, match_key, target))
                    written[scheme.key] = target.id
        except _Refused as exc:
            fail(f"{exc.scheme.kind}: {exc.message}")
            continue
        for scheme, match, match_key, target in done:
            if match is not None:
                updated += 1
            else:
                new += 1
                existing[match_key] = target
            written_by[match_key] = scheme.rows[0]
            target_nos[scheme.key] = target.target_no

    out_rows = []
    for res in parsed.rows:
        out_rows.append(
            {
                "sheet": res.sheet,
                "row": res.row,
                "pick_key": res.pick_key,
                "code": res.code,
                "name_in_file": res.name_in_file,
                "account_in_file": res.account_in_file,
                "customer": res.customer if res.status in ("mapped", "not_representable") else None,
                "match": res.match,
                "suggestions": res.suggestions,
                "account_pick": res.account_pick if res.status == "needs_pick" else None,
                "status": res.status,
                "reason": res.reason,
                "warnings": res.warnings,
                "target_keys": res.target_keys if res.status == "mapped" else [],
                "target_nos": (
                    []
                    if dry_run or res.status != "mapped"
                    else [target_nos[k] for k in res.target_keys if k in target_nos]
                ),
            }
        )

    def count(status: str) -> int:
        return sum(1 for r in out_rows if r["status"] == status)

    return {
        "dry_run": dry_run,
        "summary": {
            "rows": len(out_rows),
            "mapped": count("mapped"),
            "auto_mapped": sum(1 for r in out_rows if r["status"] == "mapped" and r["match"] == "auto"),
            "needs_pick": count("needs_pick"),
            "skipped": count("skipped"),
            "not_representable": count("not_representable"),
            "targets_new": new,
            "targets_updated": updated,
        },
        "rows": out_rows,
        "notes": parsed.notes,
    }


# --------------------------------------------------------------------------------------
# the template: the real layout plus one CUSTOMER CODE column (gate-fail round item 2)
# --------------------------------------------------------------------------------------

TEMPLATE_TIERS = 3


def template_workbook(year: int) -> bytes:
    """An empty workbook in the layout the dealer target files are delivered in, with a
    CUSTOMER CODE column beside the dealer name. A code there wins over the name match."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    yy = f"{year % 100:02d}"
    wb = Workbook()
    ws = wb.active
    ws.title = f"DEALER TARGETS {year}"
    bold = Font(bold=True)
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def head(row, col, text, across=1, down=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font, cell.alignment = bold, centre
        if across > 1 or down > 1:
            ws.merge_cells(start_row=row, start_column=col, end_row=row + down - 1, end_column=col + across - 1)

    def tiers(row, col, count=TEMPLATE_TIERS):
        for n in range(1, count + 1):
            head(row, col, f"Tier {n}")
            head(row, col + 1, f"Tier {n} Incentive")
            col += 2
        return col

    head(1, 1, "NO", down=3)
    head(1, 2, "DEALER NAME", down=3)
    head(1, 3, "CUSTOMER CODE", down=3)
    head(1, 4, "ACCOUNT", down=3)
    col = 5
    for q, months in zip(("Q1", "Q2", "Q3", "Q4"), ("JAN-MAR", "APR-JUN", "JUL-SEP", "OCT-DEC")):
        head(1, col, q, across=2 * TEMPLATE_TIERS)
        head(2, col, months, across=2 * TEMPLATE_TIERS)
        col = tiers(3, col)
    head(1, col, "YEARLY", across=1 + 2 * TEMPLATE_TIERS)
    head(2, col, "PERIOD", down=2)
    head(2, col + 1, f"JAN-DEC'{yy}", across=2 * TEMPLATE_TIERS)
    col = tiers(3, col + 1)
    head(1, col, "Monthly", across=24)
    for m in ("JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"):
        head(2, col, m, across=2)
        col = tiers(3, col, count=1)
    head(1, col, "PAY TO", down=3)
    head(1, col + 1, "ROUND", down=3)
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 16
    ws.freeze_panes = "E4"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
