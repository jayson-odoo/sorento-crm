"""AutoCount pull lifecycle (PLAN-autocount-pull-review.md, SR1: start / current / status).

A pull is one `import_jobs` row (`job_type` `autocount_products_pull` /
`autocount_stock_pull`). It stays `pending` while FoundryX builds the snapshot - there is
no RQ job for that half, so the orphan sweep (which only ever looks at `queued` and
`started`) never touches it. When FoundryX reports the snapshot ready, the SAME row flips
`pending` -> `queued` and the preview task (`app.tasks.autocount_pull_tasks`) takes it the
rest of the way to `finished`. Confirm (SR3) creates a SECOND row for the apply.

Everything the FE needs about a pull's state - its phase, progress, header facts, counts,
compare summary, confirm gate, apply job id - lives in ONE JSONB column
(`import_jobs.metadata["autocount_pull"]`), never a new table: the plan's "no staging
table, two job rows per pull" (P1/P2).

`enqueue_job` is imported under this exact name (not `app.services.queue_service.enqueue_job`
qualified inline) because the SR1 tests patch `app.services.autocount_pull_service.enqueue_job`
directly - see the test file's module docstring.
"""
from __future__ import annotations

import math
import uuid
from datetime import datetime, timedelta
from typing import Optional
from zoneinfo import ZoneInfo

from sqlalchemy import text, update
from sqlalchemy.orm import Session

from app.models.job import ImportJob, JobStatus
from app.services.foundryx_autocount_client import FoundryxAutocountClient, FoundryxPullError
from app.services.queue_service import enqueue_job

#: Entity name (as the FE/route spells it) -> the permission slug that gates it.
ENTITY_PERMISSIONS = {
    "products": "master_data.products.autocount_pull",
    "stock_balances": "inventory.stock.autocount_pull",
}

#: Entity name -> the `import_jobs.job_type` a pull of it is stored under.
JOB_TYPES = {
    "products": "autocount_products_pull",
    "stock_balances": "autocount_stock_pull",
}

#: D3 (small-fix track, browser e2e run 3): "date of the apply" is the LOCAL calendar day,
#: same convention every other module in this file's neighbourhood uses for a user-facing
#: date rather than the server's UTC one (`marketing_service.py`, `pdf_render.py`, etc. each
#: define their own `ZoneInfo("Asia/Kuala_Lumpur")` - there is no single shared helper).
_MY_TZ = ZoneInfo("Asia/Kuala_Lumpur")

#: Entity name -> the `import_jobs.job_type` Confirm's SECOND row is stored under (SR3).
APPLY_JOB_TYPES = {
    "products": "autocount_products_apply",
    "stock_balances": "autocount_stock_apply",
}

#: A pull is visible on the review page (rows / download / compare) once the preview has
#: run, and stays visible after Confirm - the checker keeps reading what they confirmed.
_ROWS_VISIBLE_PHASES = ("review", "confirmed")

#: Phases that make a pull "open" - reused by a second start, and what `current` looks for.
_OPEN_PHASES = ("building", "previewing", "review")

#: A pull still `building` past this age is abandoned server-side (AC-BD-4): FoundryX
#: itself never guarantees a snapshot build finishes, and nothing else times this out.
BUILD_EXPIRY = timedelta(minutes=60)

#: Header keys never persisted onto the job row. Both can be large (a stock snapshot's
#: `negativePairList` runs into the thousands) and both are FoundryX-owned lists the
#: preview/apply tasks read straight off a fresh snapshot read, not off this cache.
_HEADER_KEYS_NOT_STORED = ("excludedRows", "negativePairList")


class UnknownEntity(ValueError):
    """Raised for an `entity` this lane does not know - the route turns it into a 404."""


def _job_type_for(entity: str) -> str:
    try:
        return JOB_TYPES[entity]
    except KeyError:
        raise UnknownEntity(entity)


def _pull_meta(job: ImportJob) -> dict:
    meta = job.job_metadata or {}
    return dict(meta.get("autocount_pull") or {})


def _with_pull(job: ImportJob, pull: dict) -> dict:
    meta = dict(job.job_metadata or {})
    meta["autocount_pull"] = pull
    return meta


def _is_build_expired(job: ImportJob) -> bool:
    if not job.created_at:
        return False
    return datetime.utcnow() - job.created_at > BUILD_EXPIRY


def entity_of(job: ImportJob) -> str:
    return _pull_meta(job).get("entity", "")


def phase_of(job: ImportJob) -> str:
    return _pull_meta(job).get("phase", "")


class PullNotReadyForConfirm(RuntimeError):
    """The pull is not `review` and carries no already-confirmed apply job."""


class PullNotDiscardable(RuntimeError):
    """The pull is not open (`building`/`previewing`/`review`) and not already
    `discarded` - a `confirmed`, `failed` or `expired` pull cannot be thrown away."""


class PullRowsNotAvailable(RuntimeError):
    """The pull has not reached `review` yet - there is nothing to show."""


def _company_code(db: Session, company_id: str) -> str:
    from app.models.company import Company

    company = db.query(Company).filter(Company.id == company_id).first()
    return company.code if company is not None else ""


def find_open_pull(
    db: Session, *, user_id: str, company_id: Optional[str], entity: str
) -> Optional[ImportJob]:
    """The caller's own open pull for this company + entity, else None (AC-PL-4)."""
    if not company_id:
        return None
    job_type = _job_type_for(entity)
    candidates = (
        db.query(ImportJob)
        .filter(
            ImportJob.job_type == job_type,
            ImportJob.user_id == user_id,
            ImportJob.company_id == company_id,
        )
        .order_by(ImportJob.created_at.desc())
        .all()
    )
    for job in candidates:
        if job.status in (JobStatus.FAILED.value, JobStatus.CANCELLED.value):
            # The orphan sweep (or a dead-worker task) marked the `import_jobs` row
            # itself failed while the pull's own stored phase is stale (F-11) - not
            # "open" whatever the metadata still says. `cancelled` is the same story
            # for Discard (AC-DS-8b): the generic Cancel button on the job page can
            # set `status = cancelled` on a pull whose stored phase still says
            # `building`, and Discard's own `discard_pull` always sets both together
            # - so a `cancelled` row is never "open" regardless of its stored phase.
            continue
        pull = _pull_meta(job)
        if pull.get("phase") not in _OPEN_PHASES:
            continue
        if pull.get("phase") == "building" and _is_build_expired(job):
            continue
        return job
    return None


def get_owned_pull(db: Session, *, job_id: str, user_id: str) -> Optional[ImportJob]:
    """The job, owner-checked - never revealed to anyone but its owner (AC-BD-7, P12)."""
    try:
        job = db.query(ImportJob).filter(ImportJob.id == job_id).first()
    except Exception:
        # A malformed id (not a UUID) is "not found", not a server error - the DB
        # column is a real UUID, so the driver raises rather than returning no rows.
        db.rollback()
        return None
    if job is None or job.user_id != user_id:
        return None
    return job


def start_pull(db: Session, *, user_id: str, company_id: str, entity: str) -> ImportJob:
    """Reuses an open pull; otherwise asks FoundryX to build a snapshot and creates the
    pending row (AC-PL-3, AC-PL-4). Raises `FoundryxPullError` on any FoundryX refusal -
    the caller is responsible for leaving no job row behind on that path, which is true
    here because the row is only ever constructed AFTER `client.build` returns.
    """
    existing = find_open_pull(db, user_id=user_id, company_id=company_id, entity=entity)
    if existing is not None:
        return existing

    company_code = _company_code(db, company_id)
    client = FoundryxAutocountClient(db)  # raises FoundryxPullError(NOT_CONFIGURED) if unset
    body = client.build(company_code, entity)

    pull = {
        "entity": entity,
        "company_code": company_code,
        "snapshot_id": body.get("snapshotId"),
        "phase": "building",
        "progress": None,
        "header": None,
        "counts": {},
        "confirm_blocked_reason": None,
        "compare": None,
        "apply_job_id": None,
        # Set by the preview/apply task from the A5 contentHash recheck (currently
        # the only warning code); absent or empty means a clean match.
        "warnings": [],
    }
    job = ImportJob(
        job_id=str(uuid.uuid4()),
        job_type=_job_type_for(entity),
        status=JobStatus.PENDING.value,
        user_id=user_id,
        company_id=company_id,
        job_metadata={"autocount_pull": pull},
    )
    db.add(job)
    db.commit()
    db.refresh(job)
    return job


def _claim_and_enqueue_preview(db: Session, job: ImportJob, pull: dict, ready_header: dict) -> None:
    """Flips a still-`pending` pull to `queued` and enqueues the preview task EXACTLY
    ONCE (AC-BD-2). The `status == PENDING` clause in the UPDATE is the conditional
    update: a second caller racing this one finds zero rows to update and enqueues
    nothing, whether the race is two threads or two sequential polls."""
    stored_header = {k: v for k, v in ready_header.items() if k not in _HEADER_KEYS_NOT_STORED}
    pull["header"] = stored_header
    pull["phase"] = "previewing"
    pull["progress"] = None
    new_meta = _with_pull(job, pull)

    result = db.execute(
        update(ImportJob)
        .where(ImportJob.id == job.id, ImportJob.status == JobStatus.PENDING.value)
        .values(status=JobStatus.QUEUED.value, job_metadata=new_meta, updated_at=datetime.utcnow())
        .execution_options(synchronize_session=False)
    )
    db.commit()
    if result.rowcount != 1:
        return

    from app.tasks.autocount_pull_tasks import preview_autocount_pull

    enqueue_job(
        preview_autocount_pull,
        str(job.id),
        queue_name="imports",
        job_timeout=3600,
        job_id=str(job.job_id),
    )


def _mark_failed(db: Session, job: ImportJob, pull: dict, *, phase: str, error: str) -> None:
    pull["phase"] = phase
    job.job_metadata = _with_pull(job, pull)
    job.status = JobStatus.FAILED.value
    job.error = error
    job.completed_at = datetime.utcnow()
    job.updated_at = datetime.utcnow()
    db.commit()


def refresh_pull_status(db: Session, job: ImportJob) -> dict:
    """AC-BD-1..4: while `building`, reads FoundryX and advances the row; every other
    phase is a plain read of what is already stored (no outbound call)."""
    pull = _pull_meta(job)

    if pull.get("phase") == "building":
        if _is_build_expired(job):
            _mark_failed(db, job, pull, phase="expired", error="pull again")
            db.refresh(job)
            return serialize(job, db)

        client = FoundryxAutocountClient(db)
        body = client.status(pull.get("snapshot_id"))
        foundryx_status = body.get("status")

        if foundryx_status == "ready":
            _claim_and_enqueue_preview(db, job, pull, body)
            db.refresh(job)
            return serialize(job, db)

        if foundryx_status == "failed":
            error = body.get("error") or {}
            code = error.get("code", "BUILD_FAILED")
            message = error.get("message", "")
            _mark_failed(db, job, pull, phase="failed", error=f"{code}: {message}".rstrip(": "))
            db.refresh(job)
            return serialize(job, db)

        # Still building - store progress when FoundryX sent one.
        if body.get("progress") is not None:
            pull["progress"] = body["progress"]
            job.job_metadata = _with_pull(job, pull)
            job.updated_at = datetime.utcnow()
            db.commit()
            db.refresh(job)

    return serialize(job, db)


def serialize(job: ImportJob, db: Session) -> dict:
    """`db` reads the apply job's own `status` (fix-round follow-up, F-13): the
    `stock_list_not_archived` warning the apply task can append is written onto the
    PULL job's own metadata (`_append_pull_warning`) only AFTER the apply task starts,
    and the FE stops polling the instant `phase` reaches `confirmed` - `apply_status`
    is what lets it know whether the apply job is still running (`queued`/`started`) or
    has reached a terminal state (`finished`/`failed`), so it can keep polling until
    that warning (or anything else the apply job settles) has had a chance to land.
    `None` when there is no apply job yet."""
    pull = _pull_meta(job)
    phase = pull.get("phase")
    if job.status == JobStatus.FAILED.value and phase not in ("failed", "expired"):
        # A worker death (or the orphan sweep) can leave `import_jobs.status` `failed`
        # while the pull's own stored phase is still whatever it was mid-task (F-11) - the
        # job row's own status is the honest answer here, never a stale "previewing" the
        # task never got the chance to overwrite. `failed`/`expired` are already the
        # code's own deliberate answer (`_mark_failed` sets both together) and stay as-is.
        phase = "failed"
    apply_job_id = pull.get("apply_job_id")
    apply_status = None
    if apply_job_id:
        apply_status = db.execute(
            text("SELECT status FROM import_jobs WHERE id = :id"), {"id": apply_job_id}
        ).scalar()
    # B3 (small-fix track): `import_jobs.processed_rows`/`total_rows` - published by the
    # preview task through its own fresh session (see `_publish_preview_progress` in
    # `autocount_pull_tasks.py`; the preview's own session holds the whole dry-run ingest
    # in one open transaction, so it cannot publish through itself). `null` before the
    # task has set a total yet, and outside `previewing` - nothing left to show once the
    # phase has moved on.
    preview_progress = (
        {"processed": job.processed_rows, "total": job.total_rows}
        if phase == "previewing" and job.total_rows
        else None
    )
    return {
        "job_id": str(job.id),
        "entity": pull.get("entity"),
        "company_code": pull.get("company_code"),
        "phase": phase,
        "progress": pull.get("progress"),
        "preview_progress": preview_progress,
        "header": pull.get("header"),
        "counts": pull.get("counts") or {},
        "confirm_blocked_reason": pull.get("confirm_blocked_reason"),
        "compare": pull.get("compare"),
        "apply_job_id": apply_job_id,
        "apply_status": apply_status,
        "warnings": pull.get("warnings") or [],
        "error": job.error,
    }


# ==================================================================== SR3 - review page


def require_rows_available(job: ImportJob) -> None:
    """AC-RV-3/5: the review data (rows, download, compare) exists only once the preview
    has finished - 409 while building/previewing/failed/expired."""
    if phase_of(job) not in _ROWS_VISIBLE_PHASES:
        raise PullRowsNotAvailable(f"Pull is in phase {phase_of(job)!r}; nothing to show yet.")


def fetch_snapshot_rows(db: Session, job: ImportJob) -> list[dict]:
    """The RAW snapshot rows, fetched fresh through the FoundryX client every call - no
    cache (plan: "Simplest thing that works; no cache"). Shared by `/rows`, `/download.xlsx`
    and `/compare` so none of the three re-implements paging on its own."""
    pull = _pull_meta(job)
    client = FoundryxAutocountClient(db)
    return client.all_rows(pull.get("snapshot_id"))


def map_product_row(row: dict) -> dict:
    """AC-RV-3: the pull's Excel-view shape for one raw snapshot row.

    Captain ruling (SR3 fix round): Desc 2 must ROUND-TRIP through the manual join
    formula (`f"{description} {desc2}".strip()`, `product_service.
    join_description_and_desc2`) - the downloaded file, re-imported by hand, has to
    store exactly what the pull stores. So Desc 2 is the remainder of `description`
    after `name` with EXACTLY ONE separator space removed - any FURTHER leading
    whitespace (a real double space in the source data) is KEPT, never stripped, because
    that extra space is exactly what the join's own inserted space needs to reproduce.
    `description == name` -> Desc 2 empty. `description` not even starting with `name`
    (no committed fixture row hits this, but the rule must still answer something) -
    there is no boundary to split on, so the Description cell carries the full raw text
    and Desc 2 is empty.
    """
    name = row.get("name") or ""
    description = row.get("description") or ""
    if description == name:
        view_description, desc2 = name, ""
    elif not description.startswith(name):
        view_description, desc2 = description, ""
    else:
        remainder = description[len(name):]
        if remainder.startswith(" "):
            remainder = remainder[1:]
        view_description, desc2 = name, remainder
    return {
        "item_code": row.get("code"),
        "description": view_description,
        "desc_2": desc2,
        "item_group": row.get("category_code"),
        "item_brand": row.get("brand_code"),
        "price": _row_price(row.get("list_price")),
        "is_active": row.get("is_active"),
    }


def _row_price(raw_list_price) -> float:
    """AC-RV-3 (Phase 3 fix round, F-9): `price` is a real number, not the raw FoundryX
    string, so the FE never parses it and the download's Price cell lands numeric. A blank
    or unparseable value reads as 0 - the same "no price" sentinel the manual import's own
    `parse_manual_list_price` uses, but here as a float (compare keeps the raw pull rows,
    never this mapped view, so nothing else reads this conversion). `nan`/`inf` (fix-round
    follow-up, item 5) are floats `float()` parses without error but that are not a price -
    `math.isfinite` catches both, same fallback to 0."""
    if raw_list_price is None or str(raw_list_price).strip() == "":
        return 0.0
    try:
        value = float(raw_list_price)
    except (TypeError, ValueError):
        return 0.0
    return value if math.isfinite(value) else 0.0


def paginate_rows(mapped_rows: list[dict], *, page: int, limit: int, query: Optional[str]) -> dict:
    """AC-RV-3: `query` filters by item code, contains, case-insensitive - applied before
    paging, over the whole assembled set (plan: "no cache", the FoundryX side is the one
    that pages 1000 at a time; this is the already-assembled view)."""
    rows = mapped_rows
    needle = (query or "").strip().lower()
    if needle:
        rows = [r for r in rows if needle in (r.get("item_code") or "").lower()]
    total = len(rows)
    total_pages = max(1, (total + limit - 1) // limit)
    start = (page - 1) * limit
    page_rows = rows[start:start + limit]
    return {
        "data": page_rows,
        "pagination": {"page": page, "limit": limit, "total": total, "total_pages": total_pages},
    }


_PRODUCTS_TEMPLATE_HEADER = (
    "Item Code", "Description", "Desc 2", "Item Group", "Item Brand", "Price", "Is Active", "UOM",
)


def stock_list_archive_filename(company_code: str, *, when: Optional[datetime] = None) -> str:
    """D3 (small-fix track, browser e2e run 3): the archived Stock List used to be named
    `autocount-pull-pull.xlsx` - the apply job's OWN `job_metadata` is keyed
    `autocount_apply`, not `autocount_pull`, so `entity_of(apply_job)` (which only ever
    reads the latter) silently fell back to "pull" for BOTH halves of `download_filename`'s
    old `f"autocount-{entity}-pull.xlsx"`. Company code and the LOCAL (Asia/Kuala_Lumpur)
    calendar date of the apply instead, so two companies - or the same company on two
    different days - never collide on the one install-wide file name history
    (`replace_latest_stock_list` itself is unscoped, see the plan's accepted risks).

    The `or "company"` fallback (N2, opus review, fix round 3) lives HERE, the one place
    both callers (`download_filename` below, and the apply task's own archive step in
    `autocount_pull_tasks.py`) go through - an unresolvable `company_id` at either call
    site (`_company_code`'s own "" default) used to reach this function bare, producing
    `autocount-stock-list--<date>.xlsx`.
    """
    when = when or datetime.now(_MY_TZ)
    company_code = company_code or "company"
    return f"autocount-stock-list-{company_code.lower()}-{when.strftime('%Y%m%d')}.xlsx"


def download_filename(job: ImportJob) -> str:
    entity = entity_of(job) or "pull"
    if entity == "stock_balances":
        # Same name the archive gets (D3) - the pull job's OWN metadata carries
        # `company_code` (set once, at `start_pull`), so no extra DB read is needed here
        # the way the apply job (no `autocount_pull` metadata at all) requires.
        return stock_list_archive_filename(_pull_meta(job).get("company_code"))
    return f"autocount-{entity}-pull.xlsx"


def _neutralize_formula_cells(sheet) -> None:
    """Formula-injection guard (Phase 3 fix round, F-4): AutoCount item text (a
    Description, an Item Description) is untrusted input that lands straight in a
    workbook a checker opens in Excel - a value starting with `=` would otherwise
    execute as a formula the moment the file opens. Every string cell openpyxl would
    itself have classified as a formula (`data_type == "f"`) is forced back to a
    literal string cell, value unchanged - shared by both workbook builders so neither
    can forget it. (`+`/`-`/`@` are a CSV-open concern, not an xlsx one - openpyxl
    never classifies them as `data_type == "f"`, so there is nothing here for them to
    catch.)"""
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.data_type == "f":
                cell.data_type = "s"


def build_products_workbook(mapped_rows: list[dict]) -> bytes:
    """AC-RV-5: the template header row exactly, with a trailing blank UOM column."""
    import io

    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "AutoCount pull"
    sheet.append(list(_PRODUCTS_TEMPLATE_HEADER))
    for row in mapped_rows:
        sheet.append([
            row["item_code"], row["description"], row["desc_2"], row["item_group"],
            row["item_brand"], row["price"], bool(row["is_active"]), None,
        ])
    _neutralize_formula_cells(sheet)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# =================================================================== SR4 - stock rows


def _stock_template_row(row: dict) -> dict:
    """The manual-template shape for one raw snapshot stock row - what a checker's own
    Excel file, `StockService.bulk_import_stock` and `compare_stock` all speak. The ONE
    shape every `classify_stock_rows` bucket uses (fed, inactive, unknown alike)."""
    return {
        "Item Code": row.get("item_code"),
        "Item Description": row.get("item_description") or "",
        "Location": row.get("location_code"),
        "On Hand Qty": row.get("qty"),
    }


def classify_stock_rows(db: Session, company_id: str, rows: list[dict]) -> dict:
    """AC-SP-2: classifies every pulled stock row by its `location_code`, matched to
    `warehouses.warehouse_code` trimmed + case-insensitive, WITHIN `company_id` - an
    EXPLICIT argument, never ambient session scope (sp_2a proves this on purpose: it
    sets ambient scope to two companies at once and still expects the right answer).
    ONE query for the company's warehouses. Returns `{"fed", "inactive", "unknown"}`,
    every bucket in the same manual-template shape."""
    warehouse_rows = db.execute(
        text("SELECT warehouse_code, is_active FROM warehouses WHERE company_id = :cid"),
        {"cid": company_id},
    ).mappings().all()
    active_codes = {
        str(w["warehouse_code"]).strip().upper() for w in warehouse_rows if w["is_active"]
    }
    known_codes = {str(w["warehouse_code"]).strip().upper() for w in warehouse_rows}

    fed: list[dict] = []
    inactive: list[dict] = []
    unknown: list[dict] = []
    for row in rows:
        location = str(row.get("location_code") or "").strip().upper()
        template = _stock_template_row(row)
        if location in active_codes:
            fed.append(template)
        elif location in known_codes:
            inactive.append(template)
        else:
            unknown.append(template)
    return {"fed": fed, "inactive": inactive, "unknown": unknown}


def map_stock_row(template_row: dict) -> dict:
    """AC-RV-3 (stock half): the `/rows` view shape, built from a `classify_stock_rows`
    FED-bucket template row - one source of truth for what counts as FED, never a
    second re-derivation."""
    return {
        "item_code": template_row.get("Item Code"),
        "item_description": template_row.get("Item Description"),
        "location": template_row.get("Location"),
        "on_hand_qty": template_row.get("On Hand Qty"),
    }


_STOCK_TEMPLATE_HEADER = ("Item Code", "Item Description", "Location", "On Hand Qty")


def build_stock_workbook(template_rows: list[dict]) -> bytes:
    """AC-RV-5 (stock half) / AC-SC-4: ONE builder, used by `/download.xlsx` AND the
    apply task's Stock List archive - the template's own header row, no trailing
    column (unlike products, which adds a blank UOM)."""
    import io

    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Stock List"
    sheet.append(list(_STOCK_TEMPLATE_HEADER))
    for row in template_rows:
        sheet.append([
            row.get("Item Code"), row.get("Item Description"),
            row.get("Location"), row.get("On Hand Qty"),
        ])
    _neutralize_formula_cells(sheet)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def store_compare_summary(db: Session, job: ImportJob, *, filename: str, result: dict) -> ImportJob:
    """AC-CM-5: only the summary is kept on the job - the uploaded rows and the
    difference list are returned to the browser and never stored.

    `qty_total_excel`/`qty_total_pull` (AC-CM-3, stock only) come straight through from
    `compare_stock`'s own summary; absent (products) they stay `None`, matching the
    plan's metadata shape.
    """
    pull = _pull_meta(job)
    summary = result.get("summary") or {}
    pull["compare"] = {
        "filename": filename,
        "compared_at": datetime.utcnow().isoformat(),
        "total": summary.get("total", 0),
        "matched": summary.get("matched", 0),
        "different": summary.get("different", 0),
        "only_in_excel": len(result.get("only_in_excel") or []),
        "only_in_pull": len(result.get("only_in_pull") or []),
        "qty_total_excel": summary.get("qty_total_excel"),
        "qty_total_pull": summary.get("qty_total_pull"),
    }
    job.job_metadata = _with_pull(job, pull)
    job.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(job)
    return job


# ======================================================================== SR3 - confirm


def confirm_pull(db: Session, job: ImportJob, *, user_id: str) -> dict:
    """AC-PC-1: creates the SECOND `import_jobs` row (the apply) once, enqueues it once,
    and marks the pull `confirmed`. A second Confirm - on a pull already `confirmed` with
    an apply job - returns the same apply job and enqueues nothing.

    The once-guard is a conditional UPDATE on the pull's own stored phase (`review` ->
    `confirmed`), the same style `_claim_and_enqueue_preview` uses on `status` - run
    FIRST, before anything about the apply job is written, and its rowcount is checked
    BEFORE any commit. A caller that loses the race (Phase 3 fix round, F-7) rolls back
    (nothing of its own was ever committed) and re-reads the winner's state instead,
    leaving no orphan `import_jobs` row behind.

    ONE transaction, tightened further (fix-round follow-up): the phase-flip UPDATE and
    the apply job INSERT commit TOGETHER, in the SAME `db.commit()` call - the pull can
    never be stored `confirmed` without its own apply row existing. If that commit fails
    for any reason, the whole transaction (phase flip included) rolls back and the
    exception propagates - the pull is left exactly as it was (`review`, no
    `apply_job_id`), so a later Confirm starts over cleanly rather than pointing at an
    apply job that was never actually written.
    """
    pull = _pull_meta(job)
    if pull.get("phase") == "confirmed" and pull.get("apply_job_id"):
        return serialize(job, db)
    if pull.get("phase") != "review":
        raise PullNotReadyForConfirm(
            f"Pull is in phase {pull.get('phase')!r}; only a pull in review can be confirmed."
        )
    if pull.get("confirm_blocked_reason"):
        # AC-SP-1: entity-agnostic - products never sets this, stock does when the
        # fetched header reports excludedNonzeroCount > 0.
        raise PullNotReadyForConfirm(str(pull["confirm_blocked_reason"]))

    entity = pull.get("entity")
    # Generated up front so the ONE conditional UPDATE can write phase + apply_job_id
    # together, atomically - the apply job row itself is inserted only once that UPDATE
    # has proven this caller won the race.
    apply_job_id = str(uuid.uuid4())
    apply_job_rq_id = str(uuid.uuid4())

    pull["phase"] = "confirmed"
    pull["apply_job_id"] = apply_job_id
    new_meta = _with_pull(job, pull)

    result = db.execute(
        update(ImportJob)
        .where(
            ImportJob.id == job.id,
            ImportJob.job_metadata["autocount_pull"]["phase"].astext == "review",
        )
        .values(job_metadata=new_meta, updated_at=datetime.utcnow())
        .execution_options(synchronize_session=False)
    )
    if result.rowcount != 1:
        # Someone else confirmed it first between our read and this write - the UPDATE
        # above was never committed, so a plain rollback discards it; nothing of ours
        # was ever visible to anyone. The winner's apply job (and its own apply_job_id)
        # is the one that counts.
        db.rollback()
        db.refresh(job)
        return serialize(job, db)

    apply_job = ImportJob(
        id=apply_job_id,
        job_id=apply_job_rq_id,
        job_type=APPLY_JOB_TYPES[entity],
        status=JobStatus.QUEUED.value,
        user_id=user_id,
        company_id=job.company_id,
        job_metadata={
            "autocount_apply": {
                "pull_job_id": str(job.id),
                "snapshot_id": pull.get("snapshot_id"),
                "entity": entity,
            }
        },
    )
    db.add(apply_job)
    try:
        db.commit()
    except Exception:
        # The apply row failed to write - the pull must never end up `confirmed`
        # without it, so the whole transaction (the phase flip included) is undone.
        db.rollback()
        raise

    from app.tasks.autocount_pull_tasks import apply_autocount_pull

    enqueue_job(
        apply_autocount_pull,
        str(apply_job.id),
        queue_name="imports",
        job_timeout=3600,
        job_id=str(apply_job.job_id),
    )
    db.refresh(job)
    return serialize(job, db)


# ======================================================================== discard


def discard_pull(db: Session, job: ImportJob) -> dict:
    """AC-DS-1..5: throws away an open pull - no FoundryX call (the snapshot just
    expires on its own), no product/stock row ever touched. Allowed from `building`,
    `previewing` and `review` (`_OPEN_PHASES`, the same set `find_open_pull` treats as
    "open"); already `discarded` returns the same body unchanged (idempotent, AC-DS-4);
    anything else (`confirmed`, `failed`, `expired`) raises `PullNotDiscardable` - the
    route turns that into a 409.
    """
    pull = _pull_meta(job)
    phase = pull.get("phase")
    if phase == "discarded":
        return serialize(job, db)
    if phase not in _OPEN_PHASES:
        raise PullNotDiscardable(
            f"Pull is in phase {phase!r}; only an open pull can be discarded."
        )
    pull["phase"] = "discarded"
    job.job_metadata = _with_pull(job, pull)
    job.status = JobStatus.CANCELLED.value
    job.completed_at = datetime.utcnow()
    job.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(job)
    return serialize(job, db)
